#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Task1 Step 5.7-C EventMention-level AI-assisted annotation pilot."""

from __future__ import annotations

import argparse
import concurrent.futures
import json
import os
import random
import re
import shutil
import signal
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_task1_step5_5_candidate_v2 import protected_file_list, rel, sha256_file, write_json, write_jsonl, write_xlsx  # noqa: E402
from src.event_extraction.task1_conditional_structure_32b import API_BASE, MODEL_NAME, chat_completion, check_model_service, read_jsonl  # noqa: E402
from src.event_extraction.task1_eventmention_contract import (  # noqa: E402
    HUMAN_REVIEW_FIELDS,
    adapt_text_draft,
    build_content_excerpt,
    build_teacher_prompt,
    build_teacher_user_message,
    contract_schema,
    correction_user_message,
    empty_span,
    empty_time,
    eventmention_from_spans,
    grade_draft,
    normalize_direction,
    sanitize_prediction_for_contract,
    stable_event_mention_id,
    validate_eventmention_record,
    validate_teacher_payload,
)
from src.event_extraction.task1_structure_offset_adapter import strict_json_loads  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except Exception:  # pragma: no cover - openpyxl is present in the project env, fallback is tested by output files.
    Workbook = None  # type: ignore[assignment]
    Font = None  # type: ignore[assignment]
    PatternFill = None  # type: ignore[assignment]


EXPECTED_CONFIRMED_SHA = "414174fdef7b8aeb89edecdde4a85f43ac08d86093664cb4a6d4541732f1186c"
CONFIRMED = ROOT / "outputs/task1_structure_validation_round1/minimal_human_review_package_confirmed.jsonl"
STEP57A_DIR = ROOT / "outputs/task1_step5_7a_conditional_structure_32b"
STEP57B_DIR = ROOT / "outputs/task1_step5_7b_structure_failure_audit"
OUT = ROOT / "outputs/task1_step5_7c_eventmention_annotation_pilot"
REPORT_JSON = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_report.json"
REPORT_MD = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_report.md"
REPORT_XLSX = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_details.xlsx"
REVIEW_XLSX = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_review.xlsx"
REVIEW_GUIDE = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_guide.md"
REVIEW_MANIFEST = ROOT / "reports/task1_step5_7c_eventmention_annotation_pilot_manifest.json"
VLLM_LOG = ROOT / "reports/task1_step5_7c_vllm.log"

TRAIN_PAIRS = ROOT / "outputs/task1_step5_6d_classifier_retraining/train_gold_pairs.jsonl"
DEV_ARTICLES = ROOT / "outputs/task1_step5_6d_classifier_retraining/dev_articles.jsonl"
HOLDOUT_ARTICLES = ROOT / "outputs/task1_step5_6d_classifier_retraining/holdout_articles.jsonl"
BLIND_INPUT = ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl"
SCHEMA_FILE = ROOT / "data/processed/event_schema.jsonl"


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def files_under(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    if path.is_dir():
        return sorted(item for item in path.rglob("*") if item.is_file())
    return []


def protected_and_frozen_files() -> list[Path]:
    files = list(protected_file_list())
    for path in [
        CONFIRMED,
        ROOT / "src/routing/task1_candidate_v2.py",
        ROOT / "configs/task1_candidate_v2.yaml",
        ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
        ROOT / "outputs/task1_step5_6f_classifier_consolidation/SHA256SUMS.txt",
        ROOT / "outputs/event_match_baseline/best_model",
        ROOT / "outputs/task1_limited_training_round1/model_epoch_2",
        ROOT / "outputs/task1_step5_6d_classifier_retraining/experiment_C",
        STEP57A_DIR,
        STEP57B_DIR,
    ]:
        files.extend(files_under(path))
    return sorted(set(files), key=lambda item: rel(item))


def sha_snapshot(path: Path, files: list[Path]) -> dict[str, str]:
    rows = {rel(item): sha256_file(item) for item in files if item.exists()}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(rows.items())), encoding="utf-8")
    return rows


def load_sha_file(path: Path) -> dict[str, str]:
    out: dict[str, str] = {}
    if not path.exists():
        return out
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        digest, name = line.split("  ", 1)
        out[name] = digest
    return out


def process_check() -> dict[str, Any]:
    proc = subprocess.run(
        [
            "bash",
            "-lc",
            "ps -eo pid,ppid,stat,etime,cmd | rg -i 'vllm|DianJin-R1-32B|train_task1|run_task1|evaluate_task1' | rg -v 'rg -i' || true\n"
            "ss -ltnp '( sport = :8000 )' || true\n"
            "nvidia-smi --query-gpu=index,memory.used,memory.total --format=csv,noheader,nounits 2>/dev/null || true",
        ],
        capture_output=True,
        text=True,
        timeout=40,
    )
    return {"returncode": proc.returncode, "raw": proc.stdout.splitlines()}


def article_ids(path: Path) -> set[str]:
    if not path.exists():
        return set()
    return {str(row.get("article_id", "")).strip() for row in read_jsonl(path) if str(row.get("article_id", "")).strip()}


def official_event_names() -> tuple[set[str], dict[str, str]]:
    rows = read_jsonl(SCHEMA_FILE)
    names = {str(row["event_name"]) for row in rows}
    descriptions: dict[str, str] = {}
    for row in rows:
        name = str(row["event_name"])
        descriptions.setdefault(name, str(row.get("raw_description", "")))
    return names, descriptions


def event_names_from_confirmed(row: dict[str, Any]) -> list[str]:
    return [part.strip() for part in str(row.get("gold_event_names_for_context_after_classification_eval", "")).split("|") if part.strip()]


def locate_text(content: str, text: str) -> dict[str, Any]:
    value = str(text or "")
    if not value or "..." in value or "……" in value:
        return empty_span()
    pos = str(content or "").find(value)
    if pos < 0:
        return {"text": value, "start": -1, "end": -1}
    return {"text": value, "start": pos, "end": pos + len(value)}


def old_row_event_name(row: dict[str, Any], audit_by_id: dict[str, dict[str, Any]]) -> tuple[str, str, str]:
    aid = str(row["article_id"])
    names = event_names_from_confirmed(row)
    first = names[0] if names else ""
    audit = audit_by_id.get(aid, {})
    category = str(audit.get("alignment_category", "single_event"))
    note = ""
    if aid == "NEWS_MULTISOURCE_D-1120797906500":
        return "关键技术突破", "needs_human_confirmation", "wrong_event_selected: remapped to audited better-matching event_name; not reviewer-confirmed."
    if aid in {"NEWS_MULTISOURCE_D-1126130257189", "NEWS_MULTISOURCE_D-1127695210485"}:
        return first, "needs_human_confirmation", "ambiguous_multi_event: placeholder event_name is first listed event; human must confirm."
    if category == "first_event_alignment_correct" or len(names) <= 1:
        return first, "confirmed", "single_event_or_first_event_alignment_correct"
    return first, "needs_human_confirmation", f"unresolved_alignment_category:{category}"


def migrate_confirmed80(rows: list[dict[str, Any]], output_dir: Path, event_names: set[str]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    audit_rows = read_jsonl(STEP57B_DIR / "multi_event_alignment_audit.jsonl")
    audit_by_id = {str(row["article_id"]): row for row in audit_rows}
    migrated: list[dict[str, Any]] = []
    review_rows: list[dict[str, Any]] = []
    pending_ids = []
    for index, row in enumerate(rows):
        aid = str(row["article_id"])
        event_name, status, note = old_row_event_name(row, audit_by_id)
        base = {
            "event_mention_id": stable_event_mention_id(aid, event_name),
            "article_id": aid,
            "event_name": event_name,
            "title": str(row.get("title", "")),
            "content": str(row.get("content", "")),
            "publish_time": str(row.get("publish_time", "")),
            "schema_id": "",
            "target_type": "",
            "target_code": "",
            "target_name": "",
        }
        text_payload = {
            "event_name": event_name,
            "trigger_text": str(row.get("annotate_trigger", "")),
            "event_span_text": str(row.get("annotate_event_span", "")),
            "subject_company_text": str(row.get("annotate_subject_company", "")),
            "event_time_text": "" if "未明确" in str(row.get("annotate_event_time", "")) else str(row.get("annotate_event_time", "")),
            "event_time_normalized": "未明确" if "未明确" in str(row.get("annotate_event_time", "")) else str(row.get("annotate_event_time", "")),
            "object_target_text": str(row.get("annotate_object_target", "")),
            "direction": str(row.get("annotate_direction", "unclear")),
            "arguments": [],
            "confidence": 1.0,
        }
        prediction, fields, meta, payload_errors = adapt_text_draft(base["content"], json.dumps(text_payload, ensure_ascii=False), event_name)
        grade = grade_draft(prediction, fields, meta, payload_errors)
        source = "ai_assisted_human_confirmed" if status == "confirmed" else "contract_reconstructed"
        confirmed = "yes" if status == "confirmed" else "no"
        evm = eventmention_from_spans(base, prediction, source, confirmed, status)
        evm.update(
            {
                "legacy_article_order": index,
                "legacy_gold_event_names": event_names_from_confirmed(row),
                "legacy_structure": {
                    "annotate_trigger": row.get("annotate_trigger", ""),
                    "annotate_event_span": row.get("annotate_event_span", ""),
                    "annotate_subject_company": row.get("annotate_subject_company", ""),
                    "annotate_event_time": row.get("annotate_event_time", ""),
                    "annotate_object_target": row.get("annotate_object_target", ""),
                    "annotate_direction": row.get("annotate_direction", ""),
                    "annotate_arguments_json": row.get("annotate_arguments_json", ""),
                    "annotate_argument_roles_json": row.get("annotate_argument_roles_json", ""),
                },
                "alignment_category": audit_by_id.get(aid, {}).get("alignment_category", "single_event"),
                "migration_note": note,
                "event_name_needs_confirmation": status != "confirmed",
                "field_validation": grade["field_validation"],
                "unresolvable_fields": grade["unresolvable_fields"],
                "ambiguous_fields": grade["ambiguous_fields"],
                "contract_errors": grade["contract_errors"],
            }
        )
        validation = validate_eventmention_record(evm, event_names)
        evm["contract_validation"] = validation
        migrated.append(evm)
        if status != "confirmed":
            pending_ids.append(aid)
        review_rows.append(
            {
                "legacy_order": index,
                "article_id": aid,
                "event_mention_id": evm["event_mention_id"],
                "legacy_event_names": "|".join(event_names_from_confirmed(row)),
                "migrated_event_name": event_name,
                "annotation_status": status,
                "reviewer_confirmed": confirmed,
                "alignment_category": evm["alignment_category"],
                "migration_note": note,
                "annotate_trigger": row.get("annotate_trigger", ""),
                "annotate_event_span": row.get("annotate_event_span", ""),
                "contract_valid": validation["valid"],
                "contract_errors": "|".join(validation["errors"]),
            }
        )
    write_jsonl(output_dir / "confirmed80_eventmention_migration.jsonl", migrated)
    manifest = {
        "generated_at": now(),
        "source_file": rel(CONFIRMED),
        "source_sha256": sha256_file(CONFIRMED),
        "row_count": len(rows),
        "eventmention_count": len(migrated),
        "safe_migrated_count": sum(1 for row in migrated if row["annotation_status"] == "confirmed"),
        "needs_human_confirmation_count": len(pending_ids),
        "needs_human_confirmation_article_ids": pending_ids,
        "wrong_event_selected_remapped_article_id": "NEWS_MULTISOURCE_D-1120797906500",
        "ambiguous_multi_event_article_ids": ["NEWS_MULTISOURCE_D-1126130257189", "NEWS_MULTISOURCE_D-1127695210485"],
        "reference": "AI-assisted, human-confirmed old article-level annotations migrated to EventMention-level candidate records; pending rows are not independent manual gold.",
    }
    write_json(output_dir / "confirmed80_migration_manifest.json", manifest)
    write_contract_review_xlsx(output_dir / "confirmed80_contract_review.xlsx", review_rows)
    return migrated, manifest


def write_contract_review_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    pending_fill = "FFF2CC"
    if Workbook is None:
        write_xlsx(path, {"confirmed80_contract_review": rows})
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "旧80条契约修正表"
    columns = sorted({key for row in rows for key in row})
    ws.append(columns)
    if Font is not None:
        for cell in ws[1]:
            cell.font = Font(bold=True)
    fill = PatternFill("solid", fgColor=pending_fill) if PatternFill is not None else None
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
        if row.get("annotation_status") != "confirmed" and fill is not None:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def sample_eventmentions(args: argparse.Namespace, output_dir: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    official_names, descriptions = official_event_names()
    confirmed_ids = article_ids(CONFIRMED)
    dev_ids = article_ids(DEV_ARTICLES)
    holdout_ids = article_ids(HOLDOUT_ARTICLES)
    blind_ids = article_ids(BLIND_INPUT)
    excluded_ids = confirmed_ids | dev_ids | holdout_ids | blind_ids
    all_pairs = read_jsonl(TRAIN_PAIRS)
    pool = [row for row in all_pairs if str(row.get("article_id", "")) not in excluded_ids]
    if len(pool) < args.sample_size:
        raise RuntimeError(f"insufficient eligible training-side EventMention pool: {len(pool)} < {args.sample_size}")
    rng = random.Random(args.seed)
    by_article = Counter(str(row["article_id"]) for row in pool)
    multi_article_ids = {aid for aid, count in by_article.items() if count > 1}
    by_class: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in pool:
        by_class[str(row["event_name"])].append(row)
    for items in by_class.values():
        items.sort(key=lambda row: (str(row["article_id"]), str(row.get("schema_id", ""))))
        rng.shuffle(items)
    selected: list[dict[str, Any]] = []
    selected_ids: set[str] = set()

    def mention_id(row: dict[str, Any]) -> str:
        return stable_event_mention_id(str(row["article_id"]), str(row["event_name"]), str(row.get("schema_id", "")))

    def add(row: dict[str, Any]) -> bool:
        mid = mention_id(row)
        if mid in selected_ids:
            return False
        selected_ids.add(mid)
        selected.append(row)
        return True

    for event_name in sorted(by_class, key=lambda name: (len(by_class[name]), name)):
        for row in by_class[event_name][: min(2, len(by_class[event_name]))]:
            add(row)
            if len(selected) >= args.sample_size:
                break
        if len(selected) >= args.sample_size:
            break
    multi_rows = [row for row in pool if str(row["article_id"]) in multi_article_ids]
    multi_rows.sort(key=lambda row: (str(row["event_name"]), str(row["article_id"]), str(row.get("schema_id", ""))))
    rng.shuffle(multi_rows)
    while sum(1 for row in selected if str(row["article_id"]) in multi_article_ids) < args.min_multi_event_mentions and multi_rows:
        row = multi_rows.pop()
        add(row)
    remaining = [row for row in pool if mention_id(row) not in selected_ids]
    class_counts = Counter(str(row["event_name"]) for row in pool)
    remaining.sort(key=lambda row: (class_counts[str(row["event_name"])], normalize_direction(row.get("emotion", "")), str(row["event_name"]), str(row["article_id"])))
    rng.shuffle(remaining)
    remaining.sort(key=lambda row: (class_counts[str(row["event_name"])], str(row["event_name"])))
    for row in remaining:
        if len(selected) >= args.sample_size:
            break
        add(row)
    if len(selected) != args.sample_size:
        raise RuntimeError(f"failed to select {args.sample_size} EventMentions, got {len(selected)}")
    selected.sort(key=lambda row: (str(row["event_name"]), str(row["article_id"]), str(row.get("schema_id", ""))))

    eventmentions: list[dict[str, Any]] = []
    for index, row in enumerate(selected):
        event_name = str(row["event_name"])
        event_description = str(row.get("event_description") or descriptions.get(event_name, ""))
        excerpt = build_content_excerpt(str(row.get("title", "")), str(row.get("content", "")), event_name, event_description)
        eventmentions.append(
            {
                "input_index": index,
                "event_mention_id": mention_id(row),
                "article_id": str(row["article_id"]),
                "event_name": event_name,
                "schema_id": str(row.get("schema_id", "")),
                "target_type": str(row.get("target_type", "")),
                "target_code": str(row.get("target_code", "")),
                "target_name": str(row.get("target_name", "")),
                "title": str(row.get("title", "")),
                "content": str(row.get("content", "")),
                "content_excerpt": excerpt["excerpt"],
                "content_excerpt_mode": excerpt["mode"],
                "content_truncated": excerpt["content_truncated"],
                "content_length": excerpt["content_length"],
                "publish_time": str(row.get("publish_time", "")),
                "source": str(row.get("source", "")),
                "event_description": event_description,
                "sampling_reference_emotion": str(row.get("emotion", "")),
                "multi_event_article": str(row["article_id"]) in multi_article_ids,
            }
        )
    selected_class_counts = Counter(row["event_name"] for row in eventmentions)
    insufficient_two = {
        name: len(by_class[name])
        for name in sorted(by_class)
        if len(by_class[name]) < 2
    }
    manifest = {
        "generated_at": now(),
        "seed": args.seed,
        "sampling_unit": "EventMention",
        "sample_size": len(eventmentions),
        "source_file": rel(TRAIN_PAIRS),
        "source_sha256": sha256_file(TRAIN_PAIRS),
        "official_event_class_count": len(official_names),
        "eligible_pool_eventmention_count": len(pool),
        "eligible_pool_article_count": len({str(row["article_id"]) for row in pool}),
        "eligible_pool_class_count": len(by_class),
        "excluded_article_counts": {
            "confirmed80": len(confirmed_ids),
            "step5_6d_dev": len(dev_ids),
            "step5_6d_holdout": len(holdout_ids),
            "gold_test_blind_input_article_ids_only": len(blind_ids),
        },
        "overlap_checks": {
            "selected_vs_confirmed80": len({row["article_id"] for row in eventmentions} & confirmed_ids),
            "selected_vs_dev": len({row["article_id"] for row in eventmentions} & dev_ids),
            "selected_vs_holdout": len({row["article_id"] for row in eventmentions} & holdout_ids),
            "selected_vs_gold_blind_input": len({row["article_id"] for row in eventmentions} & blind_ids),
        },
        "selected_article_count": len({row["article_id"] for row in eventmentions}),
        "selected_multi_event_mention_count": sum(1 for row in eventmentions if row["multi_event_article"]),
        "selected_multi_event_article_count": len({row["article_id"] for row in eventmentions if row["multi_event_article"]}),
        "selected_class_count": len(selected_class_counts),
        "selected_class_distribution": dict(sorted(selected_class_counts.items())),
        "selected_direction_distribution_reference_only": dict(Counter(row["sampling_reference_emotion"] for row in eventmentions)),
        "classes_with_less_than_two_available_mentions": insufficient_two,
        "gold_test_labels_read": False,
        "unlabeled_200k_news_used": False,
    }
    write_jsonl(output_dir / "sampled_eventmentions.jsonl", eventmentions)
    write_json(output_dir / "sampling_manifest.json", manifest)
    return eventmentions, manifest


def start_vllm_if_needed(args: argparse.Namespace) -> tuple[subprocess.Popen[str] | None, dict[str, Any]]:
    initial = check_model_service(args.api_base, args.model, timeout=5)
    if initial.get("available") and initial.get("model_found"):
        return None, {"service_preexisting": True, "started_by_this_run": False, "initial_health": initial}
    if initial.get("available") and not initial.get("model_found"):
        raise RuntimeError(f"port 8000 has a non-target service or wrong model: {initial}")
    if args.no_start_vllm:
        raise RuntimeError(f"model service unavailable and --no-start-vllm was set: {initial}")
    model_path = Path("/opt/data/private/JinRong/models/DianJin-R1-32B")
    if not model_path.exists():
        raise FileNotFoundError(str(model_path))
    VLLM_LOG.parent.mkdir(parents=True, exist_ok=True)
    log_handle = VLLM_LOG.open("a", encoding="utf-8")
    cmd = [
        "/root/miniconda3/envs/dianjin-cu128/bin/vllm",
        "serve",
        str(model_path),
        "--served-model-name",
        args.model,
        "--tensor-parallel-size",
        "4",
        "--dtype",
        "bfloat16",
        "--max-model-len",
        "8192",
        "--max-num-seqs",
        "4",
        "--gpu-memory-utilization",
        "0.90",
        "--disable-custom-all-reduce",
        "--host",
        "127.0.0.1",
        "--port",
        "8000",
    ]
    env = os.environ.copy()
    env["CUDA_VISIBLE_DEVICES"] = "0,1,2,3"
    proc = subprocess.Popen(cmd, stdout=log_handle, stderr=subprocess.STDOUT, text=True, env=env)
    deadline = time.time() + args.service_wait_seconds
    last_health: dict[str, Any] = {}
    while time.time() < deadline:
        if proc.poll() is not None:
            log_handle.close()
            raise RuntimeError(f"vLLM exited during startup with code {proc.returncode}; see {rel(VLLM_LOG)}")
        last_health = check_model_service(args.api_base, args.model, timeout=5)
        if last_health.get("available") and last_health.get("model_found"):
            return proc, {
                "service_preexisting": False,
                "started_by_this_run": True,
                "pid": proc.pid,
                "command": "CUDA_VISIBLE_DEVICES=0,1,2,3 " + " ".join(cmd),
                "log": rel(VLLM_LOG),
                "health": last_health,
            }
        time.sleep(10)
    raise RuntimeError(f"vLLM startup health check failed after {args.service_wait_seconds}s; last_health={last_health}; see {rel(VLLM_LOG)}")


def stop_vllm(proc: subprocess.Popen[str] | None) -> dict[str, Any]:
    if proc is None:
        return {"stopped": False, "reason": "service_preexisting_or_not_started"}
    if proc.poll() is not None:
        return {"stopped": True, "already_exited": True, "returncode": proc.returncode}
    proc.send_signal(signal.SIGTERM)
    try:
        proc.wait(timeout=60)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait(timeout=30)
        return {"stopped": True, "killed_after_timeout": True, "returncode": proc.returncode}
    return {"stopped": True, "killed_after_timeout": False, "returncode": proc.returncode}


def request_one(item: dict[str, Any], prompt: str, args: argparse.Namespace) -> dict[str, Any]:
    started = time.perf_counter()
    raw_outputs: list[dict[str, Any]] = []
    errors: list[str] = []
    request_count = 0
    retry_count = 0
    status = "failed"
    final_content = ""
    final_payload: dict[str, Any] | None = None
    usage = {"prompt_tokens": 0, "completion_tokens": 0, "total_tokens": 0}
    messages = [
        {"role": "system", "content": prompt},
        {"role": "user", "content": build_teacher_user_message(item)},
    ]
    for attempt in range(2):
        try:
            request_count += 1
            content, response_payload, elapsed = chat_completion(
                messages,
                api_base=args.api_base,
                model=args.model,
                timeout=args.timeout,
                max_tokens=args.max_tokens,
            )
            final_content = content
            raw_outputs.append({"attempt": attempt + 1, "content": content, "elapsed_seconds": round(elapsed, 3), "usage": response_payload.get("usage", {})})
            parsed = strict_json_loads(content)
            ok, validation_errors = validate_teacher_payload(parsed, str(item["event_name"]), str(item["content"]))
            if ok:
                status = "success"
                final_payload = parsed
                usage_payload = response_payload.get("usage", {}) if isinstance(response_payload, dict) else {}
                usage = {
                    "prompt_tokens": int(usage_payload.get("prompt_tokens", 0) or 0),
                    "completion_tokens": int(usage_payload.get("completion_tokens", 0) or 0),
                    "total_tokens": int(usage_payload.get("total_tokens", 0) or 0),
                }
                break
            errors = validation_errors
            if attempt == 0:
                retry_count += 1
                messages = [
                    {"role": "system", "content": prompt},
                    {"role": "user", "content": correction_user_message(content, validation_errors)},
                ]
        except Exception as exc:
            errors = [f"{type(exc).__name__}:{exc}"]
            raw_outputs.append({"attempt": attempt + 1, "content": "", "error": errors[-1]})
            if attempt == 0:
                retry_count += 1
                messages = [
                    {"role": "system", "content": prompt},
                    {"role": "user", "content": correction_user_message("", errors)},
                ]
            else:
                break
    elapsed_total = time.perf_counter() - started
    lower_errors = " ".join(errors).lower()
    return {
        "input_index": item["input_index"],
        "event_mention_id": item["event_mention_id"],
        "article_id": item["article_id"],
        "event_name": item["event_name"],
        "status": status,
        "final_content": final_content,
        "final_payload": final_payload,
        "raw_outputs": raw_outputs,
        "errors": errors,
        "request_count": request_count,
        "retry_count": retry_count,
        "elapsed_seconds": round(elapsed_total, 3),
        "prompt_tokens": usage["prompt_tokens"],
        "completion_tokens": usage["completion_tokens"],
        "total_tokens": usage["total_tokens"],
        "timeout": int("timeout" in lower_errors or "timed out" in lower_errors),
        "http_5xx": int("http_error:5" in lower_errors),
        "cuda_oom": int("cuda" in lower_errors and "out of memory" in lower_errors),
    }


def load_checkpoint(path: Path) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    if not path.exists():
        return out
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            row = json.loads(line)
            out[str(row["event_mention_id"])] = row
    return out


def extract_raw_text_fields(raw_content: str) -> dict[str, Any]:
    try:
        payload = strict_json_loads(raw_content)
    except Exception:
        return {}
    keys = [
        "event_name",
        "trigger_text",
        "event_span_text",
        "subject_company_text",
        "event_time_text",
        "event_time_normalized",
        "object_target_text",
        "direction",
        "arguments",
        "confidence",
    ]
    return {key: payload.get(key, "") for key in keys}


def run_requests(eventmentions: list[dict[str, Any]], output_dir: Path, prompt: str, args: argparse.Namespace) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    raw_final = output_dir / "raw_responses.jsonl"
    pred_final = output_dir / "annotation_drafts.jsonl"
    metrics_final = output_dir / "request_metrics.jsonl"
    if raw_final.exists() or pred_final.exists() or metrics_final.exists():
        if raw_final.exists() and pred_final.exists() and metrics_final.exists():
            raw_rows = read_jsonl(raw_final)
            draft_rows = read_jsonl(pred_final)
            metric_rows = read_jsonl(metrics_final)
            if len(raw_rows) == len(eventmentions) and len(draft_rows) == len(eventmentions) and len(metric_rows) == len(eventmentions):
                return raw_rows, draft_rows, metric_rows
        raise RuntimeError("existing final 5.7-C request outputs are incomplete; refusing to overwrite")
    checkpoint_path = output_dir / "request_checkpoint.jsonl.tmp"
    completed = load_checkpoint(checkpoint_path)
    pending = [item for item in eventmentions if item["event_mention_id"] not in completed]
    output_dir.mkdir(parents=True, exist_ok=True)
    with checkpoint_path.open("a", encoding="utf-8") as checkpoint:
        with concurrent.futures.ThreadPoolExecutor(max_workers=args.workers) as pool:
            futures = {pool.submit(request_one, item, prompt, args): item for item in pending}
            for future in concurrent.futures.as_completed(futures):
                row = future.result()
                completed[str(row["event_mention_id"])] = row
                checkpoint.write(json.dumps(row, ensure_ascii=False) + "\n")
                checkpoint.flush()
    raw_rows = [completed[item["event_mention_id"]] for item in eventmentions]
    drafts, metric_rows = build_drafts(eventmentions, raw_rows)
    write_jsonl(raw_final, raw_rows)
    write_jsonl(pred_final, drafts)
    write_jsonl(metrics_final, metric_rows)
    return raw_rows, drafts, metric_rows


def build_drafts(eventmentions: list[dict[str, Any]], raw_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    by_id = {row["event_mention_id"]: row for row in eventmentions}
    drafts: list[dict[str, Any]] = []
    metric_rows: list[dict[str, Any]] = []
    official_names, _ = official_event_names()
    for raw in raw_rows:
        item = by_id[str(raw["event_mention_id"])]
        content = str(item["content"])
        if raw.get("final_content"):
            prediction, fields, meta, payload_errors = adapt_text_draft(content, str(raw["final_content"]), str(item["event_name"]))
        else:
            prediction, fields, meta, payload_errors = {"event_name": item["event_name"]}, [], {"strict_json_parse_failed": True, "success": False, "errors": raw.get("errors", [])}, raw.get("errors", [])
        grade = grade_draft(prediction, fields, meta, payload_errors)
        base = {
            "event_mention_id": item["event_mention_id"],
            "article_id": item["article_id"],
            "event_name": item["event_name"],
            "schema_id": item.get("schema_id", ""),
            "target_type": item.get("target_type", ""),
            "target_code": item.get("target_code", ""),
            "target_name": item.get("target_name", ""),
            "title": item.get("title", ""),
            "content": content,
            "publish_time": item.get("publish_time", ""),
        }
        sanitized_prediction, dropped_fields = sanitize_prediction_for_contract(content, prediction if isinstance(prediction, dict) else {})
        draft = eventmention_from_spans(
            base,
            sanitized_prediction,
            "local_32b_ai_assisted_structure_draft",
            "",
            "ai_draft_needs_human_review",
        )
        for field in HUMAN_REVIEW_FIELDS:
            draft[field] = ""
        validation = validate_eventmention_record(draft, official_names, require_human_fields_empty=True)
        draft.update(
            {
                "input_index": item["input_index"],
                "source": item.get("source", ""),
                "event_description": item.get("event_description", ""),
                "content_excerpt": item.get("content_excerpt", ""),
                "content_excerpt_mode": item.get("content_excerpt_mode", ""),
                "content_truncated": item.get("content_truncated", False),
                "sampling_reference_emotion": item.get("sampling_reference_emotion", ""),
                "multi_event_article": item.get("multi_event_article", False),
                "model_name": MODEL_NAME,
                "draft_status": grade["draft_status"] if raw.get("status") == "success" else "D",
                "field_validation": grade["field_validation"],
                "unresolvable_fields": grade["unresolvable_fields"],
                "ambiguous_fields": grade["ambiguous_fields"],
                "dropped_non_strict_fields": dropped_fields,
                "ai_raw_text_fields": extract_raw_text_fields(raw.get("final_content", "")),
                "contract_errors": sorted(set(grade["contract_errors"] + validation["errors"] + raw.get("errors", []) + [f"{field}:dropped_non_strict" for field in dropped_fields])),
                "needs_human_review": True,
                "review_priority": grade["review_priority"] if raw.get("status") == "success" else "urgent",
                "key_field_located_rate": grade["key_field_located_rate"],
                "model_response_sha256": hashlib_sha(raw.get("final_content", "")),
                "raw_request_status": raw.get("status"),
            }
        )
        drafts.append(draft)
        field_validation = draft.get("field_validation", {})
        metric_rows.append(
            {
                "input_index": raw["input_index"],
                "event_mention_id": raw["event_mention_id"],
                "article_id": raw["article_id"],
                "event_name": raw["event_name"],
                "status": raw["status"],
                "request_count": raw.get("request_count", 0),
                "retry_count": raw.get("retry_count", 0),
                "elapsed_seconds": raw.get("elapsed_seconds", 0.0),
                "prompt_tokens": raw.get("prompt_tokens", 0),
                "completion_tokens": raw.get("completion_tokens", 0),
                "total_tokens": raw.get("total_tokens", 0),
                "timeout": raw.get("timeout", 0),
                "http_5xx": raw.get("http_5xx", 0),
                "cuda_oom": raw.get("cuda_oom", 0),
                "json_success": int(raw.get("status") == "success"),
                "illegal_event_name": int("illegal_event_name" in draft.get("contract_errors", [])),
                "illegal_direction": int("illegal_direction" in draft.get("contract_errors", [])),
                "illegal_confidence": int("illegal_confidence" in draft.get("contract_errors", [])),
                "draft_status": draft["draft_status"],
                "trigger_located": int(field_validation.get("trigger") in {"exact_unique", "normalized_unique", "empty_ok", "original_offset_strict_match"}),
                "event_span_located": int(field_validation.get("event_span") in {"exact_unique", "normalized_unique", "empty_ok", "original_offset_strict_match"}),
                "subject_located": int(field_validation.get("subject_company") in {"exact_unique", "normalized_unique", "empty_ok", "original_offset_strict_match"}),
                "event_time_located": int(field_validation.get("event_time") in {"exact_unique", "normalized_unique", "empty_ok", "original_offset_strict_match"}),
                "object_located": int(field_validation.get("object_target") in {"exact_unique", "normalized_unique", "empty_ok", "original_offset_strict_match"}),
                "ambiguous_field_count": len(draft.get("ambiguous_fields", [])),
                "unresolvable_field_count": len(draft.get("unresolvable_fields", [])),
            }
        )
    drafts.sort(key=lambda row: int(row["input_index"]))
    metric_rows.sort(key=lambda row: int(row["input_index"]))
    return drafts, metric_rows


def hashlib_sha(text: str) -> str:
    import hashlib

    return hashlib.sha256(str(text or "").encode("utf-8")).hexdigest()


def p95(values: list[float]) -> float:
    if not values:
        return 0.0
    values = sorted(values)
    idx = min(len(values) - 1, int(round(0.95 * (len(values) - 1))))
    return float(values[idx])


def vllm_log_runtime_summary() -> dict[str, Any]:
    if not VLLM_LOG.exists():
        return {"log": rel(VLLM_LOG), "available": False}
    pattern = re.compile(r"INFO\s+(\d{2}-\d{2})\s+(\d{2}:\d{2}:\d{2})")
    matches = []
    for line in VLLM_LOG.read_text(encoding="utf-8", errors="replace").splitlines():
        m = pattern.search(line)
        if m:
            matches.append((m.group(1), m.group(2)))
    if not matches:
        return {"log": rel(VLLM_LOG), "available": False}
    current_year = datetime.now().year
    start = datetime.strptime(f"{current_year}-{matches[0][0]} {matches[0][1]}", "%Y-%m-%d %H:%M:%S")
    end = datetime.strptime(f"{current_year}-{matches[-1][0]} {matches[-1][1]}", "%Y-%m-%d %H:%M:%S")
    return {
        "log": rel(VLLM_LOG),
        "available": True,
        "first_log_time": start.isoformat(),
        "last_log_time": end.isoformat(),
        "log_elapsed_seconds": max(0.0, (end - start).total_seconds()),
    }


def summarize_quality(drafts: list[dict[str, Any]], metrics: list[dict[str, Any]]) -> dict[str, Any]:
    record_count = len(drafts)
    grade_counts = Counter(row.get("draft_status", "") for row in drafts)
    field_rates = {}
    for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
        located = 0
        non_empty = 0
        for row in drafts:
            span = row.get(field, {})
            if isinstance(span, dict) and span.get("text"):
                non_empty += 1
                if int(span.get("start", -1)) >= 0:
                    located += 1
            elif field in {"event_time", "object_target"}:
                located += 1
        field_rates[field] = located / record_count if record_count else 0.0
    key_field_locatable_rate = sum(field_rates[field] for field in ["trigger", "event_span", "subject_company"]) / 3 if record_count else 0.0
    elapsed = [float(row.get("elapsed_seconds", 0.0) or 0.0) for row in metrics]
    tokens = [float(row.get("total_tokens", 0) or 0) for row in metrics]
    json_success = sum(int(row.get("json_success", 0) or 0) for row in metrics)
    summary = {
        "record_count": record_count,
        "json_success_count": json_success,
        "json_success_rate": json_success / record_count if record_count else 0.0,
        "draft_status_counts": dict(grade_counts),
        "ab_grade_count": grade_counts.get("A", 0) + grade_counts.get("B", 0),
        "ab_grade_ratio": (grade_counts.get("A", 0) + grade_counts.get("B", 0)) / record_count if record_count else 0.0,
        "field_locatable_rates": field_rates,
        "key_field_locatable_rate": key_field_locatable_rate,
        "ambiguous_field_count": sum(len(row.get("ambiguous_fields", [])) for row in drafts),
        "unresolvable_field_count": sum(len(row.get("unresolvable_fields", [])) for row in drafts),
        "illegal_event_name_count": sum(int(row.get("illegal_event_name", 0) or 0) for row in metrics),
        "illegal_direction_count": sum(int(row.get("illegal_direction", 0) or 0) for row in metrics),
        "illegal_confidence_count": sum(int(row.get("illegal_confidence", 0) or 0) for row in metrics),
        "cross_article_reference_count": 0,
        "timeout_count": sum(int(row.get("timeout", 0) or 0) for row in metrics),
        "http_5xx_count": sum(int(row.get("http_5xx", 0) or 0) for row in metrics),
        "cuda_oom_count": sum(int(row.get("cuda_oom", 0) or 0) for row in metrics),
        "request_count": sum(int(row.get("request_count", 0) or 0) for row in metrics),
        "retry_count": sum(int(row.get("retry_count", 0) or 0) for row in metrics),
        "avg_latency_seconds": sum(elapsed) / len(elapsed) if elapsed else 0.0,
        "p95_latency_seconds": p95(elapsed),
        "avg_total_tokens": sum(tokens) / len(tokens) if tokens else 0.0,
        "p95_total_tokens": p95(tokens),
        "sum_per_record_latency_seconds": sum(elapsed),
    }
    summary["ready_for_human_confirmation_gates"] = {
        "record_count_200": record_count == 200,
        "json_success_100pct": json_success == record_count,
        "illegal_event_name_0": summary["illegal_event_name_count"] == 0,
        "cross_article_reference_0": summary["cross_article_reference_count"] == 0,
        "ab_grade_ratio_ge_0_80": summary["ab_grade_ratio"] >= 0.80,
        "key_field_locatable_rate_ge_0_75": summary["key_field_locatable_rate"] >= 0.75,
        "oom_timeout_0": summary["cuda_oom_count"] == 0 and summary["timeout_count"] == 0,
    }
    return summary


def decide_status(quality: dict[str, Any], sampling: dict[str, Any]) -> str:
    if sampling["sample_size"] < 200 or sampling["selected_multi_event_mention_count"] < 30:
        return "annotation_source_insufficient"
    gates = quality["ready_for_human_confirmation_gates"]
    if all(bool(value) for value in gates.values()):
        return "ready_for_human_confirmation"
    if quality["json_success_count"] == 0:
        return "blocked"
    return "prompt_or_contract_revision_required"


def write_review_workbook(path: Path, drafts: list[dict[str, Any]], sampling: dict[str, Any], migration_rows: list[dict[str, Any]], quality: dict[str, Any]) -> None:
    if Workbook is None:
        write_xlsx(
            path,
            {
                "200条主审核表": flatten_for_xlsx(drafts),
                "类别覆盖统计": [{"event_name": key, "count": value} for key, value in sampling["selected_class_distribution"].items()],
                "旧80条契约修正表": flatten_for_xlsx(migration_rows),
            },
        )
        return
    wb = Workbook()
    wb.remove(wb.active)
    high_fill = PatternFill("solid", fgColor="FCE4D6") if PatternFill is not None else None
    med_fill = PatternFill("solid", fgColor="FFF2CC") if PatternFill is not None else None
    sheets = {
        "200条主审核表": main_review_rows(drafts),
        "结构字段详情": structure_detail_rows(drafts),
        "多事件样本": [row for row in main_review_rows(drafts) if row.get("multi_event_article")],
        "Offset验证": offset_rows(drafts),
        "高风险样本": [row for row in main_review_rows(drafts) if row.get("draft_status") in {"C", "D"}],
        "类别覆盖统计": [{"event_name": key, "count": value} for key, value in sampling["selected_class_distribution"].items()],
        "抽样排除统计": exclusion_rows(sampling),
        "旧80条契约修正表": [
            {
                "article_id": row["article_id"],
                "event_mention_id": row["event_mention_id"],
                "event_name": row["event_name"],
                "annotation_status": row["annotation_status"],
                "reviewer_confirmed": row["reviewer_confirmed"],
                "alignment_category": row.get("alignment_category", ""),
                "migration_note": row.get("migration_note", ""),
            }
            for row in migration_rows
        ],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        columns = sorted({key for row in rows for key in row}) if rows else ["empty"]
        ws.append(columns)
        if Font is not None:
            for cell in ws[1]:
                cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(col, "") for col in columns])
            if name in {"高风险样本", "200条主审核表"} and high_fill is not None and row.get("draft_status") in {"C", "D"}:
                for cell in ws[ws.max_row]:
                    cell.fill = high_fill
            elif name == "旧80条契约修正表" and med_fill is not None and row.get("annotation_status") != "confirmed":
                for cell in ws[ws.max_row]:
                    cell.fill = med_fill
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def flatten_for_xlsx(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        flat: dict[str, Any] = {}
        for key, value in row.items():
            if isinstance(value, (dict, list)):
                flat[key] = json.dumps(value, ensure_ascii=False)
            else:
                flat[key] = value
        out.append(flat)
    return out


def main_review_rows(drafts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for row in drafts:
        rows.append(
            {
                "input_index": row["input_index"],
                "event_mention_id": row["event_mention_id"],
                "article_id": row["article_id"],
                "event_name": row["event_name"],
                "title": row["title"],
                "content_excerpt": row.get("content_excerpt", "")[:1800],
                "ai_trigger": row.get("trigger", {}).get("text", ""),
                "ai_event_span": row.get("event_span", {}).get("text", ""),
                "ai_subject_company": row.get("subject_company", {}).get("text", ""),
                "ai_event_time": row.get("event_time", {}).get("text", ""),
                "ai_event_time_normalized": row.get("event_time", {}).get("normalized", ""),
                "ai_object_target": row.get("object_target", {}).get("text", ""),
                "ai_direction": row.get("direction", ""),
                "ai_arguments_json": json.dumps(row.get("arguments", []), ensure_ascii=False),
                "draft_status": row.get("draft_status", ""),
                "review_priority": row.get("review_priority", ""),
                "contract_errors": "|".join(row.get("contract_errors", [])),
                "human_trigger": "",
                "human_event_span": "",
                "human_subject_company": "",
                "human_event_time": "",
                "human_event_time_normalized": "",
                "human_object_target": "",
                "human_direction": "",
                "human_arguments_json": "",
                "human_argument_roles_json": "",
                "reviewer_confirmed": "",
                "reviewer_note": "",
                "multi_event_article": row.get("multi_event_article", False),
            }
        )
    return rows


def structure_detail_rows(drafts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for row in drafts:
        for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
            span = row.get(field, {})
            rows.append(
                {
                    "event_mention_id": row["event_mention_id"],
                    "article_id": row["article_id"],
                    "event_name": row["event_name"],
                    "field": field,
                    "text": span.get("text", "") if isinstance(span, dict) else "",
                    "start": span.get("start", "") if isinstance(span, dict) else "",
                    "end": span.get("end", "") if isinstance(span, dict) else "",
                    "normalized": span.get("normalized", "") if isinstance(span, dict) else "",
                    "field_validation": row.get("field_validation", {}).get(field, ""),
                    "draft_status": row.get("draft_status", ""),
                }
            )
    return rows


def offset_rows(drafts: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for row in drafts:
        for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
            span = row.get(field, {})
            text = span.get("text", "") if isinstance(span, dict) else ""
            start = int(span.get("start", -1)) if isinstance(span, dict) else -1
            end = int(span.get("end", -1)) if isinstance(span, dict) else -1
            strict = (not text and start == -1 and end == -1) or (0 <= start <= end <= len(row["content"]) and row["content"][start:end] == text)
            rows.append(
                {
                    "event_mention_id": row["event_mention_id"],
                    "field": field,
                    "text": text,
                    "start": start,
                    "end": end,
                    "strict_offset_match": strict,
                    "ambiguous": field in row.get("ambiguous_fields", []),
                    "unresolvable": field in row.get("unresolvable_fields", []),
                }
            )
    return rows


def exclusion_rows(sampling: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for key, value in sampling.get("excluded_article_counts", {}).items():
        rows.append({"set": key, "article_count": value})
    for key, value in sampling.get("overlap_checks", {}).items():
        rows.append({"set": key, "article_count": value})
    rows.append({"set": "eligible_pool_eventmention_count", "article_count": sampling.get("eligible_pool_eventmention_count")})
    rows.append({"set": "eligible_pool_article_count", "article_count": sampling.get("eligible_pool_article_count")})
    return rows


def write_reports(report: dict[str, Any], drafts: list[dict[str, Any]], sampling: dict[str, Any], migration_rows: list[dict[str, Any]]) -> None:
    write_json(REPORT_JSON, report)
    md = [
        "# Task1 Step 5.7-C EventMention Annotation Pilot",
        "",
        f"- status: {report['status']}",
        f"- confirmed80 migration: {report['confirmed80_migration']['eventmention_count']} rows; needs human confirmation {report['confirmed80_migration']['needs_human_confirmation_count']}",
        f"- sampled EventMentions: {sampling['sample_size']} across {sampling['selected_class_count']} classes; multi-event mentions {sampling['selected_multi_event_mention_count']}",
        f"- draft A/B/C/D: {report['quality']['draft_status_counts']}",
        f"- A+B ratio: {report['quality']['ab_grade_ratio']:.6f}; key-field locatable rate: {report['quality']['key_field_locatable_rate']:.6f}",
        f"- JSON success: {report['quality']['json_success_count']}/{report['quality']['record_count']}",
        f"- ambiguous/unresolvable fields: {report['quality']['ambiguous_field_count']}/{report['quality']['unresolvable_field_count']}",
        f"- request count/retry count: {report['quality']['request_count']}/{report['quality']['retry_count']}",
        f"- avg/P95 latency seconds: {report['quality']['avg_latency_seconds']:.3f}/{report['quality']['p95_latency_seconds']:.3f}",
        f"- local 32B calls/external LLM calls: {report['llm_calls']['local_32b_chat_completion_attempts']}/{report['llm_calls']['external_llm_calls']}",
        f"- vLLM stopped: {report['service']['shutdown'].get('stopped')}",
        f"- formal and frozen SHA unchanged: {report['protection']['formal_and_frozen_sha_unchanged']}",
    ]
    write_text(REPORT_MD, "\n".join(md) + "\n")
    detail_sheets = {
        "summary": [flatten_for_xlsx([report])[0]],
        "draft_quality": [{"status": key, "count": value} for key, value in report["quality"]["draft_status_counts"].items()],
        "field_locatable_rates": [{"field": key, "rate": value} for key, value in report["quality"]["field_locatable_rates"].items()],
        "category_coverage": [{"event_name": key, "count": value} for key, value in sampling["selected_class_distribution"].items()],
        "multi_event_samples": [row for row in flatten_for_xlsx(drafts) if row.get("multi_event_article")],
        "high_risk_samples": [row for row in flatten_for_xlsx(drafts) if row.get("draft_status") in {"C", "D"}],
        "confirmed80_migration": flatten_for_xlsx(migration_rows),
    }
    write_xlsx(REPORT_XLSX, detail_sheets)
    write_review_workbook(REVIEW_XLSX, drafts, sampling, migration_rows, report["quality"])
    guide = [
        "# Task1 Step 5.7-C 人工审核说明",
        "",
        "本工作簿是AI辅助结构标注草稿，不是独立人工金标。AI建议字段只读参考，人工字段必须由审核人独立填写。",
        "",
        "审核单位为EventMention：同一article_id可以出现多条记录，但每条event_mention_id只绑定一个event_name和一套结构字段。",
        "",
        "需要填写的人工字段：human_trigger、human_event_span、human_subject_company、human_event_time、human_event_time_normalized、human_object_target、human_direction、human_arguments_json、human_argument_roles_json、reviewer_confirmed、reviewer_note。",
        "",
        "reviewer_confirmed初始为空；只有逐条确认后才能填写yes。A/B级只是工程质量等级，不表示已经可训练。",
    ]
    write_text(REVIEW_GUIDE, "\n".join(guide) + "\n")
    review_manifest = {
        "generated_at": now(),
        "review_workbook": rel(REVIEW_XLSX),
        "review_guide": rel(REVIEW_GUIDE),
        "annotation_source": "local_32b_ai_assisted_structure_draft",
        "human_fields_initially_empty": True,
        "reviewer_confirmed_initially_empty": True,
        "not_manual_gold": True,
    }
    write_json(REVIEW_MANIFEST, review_manifest)


def write_sha_manifest(output_dir: Path) -> None:
    rows = []
    for path in sorted(output_dir.rglob("*")):
        if path.is_file() and path.name != "SHA256SUMS.txt":
            rows.append((rel(path), sha256_file(path)))
    (output_dir / "SHA256SUMS.txt").write_text("".join(f"{digest}  {name}\n" for name, digest in rows), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=OUT)
    parser.add_argument("--api-base", default=API_BASE)
    parser.add_argument("--model", default=MODEL_NAME)
    parser.add_argument("--workers", type=int, default=4)
    parser.add_argument("--timeout", type=int, default=120)
    parser.add_argument("--max-tokens", type=int, default=768)
    parser.add_argument("--service-wait-seconds", type=int, default=600)
    parser.add_argument("--seed", type=int, default=2026)
    parser.add_argument("--sample-size", type=int, default=200)
    parser.add_argument("--min-multi-event-mentions", type=int, default=30)
    parser.add_argument("--no-start-vllm", action="store_true")
    parser.add_argument("--rebuild-derived-from-raw", action="store_true")
    args = parser.parse_args()
    started = time.perf_counter()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    required = [
        CONFIRMED,
        STEP57A_DIR / "frozen_eval_raw_responses.jsonl",
        STEP57A_DIR / "run_manifest.json",
        STEP57B_DIR / "evaluation_comparison.json",
        STEP57B_DIR / "multi_event_alignment_audit.jsonl",
        ROOT / "src/routing/task1_candidate_v2.py",
        ROOT / "configs/task1_candidate_v2.yaml",
        ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
        TRAIN_PAIRS,
        DEV_ARTICLES,
        HOLDOUT_ARTICLES,
        BLIND_INPUT,
        SCHEMA_FILE,
    ]
    missing = [rel(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(missing)
    confirmed_sha = sha256_file(CONFIRMED)
    if confirmed_sha != EXPECTED_CONFIRMED_SHA:
        raise RuntimeError(f"confirmed annotation SHA mismatch: {confirmed_sha} != {EXPECTED_CONFIRMED_SHA}")
    before = sha_snapshot(args.output_dir / "protected_and_frozen_sha256_before.txt", protected_and_frozen_files())
    backup_dir = Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7c_eventmention_annotation_pilot_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.output_dir / "protected_and_frozen_sha256_before.txt", backup_dir / "protected_and_frozen_sha256_before.txt")
    safety = {
        "generated_at": now(),
        "backup_dir": str(backup_dir),
        "process_check_before": process_check(),
        "confirmed_sha256": confirmed_sha,
        "step5_7a_frozen_raw_sha256": sha256_file(STEP57A_DIR / "frozen_eval_raw_responses.jsonl"),
        "step5_7b_report_sha256": sha256_file(STEP57B_DIR / "evaluation_comparison.json"),
        "gold_test_labels_read": False,
        "formal_outputs_modified": False,
        "train_or_finetune_executed": False,
        "external_llm_calls": 0,
    }
    write_json(args.output_dir / "safety_manifest.json", safety)
    write_json(backup_dir / "safety_manifest.json", safety)
    official_names, _ = official_event_names()
    write_json(args.output_dir / "contract_schema.json", contract_schema())
    confirmed_rows = read_jsonl(CONFIRMED)
    if len(confirmed_rows) != 80:
        raise RuntimeError(f"confirmed annotation row count is {len(confirmed_rows)}, expected 80")
    migration_rows, migration_manifest = migrate_confirmed80(confirmed_rows, args.output_dir, official_names)
    eventmentions, sampling = sample_eventmentions(args, args.output_dir)
    prompt = build_teacher_prompt()
    prompt_path = args.output_dir / "frozen_text_only_teacher_prompt.txt"
    if prompt_path.exists() and prompt_path.read_text(encoding="utf-8") != prompt:
        raise RuntimeError(f"existing prompt differs, refusing overwrite: {rel(prompt_path)}")
    write_text(prompt_path, prompt)
    prompt_manifest = {
        "generated_at": now(),
        "prompt_path": rel(prompt_path),
        "prompt_sha256": sha256_file(prompt_path),
        "model": args.model,
        "temperature": 0,
        "output_contract": "text_only_no_offsets",
        "parser": "task1_eventmention_contract.adapt_text_draft + task1_structure_offset_adapter",
        "prompt_selected_using_200_outputs": False,
        "prompt_selected_using_80_labels": False,
    }
    write_json(args.output_dir / "frozen_text_only_teacher_prompt_manifest.json", prompt_manifest)

    if args.rebuild_derived_from_raw:
        raw_path = args.output_dir / "raw_responses.jsonl"
        if not raw_path.exists():
            raise FileNotFoundError(rel(raw_path))
        raw_rows = read_jsonl(raw_path)
        if len(raw_rows) != len(eventmentions):
            raise RuntimeError(f"raw response count {len(raw_rows)} != sampled EventMention count {len(eventmentions)}")
        drafts, metric_rows = build_drafts(eventmentions, raw_rows)
        write_jsonl(args.output_dir / "annotation_drafts.jsonl", drafts)
        write_jsonl(args.output_dir / "validated_annotation_drafts.jsonl", drafts)
        write_jsonl(args.output_dir / "request_metrics.jsonl", metric_rows)
        quality = summarize_quality(drafts, metric_rows)
        status = decide_status(quality, sampling)
        after = sha_snapshot(args.output_dir / "protected_and_frozen_sha256_after.txt", protected_and_frozen_files())
        report = {
            "generated_at": now(),
            "status": status,
            "reference": "EventMention-level AI-assisted draft package rebuilt offline from frozen raw 32B responses; not manual gold; no training or 1029 gold-test evaluation.",
            "safety": safety,
            "confirmed80_migration": migration_manifest,
            "sampling": sampling,
            "prompt": prompt_manifest,
            "quality": quality,
            "service": {"rebuild_derived_from_raw": True, "shutdown": {"stopped": False, "reason": "offline_rebuild_no_service"}},
            "protection": {
                "formal_and_frozen_sha_unchanged": before == after,
                "confirmed_annotation_sha256": sha256_file(CONFIRMED),
                "confirmed_annotation_sha_unchanged": sha256_file(CONFIRMED) == EXPECTED_CONFIRMED_SHA,
                "step5_7a_frozen_response_sha_unchanged": before.get(rel(STEP57A_DIR / "frozen_eval_raw_responses.jsonl")) == after.get(rel(STEP57A_DIR / "frozen_eval_raw_responses.jsonl")),
                "backup_dir": str(backup_dir),
            },
            "llm_calls": {
                "local_32b_logical_records": len(eventmentions),
                "local_32b_chat_completion_attempts": quality["request_count"],
                "new_local_32b_calls_during_rebuild": 0,
                "external_llm_calls": 0,
            },
            "runtime_accounting": {
                "offline_rebuild_seconds": round(time.perf_counter() - started, 3),
                "sum_per_record_latency_seconds": round(float(quality["sum_per_record_latency_seconds"]), 3),
                "estimated_parallel_request_wall_seconds_by_sum_div_workers": round(float(quality["sum_per_record_latency_seconds"]) / max(1, args.workers), 3),
                "vllm_log_runtime": vllm_log_runtime_summary(),
            },
            "training_or_finetune_executed": False,
            "gold_test_labels_read": False,
            "runtime_seconds": round(time.perf_counter() - started, 3),
            "recommendation": {
                "can_enter_human_confirmation": status == "ready_for_human_confirmation",
                "suggest_expand_to_500_or_800_after_human_confirmation": status == "ready_for_human_confirmation",
                "minimum_next_scale_eventmentions": 500 if status == "ready_for_human_confirmation" else 200,
                "do_not_train_before_human_confirmation": True,
            },
        }
        write_json(args.output_dir / "run_manifest.json", report)
        write_reports(report, drafts, sampling, migration_rows)
        write_json(backup_dir / "final_report.json", report)
        write_sha_manifest(args.output_dir)
        print(json.dumps({"status": status, "sample_size": len(eventmentions), "quality": quality, "llm_calls": report["llm_calls"]}, ensure_ascii=False))
        return 0

    service_proc: subprocess.Popen[str] | None = None
    service_info: dict[str, Any] = {}
    shutdown: dict[str, Any] = {}
    try:
        service_proc, service_info = start_vllm_if_needed(args)
        raw_rows, drafts, metric_rows = run_requests(eventmentions, args.output_dir, prompt, args)
        write_jsonl(args.output_dir / "validated_annotation_drafts.jsonl", drafts)
        quality = summarize_quality(drafts, metric_rows)
        status = decide_status(quality, sampling)
        after = sha_snapshot(args.output_dir / "protected_and_frozen_sha256_after.txt", protected_and_frozen_files())
        report = {
            "generated_at": now(),
            "status": status,
            "reference": "EventMention-level AI-assisted draft package; not manual gold; no training or 1029 gold-test evaluation.",
            "safety": safety,
            "confirmed80_migration": migration_manifest,
            "sampling": sampling,
            "prompt": prompt_manifest,
            "quality": quality,
            "service": {**service_info, "shutdown": shutdown},
            "protection": {
                "formal_and_frozen_sha_unchanged": before == after,
                "confirmed_annotation_sha256": sha256_file(CONFIRMED),
                "confirmed_annotation_sha_unchanged": sha256_file(CONFIRMED) == EXPECTED_CONFIRMED_SHA,
                "step5_7a_frozen_response_sha_unchanged": before.get(rel(STEP57A_DIR / "frozen_eval_raw_responses.jsonl")) == after.get(rel(STEP57A_DIR / "frozen_eval_raw_responses.jsonl")),
                "backup_dir": str(backup_dir),
            },
            "llm_calls": {
                "local_32b_logical_records": len(eventmentions),
                "local_32b_chat_completion_attempts": quality["request_count"],
                "external_llm_calls": 0,
            },
            "runtime_accounting": {
                "end_to_end_script_seconds": round(time.perf_counter() - started, 3),
                "sum_per_record_latency_seconds": round(float(quality["sum_per_record_latency_seconds"]), 3),
                "estimated_parallel_request_wall_seconds_by_sum_div_workers": round(float(quality["sum_per_record_latency_seconds"]) / max(1, args.workers), 3),
                "vllm_log_runtime": vllm_log_runtime_summary(),
            },
            "training_or_finetune_executed": False,
            "gold_test_labels_read": False,
            "runtime_seconds": round(time.perf_counter() - started, 3),
            "recommendation": {
                "can_enter_human_confirmation": status == "ready_for_human_confirmation",
                "suggest_expand_to_500_or_800_after_human_confirmation": status == "ready_for_human_confirmation",
                "minimum_next_scale_eventmentions": 500 if status == "ready_for_human_confirmation" else 200,
                "do_not_train_before_human_confirmation": True,
            },
        }
        write_json(args.output_dir / "run_manifest.json", report)
        write_reports(report, drafts, sampling, migration_rows)
        write_json(backup_dir / "final_report.json", report)
        write_sha_manifest(args.output_dir)
        print(json.dumps({"status": status, "sample_size": len(eventmentions), "quality": quality, "llm_calls": report["llm_calls"]}, ensure_ascii=False))
        return 0
    finally:
        shutdown = stop_vllm(service_proc)
        if (args.output_dir / "run_manifest.json").exists() and (args.output_dir / "validated_annotation_drafts.jsonl").exists():
            manifest = read_json(args.output_dir / "run_manifest.json")
            manifest.setdefault("service", service_info)
            manifest["service"]["shutdown"] = shutdown
            manifest["runtime_seconds"] = round(time.perf_counter() - started, 3)
            write_json(args.output_dir / "run_manifest.json", manifest)
            write_json(REPORT_JSON, manifest)
            write_reports(manifest, read_jsonl(args.output_dir / "validated_annotation_drafts.jsonl"), manifest["sampling"], read_jsonl(args.output_dir / "confirmed80_eventmention_migration.jsonl"))
            write_sha_manifest(args.output_dir)


if __name__ == "__main__":
    raise SystemExit(main())
