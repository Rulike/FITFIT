from __future__ import annotations

"""CSMAR file adapter and field-mapping template for task 2 financial checks.

The repository currently does not ship any local CSMAR data files. This module
therefore focuses on a safe, explicit adapter boundary:

- load records only from caller-provided files
- canonicalize common CSMAR-like field names
- expose a field mapping template for future data drops
- never fabricate missing values
"""

import csv
import hashlib
import json
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Mapping

import pandas as pd
from openpyxl import load_workbook


EXPECTED_TABLES = {
    "FI_T1": "liquidity_and_solvency",
    "FI_T5": "profitability_and_rd",
    "FI_T6": "cash_flow",
    "FI_T8": "growth_and_industry",
    "FI_T10": "valuation",
}


CANONICAL_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "table_name": ("table_name", "table", "sheet", "sheet_name", "source_table"),
    "record_id": ("record_id", "source_record_id", "row_id", "id", "fid"),
    "company_name": ("company_name", "comp_name", "corp_name", "S_INFO_COMPNAME", "证券简称", "公司名称"),
    "stock_code": ("stock_code", "stkcd", "code", "S_INFO_WINDCODE", "证券代码", "股票代码"),
    "industry_name": ("industry_name", "industry", "ind_name", "行业名称"),
    "industry_code": ("industry_code", "ind_code", "行业代码"),
    "report_period": ("report_period", "report_date", "rpt_date", "end_date", "period", "报告期"),
    "publish_date": ("publish_date", "ann_date", "publish_time", "公告日期", "发布日期"),
    "year": ("year", "fiscal_year", "report_year", "年度"),
    "quarter": ("quarter", "fiscal_quarter", "report_quarter", "季度"),
    "roe": ("roe", "ROE", "return_on_equity", "净资产收益率"),
    "roe_yoy": ("roe_yoy", "roe_change", "roe_growth", "ROE同比", "ROE增长率"),
    "roa": ("roa", "ROA", "return_on_assets", "总资产收益率"),
    "roa_yoy": ("roa_yoy", "roa_change", "roa_growth", "ROA同比", "ROA增长率"),
    "revenue": ("revenue", "operating_revenue", "营业收入", "营收"),
    "revenue_yoy": ("revenue_yoy", "revenue_growth", "营业收入同比", "营收同比", "收入同比"),
    "net_profit": ("net_profit", "np", "profit", "归母净利润", "净利润"),
    "net_profit_yoy": ("net_profit_yoy", "net_profit_growth", "净利润同比", "归母净利润同比"),
    "ocf": ("ocf", "operating_cash_flow", "经营活动现金流净额", "经营现金流净额"),
    "ocf_yoy": ("ocf_yoy", "ocf_growth", "经营活动现金流同比", "经营现金流同比"),
    "gross_margin": ("gross_margin", "gross_profit_margin", "毛利率"),
    "gross_margin_yoy": ("gross_margin_yoy", "gross_margin_change", "毛利率同比"),
    "current_ratio": ("current_ratio", "流动比率"),
    "current_ratio_yoy": ("current_ratio_yoy", "current_ratio_change", "流动比率同比"),
    "quick_ratio": ("quick_ratio", "速动比率"),
    "quick_ratio_yoy": ("quick_ratio_yoy", "quick_ratio_change", "速动比率同比"),
    "debt_ratio": ("debt_ratio", "asset_liability_ratio", "资产负债率"),
    "debt_ratio_yoy": ("debt_ratio_yoy", "debt_ratio_change", "资产负债率同比"),
    "rd_expense": ("rd_expense", "research_and_development_expense", "研发费用", "研发投入"),
    "rd_expense_yoy": ("rd_expense_yoy", "rd_expense_growth", "研发费用同比"),
    "rd_expense_ratio": ("rd_expense_ratio", "rd_ratio", "研发费用率", "研发投入占比"),
    "rd_expense_ratio_yoy": ("rd_expense_ratio_yoy", "rd_expense_ratio_change", "研发费用率同比"),
    "accounts_receivable": ("accounts_receivable", "ar", "应收账款"),
    "accounts_receivable_yoy": ("accounts_receivable_yoy", "ar_growth", "应收账款同比"),
    "inventory": ("inventory", "存货"),
    "inventory_yoy": ("inventory_yoy", "inventory_growth", "存货同比"),
    "non_recurring_profit": ("non_recurring_profit", "non_recurring_gains", "non_recurring_loss", "非经常性损益"),
    "non_recurring_profit_ratio": ("non_recurring_profit_ratio", "non_recurring_profit_share", "非经常性损益占比"),
    "pe_ratio": ("pe_ratio", "PE", "pe", "市盈率"),
    "industry_mean_pe": ("industry_mean_pe", "industry_pe", "行业均值PE", "行业平均PE"),
    "industry_revenue_growth_mean": ("industry_revenue_growth_mean", "industry_mean_revenue_growth", "行业均值营收增速"),
}


FIELD_MAPPING_TEMPLATE = {
    "tables": EXPECTED_TABLES,
    "canonical_fields": {key: list(value) for key, value in CANONICAL_FIELD_ALIASES.items()},
    "notes": [
        "The adapter expects caller-supplied files.",
        "No local CSMAR data files are bundled with the repository.",
        "If the actual vendor export uses different column names, map them explicitly before running task 2 financial checks.",
    ],
}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def first_nonempty(record: Mapping[str, Any], names: Iterable[str]) -> Any:
    for name in names:
        if name in record and normalize_text(record.get(name)):
            return record.get(name)
    return None


def canonical_record(record: Mapping[str, Any], table_name: str | None = None, row_index: int | None = None) -> dict[str, Any]:
    out = dict(record)
    if table_name:
        out.setdefault("table_name", table_name)
    if row_index is not None:
        out.setdefault("row_index", row_index)
    for canonical, aliases in CANONICAL_FIELD_ALIASES.items():
        value = first_nonempty(record, aliases)
        if value is not None and normalize_text(value) != "":
            out[canonical] = value
    if not normalize_text(out.get("record_id")):
        prefix = normalize_text(out.get("table_name")) or "CSMAR"
        out["record_id"] = f"{prefix}:{row_index if row_index is not None else len(out)}"
    return out


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            obj = json.loads(line)
            if isinstance(obj, dict):
                rows.append(obj)
    return rows


def _load_json(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if isinstance(payload, dict):
        rows = payload.get("rows") or payload.get("data") or []
        if isinstance(rows, list):
            return [item for item in rows if isinstance(item, dict)]
    return []


def _load_tabular(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix in {".parquet", ".pq"}:
        return pd.read_parquet(path)
    if suffix in {".csv", ".tsv"}:
        return pd.read_csv(path, sep="\t" if suffix == ".tsv" else ",")
    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    raise ValueError(f"unsupported CSMAR file type: {path.suffix}")


def load_records_from_path(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        return _load_jsonl(path)
    if suffix == ".json":
        return _load_json(path)
    if suffix in {".csv", ".tsv", ".parquet", ".pq", ".xlsx", ".xls"}:
        frame = _load_tabular(path)
        return frame.to_dict("records")
    raise ValueError(f"unsupported CSMAR file type: {path.suffix}")


def discover_source_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    if not path.exists():
        return []
    patterns = ("*.parquet", "*.pq", "*.csv", "*.tsv", "*.jsonl", "*.json", "*.xlsx", "*.xls")
    files: list[Path] = []
    for pattern in patterns:
        files.extend(sorted(path.glob(pattern)))
    return files


def load_csmar_records(source: str | Path | Iterable[str | Path], table_name: str | None = None) -> list[dict[str, Any]]:
    if source is None:
        return []
    if isinstance(source, (str, Path)):
        candidates = discover_source_files(Path(source))
    else:
        candidates = []
        for item in source:
            candidates.extend(discover_source_files(Path(item)))

    records: list[dict[str, Any]] = []
    for path in candidates:
        loaded = load_records_from_path(path)
        for idx, row in enumerate(loaded):
            if not isinstance(row, dict):
                continue
            canonical = canonical_record(row, table_name=table_name or path.stem, row_index=idx)
            canonical["source_path"] = str(path)
            records.append(canonical)
    return records


def normalize_alias_map(alias_map: Mapping[str, Any] | None = None) -> dict[str, tuple[str, ...]]:
    if not alias_map:
        return {key: value for key, value in CANONICAL_FIELD_ALIASES.items()}
    normalized: dict[str, tuple[str, ...]] = {}
    for key, value in alias_map.items():
        if isinstance(value, str):
            normalized[key] = (value,)
        elif isinstance(value, Iterable):
            normalized[key] = tuple(normalize_text(item) for item in value if normalize_text(item))
        else:
            normalized[key] = CANONICAL_FIELD_ALIASES.get(key, (key,))
    for key, aliases in CANONICAL_FIELD_ALIASES.items():
        normalized.setdefault(key, aliases)
    return normalized


def apply_alias_map(record: Mapping[str, Any], alias_map: Mapping[str, Any] | None = None) -> dict[str, Any]:
    mapping = normalize_alias_map(alias_map)
    out = dict(record)
    for canonical, aliases in mapping.items():
        value = first_nonempty(record, aliases)
        if value is not None and normalize_text(value) != "":
            out[canonical] = value
    if not normalize_text(out.get("record_id")):
        prefix = normalize_text(out.get("table_name")) or "CSMAR"
        row_index = out.get("row_index")
        out["record_id"] = f"{prefix}:{row_index if row_index is not None else 0}"
    return out


def build_mapping_template() -> dict[str, Any]:
    return {
        "expected_tables": EXPECTED_TABLES,
        "canonical_fields": {key: list(value) for key, value in CANONICAL_FIELD_ALIASES.items()},
        "example_unit_policy": {
            "money_fields": "prefer explicit yuan/wan-yuan/yi-yuan metadata; otherwise require an adapter-side unit map",
            "ratio_fields": "store as percentage or decimal consistently and document the convention in the mapping file",
        },
    }


def missing_expected_files(base_dir: str | Path, expected_files: Iterable[str]) -> list[str]:
    root = Path(base_dir)
    missing: list[str] = []
    for name in expected_files:
        if not (root / name).exists():
            missing.append(name)
    return missing


def inventory_sources(base_dir: str | Path) -> list[str]:
    root = Path(base_dir)
    return [str(path.relative_to(root)) for path in discover_source_files(root)]


CSMAR_FINANCIAL_V2_VERSION = "csmar_financial_cache_v2"

CSMAR_FINANCIAL_SOURCE_SHA256 = {
    "FI_T1.xlsx": "cbd82da0052b17764d3c256e6207ea9e7510b00fc5592db451aedfdc1c835e52",
    "FI_T5.xlsx": "af9dd9681e07f5bace22e34d65a32812c6ca9c79190dbd120fad597e0e3ea180",
    "FI_T6.xlsx": "59eb251564507a29cdce6cba67e26ad5584d650efd037fd7827cca8791c446b1",
    "FI_T8.xlsx": "00aab96eb7282fa7bbb148e597511b028b39a5db8f444a67ae9c8dc662c749e7",
    "FI_T10.xlsx": "a6e589a41a7ef89b6d9e72c47f57c477b9ae51bef13ec441315895ceb5443d7d",
}

BASIC_CSMAR_COLUMNS = ["Stkcd", "ShortName", "Accper", "Typrep", "Source", "Indcd", "Indnme", "Indcd1", "Indnme1"]

CSMAR_FINANCIAL_FIELD_SPECS: dict[str, dict[str, tuple[str, str]]] = {
    "FI_T1": {
        "current_ratio": ("F010101A", "流动比率"),
        "quick_ratio": ("F010201A", "速动比率"),
        "debt_ratio": ("F011201A", "资产负债率"),
    },
    "FI_T5": {
        "roa": ("F050204C", "总资产净利润率(ROA)TTM"),
        "roe": ("F050504C", "净资产收益率(ROE)TTM"),
        "parent_roe": ("F053004C", "归属于母公司净资产收益率(ROE)TTM"),
        "gross_margin": ("F053301B", "营业毛利率"),
        "gross_margin_ttm": ("F053301C", "营业毛利率TTM"),
        "rd_expense_ratio": ("F053401B", "研发费用率"),
    },
    "FI_T6": {
        "net_profit_cash_content": ("F060101B", "净利润现金净含量"),
        "company_cash_flow_1": ("F061301B", "公司现金流1"),
        "company_cash_flow_2": ("F061302B", "公司现金流2"),
        "free_cash_flow": ("F061501B", "自由现金流"),
    },
    "FI_T8": {
        "revenue_yoy": ("F081603B", "营业收入增长率C"),
        "revenue_yoy_b": ("F081602C", "营业收入增长率B"),
        "revenue_qoq": ("F081601B", "营业收入增长率A"),
        "ocf_yoy": ("F082302B", "经营活动产生的净流量增长率B"),
        "ocf_qoq": ("F082301B", "经营活动产生的净流量增长率A"),
    },
    "FI_T10": {
        "pe_ratio": ("F100103C", "市盈率(PE)TTM"),
        "pe_ratio_1": ("F100101B", "市盈率(PE)1"),
        "pe_ratio_2": ("F100102B", "市盈率(PE)2"),
        "parent_pe_ratio": ("F100603C", "市盈率母公司(PE)TTM"),
    },
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_stock_code(value: Any) -> str:
    text = normalize_text(value)
    match = re.search(r"(\d{6})", text)
    if match:
        return match.group(1)
    digits = re.sub(r"\D+", "", text)
    if not digits:
        return ""
    return digits.zfill(6)[-6:]


def parse_accper(value: Any) -> date | None:
    text = normalize_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10] if fmt != "%Y%m%d" else text[:8], fmt).date()
        except Exception:
            continue
    match = re.search(r"(20\d{2})[-年/](\d{1,2})[-月/](\d{1,2})", text)
    if match:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    return None


def estimated_financial_available_date(accper: Any) -> str:
    """Conservative public availability estimate when no announcement date exists."""

    text = normalize_text(accper)
    if not text:
        return ""
    end = parse_accper(text)
    if end is None:
        match = re.fullmatch(r"(20\d{2})(Q1|H1|Q3|Q4|FY)", text)
        if match:
            year = int(match.group(1))
            period = match.group(2)
            if period == "Q1":
                return date(year, 4, 30).isoformat()
            if period == "H1":
                return date(year, 8, 31).isoformat()
            if period == "Q3":
                return date(year, 10, 31).isoformat()
            return date(year + 1, 4, 30).isoformat()
        return ""
    if end.month == 12 and end.day == 31:
        return date(end.year + 1, 4, 30).isoformat()
    if end.month == 3 and end.day == 31:
        return date(end.year, 4, 30).isoformat()
    if end.month == 6 and end.day == 30:
        return date(end.year, 8, 31).isoformat()
    if end.month == 9 and end.day == 30:
        return date(end.year, 10, 31).isoformat()
    return (end + timedelta(days=120)).isoformat()


def _safe_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    try:
        parsed = float(value)
    except Exception:
        return None
    if not math.isfinite(parsed):
        return None
    return parsed


def _read_csmar_excel_subset(path: Path, table_name: str, needed_columns: list[str]) -> tuple[pd.DataFrame, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    before_dimension = sheet.calculate_dimension()
    reset_applied = False
    if hasattr(sheet, "reset_dimensions"):
        sheet.reset_dimensions()
        reset_applied = True
    after_dimension = sheet.calculate_dimension(force=True)
    header = [normalize_text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    field_names = [normalize_text(value) for value in next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))]
    units = [normalize_text(value) for value in next(sheet.iter_rows(min_row=3, max_row=3, values_only=True))]
    header_index = {name: idx for idx, name in enumerate(header) if name}
    selected = [column for column in needed_columns if column in header_index]
    missing = [column for column in needed_columns if column not in header_index]
    selected_indices = [header_index[column] for column in selected]
    rows: list[list[Any]] = []
    for raw in sheet.iter_rows(min_row=4, values_only=True):
        if not raw or all(value in (None, "") for value in raw):
            continue
        rows.append([raw[idx] if idx < len(raw) else None for idx in selected_indices])
    frame = pd.DataFrame(rows, columns=selected)
    if "Stkcd" in frame.columns:
        frame["Stkcd"] = frame["Stkcd"].map(normalize_stock_code)
    for column in ("ShortName", "Accper", "Typrep", "Indcd", "Indnme", "Indcd1", "Indnme1"):
        if column in frame.columns:
            frame[column] = frame[column].map(normalize_text)
    metadata = {
        "table": table_name,
        "sheet": sheet.title,
        "dimension_before_reset": before_dimension,
        "dimension_after_reset": after_dimension,
        "reset_dimensions_applied": reset_applied,
        "excel_dimension_recovered": bool(reset_applied and len(selected) > 1 and len(header) > 1),
        "header_row": 1,
        "description_row": 2,
        "unit_row": 3,
        "data_start_row": 4,
        "source_columns": len(header),
        "selected_columns": selected,
        "missing_requested_columns": missing,
        "field_descriptions": {column: field_names[header_index[column]] for column in selected if header_index[column] < len(field_names)},
        "field_units": {column: units[header_index[column]] for column in selected if header_index[column] < len(units)},
        "rows": int(len(frame)),
    }
    workbook.close()
    return frame, metadata


def _table_needed_columns(table_name: str) -> list[str]:
    fields = [field_code for field_code, _field_name in CSMAR_FINANCIAL_FIELD_SPECS[table_name].values()]
    basic = BASIC_CSMAR_COLUMNS if table_name != "FI_T10" else [column for column in BASIC_CSMAR_COLUMNS if column != "Typrep"]
    return list(dict.fromkeys(basic + fields))


def _normalize_table_frame(frame: pd.DataFrame, table_name: str, source_file: str, source_sha256: str) -> pd.DataFrame:
    specs = CSMAR_FINANCIAL_FIELD_SPECS[table_name]
    out = pd.DataFrame()
    out["stock_code"] = frame.get("Stkcd", pd.Series(dtype=str)).map(normalize_stock_code)
    out["company_name"] = frame.get("ShortName", pd.Series(dtype=str)).map(normalize_text)
    out["report_period"] = frame.get("Accper", pd.Series(dtype=str)).map(normalize_text)
    out["Accper"] = out["report_period"]
    if "Typrep" in frame.columns:
        out["Typrep"] = frame["Typrep"].map(normalize_text)
    else:
        out["Typrep"] = "not_applicable"
    out["Source"] = frame.get("Source", pd.Series(dtype=object))
    for column in ("Indcd", "Indnme", "Indcd1", "Indnme1"):
        out[column] = frame[column].map(normalize_text) if column in frame.columns else ""
    out["industry_code"] = out["Indcd1"].where(out["Indcd1"] != "", out["Indcd"])
    out["industry_name"] = out["Indnme1"].where(out["Indnme1"] != "", out["Indnme"])
    out["estimated_available_date"] = out["Accper"].map(estimated_financial_available_date)
    out["temporal_availability"] = "temporal_availability_uncertain"
    out["source_table"] = table_name
    out["source_file"] = source_file
    out["source_sha256"] = source_sha256
    for metric, (field_code, field_name) in specs.items():
        if field_code in frame.columns:
            out[metric] = frame[field_code].map(_safe_number)
        else:
            out[metric] = None
        out[f"{metric}_field_code"] = field_code
        out[f"{metric}_field_name"] = field_name
        out[f"{metric}_source_table"] = table_name
        out[f"{metric}_source_file"] = source_file
        out[f"{metric}_source_sha256"] = source_sha256
    out = out[out["stock_code"] != ""].copy()
    return out


def _merge_table_frames(frames: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    keyed_tables = [name for name in ("FI_T1", "FI_T5", "FI_T6", "FI_T8") if name in frames and not frames[name].empty]
    combined: pd.DataFrame | None = None
    keys = ["stock_code", "report_period", "Typrep"]
    base_columns = [
        "stock_code",
        "company_name",
        "report_period",
        "Accper",
        "Typrep",
        "Source",
        "Indcd",
        "Indnme",
        "Indcd1",
        "Indnme1",
        "industry_code",
        "industry_name",
        "estimated_available_date",
        "temporal_availability",
    ]
    for name in keyed_tables:
        frame = frames[name].copy()
        metric_columns = [column for column in frame.columns if column not in base_columns + ["source_table", "source_file", "source_sha256"]]
        keep = base_columns + metric_columns
        table_frame = frame[keep]
        if combined is None:
            combined = table_frame
        else:
            combined = combined.merge(table_frame, on=keys, how="outer", suffixes=("", f"_{name}"))
            for column in base_columns:
                if column in keys:
                    continue
                alt = f"{column}_{name}"
                if alt in combined.columns:
                    combined[column] = combined[column].where(combined[column].map(normalize_text) != "", combined[alt])
                    combined = combined.drop(columns=[alt])
    if combined is None:
        combined = pd.DataFrame(columns=base_columns)
    if "FI_T10" in frames and not frames["FI_T10"].empty:
        t10 = frames["FI_T10"].copy()
        t10 = t10.drop(columns=["Typrep"], errors="ignore")
        t10_metric_columns = [
            column
            for column in t10.columns
            if column not in [c for c in base_columns if c != "Typrep"] + ["source_table", "source_file", "source_sha256"]
        ]
        combined = combined.merge(
            t10[["stock_code", "report_period"] + t10_metric_columns],
            on=["stock_code", "report_period"],
            how="outer",
            suffixes=("", "_FI_T10"),
        )
    for column in ("company_name", "Accper", "Source", "Indcd", "Indnme", "Indcd1", "Indnme1", "industry_code", "industry_name", "estimated_available_date", "temporal_availability"):
        alt = f"{column}_FI_T10"
        if alt in combined.columns:
            if column in combined.columns:
                combined[column] = combined[column].where(combined[column].map(normalize_text) != "", combined[alt])
            else:
                combined[column] = combined[alt]
            combined = combined.drop(columns=[alt])
    if "Typrep" not in combined.columns:
        combined["Typrep"] = "not_applicable"
    combined["Typrep"] = combined["Typrep"].map(normalize_text).replace("", "not_applicable")
    combined["record_id"] = (
        "CSMAR_FINANCIAL_V2:"
        + combined["stock_code"].map(normalize_text)
        + ":"
        + combined["report_period"].map(normalize_text)
        + ":"
        + combined["Typrep"].map(normalize_text)
    )
    combined["table_name"] = "CSMAR_FINANCIAL_V2"
    combined["calculation_version"] = CSMAR_FINANCIAL_V2_VERSION
    combined = combined.sort_values(["stock_code", "report_period", "Typrep"]).reset_index(drop=True)
    return combined


def _add_industry_benchmarks(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    industry_keys = ["report_period", "Typrep", "industry_code"]
    if "revenue_yoy" in out.columns:
        valid = out["revenue_yoy"].map(_safe_number).notna() & (out["industry_code"].map(normalize_text) != "")
        grouped = out.loc[valid].groupby(industry_keys, dropna=False)["revenue_yoy"].agg(["mean", "count"]).reset_index()
        grouped = grouped.rename(columns={"mean": "industry_revenue_growth_mean", "count": "industry_revenue_growth_sample_count"})
        out = out.merge(grouped, on=industry_keys, how="left")
    if "pe_ratio" in out.columns:
        pe = out["pe_ratio"].map(_safe_number)
        valid = pe.notna() & (pe > 0) & (out["industry_code"].map(normalize_text) != "")
        grouped = out.loc[valid].groupby(["report_period", "industry_code"], dropna=False)["pe_ratio"].agg(["mean", "count"]).reset_index()
        grouped = grouped.rename(columns={"mean": "industry_mean_pe", "count": "industry_pe_sample_count"})
        out = out.merge(grouped, on=["report_period", "industry_code"], how="left")
    for column in ("industry_revenue_growth_sample_count", "industry_pe_sample_count"):
        if column in out.columns:
            out[column] = out[column].fillna(0).astype(int)
    return out


def build_csmar_financial_cache_v2(
    source_dir: str | Path,
    cache_dir: str | Path,
    force: bool = False,
) -> dict[str, Any]:
    source_root = Path(source_dir)
    cache_root = Path(cache_dir)
    cache_root.mkdir(parents=True, exist_ok=True)
    manifest_path = cache_root / "csmar_cache_manifest.json"
    records_path = cache_root / "financial_records.parquet"
    sources: dict[str, dict[str, Any]] = {}
    for filename, expected_sha in CSMAR_FINANCIAL_SOURCE_SHA256.items():
        path = source_root / filename
        actual_sha = file_sha256(path) if path.exists() else "missing"
        sources[filename] = {
            "path": str(path),
            "expected_sha256": expected_sha,
            "actual_sha256": actual_sha,
            "sha256_match": actual_sha == expected_sha,
            "exists": path.exists(),
        }
    if (
        not force
        and manifest_path.exists()
        and records_path.exists()
        and all(item.get("sha256_match") for item in sources.values())
    ):
        previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        previous_sources = previous.get("sources") or {}
        if all((previous_sources.get(name) or {}).get("actual_sha256") == item["actual_sha256"] for name, item in sources.items()):
            previous["cache_reused"] = True
            return previous
    if not all(item.get("exists") and item.get("sha256_match") for item in sources.values()):
        raise ValueError(f"CSMAR source SHA check failed: {sources}")

    table_frames: dict[str, pd.DataFrame] = {}
    table_metadata: dict[str, Any] = {}
    for filename in CSMAR_FINANCIAL_SOURCE_SHA256:
        table_name = Path(filename).stem
        frame, metadata = _read_csmar_excel_subset(source_root / filename, table_name, _table_needed_columns(table_name))
        normalized = _normalize_table_frame(frame, table_name, f"data/external/csmar/财务指标/{filename}", sources[filename]["actual_sha256"])
        table_frames[table_name] = normalized
        table_cache_path = cache_root / f"{table_name}.parquet"
        normalized.to_parquet(table_cache_path, index=False)
        metadata.update(
            {
                "normalized_rows": int(len(normalized)),
                "cache_path": str(table_cache_path),
                "cache_sha256": file_sha256(table_cache_path),
                "Typrep_distribution": normalized["Typrep"].value_counts(dropna=False).to_dict() if "Typrep" in normalized.columns else {},
            }
        )
        table_metadata[table_name] = metadata

    records = _add_industry_benchmarks(_merge_table_frames(table_frames))
    records.to_parquet(records_path, index=False)
    manifest = {
        "version": CSMAR_FINANCIAL_V2_VERSION,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_dir": str(source_root),
        "cache_dir": str(cache_root),
        "sources": sources,
        "source_sha_match": all(item.get("sha256_match") for item in sources.values()),
        "excel_dimension_recovered": all(meta.get("excel_dimension_recovered") for meta in table_metadata.values()),
        "stock_code_leading_zero_preserved": bool(records["stock_code"].astype(str).str.match(r"^\d{6}$").all()) if not records.empty else False,
        "field_mapping": CSMAR_FINANCIAL_FIELD_SPECS,
        "typrep_definition": {
            "FI_T1_FI_T5_FI_T6_FI_T8": "A：合并报表；B：母公司报表；公司未公布合并报表时以单一报表数据添列。",
            "FI_T10": "source table has no Typrep column; valuation metrics are merged by stock_code+Accper and marked not_applicable in table cache.",
            "preferred_policy": "prefer Typrep=A for company financial checks; prior-period comparison keeps same Typrep.",
        },
        "temporal_policy": {
            "announcement_date_available": False,
            "policy": "Conservative estimated availability from Accper: annual Apr-30 next year, Q1 Apr-30, H1 Aug-31, Q3 Oct-31; all such uses marked temporal_availability_uncertain.",
        },
        "missing_raw_fields": {
            "ocf_amount": "not_available in FI_T6; FI_T6 contains cash-content/corporate/free-cash-flow derived indicators, not raw operating cash-flow net amount.",
            "net_profit": "not_available in the five fixed FI tables",
            "accounts_receivable": "not_available in the five fixed FI tables",
            "inventory": "not_available in the five fixed FI tables",
            "non_recurring_profit": "not_available in the five fixed FI tables",
        },
        "table_metadata": table_metadata,
        "outputs": {
            "financial_records": str(records_path),
            "financial_records_sha256": file_sha256(records_path),
            "financial_records_rows": int(len(records)),
            "financial_records_columns": list(records.columns),
        },
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return manifest


def load_csmar_financial_records_v2(cache_dir: str | Path) -> list[dict[str, Any]]:
    records_path = Path(cache_dir) / "financial_records.parquet"
    if not records_path.exists():
        return []
    frame = pd.read_parquet(records_path)
    return frame.where(pd.notna(frame), None).to_dict("records")
