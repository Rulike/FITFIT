#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build targeted Task1 event_emotion and duplication_id annotation packages."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import pickle
import random
import re
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs/task1_step5_8g_targeted_annotation"
EMOTION_XLSX = ROOT / "reports/task1_step5_8g_emotion_human_review.xlsx"
DUP_XLSX = ROOT / "reports/task1_step5_8g_duplication_human_review.xlsx"
GUIDE_MD = ROOT / "reports/task1_step5_8g_annotation_guide.md"
MANIFEST_JSON = ROOT / "reports/task1_step5_8g_annotation_manifest.json"
RAW_TRAIN = ROOT / "data/raw/news_with_tags_train.xlsx"
SHADOW_V4 = ROOT / "outputs/task1_step5_8f_final_candidate/task1_official_delivery_candidate_v4.jsonl"
EVENT_MODEL = ROOT / "outputs/task1_step5_8e_final_optimization/frozen_best_candidate/event_name_model.pkl"
EVENT_MODEL_EXPECTED_SHA = "f37ba4902784299ac17976f0fe059403d02befc15b18caa6d54ac4c5793b7c6e"
EMOTION_MODEL = ROOT / "outputs/task1_step5_8f_final_candidate/frozen_emotion_duplication/event_emotion_model.pkl"
FROZEN60 = ROOT / "outputs/task1_step5_7j_structure_confirmation/independent_frozen_test_confirmed.jsonl"
GOLD_FREE_1029_INPUTS = [
    ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl",
    ROOT / "outputs/task1_gold_shadow_smoke_v1/blind_input.jsonl",
]
FORMAL_PROTECTED = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "reports/TASK1_TASK2_FINAL_ACCEPTANCE.md",
    ROOT / "reports/task1_task2_final_manifest.json",
]

NEG_WORDS = ["风险", "处罚", "亏损", "下降", "下滑", "违约", "终止", "冻结", "减持", "损失", "事故", "破产", "撤回"]
NEU_WORDS = ["计划", "拟", "可能", "预计", "问询", "传闻", "尚未", "不确定", "或将", "投资者"]
POS_WORDS = ["突破", "获批", "中标", "增长", "扩产", "签约", "合作", "投产", "提升", "盈利", "增持", "回购"]
EMOTIONS = ["正面", "中性", "负面"]


def import_step5_8e_module():
    spec = importlib.util.spec_from_file_location("step5_8e", ROOT / "scripts/run_task1_step5_8e_final_optimization.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("cannot import Step 5.8-E module")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def import_step5_8f_module():
    spec = importlib.util.spec_from_file_location("step5_8f", ROOT / "scripts/run_task1_step5_8f_finalization.py")
    if spec is None or spec.loader is None:
        raise RuntimeError("cannot import Step 5.8-F module")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


E = import_step5_8e_module()
F = import_step5_8f_module()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def setup() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = ROOT.parent / "backups" / f"task1_step5_8g_targeted_annotation_{ts}"
    backup.mkdir(parents=True, exist_ok=True)
    write_json(OUT_DIR / "run_state.json", {"stage": "initialized", "backup_dir": str(backup), "updated_at": now_iso(), "llm_calls": 0})
    return backup


def protected_sha(path: Path) -> dict[str, str]:
    paths = [p for p in FORMAL_PROTECTED if p.exists()] + [SHADOW_V4, EVENT_MODEL, EMOTION_MODEL]
    if FROZEN60.exists():
        paths.append(FROZEN60)
    out = {str(p): sha256_file(p) for p in sorted(set(paths)) if p.exists()}
    path.write_text("\n".join(f"{v}  {k}" for k, v in sorted(out.items())) + "\n", encoding="utf-8")
    return out


def official_train_ids() -> set[str]:
    ids: set[str] = set()
    for sheet in ["个股新闻", "行业新闻"]:
        df = pd.read_excel(RAW_TRAIN, sheet_name=sheet, usecols=["article_file_id"])
        ids.update(str(x) for x in df["article_file_id"].dropna())
    return ids


def history_dev_holdout_ids() -> set[str]:
    rows = E.load_official_rows()
    split_ids, _ = E.split_official_rows(rows)
    by_id = {r["row_id"]: r["article_file_id"] for r in rows}
    ids = set()
    for split in ["model_selection_dev", "untouched_internal_holdout"]:
        ids.update(by_id[row_id] for row_id in split_ids[split])
    return ids


def frozen60_ids() -> set[str]:
    ids = set()
    for row in read_jsonl(FROZEN60):
        value = row.get("article_id") or row.get("article_file_id")
        if value:
            ids.add(str(value))
    return ids


def protected_1029_ids() -> set[str]:
    ids = set()
    for path in GOLD_FREE_1029_INPUTS:
        for row in read_jsonl(path):
            value = row.get("article_id") or row.get("article_file_id")
            if value:
                ids.add(str(value))
    return ids


def target_name(row: dict[str, Any]) -> str:
    return str(row.get("secu_abbr") or row.get("industry_name") or "")


def norm_text(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def title_sim(a: str, b: str) -> float:
    def grams(x: str) -> set[str]:
        x = norm_text(x)
        return {x[i : i + 2] for i in range(max(0, len(x) - 1))}
    ga, gb = grams(a), grams(b)
    return len(ga & gb) / len(ga | gb) if ga and gb else 0.0


def token_overlap(a: str, b: str) -> float:
    pattern = r"[\u4e00-\u9fffA-Za-z0-9]{2,}"
    sa = set(re.findall(pattern, a))
    sb = set(re.findall(pattern, b))
    return len(sa & sb) / len(sa | sb) if sa and sb else 0.0


def date_key(row: dict[str, Any]) -> str:
    return str(row.get("article_publish_time") or "")[:10]


def load_current_rows(exclude_ids: set[str]) -> tuple[list[dict[str, Any]], dict[str, int]]:
    rows = []
    excluded = Counter()
    for idx, row in enumerate(read_jsonl(SHADOW_V4)):
        aid = str(row.get("article_file_id") or "")
        if not aid:
            excluded["missing_article_id"] += 1
            continue
        if aid in exclude_ids:
            excluded["protected_or_seen_article_id"] += 1
            continue
        r = dict(row)
        r["candidate_index"] = idx
        r["target_type"] = "company" if str(r.get("trading_code") or "") else "industry"
        r["target_name"] = target_name(r)
        r["content_sha256"] = E.sha256_text(str(r.get("content") or ""))
        rows.append(r)
    return rows, dict(excluded)


def load_emotion_model():
    with EMOTION_MODEL.open("rb") as handle:
        return pickle.load(handle)


def predict_emotion_with_conf(model_obj: Any, train_rows: list[dict[str, Any]], rows: list[dict[str, Any]]) -> tuple[list[str], list[float]]:
    # Reconstruct the fields expected by Step 5.8-F emotion features.
    for r in rows:
        r["pred_event_name"] = r.get("event_name")
        r.setdefault("row_id", str(r.get("article_file_id")))
        r.setdefault("target_type", "company" if str(r.get("trading_code") or "") else "industry")
    preds = F.predict_best_emotion(model_obj, train_rows, rows)
    confs = []
    if isinstance(model_obj, dict) and model_obj.get("kind") in {"model", "fusion"}:
        model = model_obj["model"]
        try:
            probs = model.predict_proba([F.emotion_text(r) for r in rows])
            labels = list(model.named_steps["clf"].classes_)
            for pred, prob in zip(preds, probs):
                confs.append(float(prob[labels.index(pred)]) if pred in labels else float(np.max(prob)))
        except Exception:
            confs = [0.0 for _ in rows]
    else:
        confs = [0.0 for _ in rows]
    return preds, confs


def emotion_prior_map() -> dict[str, str]:
    priors: dict[str, Counter[str]] = defaultdict(Counter)
    for sheet in ["个股新闻", "行业新闻"]:
        df = pd.read_excel(RAW_TRAIN, sheet_name=sheet).fillna("")
        for _, r in df.iterrows():
            priors[str(r["event_name"])][str(r["event_emotion"])] += 1
    return {k: v.most_common(1)[0][0] for k, v in priors.items()}


def lex_counts(row: dict[str, Any]) -> dict[str, int]:
    text = str(row.get("article_title") or "") + "\n" + str(row.get("content") or "")
    return {
        "negative": sum(w in text for w in NEG_WORDS),
        "neutral": sum(w in text for w in NEU_WORDS),
        "positive": sum(w in text for w in POS_WORDS),
    }


def emotion_stratum(row: dict[str, Any]) -> str:
    counts = lex_counts(row)
    if row["model_disagreement"] or row["current_confidence"] < 0.45:
        return "low_confidence_or_disagreement"
    if row["current_emotion_prediction"] == "负面" or counts["negative"] >= 2:
        return "negative_candidate"
    if row["current_emotion_prediction"] == "中性" or counts["neutral"] >= 2:
        return "neutral_candidate"
    return "positive_candidate"


def build_emotion_candidates(rows: list[dict[str, Any]], train_rows: list[dict[str, Any]], emotion_model: Any) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    preds, confs = predict_emotion_with_conf(emotion_model, train_rows, rows)
    priors = emotion_prior_map()
    prepared = []
    for i, (row, pred, conf) in enumerate(zip(rows, preds, confs), 1):
        prior = priors.get(str(row.get("event_name") or ""), "正面")
        quote = E.evidence_quote(row)
        counts = lex_counts(row)
        out = {
            "candidate_record_id": f"EMO_5_8G_{i:05d}",
            "article_file_id": row["article_file_id"],
            "article_title": row["article_title"],
            "content": row["content"],
            "target_name": row["target_name"],
            "industry_name": row.get("industry_name", ""),
            "frozen_event_name": row.get("event_name", ""),
            "current_emotion_prediction": pred,
            "current_confidence": round(float(conf), 6),
            "prior_emotion_prediction": prior,
            "model_disagreement": pred != prior,
            "ai_suggested_emotion": pred if pred in EMOTIONS else "",
            "ai_reason": f"基于证据句和情感词信号：负面{counts['negative']}、中性{counts['neutral']}、正面{counts['positive']}；建议仅供人工确认。",
            "evidence_quote": quote,
            "human_event_emotion": "",
            "reviewer_confirmed": "",
            "reviewer_note": "",
            "annotation_source": "ai_assisted_unconfirmed",
            "sampling_stratum": "",
        }
        out["sampling_stratum"] = emotion_stratum(out)
        prepared.append(out)

    rng = random.Random(2026)
    for bucket in ["negative_candidate", "neutral_candidate", "low_confidence_or_disagreement", "positive_candidate"]:
        rng.shuffle(prepared)
    selected: list[dict[str, Any]] = []
    used: set[str] = set()

    def take(stratum: str, n: int) -> None:
        for row in prepared:
            if len([x for x in selected if x["sampling_stratum"] == stratum]) >= n:
                break
            if row["sampling_stratum"] == stratum and row["article_file_id"] not in used:
                selected.append(row)
                used.add(row["article_file_id"])

    take("negative_candidate", 150)
    take("neutral_candidate", 150)
    take("low_confidence_or_disagreement", 100)
    take("positive_candidate", 100)
    for row in prepared:
        if len(selected) >= 500:
            break
        if row["article_file_id"] not in used:
            selected.append(row)
            used.add(row["article_file_id"])
    selected = selected[:500]
    summary = {
        "available_rows": len(rows),
        "selected_rows": len(selected),
        "stratum_counts": dict(Counter(r["sampling_stratum"] for r in selected)),
        "current_prediction_counts": dict(Counter(r["current_emotion_prediction"] for r in selected)),
        "model_disagreement_count": sum(1 for r in selected if r["model_disagreement"]),
        "low_confidence_count": sum(1 for r in selected if r["current_confidence"] < 0.45),
        "evidence_quote_not_in_content": sum(1 for r in selected if r["evidence_quote"] and r["evidence_quote"] not in r["content"]),
    }
    return selected, summary


def build_duplication_pairs(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    by_target_event: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    by_target: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        t = str(r.get("trading_code") or r.get("industry_code") or "")
        by_target_event[(t, str(r.get("event_name") or ""))].append(r)
        by_target[t].append(r)

    candidates: dict[str, list[tuple[float, dict[str, Any], dict[str, Any]]]] = defaultdict(list)
    for group in by_target_event.values():
        if len(group) < 2:
            continue
        for i in range(len(group)):
            for j in range(i + 1, len(group)):
                a, b = group[i], group[j]
                sim = title_sim(a.get("article_title", ""), b.get("article_title", ""))
                overlap = token_overlap(str(a.get("content") or "")[:1200], str(b.get("content") or "")[:1200])
                same_day = date_key(a) == date_key(b)
                score = 0.7 * sim + 0.3 * overlap + (0.1 if same_day else 0)
                if score >= 0.72:
                    bucket = "high_probability_duplicate"
                elif score >= 0.42:
                    bucket = "medium_similarity_candidate"
                else:
                    bucket = "hard_negative_candidate"
                candidates[bucket].append((score, a, b))
    for group in by_target.values():
        if len(group) < 2:
            continue
        for i in range(min(len(group), 80)):
            for j in range(i + 1, min(len(group), 80)):
                a, b = group[i], group[j]
                if a.get("event_name") == b.get("event_name"):
                    continue
                sim = title_sim(a.get("article_title", ""), b.get("article_title", ""))
                if sim >= 0.28 or date_key(a) == date_key(b):
                    candidates["hard_negative_candidate"].append((sim, a, b))
    rng = random.Random(2026)
    for bucket in candidates:
        candidates[bucket].sort(key=lambda x: x[0], reverse=True)
    all_rows = rows[:]
    rng.shuffle(all_rows)
    random_pairs = []
    for i in range(len(all_rows)):
        for j in range(i + 1, min(len(all_rows), i + 80)):
            a, b = all_rows[i], all_rows[j]
            if (a.get("trading_code") or a.get("industry_code")) != (b.get("trading_code") or b.get("industry_code")):
                random_pairs.append((0.0, a, b))
            if len(random_pairs) >= 200:
                break
        if len(random_pairs) >= 200:
            break
    candidates["random_negative"].extend(random_pairs)

    targets = {
        "high_probability_duplicate": 250,
        "medium_similarity_candidate": 150,
        "hard_negative_candidate": 150,
        "random_negative": 50,
    }
    selected = []
    used_pairs: set[tuple[str, str]] = set()
    for bucket, n in targets.items():
        for score, a, b in candidates[bucket]:
            if len([x for x in selected if x["sampling_stratum"] == bucket]) >= n:
                break
            left, right = sorted([str(a["article_file_id"]), str(b["article_file_id"])])
            if left == right or (left, right) in used_pairs:
                continue
            used_pairs.add((left, right))
            overlap = token_overlap(str(a.get("content") or "")[:1500], str(b.get("content") or "")[:1500])
            selected.append(
                {
                    "pair_id": f"DUP_5_8G_{len(selected)+1:05d}",
                    "left_article_file_id": left,
                    "right_article_file_id": right,
                    "left_title": a["article_title"] if str(a["article_file_id"]) == left else b["article_title"],
                    "right_title": b["article_title"] if str(b["article_file_id"]) == right else a["article_title"],
                    "left_content_or_evidence": (a["content"] if str(a["article_file_id"]) == left else b["content"])[:900],
                    "right_content_or_evidence": (b["content"] if str(b["article_file_id"]) == right else a["content"])[:900],
                    "left_target": a["target_name"] if str(a["article_file_id"]) == left else b["target_name"],
                    "right_target": b["target_name"] if str(b["article_file_id"]) == right else a["target_name"],
                    "left_event_name": a["event_name"] if str(a["article_file_id"]) == left else b["event_name"],
                    "right_event_name": b["event_name"] if str(b["article_file_id"]) == right else a["event_name"],
                    "left_publish_time": a.get("article_publish_time", "") if str(a["article_file_id"]) == left else b.get("article_publish_time", ""),
                    "right_publish_time": b.get("article_publish_time", "") if str(b["article_file_id"]) == right else a.get("article_publish_time", ""),
                    "text_similarity": round(float(score), 6),
                    "time_difference": "same_day" if date_key(a) == date_key(b) else "different_day",
                    "entity_overlap": round(float(overlap), 6),
                    "ai_suggested_same_event": "yes" if bucket == "high_probability_duplicate" else "no" if bucket in {"hard_negative_candidate", "random_negative"} else "unclear",
                    "ai_reason": "基于target、event_name、发布时间、标题/正文相似度和实体重合度的规则建议，仅供人工确认。",
                    "shared_fact_evidence": "",
                    "human_same_event": "",
                    "human_relation": "",
                    "human_cluster_id": "",
                    "reviewer_confirmed": "",
                    "reviewer_note": "",
                    "annotation_source": "ai_assisted_unconfirmed",
                    "sampling_stratum": bucket,
                }
            )
    summary = {
        "available_rows": len(rows),
        "selected_pairs": len(selected),
        "stratum_counts": dict(Counter(r["sampling_stratum"] for r in selected)),
        "unique_unordered_pairs": len({tuple(sorted([r["left_article_file_id"], r["right_article_file_id"]])) for r in selected}),
        "same_target_pairs": sum(1 for r in selected if r["left_target"] == r["right_target"]),
        "same_event_name_pairs": sum(1 for r in selected if r["left_event_name"] == r["right_event_name"]),
        "same_day_pairs": sum(1 for r in selected if r["time_difference"] == "same_day"),
    }
    return selected, summary


def write_workbooks(emotion_rows: list[dict[str, Any]], dup_rows: list[dict[str, Any]]) -> None:
    blue = PatternFill("solid", fgColor="D9EAF7")
    orange = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(bold=True)
    for path, rows, sheet in [(EMOTION_XLSX, emotion_rows, "emotion_review"), (DUP_XLSX, dup_rows, "duplication_review")]:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame([{"说明": "AI辅助审核模板，不是人工金标；reviewer_confirmed=yes 后才可进入后续训练或评估。"}]).to_excel(writer, sheet_name="使用说明", index=False)
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet, index=False)
        wb = load_workbook(path)
        ws = wb[sheet]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = header_font
            if str(cell.value).startswith("ai_") or str(cell.value).startswith("current_") or str(cell.value).startswith("prior_"):
                cell.fill = blue
            if str(cell.value).startswith("human_") or str(cell.value).startswith("reviewer_"):
                cell.fill = orange
        wb.save(path)


def validate(emotion_rows: list[dict[str, Any]], dup_rows: list[dict[str, Any]], excluded_ids: set[str]) -> dict[str, Any]:
    emotion_ids = [r["article_file_id"] for r in emotion_rows]
    pair_ids = [r["pair_id"] for r in dup_rows]
    unordered = [tuple(sorted([r["left_article_file_id"], r["right_article_file_id"]])) for r in dup_rows]
    return {
        "emotion": {
            "count": len(emotion_rows),
            "unique_article_ids": len(set(emotion_ids)),
            "duplicate_article_ids": len(emotion_ids) - len(set(emotion_ids)),
            "stratum_counts": dict(Counter(r["sampling_stratum"] for r in emotion_rows)),
            "human_fields_empty": all(not r["human_event_emotion"] and not r["reviewer_confirmed"] and not r["reviewer_note"] for r in emotion_rows),
            "legal_ai_suggestion": all((not r["ai_suggested_emotion"]) or r["ai_suggested_emotion"] in EMOTIONS for r in emotion_rows),
            "evidence_quote_valid": all((not r["evidence_quote"]) or r["evidence_quote"] in r["content"] for r in emotion_rows),
            "protected_intersection": len(set(emotion_ids) & excluded_ids),
        },
        "duplication": {
            "count": len(dup_rows),
            "unique_pair_ids": len(set(pair_ids)),
            "duplicate_pair_ids": len(pair_ids) - len(set(pair_ids)),
            "duplicate_unordered_pairs": len(unordered) - len(set(unordered)),
            "self_pairs": sum(1 for r in dup_rows if r["left_article_file_id"] == r["right_article_file_id"]),
            "stratum_counts": dict(Counter(r["sampling_stratum"] for r in dup_rows)),
            "human_fields_empty": all(not r["human_same_event"] and not r["human_relation"] and not r["human_cluster_id"] and not r["reviewer_confirmed"] and not r["reviewer_note"] for r in dup_rows),
            "protected_intersection": len({x for pair in unordered for x in pair} & excluded_ids),
        },
    }


def main() -> None:
    start = time.time()
    backup = setup()
    before = protected_sha(backup / "protected_sha256_before.txt")
    event_sha = sha256_file(EVENT_MODEL)
    if event_sha != EVENT_MODEL_EXPECTED_SHA:
        raise RuntimeError(f"event_name model SHA mismatch: {event_sha}")

    official_ids = official_train_ids()
    history_ids = history_dev_holdout_ids()
    frozen_ids = frozen60_ids()
    gold_ids = protected_1029_ids()
    excluded_ids = official_ids | history_ids | frozen_ids | gold_ids
    rows, excluded_report = load_current_rows(excluded_ids)
    if len(rows) < 500:
        status = "blocked_by_candidate_data_insufficiency"
    else:
        status = "task1_targeted_annotation_packages_ready"

    train_rows = E.load_official_rows()
    split_ids, _ = E.split_official_rows(train_rows)
    train_split = E.rows_by_split(train_rows, split_ids)["train"]
    # Add predicted event_name field for emotion prior compatibility.
    for r in train_split:
        r["pred_event_name"] = r["event_name"]
    emotion_model = load_emotion_model()
    emotion_rows, emotion_summary = build_emotion_candidates(rows, train_split, emotion_model) if status == "task1_targeted_annotation_packages_ready" else ([], {})
    dup_rows, dup_summary = build_duplication_pairs(rows) if status == "task1_targeted_annotation_packages_ready" else ([], {})
    if status == "task1_targeted_annotation_packages_ready" and (len(emotion_rows) < 450 or len(dup_rows) < 450):
        status = "blocked_by_candidate_data_insufficiency"
    validation = validate(emotion_rows, dup_rows, excluded_ids)
    if status == "task1_targeted_annotation_packages_ready":
        if validation["emotion"]["protected_intersection"] or validation["duplication"]["protected_intersection"]:
            status = "blocked_by_integrity_failure"
        write_jsonl(OUT_DIR / "emotion_candidates.jsonl", emotion_rows)
        write_jsonl(OUT_DIR / "duplication_candidates.jsonl", dup_rows)
        write_workbooks(emotion_rows, dup_rows)
    after = protected_sha(backup / "protected_sha256_after.txt")
    exclusion = {
        "official_train_ids": len(official_ids),
        "history_dev_holdout_ids": len(history_ids),
        "frozen60_ids": len(frozen_ids),
        "protected_1029_gold_free_input_ids": len(gold_ids),
        "total_excluded_ids": len(excluded_ids),
        "excluded_from_shadow_v4": excluded_report,
        "remaining_candidate_rows": len(rows),
    }
    write_json(OUT_DIR / "protected_exclusion_report.json", exclusion)
    manifest = {
        "status": status,
        "generated_at": now_iso(),
        "event_name_model_sha256": event_sha,
        "event_name_model_sha_matches_expected": event_sha == EVENT_MODEL_EXPECTED_SHA,
        "source_shadow_v4": str(SHADOW_V4),
        "source_rows_after_exclusion": len(rows),
        "emotion_summary": emotion_summary,
        "duplication_summary": dup_summary,
        "validation": validation,
        "annotation_source": "ai_assisted_unconfirmed",
        "human_confirmation_required": True,
        "llm_calls": 0,
        "tokens": 0,
        "latency_seconds": 0,
        "retry_count": 0,
        "formal_sha_match": before == after,
        "backup_dir": str(backup),
        "next_step": "Step 5.8-H may train only after reviewer_confirmed=yes labels are supplied.",
        "wall_time_seconds": round(time.time() - start, 3),
    }
    write_json(MANIFEST_JSON, manifest)
    GUIDE_MD.write_text(
        "# Task1 Step 5.8-G Annotation Guide\n\n"
        "本包是 AI辅助、未确认审核模板，不是人工金标。只有 `reviewer_confirmed=yes` 的记录才允许进入后续 Step 5.8-H 训练或评估。\n\n"
        "Emotion 合法值：正面、中性、负面。请依据正文证据句判断，不得依据股价行情直接推断。\n\n"
        "Duplication 合法关系：same_event、related_but_distinct、unrelated、unclear。不得仅因公司和 event_name 相同判定为同一事件。\n",
        encoding="utf-8",
    )
    # SHA sums.
    files = [OUT_DIR / "emotion_candidates.jsonl", OUT_DIR / "duplication_candidates.jsonl", OUT_DIR / "protected_exclusion_report.json", EMOTION_XLSX, DUP_XLSX, GUIDE_MD, MANIFEST_JSON, ROOT / "scripts/build_task1_step5_8g_targeted_annotation.py"]
    (OUT_DIR / "SHA256SUMS.txt").write_text("\n".join(f"{sha256_file(p)}  {p}" for p in files if p.exists()) + "\n", encoding="utf-8")
    print(json.dumps({"status": status, "emotion": emotion_summary, "duplication": dup_summary, "validation": validation, "formal_sha_match": before == after, "backup_dir": str(backup)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
