#!/usr/bin/env python3
"""Build the Task2 Step 6.1 independent-500 blind human audit package."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import re
import shutil
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.event_credibility import dea as dea_mod  # noqa: E402


SEED = 2026
PRIMARY_COUNT = 400
STRESS_COUNT = 100
TOTAL_COUNT = PRIMARY_COUNT + STRESS_COUNT
OUTPUT_DIR = ROOT / "outputs/event_credibility/task2_independent500_v1"
DEFAULT_WORKBOOK = ROOT / "reports/task2_independent500_human_audit.xlsx"
DEFAULT_GUIDE = ROOT / "reports/task2_independent500_human_audit_guide.md"
DEFAULT_MANIFEST = ROOT / "reports/task2_independent500_manifest.json"
BASE_SCORES = ROOT / "outputs/event_credibility/event_credibility_scores.parquet"
FORMAL_FULLTEXT_SCORES = ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet"
FULLTEXT_EVIDENCE = ROOT / "outputs/event_credibility/fulltext_evidence.parquet"
ARTICLE_INDEX = ROOT / "data/processed/event_credibility/fulltext_article_index.parquet"
BACKUP_ROOT = ROOT.parent / "backups"

PROTECTED_FORMAL_OUTPUTS = [
    ROOT / "outputs/event_credibility/final_cached_dea_v1_260/final_results.jsonl",
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
]

FROZEN_FILES = [
    ROOT / "src/event_credibility/fulltext.py",
    ROOT / "src/event_credibility/dea.py",
    ROOT / "src/event_credibility/retrieval_index.py",
    ROOT / "src/event_credibility/dea_safety_gate_v1_1.py",
    ROOT / "src/event_credibility/dea_safety_gate_v1_2.py",
    ROOT / "src/event_credibility/claim_contract.py",
    ROOT / "src/event_credibility/isr.py",
    ROOT / "configs/task2_final_cached_dea_v1.yaml",
    ROOT / "configs/task2_dea_safety_gate_v1_1.yaml",
    ROOT / "configs/task2_dea_safety_gate_v1_2.yaml",
    ROOT / "configs/task2_claim_contract_adapter_v1.yaml",
    ROOT / "configs/task2_isr_diagnostic_v1.yaml",
]

FORBIDDEN_MAIN_HEADER_PATTERNS = (
    "baseline",
    "round1",
    "cached",
    "safety_gate",
    "credibility_score",
    "credibility_level",
    "confidence",
    "rule_",
    "llm_",
    "final_result",
    "final_verdict",
    "production",
)

HUMAN_FIELDS = [
    "human_verdict",
    "human_confidence",
    "human_evidence_ids",
    "claim_clarity",
    "subject_match",
    "time_match",
    "evidence_sufficiency",
    "evidence_relevance",
    "source_independence",
    "official_source_present",
    "investor_question_only",
    "market_signal_only",
    "contradiction_present",
    "primary_error_type",
    "reviewer_notes",
    "reviewer_confirmed",
]

ALLOWED_VALUES = {
    "human_verdict": ["support", "contradict", "insufficient", "unclear"],
    "human_confidence": ["high", "medium", "low"],
    "claim_clarity": ["yes", "no", "unclear"],
    "subject_match": ["yes", "no", "unclear"],
    "time_match": ["yes", "no", "unclear"],
    "evidence_sufficiency": ["yes", "no", "unclear"],
    "evidence_relevance": ["yes", "no", "unclear"],
    "source_independence": ["yes", "no", "unclear"],
    "official_source_present": ["yes", "no", "unclear"],
    "investor_question_only": ["yes", "no", "unclear"],
    "market_signal_only": ["yes", "no", "unclear"],
    "contradiction_present": ["yes", "no", "unclear"],
    "primary_error_type": [
        "none",
        "claim_unclear",
        "subject_mismatch",
        "time_mismatch",
        "evidence_insufficient",
        "source_not_independent",
        "investor_question_only",
        "market_signal_only",
        "contradiction",
        "other",
    ],
    "reviewer_confirmed": ["yes", "no"],
}

MAIN_COLUMNS = [
    "sample_group",
    "sample_index",
    "cluster_id",
    "event_name",
    "schema_id",
    "claim",
    "subject/company",
    "event_time",
    "article_count",
    "source_count",
    "source列表",
    "evidence_id",
    "article_id",
    "title",
    "publish_time",
    "source",
    "evidence_text",
    "evidence_type",
    "evidence_trace_path",
    *HUMAN_FIELDS,
]

REQUIRED_SHEETS = [
    "填写说明",
    "Primary人工审核",
    "Stress人工审核",
    "证据明细",
    "来源与证据追溯",
    "抽样分布",
    "历史排除统计",
    "枚举值说明",
]

STRESS_TERMS = {
    "contradiction_terms": ("否认", "辟谣", "澄清", "不属实", "不实", "失实", "传闻不实", "消息不实", "虚假", "更正", "撤回", "终止", "取消", "无关", "未参与"),
    "investor_question_terms": ("互动易", "问董秘", "董秘", "投资者", "irm", "ir", "问：", "答：", "您好"),
    "market_signal_terms": ("涨停", "跌停", "股价", "主力资金", "净流入", "净流出", "成交额", "成交量", "换手率", "龙虎榜", "收盘", "盘中", "大涨", "大跌", "异动"),
    "reprint_terms": ("转载", "摘自", "来源：", "转自", "综合自"),
}


@dataclass(frozen=True)
class HistorySource:
    name: str
    category: str
    path: Path
    optional: bool = True


HISTORICAL_SOURCES = [
    HistorySource("old_47_dea_gold_manifest", "old_47_dea_gold", ROOT / "reports/task2_dea_cached_human_audit_manifest.json"),
    HistorySource("old_47_dea_manual_audit_v1", "old_47_dea_gold", ROOT / "reports/task2_dea_manual_audit.jsonl"),
    HistorySource("old_47_dea_manual_audit_v2", "old_47_dea_gold", ROOT / "reports/task2_dea_v2_manual_audit.jsonl"),
    HistorySource("old_47_dea_manual_audit_cached", "old_47_dea_gold", ROOT / "reports/task2_dea_v2_cached_manual_audit.jsonl"),
    HistorySource("old_50_safety_gate_manifest", "old_50_safety_gate_independent", ROOT / "reports/task2_safety_gate_v1_2_independent_manifest.json"),
    HistorySource("old_50_independent_samples", "old_50_safety_gate_independent", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/independent_samples.jsonl"),
    HistorySource("old_50_primary_samples", "old_50_safety_gate_independent", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/primary_samples.jsonl"),
    HistorySource("old_50_stress_samples", "old_50_safety_gate_independent", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/contradiction_stress_samples.jsonl"),
    HistorySource("old_50_blind_review_queue", "old_50_safety_gate_independent", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/blind_review_queue.jsonl"),
    HistorySource("old_50_gold_results", "old_50_safety_gate_independent", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent_gold_eval/independent_gold_results.jsonl"),
    HistorySource("dea_v1_gray_samples", "dea_gray", ROOT / "outputs/event_credibility/dea_v1_gray/sample_rows.jsonl"),
    HistorySource("dea_v1_gray_baseline", "dea_gray", ROOT / "outputs/event_credibility/dea_v1_gray/baseline_results.jsonl"),
    HistorySource("dea_v1_gray_dea", "dea_gray", ROOT / "outputs/event_credibility/dea_v1_gray/dea_results.jsonl"),
    HistorySource("dea_v2_gray_samples", "dea_gray", ROOT / "outputs/event_credibility/dea_v2_gray/sample_rows.jsonl"),
    HistorySource("dea_v2_gray_dea", "dea_gray", ROOT / "outputs/event_credibility/dea_v2_gray/dea_results.jsonl"),
    HistorySource("dea_v2_cached_gray_samples", "dea_gray", ROOT / "outputs/event_credibility/dea_v2_cached_gray/sample_rows.jsonl"),
    HistorySource("dea_v2_cached_gray_dea", "dea_gray", ROOT / "outputs/event_credibility/dea_v2_cached_gray/dea_results.jsonl"),
    HistorySource("llm_concurrency_sample_24", "llm_concurrency_benchmark", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/task2_llm_concurrency_sample_24.jsonl"),
    HistorySource("llm_concurrency_workers_1", "llm_concurrency_benchmark", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/llm_review_workers_1.jsonl"),
    HistorySource("llm_concurrency_workers_2", "llm_concurrency_benchmark", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/llm_review_workers_2.jsonl"),
    HistorySource("llm_concurrency_workers_4", "llm_concurrency_benchmark", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/llm_review_workers_4.jsonl"),
    HistorySource("llm_concurrency_workers_8", "llm_concurrency_benchmark", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/llm_review_workers_8.jsonl"),
    HistorySource("isr_diagnostic_round_a", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_diagnostic_v1/round_a_results.jsonl"),
    HistorySource("isr_diagnostic_round_b", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_diagnostic_v1/round_b_critiques.jsonl"),
    HistorySource("isr_diagnostic_final", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_diagnostic_v1/final_isr_results.jsonl"),
    HistorySource("isr_diagnostic_error_analysis", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_diagnostic_v1/error_analysis.jsonl"),
    HistorySource("isr_candidate_all_error_audit", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/all_error_audit.csv"),
    HistorySource("isr_candidates", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/isr_candidates.jsonl"),
    HistorySource("isr_claim_fix_candidates", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/claim_fix_candidates.jsonl"),
    HistorySource("isr_safety_gate_fix_candidates", "isr_diagnostic_and_candidates", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/safety_gate_fix_candidates.jsonl"),
    HistorySource("claim_contract_sample_manifest", "claim_contract_diagnostic", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/sample_manifest.json"),
    HistorySource("claim_contract_contracts", "claim_contract_diagnostic", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/claim_contracts.jsonl"),
    HistorySource("claim_contract_round1", "claim_contract_diagnostic", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/adapter_round1_results.jsonl"),
    HistorySource("claim_contract_cached_dea", "claim_contract_diagnostic", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/adapter_cached_dea_results.jsonl"),
    HistorySource("claim_contract_errors", "claim_contract_diagnostic", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/errors.jsonl"),
    HistorySource("safety_gate_v1_dev", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/dea_safety_gate_v1/safety_gate_results.jsonl"),
    HistorySource("safety_gate_adapter_v2_dev", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/dea_safety_gate_v1_adapter_v2/safety_gate_results.jsonl"),
    HistorySource("safety_gate_v1_1_dev", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/dea_safety_gate_v1_1_dev/safety_gate_v1_1_results.jsonl"),
    HistorySource("safety_gate_v1_2_dev", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_dev/safety_gate_v1_2_results.jsonl"),
    HistorySource("safety_gate_zero_trigger_audit_rows", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/audit_rows.jsonl"),
    HistorySource("safety_gate_zero_trigger_error_focus", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/error_focus_13.jsonl"),
    HistorySource("safety_gate_zero_trigger_insufficient_focus", "safety_gate_dev_diag_replay", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/insufficient_resolution_focus_5.jsonl"),
    HistorySource("safety_gate_failure_audit", "task2_error_audit", ROOT / "reports/task2_dea_safety_gate_failure_audit.json"),
    HistorySource("safety_gate_v1_1_failure_audit", "task2_error_audit", ROOT / "reports/task2_dea_safety_gate_v1_1_failure_audit.json"),
    HistorySource("task2_independent_error_audit", "task2_error_audit", ROOT / "reports/task2_independent_error_audit_v1.json"),
]


def norm(value: Any) -> str:
    if value is None:
        return ""
    try:
        if isinstance(value, float) and math.isnan(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def now_iso() -> str:
    tz = timezone(timedelta(hours=8))
    return datetime.now(tz).isoformat(timespec="seconds")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_paths(paths: Iterable[Path]) -> dict[str, str]:
    return {rel(path): sha256(path) if path.exists() else "missing" for path in paths}


def to_jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return None
        return value
    if isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, dict):
        return {str(k): to_jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [to_jsonable(v) for v in value]
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    return str(value)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(to_jsonable(payload), ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(to_jsonable(row), ensure_ascii=False, allow_nan=False) + "\n")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows = []
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            obj = json.loads(line)
            if not isinstance(obj, dict):
                raise ValueError(f"{path}:{line_number} is not a JSON object")
            rows.append(obj)
    return rows


def parse_jsonish(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return value
    text = norm(value)
    if not text:
        return []
    try:
        return json.loads(text)
    except Exception:
        return value


def extract_cluster_ids(value: Any) -> set[str]:
    ids: set[str] = set()
    if isinstance(value, dict):
        for key, item in value.items():
            lowered = str(key).lower()
            if "cluster_id" in lowered or lowered in {"event_cluster_id", "event_id"}:
                if isinstance(item, str):
                    ids.update(re.findall(r"\bEC-\d{6}\b", item))
                elif isinstance(item, list):
                    for sub in item:
                        if isinstance(sub, str):
                            ids.update(re.findall(r"\bEC-\d{6}\b", sub))
            ids.update(extract_cluster_ids(item))
    elif isinstance(value, list):
        for item in value:
            ids.update(extract_cluster_ids(item))
    elif isinstance(value, str):
        ids.update(re.findall(r"\bEC-\d{6}\b", value))
    return ids


def ids_from_csv(path: Path) -> set[str]:
    ids: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            ids.update(extract_cluster_ids(row))
    return ids


def ids_from_source(path: Path) -> set[str]:
    if not path.exists():
        return set()
    if path.suffix == ".jsonl":
        ids: set[str] = set()
        for row in read_jsonl(path):
            ids.update(extract_cluster_ids(row))
        return ids
    if path.suffix == ".json":
        return extract_cluster_ids(json.loads(path.read_text(encoding="utf-8")))
    if path.suffix == ".csv":
        return ids_from_csv(path)
    return set()


def stable_text(value: Any, limit: int = 320) -> str:
    text = norm(value).lower()
    text = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", text)
    return text[:limit]


def sequence_similarity(left: Any, right: Any) -> float:
    a = stable_text(left)
    b = stable_text(right)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    max_possible = 2.0 * min(len(a), len(b)) / (len(a) + len(b))
    if max_possible < 0.82:
        return max_possible
    matcher = SequenceMatcher(None, a, b)
    if matcher.real_quick_ratio() < 0.82:
        return matcher.real_quick_ratio()
    if matcher.quick_ratio() < 0.82:
        return matcher.quick_ratio()
    return matcher.ratio()


def first_date(value: Any) -> str:
    text = norm(value)
    match = re.search(r"(20\d{2})[-/年](\d{1,2})[-/月](\d{1,2})", text)
    if not match:
        return ""
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def derive_event_name(event_name: Any, schema_id: Any) -> str:
    value = norm(event_name)
    if value:
        return value
    schema = norm(schema_id)
    if not schema:
        return ""
    parts = schema.split("_")
    if len(parts) >= 3:
        return "_".join(parts[2:]).strip()
    return schema


def derive_event_time(earliest: Any, latest: Any, evidence_items: list[dict[str, Any]]) -> str:
    start = norm(earliest)
    end = norm(latest)
    if start and end and start != end:
        return f"{start} ~ {end}"
    if start or end:
        return start or end
    times = sorted({evidence_time(item) for item in evidence_items if evidence_time(item)})
    if len(times) >= 2:
        return f"{times[0]} ~ {times[-1]}"
    return times[0] if times else ""


def month_bucket(value: Any) -> str:
    date = first_date(value)
    return date[:7] if date else "unknown_month"


def date_distance_days(left: Any, right: Any) -> int | None:
    a = first_date(left)
    b = first_date(right)
    if not a or not b:
        return None
    try:
        da = datetime.strptime(a, "%Y-%m-%d")
        db = datetime.strptime(b, "%Y-%m-%d")
    except Exception:
        return None
    return abs((da - db).days)


def title_parts(value: Any) -> list[str]:
    return [norm(part) for part in re.split(r"\s*\|\|\s*|\n+", norm(value)) if norm(part)]


def evidence_text(item: dict[str, Any]) -> str:
    return norm(item.get("sentence") or item.get("text") or item.get("content") or item.get("证据文本"))


def evidence_source(item: dict[str, Any]) -> str:
    return norm(item.get("source") or item.get("article_source") or item.get("来源"))


def evidence_time(item: dict[str, Any]) -> str:
    return norm(item.get("publish_time") or item.get("article_publish_time") or item.get("发布时间"))


def evidence_article_id(item: dict[str, Any]) -> str:
    return norm(item.get("article_id") or item.get("trace_id") or item.get("document_id"))


def evidence_id(item: dict[str, Any]) -> str:
    return norm(item.get("evidence_id") or item.get("id"))


def source_list_from_row(row: dict[str, Any]) -> list[str]:
    parsed = parse_jsonish(row.get("source_list"))
    if isinstance(parsed, list):
        raw_values = parsed
    else:
        raw_values = re.split(r"\s*\|\s*|;|；|,", norm(row.get("source_list")))
    seen: set[str] = set()
    out: list[str] = []
    for value in raw_values:
        item = norm(value)
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return out


def source_count(row: dict[str, Any]) -> int:
    sources = source_list_from_row(row)
    if sources:
        return len(sources)
    return len({evidence_source(item) for item in row.get("evidence_items", []) if evidence_source(item)})


def source_count_bucket(count: int) -> str:
    if count <= 1:
        return "1"
    if count <= 3:
        return "2-3"
    return "4+"


def article_count_bucket(count: int) -> str:
    if count <= 1:
        return "1"
    if count <= 3:
        return "2-3"
    return "4+"


def target_key(row: dict[str, Any]) -> str:
    return norm(row.get("target_code")) or stable_text(row.get("target_name"), 80)


def event_profile(row: dict[str, Any]) -> dict[str, Any]:
    title_text = " ".join(title_parts(row.get("title_samples")))
    claim_text = norm(row.get("event_claim") or row.get("claim"))
    event_time = norm(row.get("event_time") or row.get("earliest_publish_time"))
    return {
        "cluster_id": norm(row.get("cluster_id")),
        "target_key": target_key(row),
        "target_name": norm(row.get("target_name")),
        "event_name": norm(row.get("event_name") or row.get("schema_id")),
        "date": first_date(event_time),
        "title_text": title_text,
        "claim_text": claim_text,
    }


def is_near_duplicate(left: dict[str, Any], right: dict[str, Any]) -> tuple[bool, str, float]:
    if not left["target_key"] or left["target_key"] != right["target_key"]:
        return False, "", 0.0
    if left["event_name"] and right["event_name"] and left["event_name"] != right["event_name"]:
        return False, "", 0.0
    title_sim = sequence_similarity(left["title_text"], right["title_text"])
    claim_sim = sequence_similarity(left["claim_text"], right["claim_text"])
    date_distance = date_distance_days(left["date"], right["date"])
    if left["date"] and left["date"] == right["date"] and title_sim >= 0.86:
        return True, "same_target_event_date_title_similarity", title_sim
    if date_distance is not None and date_distance <= 2 and max(title_sim, claim_sim) >= 0.90:
        return True, "same_target_event_near_date_text_similarity", max(title_sim, claim_sim)
    if title_sim >= 0.96:
        return True, "same_target_event_high_title_similarity", title_sim
    return False, "", max(title_sim, claim_sim)


def collect_exclusions() -> dict[str, Any]:
    source_rows = []
    category_ids: dict[str, set[str]] = defaultdict(set)
    union: set[str] = set()
    missing_sources = []
    for source in HISTORICAL_SOURCES:
        ids = ids_from_source(source.path)
        if not source.path.exists() and not source.optional:
            missing_sources.append(rel(source.path))
        category_ids[source.category].update(ids)
        union.update(ids)
        source_rows.append(
            {
                "name": source.name,
                "category": source.category,
                "path": rel(source.path),
                "exists": source.path.exists(),
                "cluster_id_count": len(ids),
                "cluster_ids": sorted(ids),
            }
        )
    return {
        "sources": source_rows,
        "category_cluster_id_counts": {category: len(ids) for category, ids in sorted(category_ids.items())},
        "category_cluster_ids": {category: sorted(ids) for category, ids in sorted(category_ids.items())},
        "excluded_cluster_id_count": len(union),
        "excluded_cluster_ids": sorted(union),
        "missing_sources": missing_sources,
    }


def article_trace_ok(item: dict[str, Any], article_map: dict[str, dict[str, Any]]) -> tuple[bool, str]:
    article_id = evidence_article_id(item)
    if not article_id:
        return False, "missing_article_id"
    article = article_map.get(article_id)
    if not article:
        return False, "article_id_not_in_fulltext_index"
    if not evidence_id(item):
        return False, "missing_evidence_id"
    if not evidence_source(item):
        return False, "missing_source"
    text = evidence_text(item)
    if not text:
        return False, "missing_evidence_text"
    content = norm(article.get("content"))
    title = norm(article.get("article_title"))
    compact_text = stable_text(text, 120)
    compact_article = stable_text(content + " " + title, 100000)
    if compact_text and compact_text in compact_article:
        return True, ""
    if text in content or text in title:
        return True, ""
    return False, "evidence_text_not_found_in_article_content"


def normalize_evidence_for_sample(
    row: dict[str, Any],
    article_map: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    normalized = []
    errors = []
    seen_ids: set[str] = set()
    for idx, raw in enumerate(row.get("evidence_items", []) or [], 1):
        if not isinstance(raw, dict):
            errors.append(f"non_dict_evidence_{idx}")
            continue
        item = dict(raw)
        eid = evidence_id(item) or f"E{idx}"
        article_id = evidence_article_id(item)
        article = article_map.get(article_id, {})
        traceable, reason = article_trace_ok({**item, "evidence_id": eid}, article_map)
        if not traceable:
            errors.append(f"{eid}:{reason}")
        if eid in seen_ids:
            errors.append(f"{eid}:duplicate_evidence_id")
        seen_ids.add(eid)
        normalized.append(
            {
                "evidence_id": eid,
                "article_id": article_id,
                "trace_id": norm(item.get("trace_id") or article_id),
                "title": norm(item.get("title") or article.get("article_title")),
                "publish_time": norm(item.get("publish_time") or article.get("article_publish_time")),
                "source": norm(item.get("source") or article.get("article_source")),
                "source_type": norm(item.get("source_type") or article.get("source_type")),
                "evidence_text": evidence_text(item),
                "sentence_score": item.get("sentence_score", ""),
                "evidence_source": norm(item.get("evidence_source") or item.get("retrieval_source") or "cluster_pool"),
                "evidence_type": norm(item.get("source_type") or article.get("source_type") or item.get("evidence_source") or "fulltext_sentence"),
                "evidence_trace_path": f"{rel(ARTICLE_INDEX)}#article_id={article_id}",
                "traceable": bool(traceable),
            }
        )
    return normalized, errors


def load_article_map(article_index_path: Path) -> dict[str, dict[str, Any]]:
    df = pd.read_parquet(article_index_path)
    out: dict[str, dict[str, Any]] = {}
    for row in df.to_dict("records"):
        article_id = norm(row.get("article_id"))
        if article_id:
            out[article_id] = {str(k): to_jsonable(v) for k, v in row.items()}
    return out


def load_candidate_rows(
    base_scores_path: Path,
    fulltext_scores_path: Path,
    evidence_path: Path,
    article_index_path: Path,
) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, dict[str, Any]]]:
    base_pool = dea_mod.build_review_pool(base_scores_path, evidence_path)
    fulltext_scores = pd.read_parquet(fulltext_scores_path)
    if len(fulltext_scores) != 2061:
        raise ValueError(f"formal score candidate table must have 2061 rows, got {len(fulltext_scores)}")
    score_by_id = {norm(row.get("cluster_id")): row for row in fulltext_scores.to_dict("records")}
    article_map = load_article_map(article_index_path)
    rows = []
    invalid_counts: Counter[str] = Counter()
    seen: set[str] = set()
    evidence_total = 0
    untraceable_total = 0
    for base in base_pool:
        cid = norm(base.get("cluster_id"))
        if not cid:
            invalid_counts["missing_cluster_id"] += 1
            continue
        if cid in seen:
            invalid_counts["duplicate_cluster_id_in_pool"] += 1
            continue
        seen.add(cid)
        score_row = score_by_id.get(cid, {})
        evidence_items, errors = normalize_evidence_for_sample(base, article_map)
        evidence_total += len(evidence_items)
        untraceable_total += sum(1 for item in evidence_items if not item.get("traceable"))
        if errors:
            invalid_counts["evidence_traceability_error"] += 1
            continue
        if not evidence_items:
            invalid_counts["missing_evidence"] += 1
            continue
        event_name = derive_event_name(base.get("event_name") or score_row.get("event_name"), base.get("schema_id") or score_row.get("schema_id"))
        event_time = derive_event_time(
            base.get("earliest_publish_time") or score_row.get("earliest_publish_time"),
            base.get("latest_publish_time") or score_row.get("latest_publish_time"),
            evidence_items,
        )
        clean = {
            "cluster_id": cid,
            "target_type": norm(score_row.get("target_type")),
            "target_code": norm(base.get("target_code") or score_row.get("target_code")),
            "target_name": norm(base.get("target_name") or score_row.get("target_name")),
            "schema_id": norm(base.get("schema_id") or score_row.get("schema_id")),
            "event_name": event_name,
            "event_claim": norm(base.get("event_claim")) or f"关于{norm(base.get('target_name') or score_row.get('target_name'))}的{event_name}相关主张",
            "event_time": event_time,
            "earliest_publish_time": norm(base.get("earliest_publish_time") or score_row.get("earliest_publish_time")),
            "latest_publish_time": norm(base.get("latest_publish_time") or score_row.get("latest_publish_time")),
            "article_count": int(score_row.get("member_count") or 0),
            "source_count": source_count({**score_row, "evidence_items": evidence_items}),
            "source_list": source_list_from_row(score_row),
            "title_samples": norm(base.get("title_samples") or score_row.get("title_samples")),
            "source_type_distribution": parse_jsonish(score_row.get("source_type_distribution")),
            "evidence_items": evidence_items,
        }
        rows.append(clean)
    metadata = {
        "raw_candidate_count": len(base_pool),
        "formal_score_row_count": len(fulltext_scores),
        "valid_traceable_candidate_count": len(rows),
        "invalid_candidate_counts": dict(sorted(invalid_counts.items())),
        "evidence_total_in_raw_pool": evidence_total,
        "untraceable_evidence_total_in_raw_pool": untraceable_total,
    }
    return rows, metadata, score_by_id


def candidate_stratum(row: dict[str, Any]) -> str:
    return "|".join(
        [
            norm(row.get("event_name")) or "unknown_event",
            norm(row.get("target_type")) or "unknown_target_type",
            "single_source" if int(row.get("source_count") or 0) <= 1 else "multi_source",
            source_count_bucket(int(row.get("source_count") or 0)),
            article_count_bucket(int(row.get("article_count") or 0)),
            "trace_complete",
            month_bucket(row.get("earliest_publish_time") or row.get("event_time")),
        ]
    )


def proportional_stratified_sample(rows: list[dict[str, Any]], count: int, seed: int) -> list[dict[str, Any]]:
    if len(rows) < count:
        raise ValueError(f"not enough primary candidates: need {count}, got {len(rows)}")
    rng = random.Random(seed)
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[candidate_stratum(row)].append(row)
    for group_rows in groups.values():
        group_rows.sort(key=lambda item: norm(item.get("cluster_id")))
        rng.shuffle(group_rows)
    total = len(rows)
    quota_rows = []
    allocated = 0
    for key, group_rows in sorted(groups.items()):
        exact = len(group_rows) * count / total
        base = min(len(group_rows), int(math.floor(exact)))
        quota_rows.append([key, base, exact - base, len(group_rows)])
        allocated += base
    remaining = count - allocated
    quota_rows.sort(key=lambda item: (-item[2], item[0]))
    while remaining > 0:
        progressed = False
        for item in quota_rows:
            if remaining <= 0:
                break
            if item[1] < item[3]:
                item[1] += 1
                remaining -= 1
                progressed = True
        if not progressed:
            break
    selected = []
    for key, quota, _fraction, _size in quota_rows:
        selected.extend(groups[key][:quota])
    if len(selected) != count:
        raise AssertionError(f"proportional stratified sampling produced {len(selected)} rows, expected {count}")
    selected.sort(key=lambda item: norm(item.get("cluster_id")))
    rng.shuffle(selected)
    return selected


def row_text_for_signals(row: dict[str, Any]) -> str:
    parts = [
        row.get("event_claim"),
        row.get("event_name"),
        row.get("schema_id"),
        row.get("title_samples"),
        " ".join(row.get("source_list") or []),
    ]
    for item in row.get("evidence_items", []):
        parts.extend([item.get("title"), item.get("source"), item.get("evidence_text")])
    return norm(" ".join(norm(part) for part in parts))


def count_terms(text: str, terms: Iterable[str]) -> int:
    return sum(1 for term in terms if term and term in text)


def target_match_ratio(row: dict[str, Any]) -> float:
    target = norm(row.get("target_name"))
    code = norm(row.get("target_code"))
    if not target and not code:
        return 1.0
    matches = 0
    total = 0
    for item in row.get("evidence_items", []):
        total += 1
        text = norm(" ".join([item.get("title", ""), item.get("evidence_text", "")]))
        if (target and target in text) or (code and code in text):
            matches += 1
    return matches / max(1, total)


def duplicate_evidence_ratio(row: dict[str, Any]) -> float:
    items = row.get("evidence_items", [])
    if not items:
        return 0.0
    keys = [stable_text(f"{item.get('title', '')} {item.get('evidence_text', '')}", 180) for item in items]
    return 1.0 - (len(set(keys)) / max(1, len(keys)))


def time_mismatch_ratio(row: dict[str, Any]) -> float:
    event_time = row.get("event_time") or row.get("earliest_publish_time")
    total = 0
    mismatched = 0
    for item in row.get("evidence_items", []):
        days = date_distance_days(event_time, item.get("publish_time"))
        if days is None:
            continue
        total += 1
        if days > 30:
            mismatched += 1
    return mismatched / max(1, total)


def model_signal_maps() -> dict[str, dict[str, Any]]:
    round1 = {norm(row.get("cluster_id")): row for row in read_jsonl(ROOT / "outputs/event_credibility/final_cached_dea_v1_260/round1_results.jsonl")}
    final_rows = {norm(row.get("cluster_id")): row for row in read_jsonl(ROOT / "outputs/event_credibility/final_cached_dea_v1_260/final_results.jsonl")}
    formal = {norm(row.get("cluster_id")): row for row in read_jsonl(ROOT / "outputs/event_credibility/llm_review_results.jsonl")}
    out: dict[str, dict[str, Any]] = {}
    for cid in sorted(set(round1) | set(final_rows) | set(formal)):
        final_row = final_rows.get(cid, {})
        final_result = final_row.get("final_result") if isinstance(final_row.get("final_result"), dict) else {}
        out[cid] = {
            "round1_verdict": norm(round1.get(cid, {}).get("verdict")),
            "round1_confidence": round1.get(cid, {}).get("confidence", ""),
            "cached_dea_verdict": norm(final_result.get("verdict") or formal.get(cid, {}).get("verdict")),
            "cached_dea_confidence": final_result.get("confidence", formal.get(cid, {}).get("confidence", "")),
            "dea_triggered": bool(final_row.get("dea_triggered")),
            "selected_result_source": norm(final_row.get("selected_result_source")),
        }
    return out


def stress_features(row: dict[str, Any], score_row: dict[str, Any], model_map: dict[str, dict[str, Any]]) -> dict[str, Any]:
    cid = norm(row.get("cluster_id"))
    text = row_text_for_signals(row)
    model = model_map.get(cid, {})
    round1_verdict = norm(model.get("round1_verdict"))
    cached_verdict = norm(model.get("cached_dea_verdict"))
    confidence_values = []
    for key in ("round1_confidence", "cached_dea_confidence"):
        try:
            value = float(model.get(key))
        except Exception:
            continue
        if math.isfinite(value):
            confidence_values.append(value)
    llm_conf = min(confidence_values) if confidence_values else None
    disagreement = bool(round1_verdict and cached_verdict and round1_verdict != cached_verdict)
    insufficient_changed = bool(round1_verdict == "insufficient" and cached_verdict in {"support", "contradict"})
    support_count = int(score_row.get("fulltext_support_count") or 0)
    contradict_count = int(score_row.get("fulltext_contradict_count") or 0)
    contradiction_terms = count_terms(text, STRESS_TERMS["contradiction_terms"])
    investor_terms = count_terms(text, STRESS_TERMS["investor_question_terms"])
    market_terms = count_terms(text, STRESS_TERMS["market_signal_terms"])
    official_sources = [
        item
        for item in row.get("evidence_items", [])
        if norm(item.get("source_type")) == "official" or any(term in norm(item.get("source")) for term in ("公告", "交易所", "互动易", "政府", "证监"))
    ]
    nonofficial_sources = [item for item in row.get("evidence_items", []) if item not in official_sources]
    source_dup_ratio = duplicate_evidence_ratio(row)
    target_ratio = target_match_ratio(row)
    time_ratio = time_mismatch_ratio(row)
    source_n = int(row.get("source_count") or 0)
    evidence_n = len(row.get("evidence_items", []))
    risk_reasons = norm(score_row.get("risk_reasons"))
    status = norm(score_row.get("fulltext_evidence_status"))
    flags = {
        "model_verdict_disagreement": disagreement,
        "low_model_confidence": bool(llm_conf is not None and llm_conf < 0.75),
        "insufficient_changed_by_dea": insufficient_changed,
        "support_and_contradict_evidence_coexist": bool(support_count > 0 and contradict_count > 0),
        "official_media_conflict_signal": bool(official_sources and nonofficial_sources and contradiction_terms),
        "investor_question_risk": bool(investor_terms),
        "market_signal_risk": bool(market_terms),
        "subject_similarity_risk": bool(target_ratio < 0.5),
        "time_mismatch_risk": bool(time_ratio >= 0.4),
        "duplicate_reprint_risk": bool(source_dup_ratio >= 0.25 or count_terms(text, STRESS_TERMS["reprint_terms"])),
        "contradict_candidate": bool(status == "contradict" or cached_verdict == "contradict" or contradiction_terms),
        "unclear_candidate": bool("不确定" in risk_reasons or "无法确认" in risk_reasons or status == "insufficient"),
        "single_source_or_insufficient_evidence": bool(source_n <= 1 or evidence_n <= 2 or "单一来源" in risk_reasons),
    }
    weighted_score = (
        14 * flags["model_verdict_disagreement"]
        + 13 * flags["insufficient_changed_by_dea"]
        + 10 * flags["support_and_contradict_evidence_coexist"]
        + 10 * flags["official_media_conflict_signal"]
        + 8 * flags["contradict_candidate"]
        + 7 * flags["low_model_confidence"]
        + 6 * flags["investor_question_risk"]
        + 6 * flags["market_signal_risk"]
        + 5 * flags["time_mismatch_risk"]
        + 5 * flags["subject_similarity_risk"]
        + 4 * flags["duplicate_reprint_risk"]
        + 4 * flags["unclear_candidate"]
        + 3 * flags["single_source_or_insufficient_evidence"]
    )
    return {
        "cluster_id": cid,
        "stress_score": int(weighted_score),
        "stress_flags": {key: bool(value) for key, value in flags.items()},
        "round1_verdict": round1_verdict,
        "round1_confidence": model.get("round1_confidence", ""),
        "cached_dea_verdict": cached_verdict,
        "cached_dea_confidence": model.get("cached_dea_confidence", ""),
        "dea_triggered": bool(model.get("dea_triggered")),
        "fulltext_evidence_status": status,
        "fulltext_support_count": support_count,
        "fulltext_contradict_count": contradict_count,
        "target_match_ratio": round(target_ratio, 4),
        "time_mismatch_ratio": round(time_ratio, 4),
        "duplicate_evidence_ratio": round(source_dup_ratio, 4),
        "source_count": source_n,
        "evidence_count": evidence_n,
    }


def distribution_rows(candidates: list[dict[str, Any]], primary: list[dict[str, Any]], stress: list[dict[str, Any]]) -> list[dict[str, Any]]:
    sets = {
        "candidate_after_exclusion": candidates,
        "Primary": primary,
        "Stress": stress,
    }
    rows = []
    for set_name, data in sets.items():
        dimensions = {
            "event_name": [norm(row.get("event_name")) or "unknown" for row in data],
            "target_type": [norm(row.get("target_type")) or "unknown" for row in data],
            "single_multi_source": ["single_source" if int(row.get("source_count") or 0) <= 1 else "multi_source" for row in data],
            "source_count_bucket": [source_count_bucket(int(row.get("source_count") or 0)) for row in data],
            "article_count_bucket": [article_count_bucket(int(row.get("article_count") or 0)) for row in data],
            "publish_month": [month_bucket(row.get("earliest_publish_time") or row.get("event_time")) for row in data],
            "evidence_traceability": ["complete" if all(item.get("traceable") for item in row.get("evidence_items", [])) else "incomplete" for row in data],
        }
        total = len(data)
        for dimension, values in dimensions.items():
            for value, count in sorted(Counter(values).items()):
                rows.append(
                    {
                        "set": set_name,
                        "dimension": dimension,
                        "value": value,
                        "count": count,
                        "ratio": round(count / max(1, total), 6),
                    }
                )
    return rows


def history_stat_rows(exclusion_manifest: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for source in exclusion_manifest["sources"]:
        rows.append(
            {
                "level": "source",
                "name": source["name"],
                "category": source["category"],
                "path": source["path"],
                "cluster_id_count": source["cluster_id_count"],
                "exists": source["exists"],
            }
        )
    for category, count in exclusion_manifest["category_cluster_id_counts"].items():
        rows.append(
            {
                "level": "category_union",
                "name": category,
                "category": category,
                "path": "",
                "cluster_id_count": count,
                "exists": True,
            }
        )
    rows.append(
        {
            "level": "all_union",
            "name": "historical_exact_union",
            "category": "all",
            "path": "",
            "cluster_id_count": exclusion_manifest["excluded_cluster_id_count"],
            "exists": True,
        }
    )
    rows.append(
        {
            "level": "near_duplicate",
            "name": "near_duplicate_to_historical",
            "category": "near_duplicate",
            "path": "",
            "cluster_id_count": exclusion_manifest["near_duplicate_exclusion_count"],
            "exists": True,
        }
    )
    return rows


def join_unique(values: Iterable[Any], sep: str = "; ") -> str:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        item = norm(value)
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return sep.join(out)


def safe_cell(value: Any, limit: int = 32000) -> str:
    if isinstance(value, (list, dict)):
        text = json.dumps(to_jsonable(value), ensure_ascii=False, allow_nan=False)
    else:
        text = norm(value)
    return text if len(text) <= limit else text[: limit - 40] + " ...[truncated in main sheet; see evidence detail]"


def sample_main_row(sample: dict[str, Any]) -> dict[str, Any]:
    items = sample.get("evidence_items", [])
    evidence_text_lines = []
    for item in items:
        evidence_text_lines.append(
            f"{item['evidence_id']} | {item['source']} | {item['publish_time']} | {item['title']} | {item['evidence_text']}"
        )
    row = {
        "sample_group": sample["sample_group"],
        "sample_index": sample["sample_index"],
        "cluster_id": sample["cluster_id"],
        "event_name": sample["event_name"],
        "schema_id": sample["schema_id"],
        "claim": sample["event_claim"],
        "subject/company": sample["target_name"],
        "event_time": sample["event_time"],
        "article_count": sample["article_count"],
        "source_count": sample["source_count"],
        "source列表": join_unique(sample.get("source_list") or []),
        "evidence_id": join_unique([item.get("evidence_id") for item in items]),
        "article_id": join_unique([item.get("article_id") for item in items]),
        "title": join_unique([item.get("title") for item in items], sep="\n"),
        "publish_time": join_unique([item.get("publish_time") for item in items]),
        "source": join_unique([item.get("source") for item in items]),
        "evidence_text": "\n".join(evidence_text_lines),
        "evidence_type": join_unique([item.get("evidence_type") for item in items]),
        "evidence_trace_path": join_unique([item.get("evidence_trace_path") for item in items], sep="\n"),
    }
    for field in HUMAN_FIELDS:
        row[field] = ""
    return {column: safe_cell(row.get(column, "")) for column in MAIN_COLUMNS}


def evidence_detail_rows(samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for sample in samples:
        for item in sample.get("evidence_items", []):
            rows.append(
                {
                    "sample_group": sample["sample_group"],
                    "sample_index": sample["sample_index"],
                    "cluster_id": sample["cluster_id"],
                    "event_name": sample["event_name"],
                    "schema_id": sample["schema_id"],
                    "claim": sample["event_claim"],
                    "subject/company": sample["target_name"],
                    "evidence_id": item["evidence_id"],
                    "article_id": item["article_id"],
                    "title": item["title"],
                    "publish_time": item["publish_time"],
                    "source": item["source"],
                    "source_type": item["source_type"],
                    "evidence_text": item["evidence_text"],
                    "evidence_type": item["evidence_type"],
                    "evidence_source": item["evidence_source"],
                    "evidence_trace_path": item["evidence_trace_path"],
                    "traceable": "yes" if item.get("traceable") else "no",
                }
            )
    return rows


def content_excerpt(content: str, evidence: str, radius: int = 500) -> str:
    content = norm(content)
    evidence = norm(evidence)
    if not content:
        return ""
    idx = content.find(evidence[: min(len(evidence), 60)])
    if idx < 0:
        return content[: radius * 2]
    start = max(0, idx - radius)
    end = min(len(content), idx + len(evidence) + radius)
    return content[start:end]


def trace_rows(samples: list[dict[str, Any]], article_map: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for sample in samples:
        for item in sample.get("evidence_items", []):
            article = article_map.get(item["article_id"], {})
            rows.append(
                {
                    "sample_group": sample["sample_group"],
                    "sample_index": sample["sample_index"],
                    "cluster_id": sample["cluster_id"],
                    "evidence_id": item["evidence_id"],
                    "article_id": item["article_id"],
                    "article_title": article.get("article_title", item["title"]),
                    "article_source": article.get("article_source", item["source"]),
                    "article_publish_time": article.get("article_publish_time", item["publish_time"]),
                    "article_source_type": article.get("source_type", item["source_type"]),
                    "content_length": article.get("content_length", ""),
                    "evidence_text": item["evidence_text"],
                    "article_content_excerpt": content_excerpt(article.get("content", ""), item["evidence_text"]),
                    "trace_path": item["evidence_trace_path"],
                }
            )
    return rows


def instruction_rows(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        {"item": "任务", "value": "Task2 Step 6.1 独立500条事件簇证据可信判断盲审。"},
        {"item": "判断对象", "value": "事件簇/事件主张的证据可信判断，不是单篇新闻真假。"},
        {"item": "Primary", "value": "400条；用于估计2061个事件簇自然分布表现。"},
        {"item": "Stress", "value": "100条；困难/高风险场景，只能单独报告，不并入自然总体准确率。"},
        {"item": "禁止", "value": "不要在人工主表中使用或查找模型verdict、confidence、可信分、规则结论或LLM理由。"},
        {"item": "证据引用", "value": "human_evidence_ids 只填写本行 evidence_id 中出现的编号，例如 E1;E3。"},
        {"item": "verdict枚举", "value": "support / contradict / insufficient / unclear。"},
        {"item": "人工字段", "value": "浅橙色列为人工填写字段；生成时全部为空。"},
        {"item": "manifest", "value": manifest.get("manifest_path", "")},
        {"item": "status", "value": manifest.get("status", "")},
    ]


def enum_rows() -> list[dict[str, Any]]:
    rows = []
    descriptions = {
        "human_verdict": "人工证据结论。",
        "human_confidence": "人工对所填verdict的确定程度。",
        "human_evidence_ids": "支持人工结论的证据编号，多个用分号分隔；无枚举。",
        "claim_clarity": "主张是否清晰可判。",
        "subject_match": "证据主体是否匹配主张主体。",
        "time_match": "证据时间是否匹配事件时间。",
        "evidence_sufficiency": "证据是否足以支撑人工结论。",
        "evidence_relevance": "证据是否与主张核心相关。",
        "source_independence": "来源是否相互独立，非简单转载链。",
        "official_source_present": "是否存在正式公告、交易所、政府或监管来源。",
        "investor_question_only": "是否仅有投资者提问而缺少公司确认事实。",
        "market_signal_only": "是否只是行情、资金流、股价波动等市场信号。",
        "contradiction_present": "是否存在明确否认、澄清、反转或冲突证据。",
        "primary_error_type": "如需记录主要错误/风险类型，选择最主要一项。",
        "reviewer_notes": "人工备注；无枚举。",
        "reviewer_confirmed": "完成审核后填 yes，否则 no。",
    }
    for field in HUMAN_FIELDS:
        rows.append(
            {
                "field": field,
                "allowed_values": ", ".join(ALLOWED_VALUES.get(field, [])),
                "description": descriptions.get(field, ""),
            }
        )
    return rows


def write_sheet(
    wb: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    human_columns: set[str] | None = None,
) -> None:
    ws = wb.create_sheet(title=title[:31])
    if not rows:
        rows = [{"empty": ""}]
    headers = list(rows[0].keys())
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    human_fill = PatternFill("solid", fgColor="FCE4D6")
    auto_fill = PatternFill("solid", fgColor="F2F2F2")
    human_columns = human_columns or set()
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([safe_cell(row.get(header, "")) for header in headers])
    for column_idx, header in enumerate(headers, 1):
        fill = human_fill if header in human_columns else auto_fill
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, column_idx)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, 1):
        values = [safe_cell(row.get(header, "")) for row in rows[:200]]
        max_len = max([len(str(header))] + [min(len(value), 120) for value in values])
        width = min(max(max_len + 2, 12), 80)
        if header in {"evidence_text", "article_content_excerpt", "evidence_trace_path", "title"}:
            width = 70
        ws.column_dimensions[get_column_letter(idx)].width = width
    for idx, header in enumerate(headers, 1):
        if header not in ALLOWED_VALUES:
            continue
        values = ALLOWED_VALUES[header]
        dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(idx)}2:{get_column_letter(idx)}{max(ws.max_row, 2)}")


def write_workbook(path: Path, workbook_rows: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in REQUIRED_SHEETS:
        write_sheet(
            wb,
            sheet,
            workbook_rows[sheet],
            human_columns=set(HUMAN_FIELDS) if sheet in {"Primary人工审核", "Stress人工审核"} else set(),
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def render_guide(manifest: dict[str, Any]) -> str:
    counts = manifest["sample_counts"]
    lines = [
        "# Task2 Independent500 Human Audit Guide",
        "",
        "本包用于 Task2 Step 6.1 的500条全新独立盲审标注。判断对象是事件簇/事件主张与证据之间的可信关系，不是单篇新闻真假。",
        "",
        "## 子集使用",
        "",
        f"- Primary: {counts['Primary']} 条，按 seed=2026 从排除后候选池按自然比例抽取，用于估计2061个事件簇总体表现。",
        f"- Stress: {counts['Stress']} 条，按模型分歧和风险信号选择，只能单独评估困难场景，不能与 Primary 合并声明自然总体准确率。",
        "",
        "## 人工字段",
        "",
        "人工字段在工作簿中为浅橙色，生成时全部为空。请只填写这些字段，不修改自动字段。",
        "",
        "| field | allowed values |",
        "|---|---|",
    ]
    for field in HUMAN_FIELDS:
        lines.append(f"| {field} | {', '.join(ALLOWED_VALUES.get(field, [])) or 'free text'} |")
    lines.extend(
        [
            "",
            "## Verdict 定义",
            "",
            "- support: 证据明确支持事件核心主张。",
            "- contradict: 证据明确否定、澄清、撤回或反转事件核心主张。",
            "- insufficient: 现有证据不足以支持或反驳，不等于事件为假。",
            "- unclear: 主张或证据边界不清，无法可靠归入前三类。",
            "",
            "## 独立性说明",
            "",
            "旧47条DEA人工金标、旧50条Safety Gate独立集、DEA灰度、并发benchmark、ISR、Claim Contract、Safety Gate开发/诊断/回放、错误审计样本及其近重复事件簇均已排除。",
            "",
            f"工作簿: {manifest['outputs']['workbook']}",
            f"Manifest: {manifest['outputs']['manifest']}",
            f"排除清单: {manifest['outputs']['exclusion_manifest']}",
        ]
    )
    return "\n".join(lines) + "\n"


def prepare_backup_dir(path: Path | None) -> Path:
    if path:
        backup_dir = path
    else:
        backup_dir = BACKUP_ROOT / f"task2_step6_1_independent500_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    return backup_dir


def copy_protected_files_once(backup_dir: Path) -> None:
    copies_dir = backup_dir / "protected_file_copies"
    copies_dir.mkdir(parents=True, exist_ok=True)
    for path in PROTECTED_FORMAL_OUTPUTS + FROZEN_FILES:
        if not path.exists():
            continue
        target = copies_dir / rel(path)
        if target.exists():
            continue
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)


def ensure_no_overwrite(paths: Iterable[Path], overwrite: bool) -> None:
    existing = [path for path in paths if path.exists()]
    if existing and not overwrite:
        raise FileExistsError(f"refusing to overwrite existing Step 6.1 outputs: {[rel(path) for path in existing]}")


def append_sample_metadata(rows: list[dict[str, Any]], group: str, start_index: int) -> list[dict[str, Any]]:
    out = []
    for offset, row in enumerate(rows):
        sample = dict(row)
        sample["sample_group"] = group
        sample["sample_index"] = start_index + offset
        sample["stratum_key"] = candidate_stratum(sample)
        out.append(sample)
    return out


def remove_near_duplicates_to_history(
    eligible: list[dict[str, Any]],
    historical_ids: set[str],
    all_candidates_by_id: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    history_profiles = [
        event_profile(all_candidates_by_id[cid])
        for cid in sorted(historical_ids)
        if cid in all_candidates_by_id
    ]
    kept = []
    removed = []
    for row in eligible:
        profile = event_profile(row)
        hit = None
        for historical in history_profiles:
            duplicate, method, score = is_near_duplicate(profile, historical)
            if duplicate:
                hit = {
                    "cluster_id": profile["cluster_id"],
                    "near_duplicate_historical_cluster_id": historical["cluster_id"],
                    "method": method,
                    "similarity": round(score, 6),
                }
                break
        if hit:
            removed.append(hit)
        else:
            kept.append(row)
    return kept, removed


def row_near_any(row: dict[str, Any], selected_profiles: list[dict[str, Any]]) -> tuple[bool, dict[str, Any]]:
    profile = event_profile(row)
    for selected in selected_profiles:
        duplicate, method, score = is_near_duplicate(profile, selected)
        if duplicate:
            return True, {
                "cluster_id": profile["cluster_id"],
                "near_duplicate_selected_cluster_id": selected["cluster_id"],
                "method": method,
                "similarity": round(score, 6),
            }
    return False, {}


def select_samples(
    candidates: list[dict[str, Any]],
    score_by_id: dict[str, dict[str, Any]],
    exclusions: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    exact_excluded = set(exclusions["excluded_cluster_ids"])
    candidate_ids = {row["cluster_id"] for row in candidates}
    all_by_id = {row["cluster_id"]: row for row in candidates}
    exact_eligible = [row for row in candidates if row["cluster_id"] not in exact_excluded]
    near_eligible, near_removed = remove_near_duplicates_to_history(exact_eligible, exact_excluded, all_by_id)
    if len(near_eligible) < TOTAL_COUNT:
        raise RuntimeError("blocked_by_insufficient_independent_candidates")

    primary = proportional_stratified_sample(near_eligible, PRIMARY_COUNT, SEED)
    primary_ids = {row["cluster_id"] for row in primary}
    primary_profiles = [event_profile(row) for row in primary]
    model_map = model_signal_maps()
    signals = []
    stress_pool = []
    cross_group_near_dupes = []
    for row in near_eligible:
        if row["cluster_id"] in primary_ids:
            continue
        is_dup, detail = row_near_any(row, primary_profiles)
        if is_dup:
            cross_group_near_dupes.append(detail)
            continue
        signal = stress_features(row, score_by_id.get(row["cluster_id"], {}), model_map)
        signals.append(signal)
        if signal["stress_score"] > 0:
            stress_pool.append((row, signal))
    if len(stress_pool) < STRESS_COUNT:
        raise RuntimeError("blocked_by_insufficient_independent_candidates")
    stress_pool.sort(key=lambda pair: (-pair[1]["stress_score"], pair[0]["cluster_id"]))
    selected_stress = [row for row, _signal in stress_pool[:STRESS_COUNT]]
    stress_ids = {row["cluster_id"] for row in selected_stress}
    selected_signals = [signal for _row, signal in stress_pool[:STRESS_COUNT]]
    nonselected_signals = [signal for _row, signal in stress_pool[STRESS_COUNT:]]

    samples = append_sample_metadata(primary, "Primary", 1) + append_sample_metadata(selected_stress, "Stress", PRIMARY_COUNT + 1)
    summary = {
        "original_candidate_cluster_count": len(candidates),
        "historical_exact_exclusion_count": len(exact_excluded & candidate_ids),
        "historical_exact_exclusion_union_count": exclusions["excluded_cluster_id_count"],
        "candidate_count_after_exact_exclusion": len(exact_eligible),
        "near_duplicate_exclusion_count": len(near_removed),
        "candidate_count_after_exact_and_near_duplicate_exclusion": len(near_eligible),
        "primary_count": len(primary),
        "stress_pool_count": len(stress_pool),
        "stress_count": len(selected_stress),
        "cross_group_near_duplicate_rejected_for_stress_count": len(cross_group_near_dupes),
        "primary_stress_cluster_intersection_count": len(primary_ids & stress_ids),
        "new500_duplicate_cluster_count": len(samples) - len({row["cluster_id"] for row in samples}),
        "old47_intersection_count": len(set(exclusions["category_cluster_ids"].get("old_47_dea_gold", [])) & {row["cluster_id"] for row in samples}),
        "old50_intersection_count": len(set(exclusions["category_cluster_ids"].get("old_50_safety_gate_independent", [])) & {row["cluster_id"] for row in samples}),
        "all_history_intersection_count": len(exact_excluded & {row["cluster_id"] for row in samples}),
        "near_duplicate_method": {
            "features": ["target_code/target_name", "event_name", "event date", "title_samples", "claim"],
            "rules": [
                "same target+event+date and title similarity >= 0.86",
                "same target+event within 2 days and title/claim similarity >= 0.90",
                "same target+event and title similarity >= 0.96",
            ],
            "similarity": "difflib.SequenceMatcher over normalized alnum/CJK text",
        },
        "stress_selected_flag_counts": dict(
            Counter(flag for signal in selected_signals for flag, value in signal["stress_flags"].items() if value)
        ),
    }
    hidden_signals = selected_signals + nonselected_signals
    return samples, near_removed, summary, hidden_signals


def build_workbook_rows(
    samples: list[dict[str, Any]],
    candidates_after_exclusion: list[dict[str, Any]],
    exclusion_manifest: dict[str, Any],
    sample_summary: dict[str, Any],
    article_map: dict[str, dict[str, Any]],
    manifest_stub: dict[str, Any],
) -> dict[str, list[dict[str, Any]]]:
    primary = [row for row in samples if row["sample_group"] == "Primary"]
    stress = [row for row in samples if row["sample_group"] == "Stress"]
    dist = distribution_rows(candidates_after_exclusion, primary, stress)
    dist.extend({"set": "summary", "dimension": key, "value": "", "count": value, "ratio": ""} for key, value in sample_summary.items() if isinstance(value, int))
    return {
        "填写说明": instruction_rows(manifest_stub),
        "Primary人工审核": [sample_main_row(row) for row in primary],
        "Stress人工审核": [sample_main_row(row) for row in stress],
        "证据明细": evidence_detail_rows(samples),
        "来源与证据追溯": trace_rows(samples, article_map),
        "抽样分布": dist,
        "历史排除统计": history_stat_rows(exclusion_manifest),
        "枚举值说明": enum_rows(),
    }


def validate_before_write(samples: list[dict[str, Any]], exclusions: dict[str, Any]) -> None:
    ids = [row["cluster_id"] for row in samples]
    if len(samples) != TOTAL_COUNT:
        raise RuntimeError("blocked_by_insufficient_independent_candidates")
    if len(set(ids)) != len(ids):
        raise RuntimeError("blocked_by_data_integrity_failure")
    primary_ids = {row["cluster_id"] for row in samples if row["sample_group"] == "Primary"}
    stress_ids = {row["cluster_id"] for row in samples if row["sample_group"] == "Stress"}
    if len(primary_ids) != PRIMARY_COUNT or len(stress_ids) != STRESS_COUNT or primary_ids & stress_ids:
        raise RuntimeError("blocked_by_data_integrity_failure")
    if set(ids) & set(exclusions["excluded_cluster_ids"]):
        raise RuntimeError("blocked_by_data_integrity_failure")
    for sample in samples:
        items = sample.get("evidence_items", [])
        if not items or not all(item.get("traceable") for item in items):
            raise RuntimeError("blocked_by_evidence_traceability_failure")
        evidence_ids = [item["evidence_id"] for item in items]
        if len(evidence_ids) != len(set(evidence_ids)):
            raise RuntimeError("blocked_by_data_integrity_failure")


def build_package(args: argparse.Namespace) -> dict[str, Any]:
    output_dir = args.output_dir
    workbook = args.workbook
    guide = args.guide
    manifest_path = args.manifest
    backup_dir = prepare_backup_dir(args.backup_dir)
    ensure_no_overwrite(
        [
            output_dir / "independent_samples.jsonl",
            output_dir / "primary_samples.jsonl",
            output_dir / "stress_samples.jsonl",
            output_dir / "exclusion_manifest.json",
            output_dir / "evaluation_stage_signals.jsonl",
            workbook,
            guide,
            manifest_path,
        ],
        args.overwrite,
    )
    copy_protected_files_once(backup_dir)
    protected_before = hash_paths(PROTECTED_FORMAL_OUTPUTS + FROZEN_FILES)
    write_json(backup_dir / "protected_sha256_before_step6_1.json", protected_before)

    candidates, candidate_metadata, score_by_id = load_candidate_rows(args.base_scores, args.fulltext_scores, args.evidence, args.article_index)
    article_map = load_article_map(args.article_index)
    exclusions = collect_exclusions()
    samples, near_removed, sample_summary, hidden_signals = select_samples(candidates, score_by_id, exclusions)
    validate_before_write(samples, exclusions)
    sample_ids = {row["cluster_id"] for row in samples}
    candidates_after_exclusion = [
        row
        for row in candidates
        if row["cluster_id"] not in set(exclusions["excluded_cluster_ids"])
        and row["cluster_id"] not in {item["cluster_id"] for item in near_removed}
    ]
    evidence_count = sum(len(row.get("evidence_items", [])) for row in samples)
    untraceable_count = sum(1 for row in samples for item in row.get("evidence_items", []) if not item.get("traceable"))
    protected_after = hash_paths(PROTECTED_FORMAL_OUTPUTS + FROZEN_FILES)

    exclusion_manifest = {
        **exclusions,
        "near_duplicate_exclusion_count": len(near_removed),
        "near_duplicate_excluded_cluster_ids": sorted(item["cluster_id"] for item in near_removed),
        "near_duplicate_excluded_details": near_removed,
        "remaining_candidate_count_after_exclusions": sample_summary["candidate_count_after_exact_and_near_duplicate_exclusion"],
        "cluster_id_intersection_check": {
            "new500_vs_old47": sample_summary["old47_intersection_count"],
            "new500_vs_old50": sample_summary["old50_intersection_count"],
            "new500_vs_all_history": sample_summary["all_history_intersection_count"],
        },
        "near_duplicate_check_method": sample_summary["near_duplicate_method"],
    }

    manifest = {
        "status": "task2_independent500_ready_for_human_review",
        "generated_at": now_iso(),
        "step": "Task2 Step 6.1 independent500_v1",
        "project_root": str(ROOT),
        "backup_dir": str(backup_dir),
        "sample_seed": SEED,
        "sample_counts": {"Primary": PRIMARY_COUNT, "Stress": STRESS_COUNT, "total": TOTAL_COUNT},
        "input_paths": {
            "base_scores": rel(args.base_scores),
            "formal_fulltext_scores": rel(args.fulltext_scores),
            "fulltext_evidence": rel(args.evidence),
            "article_index": rel(args.article_index),
        },
        "candidate_metadata": candidate_metadata,
        "exclusion_summary": {
            "historical_exact_union_count": exclusions["excluded_cluster_id_count"],
            "historical_exact_in_candidate_count": sample_summary["historical_exact_exclusion_count"],
            "near_duplicate_exclusion_count": len(near_removed),
            "remaining_candidate_count": sample_summary["candidate_count_after_exact_and_near_duplicate_exclusion"],
            "category_cluster_id_counts": exclusions["category_cluster_id_counts"],
        },
        "sampling_summary": sample_summary,
        "intersection_checks": {
            "primary_stress_intersection_count": sample_summary["primary_stress_cluster_intersection_count"],
            "new500_duplicate_cluster_count": sample_summary["new500_duplicate_cluster_count"],
            "new500_vs_old47": sample_summary["old47_intersection_count"],
            "new500_vs_old50": sample_summary["old50_intersection_count"],
            "new500_vs_all_history": sample_summary["all_history_intersection_count"],
        },
        "evidence_traceability": {
            "selected_evidence_count": evidence_count,
            "selected_untraceable_evidence_count": untraceable_count,
            "trace_path_template": f"{rel(args.article_index)}#article_id=<article_id>",
            "cross_cluster_reference_count": 0,
        },
        "human_fields": HUMAN_FIELDS,
        "allowed_values": ALLOWED_VALUES,
        "human_fields_blank_on_generation": True,
        "ai_assistance": {
            "generated": False,
            "model_calls": 0,
            "reason": "blank_human_audit_package_only",
        },
        "protected_sha256_before": protected_before,
        "protected_sha256_after": protected_after,
        "protected_sha256_unchanged": protected_before == protected_after,
        "outputs": {
            "output_dir": rel(output_dir),
            "workbook": rel(workbook),
            "guide": rel(guide),
            "manifest": rel(manifest_path),
            "exclusion_manifest": rel(output_dir / "exclusion_manifest.json"),
            "independent_samples": rel(output_dir / "independent_samples.jsonl"),
            "primary_samples": rel(output_dir / "primary_samples.jsonl"),
            "stress_samples": rel(output_dir / "stress_samples.jsonl"),
            "evaluation_stage_signals": rel(output_dir / "evaluation_stage_signals.jsonl"),
        },
    }
    manifest["manifest_path"] = rel(manifest_path)
    workbook_rows = build_workbook_rows(samples, candidates_after_exclusion, exclusion_manifest, sample_summary, article_map, manifest)

    output_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl(output_dir / "independent_samples.jsonl", samples)
    write_jsonl(output_dir / "primary_samples.jsonl", [row for row in samples if row["sample_group"] == "Primary"])
    write_jsonl(output_dir / "stress_samples.jsonl", [row for row in samples if row["sample_group"] == "Stress"])
    write_jsonl(output_dir / "evaluation_stage_signals.jsonl", hidden_signals)
    write_json(output_dir / "exclusion_manifest.json", exclusion_manifest)
    write_workbook(workbook, workbook_rows)
    manifest["outputs"]["workbook_sha256"] = sha256(workbook)
    guide.write_text(render_guide(manifest), encoding="utf-8")
    manifest["outputs"]["guide_sha256"] = sha256(guide)
    write_json(manifest_path, manifest)
    write_json(backup_dir / "step6_1_manifest_snapshot.json", manifest)
    return manifest


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-scores", type=Path, default=BASE_SCORES)
    parser.add_argument("--fulltext-scores", type=Path, default=FORMAL_FULLTEXT_SCORES)
    parser.add_argument("--evidence", type=Path, default=FULLTEXT_EVIDENCE)
    parser.add_argument("--article-index", type=Path, default=ARTICLE_INDEX)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--guide", type=Path, default=DEFAULT_GUIDE)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--backup-dir", type=Path, default=None)
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    args.base_scores = args.base_scores if args.base_scores.is_absolute() else ROOT / args.base_scores
    args.fulltext_scores = args.fulltext_scores if args.fulltext_scores.is_absolute() else ROOT / args.fulltext_scores
    args.evidence = args.evidence if args.evidence.is_absolute() else ROOT / args.evidence
    args.article_index = args.article_index if args.article_index.is_absolute() else ROOT / args.article_index
    args.output_dir = args.output_dir if args.output_dir.is_absolute() else ROOT / args.output_dir
    args.workbook = args.workbook if args.workbook.is_absolute() else ROOT / args.workbook
    args.guide = args.guide if args.guide.is_absolute() else ROOT / args.guide
    args.manifest = args.manifest if args.manifest.is_absolute() else ROOT / args.manifest
    if args.backup_dir is not None and not args.backup_dir.is_absolute():
        args.backup_dir = ROOT / args.backup_dir
    manifest = build_package(args)
    print(
        json.dumps(
            {
                "status": manifest["status"],
                "sample_counts": manifest["sample_counts"],
                "candidate_metadata": manifest["candidate_metadata"],
                "exclusion_summary": manifest["exclusion_summary"],
                "intersection_checks": manifest["intersection_checks"],
                "evidence_traceability": manifest["evidence_traceability"],
                "protected_sha256_unchanged": manifest["protected_sha256_unchanged"],
                "outputs": manifest["outputs"],
                "backup_dir": manifest["backup_dir"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
