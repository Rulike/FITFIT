#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import math
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
RAW_XLSX = ROOT / "data/raw/news_with_tags_test.xlsx"
INFER_DIR = ROOT / "outputs/task3_kg_v4_official1799_inference_v1"
STRUCTURED_EVENTS = INFER_DIR / "official1799_structured_events.jsonl"
FUNNEL_SOURCE = INFER_DIR / "official1799_selection_funnel.json"
OUT = ROOT / "outputs/task34_official1799_full_benchmark_v1"
SEED = 20260825

ALLOWED_INPUT_COLUMNS = [
    "article_file_id",
    "article_title",
    "article_publish_time",
    "article_source",
    "content",
]
FORBIDDEN_INPUT_COLUMNS = {
    "trading_code",
    "secu_abbr",
    "industry_code",
    "industry_name",
    "event_name",
    "event_emotion",
    "event_impact_analysis",
    "duplication_id",
}
FORBIDDEN_ANNOTATION_COLUMNS = {
    "trading_code",
    "secu_abbr",
    "industry_code",
    "industry_name",
    "event_name",
    "official_event",
    "official_event_name",
    "official_target",
    "official_target_name",
    "event_emotion",
    "event_impact_analysis",
    "duplication_id",
    "resolver_result",
    "resolver_decision",
    "resolver_confidence",
    "model_decision",
    "model_score",
    "model_probability",
    "probability",
    "probabilities",
    "candidate_bucket",
    "rule_hits",
    "route",
    "abstain",
    "suppress",
    "severity_prediction",
}

TASK3_HUMAN_FIELDS = [
    "structured_event_valid",
    "target_grounded",
    "event_grounded",
    "evidence_support",
    "human_canonical_event_id",
    "human_event_status",
    "is_singleton",
    "matching_article_ids",
    "annotation_confidence",
    "annotation_reason",
]
TASK4_HUMAN_FIELDS = [
    "task4_applicable",
    "target_grounded",
    "event_grounded",
    "evidence_support",
    "validity",
    "direction_gold",
    "risk_intensity",
    "severity_label",
    "rule_applicability",
    "financial_evidence",
    "alert_gold",
    "annotation_confidence",
    "annotation_basis",
    "invalid_reason",
]
TASK3_ALLOWED = {
    "structured_event_valid": ["yes", "no", "uncertain"],
    "target_grounded": ["yes", "no", "uncertain"],
    "event_grounded": ["yes", "no", "uncertain"],
    "evidence_support": ["yes", "partial", "no"],
    "human_event_status": ["fact", "rumor", "forecast", "uncertain"],
    "is_singleton": ["yes", "no", "uncertain"],
    "annotation_confidence": ["high", "medium", "low"],
}
TASK4_ALLOWED = {
    "task4_applicable": ["yes", "no", "review"],
    "target_grounded": ["yes", "no", "uncertain"],
    "event_grounded": ["yes", "no", "uncertain"],
    "evidence_support": ["yes", "partial", "no"],
    "validity": ["valid", "invalid", "uncertain"],
    "direction_gold": ["positive", "negative", "neutral", "mixed", "uncertain"],
    "risk_intensity": ["1", "2", "3", "4", "5"],
    "severity_label": ["blue", "yellow", "orange", "red", "uncertain"],
    "rule_applicability": ["applicable", "not_applicable", "uncertain"],
    "financial_evidence": ["yes", "no", "uncertain"],
    "alert_gold": ["alert", "no_alert", "review"],
    "annotation_confidence": ["high", "medium", "low"],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HUMAN_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2F0D9")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null", "nat"} else text


def compact(value: Any) -> str:
    return "".join(clean(value).split())


def stable_float(*parts: Any) -> float:
    payload = "||".join([str(SEED), *[clean(p) for p in parts]])
    digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]
    return int(digest, 16) / float(16**16 - 1)


def json_ready(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {clean(k): json_ready(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [json_ready(v) for v in obj]
    if isinstance(obj, pd.Timestamp):
        return obj.isoformat()
    if hasattr(obj, "item"):
        try:
            return json_ready(obj.item())
        except Exception:
            pass
    try:
        if pd.isna(obj):
            return None
    except Exception:
        pass
    if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
        return None
    return obj


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(json_ready(payload), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def inspect_headers(path: Path) -> dict[str, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    headers: dict[str, list[str]] = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        headers[sheet] = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def load_allowed_raw() -> tuple[pd.DataFrame, dict[str, Any]]:
    headers = inspect_headers(RAW_XLSX)
    rows: list[dict[str, Any]] = []
    sheet_counts: dict[str, int] = {}
    forbidden_present: dict[str, list[str]] = {}
    for sheet, cols in headers.items():
        missing = [col for col in ALLOWED_INPUT_COLUMNS if col not in cols]
        if missing:
            raise RuntimeError(f"missing allowed columns in {sheet}: {missing}")
        forbidden_present[sheet] = [col for col in cols if col in FORBIDDEN_INPUT_COLUMNS]
        frame = pd.read_excel(RAW_XLSX, sheet_name=sheet, usecols=ALLOWED_INPUT_COLUMNS, dtype=object).fillna("")
        sheet_counts[sheet] = int(len(frame))
        for idx, row in frame.iterrows():
            rows.append(
                {
                    "official_row_id": f"OFF1799_{len(rows) + 1:04d}",
                    "sheet_name": sheet,
                    "sheet_row_index": int(idx) + 2,
                    "article_id": clean(row["article_file_id"]),
                    "title": clean(row["article_title"]),
                    "source": clean(row["article_source"]),
                    "publish_time": clean(row["article_publish_time"]),
                    "content_hash": hashlib.sha256(compact(row["content"]).encode("utf-8")).hexdigest(),
                }
            )
    raw = pd.DataFrame(rows)
    manifest = {
        "source_file": str(RAW_XLSX),
        "source_sha256": sha256_file(RAW_XLSX),
        "allowed_columns_read": ALLOWED_INPUT_COLUMNS,
        "forbidden_columns_present_but_not_read": forbidden_present,
        "forbidden_columns_read": [],
        "official_labels_read": False,
        "read_method": "pd.read_excel(..., usecols=ALLOWED_INPUT_COLUMNS)",
        "sheet_counts": sheet_counts,
        "input_rows": int(len(raw)),
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
    }
    return raw, manifest


def visible_event(row: dict[str, Any], raw_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    official_row_id = clean(row.get("official_row_id"))
    raw = raw_by_id.get(official_row_id, {})
    target = row.get("target") or {}
    event = row.get("event") or {}
    evidence = row.get("evidence") or {}
    provenance = row.get("provenance") or {}
    source_row = row.get("source_row") or {}
    offset = {
        "trigger": event.get("trigger_span") or {},
        "event_time": event.get("event_time_span") or {},
        "evidence": evidence.get("evidence_span") or {},
        "local_context": evidence.get("local_context_span") or {},
    }
    prov = {
        "official_row_id": official_row_id,
        "sample_id": source_row.get("sample_id"),
        "document_hash": provenance.get("document_hash"),
        "sentence_id": provenance.get("sentence_id"),
        "source_file": provenance.get("source_file"),
    }
    return {
        "annotation_sample_id": "FULL_" + official_row_id,
        "official_row_id": official_row_id,
        "article_id": clean(raw.get("article_id") or row.get("article_id")),
        "title": clean(raw.get("title")),
        "source": clean(raw.get("source") or (row.get("source") or {}).get("source_name")),
        "publish_time": clean(raw.get("publish_time") or row.get("publish_time")),
        "predicted_target": clean(target.get("target_name")),
        "event_type": clean(event.get("event_type")),
        "event_fact": clean(event.get("event_fact")),
        "trigger": clean(event.get("trigger_text")),
        "event_time": clean(event.get("event_time")),
        "arguments": json.dumps(json_ready(event.get("arguments") or []), ensure_ascii=False),
        "evidence": clean(evidence.get("evidence_sentence")),
        "provenance": json.dumps(json_ready(prov), ensure_ascii=False),
        "offset": json.dumps(json_ready(offset), ensure_ascii=False),
    }


def build_task3_frame(structured: list[dict[str, Any]], raw: pd.DataFrame) -> pd.DataFrame:
    raw_by_id = raw.set_index("official_row_id").to_dict(orient="index")
    rows = [visible_event(row, raw_by_id) for row in structured]
    out = pd.DataFrame(rows)
    cols = [
        "annotation_sample_id",
        "official_row_id",
        "article_id",
        "title",
        "source",
        "publish_time",
        "predicted_target",
        "event_type",
        "event_fact",
        "trigger",
        "event_time",
        "arguments",
        "evidence",
        "offset",
        "provenance",
    ]
    out = out[cols]
    for field in TASK3_HUMAN_FIELDS:
        out[field] = ""
    return out


def build_task4_frame(structured: list[dict[str, Any]], raw: pd.DataFrame) -> pd.DataFrame:
    raw_by_id = raw.set_index("official_row_id").to_dict(orient="index")
    rows = [visible_event(row, raw_by_id) for row in structured]
    out = pd.DataFrame(rows)
    cols = [
        "annotation_sample_id",
        "official_row_id",
        "article_id",
        "title",
        "source",
        "publish_time",
        "predicted_target",
        "event_type",
        "event_fact",
        "trigger",
        "event_time",
        "arguments",
        "evidence",
        "offset",
        "provenance",
    ]
    out = out[cols]
    for field in TASK4_HUMAN_FIELDS:
        out[field] = ""
    return out


def shuffled(frame: pd.DataFrame, version: str) -> pd.DataFrame:
    out = frame.copy()
    out["_sort_key"] = out["annotation_sample_id"].map(lambda x: stable_float(version, x))
    out = out.sort_values("_sort_key").drop(columns=["_sort_key"]).reset_index(drop=True)
    out.insert(0, "annotation_order", range(1, len(out) + 1))
    return out


def add_dropdowns(ws, allowed: dict[str, list[str]], row_count: int) -> None:
    headers = [clean(cell.value) for cell in ws[1]]
    for name, values in allowed.items():
        if name not in headers:
            continue
        col = headers.index(name) + 1
        letter = get_column_letter(col)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{row_count + 1}")


def write_workbook(path: Path, frame: pd.DataFrame, task: str) -> None:
    wb = Workbook()
    inst = wb.active
    inst.title = "instructions"
    if task == "task3":
        instructions = [
            ["字段", "说明"],
            ["任务", "Task3 官方1799完整下游双审Benchmark。每行是一篇官方新闻的冻结结构化事件。"],
            ["核心规则", "属于同一个现实事件的多篇报道填写相同 human_canonical_event_id；不同时间、金额、项目或对象的事件填写不同ID；无有效事件时 canonical ID 留空。"],
            ["A/B ID规则", "A/B 允许使用不同自定义ID，后续按事件成员集合对齐，并自动推导 candidate pair same/different、coverage gap 和 cluster 指标。"],
            ["合成示例", "同一公司、同一天、相同产品和订单事实，仅报道来源不同 -> same_event / 同一 human_canonical_event_id"],
        ]
        human_fields = TASK3_HUMAN_FIELDS
        allowed = TASK3_ALLOWED
    else:
        instructions = [
            ["字段", "说明"],
            ["任务", "Task4 官方1799完整下游双审Benchmark。所有1799行都必须填写 task4_applicable。"],
            ["核心规则", "只有 applicable=yes 且事件证据成立的样本填写 risk_intensity；不适用、无事件、主体错误或证据不足时风险分留空，并填写 invalid_reason；不得为了完整率强行打1分或blue。"],
            ["合成示例", "公司公告确认大额债务只能偿付40%，财务指标显示偿债压力 -> validity=valid; direction=negative; risk_intensity=4; alert_gold=alert"],
        ]
        human_fields = TASK4_HUMAN_FIELDS
        allowed = TASK4_ALLOWED
    for row in instructions:
        inst.append(row)
    for cell in inst[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    inst.column_dimensions["A"].width = 24
    inst.column_dimensions["B"].width = 120
    for row in inst.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            if clean(cell.value).startswith("同一公司") or clean(cell.value).startswith("公司公告"):
                cell.fill = EXAMPLE_FILL

    ws = wb.create_sheet("annotation")
    ws.append(list(frame.columns))
    for _, row in frame.iterrows():
        ws.append([row[col] for col in frame.columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    human_set = set(human_fields)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = TOP
            if clean(ws.cell(1, cell.column).value) in human_set:
                cell.fill = HUMAN_FILL
    widths = {
        "annotation_order": 10,
        "annotation_sample_id": 20,
        "official_row_id": 14,
        "article_id": 32,
        "title": 44,
        "source": 18,
        "publish_time": 22,
        "predicted_target": 20,
        "event_type": 42,
        "event_fact": 80,
        "trigger": 32,
        "event_time": 18,
        "arguments": 44,
        "evidence": 80,
        "offset": 44,
        "provenance": 52,
        "matching_article_ids": 44,
        "annotation_reason": 54,
        "annotation_basis": 54,
        "invalid_reason": 44,
    }
    for idx, cell in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(clean(cell.value), 18)
    add_dropdowns(ws, allowed, len(frame))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def human_empty(frame: pd.DataFrame, fields: list[str]) -> bool:
    return all(frame[field].map(clean).eq("").all() for field in fields)


def verify(task3_a: pd.DataFrame, task3_b: pd.DataFrame, task4_a: pd.DataFrame, task4_b: pd.DataFrame) -> dict[str, Any]:
    columns = set(task3_a.columns) | set(task4_a.columns)
    leaked = sorted(col for col in columns if col in FORBIDDEN_ANNOTATION_COLUMNS)
    return {
        "task3_A_rows": int(len(task3_a)),
        "task3_B_rows": int(len(task3_b)),
        "task4_A_rows": int(len(task4_a)),
        "task4_B_rows": int(len(task4_b)),
        "all_files_1799_rows": all(len(df) == 1799 for df in [task3_a, task3_b, task4_a, task4_b]),
        "task3_ab_same_sample_set": set(task3_a["annotation_sample_id"]) == set(task3_b["annotation_sample_id"]),
        "task3_ab_different_order": list(task3_a["annotation_sample_id"]) != list(task3_b["annotation_sample_id"]),
        "task4_ab_same_sample_set": set(task4_a["annotation_sample_id"]) == set(task4_b["annotation_sample_id"]),
        "task4_ab_different_order": list(task4_a["annotation_sample_id"]) != list(task4_b["annotation_sample_id"]),
        "task3_human_fields_empty": human_empty(task3_a, TASK3_HUMAN_FIELDS) and human_empty(task3_b, TASK3_HUMAN_FIELDS),
        "task4_human_fields_empty": human_empty(task4_a, TASK4_HUMAN_FIELDS) and human_empty(task4_b, TASK4_HUMAN_FIELDS),
        "forbidden_annotation_columns_present": leaked,
        "label_or_model_answer_leakage": bool(leaked),
        "estimated_human_annotation_assignments": 1799 * 4,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    raw, input_manifest = load_allowed_raw()
    structured = read_jsonl(STRUCTURED_EVENTS)
    if len(raw) != 1799:
        raise RuntimeError(f"expected raw 1799 rows, got {len(raw)}")
    if len(structured) != 1799:
        raise RuntimeError(f"expected structured 1799 rows, got {len(structured)}")

    task3 = build_task3_frame(structured, raw)
    task4 = build_task4_frame(structured, raw)
    task3_a = shuffled(task3, "task3_A")
    task3_b = shuffled(task3, "task3_B")
    task4_a = shuffled(task4, "task4_A")
    task4_b = shuffled(task4, "task4_B")

    paths = {
        "task3_A": OUT / "task3_official1799_full_A.xlsx",
        "task3_B": OUT / "task3_official1799_full_B.xlsx",
        "task4_A": OUT / "task4_official1799_full_A.xlsx",
        "task4_B": OUT / "task4_official1799_full_B.xlsx",
    }
    write_workbook(paths["task3_A"], task3_a, "task3")
    write_workbook(paths["task3_B"], task3_b, "task3")
    write_workbook(paths["task4_A"], task4_a, "task4")
    write_workbook(paths["task4_B"], task4_b, "task4")

    task3_a.to_parquet(OUT / "task3_official1799_full_A.parquet", index=False)
    task3_b.to_parquet(OUT / "task3_official1799_full_B.parquet", index=False)
    task4_a.to_parquet(OUT / "task4_official1799_full_A.parquet", index=False)
    task4_b.to_parquet(OUT / "task4_official1799_full_B.parquet", index=False)

    validation = verify(task3_a, task3_b, task4_a, task4_b)
    if not all(
        [
            validation["all_files_1799_rows"],
            validation["task3_ab_same_sample_set"],
            validation["task3_ab_different_order"],
            validation["task4_ab_same_sample_set"],
            validation["task4_ab_different_order"],
            validation["task3_human_fields_empty"],
            validation["task4_human_fields_empty"],
            not validation["label_or_model_answer_leakage"],
        ]
    ):
        raise RuntimeError(f"validation failed: {validation}")

    source_funnel = read_json(FUNNEL_SOURCE) if FUNNEL_SOURCE.exists() else {}
    quality_counts = Counter(clean((row.get("quality") or {}).get("contract_status")) for row in structured)
    target_grounded = sum(bool((row.get("quality") or {}).get("target_grounded")) for row in structured)
    event_grounded = sum(bool((row.get("quality") or {}).get("event_grounded")) for row in structured)
    evidence_grounded = sum(bool((row.get("quality") or {}).get("evidence_grounded")) for row in structured)
    funnel = {
        "input_news_rows": int(len(raw)),
        "frozen_structured_event_rows": int(len(structured)),
        "task1_contract_status_counts": dict(quality_counts),
        "task1_contract_pass": int(quality_counts.get("pass", 0)),
        "task1_target_grounded": int(target_grounded),
        "task1_event_grounded": int(event_grounded),
        "task1_evidence_grounded": int(evidence_grounded),
        "task1_grounding_pass": int(sum(bool((row.get("quality") or {}).get("grounding_pass")) for row in structured)),
        "task3_full_annotation_rows_per_reviewer": int(len(task3_a)),
        "task3_candidate_pair_derivation_planned": "derive 3377 candidate-pair labels from human canonical member sets after A/B completion",
        "task3_candidate_pairs_from_frozen_funnel": int((source_funnel.get("resolver_v2") or {}).get("candidate_pairs") or 3377),
        "task4_full_annotation_rows_per_reviewer": int(len(task4_a)),
        "estimated_human_annotation_assignments": 1799 * 4,
        "training_performed": False,
        "tuning_performed": False,
        "official_labels_read": False,
        "task34_gold_or_severity_gold_read": False,
        "resolver_pair_decisions_read": False,
        "strict_scoring_or_model_scores_read": False,
        "frozen_files_modified": False,
    }
    sampling_meta = {
        "seed": SEED,
        "scheme": "full_1799_no_sampling",
        "task3_annotation_sample_ids": sorted(task3["annotation_sample_id"].tolist()),
        "task4_annotation_sample_ids": sorted(task4["annotation_sample_id"].tolist()),
        "task3_A_order_sha256": hashlib.sha256("|".join(task3_a["annotation_sample_id"]).encode("utf-8")).hexdigest(),
        "task3_B_order_sha256": hashlib.sha256("|".join(task3_b["annotation_sample_id"]).encode("utf-8")).hexdigest(),
        "task4_A_order_sha256": hashlib.sha256("|".join(task4_a["annotation_sample_id"]).encode("utf-8")).hexdigest(),
        "task4_B_order_sha256": hashlib.sha256("|".join(task4_b["annotation_sample_id"]).encode("utf-8")).hexdigest(),
    }
    guide = "\n".join(
        [
            "# Task34 Official1799 Full Benchmark v1",
            "",
            "This is a full downstream double-review package over all 1799 official news rows.",
            "",
            "Task3: assign human_canonical_event_id by real-world event membership. Same company/day/product/order fact should share an ID; different time, amount, project, or object should use different IDs.",
            "Task3 synthetic example: 同一公司、同一天、相同产品和订单事实，仅报道来源不同 -> same_event.",
            "",
            "Task4: fill task4_applicable for every row. Fill risk_intensity only when applicable=yes and target/event/evidence are valid.",
            "Task4 synthetic example: 公司公告确认大额债务只能偿付40%，财务指标显示偿债压力 -> validity=valid; direction=negative; risk_intensity=4; alert_gold=alert.",
            "",
            "The annotation files do not expose official tags, resolver decisions, model scores, probabilities, route/rule hits, candidate buckets, or severity predictions.",
        ]
    )
    write_json(OUT / "input_manifest.json", input_manifest)
    write_json(OUT / "sampling_meta.json", sampling_meta)
    write_json(OUT / "funnel.json", funnel)
    write_json(OUT / "validation_report.json", validation)
    (OUT / "annotation_guide.md").write_text(guide + "\n", encoding="utf-8")

    print(json.dumps({"paths": {k: str(v) for k, v in paths.items()}, "funnel": funnel, "validation": validation}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
