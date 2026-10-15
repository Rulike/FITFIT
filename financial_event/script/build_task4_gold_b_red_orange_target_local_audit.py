#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "outputs" / "task4 golddata" / "task4_severity_training_data_v1.xlsx"
SOURCE_SHEET = "Gold_B_Legacy"
OUT_DIR = ROOT / "outputs" / "task4_gold_b_red_orange_target_local_audit"
OUT_XLSX = OUT_DIR / "task4_gold_b_red_orange_target_local_audit.xlsx"
SUMMARY_JSON = OUT_DIR / "summary.json"

EXPECTED_TOTAL = 72
EXPECTED_RED = 17
EXPECTED_ORANGE = 55

ORIGINAL_COLUMNS = [
    "sample_id",
    "target_name",
    "target_type",
    "event_name",
    "event_type",
    "title",
    "evidence_sentence",
    "local_context",
    "financial_facts",
    "publish_time",
    "severity_label",
    "risk_intensity",
    "annotation_confidence",
    "annotation_basis",
    "label_tier",
    "label_source",
    "target_local_status",
    "split_status",
    "provenance",
    "provenance_status",
    "supervision_eligible",
]

HUMAN_COLUMNS = [
    "human_target_grounded",
    "human_event_grounded",
    "human_evidence_grounded",
    "human_severity_judgment",
    "human_corrected_severity",
    "human_corrected_risk_intensity",
    "human_invalid_reason",
    "human_comment",
    "reviewer",
    "review_time",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HUMAN_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return str(value).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        try:
            value = value.item()
        except Exception:
            pass
    return value


def add_dropdown(ws, column_name: str, values: list[str], end_row: int) -> None:
    if not values:
        return
    headers = [clean_text(cell.value) for cell in ws[1]]
    if column_name not in headers:
        return
    col_idx = headers.index(column_name) + 1
    col_letter = get_column_letter(col_idx)
    formula = '"' + ",".join(values) + '"'
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{end_row}")


def set_widths(ws, widths: dict[str, int], default: int = 18) -> None:
    headers = [clean_text(cell.value) for cell in ws[1]]
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(header, default)


def style_annotation_sheet(ws, row_count: int) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    human_headers = set(HUMAN_COLUMNS)
    for row in ws.iter_rows(min_row=2, max_row=row_count + 1):
        for cell in row:
            cell.alignment = TOP
            header = clean_text(ws.cell(1, cell.column).value)
            if header in human_headers:
                cell.fill = HUMAN_FILL

    widths = {
        "sample_id": 30,
        "target_name": 18,
        "target_type": 14,
        "event_name": 24,
        "event_type": 18,
        "title": 42,
        "evidence_sentence": 90,
        "local_context": 100,
        "financial_facts": 80,
        "publish_time": 20,
        "severity_label": 12,
        "risk_intensity": 12,
        "annotation_confidence": 16,
        "annotation_basis": 18,
        "label_tier": 12,
        "label_source": 14,
        "target_local_status": 16,
        "split_status": 12,
        "provenance": 60,
        "provenance_status": 16,
        "supervision_eligible": 14,
        "human_target_grounded": 18,
        "human_event_grounded": 18,
        "human_evidence_grounded": 20,
        "human_severity_judgment": 20,
        "human_corrected_severity": 20,
        "human_corrected_risk_intensity": 24,
        "human_invalid_reason": 24,
        "human_comment": 36,
        "reviewer": 16,
        "review_time": 20,
    }
    set_widths(ws, widths, default=18)

    add_dropdown(ws, "human_target_grounded", ["yes", "no"], row_count + 1)
    add_dropdown(ws, "human_event_grounded", ["yes", "no"], row_count + 1)
    add_dropdown(ws, "human_evidence_grounded", ["yes", "no"], row_count + 1)
    add_dropdown(ws, "human_severity_judgment", ["correct", "revise", "invalid", "uncertain"], row_count + 1)
    add_dropdown(ws, "human_corrected_severity", ["red", "orange", "yellow", "blue", "review"], row_count + 1)
    add_dropdown(ws, "human_corrected_risk_intensity", ["1", "2", "3", "4", "5"], row_count + 1)
    add_dropdown(
        ws,
        "human_invalid_reason",
        ["wrong_target", "wrong_event", "insufficient_evidence", "article_global_noise", "duplicate", "other"],
        row_count + 1,
    )


def write_instructions_sheet(wb: Workbook, total: int, red: int, orange: int) -> None:
    ws = wb.active
    ws.title = "instructions"
    rows = [
        ["字段", "说明"],
        ["任务", "基于 task4_severity_training_data_v1.xlsx 的 Gold_B_Legacy 生成人工复核包。"],
        ["筛选条件", "只保留 severity_label = red 或 orange 的样本。"],
        ["数量检查", f"预期总数 {total} 条，其中 red {red} 条、orange {orange} 条。"],
        ["编辑原则", "只填写右侧浅黄色 human_* 列，原有 21 个字段保持不变，不自动修改任何人工标签。"],
        ["human_target_grounded", "yes / no"],
        ["human_event_grounded", "yes / no"],
        ["human_evidence_grounded", "yes / no"],
        ["human_severity_judgment", "correct / revise / invalid / uncertain"],
        ["human_corrected_severity", "red / orange / yellow / blue / review；无需修改时留空"],
        ["human_corrected_risk_intensity", "1-5；无需修改时留空"],
        ["human_invalid_reason", "wrong_target / wrong_event / insufficient_evidence / article_global_noise / duplicate / other"],
        ["human_comment", "自由文本"],
        ["reviewer", "自由文本"],
        ["review_time", "自由文本或时间字符串"],
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"字段": 28, "说明": 110}
    set_widths(ws, widths, default=24)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def main() -> int:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"Missing input workbook: {INPUT_XLSX}")

    input_sha256 = sha256_file(INPUT_XLSX)

    df = pd.read_excel(INPUT_XLSX, sheet_name=SOURCE_SHEET, dtype=object)
    missing_cols = [col for col in ORIGINAL_COLUMNS if col not in df.columns]
    if missing_cols:
        raise RuntimeError(f"Missing expected columns in {SOURCE_SHEET}: {missing_cols}")

    df = df.loc[:, ORIGINAL_COLUMNS].copy()
    df["severity_label"] = df["severity_label"].map(clean_text)
    selected = df[df["severity_label"].isin(["red", "orange"])].copy()
    selected = selected.reset_index(drop=True)

    counts = selected["severity_label"].value_counts().to_dict()
    red_count = int(counts.get("red", 0))
    orange_count = int(counts.get("orange", 0))
    total_count = int(len(selected))
    duplicate_sample_id = bool(selected["sample_id"].astype(str).duplicated().any())

    if total_count != EXPECTED_TOTAL or red_count != EXPECTED_RED or orange_count != EXPECTED_ORANGE:
        raise RuntimeError(
            "Filtered counts mismatch: "
            f"total={total_count}, red={red_count}, orange={orange_count}; "
            f"expected total={EXPECTED_TOTAL}, red={EXPECTED_RED}, orange={EXPECTED_ORANGE}"
        )
    if duplicate_sample_id:
        raise RuntimeError("Filtered sample_id contains duplicates")

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    write_instructions_sheet(wb, EXPECTED_TOTAL, EXPECTED_RED, EXPECTED_ORANGE)
    ws = wb.create_sheet("annotation")

    headers = ORIGINAL_COLUMNS + HUMAN_COLUMNS
    ws.append(headers)
    for _, row in selected.iterrows():
        values = [normalize_cell(row[col]) for col in ORIGINAL_COLUMNS] + [None] * len(HUMAN_COLUMNS)
        ws.append(values)

    style_annotation_sheet(ws, total_count)
    wb.save(OUT_XLSX)

    summary = {
        "input_file": str(INPUT_XLSX),
        "input_file_sha256": input_sha256,
        "source_sheet": SOURCE_SHEET,
        "total_samples": total_count,
        "red_count": red_count,
        "orange_count": orange_count,
        "sample_id_duplicate": duplicate_sample_id,
        "original_file_modified": False,
        "output_file": str(OUT_XLSX),
    }
    SUMMARY_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
