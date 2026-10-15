#!/usr/bin/env python3
"""Build Task2 independent500-v2 blind human annotation package.

The sampler uses only source/event/target/time metadata and raw text features.
It does not train, tune, run evaluation, read current model predictions, or use
old gold labels for sampling.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone, timedelta
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

SEED = 20260826
TOTAL_COUNT = 500
OUTPUT_DIR = ROOT / "outputs/event_credibility/task2_independent500_v2"
LOCKED_IDS_CONFIG = ROOT / "configs/task2_independent500_v2_locked_ids.json"
WORKBOOK = ROOT / "reports/task2_independent500_v2_human_annotation.xlsx"
GUIDE = ROOT / "reports/task2_independent500_v2_annotation_guide.md"
ALL_EVENTS = ROOT / "data/processed/event_clusters/all_final_events_with_sources.jsonl"
ARTICLE_INDEX = ROOT / "data/processed/event_credibility/fulltext_article_index.parquet"

HUMAN_FIELDS = [
    "human_label",
    "human_confidence",
    "human_reason",
    "selected_evidence_ids",
    "needs_second_review",
    "adjudicated_label",
]
LABEL_VALUES = ["support", "contradict", "insufficient", "unclear"]
CONFIDENCE_VALUES = ["high", "medium", "low"]
REVIEW_VALUES = ["yes", "no"]

ANNOTATION_COLUMNS = [
    "sample_id",
    "cluster_id",
    "article_file_id",
    "target",
    "event_name",
    "claim_title",
    "claim_text",
    "publish_time",
    "evidence_list",
    "evidence_sources",
    "evidence_times",
    "financial_context",
    "market_context",
    *HUMAN_FIELDS,
]

FORBIDDEN_VISIBLE_PATTERNS = (
    "model_prediction",
    "prediction",
    "predicted_label",
    "silver_label",
    "final_label",
    "gold_label",
    "human_verdict",
    "llm_verdict",
    "cached_dea",
    "baseline_verdict",
    "credibility_score",
    "confidence_score",
    "market_consistency_verdict",
    "financial_consistency_verdict",
    "rule_verdict",
    "verdict=",
)

MARKET_TERMS = (
    "涨停",
    "跌停",
    "股价",
    "股票",
    "收涨",
    "收跌",
    "上涨",
    "下跌",
    "大涨",
    "大跌",
    "盘中",
    "收盘",
    "开盘",
    "换手率",
    "成交额",
)
FINANCIAL_TERMS = (
    "营收",
    "营业收入",
    "净利润",
    "归母净利润",
    "扣非净利润",
    "现金流",
    "盈利",
    "亏损",
    "偿债",
    "负债",
    "财报",
    "年报",
    "季报",
    "业绩",
)


@dataclass(frozen=True)
class HistorySource:
    name: str
    category: str
    path: Path
    optional: bool = True


HISTORY_SOURCES = [
    HistorySource("old_locked500_samples", "retired_independent500_v1", ROOT / "outputs/event_credibility/task2_independent500_v1/independent_samples.jsonl"),
    HistorySource("clean_evidence_1500", "clean_evidence_1500", ROOT / "outputs/event_credibility/task2_hybrid_v3/clean_evidence_v1/clean_evidence_1500.jsonl"),
    HistorySource("heterogeneous_silver_final", "heterogeneous_silver_v2", ROOT / "outputs/event_credibility/task2_hybrid_v3/heterogeneous_silver_v2/silver_final_1500.jsonl"),
    HistorySource("heterogeneous_silver_training_eligible", "heterogeneous_silver_v2", ROOT / "outputs/event_credibility/task2_hybrid_v3/heterogeneous_silver_v2/silver_training_eligible.jsonl"),
    HistorySource("heterogeneous_silver_needs_review", "heterogeneous_silver_v2", ROOT / "outputs/event_credibility/task2_hybrid_v3/heterogeneous_silver_v2/silver_needs_review.jsonl"),
    HistorySource("fusion_train", "fusion_model_v1_split", ROOT / "outputs/event_credibility/task2_hybrid_v3/fusion_model_v1/train.jsonl"),
    HistorySource("fusion_dev", "fusion_model_v1_split", ROOT / "outputs/event_credibility/task2_hybrid_v3/fusion_model_v1/dev.jsonl"),
    HistorySource("fusion_internal_test", "fusion_model_v1_split", ROOT / "outputs/event_credibility/task2_hybrid_v3/fusion_model_v1/internal_test.jsonl"),
    HistorySource("old_47_dea_gold_manifest", "old_47_dea_gold_identifiers", ROOT / "reports/task2_dea_cached_human_audit_manifest.json"),
    HistorySource("old_47_dea_manual_audit_v1", "old_47_dea_gold_identifiers", ROOT / "reports/task2_dea_manual_audit.jsonl"),
    HistorySource("old_47_dea_manual_audit_v2", "old_47_dea_gold_identifiers", ROOT / "reports/task2_dea_v2_manual_audit.jsonl"),
    HistorySource("old_47_dea_manual_audit_cached", "old_47_dea_gold_identifiers", ROOT / "reports/task2_dea_v2_cached_manual_audit.jsonl"),
    HistorySource("old_50_safety_gate_manifest", "old_50_safety_gate_identifiers", ROOT / "reports/task2_safety_gate_v1_2_independent_manifest.json"),
    HistorySource("old_50_independent_samples", "old_50_safety_gate_identifiers", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/independent_samples.jsonl"),
    HistorySource("old_50_primary_samples", "old_50_safety_gate_identifiers", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/primary_samples.jsonl"),
    HistorySource("old_50_stress_samples", "old_50_safety_gate_identifiers", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/contradiction_stress_samples.jsonl"),
    HistorySource("old_50_blind_review_queue", "old_50_safety_gate_identifiers", ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/blind_review_queue.jsonl"),
    HistorySource("dea_v1_gray_samples", "task2_historical_samples", ROOT / "outputs/event_credibility/dea_v1_gray/sample_rows.jsonl"),
    HistorySource("dea_v2_gray_samples", "task2_historical_samples", ROOT / "outputs/event_credibility/dea_v2_gray/sample_rows.jsonl"),
    HistorySource("dea_v2_cached_gray_samples", "task2_historical_samples", ROOT / "outputs/event_credibility/dea_v2_cached_gray/sample_rows.jsonl"),
    HistorySource("llm_concurrency_sample_24", "task2_historical_samples", ROOT / "reports/task2_llm_concurrency_benchmark_runs/20260807_163324/task2_llm_concurrency_sample_24.jsonl"),
    HistorySource("isr_candidates", "task2_historical_samples", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/isr_candidates.jsonl"),
    HistorySource("isr_claim_fix_candidates", "task2_historical_samples", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/claim_fix_candidates.jsonl"),
    HistorySource("isr_safety_gate_fix_candidates", "task2_historical_samples", ROOT / "outputs/event_credibility/isr_candidate_audit_v1/safety_gate_fix_candidates.jsonl"),
    HistorySource("claim_contract_sample_manifest", "task2_historical_samples", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/sample_manifest.json"),
    HistorySource("claim_contract_contracts", "task2_historical_samples", ROOT / "outputs/event_credibility/claim_contract_adapter_v1_gray/claim_contracts.jsonl"),
    HistorySource("safety_gate_zero_trigger_audit_rows", "task2_historical_samples", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/audit_rows.jsonl"),
    HistorySource("safety_gate_zero_trigger_error_focus", "task2_historical_samples", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/error_focus_13.jsonl"),
    HistorySource("safety_gate_zero_trigger_insufficient_focus", "task2_historical_samples", ROOT / "outputs/event_credibility/safety_gate_v1_2_zero_trigger_audit_v1/insufficient_resolution_focus_5.jsonl"),
    HistorySource("task2_independent_error_audit", "task2_historical_samples", ROOT / "reports/task2_independent_error_audit_v1.json"),
    HistorySource("task1_structure_independent60_source", "task1_task3_same_source_identifiers", ROOT / "outputs/task1_step5_7i_structure_expansion/independent_frozen_test_60_source.jsonl"),
    HistorySource("task1_structure_confirmed60", "task1_task3_same_source_identifiers", ROOT / "outputs/task1_step5_7j_structure_confirmation/independent_frozen_test_confirmed.jsonl"),
    HistorySource("task3_split_test", "task1_task3_same_source_identifiers", ROOT / "outputs/task3/splits/test.jsonl"),
]


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def now_iso() -> str:
    return datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")


def norm(value: Any) -> str:
    if value is None:
        return ""
    try:
        if isinstance(value, float) and math.isnan(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def to_jsonable(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, dict):
        return {str(k): to_jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [to_jsonable(item) for item in value]
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    return str(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_text(value: Any, limit: int = 100000) -> str:
    text = norm(value).lower()
    text = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", text)
    return text[:limit]


def text_sha256(value: Any) -> str:
    text = stable_text(value)
    if not text:
        return ""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def deterministic_key(*parts: Any, seed: int = SEED) -> str:
    payload = "|".join(norm(part) for part in parts) + f"|seed={seed}"
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def first_date(value: Any) -> str:
    text = norm(value)
    match = re.search(r"(20\d{2})[-/年.](\d{1,2})[-/月.](\d{1,2})", text)
    if not match:
        return ""
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def month_bucket(value: Any) -> str:
    date = first_date(value)
    return date[:7] if date else "unknown_month"


def parse_date_value(value: Any) -> datetime | None:
    date = first_date(value)
    if not date:
        return None
    try:
        return datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        return None


def date_distance_days(left: Any, right: Any) -> int | None:
    a = parse_date_value(left)
    b = parse_date_value(right)
    if a is None or b is None:
        return None
    return abs((a - b).days)


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows = []
    with path.open(encoding="utf-8-sig") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            obj = json.loads(line)
            if not isinstance(obj, dict):
                raise ValueError(f"{rel(path)}:{line_number} is not a JSON object")
            rows.append(obj)
    return rows


def read_csv(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(to_jsonable(payload), ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(to_jsonable(row), ensure_ascii=False, sort_keys=True, allow_nan=False) + "\n")


def parse_jsonish(value: Any) -> Any:
    if isinstance(value, (list, dict)):
        return value
    text = norm(value)
    if not text:
        return []
    try:
        return json.loads(text)
    except Exception:
        return value


def title_parts(value: Any) -> list[str]:
    return [norm(part) for part in re.split(r"\s*\|\|\s*|\n+", norm(value)) if norm(part)]


def source_list_from_value(value: Any) -> list[str]:
    parsed = parse_jsonish(value)
    raw_values = parsed if isinstance(parsed, list) else re.split(r"\s*\|\s*|;|；|,", norm(value))
    out: list[str] = []
    seen: set[str] = set()
    for raw in raw_values:
        item = norm(raw)
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return out


def collect_identifiers(
    value: Any,
    identifiers: dict[str, set[str]],
) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            lowered = str(key).lower()
            if lowered in {"sample_id", "cluster_id", "event_id", "candidate_id", "source_event_id"} or "cluster_id" in lowered:
                if isinstance(item, str) and item:
                    identifiers["ids"].add(item)
                    for match in re.findall(r"EC-\d{6}(?:__v\d+)?", item):
                        identifiers["legacy_cluster_ids"].add(match.split("__", 1)[0])
            if lowered in {"article_id", "article_file_id", "document_id", "canonical_article_id", "trace_id", "claim_article_id"} or "article_id" in lowered or "article_file_id" in lowered:
                if isinstance(item, str) and item:
                    identifiers["article_ids"].add(item)
            if "normalized_title_sha256" in lowered and isinstance(item, str) and item:
                identifiers["title_hashes"].add(item)
            if "normalized_content_sha256" in lowered and isinstance(item, str) and item:
                identifiers["content_hashes"].add(item)
            if lowered in {"title", "claim_title", "article_title"} and isinstance(item, str):
                digest = text_sha256(item)
                if digest:
                    identifiers["title_hashes"].add(digest)
            if lowered == "title_samples" and isinstance(item, str):
                for title in title_parts(item):
                    digest = text_sha256(title)
                    if digest:
                        identifiers["title_hashes"].add(digest)
            if lowered == "content" and isinstance(item, str):
                digest = text_sha256(item)
                if digest:
                    identifiers["content_hashes"].add(digest)
            collect_identifiers(item, identifiers)
    elif isinstance(value, list):
        for item in value:
            collect_identifiers(item, identifiers)


def target_key(row: dict[str, Any]) -> str:
    claim = row.get("claim") if isinstance(row.get("claim"), dict) else {}
    code = norm(row.get("target_code") or claim.get("target_code") or row.get("target") or claim.get("target"))
    match = re.search(r"(?<!\d)(\d{6})(?!\d)", code)
    if match:
        return match.group(1)
    name = norm(row.get("target_name") or row.get("target_entity") or claim.get("target") or row.get("target"))
    return stable_text(name, 80)


def profile_from_row(row: dict[str, Any], source_name: str = "") -> dict[str, Any] | None:
    claim = row.get("claim") if isinstance(row.get("claim"), dict) else {}
    titles = " ".join(
        part
        for part in (
            row.get("claim_title"),
            row.get("title"),
            claim.get("title"),
            row.get("title_samples"),
        )
        if norm(part)
    )
    if "||" in titles:
        titles = " ".join(title_parts(titles))
    claim_text = norm(row.get("claim_text") or row.get("event_claim") or claim.get("event_claim") or row.get("claim"))
    event_name = norm(row.get("event_name") or row.get("event_type") or row.get("schema_id") or claim.get("event_type"))
    event_time = norm(row.get("publish_time") or row.get("event_time") or row.get("earliest_publish_time") or claim.get("publish_time"))
    key = target_key(row)
    if not (titles or claim_text or key):
        return None
    return {
        "source_name": source_name,
        "id": norm(row.get("sample_id") or row.get("cluster_id") or row.get("event_id") or row.get("candidate_id")),
        "target_key": key,
        "target_name": norm(row.get("target_name") or row.get("target_entity") or claim.get("target")),
        "event_name": event_name,
        "date": first_date(event_time),
        "title_text": titles,
        "claim_text": claim_text,
        "title_sha256": text_sha256(titles),
    }


def load_history_rows(source: HistorySource) -> list[dict[str, Any]]:
    path = source.path
    if not path.exists():
        return []
    if path.suffix == ".jsonl":
        return read_jsonl(path)
    if path.suffix == ".json":
        payload = read_json(path)
        return [payload] if isinstance(payload, dict) else [item for item in payload if isinstance(item, dict)]
    if path.suffix == ".csv":
        return read_csv(path)
    return []


def collect_history() -> dict[str, Any]:
    identifiers: dict[str, set[str]] = {
        "ids": set(),
        "article_ids": set(),
        "title_hashes": set(),
        "content_hashes": set(),
        "legacy_cluster_ids": set(),
    }
    profiles: list[dict[str, Any]] = []
    source_stats = []
    category_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for source in HISTORY_SOURCES:
        rows = load_history_rows(source)
        before = {key: len(value) for key, value in identifiers.items()}
        for row in rows:
            collect_identifiers(row, identifiers)
            profile = profile_from_row(row, source.name)
            if profile:
                profiles.append(profile)
        after = {key: len(value) for key, value in identifiers.items()}
        delta = {key: after[key] - before[key] for key in before}
        for key, value in delta.items():
            category_counts[source.category][key] += value
        source_stats.append(
            {
                "name": source.name,
                "category": source.category,
                "path": rel(source.path),
                "exists": source.path.exists(),
                "row_count": len(rows),
                "new_identifier_counts": delta,
            }
        )
    return {
        "identifiers": identifiers,
        "profiles": profiles,
        "source_stats": source_stats,
        "category_identifier_counts": {
            category: dict(counter)
            for category, counter in sorted(category_counts.items())
        },
    }


def sequence_similarity(left: Any, right: Any) -> float:
    a = stable_text(left, 320)
    b = stable_text(right, 320)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    max_possible = 2.0 * min(len(a), len(b)) / (len(a) + len(b))
    if max_possible < 0.82:
        return max_possible
    matcher = SequenceMatcher(None, a, b)
    if matcher.real_quick_ratio() < 0.82:
        return matcher.real_quick_ratio()
    if matcher.quick_ratio() < 0.82:
        return matcher.quick_ratio()
    return matcher.ratio()


def is_near_duplicate(left: dict[str, Any], right: dict[str, Any]) -> tuple[bool, str, float]:
    if left.get("target_key") and right.get("target_key") and left["target_key"] != right["target_key"]:
        return False, "", 0.0
    if left.get("event_name") and right.get("event_name") and left["event_name"] != right["event_name"]:
        return False, "", 0.0
    title_sim = sequence_similarity(left.get("title_text"), right.get("title_text"))
    claim_sim = sequence_similarity(left.get("claim_text"), right.get("claim_text"))
    date_distance = date_distance_days(left.get("date"), right.get("date"))
    if left.get("date") and left.get("date") == right.get("date") and title_sim >= 0.90:
        return True, "same_target_event_date_title_similarity_ge_0.90", title_sim
    if date_distance is not None and date_distance <= 2 and max(title_sim, claim_sim) >= 0.92:
        return True, "same_target_event_near_date_text_similarity_ge_0.92", max(title_sim, claim_sim)
    if title_sim >= 0.97:
        return True, "same_target_event_high_title_similarity_ge_0.97", title_sim
    return False, "", max(title_sim, claim_sim)


def history_profile_buckets(profiles: list[dict[str, Any]]) -> dict[tuple[str, str], list[dict[str, Any]]]:
    buckets: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for profile in profiles:
        buckets[(profile.get("target_key", ""), profile.get("event_name", ""))].append(profile)
    return buckets


def candidate_profile(candidate: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": candidate["source_event_id"],
        "target_key": candidate["target_code"],
        "target_name": candidate["target_name"],
        "event_name": candidate["event_name"],
        "date": first_date(candidate["publish_time"]),
        "title_text": candidate["claim_title"],
        "claim_text": candidate["claim_text"],
        "title_sha256": candidate["normalized_title_sha256"],
    }


def near_duplicate_to_history(candidate: dict[str, Any], buckets: dict[tuple[str, str], list[dict[str, Any]]]) -> tuple[bool, dict[str, Any]]:
    profile = candidate_profile(candidate)
    candidates = buckets.get((profile["target_key"], profile["event_name"]), [])
    for historical in candidates:
        duplicate, method, score = is_near_duplicate(profile, historical)
        if duplicate:
            return True, {
                "source_event_id": candidate["source_event_id"],
                "historical_id": historical.get("id"),
                "historical_source": historical.get("source_name"),
                "method": method,
                "similarity": round(score, 6),
            }
    return False, {}


def row_near_selected(candidate: dict[str, Any], selected_profiles: list[dict[str, Any]]) -> tuple[bool, dict[str, Any]]:
    profile = candidate_profile(candidate)
    for selected in selected_profiles:
        duplicate, method, score = is_near_duplicate(profile, selected)
        if duplicate:
            return True, {
                "source_event_id": candidate["source_event_id"],
                "near_duplicate_selected_source_event_id": selected.get("id"),
                "method": method,
                "similarity": round(score, 6),
            }
    return False, {}


def source_type_bucket(value: Any) -> str:
    item = norm(value) or "unknown_source_type"
    return item


def target_bucket(code: str) -> str:
    if not code:
        return "unknown_target"
    return f"target_hash_{int(hashlib.sha256(code.encode('utf-8')).hexdigest()[:4], 16) % 12:02d}"


def event_bucket(event_name: str) -> str:
    return norm(event_name) or "unknown_event"


def candidate_stratum(candidate: dict[str, Any]) -> str:
    return "|".join(
        [
            source_type_bucket(candidate.get("source_type")),
            event_bucket(candidate.get("event_name")),
            month_bucket(candidate.get("publish_time")),
            target_bucket(candidate.get("target_code")),
        ]
    )


def raw_context(title: str, content: str, terms: Iterable[str], max_snippets: int = 3) -> str:
    text = norm(f"{title}。{content}")
    snippets = []
    for sentence in re.split(r"(?<=[。！？!?；;])\s*|\n+", text):
        sentence = norm(sentence)
        if not sentence:
            continue
        if any(term in sentence for term in terms):
            snippets.append(sentence[:220])
        if len(snippets) >= max_snippets:
            break
    if not snippets:
        return ""
    return " | ".join(snippets)


def evidence_sentence(article: dict[str, Any], target_name: str, event_name: str) -> str:
    content = norm(article.get("content"))
    title = norm(article.get("article_title"))
    terms = [term for term in (target_name, event_name) if term]
    for sentence in re.split(r"(?<=[。！？!?；;])\s*|\n+", content):
        sentence = norm(sentence)
        if sentence and any(term in sentence for term in terms):
            return sentence[:360]
    return (content or title)[:360]


def load_article_map(path: Path) -> dict[str, dict[str, Any]]:
    df = pd.read_parquet(
        path,
        columns=["article_id", "article_title", "article_publish_time", "article_source", "source_type", "content", "content_length"],
    )
    out: dict[str, dict[str, Any]] = {}
    for row in df.to_dict("records"):
        article_id = norm(row.get("article_id"))
        if article_id:
            out[article_id] = {str(key): to_jsonable(value) for key, value in row.items()}
    return out


def normalize_candidate_row(row: dict[str, Any], article_map: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
    article_id = norm(row.get("article_id"))
    article = article_map.get(article_id)
    if not article:
        return None
    title = norm(row.get("title") or article.get("article_title"))
    content = norm(article.get("content"))
    target_code = re.sub(r"\D", "", norm(row.get("target_code")))[:6]
    target_name = norm(row.get("target_name"))
    event_name = norm(row.get("event_name"))
    publish_time = norm(row.get("publish_time") or article.get("article_publish_time"))
    source_type = norm(row.get("source_type") or article.get("source_type"))
    source = norm(row.get("source") or article.get("article_source"))
    if not article_id or not title or not content or not target_code or norm(row.get("target_type")) != "company":
        return None
    claim_text = f"关于{target_name}的{event_name or '相关事件'}主张：{title}"
    return {
        "source_event_id": norm(row.get("event_id")),
        "claim_article_id": article_id,
        "article_file_id": article_id,
        "target_code": target_code,
        "target_name": target_name,
        "target": f"{target_name} / {target_code}".strip(" /"),
        "schema_id": norm(row.get("schema_id")),
        "event_name": event_name,
        "claim_title": title,
        "claim_text": claim_text,
        "publish_time": publish_time,
        "source": source,
        "source_type": source_type,
        "dataset_split": norm(row.get("dataset_split")),
        "quality_status": norm(row.get("quality_status")),
        "normalized_title_sha256": text_sha256(title),
        "normalized_content_sha256": text_sha256(content),
        "source_event_group_key": "|".join([target_code, event_name, first_date(publish_time), stable_text(title, 96)]),
        "raw_financial_context": raw_context(title, content, FINANCIAL_TERMS),
        "raw_market_context": raw_context(title, content, MARKET_TERMS),
        "has_raw_financial_terms": any(term in f"{title} {content[:1000]}" for term in FINANCIAL_TERMS),
        "has_raw_market_terms": any(term in f"{title} {content[:1000]}" for term in MARKET_TERMS),
    }


def build_evidence_items(
    candidate: dict[str, Any],
    all_candidates: list[dict[str, Any]],
    article_map: dict[str, dict[str, Any]],
    excluded_article_ids: set[str],
    limit: int = 6,
) -> list[dict[str, Any]]:
    claim_article_id = candidate["claim_article_id"]
    same_event = [
        row
        for row in all_candidates
        if row["claim_article_id"] != claim_article_id
        and row["claim_article_id"] not in excluded_article_ids
        and row["target_code"] == candidate["target_code"]
        and row["event_name"] == candidate["event_name"]
    ]
    same_target = [
        row
        for row in all_candidates
        if row["claim_article_id"] != claim_article_id
        and row["claim_article_id"] not in excluded_article_ids
        and row["target_code"] == candidate["target_code"]
        and row["event_name"] != candidate["event_name"]
    ]
    rows = same_event or same_target
    rows.sort(
        key=lambda row: (
            date_distance_days(candidate["publish_time"], row["publish_time"]) if date_distance_days(candidate["publish_time"], row["publish_time"]) is not None else 9999,
            0 if row["source"] != candidate["source"] else 1,
            deterministic_key(row["source_event_id"], row["claim_article_id"]),
        )
    )
    selected = []
    seen_articles: set[str] = set()
    seen_sources: set[str] = set()
    for prefer_distinct_source in (True, False):
        for row in rows:
            article_id = row["claim_article_id"]
            if article_id in seen_articles:
                continue
            if prefer_distinct_source and row["source"] in seen_sources:
                continue
            article = article_map.get(article_id, {})
            if not norm(article.get("content")):
                continue
            selected.append(row)
            seen_articles.add(article_id)
            seen_sources.add(row["source"])
            if len(selected) >= limit:
                break
        if len(selected) >= limit:
            break
    evidence = []
    for idx, row in enumerate(selected, 1):
        article = article_map.get(row["claim_article_id"], {})
        evidence.append(
            {
                "evidence_id": f"E{idx}",
                "article_id": row["claim_article_id"],
                "title": row["claim_title"],
                "publish_time": row["publish_time"],
                "source": row["source"],
                "source_type": row["source_type"],
                "evidence_text": evidence_sentence(article, candidate["target_name"], candidate["event_name"]),
                "evidence_relation": "same_target_event" if row["event_name"] == candidate["event_name"] else "same_target_context",
                "trace_path": f"{rel(ARTICLE_INDEX)}#article_id={row['claim_article_id']}",
            }
        )
    return evidence


def exact_exclusion_reason(candidate: dict[str, Any], identifiers: dict[str, set[str]]) -> str:
    if candidate["source_event_id"] in identifiers["ids"]:
        return "source_event_id_exact_overlap"
    if candidate["claim_article_id"] in identifiers["article_ids"]:
        return "article_id_exact_overlap"
    if candidate["normalized_title_sha256"] in identifiers["title_hashes"]:
        return "title_sha256_exact_overlap"
    if candidate["normalized_content_sha256"] in identifiers["content_hashes"]:
        return "content_sha256_exact_overlap"
    return ""


def load_candidates(all_events_path: Path, article_map: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows = read_jsonl(all_events_path)
    candidates = []
    seen_event_article: set[tuple[str, str]] = set()
    for row in rows:
        candidate = normalize_candidate_row(row, article_map)
        if not candidate:
            continue
        key = (candidate["source_event_id"], candidate["claim_article_id"])
        if key in seen_event_article:
            continue
        seen_event_article.add(key)
        candidates.append(candidate)
    return candidates


def filter_candidates(
    candidates: list[dict[str, Any]],
    article_map: dict[str, dict[str, Any]],
    history: dict[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    identifiers = history["identifiers"]
    exact_reasons: Counter[str] = Counter()
    exact_eligible = []
    for candidate in candidates:
        reason = exact_exclusion_reason(candidate, identifiers)
        if reason:
            exact_reasons[reason] += 1
        else:
            exact_eligible.append(candidate)
    buckets = history_profile_buckets(history["profiles"])
    near_removed = []
    near_eligible = []
    for candidate in exact_eligible:
        duplicate, detail = near_duplicate_to_history(candidate, buckets)
        if duplicate:
            near_removed.append(detail)
        else:
            near_eligible.append(candidate)
    with_evidence = []
    evidence_missing = 0
    all_articles_excluded = set(identifiers["article_ids"])
    for candidate in near_eligible:
        evidence = build_evidence_items(candidate, near_eligible, article_map, all_articles_excluded)
        if not evidence:
            evidence_missing += 1
            continue
        candidate["evidence_items"] = evidence
        candidate["evidence_count"] = len(evidence)
        candidate["distinct_evidence_source_count"] = len({item["source"] for item in evidence if item["source"]})
        with_evidence.append(candidate)
    summary = {
        "raw_candidate_count": len(candidates),
        "candidate_count_after_exact_exclusion": len(exact_eligible),
        "candidate_count_after_near_duplicate_exclusion": len(near_eligible),
        "candidate_count_with_independent_evidence": len(with_evidence),
        "exact_exclusion_reasons": dict(exact_reasons),
        "near_duplicate_exclusion_count": len(near_removed),
        "evidence_missing_after_exclusion_count": evidence_missing,
    }
    return with_evidence, summary, near_removed


def select_samples(candidates: list[dict[str, Any]], count: int, seed: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    if len(candidates) < count:
        raise RuntimeError(f"blocked_by_insufficient_independent_candidates: need {count}, got {len(candidates)}")
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for candidate in candidates:
        groups[candidate_stratum(candidate)].append(candidate)
    for key in groups:
        groups[key].sort(key=lambda row: deterministic_key(row["source_event_id"], row["claim_article_id"], seed=seed))
    quota_rows = []
    allocated = 0
    total = len(candidates)
    for key, rows in sorted(groups.items()):
        exact = len(rows) * count / total
        quota = min(len(rows), int(math.floor(exact)))
        allocated += quota
        quota_rows.append([key, quota, exact - quota, len(rows)])
    remaining = count - allocated
    for item in sorted(quota_rows, key=lambda row: (-row[2], row[0])):
        if remaining <= 0:
            break
        if item[1] < item[3]:
            item[1] += 1
            remaining -= 1

    selected: list[dict[str, Any]] = []
    selected_profiles: list[dict[str, Any]] = []
    used_group_keys: set[str] = set()
    near_selected_rejections = []

    def try_add(row: dict[str, Any]) -> bool:
        if row["source_event_group_key"] in used_group_keys:
            return False
        duplicate, detail = row_near_selected(row, selected_profiles)
        if duplicate:
            near_selected_rejections.append(detail)
            return False
        selected.append(row)
        selected_profiles.append(candidate_profile(row))
        used_group_keys.add(row["source_event_group_key"])
        return True

    quota_by_key = {key: quota for key, quota, _fraction, _size in quota_rows}
    for key, quota in quota_by_key.items():
        added = 0
        for row in groups[key]:
            if added >= quota:
                break
            if try_add(row):
                added += 1
    if len(selected) < count:
        selected_ids = {row["source_event_id"] for row in selected}
        fill_pool = [row for row in candidates if row["source_event_id"] not in selected_ids]
        fill_pool.sort(key=lambda row: deterministic_key("fill", row["source_event_id"], row["claim_article_id"], seed=seed))
        for row in fill_pool:
            if len(selected) >= count:
                break
            try_add(row)
    if len(selected) != count:
        raise RuntimeError(f"blocked_by_near_duplicate_constraints: selected {len(selected)}, need {count}")
    selected.sort(key=lambda row: deterministic_key("final_order", row["source_event_id"], row["claim_article_id"], seed=seed))
    for index, row in enumerate(selected, 1):
        row["sample_id"] = f"I500V2-{index:06d}"
        row["cluster_id"] = row["sample_id"]
        row["sample_index"] = index
        row["stratum_key"] = candidate_stratum(row)
        for field in HUMAN_FIELDS:
            row[field] = ""
    summary = {
        "stratum_count": len(groups),
        "selected_count": len(selected),
        "near_duplicate_rejected_during_selection_count": len(near_selected_rejections),
        "duplicate_source_event_group_key_count": len(selected) - len({row["source_event_group_key"] for row in selected}),
    }
    return selected, summary


def annotation_row(sample: dict[str, Any]) -> dict[str, Any]:
    evidence_lines = [
        f"{item['evidence_id']} | {item['source']} | {item['publish_time']} | {item['title']} | {item['evidence_text']}"
        for item in sample.get("evidence_items", [])
    ]
    row = {
        "sample_id": sample["sample_id"],
        "cluster_id": sample["cluster_id"],
        "article_file_id": sample["article_file_id"],
        "target": sample["target"],
        "event_name": sample["event_name"],
        "claim_title": sample["claim_title"],
        "claim_text": sample["claim_text"],
        "publish_time": sample["publish_time"],
        "evidence_list": "\n".join(evidence_lines),
        "evidence_sources": "; ".join(dict.fromkeys(item["source"] for item in sample.get("evidence_items", []) if item["source"])),
        "evidence_times": "; ".join(dict.fromkeys(item["publish_time"] for item in sample.get("evidence_items", []) if item["publish_time"])),
        "financial_context": sample["raw_financial_context"],
        "market_context": sample["raw_market_context"],
    }
    for field in HUMAN_FIELDS:
        row[field] = ""
    return row


def evidence_detail_rows(samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for sample in samples:
        for item in sample.get("evidence_items", []):
            rows.append(
                {
                    "sample_id": sample["sample_id"],
                    "cluster_id": sample["cluster_id"],
                    "article_file_id": sample["article_file_id"],
                    "source_event_id": sample["source_event_id"],
                    "target": sample["target"],
                    "event_name": sample["event_name"],
                    "evidence_id": item["evidence_id"],
                    "evidence_article_id": item["article_id"],
                    "evidence_relation": item["evidence_relation"],
                    "title": item["title"],
                    "publish_time": item["publish_time"],
                    "source": item["source"],
                    "source_type": item["source_type"],
                    "evidence_text": item["evidence_text"],
                    "trace_path": item["trace_path"],
                }
            )
    return rows


def distribution_rows(samples: list[dict[str, Any]]) -> list[dict[str, Any]]:
    dimensions = {
        "source_type": [sample["source_type"] or "unknown" for sample in samples],
        "event_name": [sample["event_name"] or "unknown" for sample in samples],
        "target_company": [sample["target"] or "unknown" for sample in samples],
        "publish_month": [month_bucket(sample["publish_time"]) for sample in samples],
        "dataset_split": [sample["dataset_split"] or "unknown" for sample in samples],
        "raw_market_claim_context": ["yes" if sample.get("has_raw_market_terms") else "no" for sample in samples],
        "raw_financial_claim_context": ["yes" if sample.get("has_raw_financial_terms") else "no" for sample in samples],
        "evidence_count": [str(sample.get("evidence_count") or 0) for sample in samples],
    }
    rows = []
    total = len(samples)
    for dimension, values in dimensions.items():
        for value, count in sorted(Counter(values).items()):
            rows.append({"dimension": dimension, "value": value, "count": count, "ratio": round(count / max(1, total), 6)})
    return rows


def safe_cell(value: Any, limit: int = 32000) -> str:
    if isinstance(value, (list, dict)):
        text = json.dumps(to_jsonable(value), ensure_ascii=False, allow_nan=False)
    else:
        text = norm(value)
    return text if len(text) <= limit else text[: limit - 40] + " ...[truncated]"


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]], human_columns: set[str] | None = None) -> None:
    ws = wb.create_sheet(title=title[:31])
    if not rows:
        rows = [{"empty": ""}]
    headers = list(rows[0].keys())
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    human_fill = PatternFill("solid", fgColor="FCE4D6")
    auto_fill = PatternFill("solid", fgColor="F2F2F2")
    human_columns = human_columns or set()
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([safe_cell(row.get(header, "")) for header in headers])
    for column_idx, header in enumerate(headers, 1):
        fill = human_fill if header in human_columns else auto_fill
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row_idx, column_idx)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "claim_title": 52,
        "claim_text": 62,
        "evidence_list": 90,
        "financial_context": 60,
        "market_context": 60,
        "evidence_text": 90,
        "trace_path": 60,
    }
    for idx, header in enumerate(headers, 1):
        width = widths.get(header)
        if width is None:
            values = [safe_cell(row.get(header, "")) for row in rows[:100]]
            width = min(max(max([len(str(header))] + [min(len(value), 80) for value in values]) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(idx)].width = width
    validations = {
        "human_label": LABEL_VALUES,
        "adjudicated_label": LABEL_VALUES,
        "human_confidence": CONFIDENCE_VALUES,
        "needs_second_review": REVIEW_VALUES,
    }
    for idx, header in enumerate(headers, 1):
        if header not in validations:
            continue
        dv = DataValidation(type="list", formula1=f'"{",".join(validations[header])}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(idx)}2:{get_column_letter(idx)}{max(ws.max_row, 2)}")


def write_workbook(path: Path, samples: list[dict[str, Any]], sampling_manifest: dict[str, Any], leakage_audit: dict[str, Any]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "annotation", [annotation_row(sample) for sample in samples], human_columns=set(HUMAN_FIELDS))
    write_sheet(wb, "evidence_detail", evidence_detail_rows(samples))
    write_sheet(wb, "sampling_distribution", distribution_rows(samples))
    write_sheet(
        wb,
        "protocol",
        [
            {"item": "dataset_id", "value": "task2_independent500_v2"},
            {"item": "label_contract", "value": ", ".join(LABEL_VALUES)},
            {"item": "human_fields_blank", "value": "true"},
            {"item": "model_predictions_visible", "value": "false"},
            {"item": "model_predictions_read_for_sampling", "value": str(sampling_manifest["constraints"]["model_predictions_read"]).lower()},
            {"item": "old_gold_label_values_used_for_sampling", "value": str(sampling_manifest["constraints"]["old_gold_used_for_sampling"]).lower()},
            {"item": "leakage_audit_exact_overlap_zero", "value": str(leakage_audit["all_required_overlap_counts_zero"]).lower()},
        ],
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def render_guide() -> str:
    return "\n".join(
        [
            "# Task2 Independent500-v2 Annotation Guide",
            "",
            "本工作簿用于新的Task2独立500条人工盲标。判断对象是事件主张与可见独立证据之间的可信关系。",
            "",
            "## 标签契约",
            "",
            "- support：证据明确支持主张核心事实。",
            "- contradict：证据明确否定、澄清、撤回或反转主张核心事实。",
            "- insufficient：可见证据不足以支持或反驳。",
            "- unclear：主张或证据边界不清，无法可靠归入前三类。",
            "",
            "## 填写字段",
            "",
            "- human_label：从 support / contradict / insufficient / unclear 中选择。",
            "- human_confidence：high / medium / low。",
            "- human_reason：简要说明人工判断依据。",
            "- selected_evidence_ids：填写使用的证据编号，如 E1;E3。",
            "- needs_second_review：yes / no。",
            "- adjudicated_label：二审或仲裁后填写，初始为空。",
            "",
            "自动字段仅提供原始新闻、来源、时间和证据文本。工作簿不包含模型预测、旧标签、银标、规则verdict或置信分。",
        ]
    ) + "\n"


def validate_no_visible_answer_leak(samples: list[dict[str, Any]], workbook: Path | None = None) -> dict[str, Any]:
    non_human_payloads = []
    for sample in samples:
        row = annotation_row(sample)
        for key, value in row.items():
            if key in HUMAN_FIELDS:
                continue
            non_human_payloads.append((sample["sample_id"], key, safe_cell(value)))
        for item in sample.get("evidence_items", []):
            for key, value in item.items():
                non_human_payloads.append((sample["sample_id"], f"evidence.{key}", safe_cell(value)))
    leaks = []
    for sample_id, key, value in non_human_payloads:
        lowered = value.lower()
        if any(pattern in lowered for pattern in FORBIDDEN_VISIBLE_PATTERNS):
            leaks.append({"sample_id": sample_id, "field": key, "pattern": next(pattern for pattern in FORBIDDEN_VISIBLE_PATTERNS if pattern in lowered)})
    human_nonempty = []
    for sample in samples:
        for field in HUMAN_FIELDS:
            if norm(sample.get(field)):
                human_nonempty.append({"sample_id": sample["sample_id"], "field": field})
    return {
        "human_fields_blank": len(human_nonempty) == 0,
        "human_nonempty_examples": human_nonempty[:20],
        "visible_answer_leak_count": len(leaks),
        "visible_answer_leak_examples": leaks[:20],
        "workbook_checked": rel(workbook) if workbook else "",
    }


def intersection_counts(samples: list[dict[str, Any]], identifiers: dict[str, set[str]]) -> dict[str, int]:
    sample_ids = {sample["sample_id"] for sample in samples}
    cluster_ids = {sample["cluster_id"] for sample in samples}
    source_event_ids = {sample["source_event_id"] for sample in samples}
    article_ids = {sample["claim_article_id"] for sample in samples}
    evidence_article_ids = {item["article_id"] for sample in samples for item in sample.get("evidence_items", [])}
    title_hashes = {sample["normalized_title_sha256"] for sample in samples if sample["normalized_title_sha256"]}
    content_hashes = {sample["normalized_content_sha256"] for sample in samples if sample["normalized_content_sha256"]}
    return {
        "sample_id_exact_overlap": len(sample_ids & identifiers["ids"]),
        "cluster_id_exact_overlap": len(cluster_ids & identifiers["ids"]),
        "source_event_id_exact_overlap": len(source_event_ids & identifiers["ids"]),
        "claim_article_id_exact_overlap": len(article_ids & identifiers["article_ids"]),
        "evidence_article_id_exact_overlap": len(evidence_article_ids & identifiers["article_ids"]),
        "title_sha256_exact_overlap": len(title_hashes & identifiers["title_hashes"]),
        "content_sha256_exact_overlap": len(content_hashes & identifiers["content_hashes"]),
        "legacy_cluster_overlap": len({sample["source_event_id"] for sample in samples if sample["source_event_id"].startswith("EC-")} & identifiers["legacy_cluster_ids"]),
    }


def selected_near_duplicate_overlap(samples: list[dict[str, Any]], history_profiles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets = history_profile_buckets(history_profiles)
    overlaps = []
    for sample in samples:
        duplicate, detail = near_duplicate_to_history(sample, buckets)
        if duplicate:
            detail["sample_id"] = sample["sample_id"]
            overlaps.append(detail)
    return overlaps


def sample_distribution_summary(samples: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "source_type_distribution": dict(Counter(sample["source_type"] or "unknown" for sample in samples)),
        "event_type_top20": dict(Counter(sample["event_name"] or "unknown" for sample in samples).most_common(20)),
        "target_company_top20": dict(Counter(sample["target"] or "unknown" for sample in samples).most_common(20)),
        "publish_month_distribution": dict(Counter(month_bucket(sample["publish_time"]) for sample in samples)),
        "dataset_split_distribution": dict(Counter(sample["dataset_split"] or "unknown" for sample in samples)),
        "raw_market_context_count": sum(1 for sample in samples if sample.get("has_raw_market_terms")),
        "raw_financial_context_count": sum(1 for sample in samples if sample.get("has_raw_financial_terms")),
        "evidence_count_distribution": dict(Counter(str(sample.get("evidence_count") or 0) for sample in samples)),
    }


def write_sha256s(output_dir: Path, extra_paths: list[Path]) -> None:
    paths = [path for path in output_dir.glob("*") if path.is_file() and path.name != "SHA256SUMS.txt"]
    paths.extend(path for path in extra_paths if path.exists())
    lines = [f"{sha256(path)}  {rel(path)}" for path in sorted(set(paths))]
    (output_dir / "SHA256SUMS.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def ensure_no_existing_outputs(paths: Iterable[Path]) -> None:
    existing = [rel(path) for path in paths if path.exists()]
    if existing:
        raise FileExistsError(f"refusing to overwrite frozen or generated independent500-v2 outputs: {existing}")


def build_package(args: argparse.Namespace) -> dict[str, Any]:
    output_dir = args.output_dir
    locked_ids_config = args.locked_ids_config
    workbook = args.workbook
    guide = args.guide
    ensure_no_existing_outputs(
        [
            locked_ids_config,
            output_dir / "independent_samples.jsonl",
            output_dir / "leakage_audit.json",
            output_dir / "sampling_manifest.json",
            output_dir / "run_state.json",
            workbook,
            guide,
        ]
    )
    history = collect_history()
    article_map = load_article_map(args.article_index)
    raw_candidates = load_candidates(args.all_events, article_map)
    eligible, filter_summary, near_removed = filter_candidates(raw_candidates, article_map, history)
    samples, selection_summary = select_samples(eligible, TOTAL_COUNT, args.seed)
    near_overlaps = selected_near_duplicate_overlap(samples, history["profiles"])
    intersections = intersection_counts(samples, history["identifiers"])
    leak_visible = validate_no_visible_answer_leak(samples)

    leakage_audit = {
        "generated_at": now_iso(),
        "dataset_id": "task2_independent500_v2",
        "sample_count": len(samples),
        "unique_sample_id_count": len({sample["sample_id"] for sample in samples}),
        "unique_cluster_id_count": len({sample["cluster_id"] for sample in samples}),
        "unique_source_event_group_key_count": len({sample["source_event_group_key"] for sample in samples}),
        "exact_overlap_counts": intersections,
        "all_required_overlap_counts_zero": all(value == 0 for value in intersections.values()),
        "near_duplicate_overlap_count": len(near_overlaps),
        "near_duplicate_overlap_examples": near_overlaps[:20],
        "human_fields_blank": leak_visible["human_fields_blank"],
        "visible_answer_leak_count": leak_visible["visible_answer_leak_count"],
        "visible_answer_leak_examples": leak_visible["visible_answer_leak_examples"],
        "model_predictions_read": False,
        "old_gold_used_for_sampling": False,
        "old_gold_label_values_read_for_sampling": False,
        "training_performed": False,
        "tuning_performed": False,
        "locked500_gold_read_for_sampling": False,
        "independent500_v2_ids_frozen": True,
    }
    if not leakage_audit["all_required_overlap_counts_zero"] or leakage_audit["near_duplicate_overlap_count"] != 0:
        raise RuntimeError(f"leakage audit failed before freeze: {leakage_audit}")
    if not leakage_audit["human_fields_blank"] or leakage_audit["visible_answer_leak_count"] != 0:
        raise RuntimeError(f"annotation leak audit failed before freeze: {leakage_audit}")

    sampling_manifest = {
        "status": "independent500_v2_ready_for_blind_human_annotation",
        "generated_at": now_iso(),
        "project_root": str(ROOT),
        "dataset_id": "task2_independent500_v2",
        "seed": args.seed,
        "sample_count": len(samples),
        "source_inputs": {
            "all_events": rel(args.all_events),
            "article_index": rel(args.article_index),
        },
        "history_sources": history["source_stats"],
        "history_identifier_counts": {key: len(value) for key, value in history["identifiers"].items()},
        "history_category_identifier_counts": history["category_identifier_counts"],
        "candidate_filter_summary": filter_summary,
        "near_duplicate_exclusion_examples": near_removed[:50],
        "selection_summary": selection_summary,
        "distribution_summary": sample_distribution_summary(samples),
        "constraints": {
            "model_predictions_read": False,
            "old_gold_used_for_sampling": False,
            "old_gold_label_values_read_for_sampling": False,
            "training_performed": False,
            "tuning_performed": False,
            "new500_predictions_generated": False,
            "new500_gold_read": False,
            "selection_used_true_credibility_labels": False,
            "selection_used_model_scores_or_predictions": False,
        },
        "leakage_audit_path": rel(output_dir / "leakage_audit.json"),
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    sample_path = output_dir / "independent_samples.jsonl"
    write_jsonl(sample_path, samples)
    independent_sha = sha256(sample_path)
    locked_payload = {
        "generated_at": now_iso(),
        "dataset_id": "task2_independent500_v2",
        "seed": args.seed,
        "count": len(samples),
        "source": rel(sample_path),
        "source_sha256": independent_sha,
        "independent500_v2_ids_frozen": True,
        "sample_ids": [sample["sample_id"] for sample in samples],
        "cluster_ids": [sample["cluster_id"] for sample in samples],
        "source_event_ids": [sample["source_event_id"] for sample in samples],
        "claim_article_ids": [sample["claim_article_id"] for sample in samples],
        "constraints": {
            "do_not_modify_id_set_after_write": True,
            "gold_not_created": True,
            "model_predictions_not_generated": True,
            "training_performed": False,
            "tuning_performed": False,
        },
    }
    write_json(locked_ids_config, locked_payload)
    write_json(output_dir / "leakage_audit.json", leakage_audit)
    write_json(output_dir / "sampling_manifest.json", sampling_manifest)
    write_workbook(workbook, samples, sampling_manifest, leakage_audit)
    guide.parent.mkdir(parents=True, exist_ok=True)
    guide.write_text(render_guide(), encoding="utf-8")
    final_visible_check = validate_no_visible_answer_leak(samples, workbook)
    run_state = {
        "status": "completed",
        "generated_at": now_iso(),
        "dataset_id": "task2_independent500_v2",
        "training_performed": False,
        "tuning_performed": False,
        "locked500_gold_read_for_sampling": False,
        "new500_gold_read": False,
        "new500_predictions_generated": False,
        "model_predictions_read": False,
        "old_gold_used_for_sampling": False,
        "independent500_v2_ids_frozen": True,
        "validation": {
            "sample_count_500": len(samples) == TOTAL_COUNT,
            "unique_sample_id_500": len({sample["sample_id"] for sample in samples}) == TOTAL_COUNT,
            "exact_overlap_zero": leakage_audit["all_required_overlap_counts_zero"],
            "near_duplicate_overlap_zero": leakage_audit["near_duplicate_overlap_count"] == 0,
            "human_fields_blank": final_visible_check["human_fields_blank"],
            "visible_answer_leak_zero": final_visible_check["visible_answer_leak_count"] == 0,
        },
        "outputs": {
            "locked_ids_config": rel(locked_ids_config),
            "independent_samples": rel(sample_path),
            "leakage_audit": rel(output_dir / "leakage_audit.json"),
            "sampling_manifest": rel(output_dir / "sampling_manifest.json"),
            "workbook": rel(workbook),
            "guide": rel(guide),
            "sha256s": rel(output_dir / "SHA256SUMS.txt"),
        },
    }
    write_json(output_dir / "run_state.json", run_state)
    write_sha256s(output_dir, [locked_ids_config, workbook, guide])
    return {
        "run_state": run_state,
        "sampling_manifest": sampling_manifest,
        "leakage_audit": leakage_audit,
        "locked_ids_config_sha256": sha256(locked_ids_config),
        "independent_samples_sha256": independent_sha,
        "workbook_sha256": sha256(workbook),
        "guide_sha256": sha256(guide),
        "sha256s_sha256": sha256(output_dir / "SHA256SUMS.txt"),
    }


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--all-events", type=Path, default=ALL_EVENTS)
    parser.add_argument("--article-index", type=Path, default=ARTICLE_INDEX)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--locked-ids-config", type=Path, default=LOCKED_IDS_CONFIG)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--guide", type=Path, default=GUIDE)
    parser.add_argument("--seed", type=int, default=SEED)
    return parser


def absolutize(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    args.all_events = absolutize(args.all_events)
    args.article_index = absolutize(args.article_index)
    args.output_dir = absolutize(args.output_dir)
    args.locked_ids_config = absolutize(args.locked_ids_config)
    args.workbook = absolutize(args.workbook)
    args.guide = absolutize(args.guide)
    result = build_package(args)
    manifest = result["sampling_manifest"]
    print(
        json.dumps(
            {
                "status": result["run_state"]["status"],
                "sample_count": manifest["sample_count"],
                "distribution_summary": manifest["distribution_summary"],
                "overlap_counts": result["leakage_audit"]["exact_overlap_counts"],
                "near_duplicate_overlap_count": result["leakage_audit"]["near_duplicate_overlap_count"],
                "outputs": result["run_state"]["outputs"],
                "sha256": {
                    "locked_ids_config": result["locked_ids_config_sha256"],
                    "independent_samples": result["independent_samples_sha256"],
                    "workbook": result["workbook_sha256"],
                    "guide": result["guide_sha256"],
                    "SHA256SUMS": result["sha256s_sha256"],
                },
                "training_performed": False,
                "tuning_performed": False,
                "model_predictions_read": False,
                "old_gold_used_for_sampling": False,
            },
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
