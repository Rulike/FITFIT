#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Step 5.8-HB traceable mapping from incremental 45 labels to current official64 labels.

This script does not train models and does not modify the original 1,000-row workbook.
It preserves each original event_name and emits a separate mapped_event_name candidate.
"""

from __future__ import annotations

import hashlib
import json
import math
import pickle
import re
import shutil
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sklearn.feature_extraction.text import TfidfVectorizer


ROOT = Path(__file__).resolve().parents[1]
INCREMENTAL_XLSX = ROOT / "reports/财经新闻标注数据8.18.xlsx"
EXPECTED_INCREMENTAL_SHA = "a58bec4b2125a5b7b68f11f303d0c4fbaded88964a6724178dce5d7af594ffb4"
OFFICIAL_TRAIN_XLSX = ROOT / "data/raw/news_with_tags_train.xlsx"
OFFICIAL_STOCK_JSON = ROOT / "data/raw/事件类型_标的.json"
OFFICIAL_INDUSTRY_JSON = ROOT / "data/raw/事件类型_行业.json"
EVENT_SCHEMA_JSONL = ROOT / "data/processed/event_schema.jsonl"
FROZEN_EVENT_MODEL = ROOT / "outputs/task1_step5_8e_final_optimization/frozen_best_candidate/event_name_model.pkl"
EXPECTED_EVENT_MODEL_SHA = "f37ba4902784299ac17976f0fe059403d02befc15b18caa6d54ac4c5793b7c6e"

OUT_DIR = ROOT / "outputs/task1_step5_8hb_event_mapping"
REPORT_JSON = ROOT / "reports/task1_step5_8hb_event_mapping_report.json"
REPORT_MD = ROOT / "reports/task1_step5_8hb_event_mapping_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_8hb_event_mapping_details.xlsx"
REVIEW_XLSX = ROOT / "reports/task1_step5_8hb_event_mapping_human_review.xlsx"
CONFIG_YAML = ROOT / "configs/task1_step5_8hb_event_mapping_v1.yaml"

PROTECTED_CANDIDATES = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "outputs/task1_step5_8e_final_optimization/task1_official_delivery_candidate_v3.xlsx",
    ROOT / "outputs/task1_step5_8e_final_optimization/task1_official_delivery_candidate_v3.jsonl",
    ROOT / "outputs/task1_step5_8e_final_optimization/task1_official_delivery_candidate_v3.parquet",
    ROOT / "outputs/task1_step5_8f_final_candidate/task1_official_delivery_candidate_v4.xlsx",
    ROOT / "outputs/task1_step5_8f_final_candidate/task1_official_delivery_candidate_v4.jsonl",
    ROOT / "outputs/task1_step5_8f_final_candidate/task1_official_delivery_candidate_v4.parquet",
    ROOT / "reports/TASK1_TASK2_FINAL_ACCEPTANCE.md",
    ROOT / "reports/TASK1_FINAL_TECHNICAL_REPORT.md",
    ROOT / "reports/TASK2_FINAL_TECHNICAL_REPORT.md",
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_expected_sha(path: Path, expected: str, label: str) -> str:
    if not path.exists():
        raise FileNotFoundError(f"{label} missing: {path}")
    actual = sha256_file(path)
    if actual != expected:
        raise RuntimeError(f"{label} SHA mismatch: expected {expected}, got {actual}")
    return actual


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def json_dumps(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True)


def write_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json_dumps(obj) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return ""
    if pd.isna(value):
        return ""
    if isinstance(value, pd.Timestamp):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, np.generic):
        return value.item()
    return value


def normalize_text(value: Any) -> str:
    value = cell_value(value)
    if value == "":
        return ""
    text = str(value)
    return re.sub(r"\s+", " ", text).strip()


def content_sha(text: Any) -> str:
    return hashlib.sha256(normalize_text(text).encode("utf-8")).hexdigest()


def row_text_for_model(row: dict[str, Any]) -> str:
    parts = [
        str(row.get("original_event_name", "")),
        str(row.get("article_title", "")),
        str(row.get("target_name", "")),
        str(row.get("industry_name", "")),
        str(row.get("article_source", "")),
        str(row.get("content", ""))[:3600],
    ]
    return "\n".join(p for p in parts if p)


def safe_sheet_name(name: str) -> str:
    return name[:31]


def format_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column_cells in ws.columns:
            length = 8
            for cell in column_cells[:200]:
                if cell.value is not None:
                    length = max(length, min(45, len(str(cell.value)) + 2))
            ws.column_dimensions[column_cells[0].column_letter].width = length
    wb.save(path)


def collect_protected_sha() -> dict[str, str]:
    out: dict[str, str] = {}
    for path in PROTECTED_CANDIDATES:
        if path.exists():
            out[str(path.relative_to(ROOT))] = sha256_file(path)
    return out


def create_backup_dir(generated_at: str, protected_sha: dict[str, str]) -> tuple[str, str]:
    external = ROOT.parent / "backups" / f"task1_step5_8hb_event_mapping_{generated_at.replace(':', '').replace('-', '').replace('T', '_')}"
    local = OUT_DIR / "backup" / f"task1_step5_8hb_event_mapping_{generated_at.replace(':', '').replace('-', '').replace('T', '_')}"
    try:
        external.mkdir(parents=True, exist_ok=True)
        target = external
        status = "external_backup_created"
    except Exception as exc:
        target = local
        target.mkdir(parents=True, exist_ok=True)
        status = f"project_local_backup_created_after_external_failure: {type(exc).__name__}"
    write_json(target / "protected_sha256_before.json", protected_sha)
    return str(target), status


def load_incremental_rows() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    actual_sha = ensure_expected_sha(INCREMENTAL_XLSX, EXPECTED_INCREMENTAL_SHA, "incremental workbook")
    xls = pd.ExcelFile(INCREMENTAL_XLSX)
    expected_sheets = {"个股新闻", "行业新闻"}
    missing = expected_sheets - set(xls.sheet_names)
    if missing:
        raise RuntimeError(f"incremental workbook missing sheets: {sorted(missing)}")
    rows: list[dict[str, Any]] = []
    sheet_stats: dict[str, Any] = {}
    expected_cols = {
        "个股新闻": [
            "article_file_id",
            "article_title",
            "article_publish_time",
            "article_source",
            "content",
            "trading_code",
            "secu_abbr",
            "industry_code",
            "industry_name",
            "event_name",
            "event_emotion",
            "event_impact_analysis",
            "duplication_id",
        ],
        "行业新闻": [
            "article_file_id",
            "article_title",
            "article_publish_time",
            "article_source",
            "content",
            "industry_code",
            "industry_name",
            "event_name",
            "event_emotion",
            "event_impact_analysis",
            "duplication_id",
        ],
    }
    for sheet in ["个股新闻", "行业新闻"]:
        df = pd.read_excel(INCREMENTAL_XLSX, sheet_name=sheet, dtype=object)
        cols = list(df.columns)
        if cols != expected_cols[sheet]:
            raise RuntimeError(f"{sheet} columns differ from expected official-format columns: {cols}")
        target_type = "company" if sheet == "个股新闻" else "industry"
        for idx, raw in df.iterrows():
            record = {col: cell_value(raw.get(col)) for col in cols}
            event_name = normalize_text(record.get("event_name"))
            target_name = normalize_text(record.get("secu_abbr" if target_type == "company" else "industry_name"))
            record.update(
                {
                    "row_number_1based": int(idx) + 2,
                    "original_sheet": sheet,
                    "target_type": target_type,
                    "target_name": target_name,
                    "original_event_name": event_name,
                    "content_sha256": content_sha(record.get("content")),
                    "mapping_input_policy": "official_incremental_1000_original_event_name_preserved",
                }
            )
            rows.append(record)
        sheet_stats[sheet] = {"rows": len(df), "columns": cols}
    article_ids = [normalize_text(r.get("article_file_id")) for r in rows]
    if len(rows) != 1000:
        raise RuntimeError(f"expected 1000 incremental rows, got {len(rows)}")
    if len(set(article_ids)) != 1000 or any(not x for x in article_ids):
        raise RuntimeError("article_file_id must be unique and non-empty for all 1000 rows")
    stats = {
        "sha256": actual_sha,
        "rows": len(rows),
        "sheet_stats": sheet_stats,
        "unique_article_file_id": len(set(article_ids)),
        "source_event_name_count": len({r["original_event_name"] for r in rows}),
        "source_event_name_distribution": dict(Counter(r["original_event_name"] for r in rows).most_common()),
        "target_type_distribution": dict(Counter(r["target_type"] for r in rows)),
        "event_emotion_distribution": dict(Counter(normalize_text(r.get("event_emotion")) for r in rows)),
    }
    return rows, stats


def load_official_catalog() -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, set[str]]]:
    schema_rows = read_jsonl(EVENT_SCHEMA_JSONL)
    official_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    allowed_by_type: dict[str, set[str]] = {"company": set(), "industry": set()}
    for row in schema_rows:
        event_name = normalize_text(row.get("event_name"))
        target_type = normalize_text(row.get("target_type"))
        official_by_name[event_name].append(row)
        if target_type in allowed_by_type:
            allowed_by_type[target_type].add(event_name)

    official_train_examples: dict[str, list[dict[str, Any]]] = defaultdict(list)
    official_train_counts: Counter[str] = Counter()
    for sheet, target_type in [("个股新闻", "company"), ("行业新闻", "industry")]:
        df = pd.read_excel(OFFICIAL_TRAIN_XLSX, sheet_name=sheet, dtype=object)
        for _, raw in df.iterrows():
            event_name = normalize_text(raw.get("event_name"))
            official_train_counts[event_name] += 1
            if len(official_train_examples[event_name]) < 5:
                official_train_examples[event_name].append(
                    {
                        "sheet": sheet,
                        "target_type": target_type,
                        "article_title": normalize_text(raw.get("article_title"))[:160],
                        "target_name": normalize_text(raw.get("secu_abbr" if target_type == "company" else "industry_name")),
                        "content_excerpt": normalize_text(raw.get("content"))[:240],
                    }
                )

    catalog: list[dict[str, Any]] = []
    for event_name in sorted(official_by_name):
        rows = official_by_name[event_name]
        descriptions = [normalize_text(r.get("raw_description")) for r in rows if normalize_text(r.get("raw_description"))]
        target_types = sorted({normalize_text(r.get("target_type")) for r in rows})
        keywords = extract_keywords(event_name + " " + " ".join(descriptions[:3]))
        catalog.append(
            {
                "event_name": event_name,
                "target_types": target_types,
                "schema_ids": [normalize_text(r.get("schema_id")) for r in rows],
                "descriptions": descriptions[:8],
                "description_excerpt": descriptions[0][:800] if descriptions else "",
                "event_examples": official_train_examples.get(event_name, []),
                "official_train_count": int(official_train_counts[event_name]),
                "keywords_and_action_words": keywords,
                "allowed_subject_type": "上市公司" if target_types == ["company"] else "行业" if target_types == ["industry"] else "上市公司或行业",
            }
        )
    if len({c["event_name"] for c in catalog}) != 64:
        raise RuntimeError(f"official catalog expected 64 names, got {len(catalog)}")
    return catalog, {"schema_rows": schema_rows, "by_name": official_by_name}, allowed_by_type


def extract_keywords(text: str) -> list[str]:
    vocab = [
        "合同", "中标", "项目", "投资", "产能", "扩产", "基地", "技术", "研发", "创新", "突破", "产品", "发布", "升级",
        "政策", "监管", "调控", "支持", "价格", "涨价", "降价", "需求", "供给", "销量", "市场", "竞争", "现金流",
        "债务", "违约", "资本", "重组", "并购", "融资", "股权", "激励", "回购", "减持", "增持", "质押", "分红",
        "出海", "海外", "全球化", "地缘", "电力", "储能", "动力电池", "光伏", "医药", "创新药", "教育", "影视",
        "农业", "风电", "水产", "OPEC", "利率", "基建", "商业航天", "低空经济",
    ]
    found = [word for word in vocab if word in text]
    return sorted(set(found), key=found.index)


def spec(
    mapping_type: str,
    semantic_relation: str,
    candidates: list[str],
    confidence: float,
    rationale: str,
    usable_for_official_classification: bool = False,
    content_scope_all_same_target: bool = False,
) -> dict[str, Any]:
    return {
        "mapping_type": mapping_type,
        "semantic_relation": semantic_relation,
        "candidates": candidates,
        "confidence": confidence,
        "rationale": rationale,
        "usable_for_official_classification": usable_for_official_classification,
        "content_scope_all_same_target": content_scope_all_same_target,
    }


def label_mapping_specs() -> dict[str, dict[str, Any]]:
    company_capital = ["资本市场运作"]
    company_growth = ["新业务拓展", "业务拓展", "重大业务拓展", "公司战略调整", "结构升级"]
    company_risk = ["现金流风险", "市场竞争加剧", "公司战略调整"]
    company_tech = ["产品技术创新", "技术突破与创新", "关键技术突破", "技术研发进展", "新产品升级发布", "医药研发突破"]
    company_project = ["签订重要合同", "重大项目投资建设", "重大业务拓展"]
    company_strategy = ["公司战略调整", "公司战略变化及对外投资", "业务战略调整", "企业转型催化"]
    industry_policy = [
        "行业政策", "行业政策推动", "行业调控政策发布", "金融政策支持", "稳增长政策", "绿色政策发布",
        "电力投资政策", "风电产业政策支持", "国内创新药支持政策", "教育行业政策调整", "影视行业政策变动",
        "农业政策变动", "工业自动化政策支持变化",
    ]
    industry_market = [
        "水产品价格变动", "OPEC+产量政策变化", "影视市场表现变化", "行业下游的需求变化", "市场竞争",
        "市场竞争加剧", "大宗商品景气与价格波动加剧", "行业政策",
    ]
    return {
        "IPO上市": spec("one_to_many_content_required", "partial_overlap", company_capital + ["新业务拓展"], 0.58, "上市是资本市场行为，但官方资本市场运作定义更偏公司自身重组、再融资、股权激励等，需逐条核对。"),
        "业绩下降": spec("one_to_many_content_required", "partial_overlap", company_risk + ["业务战略调整"], 0.48, "业绩变化不是官方64类的独立名称，只有正文体现竞争、现金流或战略调整时才可映射。", content_scope_all_same_target=True),
        "业绩亏损": spec("one_to_many_content_required", "partial_overlap", company_risk, 0.50, "亏损本身不是官方事件名，需判断是否符合现金流风险、市场竞争加剧等官方成立条件。", content_scope_all_same_target=True),
        "业绩增长": spec("one_to_many_content_required", "partial_overlap", company_growth, 0.50, "业绩增长是结果性标签，需根据正文业务扩张、结构升级等事实选择官方类别。", content_scope_all_same_target=True),
        "中标项目": spec("one_to_one_safe", "source_narrower", ["签订重要合同", "重大项目投资建设"], 0.91, "中标项目通常是重要合同或项目建设的明确子类，但合同金额、主体和项目属性需行级确认。", True),
        "产品发布": spec("one_to_many_content_required", "source_narrower", company_tech + ["动力电池产品迭代", "储能电池产品迭代", "新车发布定价"], 0.82, "产品发布与产品技术创新/产品迭代相关，但不同公司schema粒度不同，需正文和主体约束。"),
        "产品认证": spec("one_to_many_content_required", "source_narrower", company_tech, 0.86, "认证可支持技术或产品创新，但需确认认证对象和官方事件定义兼容。"),
        "产能扩张": spec("one_to_one_safe", "source_narrower", ["产能布局调整", "重大项目投资建设", "海外产能进展"], 0.90, "扩产、基地投建与产能布局/重大项目/海外产能存在明确语义子类关系。", True),
        "价格波动": spec("one_to_many_content_required", "partial_overlap", ["大宗商品景气与价格波动加剧"] + industry_market, 0.72, "价格波动可对应部分行业价格事件，但需识别具体品类和行业。"),
        "供需变化": spec("one_to_many_content_required", "partial_overlap", ["行业下游的需求变化", "市场竞争", "市场竞争加剧"] + industry_market, 0.70, "供需变化可能落入行业需求、竞争或价格类官方事件，必须逐条判定。"),
        "分红送配": spec("one_to_one_safe", "partial_overlap", company_capital, 0.72, "分红送配是资本市场相关事项，但不在官方定义示例核心范围内，需复核后低权重使用。"),
        "创历史新低": spec("no_official_equivalent", "partial_overlap", company_capital + company_risk + industry_market, 0.35, "股价或指数创新低是行情结果，不等同官方公司/行业事件。"),
        "创历史新高": spec("no_official_equivalent", "partial_overlap", company_capital + industry_market, 0.35, "股价或指数创新高是行情结果，不等同官方公司/行业事件。"),
        "合同签订": spec("one_to_one_safe", "source_narrower", ["签订重要合同", "重大业务拓展"], 0.93, "签约事实与官方签订重要合同/重大业务拓展高度相关，仍需金额、主体和合同性质约束。", True),
        "大股东减持": spec("one_to_one_safe", "partial_overlap", company_capital, 0.66, "股东减持是资本市场相关行为，但官方定义更偏公司自身资本运作，必须复核。"),
        "大股东增持": spec("one_to_one_safe", "partial_overlap", company_capital, 0.66, "股东增持是资本市场相关行为，但官方定义更偏公司自身资本运作，必须复核。"),
        "安全事故": spec("no_official_equivalent", "partial_overlap", company_risk + ["公共卫生事件"], 0.30, "事故类事件未在官方64类中形成明确对应，不能硬映射。"),
        "定增募资": spec("many_to_one_safe", "source_narrower", company_capital, 0.92, "定增募资属于公司利用资本市场工具融资，落入资本市场运作的子类。", True),
        "并购重组": spec("many_to_one_safe", "source_narrower", company_capital + ["公司战略变化及对外投资"], 0.91, "并购重组与官方资本市场运作/战略变化定义中的并购重组高度一致。", True),
        "战略合作": spec("one_to_many_content_required", "partial_overlap", ["签订重要合同", "重大业务拓展", "业务拓展", "公司战略变化及对外投资"], 0.78, "合作可能是合同、业务拓展或战略投资，需逐条按事实映射。"),
        "技术突破": spec("one_to_many_content_required", "source_narrower", company_tech + ["钢铁生产核心技术突破", "光伏新技术迭代升级", "光伏组件行业技术迭代"], 0.88, "技术突破与官方技术创新体系高度相关，但需按主体和领域细分。"),
        "政策利好": spec("one_to_many_content_required", "partial_overlap", industry_policy + ["电力投资政策"], 0.74, "政策利好是影响方向而非事件名，需逐条识别具体政策类官方事件。"),
        "政策监管": spec("one_to_many_content_required", "partial_overlap", industry_policy, 0.74, "监管政策需映射到具体行业政策/调控政策/行业变动类事件。"),
        "涨停板": spec("no_official_equivalent", "partial_overlap", company_capital + industry_market, 0.28, "涨停是行情表现，不是官方事件成立条件。"),
        "监管处罚": spec("one_to_many_content_required", "partial_overlap", company_risk + industry_policy, 0.45, "处罚可能导致风险或监管变化，但官方64类无直接处罚类别。"),
        "股东大会": spec("no_official_equivalent", "partial_overlap", company_capital + company_strategy, 0.32, "股东大会是公司治理会议，不直接等价官方资本或战略事件。"),
        "股价上涨": spec("no_official_equivalent", "partial_overlap", company_capital + industry_market, 0.25, "股价上涨是市场价格结果，不能作为官方事件名硬映射。"),
        "股价下跌": spec("no_official_equivalent", "partial_overlap", company_capital + company_risk + industry_market, 0.25, "股价下跌是市场价格结果，不能作为官方事件名硬映射。"),
        "股价异动": spec("no_official_equivalent", "partial_overlap", company_capital + industry_market, 0.25, "股价异动是行情信号，非官方事件成立条件。"),
        "股份质押": spec("one_to_many_content_required", "partial_overlap", company_capital + company_risk, 0.46, "股份质押可能与资本事项或现金流压力相关，但官方定义并非质押事件。"),
        "股权变动": spec("one_to_many_content_required", "partial_overlap", company_capital + company_strategy, 0.58, "股权变动需要区分并购重组、战略调整、资本运作等不同官方事件。"),
        "股权激励": spec("many_to_one_safe", "source_narrower", company_capital, 0.94, "股权激励在官方资本市场运作定义中明确列出。", True),
        "股票回购": spec("many_to_one_safe", "partial_overlap", company_capital, 0.76, "回购是资本市场行为，但官方定义示例未明确列回购，需复核。"),
        "融资融券": spec("one_to_many_content_required", "partial_overlap", company_capital + ["金融政策支持", "行业政策"], 0.38, "融资融券常是市场交易数据或政策环境，不能直接作为公司事件。"),
        "行业数据": spec("one_to_many_content_required", "partial_overlap", industry_market + industry_policy + company_growth + company_risk, 0.30, "行业数据是泛化观察口径，需要逐条依据正文事实选择官方行业/公司事件，低置信。", content_scope_all_same_target=True),
        "诉讼纠纷": spec("no_official_equivalent", "partial_overlap", company_risk + company_strategy, 0.36, "诉讼纠纷未在官方64类中形成独立类别，仅在产生现金流或战略影响时部分重合。"),
        "资金流入": spec("no_official_equivalent", "partial_overlap", company_capital + ["金融政策支持", "行业政策"], 0.25, "资金流入通常是行情资金面信号，非官方事件名。"),
        "资金流出": spec("no_official_equivalent", "partial_overlap", company_capital + ["金融政策支持", "行业政策"], 0.25, "资金流出通常是行情资金面信号，非官方事件名。"),
        "跌停板": spec("no_official_equivalent", "partial_overlap", company_capital + company_risk + industry_market, 0.28, "跌停是行情表现，不是官方事件成立条件。"),
        "退市风险": spec("one_to_many_content_required", "partial_overlap", company_risk + company_capital, 0.44, "退市风险可能由现金流、经营或资本问题触发，官方无直接退市风险类别。"),
        "退市风险警示": spec("one_to_many_content_required", "partial_overlap", company_risk + company_capital, 0.44, "退市风险警示是交易所状态，不等同官方64类，需要证据支持。"),
        "配售股份": spec("many_to_one_safe", "source_narrower", company_capital, 0.90, "配售股份属于资本市场融资/股本安排，与资本市场运作有明确子类关系。", True),
        "销量增长": spec("one_to_many_content_required", "partial_overlap", company_growth + industry_market, 0.54, "销量增长是结果指标，需识别是否支持业务拓展、结构升级或行业需求变化。"),
        "高管变动": spec("no_official_equivalent", "partial_overlap", company_strategy, 0.32, "高管变动不直接等价官方公司战略调整，除非正文同时给出战略事实。"),
        "龙虎榜": spec("no_official_equivalent", "partial_overlap", company_capital, 0.20, "龙虎榜是交易公开信息，不是官方事件类别。"),
    }


def build_official_candidate(spec_row: dict[str, Any], event_name: str, catalog_by_name: dict[str, dict[str, Any]]) -> dict[str, Any]:
    cat = catalog_by_name[event_name]
    return {
        "schema_id": cat["schema_ids"][0] if cat["schema_ids"] else None,
        "target_types": cat["target_types"],
        "event_name": event_name,
        "relation": spec_row["semantic_relation"],
        "confidence": spec_row["confidence"],
        "mapping_reason": spec_row["rationale"],
        "definition_evidence": cat.get("description_excerpt", "")[:500],
    }


def load_frozen_event_model(allowed_by_type: dict[str, set[str]]) -> Any:
    ensure_expected_sha(FROZEN_EVENT_MODEL, EXPECTED_EVENT_MODEL_SHA, "frozen event_name model")
    with FROZEN_EVENT_MODEL.open("rb") as f:
        payload = pickle.load(f)
    model = payload.get("model") if isinstance(payload, dict) else payload
    if not hasattr(model, "predict_proba"):
        raise RuntimeError("frozen event_name model does not expose predict_proba")
    classes = list(getattr(model, "classes_", []))
    if not classes and hasattr(model, "named_steps"):
        clf = model.named_steps.get("classifier") or model.named_steps.get("clf")
        classes = list(getattr(clf, "classes_", []))
    official_allowed = set().union(*allowed_by_type.values())
    missing = [c for c in classes if c not in official_allowed]
    if missing:
        raise RuntimeError(f"frozen event model classes contain non-official labels: {missing[:5]}")
    return model


def model_topk(model: Any, texts: list[str], target_types: list[str], allowed_by_type: dict[str, set[str]], k: int = 20) -> list[list[dict[str, Any]]]:
    probs = model.predict_proba(texts)
    classes = list(getattr(model, "classes_", []))
    if not classes and hasattr(model, "named_steps"):
        clf = model.named_steps.get("classifier") or model.named_steps.get("clf")
        classes = list(getattr(clf, "classes_", []))
    out: list[list[dict[str, Any]]] = []
    for row_probs, target_type in zip(probs, target_types):
        allowed = allowed_by_type[target_type]
        ranked = []
        for idx in np.argsort(row_probs)[::-1]:
            name = str(classes[int(idx)])
            if name in allowed:
                ranked.append({"event_name": name, "score": float(row_probs[int(idx)])})
            if len(ranked) >= k:
                break
        out.append(ranked)
    return out


def build_similarity_scores(rows: list[dict[str, Any]], catalog: list[dict[str, Any]]) -> tuple[list[list[dict[str, Any]]], dict[str, int]]:
    official_texts: list[str] = []
    names: list[str] = []
    for cat in catalog:
        names.append(cat["event_name"])
        official_texts.append(
            " ".join(
                [
                    cat["event_name"],
                    " ".join(cat.get("keywords_and_action_words") or []),
                    " ".join(cat.get("descriptions") or []),
                    " ".join(ex.get("article_title", "") + " " + ex.get("content_excerpt", "") for ex in cat.get("event_examples", [])),
                ]
            )
        )
    query_texts = [
        " ".join(
            [
                str(r.get("original_event_name", "")),
                str(r.get("article_title", "")),
                str(r.get("content", ""))[:3000],
            ]
        )
        for r in rows
    ]
    vectorizer = TfidfVectorizer(analyzer="char", ngram_range=(2, 4), min_df=1)
    matrix = vectorizer.fit_transform(official_texts + query_texts)
    official_matrix = matrix[: len(official_texts)]
    query_matrix = matrix[len(official_texts) :]
    sims = query_matrix @ official_matrix.T
    rank_by_name = {name: i for i, name in enumerate(names)}
    ranked_rows: list[list[dict[str, Any]]] = []
    for i in range(sims.shape[0]):
        dense = np.asarray(sims[i].todense()).ravel()
        order = np.argsort(dense)[::-1]
        ranked_rows.append([{"event_name": names[int(j)], "score": float(dense[int(j)])} for j in order[:20]])
    return ranked_rows, rank_by_name


def choose_mapping(
    row: dict[str, Any],
    spec_row: dict[str, Any],
    allowed_by_type: dict[str, set[str]],
    frozen_top: list[dict[str, Any]],
    sim_top: list[dict[str, Any]],
) -> dict[str, Any]:
    target_type = row["target_type"]
    allowed = allowed_by_type[target_type]
    rule_candidates = [c for c in spec_row["candidates"] if c in allowed]
    all_rule_candidates = [c for c in spec_row["candidates"]]
    target_type_conflict = bool(all_rule_candidates) and not rule_candidates and not spec_row.get("content_scope_all_same_target")
    if spec_row.get("content_scope_all_same_target"):
        candidate_pool = sorted(allowed)
        pool_reason = "all_official_events_for_same_target_type_due_to_broad_source_label"
    else:
        candidate_pool = rule_candidates
        pool_reason = "label_level_candidate_pool"
    model_scores = {x["event_name"]: float(x["score"]) for x in frozen_top}
    sim_scores = {x["event_name"]: float(x["score"]) for x in sim_top}

    selected = None
    selected_reason = ""
    if candidate_pool:
        candidates_ranked = []
        for candidate in candidate_pool:
            rule_score = 1.0 if candidate in rule_candidates else 0.15
            score = 0.44 * rule_score + 0.34 * model_scores.get(candidate, 0.0) + 0.22 * sim_scores.get(candidate, 0.0)
            if frozen_top and candidate == frozen_top[0]["event_name"]:
                score += 0.08
            if sim_top and candidate == sim_top[0]["event_name"]:
                score += 0.04
            candidates_ranked.append((score, candidate))
        candidates_ranked.sort(reverse=True)
        selected = candidates_ranked[0][1]
        selected_reason = "highest_fused_rule_model_description_score"
    elif frozen_top and not target_type_conflict:
        selected = frozen_top[0]["event_name"]
        selected_reason = "fallback_frozen_model_same_target_top1_for_review_only"

    model_top1 = frozen_top[0]["event_name"] if frozen_top else ""
    model_conf = float(frozen_top[0]["score"]) if frozen_top else 0.0
    model_gap = float(frozen_top[0]["score"] - frozen_top[1]["score"]) if len(frozen_top) > 1 else model_conf
    sim_top1 = sim_top[0]["event_name"] if sim_top else ""
    sim_selected_score = sim_scores.get(selected or "", 0.0)

    if not selected:
        status = "target_type_conflict" if target_type_conflict else "unresolved"
        confidence = 0.0
    elif target_type_conflict:
        status = "target_type_conflict"
        confidence = 0.0
    elif spec_row["mapping_type"] in {"one_to_one_safe", "many_to_one_safe"} and selected in rule_candidates and spec_row["confidence"] >= 0.90:
        status = "auto_mapped_high_confidence"
        confidence = min(0.98, max(spec_row["confidence"], 0.90 + min(model_conf, 0.07)))
    elif (
        spec_row["mapping_type"] in {"one_to_many_content_required"}
        and selected in rule_candidates
        and model_top1 == selected
        and model_conf >= 0.12
        and model_gap >= 0.03
        and spec_row["semantic_relation"] in {"source_narrower", "partial_overlap"}
        and spec_row["confidence"] >= 0.80
    ):
        status = "auto_mapped_high_confidence"
        confidence = min(0.92, 0.78 + model_conf + min(model_gap, 0.08))
    elif spec_row["mapping_type"] == "no_official_equivalent":
        status = "mapped_needs_review" if selected else "unresolved"
        confidence = min(0.55, 0.20 + model_conf + sim_selected_score)
    else:
        status = "mapped_needs_review"
        confidence = min(0.82, 0.45 + model_conf + sim_selected_score + (0.10 if selected in rule_candidates else 0.0))

    candidate_top5 = []
    seen = set()
    for source, items in [("label_rule", [{"event_name": c, "score": 1.0} for c in rule_candidates]), ("frozen_model", frozen_top), ("description_similarity", sim_top)]:
        for item in items:
            name = item["event_name"]
            if name not in allowed or name in seen:
                continue
            seen.add(name)
            candidate_top5.append({"event_name": name, "source": source, "score": float(item.get("score") or 0.0)})
            if len(candidate_top5) >= 5:
                break
        if len(candidate_top5) >= 5:
            break

    return {
        "mapped_event_name": selected,
        "mapping_status": status,
        "mapping_confidence": round(float(confidence), 6),
        "mapping_type": spec_row["mapping_type"],
        "semantic_relation": spec_row["semantic_relation"],
        "candidate_top5": candidate_top5,
        "candidate_scores": {
            "pool_reason": pool_reason,
            "frozen_model_top1": model_top1,
            "frozen_model_top1_score": model_conf,
            "frozen_model_top1_top2_gap": model_gap,
            "description_similarity_top1": sim_top1,
            "description_similarity_selected_score": sim_selected_score,
            "selected_reason": selected_reason,
        },
        "mapping_evidence": {
            "label_level_rationale": spec_row["rationale"],
            "row_title": normalize_text(row.get("article_title"))[:200],
            "row_content_excerpt": normalize_text(row.get("content"))[:360],
            "impact_analysis_excerpt": normalize_text(row.get("event_impact_analysis"))[:220],
            "target_type": target_type,
            "target_name": normalize_text(row.get("target_name")),
        },
        "classification_loss_weight": (
            1.0
            if status == "auto_mapped_high_confidence" and spec_row["semantic_relation"] == "exact_equivalent"
            else 0.5
            if status == "auto_mapped_high_confidence" and spec_row["semantic_relation"] == "source_narrower"
            else 0.0
        ),
        "usable_for_official_classification": status == "auto_mapped_high_confidence",
        "target_type_conflict": target_type_conflict,
    }


def build_label_level_mapping(
    rows: list[dict[str, Any]],
    specs: dict[str, dict[str, Any]],
    catalog: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    catalog_by_name = {c["event_name"]: c for c in catalog}
    official_names = set(catalog_by_name)
    source_counts = Counter(r["original_event_name"] for r in rows)
    sheet_dist: dict[str, Counter[str]] = defaultdict(Counter)
    emotion_dist: dict[str, Counter[str]] = defaultdict(Counter)
    sample_rows: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        label = row["original_event_name"]
        sheet_dist[label][row["original_sheet"]] += 1
        emotion_dist[label][normalize_text(row.get("event_emotion"))] += 1
        if len(sample_rows[label]) < 5:
            sample_rows[label].append(
                {
                    "article_file_id": normalize_text(row.get("article_file_id")),
                    "target_type": row["target_type"],
                    "target_name": normalize_text(row.get("target_name")),
                    "article_title": normalize_text(row.get("article_title"))[:180],
                    "content_excerpt": normalize_text(row.get("content"))[:240],
                    "impact_analysis_excerpt": normalize_text(row.get("event_impact_analysis"))[:180],
                }
            )
    missing_specs = sorted(set(source_counts) - set(specs))
    if missing_specs:
        raise RuntimeError(f"manual mapping spec missing source labels: {missing_specs}")
    label_rows: list[dict[str, Any]] = []
    for label in sorted(source_counts):
        spec_row = specs[label]
        invalid_candidates = [c for c in spec_row["candidates"] if c not in official_names]
        if invalid_candidates:
            raise RuntimeError(f"{label}: invalid official candidates {invalid_candidates}")
        candidate_rows = [build_official_candidate(spec_row, c, catalog_by_name) for c in spec_row["candidates"]]
        selected = None
        usable = False
        if spec_row["mapping_type"] in {"one_to_one_safe", "many_to_one_safe"} and len(spec_row["candidates"]) == 1 and spec_row["confidence"] >= 0.90:
            selected = spec_row["candidates"][0]
            usable = bool(spec_row["usable_for_official_classification"])
        label_rows.append(
            {
                "source_event_name": label,
                "record_count": int(source_counts[label]),
                "sheet_distribution": dict(sheet_dist[label]),
                "emotion_distribution": dict(emotion_dist[label]),
                "mapping_type": spec_row["mapping_type"],
                "semantic_relation": spec_row["semantic_relation"],
                "candidate_official_events": candidate_rows,
                "selected_schema_id": catalog_by_name[selected]["schema_ids"][0] if selected else None,
                "selected_event_name": selected,
                "usable_for_official_classification": usable,
                "usable_for_shared_extraction_pretraining": True,
                "requires_review": True,
                "confidence": spec_row["confidence"],
                "rationale": spec_row["rationale"],
                "representative_samples": sample_rows[label],
            }
        )
    return label_rows


def build_row_level_mapping(
    rows: list[dict[str, Any]],
    specs: dict[str, dict[str, Any]],
    catalog: list[dict[str, Any]],
    allowed_by_type: dict[str, set[str]],
) -> list[dict[str, Any]]:
    model = load_frozen_event_model(allowed_by_type)
    texts = [row_text_for_model(r) for r in rows]
    target_types = [r["target_type"] for r in rows]
    frozen_top = model_topk(model, texts, target_types, allowed_by_type, k=20)
    sim_top, _ = build_similarity_scores(rows, catalog)
    mapped_rows: list[dict[str, Any]] = []
    for row, top, sim in zip(rows, frozen_top, sim_top):
        spec_row = specs[row["original_event_name"]]
        chosen = choose_mapping(row, spec_row, allowed_by_type, top, sim)
        out = dict(row)
        out.update(chosen)
        out["event_name_for_64_training"] = chosen["mapped_event_name"] if chosen["usable_for_official_classification"] else None
        out["original_event_name_preserved"] = True
        mapped_rows.append(out)
    return mapped_rows


def sample_audit_rows(mapped_rows: list[dict[str, Any]], n: int = 100) -> list[dict[str, Any]]:
    rng = np.random.default_rng(2026)
    buckets: dict[tuple[str, str, str], list[int]] = defaultdict(list)
    for i, row in enumerate(mapped_rows):
        buckets[(row["mapping_status"], row["target_type"], row["original_event_name"])].append(i)
    selected: list[int] = []
    for key in sorted(buckets):
        selected.append(int(rng.choice(buckets[key])))
        if len(selected) >= n:
            break
    remaining = [i for i in range(len(mapped_rows)) if i not in set(selected)]
    if len(selected) < n and remaining:
        extra = rng.choice(remaining, size=min(n - len(selected), len(remaining)), replace=False)
        selected.extend(int(x) for x in extra)
    audit: list[dict[str, Any]] = []
    for idx in selected[:n]:
        row = mapped_rows[idx]
        cands = [x["event_name"] for x in row.get("candidate_top5") or []]
        frozen_top1 = row.get("candidate_scores", {}).get("frozen_model_top1")
        desc_top1 = row.get("candidate_scores", {}).get("description_similarity_top1")
        selected_name = row.get("mapped_event_name")
        direct_rule = row.get("mapping_type") in {"one_to_one_safe", "many_to_one_safe"} and row.get("mapping_status") == "auto_mapped_high_confidence"
        multi_signal = selected_name and selected_name in cands and (selected_name == frozen_top1 or selected_name == desc_top1)
        audit.append(
            {
                "article_file_id": row.get("article_file_id"),
                "target_type": row.get("target_type"),
                "original_event_name": row.get("original_event_name"),
                "mapped_event_name": selected_name,
                "mapping_status": row.get("mapping_status"),
                "mapping_confidence": row.get("mapping_confidence"),
                "frozen_top1": frozen_top1,
                "description_top1": desc_top1,
                "multi_method_consistent": bool(direct_rule or multi_signal),
                "manual_review_required": row.get("mapping_status") != "auto_mapped_high_confidence",
                "article_title": normalize_text(row.get("article_title"))[:180],
            }
        )
    return audit


def compute_mapping_stats(mapped_rows: list[dict[str, Any]], label_rows: list[dict[str, Any]], audit_rows: list[dict[str, Any]]) -> dict[str, Any]:
    status_counts = Counter(r["mapping_status"] for r in mapped_rows)
    target_status = defaultdict(Counter)
    for row in mapped_rows:
        target_status[row["target_type"]][row["mapping_status"]] += 1
    mapped_names = [r["mapped_event_name"] for r in mapped_rows if r.get("mapped_event_name")]
    high_conf = [r for r in mapped_rows if r["mapping_status"] == "auto_mapped_high_confidence"]
    high_names = [r["mapped_event_name"] for r in high_conf if r.get("mapped_event_name")]
    class_absorption = Counter(mapped_names).most_common(10)
    label_mapping_type_counts = Counter(row["mapping_type"] for row in label_rows)
    audit_consistency = sum(1 for r in audit_rows if r["multi_method_consistent"]) / len(audit_rows) if audit_rows else 0.0
    gate = {
        "high_confidence_mapping_records": len(high_conf),
        "high_confidence_mapping_records_required": 850,
        "stratified_auto_audit_consistency": audit_consistency,
        "stratified_auto_audit_consistency_required": 0.95,
        "illegal_mapped_event_name": 0,
        "target_type_conflict_records": status_counts.get("target_type_conflict", 0),
        "target_type_conflict_required": 0,
        "unresolved_records": status_counts.get("unresolved", 0),
        "unresolved_required_max": 50,
        "single_class_absorption_max_share": (class_absorption[0][1] / len(mapped_rows)) if class_absorption else 0.0,
        "single_class_absorption_flag": bool(class_absorption and class_absorption[0][1] / len(mapped_rows) > 0.35),
    }
    gate["passed"] = (
        gate["high_confidence_mapping_records"] >= gate["high_confidence_mapping_records_required"]
        and gate["stratified_auto_audit_consistency"] >= gate["stratified_auto_audit_consistency_required"]
        and gate["illegal_mapped_event_name"] == 0
        and gate["target_type_conflict_records"] == 0
        and gate["unresolved_records"] <= gate["unresolved_required_max"]
        and not gate["single_class_absorption_flag"]
    )
    return {
        "status_counts": dict(status_counts),
        "status_by_target_type": {k: dict(v) for k, v in target_status.items()},
        "label_mapping_type_counts": dict(label_mapping_type_counts),
        "mapped_official_class_distribution_top10": class_absorption,
        "high_confidence_official_class_distribution": dict(Counter(high_names).most_common()),
        "mapped_official_class_count": len(set(mapped_names)),
        "high_confidence_official_class_count": len(set(high_names)),
        "labels_per_official_class": {k: v for k, v in Counter(r["mapped_event_name"] for r in mapped_rows if r.get("mapped_event_name")).most_common()},
        "mapping_gate": gate,
    }


def write_config(label_rows: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    CONFIG_YAML.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "version: task1_step5_8hb_event_mapping_v1",
        "source_workbook: reports/财经新闻标注数据8.18.xlsx",
        f"source_workbook_sha256: {EXPECTED_INCREMENTAL_SHA}",
        "official_label_space:",
        "  source: data/processed/event_schema.jsonl",
        "  official_event_name_count: 64",
        "policy:",
        "  preserve_original_event_name: true",
        "  mapped_event_name_field: mapped_event_name",
        "  do_not_extend_formal_label_space: true",
        "  high_confidence_min_records_for_training_gate: 850",
        "  stratified_auto_audit_min_consistency: 0.95",
        "  unresolved_max_records_for_training_gate: 50",
        "  llm_calls: 0",
        "status:",
        f"  mapping_gate_passed: {str(stats['mapping_gate']['passed']).lower()}",
        "mappings:",
    ]
    for row in label_rows:
        lines.append(f"  - source_event_name: {row['source_event_name']}")
        lines.append(f"    mapping_type: {row['mapping_type']}")
        lines.append(f"    semantic_relation: {row['semantic_relation']}")
        lines.append(f"    confidence: {row['confidence']}")
        lines.append("    candidate_official_event_names:")
        for cand in row["candidate_official_events"]:
            lines.append(f"      - {cand['event_name']}")
    CONFIG_YAML.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_workbooks(
    label_rows: list[dict[str, Any]],
    mapped_rows: list[dict[str, Any]],
    audit_rows: list[dict[str, Any]],
    official_catalog: list[dict[str, Any]],
    stats: dict[str, Any],
) -> None:
    flat_label_rows = []
    for row in label_rows:
        candidates = row["candidate_official_events"]
        flat_label_rows.append(
            {
                "source_event_name": row["source_event_name"],
                "record_count": row["record_count"],
                "sheet_distribution": json.dumps(row["sheet_distribution"], ensure_ascii=False),
                "emotion_distribution": json.dumps(row["emotion_distribution"], ensure_ascii=False),
                "mapping_type": row["mapping_type"],
                "semantic_relation": row["semantic_relation"],
                "candidate_official_event_names": "；".join(c["event_name"] for c in candidates),
                "candidate_relations": "；".join(c["relation"] for c in candidates),
                "confidence": row["confidence"],
                "usable_for_official_classification": row["usable_for_official_classification"],
                "requires_review": row["requires_review"],
                "rationale": row["rationale"],
                "sample_titles": "；".join(s["article_title"] for s in row["representative_samples"]),
                "human_mapped_event_name": "",
                "reviewer_confirmed": "",
                "reviewer_note": "",
            }
        )

    def view_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        cols = [
            "article_file_id", "original_sheet", "target_type", "target_name", "article_title", "original_event_name",
            "mapped_event_name", "mapping_status", "mapping_confidence", "mapping_type", "semantic_relation",
            "event_name_for_64_training", "classification_loss_weight", "event_emotion", "duplication_id",
        ]
        out = []
        for row in rows:
            item = {c: row.get(c, "") for c in cols}
            item["candidate_top5"] = json.dumps(row.get("candidate_top5", []), ensure_ascii=False)
            item["frozen_top1"] = row.get("candidate_scores", {}).get("frozen_model_top1", "")
            item["description_top1"] = row.get("candidate_scores", {}).get("description_similarity_top1", "")
            item["content_excerpt"] = normalize_text(row.get("content"))[:300]
            item["impact_analysis_excerpt"] = normalize_text(row.get("event_impact_analysis"))[:180]
            item["human_mapped_event_name"] = ""
            item["reviewer_confirmed"] = ""
            item["reviewer_note"] = ""
            out.append(item)
        return out

    status_rows = [
        {"metric": "total_rows", "value": len(mapped_rows)},
        {"metric": "source_event_name_count", "value": len(label_rows)},
        {"metric": "mapped_official_class_count", "value": stats["mapped_official_class_count"]},
        {"metric": "high_confidence_records", "value": stats["mapping_gate"]["high_confidence_mapping_records"]},
        {"metric": "needs_review_records", "value": stats["status_counts"].get("mapped_needs_review", 0)},
        {"metric": "unresolved_records", "value": stats["status_counts"].get("unresolved", 0)},
        {"metric": "target_type_conflict_records", "value": stats["status_counts"].get("target_type_conflict", 0)},
        {"metric": "stratified_auto_audit_consistency", "value": stats["mapping_gate"]["stratified_auto_audit_consistency"]},
        {"metric": "mapping_gate_passed", "value": stats["mapping_gate"]["passed"]},
    ]
    coverage_rows = [
        {
            "official_event_name": cat["event_name"],
            "target_types": "；".join(cat["target_types"]),
            "official_train_count": cat["official_train_count"],
            "mapped_record_count": sum(1 for r in mapped_rows if r.get("mapped_event_name") == cat["event_name"]),
            "high_confidence_record_count": sum(1 for r in mapped_rows if r.get("mapped_event_name") == cat["event_name"] and r.get("mapping_status") == "auto_mapped_high_confidence"),
            "schema_ids": "；".join(cat["schema_ids"][:3]),
            "description_excerpt": cat["description_excerpt"][:360],
        }
        for cat in official_catalog
    ]
    with pd.ExcelWriter(DETAILS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(status_rows).to_excel(writer, sheet_name="映射统计", index=False)
        pd.DataFrame(flat_label_rows).to_excel(writer, sheet_name="标签级映射45类", index=False)
        pd.DataFrame(view_rows(mapped_rows)).to_excel(writer, sheet_name="逐条映射1000", index=False)
        pd.DataFrame(audit_rows).to_excel(writer, sheet_name="100条抽样审计", index=False)
        pd.DataFrame(coverage_rows).to_excel(writer, sheet_name="64类覆盖统计", index=False)
    format_workbook(DETAILS_XLSX)

    high_rows = [r for r in mapped_rows if r["mapping_status"] == "auto_mapped_high_confidence"]
    review_rows = [r for r in mapped_rows if r["mapping_status"] == "mapped_needs_review"]
    unresolved_rows = [r for r in mapped_rows if r["mapping_status"] in {"unresolved", "target_type_conflict"}]
    conflict_rows = [r for r in mapped_rows if r["mapping_status"] == "target_type_conflict"]
    with pd.ExcelWriter(REVIEW_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(flat_label_rows).to_excel(writer, sheet_name="标签级映射45类", index=False)
        pd.DataFrame(view_rows(high_rows)).to_excel(writer, sheet_name="高置信自动映射", index=False)
        pd.DataFrame(view_rows(review_rows)).to_excel(writer, sheet_name="需要复核记录", index=False)
        pd.DataFrame(view_rows(unresolved_rows)).to_excel(writer, sheet_name="无法映射记录", index=False)
        pd.DataFrame(view_rows(conflict_rows)).to_excel(writer, sheet_name="主体类型冲突", index=False)
        pd.DataFrame(status_rows).to_excel(writer, sheet_name="映射统计", index=False)
        pd.DataFrame(coverage_rows).to_excel(writer, sheet_name="64类覆盖统计", index=False)
    format_workbook(REVIEW_XLSX)


def write_reports(
    generated_at: str,
    backup_dir: str,
    backup_status: str,
    protected_before: dict[str, str],
    protected_after: dict[str, str],
    incremental_stats: dict[str, Any],
    official_catalog: list[dict[str, Any]],
    label_rows: list[dict[str, Any]],
    mapped_rows: list[dict[str, Any]],
    audit_rows: list[dict[str, Any]],
    stats: dict[str, Any],
) -> dict[str, Any]:
    final_status = "event_mapping_ready_for_training" if stats["mapping_gate"]["passed"] else "event_mapping_human_confirmation_required"
    label_summary = []
    for row in label_rows:
        label_records = [r for r in mapped_rows if r["original_event_name"] == row["source_event_name"]]
        label_summary.append(
            {
                "source_event_name": row["source_event_name"],
                "record_count": row["record_count"],
                "mapping_type": row["mapping_type"],
                "semantic_relation": row["semantic_relation"],
                "candidate_official_event_names": [c["event_name"] for c in row["candidate_official_events"]],
                "row_status_counts": dict(Counter(r["mapping_status"] for r in label_records)),
                "mapped_event_name_distribution": dict(Counter(r.get("mapped_event_name") or "" for r in label_records).most_common()),
                "requires_human_review": True,
                "rationale": row["rationale"],
            }
        )
    report = {
        "generated_at": generated_at,
        "final_status": final_status,
        "training_started": False,
        "training_allowed": bool(stats["mapping_gate"]["passed"]),
        "llm_calls": 0,
        "local_32b_calls": 0,
        "external_llm_calls": 0,
        "protected_test_labels_read": False,
        "formal_outputs_modified": protected_before != protected_after,
        "formal_sha_match": protected_before == protected_after,
        "backup_dir": backup_dir,
        "backup_status": backup_status,
        "input_files": {
            "incremental_workbook": str(INCREMENTAL_XLSX.relative_to(ROOT)),
            "incremental_workbook_sha256": incremental_stats["sha256"],
            "official_train": str(OFFICIAL_TRAIN_XLSX.relative_to(ROOT)),
            "official_train_sha256": sha256_file(OFFICIAL_TRAIN_XLSX),
            "official_stock_json_sha256": sha256_file(OFFICIAL_STOCK_JSON),
            "official_industry_json_sha256": sha256_file(OFFICIAL_INDUSTRY_JSON),
            "event_schema_jsonl_sha256": sha256_file(EVENT_SCHEMA_JSONL),
            "frozen_event_model_sha256": sha256_file(FROZEN_EVENT_MODEL),
        },
        "incremental_stats": incremental_stats,
        "official64_catalog_summary": {
            "official_event_name_count": len(official_catalog),
            "company_event_name_count": sum(1 for c in official_catalog if c["target_types"] == ["company"]),
            "industry_event_name_count": sum(1 for c in official_catalog if c["target_types"] == ["industry"]),
        },
        "label_level_mapping_summary": {
            "source_event_name_count": len(label_rows),
            "mapping_type_counts": stats["label_mapping_type_counts"],
            "labels": label_summary,
        },
        "row_level_mapping_summary": stats,
        "stratified_audit_100": {
            "rows": len(audit_rows),
            "multi_method_consistent": sum(1 for r in audit_rows if r["multi_method_consistent"]),
            "consistency_rate": stats["mapping_gate"]["stratified_auto_audit_consistency"],
        },
        "gate_decision": {
            "decision": final_status,
            "reason": (
                "high-confidence mapping gate passed"
                if stats["mapping_gate"]["passed"]
                else "high-confidence rows, unresolved/conflict counts, or automatic consistency did not satisfy training gate; human confirmation is required before incremental training"
            ),
            "mapping_gate": stats["mapping_gate"],
        },
        "artifacts": {
            "official64_catalog_jsonl": str((OUT_DIR / "official64_catalog.jsonl").relative_to(ROOT)),
            "incremental1000_original_jsonl": str((OUT_DIR / "incremental1000_original.jsonl").relative_to(ROOT)),
            "high_confidence_jsonl": str((OUT_DIR / "incremental1000_mapped_high_confidence.jsonl").relative_to(ROOT)),
            "needs_review_jsonl": str((OUT_DIR / "incremental1000_needs_review.jsonl").relative_to(ROOT)),
            "unresolved_jsonl": str((OUT_DIR / "incremental1000_unresolved.jsonl").relative_to(ROOT)),
            "label_level_mapping_json": str((OUT_DIR / "label_level_mapping.json").relative_to(ROOT)),
            "row_level_mapping_jsonl": str((OUT_DIR / "row_level_mapping.jsonl").relative_to(ROOT)),
            "mapping_manifest_json": str((OUT_DIR / "mapping_manifest.json").relative_to(ROOT)),
            "review_workbook": str(REVIEW_XLSX.relative_to(ROOT)),
            "details_workbook": str(DETAILS_XLSX.relative_to(ROOT)),
            "config_yaml": str(CONFIG_YAML.relative_to(ROOT)),
        },
    }
    write_json(REPORT_JSON, report)
    markdown_lines = [
        "# Task1 Step 5.8-HB 45类到官方64类映射审计",
        "",
        f"- 生成时间：{generated_at}",
        f"- 最终状态：`{final_status}`",
        f"- 新增Excel SHA：`{incremental_stats['sha256']}`",
        f"- 冻结64类模型 SHA：`{sha256_file(FROZEN_EVENT_MODEL)}`",
        f"- 官方64类目录：64类，其中个股44类、行业20类。",
        f"- 新增源标签：{len(label_rows)}类，1000条。",
        f"- 高置信自动映射：{stats['mapping_gate']['high_confidence_mapping_records']}条。",
        f"- 需要复核：{stats['status_counts'].get('mapped_needs_review', 0)}条。",
        f"- 无法映射：{stats['status_counts'].get('unresolved', 0)}条。",
        f"- 主体类型冲突：{stats['status_counts'].get('target_type_conflict', 0)}条。",
        f"- 100条分层双方法一致性：{stats['mapping_gate']['stratified_auto_audit_consistency']:.3f}。",
        f"- 是否进入增量训练：`{bool(stats['mapping_gate']['passed'])}`。",
        f"- 正式Task1/Task2输出SHA不变：`{protected_before == protected_after}`。",
        f"- LLM/32B调用次数：0/0。",
        f"- 备份目录：`{backup_dir}`（{backup_status}）。",
        "",
        "## 结论",
        "",
        (
            "高置信映射门槛已通过，可进入 Step 5.8-HC 增量训练。"
            if stats["mapping_gate"]["passed"]
            else "当前45类与官方64类存在大量宽泛行情类、结果类和行业观察类标签；冻结64类模型Top-K只作为建议，不能直接替代人工确认。已生成审核工作簿，训练前需要人工确认 mapped_event_name。"
        ),
        "",
        "## 45类标签级映射摘要",
        "",
        "| 新增标签 | 条数 | 映射类型 | 候选官方事件 | 状态分布 |",
        "|---|---:|---|---|---|",
    ]
    for row in label_summary:
        markdown_lines.append(
            "| {source_event_name} | {record_count} | {mapping_type} | {cands} | {status} |".format(
                source_event_name=row["source_event_name"],
                record_count=row["record_count"],
                mapping_type=row["mapping_type"],
                cands="、".join(row["candidate_official_event_names"][:5]) or "无安全候选",
                status=json.dumps(row["row_status_counts"], ensure_ascii=False),
            )
        )
    markdown_lines.extend(
        [
            "",
            "## 产物",
            "",
            f"- 审核工作簿：`{REVIEW_XLSX.relative_to(ROOT)}`",
            f"- 逐条映射：`{(OUT_DIR / 'row_level_mapping.jsonl').relative_to(ROOT)}`",
            f"- 报告JSON：`{REPORT_JSON.relative_to(ROOT)}`",
        ]
    )
    REPORT_MD.write_text("\n".join(markdown_lines) + "\n", encoding="utf-8")
    return report


def write_sha256s(paths: list[Path]) -> None:
    rows = []
    for path in sorted({p for p in paths if p.exists()}):
        rows.append(f"{sha256_file(path)}  {path.relative_to(ROOT)}")
    (OUT_DIR / "SHA256SUMS.txt").write_text("\n".join(rows) + "\n", encoding="utf-8")


def main() -> None:
    generated_at = datetime.now().replace(microsecond=0).isoformat()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    protected_before = collect_protected_sha()
    backup_dir, backup_status = create_backup_dir(generated_at, protected_before)

    rows, incremental_stats = load_incremental_rows()
    official_catalog, _, allowed_by_type = load_official_catalog()
    specs = label_mapping_specs()
    label_rows = build_label_level_mapping(rows, specs, official_catalog)
    mapped_rows = build_row_level_mapping(rows, specs, official_catalog, allowed_by_type)
    audit_rows = sample_audit_rows(mapped_rows, n=100)
    stats = compute_mapping_stats(mapped_rows, label_rows, audit_rows)
    write_config(label_rows, stats)

    high_conf = [r for r in mapped_rows if r["mapping_status"] == "auto_mapped_high_confidence"]
    needs_review = [r for r in mapped_rows if r["mapping_status"] == "mapped_needs_review"]
    unresolved = [r for r in mapped_rows if r["mapping_status"] in {"unresolved", "target_type_conflict"}]

    write_jsonl(OUT_DIR / "official64_catalog.jsonl", official_catalog)
    write_jsonl(OUT_DIR / "incremental1000_original.jsonl", rows)
    write_jsonl(OUT_DIR / "incremental1000_mapped_high_confidence.jsonl", high_conf)
    write_jsonl(OUT_DIR / "incremental1000_needs_review.jsonl", needs_review)
    write_jsonl(OUT_DIR / "incremental1000_unresolved.jsonl", unresolved)
    write_json(OUT_DIR / "label_level_mapping.json", label_rows)
    write_jsonl(OUT_DIR / "row_level_mapping.jsonl", mapped_rows)

    write_workbooks(label_rows, mapped_rows, audit_rows, official_catalog, stats)

    protected_after = collect_protected_sha()
    manifest = {
        "generated_at": generated_at,
        "status": "event_mapping_ready_for_training" if stats["mapping_gate"]["passed"] else "event_mapping_human_confirmation_required",
        "training_started": False,
        "high_confidence_count": len(high_conf),
        "needs_review_count": len(needs_review),
        "unresolved_or_conflict_count": len(unresolved),
        "llm_calls": 0,
        "local_32b_calls": 0,
        "protected_test_labels_read": False,
        "formal_sha_match": protected_before == protected_after,
        "backup_dir": backup_dir,
        "backup_status": backup_status,
        "mapping_gate": stats["mapping_gate"],
        "resume_policy": "If human confirmation is completed, use row_level_mapping.jsonl plus reviewed workbook to enter Step 5.8-HC. Do not train on mapped_needs_review or unresolved rows before confirmation.",
    }
    write_json(OUT_DIR / "mapping_manifest.json", manifest)
    write_json(OUT_DIR / "run_state.json", manifest)
    write_reports(
        generated_at,
        backup_dir,
        backup_status,
        protected_before,
        protected_after,
        incremental_stats,
        official_catalog,
        label_rows,
        mapped_rows,
        audit_rows,
        stats,
    )
    write_sha256s(
        [
            OUT_DIR / "official64_catalog.jsonl",
            OUT_DIR / "incremental1000_original.jsonl",
            OUT_DIR / "incremental1000_mapped_high_confidence.jsonl",
            OUT_DIR / "incremental1000_needs_review.jsonl",
            OUT_DIR / "incremental1000_unresolved.jsonl",
            OUT_DIR / "label_level_mapping.json",
            OUT_DIR / "row_level_mapping.jsonl",
            OUT_DIR / "mapping_manifest.json",
            OUT_DIR / "run_state.json",
            REPORT_JSON,
            REPORT_MD,
            DETAILS_XLSX,
            REVIEW_XLSX,
            CONFIG_YAML,
        ]
    )
    print(
        json.dumps(
            {
                "status": manifest["status"],
                "high_confidence_count": len(high_conf),
                "needs_review_count": len(needs_review),
                "unresolved_or_conflict_count": len(unresolved),
                "mapping_gate_passed": stats["mapping_gate"]["passed"],
                "review_workbook": str(REVIEW_XLSX),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
