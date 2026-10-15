#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import math
import random
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
RAW_XLSX = ROOT / "data/raw/news_with_tags_test.xlsx"
INFER_DIR = ROOT / "outputs/task3_kg_v4_official1799_inference_v1"
OUT = ROOT / "outputs/task34_official1799_downstream_benchmark_v1"
SEED = 20260825

ALLOWED_INPUT_COLUMNS = [
    "article_file_id",
    "article_title",
    "article_publish_time",
    "article_source",
    "content",
]
FORBIDDEN_INPUT_COLUMNS = {
    "trading_code",
    "secu_abbr",
    "industry_code",
    "industry_name",
    "event_name",
    "event_emotion",
    "event_impact_analysis",
    "duplication_id",
}
FORBIDDEN_ANNOTATION_COLUMNS = {
    "trading_code",
    "secu_abbr",
    "industry_code",
    "industry_name",
    "official_event_name",
    "event_emotion",
    "event_impact_analysis",
    "duplication_id",
    "model_decision",
    "model_score",
    "model_probability",
    "probability",
    "probabilities",
    "candidate_bucket",
    "red_orange_hint",
    "resolver_decision",
    "resolver_confidence",
    "rule_hits",
    "route",
    "abstain",
    "suppress",
    "official_label",
}
TASK3_HUMAN_FIELDS = [
    "human_merge_label",
    "event_a_valid",
    "event_b_valid",
    "evidence_sufficient",
    "annotation_confidence",
    "annotation_reason",
]
TASK4_HUMAN_FIELDS = [
    "target_grounded",
    "event_grounded",
    "evidence_support",
    "validity",
    "direction_gold",
    "risk_intensity",
    "severity_label",
    "rule_applicability",
    "financial_evidence",
    "alert_gold",
    "annotation_confidence",
    "annotation_basis",
    "invalid_reason",
]
TASK3_ALLOWED = {
    "human_merge_label": ["same_event", "different_event", "uncertain"],
    "event_a_valid": ["valid", "invalid", "uncertain"],
    "event_b_valid": ["valid", "invalid", "uncertain"],
    "evidence_sufficient": ["yes", "no", "uncertain"],
    "annotation_confidence": ["high", "medium", "low"],
}
TASK4_ALLOWED = {
    "target_grounded": ["yes", "no", "uncertain"],
    "event_grounded": ["yes", "no", "uncertain"],
    "evidence_support": ["direct", "partial", "none", "uncertain"],
    "validity": ["valid", "invalid", "uncertain"],
    "direction_gold": ["positive", "negative", "neutral", "uncertain"],
    "risk_intensity": ["1", "2", "3", "4", "5"],
    "severity_label": ["blue", "yellow", "orange", "red", "uncertain"],
    "rule_applicability": ["applicable", "not_applicable", "uncertain"],
    "financial_evidence": ["yes", "no", "uncertain"],
    "alert_gold": ["alert", "no_alert", "review", "uncertain"],
    "annotation_confidence": ["high", "medium", "low"],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HUMAN_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2F0D9")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null", "nat"} else text


def compact(value: Any) -> str:
    return "".join(clean(value).split())


def stable_hash(*parts: Any, n: int = 20) -> str:
    payload = "||".join(clean(p) for p in parts)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:n]


def deterministic_float(*parts: Any) -> float:
    digest = hashlib.sha256(("||".join([str(SEED), *[clean(p) for p in parts]])).encode("utf-8")).hexdigest()[:16]
    return int(digest, 16) / float(16**16 - 1)


def json_ready(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {clean(k): json_ready(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple, set)):
        return [json_ready(v) for v in obj]
    if isinstance(obj, pd.Timestamp):
        return obj.isoformat()
    if hasattr(obj, "item"):
        try:
            return json_ready(obj.item())
        except Exception:
            pass
    try:
        if pd.isna(obj):
            return None
    except Exception:
        pass
    if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
        return None
    return obj


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(json_ready(payload), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def inspect_headers(path: Path) -> dict[str, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    out: dict[str, list[str]] = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        out[sheet] = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return out


def load_allowed_raw_manifest() -> tuple[pd.DataFrame, dict[str, Any]]:
    headers = inspect_headers(RAW_XLSX)
    rows = []
    sheet_counts = {}
    forbidden_present = {}
    for sheet, cols in headers.items():
        missing = [c for c in ALLOWED_INPUT_COLUMNS if c not in cols]
        if missing:
            raise RuntimeError(f"missing allowed columns in {sheet}: {missing}")
        forbidden_present[sheet] = [c for c in cols if c in FORBIDDEN_INPUT_COLUMNS]
        frame = pd.read_excel(RAW_XLSX, sheet_name=sheet, usecols=ALLOWED_INPUT_COLUMNS, dtype=object).fillna("")
        sheet_counts[sheet] = int(len(frame))
        for idx, row in frame.iterrows():
            rows.append(
                {
                    "official_row_id": f"OFF1799_{len(rows) + 1:04d}",
                    "sheet_name": sheet,
                    "sheet_row_index": int(idx) + 2,
                    "article_file_id": clean(row["article_file_id"]),
                    "article_title": clean(row["article_title"]),
                    "article_publish_time": clean(row["article_publish_time"]),
                    "article_source": clean(row["article_source"]),
                    "content_hash": hashlib.sha256(compact(row["content"]).encode("utf-8")).hexdigest(),
                    "content_length": len(clean(row["content"])),
                }
            )
    raw = pd.DataFrame(rows)
    manifest = {
        "source_file": str(RAW_XLSX),
        "source_sha256": sha256_file(RAW_XLSX),
        "allowed_columns_read": ALLOWED_INPUT_COLUMNS,
        "forbidden_columns_present_but_not_read": forbidden_present,
        "forbidden_columns_read": [],
        "read_method": "pd.read_excel(sheet_name=..., usecols=allowed_columns_only)",
        "sheet_counts": sheet_counts,
        "input_rows": int(len(raw)),
        "unique_article_file_id": int(raw["article_file_id"].nunique()),
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
    }
    return raw, manifest


def event_record_from_structured(row: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    target = row.get("target") or {}
    event = row.get("event") or {}
    evidence = row.get("evidence") or {}
    source_row = row.get("source_row") or {}
    provenance = row.get("provenance") or {}
    evidence_span = evidence.get("evidence_span") or {}
    trigger_span = event.get("trigger_span") or {}
    offset = {
        "trigger": trigger_span,
        "evidence": evidence_span,
        "local_context": evidence.get("local_context_span") or {},
    }
    base = {
        f"{prefix}sample_id": clean(source_row.get("sample_id")),
        f"{prefix}official_row_id": clean(row.get("official_row_id")),
        f"{prefix}article_file_id": clean(row.get("article_id")),
        f"{prefix}article_source": clean((row.get("source") or {}).get("source_name")),
        f"{prefix}article_publish_time": clean(row.get("publish_time")),
        f"{prefix}target": clean(target.get("target_name")),
        f"{prefix}target_type": clean(target.get("target_type")),
        f"{prefix}event_type": clean(event.get("event_type")),
        f"{prefix}trigger": clean(event.get("trigger_text")),
        f"{prefix}event_fact": clean(event.get("event_fact")),
        f"{prefix}event_time": clean(event.get("event_time")),
        f"{prefix}arguments": json.dumps(json_ready(event.get("arguments") or []), ensure_ascii=False),
        f"{prefix}evidence": clean(evidence.get("evidence_sentence")),
        f"{prefix}local_context": clean(evidence.get("local_context")),
        f"{prefix}offset": json.dumps(json_ready(offset), ensure_ascii=False),
        f"{prefix}provenance": json.dumps(
            json_ready(
                {
                    "document_hash": provenance.get("document_hash"),
                    "sentence_id": provenance.get("sentence_id"),
                    "source_file": provenance.get("source_file"),
                    "official_row_id": row.get("official_row_id"),
                }
            ),
            ensure_ascii=False,
        ),
    }
    return base


def build_pair_rows(pair_decisions: list[dict[str, Any]], structured_by_sample: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for pair in pair_decisions:
        a = structured_by_sample.get(clean(pair.get("sample_id_a")))
        b = structured_by_sample.get(clean(pair.get("sample_id_b")))
        if not a or not b:
            continue
        sample_id = "T3PAIR_" + stable_hash(pair.get("pair_id"), pair.get("sample_id_a"), pair.get("sample_id_b"))
        visible = {
            "annotation_sample_id": sample_id,
            "row_type": "event_pair",
            "pair_id": sample_id,
            **event_record_from_structured(a, "event_a_"),
            **event_record_from_structured(b, "event_b_"),
        }
        hidden = {
            "sample_id": sample_id,
            "row_type": "event_pair",
            "resolver_prediction_hidden": clean(pair.get("prediction")),
            "channels_hidden": clean(pair.get("channels")),
            "hard_conflict_count_hidden": int(pair.get("hard_conflict_count") or 0),
            "same_target_hidden": int((pair.get("features") or {}).get("same_target") or 0),
            "same_event_name_hidden": int((pair.get("features") or {}).get("same_event_name") or 0),
            "time_relation_hidden": clean((pair.get("features") or {}).get("time_relation")),
        }
        rows.append({"visible": visible, "hidden": hidden})
    return rows


def build_singleton_rows(structured_events: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for ev in structured_events:
        source_row = ev.get("source_row") or {}
        sample_id = "T3SINGLE_" + stable_hash(source_row.get("sample_id"), ev.get("official_row_id"))
        rows.append(
            {
                "visible": {
                    "annotation_sample_id": sample_id,
                    "row_type": "singleton",
                    "pair_id": "",
                    **event_record_from_structured(ev, "event_a_"),
                    **{f"event_b_{field}": "" for field in [
                        "sample_id",
                        "official_row_id",
                        "article_file_id",
                        "article_source",
                        "article_publish_time",
                        "target",
                        "target_type",
                        "event_type",
                        "trigger",
                        "event_fact",
                        "event_time",
                        "arguments",
                        "evidence",
                        "local_context",
                        "offset",
                        "provenance",
                    ]},
                },
                "hidden": {
                    "sample_id": sample_id,
                    "row_type": "singleton",
                    "source_sample_id_hidden": clean(source_row.get("sample_id")),
                    "target_type_hidden": clean((ev.get("target") or {}).get("target_type")),
                    "event_family_hidden": clean((ev.get("event") or {}).get("event_family")),
                },
            }
        )
    return rows


def choose_task3_samples(pair_rows: list[dict[str, Any]], singleton_rows: list[dict[str, Any]]) -> tuple[pd.DataFrame, dict[str, Any]]:
    rng = random.Random(SEED)
    by_pred: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in pair_rows:
        by_pred[row["hidden"]["resolver_prediction_hidden"]].append(row)
    for values in by_pred.values():
        values.sort(key=lambda r: deterministic_float(r["hidden"]["sample_id"]))
    quotas = {"same_event": 130, "different_event": 190, "uncertain": 80}
    selected = []
    used = set()
    for pred, quota in quotas.items():
        for row in by_pred.get(pred, [])[:quota]:
            selected.append(row)
            used.add(row["hidden"]["sample_id"])
    if len(selected) < 400:
        remainder = [r for r in pair_rows if r["hidden"]["sample_id"] not in used]
        remainder.sort(key=lambda r: deterministic_float("task3_remainder", r["hidden"]["sample_id"]))
        selected.extend(remainder[: 400 - len(selected)])
    selected = selected[:400]

    singleton_rows = sorted(singleton_rows, key=lambda r: deterministic_float("singleton", r["hidden"]["sample_id"]))[:100]
    all_rows = selected + singleton_rows
    visible_rows = []
    hidden_rows = []
    for row in all_rows:
        visible = dict(row["visible"])
        for field in TASK3_HUMAN_FIELDS:
            visible[field] = ""
        visible_rows.append(visible)
        hidden_rows.append(row["hidden"])
    meta = {
        "task3_pair_rows": 400,
        "task3_singleton_rows": 100,
        "pair_hidden_resolver_strata": dict(Counter(r["resolver_prediction_hidden"] for r in hidden_rows if r["row_type"] == "event_pair")),
        "singleton_hidden_target_type": dict(Counter(r.get("target_type_hidden", "") for r in hidden_rows if r["row_type"] == "singleton")),
        "annotation_sample_ids": [r["annotation_sample_id"] for r in visible_rows],
        "hidden_sampling_rows": hidden_rows,
    }
    return pd.DataFrame(visible_rows), meta


def choose_task4_samples(task4_rows: list[dict[str, Any]], structured_by_official: dict[str, dict[str, Any]]) -> tuple[pd.DataFrame, dict[str, Any]]:
    candidates = []
    for row in task4_rows:
        ev = structured_by_official.get(clean(row.get("official_row_id")))
        if not ev:
            continue
        source_row = ev.get("source_row") or {}
        quality = ev.get("quality") or {}
        event = ev.get("event") or {}
        evidence = ev.get("evidence") or {}
        sample_id = "T4ANN_" + stable_hash(source_row.get("sample_id"), row.get("canonical_event_id"), row.get("official_row_id"))
        candidates.append(
            {
                "visible": {
                    "annotation_sample_id": sample_id,
                    "official_row_id": clean(row.get("official_row_id")),
                    "article_file_id": clean(row.get("article_id")),
                    "article_source": clean((ev.get("source") or {}).get("source_name")),
                    "article_publish_time": clean(row.get("event_time") or ev.get("publish_time")),
                    "target": clean(row.get("target_name")),
                    "target_type": clean(row.get("target_type")),
                    "event_type": clean(row.get("event_type")),
                    "trigger": clean(row.get("trigger_text") or event.get("trigger_text")),
                    "event_fact": clean(event.get("event_fact")),
                    "event_time": clean(row.get("event_time") or event.get("event_time")),
                    "arguments": json.dumps(json_ready(event.get("arguments") or []), ensure_ascii=False),
                    "evidence": clean((row.get("evidence_span") or {}).get("text") or evidence.get("evidence_sentence")),
                    "local_context": clean(evidence.get("local_context")),
                    "offset": json.dumps(
                        json_ready(
                            {
                                "trigger": event.get("trigger_span") or {},
                                "evidence": row.get("evidence_span") or evidence.get("evidence_span") or {},
                                "local_context": evidence.get("local_context_span") or {},
                            }
                        ),
                        ensure_ascii=False,
                    ),
                    "provenance": json.dumps(
                        json_ready(
                            {
                                "document_hash": (row.get("provenance") or {}).get("document_hash"),
                                "sentence_id": (row.get("provenance") or {}).get("sentence_id"),
                                "source_file": (row.get("provenance") or {}).get("source_file"),
                                "canonical_event_id": row.get("canonical_event_id"),
                                "event_mention_id": row.get("event_mention_id"),
                            }
                        ),
                        ensure_ascii=False,
                    ),
                    "support_type": clean(row.get("support_type")),
                },
                "hidden": {
                    "sample_id": sample_id,
                    "target_type_hidden": clean(row.get("target_type")),
                    "event_family_hidden": clean(row.get("event_family")),
                    "support_type_hidden": clean(row.get("support_type")),
                    "structured_pass_hidden": clean(row.get("contract_status")) == "pass",
                    "target_grounded_hidden": bool(quality.get("target_grounded")),
                    "event_grounded_hidden": bool(quality.get("event_grounded")),
                    "source_sample_id_hidden": clean(source_row.get("sample_id")),
                },
            }
        )
    if len(candidates) <= 500:
        selected = sorted(candidates, key=lambda r: deterministic_float("task4_all", r["hidden"]["sample_id"]))
        strategy = "all_candidates"
    else:
        selected = []
        used = set()
        rep_quotas = {
            ("target_type_hidden", "company"): 286,
            ("target_type_hidden", "industry"): 34,
        }
        for (key, value), quota in rep_quotas.items():
            bucket = [r for r in candidates if r["hidden"].get(key) == value]
            bucket.sort(key=lambda r: deterministic_float("task4_rep", key, value, r["hidden"]["sample_id"]))
            for row in bucket[:quota]:
                selected.append(row)
                used.add(row["hidden"]["sample_id"])
        if len(selected) < 320:
            rest = [r for r in candidates if r["hidden"]["sample_id"] not in used]
            rest.sort(key=lambda r: deterministic_float("task4_rep_rest", r["hidden"]["sample_id"]))
            for row in rest[: 320 - len(selected)]:
                selected.append(row)
                used.add(row["hidden"]["sample_id"])

        diagnostic_pool = [r for r in candidates if r["hidden"]["sample_id"] not in used]
        diagnostic_pool.sort(
            key=lambda r: (
                not (r["hidden"]["support_type_hidden"] == "quoted_claim"),
                not (r["hidden"]["target_type_hidden"] == "industry"),
                r["hidden"]["target_grounded_hidden"],
                deterministic_float("task4_diag", r["hidden"]["sample_id"]),
            )
        )
        for row in diagnostic_pool[:180]:
            selected.append(row)
            used.add(row["hidden"]["sample_id"])
        strategy = "320_representative_plus_180_diagnostic"
    selected = selected[:500]
    visible_rows = []
    hidden_rows = []
    for row in selected:
        visible = dict(row["visible"])
        for field in TASK4_HUMAN_FIELDS:
            visible[field] = ""
        visible_rows.append(visible)
        hidden_rows.append(row["hidden"])
    meta = {
        "task4_candidate_rows": int(len(candidates)),
        "task4_selected_rows": int(len(visible_rows)),
        "sampling_strategy": strategy,
        "selected_target_type_distribution": dict(Counter(r["target_type_hidden"] for r in hidden_rows)),
        "selected_event_family_distribution": dict(Counter(r["event_family_hidden"] for r in hidden_rows)),
        "selected_support_type_distribution": dict(Counter(r["support_type_hidden"] for r in hidden_rows)),
        "annotation_sample_ids": [r["annotation_sample_id"] for r in visible_rows],
        "hidden_sampling_rows": hidden_rows,
    }
    return pd.DataFrame(visible_rows), meta


def shuffled_version(frame: pd.DataFrame, version: str) -> pd.DataFrame:
    out = frame.copy()
    out["_sort_key"] = out["annotation_sample_id"].map(lambda x: deterministic_float(version, x))
    out = out.sort_values("_sort_key").drop(columns=["_sort_key"]).reset_index(drop=True)
    out.insert(0, "annotation_order", range(1, len(out) + 1))
    return out


def add_dropdowns(ws, allowed: dict[str, list[str]], end_row: int) -> None:
    headers = [clean(c.value) for c in ws[1]]
    for col, values in allowed.items():
        if col not in headers:
            continue
        idx = headers.index(col) + 1
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{end_row}")


def write_workbook(path: Path, frame: pd.DataFrame, task: str, version: str) -> None:
    wb = Workbook()
    ws_inst = wb.active
    ws_inst.title = "instructions"
    if task == "task3":
        rows = [
            ["字段", "说明"],
            ["任务", "Task3 event pair/singleton 双人独立盲标"],
            ["合法 human_merge_label", "same_event / different_event / uncertain"],
            ["填写要求", "只填写浅黄色人工字段；不要修改样本字段。"],
            ["合成示例", "同一公司、同一天、相同产品和订单事实，仅报道来源不同 -> same_event"],
        ]
        human_fields = TASK3_HUMAN_FIELDS
        allowed = TASK3_ALLOWED
    else:
        rows = [
            ["字段", "说明"],
            ["任务", "Task4 Severity/Impact downstream 双人独立盲标"],
            ["评分目标", "判断 target/event/evidence 是否成立，再标注 validity、direction、1-5 risk_intensity、severity_label 与 alert_gold。"],
            ["填写要求", "只填写浅黄色人工字段；不要修改样本字段。"],
            ["合成示例", "公司公告确认大额债务只能偿付40%，财务指标显示偿债压力 -> validity=valid; direction=negative; risk_intensity=4; alert_gold=alert"],
        ]
        human_fields = TASK4_HUMAN_FIELDS
        allowed = TASK4_ALLOWED
    for row in rows:
        ws_inst.append(row)
    for cell in ws_inst[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    ws_inst.column_dimensions["A"].width = 24
    ws_inst.column_dimensions["B"].width = 120
    for row in ws_inst.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            if clean(cell.value).startswith("同一公司") or clean(cell.value).startswith("公司公告"):
                cell.fill = EXAMPLE_FILL

    ws = wb.create_sheet("annotation")
    ws.append(list(frame.columns))
    for _, row in frame.iterrows():
        ws.append([row[col] for col in frame.columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = WRAP
    human_set = set(human_fields)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = TOP
            header = clean(ws.cell(1, cell.column).value)
            if header in human_set:
                cell.fill = HUMAN_FILL
    width_by_name = {
        "annotation_order": 10,
        "annotation_sample_id": 26,
        "row_type": 14,
        "article_file_id": 30,
        "article_source": 18,
        "article_publish_time": 22,
        "target": 18,
        "target_type": 14,
        "event_type": 36,
        "trigger": 28,
        "event_fact": 70,
        "evidence": 70,
        "local_context": 80,
        "arguments": 42,
        "offset": 38,
        "provenance": 44,
        "annotation_reason": 52,
        "annotation_basis": 52,
        "invalid_reason": 42,
    }
    for idx, cell in enumerate(ws[1], start=1):
        name = clean(cell.value)
        width = width_by_name.get(name, 18)
        if name.startswith("event_a_") or name.startswith("event_b_"):
            width = width_by_name.get(name.replace("event_a_", "").replace("event_b_", ""), 22)
        ws.column_dimensions[get_column_letter(idx)].width = width
    add_dropdowns(ws, allowed, len(frame) + 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def verify_files(task3_a: pd.DataFrame, task3_b: pd.DataFrame, task4_a: pd.DataFrame, task4_b: pd.DataFrame) -> dict[str, Any]:
    def empty_human(frame: pd.DataFrame, fields: list[str]) -> bool:
        return all(frame[field].map(clean).eq("").all() for field in fields)

    def same_set(a: pd.DataFrame, b: pd.DataFrame) -> bool:
        return set(a["annotation_sample_id"]) == set(b["annotation_sample_id"])

    def different_order(a: pd.DataFrame, b: pd.DataFrame) -> bool:
        return list(a["annotation_sample_id"]) != list(b["annotation_sample_id"])

    annotation_cols = set(task3_a.columns) | set(task4_a.columns)
    leaked = sorted(c for c in annotation_cols if c in FORBIDDEN_ANNOTATION_COLUMNS)
    return {
        "task3_ab_same_sample_set": same_set(task3_a, task3_b),
        "task3_ab_different_order": different_order(task3_a, task3_b),
        "task4_ab_same_sample_set": same_set(task4_a, task4_b),
        "task4_ab_different_order": different_order(task4_a, task4_b),
        "task3_human_fields_empty": empty_human(task3_a, TASK3_HUMAN_FIELDS) and empty_human(task3_b, TASK3_HUMAN_FIELDS),
        "task4_human_fields_empty": empty_human(task4_a, TASK4_HUMAN_FIELDS) and empty_human(task4_b, TASK4_HUMAN_FIELDS),
        "forbidden_annotation_columns_present": leaked,
        "no_annotation_leakage": not leaked,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    raw, input_manifest = load_allowed_raw_manifest()
    if len(raw) != 1799:
        raise RuntimeError(f"expected 1799 input rows, got {len(raw)}")

    structured_events = read_jsonl(INFER_DIR / "official1799_structured_events.jsonl")
    task4_interface = read_jsonl(INFER_DIR / "task4_event_interface.jsonl")
    pair_decisions = read_jsonl(INFER_DIR / "resolver_v2_pair_decisions.jsonl")
    structured_by_sample = {clean((row.get("source_row") or {}).get("sample_id")): row for row in structured_events}
    structured_by_official = {clean(row.get("official_row_id")): row for row in structured_events}

    pair_rows = build_pair_rows(pair_decisions, structured_by_sample)
    singleton_rows = build_singleton_rows(structured_events)
    task3_frame, task3_meta = choose_task3_samples(pair_rows, singleton_rows)
    task4_frame, task4_meta = choose_task4_samples(task4_interface, structured_by_official)

    task3_a = shuffled_version(task3_frame, "task3_A")
    task3_b = shuffled_version(task3_frame, "task3_B")
    task4_a = shuffled_version(task4_frame, "task4_A")
    task4_b = shuffled_version(task4_frame, "task4_B")

    paths = {
        "task3_A": OUT / "task3_official1799_annotation_A.xlsx",
        "task3_B": OUT / "task3_official1799_annotation_B.xlsx",
        "task4_A": OUT / "task4_official1799_annotation_A.xlsx",
        "task4_B": OUT / "task4_official1799_annotation_B.xlsx",
    }
    write_workbook(paths["task3_A"], task3_a, "task3", "A")
    write_workbook(paths["task3_B"], task3_b, "task3", "B")
    write_workbook(paths["task4_A"], task4_a, "task4", "A")
    write_workbook(paths["task4_B"], task4_b, "task4", "B")

    task3_a.to_parquet(OUT / "task3_official1799_annotation_A.parquet", index=False)
    task3_b.to_parquet(OUT / "task3_official1799_annotation_B.parquet", index=False)
    task4_a.to_parquet(OUT / "task4_official1799_annotation_A.parquet", index=False)
    task4_b.to_parquet(OUT / "task4_official1799_annotation_B.parquet", index=False)

    verification = verify_files(task3_a, task3_b, task4_a, task4_b)
    if not all(
        [
            verification["task3_ab_same_sample_set"],
            verification["task3_ab_different_order"],
            verification["task4_ab_same_sample_set"],
            verification["task4_ab_different_order"],
            verification["task3_human_fields_empty"],
            verification["task4_human_fields_empty"],
            verification["no_annotation_leakage"],
        ]
    ):
        raise RuntimeError(f"verification failed: {verification}")

    sampling_meta = {
        "seed": SEED,
        "task3": task3_meta,
        "task4": task4_meta,
        "workbook_paths": {k: str(v) for k, v in paths.items()},
    }
    funnel = {
        "input_news_rows": int(len(raw)),
        "task1_structured_events": int(len(structured_events)),
        "task1_contract_pass": int(sum(clean((r.get("quality") or {}).get("contract_status")) == "pass" for r in structured_events)),
        "task1_grounding_pass": int(sum(bool((r.get("quality") or {}).get("grounding_pass")) for r in structured_events)),
        "task3_resolver_candidate_pairs": int(len(pair_decisions)),
        "task3_annotation_pairs": 400,
        "task3_annotation_singletons": 100,
        "task4_interface_candidates": int(len(task4_interface)),
        "task4_annotation_rows": int(len(task4_frame)),
        "task4_sampling_strategy": task4_meta["sampling_strategy"],
        "training_or_tuning_performed": False,
        "task34_or_severity_gold_read": False,
        "frozen_files_modified": False,
    }
    annotation_guide = "\n".join(
        [
            "# Task3/Task4 Official1799 Downstream Blind Annotation Guide",
            "",
            "Only fill yellow human fields. Do not infer labels from row order.",
            "",
            "Task3: mark event pairs as same_event, different_event, or uncertain. Singleton rows are for event validity/evidence checks.",
            "Synthetic Task3 example: 同一公司、同一天、相同产品和订单事实，仅报道来源不同 -> same_event.",
            "",
            "Task4: verify target/event grounding and evidence support first, then annotate validity, direction, risk_intensity 1-5, severity_label, rule_applicability, financial_evidence, and alert_gold.",
            "Synthetic Task4 example: 公司公告确认大额债务只能偿付40%，财务指标显示偿债压力 -> validity=valid; direction=negative; risk_intensity=4; alert_gold=alert.",
            "",
            "No official labels, model decisions, model scores, probabilities, resolver results, rule hits, routes, or alert hints are shown in the annotation workbooks.",
        ]
    )
    write_json(OUT / "input_manifest.json", input_manifest)
    write_json(OUT / "sampling_meta.json", sampling_meta)
    write_json(OUT / "validation_report.json", verification)
    write_json(OUT / "funnel.json", funnel)
    (OUT / "annotation_guide.md").write_text(annotation_guide + "\n", encoding="utf-8")

    print(
        json.dumps(
            {
                "paths": {k: str(v) for k, v in paths.items()},
                "funnel": funnel,
                "verification": verification,
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
