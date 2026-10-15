#!/usr/bin/env python3
"""Audit independent Safety Gate v1.2 errors and prepare ISR routing candidates.

This is a post-hoc diagnostic tool. It reads filled human labels only after
predictions already exist, does not call a model, and does not alter production
artifacts, gold labels, or safety-gate rules.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.validate_task2_dea_human_audit import table_records  # noqa: E402


CLASSES = {"support", "contradict", "insufficient", "unclear"}
ERROR_TYPES = {
    "claim_unclear",
    "claim_construction_error",
    "evidence_retrieval_miss",
    "evidence_ranking_error",
    "evidence_insufficient",
    "evidence_irrelevant",
    "evidence_conflict",
    "subject_mismatch",
    "time_mismatch",
    "investor_question_as_fact",
    "market_signal_as_event",
    "reasoning_error",
    "verdict_mapping_error",
    "gold_label_uncertain",
    "other",
}
ERROR_STAGES = {
    "claim",
    "retrieval",
    "ranking",
    "evidence_selection",
    "reasoning",
    "verdict",
    "safety_gate",
    "gold",
}
ROUTES = {
    "isr": "A. ISR适合",
    "retrieval": "B. 检索优先",
    "claim": "C. 主张修复优先",
    "safety_gate": "D. 门控修复优先",
    "gold_recheck": "E. 金标待复核",
}

DEFAULT_GOLD_FILE = ROOT / "reports/task2_safety_gate_v1_2_independent_human_audit_gold_filled.xlsx"
DEFAULT_EVAL_DIR = ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent_gold_eval"
DEFAULT_RUN_DIR = ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent"
DEFAULT_EVAL_REPORT = ROOT / "reports/task2_safety_gate_v1_2_independent_gold_evaluation.json"
DEFAULT_DEV_RESULTS = ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_dev/safety_gate_v1_2_results.jsonl"
DEFAULT_DEV_REPORT = ROOT / "reports/task2_dea_safety_gate_v1_2_dev_report.json"
DEFAULT_OUTPUT_DIR = ROOT / "outputs/event_credibility/isr_candidate_audit_v1"
DEFAULT_REPORT_JSON = ROOT / "reports/task2_independent_error_audit_v1.json"
DEFAULT_REPORT_MD = ROOT / "reports/task2_independent_error_audit_v1.md"
DEFAULT_REPORT_XLSX = ROOT / "reports/task2_independent_error_audit_v1.xlsx"

FORMAL_OUTPUTS = [
    Path("outputs/event_credibility/llm_review_results.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.parquet"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.xlsx"),
]
FROZEN_FILES = [
    Path("src/event_credibility/dea.py"),
    Path("src/event_credibility/dea_safety_gate.py"),
    Path("configs/task2_dea_safety_gate_v1.yaml"),
    Path("src/event_credibility/dea_safety_gate_v1_1.py"),
    Path("configs/task2_dea_safety_gate_v1_1.yaml"),
    Path("src/event_credibility/dea_safety_gate_v1_2.py"),
    Path("configs/task2_dea_safety_gate_v1_2.yaml"),
    Path("src/event_credibility/dea_safety_gate_adapter.py"),
]

INVESTOR_PATTERN = re.compile(r"(互动易|e互动|投资者|股东|请问|董秘)")
MARKET_PATTERN = re.compile(r"(股价|涨跌|收盘|成交额|成交量|换手率|资金流|主力资金|融资买入|融资偿还|板块指数|技术指标|大宗交易)")
CONTRADICTION_PATTERN = re.compile(r"(否认|不属实|澄清|未发生|不存在|终止|取消|更正|与事实不符|传闻不实|未参与|没有计划)")
ENCODED_CLAIM_PATTERN = re.compile(r"\b(?:company|industry)_\d+\b")


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return norm(value).lower() in {"true", "1", "yes", "y", "是", "ok"}


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"{path} must contain an object")
    return data


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    rows: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        data = json.loads(line)
        if not isinstance(data, dict):
            raise ValueError(f"{path}:{line_number} is not an object")
        rows.append(data)
    return rows


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def hash_files(paths: Iterable[Path]) -> dict[str, str]:
    result: dict[str, str] = {}
    for rel_path in paths:
        path = rel_path if rel_path.is_absolute() else ROOT / rel_path
        result[str(rel_path)] = sha256(path) if path.exists() else "missing"
    return result


def unique_by_cluster(rows: list[dict[str, Any]], name: str, expected_count: int | None = None) -> dict[str, dict[str, Any]]:
    ids = [norm(row.get("cluster_id")) for row in rows]
    missing = [index + 1 for index, cid in enumerate(ids) if not cid]
    duplicates = sorted({cid for cid in ids if cid and ids.count(cid) > 1})
    if missing or duplicates:
        raise ValueError(f"{name} cluster_id invalid: missing={missing[:10]}, duplicates={duplicates[:10]}")
    if expected_count is not None and len(ids) != expected_count:
        raise ValueError(f"{name} expected {expected_count} rows, got {len(ids)}")
    return {norm(row["cluster_id"]): dict(row) for row in rows}


def assert_same_cluster_ids(mappings: dict[str, dict[str, Any]], expected_count: int = 50) -> list[str]:
    expected: set[str] | None = None
    for name, mapping in mappings.items():
        ids = set(mapping)
        if expected is None:
            expected = ids
        elif ids != expected:
            raise ValueError(
                f"cluster_id mismatch for {name}: missing={sorted(expected - ids)[:10]}, "
                f"unexpected={sorted(ids - expected)[:10]}"
            )
    if expected is None or len(expected) != expected_count:
        raise ValueError(f"expected {expected_count} strict cluster_id matches, got {len(expected or set())}")
    return sorted(expected)


def parse_json_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [norm(item) for item in value if norm(item)]
    text = norm(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        return [text]
    return [norm(item) for item in parsed] if isinstance(parsed, list) else [norm(parsed)]


def evidence_text(item: dict[str, Any]) -> str:
    return norm(item.get("text") or item.get("sentence") or item.get("content") or item.get("evidence_text"))


def evidence_source(item: dict[str, Any]) -> str:
    return norm(item.get("source") or item.get("article_source") or item.get("来源"))


def evidence_time(item: dict[str, Any]) -> str:
    return norm(item.get("publish_time") or item.get("article_publish_time") or item.get("发布时间"))


def evidence_id(item: dict[str, Any]) -> str:
    return norm(item.get("evidence_id") or item.get("id"))


def combine_evidence(*containers: Any) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for container in containers:
        if not isinstance(container, list):
            continue
        for item in container:
            if not isinstance(item, dict):
                continue
            eid = evidence_id(item)
            if not eid or eid in seen:
                continue
            seen.add(eid)
            result.append(dict(item))
    return result


def evidence_serialized(items: list[dict[str, Any]]) -> list[dict[str, str]]:
    return [
        {
            "evidence_id": evidence_id(item),
            "source": evidence_source(item),
            "publish_time": evidence_time(item),
            "text": evidence_text(item),
        }
        for item in items
    ]


def load_gold(path: Path) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    primary = table_records(path, "Primary人工审核")
    stress = table_records(path, "Contradiction压力审核")
    if len(primary) != 40 or len(stress) != 10:
        raise ValueError(f"gold workbook group counts invalid: primary={len(primary)}, stress={len(stress)}")
    rows = primary + stress
    mapping = unique_by_cluster(rows, "gold workbook", expected_count=50)
    for row in rows:
        if norm(row.get("reviewer_confirmed")).lower() != "yes":
            raise ValueError(f"{row['cluster_id']} reviewer_confirmed must be yes")
        if norm(row.get("human_verdict")) not in CLASSES:
            raise ValueError(f"{row['cluster_id']} has invalid human_verdict")
    return mapping, rows


def load_inputs(args: argparse.Namespace) -> dict[str, Any]:
    gold, gold_rows = load_gold(args.gold_file)
    evaluation = unique_by_cluster(
        read_jsonl(args.eval_dir / "independent_gold_results.jsonl"),
        "independent gold evaluation",
        expected_count=50,
    )
    samples = unique_by_cluster(
        read_jsonl(args.run_dir / "independent_samples.jsonl"),
        "independent samples",
        expected_count=50,
    )
    cached = unique_by_cluster(
        read_jsonl(args.run_dir / "cached_dea_results.jsonl"),
        "cached DEA results",
        expected_count=50,
    )
    gate = unique_by_cluster(
        read_jsonl(args.run_dir / "safety_gate_v1_2_results.jsonl"),
        "Safety Gate v1.2 results",
        expected_count=50,
    )
    model = unique_by_cluster(
        read_jsonl(args.run_dir / "model_request_log.jsonl"),
        "AI assistance results",
        expected_count=50,
    )
    cluster_ids = assert_same_cluster_ids(
        {
            "gold": gold,
            "independent_gold_evaluation": evaluation,
            "samples": samples,
            "cached": cached,
            "safety_gate": gate,
            "ai_assistance": model,
        }
    )
    evaluation_report = read_json(args.eval_report)
    dev_report = read_json(args.dev_report)
    dev_results = read_jsonl(args.dev_results)
    return {
        "gold": gold,
        "gold_rows": gold_rows,
        "evaluation": evaluation,
        "samples": samples,
        "cached": cached,
        "gate": gate,
        "model": model,
        "cluster_ids": cluster_ids,
        "evaluation_report": evaluation_report,
        "dev_report": dev_report,
        "dev_results": dev_results,
    }


def validate_evaluation_consistency(inputs: dict[str, Any]) -> dict[str, Any]:
    errors: list[str] = []
    cached_wrong: set[str] = set()
    gate_wrong: set[str] = set()
    cross_cluster_id_count = 0
    invalid_verdict_count = 0
    evidence_id_sets: dict[str, set[str]] = {}

    for cid in inputs["cluster_ids"]:
        sample = inputs["samples"][cid]
        cached = inputs["cached"][cid]
        gate = inputs["gate"][cid]
        gold = inputs["gold"][cid]
        evaluation = inputs["evaluation"][cid]

        all_evidence = combine_evidence(
            sample.get("evidence_items"),
            cached.get("initial_evidence"),
            cached.get("additional_evidence"),
            cached.get("final_evidence"),
            gate.get("evidence_records"),
        )
        evidence_ids = {evidence_id(item) for item in all_evidence if evidence_id(item)}
        evidence_id_sets[cid] = evidence_ids
        cited_ids = parse_json_list(gate.get("cited_evidence_ids"))
        unmatched = sorted(set(cited_ids) - evidence_ids)
        if unmatched:
            cross_cluster_id_count += len(unmatched)
            errors.append(f"{cid} has untraceable cited evidence IDs: {unmatched}")
        if parse_json_list(gate.get("unmatched_cited_evidence_ids")):
            errors.append(f"{cid} has gate-reported unmatched evidence IDs")

        verdicts = {
            "human": norm(gold.get("human_verdict")),
            "baseline": norm(evaluation.get("baseline_verdict")),
            "round1": norm(evaluation.get("round1_verdict")),
            "cached": norm(evaluation.get("cached_verdict")),
            "safety_gate_v1_2": norm(evaluation.get("safety_gate_v1_2_verdict")),
        }
        bad = {name: value for name, value in verdicts.items() if value not in CLASSES}
        if bad:
            invalid_verdict_count += len(bad)
            errors.append(f"{cid} has invalid verdicts: {bad}")
        if verdicts["cached"] != norm(cached.get("final_verdict")):
            errors.append(f"{cid} cached verdict does not match cached DEA output")
        if verdicts["safety_gate_v1_2"] != norm(gate.get("safety_gate_v1_2_verdict")):
            errors.append(f"{cid} safety verdict does not match gate output")
        if verdicts["cached"] != verdicts["human"]:
            cached_wrong.add(cid)
        if verdicts["safety_gate_v1_2"] != verdicts["human"]:
            gate_wrong.add(cid)

    if len(cached_wrong) != 13 or len(gate_wrong) != 13:
        errors.append(f"expected 13 cached/gate errors; got cached={len(cached_wrong)}, gate={len(gate_wrong)}")
    if cached_wrong != gate_wrong:
        errors.append(
            f"cached/gate error sets differ: cached_only={sorted(cached_wrong - gate_wrong)}, "
            f"gate_only={sorted(gate_wrong - cached_wrong)}"
        )
    report_count = int(inputs["evaluation_report"].get("error_analysis", {}).get("error_count", -1))
    if report_count != 13:
        errors.append(f"evaluation report error_count must be 13, got {report_count}")
    if errors:
        raise ValueError("; ".join(errors))
    return {
        "error_cluster_ids": sorted(gate_wrong),
        "error_count": len(gate_wrong),
        "invalid_verdict_count": invalid_verdict_count,
        "cross_cluster_evidence_reference_count": cross_cluster_id_count,
        "evidence_id_sets": evidence_id_sets,
    }


def is_investor_question(item: dict[str, Any]) -> bool:
    source = evidence_source(item)
    text = evidence_text(item)
    return bool(INVESTOR_PATTERN.search(source)) and bool(INVESTOR_PATTERN.search(text))


def is_market_signal(item: dict[str, Any]) -> bool:
    return bool(MARKET_PATTERN.search(f"{evidence_source(item)} {evidence_text(item)}"))


def has_explicit_contradiction(item: dict[str, Any]) -> bool:
    return bool(CONTRADICTION_PATTERN.search(evidence_text(item)))


def assess_evidence(cited_items: list[dict[str, Any]], gate: dict[str, Any]) -> dict[str, Any]:
    investor_flags = [is_investor_question(item) for item in cited_items]
    market_flags = [is_market_signal(item) for item in cited_items]
    return {
        "cited_evidence_count": len(cited_items),
        "direct_evidence_ids": parse_json_list(gate.get("direct_evidence_ids")),
        "has_direct_evidence": bool(parse_json_list(gate.get("direct_evidence_ids"))),
        "has_explicit_contradiction": any(has_explicit_contradiction(item) for item in cited_items),
        "only_investor_questions": bool(cited_items) and all(investor_flags),
        "only_market_signals": bool(cited_items) and all(market_flags),
        "contains_any_investor_question": any(investor_flags),
        "contains_any_market_signal": any(market_flags),
        "gate_subject_match": bool_value(gate.get("cited_subject_match")),
        "gate_time_match": bool_value(gate.get("cited_time_match")),
        "gate_contains_concrete_fact": bool_value(gate.get("cited_contains_concrete_fact")),
        "gate_traceable": bool_value(gate.get("cited_traceable")),
    }


def structural_claim_issue(gold: dict[str, Any]) -> bool:
    claim = norm(gold.get("claim"))
    human_error = norm(gold.get("human_error_type"))
    return (
        norm(gold.get("human_verdict")) == "unclear"
        and (
            human_error == "claim_unclear"
            or bool(ENCODED_CLAIM_PATTERN.search(claim))
            or (not norm(gold.get("event_name")) and not norm(gold.get("event_time")))
        )
    )


def route_error(
    gold: dict[str, Any],
    evaluation: dict[str, Any],
    evidence: dict[str, Any],
) -> tuple[str, str, str, str, bool, str, str]:
    """Return route, stage, primary, secondary, ISR eligibility, action, note."""

    human = norm(gold.get("human_verdict"))
    cached = norm(evaluation.get("cached_verdict"))
    baseline = norm(evaluation.get("baseline_verdict"))
    round1 = norm(evaluation.get("round1_verdict"))
    evidence_sufficient = norm(gold.get("human_evidence_sufficient"))
    subject_match = norm(gold.get("human_subject_match"))
    time_match = norm(gold.get("human_time_match"))

    if structural_claim_issue(gold):
        return (
            "claim",
            "claim",
            "claim_unclear",
            "claim_construction_error",
            False,
            "先补全可验证的主体—动作/状态—事实—时间主张，再决定是否检索或推理。",
            "人工确认主张为unclear，且编码化/缺少具体对象或时间；ISR重复推理不能消除主张结构歧义。",
        )

    gate_miss = (
        cached in {"support", "contradict"}
        and cached != human
        and not bool_value(evaluation.get("gate_blocked"))
        and (
            evidence_sufficient == "no"
            or time_match == "no"
            or evidence["only_investor_questions"]
            or evidence["only_market_signals"]
        )
    )
    if gate_miss:
        if time_match == "no":
            primary, secondary = "time_mismatch", "evidence_irrelevant"
        elif evidence["only_investor_questions"]:
            primary, secondary = "investor_question_as_fact", "evidence_irrelevant"
        elif evidence["only_market_signals"] or evidence["contains_any_market_signal"]:
            primary, secondary = "market_signal_as_event", "evidence_irrelevant"
        else:
            primary, secondary = "evidence_irrelevant", "evidence_insufficient"
        return (
            "safety_gate",
            "safety_gate",
            primary,
            secondary,
            False,
            "在后续独立设计中审计“具体事实”与事件语义/时间的一致性；风险命中时优先回退Round1。",
            "Cached给出support/contradict但人工认为证据不足或时间不符；当前门控没有把这类泛化事实识别为高风险。",
        )

    if (
        human in {"support", "contradict"}
        and evidence_sufficient == "yes"
        and subject_match == "yes"
        and time_match in {"yes", "not_applicable"}
        and evidence["cited_evidence_count"] > 0
    ):
        if baseline == human and round1 != human:
            primary, secondary, stage = "evidence_ranking_error", "reasoning_error", "ranking"
            action = "保留现有证据，要求ISR重新排序并优先综合与事件语义直接匹配的事实证据。"
        else:
            primary, secondary, stage = "reasoning_error", "evidence_ranking_error", "reasoning"
            action = "保留现有证据，要求ISR在不新增检索的前提下做多证据综合和verdict选择复核。"
        return (
            "isr",
            stage,
            primary,
            secondary,
            True,
            action,
            "人工确认现有证据充足且主体/时间基本匹配，错误来自证据排序、综合或结论选择，适合离线ISR候选包。",
        )

    if human in {"support", "contradict"} and evidence_sufficient in {"no", "partial"}:
        return (
            "retrieval",
            "retrieval",
            "evidence_retrieval_miss",
            "evidence_insufficient",
            False,
            "优先补召回或补充权威事件证据；仅重复ISR推理预计不能解决。",
            "人工认为当前证据不足以支撑正确的support/contradict结论。",
        )

    return (
        "gold_recheck",
        "gold",
        "gold_label_uncertain",
        "other",
        False,
        "复核人工标签与证据充分性，再决定后续处理。",
        "现有字段无法在不做具体样本特例的前提下稳定归入ISR、检索、主张或门控路径。",
    )


def build_audit_records(inputs: dict[str, Any], error_ids: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for cid in error_ids:
        gold = inputs["gold"][cid]
        evaluation = inputs["evaluation"][cid]
        cached = inputs["cached"][cid]
        gate = inputs["gate"][cid]
        model = inputs["model"][cid]
        final_evidence = combine_evidence(
            cached.get("final_evidence"),
            cached.get("additional_evidence"),
            cached.get("initial_evidence"),
        )
        by_id = {evidence_id(item): item for item in final_evidence}
        cited_ids = parse_json_list(gate.get("cited_evidence_ids"))
        cited_items = [by_id[item] for item in cited_ids if item in by_id]
        evidence = assess_evidence(cited_items, gate)
        route, stage, primary, secondary, isr_eligible, action, explanation = route_error(
            gold,
            evaluation,
            evidence,
        )
        if primary not in ERROR_TYPES or secondary not in ERROR_TYPES or stage not in ERROR_STAGES:
            raise ValueError(f"invalid diagnostic enum for {cid}")
        rows.append(
            {
                "cluster_id": cid,
                "sample_group": norm(gold.get("sample_group")),
                "claim": norm(gold.get("claim")),
                "human_verdict": norm(gold.get("human_verdict")),
                "baseline_verdict": norm(evaluation.get("baseline_verdict")),
                "round1_verdict": norm(evaluation.get("round1_verdict")),
                "cached_verdict": norm(evaluation.get("cached_verdict")),
                "safety_gate_v1_2_verdict": norm(evaluation.get("safety_gate_v1_2_verdict")),
                "AI辅助建议_verdict": norm(model.get("verdict") or gold.get("AI辅助建议_verdict")),
                "AI辅助建议_confidence": model.get("confidence", gold.get("AI辅助建议_confidence", "")),
                "AI辅助建议_reason": norm(model.get("reason") or gold.get("AI辅助建议_reason")),
                "cited_evidence_ids": cited_ids,
                "final_evidence": evidence_serialized(final_evidence),
                "cited_evidence": evidence_serialized(cited_items),
                "human_subject_match": norm(gold.get("human_subject_match")),
                "human_time_match": norm(gold.get("human_time_match")),
                "gate_subject_match": evidence["gate_subject_match"],
                "gate_time_match": evidence["gate_time_match"],
                "has_direct_evidence": evidence["has_direct_evidence"],
                "direct_evidence_ids": evidence["direct_evidence_ids"],
                "has_explicit_contradiction": evidence["has_explicit_contradiction"],
                "only_market_signals": evidence["only_market_signals"],
                "only_investor_questions": evidence["only_investor_questions"],
                "contains_any_market_signal": evidence["contains_any_market_signal"],
                "contains_any_investor_question": evidence["contains_any_investor_question"],
                "human_evidence_sufficient": norm(gold.get("human_evidence_sufficient")),
                "human_error_type": norm(gold.get("human_error_type")),
                "human_notes": norm(gold.get("human_notes")),
                "gate_blocked": bool_value(evaluation.get("gate_blocked")),
                "gate_reasons": parse_json_list(evaluation.get("gate_reasons")),
                "error_stage": stage,
                "primary_error_type": primary,
                "secondary_error_type": secondary,
                "route": route,
                "route_label": ROUTES[route],
                "isr_suitable": isr_eligible,
                "isr_corrective_action": action,
                "diagnosis": explanation,
            }
        )
    return sorted(rows, key=lambda row: row["cluster_id"])


def validate_routes(records: list[dict[str, Any]], expected_error_ids: set[str]) -> None:
    route_ids: dict[str, set[str]] = {route: set() for route in ROUTES}
    ids = set()
    for row in records:
        cid = norm(row.get("cluster_id"))
        if cid in ids:
            raise ValueError(f"duplicate audited error: {cid}")
        ids.add(cid)
        route = norm(row.get("route"))
        if route not in ROUTES:
            raise ValueError(f"{cid} has invalid route: {route}")
        route_ids[route].add(cid)
        if route == "isr":
            if not row.get("isr_suitable"):
                raise ValueError(f"{cid} ISR route must be marked suitable")
            if norm(row.get("human_evidence_sufficient")) != "yes":
                raise ValueError(f"{cid} ISR route requires existing sufficient evidence")
            if norm(row.get("human_subject_match")) != "yes":
                raise ValueError(f"{cid} ISR route requires subject match")
            if norm(row.get("human_time_match")) not in {"yes", "not_applicable"}:
                raise ValueError(f"{cid} ISR route requires time match or not_applicable")
    if ids != expected_error_ids:
        raise ValueError(
            f"route assignment incomplete: missing={sorted(expected_error_ids - ids)}, "
            f"unexpected={sorted(ids - expected_error_ids)}"
        )
    if sum(len(items) for items in route_ids.values()) != len(ids):
        raise ValueError("routes are not mutually exclusive")


def gate_input_distribution(rows: list[dict[str, Any]]) -> dict[str, Any]:
    fields = (
        "gate_blocked",
        "raw_claim_incomplete",
        "evidence_resolves_claim",
        "structurally_unverifiable",
        "cited_subject_match",
        "cited_time_match",
        "cited_contains_concrete_fact",
        "cited_only_investor_questions",
        "cited_only_market_signals",
        "cited_traceable",
        "human_fields_used",
    )
    bool_counts = {
        field: {
            "true": sum(1 for row in rows if bool_value(row.get(field))),
            "false": sum(1 for row in rows if not bool_value(row.get(field))),
        }
        for field in fields
    }
    cited_sizes = [int(row.get("cited_evidence_scope_size") or 0) for row in rows]
    direct_count = sum(1 for row in rows if parse_json_list(row.get("direct_evidence_ids")))
    return {
        "sample_count": len(rows),
        "boolean_counts": bool_counts,
        "direct_evidence_sample_count": direct_count,
        "direct_evidence_id_count": sum(len(parse_json_list(row.get("direct_evidence_ids"))) for row in rows),
        "avg_cited_evidence_scope_size": round(sum(cited_sizes) / len(cited_sizes), 6) if cited_sizes else 0.0,
        "gate_reason_counts": dict(Counter(reason for row in rows for reason in parse_json_list(row.get("gate_reasons")))),
        "selected_result_source_distribution": dict(
            Counter(norm(row.get("safety_gate_v1_2_selected_source") or row.get("selected_result_source")) for row in rows)
        ),
        "unmatched_cited_evidence_id_count": sum(
            len(parse_json_list(row.get("unmatched_cited_evidence_ids"))) for row in rows
        ),
    }


def zero_trigger_diagnosis(inputs: dict[str, Any], records: list[dict[str, Any]]) -> dict[str, Any]:
    independent_rows = [inputs["gate"][cid] for cid in inputs["cluster_ids"]]
    dev_rows = inputs["dev_results"]
    independent = gate_input_distribution(independent_rows)
    development = gate_input_distribution(dev_rows)
    safety_missed = [row for row in records if row["route"] == "safety_gate"]
    return {
        "independent": independent,
        "development": development,
        "development_gate_blocked": sum(1 for row in dev_rows if bool_value(row.get("gate_blocked"))),
        "independent_gate_blocked": sum(1 for row in independent_rows if bool_value(row.get("gate_blocked"))),
        "input_field_coverage": {
            "gate_rows": len(independent_rows),
            "evidence_records_nonempty": sum(1 for row in independent_rows if isinstance(row.get("evidence_records"), list) and row["evidence_records"]),
            "cited_evidence_ids_nonempty": sum(1 for row in independent_rows if parse_json_list(row.get("cited_evidence_ids"))),
            "unmatched_cited_evidence_ids": independent["unmatched_cited_evidence_id_count"],
            "human_fields_used": independent["boolean_counts"]["human_fields_used"],
            "safety_gate_version_distribution": dict(Counter(norm(row.get("safety_gate_version")) for row in independent_rows)),
        },
        "why_zero": [
            "当前独立集的门控输入中，50/50均为 cited_traceable、主体匹配、时间匹配、包含具体事实；投资者提问-only和行情信号-only均为0。",
            "所有50条都被标记为 evidence_resolves_claim=true，structurally_unverifiable=true为0，因此没有unclear风险路由。",
            "引用证据ID无跨cluster未匹配，evidence_records和引用字段均存在，未发现字段缺失或adapter映射失败。",
            "direct evidence仅为诊断信号，独立集有命中但不单独控制最终verdict；零触发不是由direct evidence开关造成。",
            "开发集的2次阻断分别来自 cited_contains_concrete_fact=false 和contradict缺少明确反转；独立集在当前输入特征下没有这些条件。",
        ],
        "missed_gate_trigger_count": len(safety_missed),
        "missed_gate_trigger_cluster_ids": [row["cluster_id"] for row in safety_missed],
        "missed_gate_interpretation": (
            "这些是事后人工诊断的门控漏触发候选：泛化“具体事实”或时间匹配标记通过，"
            "但事实与主张事件语义不足或人工时间判断不一致。该结论不等于本轮修改门控。"
        ),
    }


def unique_contradict_analysis(inputs: dict[str, Any]) -> dict[str, Any]:
    rows = [row for row in inputs["gold_rows"] if norm(row.get("human_verdict")) == "contradict"]
    if len(rows) != 1:
        raise ValueError(f"expected exactly one contradict gold sample, got {len(rows)}")
    gold = rows[0]
    cid = norm(gold["cluster_id"])
    evaluation = inputs["evaluation"][cid]
    cached = inputs["cached"][cid]
    gate = inputs["gate"][cid]
    final_evidence = combine_evidence(
        cached.get("final_evidence"),
        cached.get("additional_evidence"),
        cached.get("initial_evidence"),
    )
    cited_ids = parse_json_list(gate.get("cited_evidence_ids"))
    by_id = {evidence_id(item): item for item in final_evidence}
    cited = [by_id[item] for item in cited_ids if item in by_id]
    return {
        "cluster_id": cid,
        "sample_group": norm(gold.get("sample_group")),
        "claim": norm(gold.get("claim")),
        "human_verdict": "contradict",
        "baseline_verdict": norm(evaluation.get("baseline_verdict")),
        "round1_verdict": norm(evaluation.get("round1_verdict")),
        "cached_verdict": norm(evaluation.get("cached_verdict")),
        "safety_gate_v1_2_verdict": norm(evaluation.get("safety_gate_v1_2_verdict")),
        "ai_assistance_verdict": norm(inputs["model"][cid].get("verdict")),
        "cited_evidence_ids": cited_ids,
        "cited_evidence": evidence_serialized(cited),
        "has_explicit_contradiction": any(has_explicit_contradiction(item) for item in cited),
        "methods_correct": {
            name: norm(evaluation.get(f"{name}_verdict")) == "contradict"
            for name in ("baseline", "round1", "cached")
        }
        | {"safety_gate_v1_2": norm(evaluation.get("safety_gate_v1_2_verdict")) == "contradict"},
        "diagnosis": (
            "引用证据含公司公告中“主营业务及基本面未发生重大变化”等明确反向事实；"
            "Baseline、Cached、Safety Gate和AI辅助建议均正确，Round1为insufficient。"
        ),
        "limitation": "独立金标仅有1条contradict样本，不能据此推断整体反证能力。",
    }


def compact_candidate(row: dict[str, Any]) -> dict[str, Any]:
    keys = (
        "cluster_id",
        "sample_group",
        "claim",
        "human_verdict",
        "baseline_verdict",
        "round1_verdict",
        "cached_verdict",
        "safety_gate_v1_2_verdict",
        "cited_evidence_ids",
        "cited_evidence",
        "human_subject_match",
        "human_time_match",
        "human_evidence_sufficient",
        "primary_error_type",
        "secondary_error_type",
        "error_stage",
        "route",
        "route_label",
        "isr_suitable",
        "isr_corrective_action",
        "diagnosis",
    )
    return {key: row.get(key) for key in keys}


def ensure_empty(path: Path, overwrite: bool) -> None:
    if path.exists() and path.is_dir() and any(path.iterdir()) and not overwrite:
        raise FileExistsError(f"{path} is non-empty; pass --overwrite to replace audit output")
    if path.exists() and path.is_file() and not overwrite:
        raise FileExistsError(f"{path} exists; pass --overwrite to replace it")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def stringify(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    return norm(value)


def write_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({header: stringify(row.get(header, "")) for header in headers})


def write_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = workbook.create_sheet(title=title[:31])
    if not rows:
        rows = [{"empty": ""}]
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([stringify(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, 1):
        longest = max([len(header)] + [len(stringify(row.get(header, ""))) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 72)


def write_xlsx(path: Path, audit: dict[str, Any]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    records = audit["error_records"]
    write_sheet(workbook, "13条错误总表", records)
    write_sheet(workbook, "ISR候选", audit["route_records"]["isr"])
    write_sheet(workbook, "检索修复候选", audit["route_records"]["retrieval"])
    write_sheet(workbook, "主张修复候选", audit["route_records"]["claim"])
    write_sheet(workbook, "门控漏触发候选", audit["route_records"]["safety_gate"])
    write_sheet(workbook, "金标待复核", audit["route_records"]["gold_recheck"])
    write_sheet(workbook, "零触发诊断", flatten_mapping(audit["zero_trigger_diagnosis"]))
    write_sheet(workbook, "唯一反证样本", [audit["unique_contradict"]])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def flatten_mapping(mapping: dict[str, Any], prefix: str = "") -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for key, value in mapping.items():
        item = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            rows.extend(flatten_mapping(value, item))
        else:
            rows.append({"item": item, "value": stringify(value)})
    return rows


def render_markdown(audit: dict[str, Any]) -> str:
    summary = audit["summary"]
    zero = audit["zero_trigger_diagnosis"]
    lines = [
        "# 独立金标错误审计与 ISR 候选路由诊断",
        "",
        "本报告为独立集事后诊断；未调用32B模型，未运行ISR，未修改DEA、Safety Gate、配置、金标或正式输出。",
        "",
        "## 输入校验",
        "",
        f"- 50条按cluster_id严格匹配: {summary['strict_match_ok']}",
        f"- Cached与Safety Gate v1.2错误集合一致: {summary['cached_gate_error_sets_identical']}",
        f"- 错误样本数: {summary['error_count']}",
        f"- 非法verdict: {summary['invalid_verdict_count']}",
        f"- 跨cluster证据引用: {summary['cross_cluster_evidence_reference_count']}",
        "",
        "## 诊断分布",
        "",
        f"- 主错误类型: `{json.dumps(summary['primary_error_type_distribution'], ensure_ascii=False)}`",
        f"- 错误阶段: `{json.dumps(summary['error_stage_distribution'], ensure_ascii=False)}`",
        f"- 路由: `{json.dumps(summary['route_distribution'], ensure_ascii=False)}`",
        "",
        "## 门控零触发",
        "",
        f"- 开发集/独立集 gate_blocked: {zero['development_gate_blocked']}/{zero['independent_gate_blocked']}",
        f"- 事后诊断门控漏触发候选: {zero['missed_gate_trigger_count']}，"
        f"`{', '.join(zero['missed_gate_trigger_cluster_ids']) or '无'}`",
    ]
    for item in zero["why_zero"]:
        lines.append(f"- {item}")
    lines.extend(
        [
            "",
            "## 唯一反证样本",
            "",
            f"- cluster_id: {audit['unique_contradict']['cluster_id']}",
            f"- 方法正确性: `{json.dumps(audit['unique_contradict']['methods_correct'], ensure_ascii=False)}`",
            f"- 结论: {audit['unique_contradict']['diagnosis']}",
            f"- 限制: {audit['unique_contradict']['limitation']}",
            "",
            "## ISR建议",
            "",
            f"- 推荐 ISR 候选: {summary['recommended_isr_candidate_count']}",
            f"- 当前不建议进入ISR灰度: {not summary['recommend_enter_isr_gray']}",
            "- 原因：独立集门控零触发，且仅有13条事后错误；应先人工复核候选包，再决定是否以不超过候选数量的离线ISR实验验证。",
            "",
        ]
    )
    return "\n".join(lines)


def build_audit(inputs: dict[str, Any]) -> dict[str, Any]:
    validation = validate_evaluation_consistency(inputs)
    records = build_audit_records(inputs, validation["error_cluster_ids"])
    validate_routes(records, set(validation["error_cluster_ids"]))
    route_records = {
        route: [compact_candidate(row) for row in records if row["route"] == route]
        for route in ROUTES
    }
    zero = zero_trigger_diagnosis(inputs, records)
    unique_contradict = unique_contradict_analysis(inputs)
    return {
        "summary": {
            "step": "4.4-I",
            "model_calls": 0,
            "isr_executed": False,
            "strict_match_ok": True,
            "cached_gate_error_sets_identical": True,
            "error_count": len(records),
            "invalid_verdict_count": validation["invalid_verdict_count"],
            "cross_cluster_evidence_reference_count": validation["cross_cluster_evidence_reference_count"],
            "primary_error_type_distribution": dict(
                Counter(row["primary_error_type"] for row in records)
            ),
            "error_stage_distribution": dict(Counter(row["error_stage"] for row in records)),
            "route_distribution": {
                ROUTES[route]: len(route_records[route]) for route in ROUTES
            },
            "recommended_isr_candidate_count": len(route_records["isr"]),
            "recommend_enter_isr_gray": False,
            "gate_blocked_independent": zero["independent_gate_blocked"],
            "gate_blocked_development": zero["development_gate_blocked"],
            "missed_gate_trigger_count": zero["missed_gate_trigger_count"],
        },
        "validation": {
            key: value for key, value in validation.items() if key != "evidence_id_sets"
        },
        "error_records": records,
        "route_records": route_records,
        "zero_trigger_diagnosis": zero,
        "unique_contradict": unique_contradict,
    }


def run_audit(args: argparse.Namespace) -> dict[str, Any]:
    formal_before = hash_files(FORMAL_OUTPUTS)
    frozen_before = hash_files(FROZEN_FILES)
    inputs = load_inputs(args)
    audit = build_audit(inputs)

    ensure_empty(args.output_dir, args.overwrite)
    for path in (args.report_json, args.report_md, args.report_xlsx):
        ensure_empty(path, args.overwrite)
    args.output_dir.mkdir(parents=True, exist_ok=True)

    write_jsonl(args.output_dir / "isr_candidates.jsonl", audit["route_records"]["isr"])
    write_jsonl(args.output_dir / "retrieval_fix_candidates.jsonl", audit["route_records"]["retrieval"])
    write_jsonl(args.output_dir / "claim_fix_candidates.jsonl", audit["route_records"]["claim"])
    write_jsonl(args.output_dir / "safety_gate_fix_candidates.jsonl", audit["route_records"]["safety_gate"])
    write_jsonl(args.output_dir / "gold_recheck_candidates.jsonl", audit["route_records"]["gold_recheck"])
    write_csv(
        args.output_dir / "all_error_audit.csv",
        audit["error_records"],
        [
            "cluster_id",
            "sample_group",
            "claim",
            "human_verdict",
            "baseline_verdict",
            "round1_verdict",
            "cached_verdict",
            "safety_gate_v1_2_verdict",
            "error_stage",
            "primary_error_type",
            "secondary_error_type",
            "route",
            "isr_suitable",
            "diagnosis",
        ],
    )

    formal_after = hash_files(FORMAL_OUTPUTS)
    frozen_after = hash_files(FROZEN_FILES)
    audit["summary"]["formal_outputs_unchanged"] = formal_before == formal_after
    audit["summary"]["frozen_files_unchanged"] = frozen_before == frozen_after
    audit["summary"]["formal_output_sha256_before"] = formal_before
    audit["summary"]["formal_output_sha256_after"] = formal_after
    audit["summary"]["frozen_file_sha256_before"] = frozen_before
    audit["summary"]["frozen_file_sha256_after"] = frozen_after
    audit["summary"]["output_dir"] = str(args.output_dir)
    audit["summary"]["report_json"] = str(args.report_json)
    audit["summary"]["report_md"] = str(args.report_md)
    audit["summary"]["report_xlsx"] = str(args.report_xlsx)
    if not audit["summary"]["formal_outputs_unchanged"]:
        raise RuntimeError("formal output SHA256 changed during independent error audit")
    if not audit["summary"]["frozen_files_unchanged"]:
        raise RuntimeError("frozen file SHA256 changed during independent error audit")

    args.report_json.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.report_md.write_text(render_markdown(audit), encoding="utf-8")
    write_xlsx(args.report_xlsx, audit)
    return audit


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--gold-file", type=Path, default=DEFAULT_GOLD_FILE)
    parser.add_argument("--eval-dir", type=Path, default=DEFAULT_EVAL_DIR)
    parser.add_argument("--run-dir", type=Path, default=DEFAULT_RUN_DIR)
    parser.add_argument("--eval-report", type=Path, default=DEFAULT_EVAL_REPORT)
    parser.add_argument("--dev-results", type=Path, default=DEFAULT_DEV_RESULTS)
    parser.add_argument("--dev-report", type=Path, default=DEFAULT_DEV_REPORT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--report-json", type=Path, default=DEFAULT_REPORT_JSON)
    parser.add_argument("--report-md", type=Path, default=DEFAULT_REPORT_MD)
    parser.add_argument("--report-xlsx", type=Path, default=DEFAULT_REPORT_XLSX)
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    audit = run_audit(args)
    print(
        json.dumps(
            {
                "error_count": audit["summary"]["error_count"],
                "route_distribution": audit["summary"]["route_distribution"],
                "gate_blocked_independent": audit["summary"]["gate_blocked_independent"],
                "missed_gate_trigger_count": audit["summary"]["missed_gate_trigger_count"],
                "model_calls": 0,
                "isr_executed": False,
                "formal_outputs_unchanged": audit["summary"]["formal_outputs_unchanged"],
                "frozen_files_unchanged": audit["summary"]["frozen_files_unchanged"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
