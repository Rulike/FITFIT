#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Task1 Step 5.6-D official-data classifier splits and pair data."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_task1_step5_5_candidate_v2 import (  # noqa: E402
    candidate_rows_for_top_k,
    load_config,
    load_frozen_split,
    protected_file_list,
    read_jsonl,
    rel,
    sha256_file,
    write_csv,
    write_json,
    write_jsonl,
)
from src.evaluation.task1_task_metrics import candidate_recall_at_k  # noqa: E402

OUTPUT_DIR = ROOT / "outputs/task1_step5_6d_classifier_retraining"
CONFIG_PATH = ROOT / "configs/task1_candidate_v2.yaml"
CONFIRMED_80 = ROOT / "outputs/task1_structure_validation_round1/minimal_human_review_package_confirmed.jsonl"
GOLD_FREE_TEST_INPUT = ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl"
FORBIDDEN_GOLD_TEST_LABEL_PATHS = {
    "data/processed/labeled/news_test.jsonl",
    "data/processed/labeled/event_pairs_test.jsonl",
    "data/processed/model_data/event_match_official_test.jsonl",
    "data/raw/news_with_tags_test.xlsx",
}


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def safe_sha(path: Path) -> str:
    return sha256_file(path) if path.exists() and path.is_file() else ""


def assert_not_forbidden(path: Path) -> None:
    try:
        path_text = rel(path)
    except Exception:
        path_text = path.as_posix()
    if path_text in FORBIDDEN_GOLD_TEST_LABEL_PATHS:
        raise RuntimeError(f"forbidden gold-test label path requested: {path_text}")


def read_jsonl_safe(path: Path) -> list[dict[str, Any]]:
    assert_not_forbidden(path)
    return read_jsonl(path)


def write_sha_manifest(output_dir: Path) -> None:
    rows = []
    for path in sorted(output_dir.rglob("*")):
        if path.is_file() and path.name != "SHA256SUMS.txt":
            rows.append((rel(path), sha256_file(path)))
    (output_dir / "SHA256SUMS.txt").write_text(
        "".join(f"{digest}  {name}\n" for name, digest in rows),
        encoding="utf-8",
    )


def protected_snapshot(path: Path) -> dict[str, str]:
    files = protected_file_list()
    rows = {rel(item): sha256_file(item) for item in files}
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(rows.items())), encoding="utf-8")
    return rows


def xlsx_shape(path: Path) -> dict[str, Any]:
    assert_not_forbidden(path)
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover
        return {"path": rel(path), "exists": path.exists(), "status": "openpyxl_missing", "error": repr(exc)}
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    result = {
        "path": rel(path),
        "exists": path.exists(),
        "sheet": ws.title,
        "row_count_excluding_header": max(0, int(ws.max_row or 0) - 1),
        "column_count": int(ws.max_column or 0),
        "headers": headers,
        "has_event_label_columns": any(str(header or "") in {"event_name", "event_emotion", "event_impact_analysis"} for header in headers),
        "sha256": safe_sha(path),
    }
    wb.close()
    return result


def load_event_names(config: dict[str, Any]) -> list[str]:
    return sorted({str(row["event_name"]) for row in read_jsonl_safe(ROOT / config["schema_path"]) if row.get("event_name")})


def schema_description_by_event(config: dict[str, Any]) -> dict[str, str]:
    out: dict[str, str] = {}
    for row in read_jsonl_safe(ROOT / config["schema_path"]):
        event_name = str(row.get("event_name") or "")
        if event_name and event_name not in out:
            out[event_name] = str(row.get("raw_description") or row.get("event_description") or "")
    return out


def label_sets(pairs: Iterable[dict[str, Any]]) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    for row in pairs:
        if int(row.get("label", 1)) == 1 and row.get("event_name"):
            out[str(row["article_id"])].add(str(row["event_name"]))
    return out


def pair_rows_by_article(pairs: Iterable[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in pairs:
        if int(row.get("label", 1)) == 1:
            out[str(row["article_id"])].append(row)
    return out


def candidate_map(rows: Iterable[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen = set()
    for row in rows:
        key = (str(row.get("article_id") or ""), str(row.get("event_name") or ""))
        if key in seen:
            continue
        seen.add(key)
        out[key[0]].append(row)
    for values in out.values():
        values.sort(key=lambda item: (int(item.get("candidate_rank", 10**9) or 10**9), -float(item.get("candidate_score", 0.0) or 0.0), str(item.get("event_name", ""))))
    return out


def make_positive_candidate(pair: dict[str, Any], article: dict[str, Any], descriptions: dict[str, str]) -> dict[str, Any]:
    event_name = str(pair["event_name"])
    return {
        "article_id": str(article["article_id"]),
        "title": article.get("title", ""),
        "event_name": event_name,
        "schema_id": pair.get("schema_id", f"event::{event_name}"),
        "target_type": pair.get("target_type", ""),
        "target_code": str(pair.get("target_code", "")),
        "target_name": pair.get("target_name", ""),
        "event_description": pair.get("event_description") or descriptions.get(event_name, ""),
        "candidate_rank": 0,
        "candidate_score": 999.0,
        "candidate_source": "gold_positive_injected_for_training",
        "source_scores": {"gold_positive_injected_for_training": 1.0},
    }


def enrich_candidate(row: dict[str, Any], descriptions: dict[str, str]) -> dict[str, Any]:
    item = dict(row)
    event_name = str(item.get("event_name") or "")
    item["event_description"] = item.get("event_description") or item.get("raw_description") or descriptions.get(event_name, "")
    return item


def parse_time(value: Any) -> str:
    return str(value or "")


def choose_holdout(train_ids: list[str], article_by_id: dict[str, dict[str, Any]], holdout_ratio: float) -> tuple[list[str], list[str]]:
    ordered = sorted(train_ids, key=lambda aid: (parse_time(article_by_id[aid].get("publish_time")), aid))
    holdout_n = max(1, int(round(len(ordered) * holdout_ratio)))
    holdout = set(ordered[-holdout_n:])
    return [aid for aid in ordered if aid not in holdout], [aid for aid in ordered if aid in holdout]


def duplication_leakage(splits: dict[str, list[str]], article_by_id: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, set[str]] = defaultdict(set)
    for split_name, ids in splits.items():
        for aid in ids:
            dup = article_by_id[aid].get("duplication_id")
            if dup not in {None, "", "None"}:
                groups[str(dup)].add(split_name)
    return [{"duplication_id": dup, "splits": sorted(values)} for dup, values in sorted(groups.items()) if len(values) > 1]


def official_data_audit(config: dict[str, Any], news: list[dict[str, Any]], pairs: list[dict[str, Any]], ids_80: set[str], gold_free_test_ids: set[str]) -> dict[str, Any]:
    labels = [label for row in news for label in (row.get("labels") or [])]
    pair_positive = [row for row in pairs if int(row.get("label", 1)) == 1]
    news_ids = {str(row["article_id"]) for row in news}
    pair_ids = {str(row["article_id"]) for row in pairs}
    no_content = sum(1 for row in news if not str(row.get("content") or "").strip())
    duplicate_articles = [aid for aid, count in Counter(str(row["article_id"]) for row in news).items() if count > 1]
    return {
        "generated_at": now(),
        "official_train_sources": {
            "processed_news_train": {
                "path": "data/processed/labeled/news_train.jsonl",
                "sha256": safe_sha(ROOT / "data/processed/labeled/news_train.jsonl"),
                "article_count": len(news),
                "unique_article_count": len(news_ids),
                "with_event_label_article_count": sum(1 for row in news if row.get("labels")),
                "no_event_article_count": sum(1 for row in news if not row.get("labels")),
                "multi_label_article_count": sum(1 for row in news if len(row.get("labels") or []) > 1),
                "max_labels_per_article": max((len(row.get("labels") or []) for row in news), default=0),
                "article_without_content_count": no_content,
            },
            "processed_event_pairs_train": {
                "path": "data/processed/labeled/event_pairs_train.jsonl",
                "sha256": safe_sha(ROOT / "data/processed/labeled/event_pairs_train.jsonl"),
                "event_pair_count": len(pairs),
                "positive_event_pair_count": len(pair_positive),
                "negative_event_pair_count": len(pairs) - len(pair_positive),
                "unique_article_count": len(pair_ids),
                "event_category_count": len({str(row["event_name"]) for row in pair_positive}),
            },
            "raw_news_with_tags_train": xlsx_shape(ROOT / "data/raw/news_with_tags_train.xlsx"),
            "raw_news_without_tags_train": xlsx_shape(ROOT / "data/raw/news_without_tags_train.xlsx"),
        },
        "official_training_pool_near_200k": {
            "has_200k_unlabeled_news_pool": True,
            "unlabeled_pool_article_count": 200000,
            "supervised_classification_usable_labeled_articles": len(news),
            "remaining_official_training_articles_have_legal_event_labels": False,
            "remaining_official_training_articles_reason": "news_without_tags_train.xlsx has article fields only and no event_name/event label columns",
        },
        "why_step5_4_used_2046_articles": (
            "Step 5.4 used the complete processed labeled official training subset: "
            "2046 articles with 2081 positive event pairs. The near-200k official train pool is unlabeled and cannot be used as supervised event-label training data without weak-label assumptions."
        ),
        "only_gold_subset_used_for_supervised_training": True,
        "labels_and_schema": {
            "event_category_count": len({str(label["event_name"]) for label in labels}),
            "event_pair_count_from_news_labels": len(labels),
            "target_type_counts": dict(Counter(str(label.get("target_type") or "") for label in labels)),
            "event_category_support": dict(Counter(str(label.get("event_name") or "") for label in labels)),
        },
        "integrity": {
            "news_ids_minus_pair_ids": len(news_ids - pair_ids),
            "pair_ids_minus_news_ids": len(pair_ids - news_ids),
            "duplicate_article_id_count": len(duplicate_articles),
            "duplicate_article_ids": duplicate_articles[:100],
            "overlap_with_80_confirmed_count": len(news_ids & ids_80),
            "overlap_with_80_confirmed_ids": sorted(news_ids & ids_80),
            "overlap_with_gold_free_test_count": len(news_ids & gold_free_test_ids),
            "overlap_with_gold_free_test_ids": sorted(news_ids & gold_free_test_ids)[:100],
        },
        "gold_test_label_usage": {
            "gold_test_label_files_read": False,
            "test_article_ids_source": "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl",
            "gold_free_test_article_count": len(gold_free_test_ids),
        },
    }


def candidate_recall(article_ids: list[str], pairs: list[dict[str, Any]], candidates: list[dict[str, Any]], event_names: list[str]) -> dict[str, Any]:
    return candidate_recall_at_k(article_ids, pairs, candidates, event_names, [1, 3, 5, 10, 20])


def build_classifier_rows(
    split_name: str,
    article_ids: list[str],
    article_by_id: dict[str, dict[str, Any]],
    positive_pairs_by_article: dict[str, list[dict[str, Any]]],
    candidates_by_article: dict[str, list[dict[str, Any]]],
    descriptions: dict[str, str],
    rng: random.Random,
    training_mode: bool,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    rows: list[dict[str, Any]] = []
    stats = Counter()
    for aid in article_ids:
        positives = positive_pairs_by_article.get(aid, [])
        gold_names = {str(pair["event_name"]) for pair in positives}
        gold_target_codes = {str(pair.get("target_code") or "") for pair in positives if pair.get("target_code")}
        candidates = [enrich_candidate(row, descriptions) for row in candidates_by_article.get(aid, [])]
        by_event = {str(row["event_name"]): row for row in candidates}
        if not training_mode:
            for cand in candidates:
                item = dict(cand)
                item.update({"label": int(str(cand["event_name"]) in gold_names), "negative_type": "" if str(cand["event_name"]) in gold_names else "candidate_negative", "split": split_name, "sample_id": f"{split_name}_cand_{aid}_{cand['event_name']}"})
                rows.append(item)
            continue
        for pair in positives:
            event_name = str(pair["event_name"])
            item = dict(by_event.get(event_name, make_positive_candidate(pair, article_by_id[aid], descriptions)))
            item.update({"label": 1, "negative_type": "", "split": split_name, "sample_id": f"{split_name}_pos_{aid}_{event_name}"})
            if event_name not in by_event:
                stats["positive_gold_injected"] += 1
            else:
                stats["positive_from_candidate_v2"] += 1
            rows.append(item)
        negative_seen: set[str] = set()

        def add_negative(cand: dict[str, Any], neg_type: str) -> None:
            event_name = str(cand.get("event_name") or "")
            if not event_name or event_name in gold_names or event_name in negative_seen:
                return
            negative_seen.add(event_name)
            item = dict(cand)
            item.update({"label": 0, "negative_type": neg_type, "split": split_name, "sample_id": f"{split_name}_{neg_type}_{aid}_{event_name}"})
            rows.append(item)
            stats[neg_type] += 1

        for cand in candidates:
            if len([row for row in rows if row.get("article_id") == aid and row.get("negative_type") == "hard_top_rank"]) >= 8:
                break
            add_negative(cand, "hard_top_rank")
        semantic = [cand for cand in candidates if str(cand.get("target_code") or "") in gold_target_codes and str(cand.get("event_name") or "") not in gold_names]
        for cand in semantic[:3]:
            add_negative(cand, "semantic_same_target_hard_negative")
        # Event-name-only classification makes same-event/different-target rows label-conflicting positives, so this bucket is reported as zero.
        stats["target_same_event_different_target_hard_negative"] += 0
        rest = [cand for cand in candidates if str(cand.get("event_name") or "") not in gold_names and str(cand.get("event_name") or "") not in negative_seen]
        rng.shuffle(rest)
        for cand in rest[:2]:
            add_negative(cand, "random_negative")
    stats["rows_total"] = len(rows)
    stats["positive_total"] = sum(1 for row in rows if int(row.get("label", 0)) == 1)
    stats["negative_total"] = sum(1 for row in rows if int(row.get("label", 0)) == 0)
    return rows, dict(stats)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--candidate-config", type=Path, default=CONFIG_PATH)
    parser.add_argument("--holdout-ratio", type=float, default=0.10)
    parser.add_argument("--seed", type=int, default=2026)
    parser.add_argument("--backup-dir", type=Path, default=None)
    args = parser.parse_args()
    rng = random.Random(args.seed)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    protected_snapshot(args.output_dir / "formal_protected_sha256_before.txt")

    config = load_config(args.candidate_config)
    split = load_frozen_split(config)
    event_names = load_event_names(config)
    descriptions = schema_description_by_event(config)
    news = read_jsonl_safe(ROOT / "data/processed/labeled/news_train.jsonl")
    pairs = read_jsonl_safe(ROOT / "data/processed/labeled/event_pairs_train.jsonl")
    article_by_id = {str(row["article_id"]): row for row in news}
    ids_80 = {str(row["article_id"]) for row in read_jsonl_safe(CONFIRMED_80)}
    gold_free_test = read_jsonl_safe(GOLD_FREE_TEST_INPUT)
    gold_free_test_ids = {str(row["article_id"]) for row in gold_free_test}
    data_audit = official_data_audit(config, news, pairs, ids_80, gold_free_test_ids)
    write_json(args.output_dir / "data_audit.json", data_audit)
    if data_audit["official_train_sources"]["processed_news_train"]["with_event_label_article_count"] == 0:
        write_json(args.output_dir / "blocked_status.json", {"status": "blocked_by_insufficient_classification_training_data", "reason": "no official labeled training articles found"})
        write_sha_manifest(args.output_dir)
        print(json.dumps({"status": "blocked_by_insufficient_classification_training_data"}, ensure_ascii=False))
        return

    excluded = ids_80 | gold_free_test_ids
    base_train_ids = [aid for aid in map(str, split["train_ids"]) if aid not in excluded]
    dev_ids = [aid for aid in map(str, split["dev_ids"]) if aid not in excluded]
    train_ids, holdout_ids = choose_holdout(base_train_ids, article_by_id, args.holdout_ratio)
    leakage = duplication_leakage({"train": train_ids, "dev": dev_ids, "holdout": holdout_ids}, article_by_id)
    if leakage:
        raise RuntimeError(f"duplicate group leakage across splits: {leakage[:5]}")

    pos_pairs_by_article = pair_rows_by_article(pairs)
    split_pairs = {
        "train": [row for aid in train_ids for row in pos_pairs_by_article.get(aid, [])],
        "dev": [row for aid in dev_ids for row in pos_pairs_by_article.get(aid, [])],
        "holdout": [row for aid in holdout_ids for row in pos_pairs_by_article.get(aid, [])],
    }
    train_candidate_rows = read_jsonl_safe(ROOT / "outputs/task1_limited_training_round1/train_candidate_v2_top20.jsonl")
    dev_candidate_rows = candidate_rows_for_top_k(read_jsonl_safe(ROOT / "outputs/task1_step5_5_candidate_v2_dev/candidate_v2_dev_all64.jsonl"), 20)
    all_candidate_rows = train_candidate_rows + dev_candidate_rows
    candidates_by_article = candidate_map(all_candidate_rows)
    classifier_rows = {}
    neg_stats = {}
    for split_name, ids in [("train", train_ids), ("dev", dev_ids), ("holdout", holdout_ids)]:
        rows, stats = build_classifier_rows(
            split_name,
            ids,
            article_by_id,
            pos_pairs_by_article,
            candidates_by_article,
            descriptions,
            rng,
            training_mode=(split_name == "train"),
        )
        classifier_rows[split_name] = rows
        neg_stats[split_name] = stats

    write_jsonl(args.output_dir / "train_articles.jsonl", [article_by_id[aid] for aid in train_ids])
    write_jsonl(args.output_dir / "dev_articles.jsonl", [article_by_id[aid] for aid in dev_ids])
    write_jsonl(args.output_dir / "holdout_articles.jsonl", [article_by_id[aid] for aid in holdout_ids])
    for split_name, rows in split_pairs.items():
        write_jsonl(args.output_dir / f"{split_name}_gold_pairs.jsonl", rows)
    for split_name, rows in classifier_rows.items():
        write_jsonl(args.output_dir / f"{split_name}_classifier_pairs.jsonl", rows)
    write_jsonl(args.output_dir / "train_candidate_v2_top20.jsonl", [row for row in train_candidate_rows if str(row.get("article_id")) in set(train_ids) | set(holdout_ids)])
    write_jsonl(args.output_dir / "dev_candidate_v2_top20.jsonl", [row for row in dev_candidate_rows if str(row.get("article_id")) in set(dev_ids)])

    split_summary = {
        "generated_at": now(),
        "seed": args.seed,
        "split_policy": "Step5.4 frozen dev retained; internal holdout carved from latest 10% of Step5.4 train by publish_time",
        "train_article_count": len(train_ids),
        "dev_article_count": len(dev_ids),
        "holdout_article_count": len(holdout_ids),
        "train_positive_event_pair_count": len(split_pairs["train"]),
        "dev_positive_event_pair_count": len(split_pairs["dev"]),
        "holdout_positive_event_pair_count": len(split_pairs["holdout"]),
        "excluded_80_overlap_count": len((set(split["train_ids"]) | set(split["dev_ids"])) & ids_80),
        "excluded_gold_free_test_overlap_count": len((set(split["train_ids"]) | set(split["dev_ids"])) & gold_free_test_ids),
        "duplicate_group_leakage_count": len(leakage),
        "article_id_overlap": {
            "train_dev": len(set(train_ids) & set(dev_ids)),
            "train_holdout": len(set(train_ids) & set(holdout_ids)),
            "dev_holdout": len(set(dev_ids) & set(holdout_ids)),
        },
        "candidate_recall": {
            "train": candidate_recall(train_ids, split_pairs["train"], [row for row in train_candidate_rows if str(row.get("article_id")) in set(train_ids)], event_names),
            "dev": candidate_recall(dev_ids, split_pairs["dev"], [row for row in dev_candidate_rows if str(row.get("article_id")) in set(dev_ids)], event_names),
            "holdout": candidate_recall(holdout_ids, split_pairs["holdout"], [row for row in train_candidate_rows if str(row.get("article_id")) in set(holdout_ids)], event_names),
        },
        "class_support": {
            split_name: dict(Counter(str(row["event_name"]) for row in rows))
            for split_name, rows in split_pairs.items()
        },
        "sha256": {
            "train_article_ids": hashlib.sha256("\n".join(train_ids).encode("utf-8")).hexdigest(),
            "dev_article_ids": hashlib.sha256("\n".join(dev_ids).encode("utf-8")).hexdigest(),
            "holdout_article_ids": hashlib.sha256("\n".join(holdout_ids).encode("utf-8")).hexdigest(),
        },
    }
    write_json(args.output_dir / "split_summary.json", split_summary)
    write_json(
        args.output_dir / "hard_negative_statistics.json",
        {
            "generated_at": now(),
            "train_pair_generation": neg_stats["train"],
            "dev_pair_rows": neg_stats["dev"],
            "holdout_pair_rows": neg_stats["holdout"],
            "negative_sampling_policy": {
                "hard_top_rank_per_article": 8,
                "semantic_same_target_per_article": 3,
                "random_negative_per_article": 2,
                "target_same_event_different_target": "not used because Task1 classification label is event_name-only; same event_name remains positive",
            },
            "positive_negative_ratio_train": neg_stats["train"].get("positive_total", 0) / max(1, neg_stats["train"].get("negative_total", 0)),
        },
    )
    write_csv(
        args.output_dir / "class_support_by_split.csv",
        [
            {
                "event_name": event,
                "train_support": split_summary["class_support"]["train"].get(event, 0),
                "dev_support": split_summary["class_support"]["dev"].get(event, 0),
                "holdout_support": split_summary["class_support"]["holdout"].get(event, 0),
            }
            for event in event_names
        ],
    )
    protected_snapshot(args.output_dir / "formal_protected_sha256_after_build.txt")
    write_json(
        args.output_dir / "build_manifest.json",
        {
            "generated_at": now(),
            "status": "training_data_ready",
            "backup_dir": str(args.backup_dir) if args.backup_dir else "",
            "output_dir": rel(args.output_dir),
            "gold_test_label_files_read": False,
            "llm_calls": {"dianjin_r1_32b_calls": 0, "external_llm_calls": 0},
        },
    )
    write_sha_manifest(args.output_dir)
    print(json.dumps({"status": "training_data_ready", "train": len(train_ids), "dev": len(dev_ids), "holdout": len(holdout_ids)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
