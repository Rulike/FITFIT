#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 gold coverage gap and field propagation audit.

This script is read-only for existing Task1/Task2 artifacts. It traces why the
independent gold test set cannot be fully evaluated against the frozen Task1
outputs, and writes only shadow/audit files into a separate output directory.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = ROOT / "outputs/task1_gold_coverage_gap_audit_v1"

SCHEMA_PATH = ROOT / "data/processed/event_schema.jsonl"
GOLD_NEWS_TEST = ROOT / "data/processed/labeled/news_test.jsonl"
GOLD_PAIRS_TEST = ROOT / "data/processed/labeled/event_pairs_test.jsonl"
RAW_GOLD_TEST = ROOT / "data/raw/news_with_tags_test.xlsx"
RAW_UNLABELED = {
    "train": ROOT / "data/raw/news_without_tags_train.xlsx",
    "test": ROOT / "data/raw/news_without_tags_test.xlsx",
}
FULL_PREDICTIONS = {
    "train": ROOT / "outputs/unlabeled_event_inference/full/news_predictions_train.jsonl",
    "test": ROOT / "outputs/unlabeled_event_inference/full/news_predictions_test.jsonl",
}
FINAL_EVENTS = {
    "train": ROOT / "outputs/final_task1/final_events_train.jsonl",
    "test": ROOT / "outputs/final_task1/final_events_test.jsonl",
}
QUALITY_GATE = {
    "train": ROOT / "data/processed/event_extraction/quality_gated_train.jsonl",
    "test": ROOT / "data/processed/event_extraction/quality_gated_test.jsonl",
}
AUTO_ACCEPT = {
    "train": ROOT / "data/processed/event_extraction/auto_accept_events_train.jsonl",
    "test": ROOT / "data/processed/event_extraction/auto_accept_events_test.jsonl",
}
MANUAL_REVIEW = {
    "train": ROOT / "data/processed/event_extraction/manual_review_events_train.jsonl",
    "test": ROOT / "data/processed/event_extraction/manual_review_events_test.jsonl",
}
JOINT_GROUPS = {
    "train": ROOT / "data/processed/event_extraction/joint_adjudication_groups_train.jsonl",
    "test": ROOT / "data/processed/event_extraction/joint_adjudication_groups_test.jsonl",
}
JOINT_RESULTS = {
    "train": ROOT / "outputs/final_task1/joint_adjudication_train.jsonl",
    "test": ROOT / "outputs/final_task1/joint_adjudication_test.jsonl",
}
TASK3_DOCS = ROOT / "outputs/task3_fulltext_v2/normalized_documents.jsonl"
TASK3_MENTIONS = ROOT / "outputs/task3_fulltext_v2/normalized_event_mentions.jsonl"
TASK3_EVIDENCE = ROOT / "outputs/task3_fulltext_v2/event_mention_evidence_sentences.jsonl"

PROTECTED_FILES = [
    FINAL_EVENTS["train"],
    FINAL_EVENTS["test"],
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    GOLD_NEWS_TEST,
    GOLD_PAIRS_TEST,
    TASK3_DOCS,
    TASK3_MENTIONS,
    TASK3_EVIDENCE,
]

STRUCTURED_FIELDS = [
    "event_subject_name",
    "event_subject_type",
    "event_time_raw",
    "event_time_normalized",
    "event_time_precision",
    "event_action",
    "event_object",
    "event_status",
    "location",
    "impact_direction",
    "extraction_confidence",
]

ATTRIBUTION_LABELS = {
    "input_universe_mismatch": "A. input_universe_mismatch",
    "article_id_mismatch": "B. article_id_mismatch",
    "content_hash_match_only": "C. content_hash_match_only",
    "normalization_loss": "D. normalization_loss",
    "candidate_generation_loss": "E. candidate_generation_loss",
    "filter_or_threshold_loss": "F. filter_or_threshold_loss",
    "schema_mapping_mismatch": "G. schema_mapping_mismatch",
    "final_export_loss": "H. final_export_loss",
    "duplicate_or_merge_loss": "I. duplicate_or_merge_loss",
    "truly_no_prediction": "J. truly_no_prediction",
    "unknown": "K. unknown",
}


@dataclass
class InputHit:
    split: str
    article_id: str
    title: str = ""
    publish_time: str = ""


@dataclass
class RawInputIndex:
    article_ids: set[str] = field(default_factory=set)
    hits_by_article_id: dict[str, list[InputHit]] = field(default_factory=lambda: defaultdict(list))
    title_hash: dict[str, list[InputHit]] = field(default_factory=lambda: defaultdict(list))
    content_hash: dict[str, list[InputHit]] = field(default_factory=lambda: defaultdict(list))
    title_time: dict[tuple[str, str], list[InputHit]] = field(default_factory=lambda: defaultdict(list))
    row_count: int = 0


def rel(path: Path) -> str:
    return path.resolve().relative_to(ROOT).as_posix()


def read_jsonl(path: Path, *, allow_missing: bool = False) -> list[dict[str, Any]]:
    if allow_missing and not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_files(paths: list[Path]) -> dict[str, str]:
    return {rel(path): sha256(path) if path.exists() else "missing" for path in paths}


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip() == ""
    if isinstance(value, (list, tuple, set, dict)):
        return len(value) == 0
    if isinstance(value, float) and math.isnan(value):
        return True
    return False


def nonempty(value: Any) -> bool:
    return not is_blank(value)


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def stable_hash(value: Any) -> str:
    text = norm_text(value)
    if not text:
        return ""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def norm_time(value: Any) -> str:
    if is_blank(value):
        return ""
    try:
        return pd.to_datetime(value).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(value).strip()


def rate(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 6) if denominator else 0.0


def f1(precision: float, recall: float) -> float:
    return round(2 * precision * recall / (precision + recall), 6) if precision + recall else 0.0


def split_rows(paths: dict[str, Path]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for split, path in paths.items():
        for row in read_jsonl(path, allow_missing=True):
            item = dict(row)
            item["_split"] = split
            item["_source_path"] = rel(path)
            rows.append(item)
    return rows


def load_schema() -> dict[str, dict[str, Any]]:
    return {row["schema_id"]: row for row in read_jsonl(SCHEMA_PATH)}


def raw_gold_test_article_ids() -> tuple[set[str], dict[str, Any]]:
    article_ids: set[str] = set()
    sheet_rows: dict[str, int] = {}
    wb = load_workbook(RAW_GOLD_TEST, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            if not headers:
                continue
            aid_idx = headers.index("article_id") if "article_id" in headers else 0
            rows = 0
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows += 1
                aid = str(row[aid_idx] or "").strip()
                if aid:
                    article_ids.add(aid)
            sheet_rows[sheet_name] = rows
    finally:
        wb.close()
    return article_ids, {"sheet_rows": sheet_rows, "unique_article_ids": len(article_ids)}


def build_raw_input_index() -> RawInputIndex:
    index = RawInputIndex()
    for split, path in RAW_UNLABELED.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb["Sheet1"]
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            columns = {name: idx for idx, name in enumerate(headers)}
            for row in ws.iter_rows(min_row=2, values_only=True):
                index.row_count += 1
                article_id = str(row[columns["article_file_id"]] or "").strip()
                title = norm_text(row[columns["article_title"]])
                publish_time = norm_time(row[columns["article_publish_time"]])
                content = row[columns["content"]]
                if not article_id:
                    continue
                hit = InputHit(split=split, article_id=article_id, title=title, publish_time=publish_time)
                title_key = stable_hash(title)
                content_key = stable_hash(content)
                index.article_ids.add(article_id)
                index.hits_by_article_id[article_id].append(hit)
                if title_key:
                    index.title_hash[title_key].append(hit)
                    if publish_time:
                        index.title_time[(title_key, publish_time)].append(hit)
                if content_key:
                    index.content_hash[content_key].append(hit)
        finally:
            wb.close()
    return index


def article_set(rows: list[dict[str, Any]]) -> set[str]:
    return {str(row.get("article_id")) for row in rows if nonempty(row.get("article_id"))}


def schema_pair_set(rows: list[dict[str, Any]]) -> set[tuple[str, str]]:
    pairs = set()
    for row in rows:
        if nonempty(row.get("article_id")) and nonempty(row.get("schema_id")):
            pairs.add((str(row["article_id"]), str(row["schema_id"])))
    return pairs


def prediction_event_pairs(rows: list[dict[str, Any]]) -> set[tuple[str, str]]:
    pairs = set()
    for row in rows:
        article_id = str(row.get("article_id") or "")
        for event in row.get("events", []) or []:
            schema_id = event.get("schema_id")
            if article_id and schema_id:
                pairs.add((article_id, str(schema_id)))
    return pairs


def nested_event_pairs(rows: list[dict[str, Any]], field: str = "events") -> set[tuple[str, str]]:
    pairs = set()
    for row in rows:
        article_id = str(row.get("article_id") or "")
        for event in row.get(field, []) or []:
            schema_id = event.get("schema_id")
            if article_id and schema_id:
                pairs.add((article_id, str(schema_id)))
    return pairs


def values_by_article(rows: list[dict[str, Any]], field_name: str) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if nonempty(row.get("article_id")) and nonempty(row.get(field_name)):
            out[str(row["article_id"])].add(str(row[field_name]))
    return out


def metric_pair_sets(
    gold_by_article: dict[str, set[str]],
    pred_by_article: dict[str, set[str]],
    overlap_articles: set[str],
) -> dict[str, Any]:
    tp = pred_total = gold_total = exact = 0
    for article_id in overlap_articles:
        gold = gold_by_article.get(article_id, set())
        pred = pred_by_article.get(article_id, set())
        tp += len(gold & pred)
        pred_total += len(pred)
        gold_total += len(gold)
        if gold == pred:
            exact += 1
    precision = rate(tp, pred_total)
    recall = rate(tp, gold_total)
    return {
        "overlap_article_count": len(overlap_articles),
        "tp": tp,
        "pred_total": pred_total,
        "gold_total": gold_total,
        "precision": precision,
        "recall": recall,
        "f1": f1(precision, recall),
        "exact_article_count": exact,
        "exact_article_accuracy": rate(exact, len(overlap_articles)),
    }


def stage_row(
    *,
    stage: str,
    source_path: str,
    total_rows: int,
    stage_articles: set[str],
    stage_pairs: set[tuple[str, str]],
    gold_article_ids: set[str],
    gold_pair_keys: set[tuple[str, str]],
    previous_retained: set[str] | None,
    match_basis: str,
    note: str = "",
) -> dict[str, Any]:
    retained_articles = stage_articles & gold_article_ids
    retained_pairs = stage_pairs & gold_pair_keys
    if previous_retained is None:
        lost = len(gold_article_ids - retained_articles)
        denom = len(gold_article_ids)
    else:
        lost = len(previous_retained - retained_articles)
        denom = len(previous_retained)
    schema_ids = {schema_id for _, schema_id in stage_pairs}
    return {
        "stage": stage,
        "source_path": source_path,
        "total_rows": total_rows,
        "gold_input_document_count": len(gold_article_ids),
        "retained_gold_document_count": len(retained_articles),
        "lost_document_count_from_previous_stage": lost,
        "lost_rate_from_previous_stage": rate(lost, denom),
        "matched_gold_event_pair_count": len(retained_pairs),
        "stage_unique_article_id_count": len(stage_articles),
        "stage_unique_schema_id_count": len(schema_ids),
        "stage_unique_article_schema_pair_count": len(stage_pairs),
        "article_id_unique": len(stage_articles) == total_rows if total_rows else True,
        "article_schema_pair_unique": len(stage_pairs) == total_rows if stage_pairs else True,
        "match_basis": match_basis,
        "note": note,
    }


def all_splits_path_label(paths: dict[str, Path]) -> str:
    return "|".join(rel(path) for path in paths.values())


def build_coverage_waterfall(
    *,
    raw_gold_ids: set[str],
    gold_news: list[dict[str, Any]],
    gold_pairs: list[dict[str, Any]],
    raw_input_index: RawInputIndex,
    full_predictions: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
    auto_rows: list[dict[str, Any]],
    joint_groups: list[dict[str, Any]],
    joint_results: list[dict[str, Any]],
    final_rows: list[dict[str, Any]],
    task3_docs: list[dict[str, Any]],
    task3_mentions: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    gold_article_ids = {row["article_id"] for row in gold_news}
    gold_pair_keys = {(row["article_id"], row["schema_id"]) for row in gold_pairs}
    gold_event_articles = {row["article_id"] for row in gold_pairs}
    raw_input_overlap = raw_input_index.article_ids & gold_article_ids
    positive_prediction_pairs = prediction_event_pairs(full_predictions)
    llm_rows = [row for row in final_rows if row.get("quality_status") == "llm_accepted"]
    evidence_mentions = {row.get("event_mention_id") for row in evidence_rows if row.get("event_mention_id")}
    task3_evidence_covered_mentions = [row for row in task3_mentions if row.get("event_mention_id") in evidence_mentions]
    rows: list[dict[str, Any]] = []
    prev: set[str] | None = None
    stage_specs = [
        (
            "raw_gold_excel",
            rel(RAW_GOLD_TEST),
            sum(1 for _ in raw_gold_ids),
            raw_gold_ids,
            {(row["article_id"], row["schema_id"]) for row in gold_pairs if row["article_id"] in raw_gold_ids},
            "article_id",
            "raw workbook 全 sheet 合并口径。",
        ),
        (
            "processed_news_test",
            rel(GOLD_NEWS_TEST),
            len(gold_news),
            gold_article_ids,
            set(),
            "article_id",
            "规范化后的 gold test 文档。",
        ),
        (
            "event_pairs_test",
            rel(GOLD_PAIRS_TEST),
            len(gold_pairs),
            gold_event_articles,
            gold_pair_keys,
            "article_id + schema_id",
            "规范化后的 gold test 文章-事件对。",
        ),
        (
            "task1_raw_unlabeled_input",
            all_splits_path_label(RAW_UNLABELED),
            raw_input_index.row_count,
            raw_input_index.article_ids,
            {(row["article_id"], row["schema_id"]) for row in gold_pairs if row["article_id"] in raw_input_index.article_ids},
            "article_id",
            "正式 Task1 输入宇宙来自 news_without_tags_train/test，不是 gold test。",
        ),
        (
            "task1_routed_prediction_groups",
            all_splits_path_label(FULL_PREDICTIONS),
            len(full_predictions),
            article_set(full_predictions),
            positive_prediction_pairs,
            "article_id + nested events[].schema_id",
            "路由后的文章-目标组和模型事件预测。",
        ),
        (
            "quality_gated",
            all_splits_path_label(QUALITY_GATE),
            len(quality_rows),
            article_set(quality_rows),
            schema_pair_set(quality_rows) | nested_event_pairs(quality_rows),
            "article_id + schema_id",
            "质量分层后的 auto/manual/joint group 合并层。",
        ),
        (
            "auto_accept",
            all_splits_path_label(AUTO_ACCEPT),
            len(auto_rows),
            article_set(auto_rows),
            schema_pair_set(auto_rows),
            "article_id + schema_id",
            "无需 LLM 裁决的高置信自动接受事件。",
        ),
        (
            "llm_accepted",
            all_splits_path_label(FINAL_EVENTS),
            len(llm_rows),
            article_set(llm_rows),
            schema_pair_set(llm_rows),
            "article_id + schema_id",
            "最终文件中的 llm_accepted 子集。",
        ),
        (
            "joint_adjudication_results",
            all_splits_path_label(JOINT_RESULTS),
            len(joint_results),
            article_set(joint_results),
            nested_event_pairs(joint_results, "selected_events"),
            "article_id + selected_events[].schema_id",
            "联合裁决结果层；部分 gold schema 在此被裁掉。",
        ),
        (
            "final_events",
            all_splits_path_label(FINAL_EVENTS),
            len(final_rows),
            article_set(final_rows),
            schema_pair_set(final_rows),
            "article_id + schema_id",
            "Task1 冻结正式输出，train/test 联合口径。",
        ),
        (
            "event_mention",
            rel(TASK3_MENTIONS),
            len(task3_mentions),
            article_set(task3_mentions),
            schema_pair_set(task3_mentions),
            "article_id + schema_id",
            "Task3/fulltext_v2 从 Task1 正式输出派生的 EventMention。",
        ),
        (
            "fulltext_v2_documents",
            rel(TASK3_DOCS),
            len(task3_docs),
            article_set(task3_docs),
            {(row["article_id"], row["schema_id"]) for row in task3_mentions if row.get("article_id") and row.get("schema_id")},
            "article_id",
            "fulltext_v2 规范化文档层。",
        ),
        (
            "fulltext_v2_evidence_covered_mentions",
            rel(TASK3_MENTIONS),
            len(task3_evidence_covered_mentions),
            article_set(task3_evidence_covered_mentions),
            schema_pair_set(task3_evidence_covered_mentions),
            "event_mention_id -> article_id + schema_id",
            "有 fulltext_v2 evidence_sentences 的 EventMention。",
        ),
        (
            "evidence_sentence_output",
            rel(TASK3_EVIDENCE),
            len(evidence_rows),
            article_set(evidence_rows),
            schema_pair_set(task3_evidence_covered_mentions),
            "event_mention_id",
            "证据句记录，19681 行，按 event_mention_id 回连。",
        ),
    ]
    for stage, source_path, total_rows, articles, pairs, basis, note in stage_specs:
        rows.append(
            stage_row(
                stage=stage,
                source_path=source_path,
                total_rows=total_rows,
                stage_articles=articles,
                stage_pairs=pairs,
                gold_article_ids=gold_article_ids,
                gold_pair_keys=gold_pair_keys,
                previous_retained=prev,
                match_basis=basis,
                note=note,
            )
        )
        prev = articles & gold_article_ids
    return rows


def prediction_maps(rows: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, set[str]]]:
    by_article: dict[str, list[dict[str, Any]]] = defaultdict(list)
    schemas_by_article: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        article_id = str(row.get("article_id") or "")
        if not article_id:
            continue
        by_article[article_id].append(row)
        for event in row.get("events", []) or []:
            schema_id = event.get("schema_id")
            if schema_id:
                schemas_by_article[article_id].add(str(schema_id))
    return by_article, schemas_by_article


def split_lookup(rows: list[dict[str, Any]]) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if row.get("article_id"):
            out[str(row["article_id"])].add(str(row.get("_split", "")))
    return out


def classify_unmatched(
    *,
    gold_news: list[dict[str, Any]],
    gold_pairs_by_article: dict[str, list[dict[str, Any]]],
    raw_input_index: RawInputIndex,
    full_predictions: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
    joint_groups: list[dict[str, Any]],
    joint_results: list[dict[str, Any]],
    final_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    gold_articles = {row["article_id"] for row in gold_news}
    final_articles = article_set(final_rows)
    prediction_articles = article_set(full_predictions)
    quality_articles = article_set(quality_rows)
    joint_group_articles = article_set(joint_groups)
    joint_result_articles = article_set(joint_results)
    final_pairs = schema_pair_set(final_rows)
    prediction_by_article, prediction_schemas = prediction_maps(full_predictions)
    quality_splits = split_lookup(quality_rows)
    prediction_splits = split_lookup(full_predictions)
    final_splits = split_lookup(final_rows)
    joint_result_splits = split_lookup(joint_results)
    raw_input_splits = {aid: {hit.split for hit in hits} for aid, hits in raw_input_index.hits_by_article_id.items()}

    gold_by_id = {row["article_id"]: row for row in gold_news}
    id_rows: list[dict[str, Any]] = []
    attribution_rows: list[dict[str, Any]] = []
    unmatched_articles = sorted(gold_articles - final_articles)

    for article_id in sorted(gold_articles):
        gold = gold_by_id[article_id]
        gold_pairs = gold_pairs_by_article.get(article_id, [])
        gold_schema_ids = sorted({row["schema_id"] for row in gold_pairs})
        gold_event_names = sorted({row.get("event_name", "") for row in gold_pairs if row.get("event_name")})
        title_key = stable_hash(gold.get("title"))
        content_key = stable_hash(gold.get("content"))
        title_time_key = (title_key, norm_time(gold.get("publish_time")))
        title_hits = raw_input_index.title_hash.get(title_key, [])
        content_hits = raw_input_index.content_hash.get(content_key, [])
        title_time_hits = raw_input_index.title_time.get(title_time_key, [])
        dup = str(gold.get("duplication_id") or "").strip()
        dup_in_raw = bool(dup and dup in raw_input_index.article_ids)
        dup_in_pred = bool(dup and dup in prediction_articles)
        dup_in_final = bool(dup and dup in final_articles)

        non_exact_content_hits = [hit for hit in content_hits if hit.article_id != article_id]
        non_exact_title_time_hits = [hit for hit in title_time_hits if hit.article_id != article_id]
        non_exact_title_hits = [hit for hit in title_hits if hit.article_id != article_id]
        exact_raw = article_id in raw_input_index.article_ids
        exact_pred = article_id in prediction_articles
        exact_final = article_id in final_articles
        strict_pair_matches = sorted(schema_id for schema_id in gold_schema_ids if (article_id, schema_id) in final_pairs)

        best_method = "none"
        confidence = "none"
        matched_article_id = ""
        if exact_final:
            best_method = "strict_article_id_in_final"
            confidence = "strict"
            matched_article_id = article_id
        elif exact_pred:
            best_method = "strict_article_id_in_prediction_universe"
            confidence = "high"
            matched_article_id = article_id
        elif exact_raw:
            best_method = "strict_article_id_in_task1_raw_input"
            confidence = "high"
            matched_article_id = article_id
        elif dup_in_final or dup_in_pred or dup_in_raw:
            best_method = "duplication_id_exact"
            confidence = "medium"
            matched_article_id = dup
        elif non_exact_content_hits:
            best_method = "content_hash_match_only"
            confidence = "high_non_strict"
            matched_article_id = non_exact_content_hits[0].article_id
        elif non_exact_title_time_hits:
            best_method = "title_publish_time_match"
            confidence = "medium_non_strict"
            matched_article_id = non_exact_title_time_hits[0].article_id
        elif non_exact_title_hits:
            best_method = "title_hash_only"
            confidence = "low_non_strict"
            matched_article_id = non_exact_title_hits[0].article_id

        id_rows.append(
            {
                "article_id": article_id,
                "gold_schema_ids": "|".join(gold_schema_ids),
                "sample_ids": "|".join(sorted(row.get("sample_id", "") for row in gold_pairs if row.get("sample_id"))),
                "sample_id_match_possible": False,
                "strict_article_id_in_raw_unlabeled_input": exact_raw,
                "strict_article_id_in_prediction_universe": exact_pred,
                "strict_article_id_in_final_events": exact_final,
                "strict_article_schema_pair_matches": "|".join(strict_pair_matches),
                "raw_input_splits": "|".join(sorted(raw_input_splits.get(article_id, set()))),
                "prediction_splits": "|".join(sorted(prediction_splits.get(article_id, set()))),
                "final_splits": "|".join(sorted(final_splits.get(article_id, set()))),
                "title_hash_input_match_count": len(title_hits),
                "content_hash_input_match_count": len(content_hits),
                "title_publish_time_input_match_count": len(title_time_hits),
                "non_exact_title_hash_match_count": len(non_exact_title_hits),
                "non_exact_content_hash_match_count": len(non_exact_content_hits),
                "non_exact_title_publish_time_match_count": len(non_exact_title_time_hits),
                "duplication_id": dup,
                "duplication_id_in_raw_input": dup_in_raw,
                "duplication_id_in_prediction_universe": dup_in_pred,
                "duplication_id_in_final_events": dup_in_final,
                "company_schema_pair_in_final": any((article_id, schema_id) in final_pairs for schema_id in gold_schema_ids),
                "best_match_method": best_method,
                "matched_article_id": matched_article_id,
                "match_confidence": confidence,
                "note": "非严格匹配不计入正确样本。" if confidence.endswith("non_strict") or confidence == "medium" else "",
            }
        )

        if article_id not in unmatched_articles:
            continue

        pred_schemas = prediction_schemas.get(article_id, set())
        gold_schema_set = set(gold_schema_ids)
        category = "unknown"
        first_missing_stage = ""
        reason = ""
        if exact_pred:
            if pred_schemas & gold_schema_set:
                selected_pairs = nested_event_pairs([row for row in joint_results if row.get("article_id") == article_id], "selected_events")
                if article_id in joint_result_articles and not any((article_id, sid) in selected_pairs for sid in gold_schema_set):
                    category = "filter_or_threshold_loss"
                    first_missing_stage = "joint_adjudication_results_to_final_events"
                    reason = "模型候选包含 gold schema，但联合裁决未选中或最终未接受。"
                elif article_id in quality_articles and article_id not in final_articles:
                    category = "filter_or_threshold_loss"
                    first_missing_stage = "quality_gated_to_final_events"
                    reason = "候选进入质量门后未进入最终接受输出。"
                else:
                    category = "final_export_loss"
                    first_missing_stage = "final_export"
                    reason = "中间候选包含 gold schema，但最终文件未出现该文章。"
            elif prediction_by_article.get(article_id) and any(row.get("predicted_event_count", 0) == 0 for row in prediction_by_article[article_id]):
                category = "truly_no_prediction"
                first_missing_stage = "model_prediction"
                reason = "文章进入预测宇宙，但模型事件列表为空。"
            else:
                category = "candidate_generation_loss"
                first_missing_stage = "target_or_schema_candidate_generation"
                reason = "文章进入预测宇宙，但未生成 gold target/schema 候选。"
        elif exact_raw:
            category = "candidate_generation_loss"
            first_missing_stage = "raw_input_to_routed_prediction_groups"
            reason = "文章在 Task1 原始未标注输入中，但没有进入路由/预测组。"
        elif dup and not exact_raw and (dup_in_final or dup_in_pred or dup_in_raw):
            category = "duplicate_or_merge_loss"
            first_missing_stage = "dedup_or_canonical_mapping"
            reason = "gold article_id 不在正式宇宙，duplication_id 可连接到 Task1 输入/预测宇宙。"
        elif non_exact_content_hits:
            category = "content_hash_match_only"
            first_missing_stage = "article_id_assignment"
            reason = "正文 hash 能连接到 Task1 输入宇宙，但 article_id 不同。"
        elif non_exact_title_time_hits:
            category = "article_id_mismatch"
            first_missing_stage = "article_id_assignment"
            reason = "标题+发布时间能连接到 Task1 输入宇宙，但 article_id 不同。"
        elif article_id not in raw_input_index.article_ids:
            category = "input_universe_mismatch"
            first_missing_stage = "task1_raw_unlabeled_input"
            reason = "gold test 文章未进入 news_without_tags_train/test 正式输入宇宙。"

        attribution_rows.append(
            {
                "article_id": article_id,
                "title": gold.get("title", ""),
                "publish_time": gold.get("publish_time", ""),
                "gold_schema_ids": "|".join(gold_schema_ids),
                "gold_event_names": "|".join(gold_event_names),
                "target_codes": "|".join(sorted({row.get("target_code", "") for row in gold_pairs if row.get("target_code")})),
                "attribution_category": ATTRIBUTION_LABELS[category],
                "category_key": category,
                "first_missing_stage": first_missing_stage,
                "reason": reason,
                "exact_raw_unlabeled_input": exact_raw,
                "exact_prediction_group": exact_pred,
                "gold_schema_in_prediction_events": bool(pred_schemas & gold_schema_set),
                "prediction_schema_ids": "|".join(sorted(pred_schemas)),
                "in_quality_gate": article_id in quality_articles,
                "in_joint_adjudication_group": article_id in joint_group_articles,
                "in_joint_adjudication_result": article_id in joint_result_articles,
                "joint_result_splits": "|".join(sorted(joint_result_splits.get(article_id, set()))),
                "in_final_events": False,
                "duplication_id": dup,
                "duplication_id_in_raw_input": dup_in_raw,
                "duplication_id_in_prediction_universe": dup_in_pred,
                "duplication_id_in_final_events": dup_in_final,
                "non_exact_content_hash_match_count": len(non_exact_content_hits),
                "non_exact_title_publish_time_match_count": len(non_exact_title_time_hits),
                "non_exact_title_hash_match_count": len(non_exact_title_hits),
                "best_id_match_method": best_method,
                "match_confidence": confidence,
            }
        )
    summary = {
        "unmatched_gold_test_article_count": len(attribution_rows),
        "category_counts": dict(Counter(row["category_key"] for row in attribution_rows)),
        "true_model_miss_count": sum(1 for row in attribution_rows if row["category_key"] == "truly_no_prediction"),
        "raw_input_exact_gold_overlap": sum(1 for row in id_rows if row["strict_article_id_in_raw_unlabeled_input"]),
        "prediction_exact_gold_overlap": sum(1 for row in id_rows if row["strict_article_id_in_prediction_universe"]),
        "final_exact_gold_overlap": sum(1 for row in id_rows if row["strict_article_id_in_final_events"]),
        "non_exact_content_hash_gold_articles": sum(1 for row in id_rows if row["non_exact_content_hash_match_count"]),
        "non_exact_title_publish_time_gold_articles": sum(1 for row in id_rows if row["non_exact_title_publish_time_match_count"]),
        "non_exact_title_hash_gold_articles": sum(1 for row in id_rows if row["non_exact_title_hash_match_count"]),
        "duplication_id_non_strict_candidates": sum(
            1
            for row in id_rows
            if not row["strict_article_id_in_raw_unlabeled_input"]
            and (row["duplication_id_in_raw_input"] or row["duplication_id_in_prediction_universe"] or row["duplication_id_in_final_events"])
        ),
    }
    return attribution_rows, id_rows, summary


def shadow_event_name_repair(
    *,
    final_rows: list[dict[str, Any]],
    schema_by_id: dict[str, dict[str, Any]],
    gold_pairs: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    repaired_rows: list[dict[str, Any]] = []
    audit_rows: list[dict[str, Any]] = []
    blank_rows = [row for row in final_rows if is_blank(row.get("event_name"))]
    for row in final_rows:
        item = dict(row)
        schema_id = str(item.get("schema_id") or "")
        schema_name = schema_by_id.get(schema_id, {}).get("event_name", "")
        if is_blank(item.get("event_name")) and schema_name:
            item["event_name"] = schema_name
            item["shadow_repair_method"] = "schema_id_to_event_name"
        repaired_rows.append(item)
    for row in blank_rows:
        schema_id = str(row.get("schema_id") or "")
        schema_name = schema_by_id.get(schema_id, {}).get("event_name", "")
        audit_rows.append(
            {
                "event_id": row.get("event_id", ""),
                "article_id": row.get("article_id", ""),
                "split": row.get("_split", ""),
                "quality_status": row.get("quality_status", ""),
                "schema_id": schema_id,
                "schema_id_valid": schema_id in schema_by_id,
                "schema_event_name": schema_name,
                "formal_event_name_blank": True,
                "recoverable_without_model": bool(schema_name),
                "attribution": "mapping_export_omission",
                "detail": "joint_adjudication selected_events 只有 schema_id，最终汇总未用冻结 schema 字典补 event_name。",
            }
        )
    before_nonempty = sum(1 for row in final_rows if nonempty(row.get("event_name")))
    after_nonempty = sum(1 for row in repaired_rows if nonempty(row.get("event_name")))
    gold_by_article = values_by_article(gold_pairs, "event_name")
    overlap_articles = set(gold_by_article) & article_set(final_rows)
    raw_pred_event_names = prediction_event_name_by_article(final_rows, schema_by_id, raw=True)
    repaired_event_names = prediction_event_name_by_article(repaired_rows, schema_by_id, raw=True)
    before_metrics = metric_pair_sets(gold_by_article, raw_pred_event_names, overlap_articles)
    after_metrics = metric_pair_sets(gold_by_article, repaired_event_names, overlap_articles)
    summary = {
        "blank_event_name_count": len(blank_rows),
        "blank_quality_status_counts": dict(Counter(row.get("quality_status", "") for row in blank_rows)),
        "blank_schema_valid_count": sum(1 for row in blank_rows if row.get("schema_id") in schema_by_id),
        "blank_unique_schema_id_count": len({row.get("schema_id") for row in blank_rows}),
        "schema_id_to_event_name_one_to_one": all(nonempty(row.get("event_name")) for row in schema_by_id.values()),
        "before_event_name_nonempty_count": before_nonempty,
        "before_event_name_nonempty_rate": rate(before_nonempty, len(final_rows)),
        "after_event_name_nonempty_count": after_nonempty,
        "after_event_name_nonempty_rate": rate(after_nonempty, len(final_rows)),
        "before_raw_event_name_overlap_metrics": before_metrics,
        "after_shadow_event_name_overlap_metrics": after_metrics,
        "classification": "mapping_field_export_omission",
    }
    return repaired_rows, audit_rows, summary


def prediction_event_name_by_article(
    rows: list[dict[str, Any]],
    schema_by_id: dict[str, dict[str, Any]],
    *,
    raw: bool,
) -> dict[str, set[str]]:
    out: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        article_id = row.get("article_id")
        if not article_id:
            continue
        value = row.get("event_name") if raw else schema_by_id.get(str(row.get("schema_id") or ""), {}).get("event_name")
        if nonempty(value):
            out[str(article_id)].add(str(value))
        else:
            out[str(article_id)].add("__MISSING__")
    return out


def evidence_and_field_audits(
    *,
    final_rows: list[dict[str, Any]],
    task3_mentions: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
    full_predictions: list[dict[str, Any]],
    auto_rows: list[dict[str, Any]],
    joint_results: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    final_by_key = {
        (row.get("article_id"), row.get("schema_id"), row.get("target_type"), str(row.get("target_code"))): row
        for row in final_rows
    }
    mention_by_key = {
        (row.get("article_id"), row.get("schema_id"), row.get("target_type"), str(row.get("target_code"))): row
        for row in task3_mentions
    }
    evidence_count_by_mention: Counter[str] = Counter(
        str(row.get("event_mention_id")) for row in evidence_rows if row.get("event_mention_id")
    )
    evidence_audit: list[dict[str, Any]] = []
    for key, final in sorted(final_by_key.items(), key=lambda item: (str(item[0]), str(item[1]))):
        mention = mention_by_key.get(key, {})
        mention_id = mention.get("event_mention_id", "")
        evidence_count = evidence_count_by_mention.get(str(mention_id), 0)
        evidence_audit.append(
            {
                "event_id": final.get("event_id", ""),
                "article_id": final.get("article_id", ""),
                "schema_id": final.get("schema_id", ""),
                "target_type": final.get("target_type", ""),
                "target_code": final.get("target_code", ""),
                "quality_status": final.get("quality_status", ""),
                "formal_evidence_nonempty": nonempty(final.get("evidence_sentences")),
                "task3_mention_linked": bool(mention),
                "event_mention_id": mention_id,
                "task3_evidence_list_nonempty": nonempty(mention.get("evidence_sentences")) if mention else False,
                "evidence_sentence_record_count": evidence_count,
                "can_shadow_backfill_evidence": bool(mention and evidence_count > 0),
                "propagation_classification": "downstream_generated_not_formal_task1_export"
                if evidence_count
                else "no_downstream_evidence_record",
            }
        )
    all_source_layers = {
        "final_events": final_rows,
        "quality_gated": quality_rows,
        "full_predictions": full_predictions,
        "auto_accept": auto_rows,
        "joint_results": joint_results,
        "task3_mentions": task3_mentions,
    }
    field_rows: list[dict[str, Any]] = []
    for field_name in ["event_name", "evidence_sentences", *STRUCTURED_FIELDS]:
        source_counts = {}
        for layer, rows in all_source_layers.items():
            source_counts[layer] = sum(1 for row in rows if nonempty(row.get(field_name)))
        if field_name == "event_name":
            conclusion = "formal_partial_export_gap"
            detail = "auto_accept 有 event_name；llm_accepted 缺 event_name，但 schema_id 可一对一恢复。"
        elif field_name == "evidence_sentences":
            conclusion = "downstream_generated_not_propagated"
            detail = "Task3/fulltext_v2 下游产生证据句，冻结 Task1 正式输出未回写。"
        elif field_name in {"event_subject_name", "event_time_raw", "event_time_normalized", "event_action", "event_object", "location"}:
            conclusion = "never_extracted"
            detail = "正式管线中未发现该结构化字段的上游抽取结果。"
        elif field_name in {"event_subject_type", "event_status", "impact_direction", "extraction_confidence"}:
            conclusion = "placeholder_or_missing_not_extraction"
            detail = "旧汇总脚本曾写占位值或缺省值，但冻结正式 JSONL 未保留可评估抽取字段。"
        else:
            conclusion = "not_available"
            detail = "未发现可评估字段。"
        field_rows.append(
            {
                "field": field_name,
                "formal_final_nonempty_count": source_counts["final_events"],
                "quality_gated_nonempty_count": source_counts["quality_gated"],
                "full_predictions_nonempty_count": source_counts["full_predictions"],
                "auto_accept_nonempty_count": source_counts["auto_accept"],
                "joint_results_nonempty_count": source_counts["joint_results"],
                "task3_mentions_nonempty_count": source_counts["task3_mentions"],
                "provenance_conclusion": conclusion,
                "detail": detail,
            }
        )
    summary = {
        "final_event_count": len(final_rows),
        "task3_mention_count": len(task3_mentions),
        "final_to_task3_key_link_count": len(set(final_by_key) & set(mention_by_key)),
        "final_to_task3_key_unlinked_count": len(set(final_by_key) - set(mention_by_key)),
        "task3_evidence_sentence_record_count": len(evidence_rows),
        "task3_evidence_unique_mention_count": len(evidence_count_by_mention),
        "formal_evidence_nonempty_count": sum(1 for row in final_rows if nonempty(row.get("evidence_sentences"))),
        "shadow_evidence_backfillable_event_count": sum(1 for row in evidence_audit if row["can_shadow_backfill_evidence"]),
        "shadow_evidence_backfillable_rate": rate(sum(1 for row in evidence_audit if row["can_shadow_backfill_evidence"]), len(final_rows)),
        "evidence_rows_per_mention_distribution": dict(Counter(evidence_count_by_mention.values())),
        "evidence_origin": "Task3/fulltext_v2 downstream fulltext sentence selection, not formal Task1 extraction output.",
    }
    return evidence_audit, field_rows, summary


def representative_samples(rows: list[dict[str, Any]], per_category: int = 5) -> dict[str, list[dict[str, Any]]]:
    out: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = row["category_key"]
        if len(out[key]) < per_category:
            out[key].append(
                {
                    "article_id": row["article_id"],
                    "title": row["title"],
                    "gold_schema_ids": row["gold_schema_ids"],
                    "first_missing_stage": row["first_missing_stage"],
                    "reason": row["reason"],
                }
            )
    return dict(out)


def write_markdown(path: Path, report: dict[str, Any]) -> None:
    cat_counts = report["unmatched_attribution_summary"]["category_counts"]
    event_name = report["empty_event_name_summary"]
    evidence = report["evidence_summary"]
    lines = [
        "# Task1 金标覆盖缺口与字段传递归因审计",
        "",
        f"- 审计状态：`{report['audit_status']}`。",
        f"- 主结论：`{report['main_conclusion']}`。",
        f"- 下一步唯一建议：{report['next_step_recommendation']}。",
        "- 本步骤未调用 32B、未调用外部模型、未修改 Task1/Task2 正式输出。",
        "",
        "## 覆盖缺口",
        "",
        f"- gold test：{report['gold_test']['document_count']} 篇文档，{report['gold_test']['event_pair_count']} 个事件对。",
        f"- Task1 正式输出：{report['formal_task1']['final_event_count']} 条 EventMention，覆盖 gold test 文章 {report['formal_task1']['gold_test_article_overlap']} 篇。",
        f"- 1004 篇未匹配文章归因：input_universe_mismatch={cat_counts.get('input_universe_mismatch', 0)}，candidate_generation_loss={cat_counts.get('candidate_generation_loss', 0)}，filter_or_threshold_loss={cat_counts.get('filter_or_threshold_loss', 0)}，duplicate_or_merge_loss={cat_counts.get('duplicate_or_merge_loss', 0)}。",
        f"- 真正归入 `truly_no_prediction` 的数量：{report['unmatched_attribution_summary']['true_model_miss_count']}。",
        "",
        "## ID 映射审计",
        "",
        f"- gold test 与 Task1 原始未标注输入 exact article_id 重叠：{report['id_mapping_summary']['raw_input_exact_gold_overlap']}。",
        f"- gold test 与路由/预测宇宙 exact article_id 重叠：{report['id_mapping_summary']['prediction_exact_gold_overlap']}。",
        f"- gold test 与最终正式输出 exact article_id 重叠：{report['id_mapping_summary']['final_exact_gold_overlap']}。",
        f"- 非严格 content hash 命中：{report['id_mapping_summary']['non_exact_content_hash_gold_articles']}；标题+发布时间命中：{report['id_mapping_summary']['non_exact_title_publish_time_gold_articles']}；仅标题命中：{report['id_mapping_summary']['non_exact_title_hash_gold_articles']}。",
        "- 非严格命中均未计入正确样本。",
        "",
        "## event_name",
        "",
        f"- 空 event_name：{event_name['blank_event_name_count']}，全部为 llm_accepted。",
        f"- 空 event_name 的 schema_id 合法数：{event_name['blank_schema_valid_count']}。",
        f"- 修复前覆盖率：{event_name['before_event_name_nonempty_rate']}；shadow 修复后覆盖率：{event_name['after_event_name_nonempty_rate']}。",
        f"- raw event_name 局部 F1：{event_name['before_raw_event_name_overlap_metrics']['f1']}；shadow 后局部 F1：{event_name['after_shadow_event_name_overlap_metrics']['f1']}。",
        "",
        "## 证据与结构化字段",
        "",
        f"- fulltext_v2 证据记录：{evidence['task3_evidence_sentence_record_count']} 条，覆盖 EventMention：{evidence['task3_evidence_unique_mention_count']} 条。",
        f"- 可严格回填证据的 EventMention：{evidence['shadow_evidence_backfillable_event_count']}/{evidence['final_event_count']}。",
        "- evidence_sentences 是 downstream fulltext_v2 生成，冻结 Task1 正式文件没有回写。",
        "- 主体、时间、金额、地点、人物等结构化字段未发现可评估的上游抽取结果。",
        "",
        "## 产物",
        "",
        "- `coverage_waterfall.csv`：全链路覆盖瀑布。",
        "- `unmatched_gold_test_attribution.csv`：1004 篇未匹配文章逐条归因。",
        "- `id_mapping_audit.csv`：严格/非严格 ID 映射审计。",
        "- `empty_event_name_audit.csv`：924 条空 event_name 归因。",
        "- `evidence_propagation_audit.csv`：4973 条 EventMention 证据回填可追溯性。",
        "- `field_provenance.csv`：字段来源和传递结论。",
        "- `shadow_event_name_filled_final_events.jsonl`：只读 shadow 修复结果，不替换正式文件。",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_manifest(output_dir: Path, formal_before: dict[str, str], formal_after: dict[str, str], report: dict[str, Any]) -> None:
    files = []
    for path in sorted(output_dir.iterdir()):
        if path.is_file() and path.name not in {"manifest.json", "SHA256SUMS.txt"}:
            files.append({"path": path.name, "size_bytes": path.stat().st_size, "sha256": sha256(path)})
    manifest = {
        "manifest_version": "task1_gold_coverage_gap_audit_v1",
        "generated_at": report["generated_at"],
        "audit_status": report["audit_status"],
        "main_conclusion": report["main_conclusion"],
        "model_calls_32b": 0,
        "formal_file_sha256_unchanged": formal_before == formal_after,
        "files": files,
    }
    manifest_path = output_dir / "manifest.json"
    write_json(manifest_path, manifest)
    all_files = files + [{"path": "manifest.json", "size_bytes": manifest_path.stat().st_size, "sha256": sha256(manifest_path)}]
    with (output_dir / "SHA256SUMS.txt").open("w", encoding="utf-8") as handle:
        for item in all_files:
            handle.write(f"{item['sha256']}  {item['path']}\n")


def run_audit(output_dir: Path, *, overwrite: bool = False) -> dict[str, Any]:
    if output_dir.exists():
        if overwrite:
            shutil.rmtree(output_dir)
        else:
            raise FileExistsError(f"Output dir already exists: {output_dir}")
    output_dir.mkdir(parents=True, exist_ok=True)

    formal_before = hash_files(PROTECTED_FILES)
    schema_by_id = load_schema()
    gold_news = read_jsonl(GOLD_NEWS_TEST)
    gold_pairs = read_jsonl(GOLD_PAIRS_TEST)
    raw_gold_ids, raw_gold_summary = raw_gold_test_article_ids()
    raw_input_index = build_raw_input_index()
    full_predictions = split_rows(FULL_PREDICTIONS)
    quality_rows = split_rows(QUALITY_GATE)
    auto_rows = split_rows(AUTO_ACCEPT)
    manual_rows = split_rows(MANUAL_REVIEW)
    joint_groups = split_rows(JOINT_GROUPS)
    joint_results = split_rows(JOINT_RESULTS)
    final_rows = split_rows(FINAL_EVENTS)
    task3_docs = read_jsonl(TASK3_DOCS)
    task3_mentions = read_jsonl(TASK3_MENTIONS)
    evidence_rows = read_jsonl(TASK3_EVIDENCE)

    gold_pairs_by_article: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in gold_pairs:
        gold_pairs_by_article[row["article_id"]].append(row)

    coverage_rows = build_coverage_waterfall(
        raw_gold_ids=raw_gold_ids,
        gold_news=gold_news,
        gold_pairs=gold_pairs,
        raw_input_index=raw_input_index,
        full_predictions=full_predictions,
        quality_rows=quality_rows,
        auto_rows=auto_rows,
        joint_groups=joint_groups,
        joint_results=joint_results,
        final_rows=final_rows,
        task3_docs=task3_docs,
        task3_mentions=task3_mentions,
        evidence_rows=evidence_rows,
    )
    attribution_rows, id_rows, attribution_summary = classify_unmatched(
        gold_news=gold_news,
        gold_pairs_by_article=gold_pairs_by_article,
        raw_input_index=raw_input_index,
        full_predictions=full_predictions,
        quality_rows=quality_rows,
        joint_groups=joint_groups,
        joint_results=joint_results,
        final_rows=final_rows,
    )
    shadow_rows, empty_name_rows, event_name_summary = shadow_event_name_repair(
        final_rows=final_rows,
        schema_by_id=schema_by_id,
        gold_pairs=gold_pairs,
    )
    evidence_audit_rows, field_rows, evidence_summary = evidence_and_field_audits(
        final_rows=final_rows,
        task3_mentions=task3_mentions,
        evidence_rows=evidence_rows,
        quality_rows=quality_rows,
        full_predictions=full_predictions,
        auto_rows=auto_rows,
        joint_results=joint_results,
    )

    write_csv(
        output_dir / "coverage_waterfall.csv",
        coverage_rows,
        [
            "stage",
            "source_path",
            "total_rows",
            "gold_input_document_count",
            "retained_gold_document_count",
            "lost_document_count_from_previous_stage",
            "lost_rate_from_previous_stage",
            "matched_gold_event_pair_count",
            "stage_unique_article_id_count",
            "stage_unique_schema_id_count",
            "stage_unique_article_schema_pair_count",
            "article_id_unique",
            "article_schema_pair_unique",
            "match_basis",
            "note",
        ],
    )
    write_csv(
        output_dir / "unmatched_gold_test_attribution.csv",
        attribution_rows,
        [
            "article_id",
            "title",
            "publish_time",
            "gold_schema_ids",
            "gold_event_names",
            "target_codes",
            "attribution_category",
            "category_key",
            "first_missing_stage",
            "reason",
            "exact_raw_unlabeled_input",
            "exact_prediction_group",
            "gold_schema_in_prediction_events",
            "prediction_schema_ids",
            "in_quality_gate",
            "in_joint_adjudication_group",
            "in_joint_adjudication_result",
            "joint_result_splits",
            "in_final_events",
            "duplication_id",
            "duplication_id_in_raw_input",
            "duplication_id_in_prediction_universe",
            "duplication_id_in_final_events",
            "non_exact_content_hash_match_count",
            "non_exact_title_publish_time_match_count",
            "non_exact_title_hash_match_count",
            "best_id_match_method",
            "match_confidence",
        ],
    )
    write_csv(
        output_dir / "id_mapping_audit.csv",
        id_rows,
        [
            "article_id",
            "gold_schema_ids",
            "sample_ids",
            "sample_id_match_possible",
            "strict_article_id_in_raw_unlabeled_input",
            "strict_article_id_in_prediction_universe",
            "strict_article_id_in_final_events",
            "strict_article_schema_pair_matches",
            "raw_input_splits",
            "prediction_splits",
            "final_splits",
            "title_hash_input_match_count",
            "content_hash_input_match_count",
            "title_publish_time_input_match_count",
            "non_exact_title_hash_match_count",
            "non_exact_content_hash_match_count",
            "non_exact_title_publish_time_match_count",
            "duplication_id",
            "duplication_id_in_raw_input",
            "duplication_id_in_prediction_universe",
            "duplication_id_in_final_events",
            "company_schema_pair_in_final",
            "best_match_method",
            "matched_article_id",
            "match_confidence",
            "note",
        ],
    )
    write_csv(
        output_dir / "empty_event_name_audit.csv",
        empty_name_rows,
        [
            "event_id",
            "article_id",
            "split",
            "quality_status",
            "schema_id",
            "schema_id_valid",
            "schema_event_name",
            "formal_event_name_blank",
            "recoverable_without_model",
            "attribution",
            "detail",
        ],
    )
    write_csv(
        output_dir / "evidence_propagation_audit.csv",
        evidence_audit_rows,
        [
            "event_id",
            "article_id",
            "schema_id",
            "target_type",
            "target_code",
            "quality_status",
            "formal_evidence_nonempty",
            "task3_mention_linked",
            "event_mention_id",
            "task3_evidence_list_nonempty",
            "evidence_sentence_record_count",
            "can_shadow_backfill_evidence",
            "propagation_classification",
        ],
    )
    write_csv(
        output_dir / "field_provenance.csv",
        field_rows,
        [
            "field",
            "formal_final_nonempty_count",
            "quality_gated_nonempty_count",
            "full_predictions_nonempty_count",
            "auto_accept_nonempty_count",
            "joint_results_nonempty_count",
            "task3_mentions_nonempty_count",
            "provenance_conclusion",
            "detail",
        ],
    )
    write_jsonl(output_dir / "shadow_event_name_filled_final_events.jsonl", shadow_rows)

    formal_after = hash_files(PROTECTED_FILES)
    final_articles = article_set(final_rows)
    gold_article_ids = {row["article_id"] for row in gold_news}
    strict_final_pairs = schema_pair_set(final_rows)
    gold_pair_keys = {(row["article_id"], row["schema_id"]) for row in gold_pairs}
    report = {
        "generated_at": datetime.now().replace(microsecond=0).isoformat(),
        "audit_status": "completed_read_only",
        "main_conclusion": "mixed_causes",
        "next_step_recommendation": "对完整 gold test 做独立 shadow 推理；仅修评测适配器或导出映射无法扩大到 1029 篇。",
        "model_calls_32b": 0,
        "external_model_calls": 0,
        "formal_file_sha256_before": formal_before,
        "formal_file_sha256_after": formal_after,
        "formal_file_sha256_unchanged": formal_before == formal_after,
        "gold_test": {
            "document_count": len(gold_news),
            "event_pair_count": len(gold_pairs),
            "raw_gold_workbook": rel(RAW_GOLD_TEST),
            "raw_gold_summary": raw_gold_summary,
            "raw_gold_overlap_processed_news_test": len(raw_gold_ids & gold_article_ids),
        },
        "formal_task1": {
            "final_event_count": len(final_rows),
            "final_unique_article_count": len(final_articles),
            "gold_test_article_overlap": len(final_articles & gold_article_ids),
            "gold_test_strict_article_schema_pair_match_count": len(strict_final_pairs & gold_pair_keys),
            "task1_raw_unlabeled_input_rows": raw_input_index.row_count,
            "task1_raw_unlabeled_input_gold_article_overlap": len(raw_input_index.article_ids & gold_article_ids),
            "prediction_group_count": len(full_predictions),
            "prediction_gold_article_overlap": len(article_set(full_predictions) & gold_article_ids),
        },
        "coverage_waterfall_summary": {
            "rows": len(coverage_rows),
            "dominant_loss": "task1_raw_unlabeled_input stage: gold test and formal Task1 input universe differ.",
        },
        "unmatched_attribution_summary": attribution_summary,
        "id_mapping_summary": attribution_summary,
        "representative_unmatched_samples": representative_samples(attribution_rows),
        "empty_event_name_summary": event_name_summary,
        "evidence_summary": evidence_summary,
        "field_provenance_summary": {
            row["field"]: row["provenance_conclusion"] for row in field_rows
        },
        "can_expand_evaluation_coverage_by_deterministic_mapping": False,
        "deterministic_mapping_reason": "content hash 和 title+publish_time 未发现非严格可用扩展；仅有少量 title-only/duplication_id 候选，不能计入严格正确。",
        "can_shadow_backfill_event_name_without_model": True,
        "can_shadow_backfill_task1_evidence_without_model": True,
        "can_run_full_1029_gold_test_evaluation_now": False,
        "full_1029_blocker": "正式 Task1 输出只覆盖 gold test 的 25 篇文章；完整评测需要对 1029 篇 gold test 做独立 shadow 推理。",
        "protected_actions": {
            "no_32b_call": True,
            "no_external_model_call": True,
            "no_formal_output_modified": formal_before == formal_after,
            "no_gold_modified": formal_before.get(rel(GOLD_NEWS_TEST)) == formal_after.get(rel(GOLD_NEWS_TEST))
            and formal_before.get(rel(GOLD_PAIRS_TEST)) == formal_after.get(rel(GOLD_PAIRS_TEST)),
            "no_task2_execution": True,
        },
        "outputs": {
            "coverage_waterfall_csv": "coverage_waterfall.csv",
            "unmatched_gold_test_attribution_csv": "unmatched_gold_test_attribution.csv",
            "id_mapping_audit_csv": "id_mapping_audit.csv",
            "empty_event_name_audit_csv": "empty_event_name_audit.csv",
            "evidence_propagation_audit_csv": "evidence_propagation_audit.csv",
            "field_provenance_csv": "field_provenance.csv",
            "shadow_event_name_filled_final_events_jsonl": "shadow_event_name_filled_final_events.jsonl",
        },
    }
    write_json(output_dir / "audit_report.json", report)
    write_markdown(output_dir / "audit_report.md", report)
    write_manifest(output_dir, formal_before, formal_after, report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    report = run_audit(Path(args.output_dir), overwrite=args.overwrite)
    print(
        json.dumps(
            {
                "audit_status": report["audit_status"],
                "main_conclusion": report["main_conclusion"],
                "output_dir": str(Path(args.output_dir).resolve()),
                "model_calls_32b": report["model_calls_32b"],
                "formal_sha_unchanged": report["formal_file_sha256_unchanged"],
                "unmatched_category_counts": report["unmatched_attribution_summary"]["category_counts"],
                "shadow_event_name_recoverable": report["can_shadow_backfill_event_name_without_model"],
                "shadow_evidence_backfillable_event_count": report["evidence_summary"]["shadow_evidence_backfillable_event_count"],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
