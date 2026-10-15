#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-N strict-fixed ChatGPT500 confirmation and training gate."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import random
import re
import shutil
import subprocess
import sys
import tarfile
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_task1_step5_7f_structure_training import evaluate_predictions  # noqa: E402
from src.event_extraction.task1_supervised_structure_extractor import (  # noqa: E402
    DIR_LABELS,
    ROLE_LABELS,
    SPAN_FIELDS,
    CandidateRankerExtractor,
    build_extractor,
    predict_batch,
    stable_json_sha,
    validate_offsets,
)


INPUT_PACKAGE = ROOT / "incoming/task1_step5_7n/task1_chatgpt_annotation_pilot_500_labeled_strict_fixed_v3.tar.gz"
EXPECTED_INPUT_SHA = "fdce5ca13ab1dd7e242382a0c41d7af10199e2fdd17b55b81b00b9e70bf84854"
OLD_FAILED_JSONL = ROOT / "outputs/task1_step5_7n_chatgpt500_confirmation/extracted/task1_pilot_500_annotated/pilot_500_chatgpt_ai_assisted_silver.jsonl"
OUT_DIR = ROOT / "outputs/task1_step5_7n_chatgpt500_confirmation"
EXTRACT_DIR = OUT_DIR / "extracted_strict_fixed_v3"
PACKAGE_ROOT = "task1_pilot_500_annotated"
NEW_JSONL_REL = f"{PACKAGE_ROOT}/pilot_500_chatgpt_ai_assisted_silver.jsonl"
REQUIRED_PACKAGE_FILES = [
    NEW_JSONL_REL,
    f"{PACKAGE_ROOT}/annotation_manifest.json",
    f"{PACKAGE_ROOT}/validation_report.json",
    f"{PACKAGE_ROOT}/SHA256SUMS.txt",
    f"{PACKAGE_ROOT}/training_eligible_high_precision.jsonl",
    f"{PACKAGE_ROOT}/requires_review_or_reject.jsonl",
    f"{PACKAGE_ROOT}/STRICT_REPAIR_MANIFEST.json",
]

REPORT_CONFIRM_JSON = ROOT / "reports/task1_step5_7n_chatgpt500_confirmation_report.json"
REPORT_CONFIRM_MD = ROOT / "reports/task1_step5_7n_chatgpt500_confirmation_report.md"
REPORT_TRAIN_JSON = ROOT / "reports/task1_step5_7n_structure_training_report.json"
REPORT_TRAIN_MD = ROOT / "reports/task1_step5_7n_structure_training_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7n_structure_training_details.xlsx"

TRAIN_LABELS = ROOT / "data/processed/labeled/news_train.jsonl"
OLD160_TRAIN = ROOT / "outputs/task1_step5_7f_structure_training_v1/train_eventmentions.jsonl"
OLD40_DEV = ROOT / "outputs/task1_step5_7f_structure_training_v1/dev_eventmentions.jsonl"
NEW240_TRAIN = ROOT / "outputs/task1_step5_7j_structure_confirmation/expansion_train_confirmed.jsonl"
FROZEN60_TEST = ROOT / "outputs/task1_step5_7j_structure_confirmation/independent_frozen_test_confirmed.jsonl"
LEGACY80 = ROOT / "outputs/task1_step5_7f_structure_training_v1/legacy80_diagnostic_eventmentions.jsonl"
EXPECTED_FROZEN60_SHA = "76efd5a27057dd177ee8688f4769bfb6e57e8539795cad1727ea29418f7c3579"
STEP57K_REPORT = ROOT / "reports/task1_step5_7k_structure_extractor_report.json"

FORMAL_PROTECTED = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
]
FROZEN_PROTECTED = [
    OLD160_TRAIN,
    OLD40_DEV,
    NEW240_TRAIN,
    FROZEN60_TEST,
    LEGACY80,
    ROOT / "outputs/task1_step5_7k_structure_extractor/frozen_best_model/model.joblib",
    ROOT / "outputs/task1_step5_7k_structure_extractor/frozen_best_model/model.pt",
    ROOT / "outputs/task1_step5_7k_structure_extractor/frozen_model_manifest.json",
    ROOT / "configs/task1_structure_extractor_v1.yaml",
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/evaluation_summary.json",
    ROOT / "configs/task1_candidate_v2.yaml",
    ROOT / "src/routing/task1_candidate_v2.py",
    STEP57K_REPORT,
]

BLOCKING_INTERSECTION_NAMES = {
    "structure_train_400",
    "structure_dev_40",
    "independent_frozen_60_article_ids_only",
    "legacy80_diagnostic",
    "historical_new200",
    "historical_new300",
    "gold1029_article_ids_only_from_blind_candidate_manifest",
    "official_unlabeled_test_100k",
}

CONFIRMATION_NOTE = "用户已逐条复核并确认标注结果无误"
ANNOTATION_SOURCE = "chatgpt_ai_assisted_human_confirmed"
VALID_DIRECTIONS = {"positive", "negative", "neutral"}
VALID_ARGUMENT_ROLES = set(ROLE_LABELS) | {"other"}

DEV_GATES_N = {
    "trigger_relaxed_f1": 0.70,
    "event_span_character_overlap_f1": 0.70,
    "subject_normalized_accuracy": 0.65,
    "event_time_normalized_date_accuracy": 0.50,
    "object_normalized_accuracy": 0.55,
    "direction_accuracy": 0.80,
    "argument_identification_f1": 0.55,
    "argument_role_end_to_end_f1": 0.50,
    "article_exact_match_accuracy": 0.20,
}
FROZEN_GATES_N = {
    "trigger_relaxed_f1": 0.60,
    "event_span_character_overlap_f1": 0.60,
    "subject_normalized_accuracy": 0.55,
    "event_time_normalized_date_accuracy": 0.40,
    "object_normalized_accuracy": 0.45,
    "direction_accuracy": 0.75,
    "argument_role_end_to_end_f1": 0.40,
    "article_exact_match_accuracy": 0.10,
}

SCHEME_CONFIGS = {
    "A": {
        "name": "A_step5_7k_scheme_c_reproduction_macbert_hybrid_span_consistency",
        "candidate_scheme": "C",
        "thresholds": {"subject_company": 0.35, "event_time": 0.38, "object_target": 0.35},
        "description": "Step 5.7-K best scheme C reproduced with the existing gold-free candidate ranker interface.",
    },
    "B": {
        "name": "B_event_conditioned_span_extractor",
        "candidate_scheme": "B",
        "thresholds": {"subject_company": 0.30, "event_time": 0.32, "object_target": 0.30},
        "description": "Event-conditioned field-query candidate ranker over generated spans.",
    },
    "C": {
        "name": "C_two_stage_span_reranker_direction_role_classifier",
        "candidate_scheme": "A",
        "thresholds": {"subject_company": 0.28, "event_time": 0.30, "object_target": 0.28},
        "description": "Sentence/span candidate generation followed by supervised span reranking and direction classification.",
    },
    "D": {
        "name": "D_improved_joint_constraints_weighted_hard_negative",
        "candidate_scheme": "C",
        "thresholds": {"subject_company": 0.20, "event_time": 0.24, "object_target": 0.20},
        "description": "Improved joint candidate ranker variant using lower no-answer thresholds and trigger/span containment post-processing.",
    },
}


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def content_hash(value: Any) -> str:
    return sha256_text(str(value or ""))


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_sha_txt(path: Path, snapshot: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(snapshot.items())), encoding="utf-8")


def sha_snapshot(paths: Iterable[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in sorted(set(paths), key=rel) if path.exists() and path.is_file()}


def process_check() -> list[str]:
    proc = subprocess.run(
        ["bash", "-lc", "ps -eo pid,ppid,stat,etime,cmd | grep -E 'task1_step5_7n|run_task1_step5_7n|train_task1_step5_7n|python.*5_7n' | grep -v grep || true"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )
    return [line for line in proc.stdout.splitlines() if "run_task1_step5_7n_chatgpt500_recovery.py" not in line]


def official_event_names() -> set[str]:
    names: set[str] = set()
    label_count = 0
    for row in read_jsonl(TRAIN_LABELS):
        for label in row.get("labels") or []:
            name = str(label.get("event_name", "")).strip()
            if name:
                names.add(name)
                label_count += 1
    if len(names) != 64 or label_count != 2081:
        raise RuntimeError(f"unexpected official event_name contract: classes={len(names)} labels={label_count}")
    return names


def safe_extract(package: Path, extract_dir: Path) -> dict[str, Any]:
    extract_dir.parent.mkdir(parents=True, exist_ok=True)
    if extract_dir.exists() and any(extract_dir.iterdir()):
        missing = [name for name in REQUIRED_PACKAGE_FILES if not (extract_dir / name).exists()]
        if missing:
            raise RuntimeError(f"existing extract dir is incomplete: {extract_dir}, missing={missing}")
        return {"status": "resumed_existing_extract_dir", "extract_dir": rel(extract_dir), "member_count": None}
    extract_dir.mkdir(parents=True, exist_ok=True)
    with tarfile.open(package, "r:gz") as tf:
        members = tf.getmembers()
        for member in members:
            name = member.name
            posix_parts = Path(name).parts
            if name.startswith("/") or ".." in posix_parts:
                raise RuntimeError(f"unsafe archive path: {name}")
            if member.issym() or member.islnk():
                raise RuntimeError(f"archive link entry is not allowed: {name}")
            target = (extract_dir / name).resolve()
            if not str(target).startswith(str(extract_dir.resolve()) + os.sep) and target != extract_dir.resolve():
                raise RuntimeError(f"archive path escapes extract dir: {name}")
            if target.exists() and not member.isdir():
                raise RuntimeError(f"archive would overwrite existing file: {target}")
        tf.extractall(extract_dir)
    missing = [name for name in REQUIRED_PACKAGE_FILES if not (extract_dir / name).exists()]
    if missing:
        raise RuntimeError(f"required package files missing after extraction: {missing}")
    return {"status": "extracted", "extract_dir": rel(extract_dir), "member_count": len(members)}


def parse_optional_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    try:
        return int(value)
    except Exception:
        return None


def parse_json_field(value: Any) -> tuple[Any, str]:
    if value is None or value == "":
        return [], ""
    try:
        return json.loads(str(value)), ""
    except Exception as exc:
        return None, f"{type(exc).__name__}:{exc}"


def strict_field_check(row: dict[str, Any], prefix: str, content: str) -> tuple[dict[str, Any], list[str]]:
    offset_prefix = {
        "trigger": "trigger",
        "event_span": "event_span",
        "subject_company": "subject",
        "event_time": "event_time",
        "object_target": "object",
    }.get(prefix, prefix)
    text = str(row.get(f"chatgpt_{prefix}", "") or "")
    start = parse_optional_int(row.get(f"chatgpt_{offset_prefix}_start"))
    end = parse_optional_int(row.get(f"chatgpt_{offset_prefix}_end"))
    if not text:
        if prefix == "event_time":
            return {"text": "", "normalized": str(row.get("chatgpt_event_time_normalized", "") or "未明确"), "start": -1, "end": -1}, []
        return {"text": "", "start": -1, "end": -1}, []
    errors: list[str] = []
    if start is None or end is None:
        errors.append(f"{prefix}:offset_missing_or_non_int")
        start, end = -1, -1
    elif not (0 <= start <= end <= len(content)) or content[start:end] != text:
        errors.append(f"{prefix}:content_quote_mismatch")
    span: dict[str, Any] = {"text": text, "start": start, "end": end}
    if prefix == "event_time":
        span["normalized"] = str(row.get("chatgpt_event_time_normalized", "") or "未明确")
    return span, errors


def locate_argument(content: str, text: str, preferred: dict[str, dict[str, Any]]) -> tuple[int, int]:
    value = str(text or "")
    if not value:
        return -1, -1
    for span in preferred.values():
        if span.get("text") == value:
            return int(span.get("start", -1)), int(span.get("end", -1))
    pos = content.find(value)
    if pos < 0:
        return -1, -1
    return pos, pos + len(value)


def validate_chatgpt_rows(rows: list[dict[str, Any]], official64: set[str]) -> dict[str, Any]:
    ids = [str(row.get("candidate_record_id", "") or "") for row in rows]
    event_exists_values = [str(row.get("chatgpt_event_exists", "") or "").strip().lower() for row in rows]
    error_counts: Counter[str] = Counter()
    error_details: list[dict[str, Any]] = []
    field_location_counts: Counter[str] = Counter()
    for idx, row in enumerate(rows, start=1):
        cid = str(row.get("candidate_record_id", "") or "")
        content = str(row.get("content", "") or "")
        row_errors: list[str] = []
        if not cid:
            row_errors.append("candidate_record_id_missing")
        if not row.get("article_id"):
            row_errors.append("article_id_missing")
        if not content:
            row_errors.append("content_missing")
        expected_hash = str(row.get("content_sha256", "") or "")
        if expected_hash and expected_hash != content_hash(content):
            row_errors.append("content_sha256_mismatch")
        event_exists = str(row.get("chatgpt_event_exists", "") or "").strip().lower()
        if event_exists not in {"yes", "no"}:
            row_errors.append("illegal_event_exists")
        event_name = str(row.get("chatgpt_event_name", "") or "").strip()
        if event_exists == "yes":
            if event_name not in official64:
                row_errors.append("illegal_event_name")
        elif event_name and event_name not in official64:
            row_errors.append("illegal_event_name")
        direction = str(row.get("chatgpt_direction", "") or "").strip()
        if event_exists == "yes" and direction not in VALID_DIRECTIONS:
            row_errors.append("illegal_direction")
        spans: dict[str, dict[str, Any]] = {}
        for field, prefix in [
            ("trigger", "trigger"),
            ("event_span", "event_span"),
            ("subject_company", "subject_company"),
            ("event_time", "event_time"),
            ("object_target", "object_target"),
        ]:
            span, field_errors = strict_field_check(row, prefix, content)
            spans[field] = span
            if span.get("text"):
                field_location_counts[f"{field}_strict"] += 1
            row_errors.extend(field_errors)
        if event_exists == "yes" and not spans["trigger"].get("text"):
            row_errors.append("missing_trigger")
        if event_exists == "yes" and not spans["event_span"].get("text"):
            row_errors.append("missing_event_span")
        if spans["trigger"].get("text") and spans["event_span"].get("text"):
            if str(spans["trigger"]["text"]) not in str(spans["event_span"]["text"]):
                row_errors.append("trigger_text_not_in_event_span")
            if not (
                int(spans["event_span"]["start"]) <= int(spans["trigger"]["start"]) < int(spans["trigger"]["end"]) <= int(spans["event_span"]["end"])
            ):
                row_errors.append("trigger_offset_not_in_event_span")
        args, err = parse_json_field(row.get("chatgpt_arguments_json", "[]"))
        roles, role_err = parse_json_field(row.get("chatgpt_argument_roles_json", "[]"))
        if err:
            row_errors.append("arguments_json_parse_failed")
        if role_err:
            row_errors.append("argument_roles_json_parse_failed")
        if not isinstance(args, list):
            row_errors.append("arguments_json_not_list")
            args = []
        if not isinstance(roles, list):
            row_errors.append("argument_roles_json_not_list")
            roles = []
        if len(args) != len(roles):
            row_errors.append("argument_role_count_mismatch")
        for arg_idx, (arg, role) in enumerate(zip(args, roles)):
            role = str(role or "")
            arg_text = str(arg or "")
            if role not in VALID_ARGUMENT_ROLES:
                row_errors.append("illegal_argument_role")
            if arg_text and arg_text not in content:
                row_errors.append(f"argument_{arg_idx}_text_not_in_content")
        for key, value in row.items():
            if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
                row_errors.append(f"{key}:nan_or_inf")
        if row_errors:
            for item in row_errors:
                error_counts[item.split(":")[0]] += 1
            error_details.append({"row_number": idx, "candidate_record_id": cid, "article_id": row.get("article_id", ""), "errors": sorted(set(row_errors))})
    validation = {
        "valid": not error_details and len(rows) == 500 and len(set(ids)) == 500 and all(ids),
        "record_count": len(rows),
        "unique_candidate_record_id_count": len(set(ids)),
        "missing_candidate_record_id": sum(1 for item in ids if not item),
        "event_exists_distribution": dict(Counter(event_exists_values)),
        "missing_trigger": error_counts.get("missing_trigger", 0),
        "trigger_text_not_in_event_span": error_counts.get("trigger_text_not_in_event_span", 0),
        "trigger_offset_not_in_event_span": error_counts.get("trigger_offset_not_in_event_span", 0),
        "content_quote_mismatch": sum(value for key, value in error_counts.items() if key.endswith("content_quote_mismatch")),
        "json_parse_failed": error_counts.get("arguments_json_parse_failed", 0) + error_counts.get("argument_roles_json_parse_failed", 0),
        "illegal_event_name": error_counts.get("illegal_event_name", 0),
        "illegal_direction": error_counts.get("illegal_direction", 0),
        "error_count": len(error_details),
        "error_counts": dict(error_counts),
        "error_details": error_details[:200],
        "field_location_counts": dict(field_location_counts),
    }
    return validation


def repair_manifest_check(manifest: dict[str, Any]) -> dict[str, Any]:
    offset_ids = set(manifest.get("offset_only_repair_ids") or [])
    trigger_ids = set((manifest.get("trigger_and_span_repair") or {}).keys())
    downgraded_ids = set(manifest.get("downgraded_to_no_event") or [])
    all_ids = offset_ids | trigger_ids | downgraded_ids
    expected = {
        "CGPT_PILOT_500_0023",
        "CGPT_PILOT_500_0038",
        "CGPT_PILOT_500_0095",
        "CGPT_PILOT_500_0104",
        "CGPT_PILOT_500_0106",
        "CGPT_PILOT_500_0139",
        "CGPT_PILOT_500_0177",
        "CGPT_PILOT_500_0257",
        "CGPT_PILOT_500_0271",
        "CGPT_PILOT_500_0274",
        "CGPT_PILOT_500_0277",
        "CGPT_PILOT_500_0321",
        "CGPT_PILOT_500_0399",
        "CGPT_PILOT_500_0145",
        "CGPT_PILOT_500_0197",
        "CGPT_PILOT_500_0310",
        "CGPT_PILOT_500_0415",
    }
    checks = {
        "repair_scope": manifest.get("repair_scope"),
        "offset_only_count_13": len(offset_ids) == 13 and manifest.get("offset_only_repair_count") == 13,
        "trigger_and_span_count_1": len(trigger_ids) == 1 and "CGPT_PILOT_500_0145" in trigger_ids,
        "downgraded_count_3": len(downgraded_ids) == 3,
        "all_17_expected_ids": all_ids == expected,
        "strict_validation_manifest_zero": (manifest.get("strict_validation") or {}).get("validation_error_count") == 0,
        "manifest_ids": sorted(all_ids),
    }
    checks["valid"] = all(value for key, value in checks.items() if key.endswith("_13") or key.endswith("_1") or key.endswith("_3") or key.endswith("_ids") or key.endswith("_zero"))
    return checks


def diff_fixed_vs_old(old_rows: list[dict[str, Any]], new_rows: list[dict[str, Any]], manifest: dict[str, Any]) -> dict[str, Any]:
    old_by_id = {str(row.get("candidate_record_id", "")): row for row in old_rows}
    new_by_id = {str(row.get("candidate_record_id", "")): row for row in new_rows}
    offset_ids = set(manifest.get("offset_only_repair_ids") or [])
    trigger_ids = set((manifest.get("trigger_and_span_repair") or {}).keys())
    downgraded_ids = set(manifest.get("downgraded_to_no_event") or [])
    allowed_ids = offset_ids | trigger_ids | downgraded_ids
    provenance_only_fields = {"chatgpt_annotated_at"}
    raw_changed: dict[str, list[str]] = {}
    changed: dict[str, list[str]] = {}
    unexpected: list[dict[str, Any]] = []
    missing: list[str] = []
    for cid, old in old_by_id.items():
        new = new_by_id.get(cid)
        if new is None:
            missing.append(cid)
            continue
        fields = sorted(set(old) | set(new))
        raw_diff_fields = [field for field in fields if old.get(field) != new.get(field)]
        if raw_diff_fields:
            raw_changed[cid] = raw_diff_fields
        diff_fields = [field for field in raw_diff_fields if field not in provenance_only_fields]
        if diff_fields:
            changed[cid] = diff_fields
            if cid not in allowed_ids:
                unexpected.append({"candidate_record_id": cid, "diff_fields": diff_fields})
            elif cid in offset_ids and not set(diff_fields).issubset({"chatgpt_trigger_start", "chatgpt_trigger_end"}):
                unexpected.append({"candidate_record_id": cid, "reason": "offset_only_row_changed_non_offset_fields", "diff_fields": diff_fields})
            elif cid in trigger_ids:
                if new.get("chatgpt_trigger") != "扩能改造" or str(new.get("chatgpt_event_exists", "")).lower() != "yes":
                    unexpected.append({"candidate_record_id": cid, "reason": "trigger_and_span_repair_content_mismatch", "diff_fields": diff_fields})
            elif cid in downgraded_ids:
                if str(new.get("chatgpt_event_exists", "")).lower() != "no":
                    unexpected.append({"candidate_record_id": cid, "reason": "downgraded_row_not_no_event", "diff_fields": diff_fields})
    unchanged_483 = 0
    for cid, old in old_by_id.items():
        new = new_by_id.get(cid, {})
        fields = sorted(set(old) | set(new))
        substantive = [field for field in fields if field not in provenance_only_fields and old.get(field) != new.get(field)]
        if not substantive:
            unchanged_483 += 1
    diff_report = {
        "old_count": len(old_rows),
        "new_count": len(new_rows),
        "raw_changed_count": len(raw_changed),
        "raw_changed_fields_summary": dict(Counter(field for fields in raw_changed.values() for field in fields)),
        "provenance_only_fields_ignored_for_substantive_diff": sorted(provenance_only_fields),
        "changed_count": len(changed),
        "changed_ids": sorted(changed),
        "changed_fields_by_id": changed,
        "manifest_allowed_ids": sorted(allowed_ids),
        "changed_ids_match_manifest": set(changed) == allowed_ids,
        "substantive_unchanged_483": unchanged_483 == 483,
        "unchanged_count_excluding_provenance_fields": unchanged_483,
        "missing_in_new": missing,
        "unexpected_changes": unexpected,
    }
    diff_report["valid"] = (
        len(old_rows) == len(new_rows) == 500
        and not missing
        and not unexpected
        and diff_report["changed_ids_match_manifest"]
        and diff_report["substantive_unchanged_483"]
    )
    return diff_report


def make_span(row: dict[str, Any], prefix: str) -> dict[str, Any]:
    content = str(row.get("content", "") or "")
    span, _ = strict_field_check(row, prefix, content)
    return span


def normalize_chatgpt_row(row: dict[str, Any], confirmed_at: str) -> dict[str, Any]:
    content = str(row.get("content", "") or "")
    event_exists = str(row.get("chatgpt_event_exists", "") or "").strip().lower()
    event_name = str(row.get("chatgpt_event_name", "") or "").strip() if event_exists == "yes" else ""
    trigger = make_span(row, "trigger")
    event_span = make_span(row, "event_span")
    subject = make_span(row, "subject_company")
    event_time = make_span(row, "event_time")
    obj = make_span(row, "object_target")
    args_raw, _ = parse_json_field(row.get("chatgpt_arguments_json", "[]"))
    roles_raw, _ = parse_json_field(row.get("chatgpt_argument_roles_json", "[]"))
    args_raw = args_raw if isinstance(args_raw, list) else []
    roles_raw = roles_raw if isinstance(roles_raw, list) else []
    preferred = {"subject_company": subject, "event_time": event_time, "object_target": obj, "trigger": trigger}
    arguments: list[dict[str, Any]] = []
    for text, role in zip(args_raw, roles_raw):
        role = str(role or "")
        text = str(text or "")
        if role not in ROLE_LABELS or not text:
            continue
        start, end = locate_argument(content, text, preferred)
        if start >= 0 and end >= 0 and content[start:end] == text:
            arguments.append({"text": text, "role": role, "start": start, "end": end})
    direction = str(row.get("chatgpt_direction", "") or "").strip()
    if event_exists == "no" and direction not in DIR_LABELS:
        direction = "neutral"
    out = dict(row)
    out.update(
        {
            "event_mention_id": row.get("candidate_record_id", ""),
            "event_exists": event_exists,
            "event_name": event_name,
            "trigger": trigger if event_exists == "yes" else {"text": "", "start": -1, "end": -1},
            "event_span": event_span if event_exists == "yes" else {"text": "", "start": -1, "end": -1},
            "subject_company": subject if event_exists == "yes" else {"text": "", "start": -1, "end": -1},
            "event_time": event_time if event_exists == "yes" else {"text": "", "normalized": "未明确", "start": -1, "end": -1},
            "object_target": obj if event_exists == "yes" else {"text": "", "start": -1, "end": -1},
            "direction": direction if direction in DIR_LABELS else "neutral",
            "arguments": arguments if event_exists == "yes" else [],
            "argument_roles": arguments if event_exists == "yes" else [],
            "reviewer_confirmed": "yes",
            "annotation_source": ANNOTATION_SOURCE,
            "confirmation_note": CONFIRMATION_NOTE,
            "confirmed_at": confirmed_at,
            "original_quality_grade": row.get("chatgpt_quality_grade", ""),
            "original_training_eligible": row.get("chatgpt_training_eligible", ""),
            "data_source_note": "AI辅助、人工确认数据，不是独立人工金标",
        }
    )
    return out


def title_norm(value: Any) -> str:
    text = re.sub(r"\s+", "", str(value or "").lower())
    return re.sub(r"[，。；：、,.!！?？（）()《》“”\"'`·\\-—_<>【】\\[\\]{}]", "", text)


def article_ids_from_jsonl(path: Path) -> set[str]:
    if not path.exists():
        return set()
    ids: set[str] = set()
    for row in read_jsonl(path):
        aid = str(row.get("article_id", "") or "").strip()
        if aid:
            ids.add(aid)
    return ids


def article_ids_only_by_regex(path: Path) -> set[str]:
    ids: set[str] = set()
    if not path.exists():
        return ids
    pattern = re.compile(r'"article_id"\s*:\s*"([^"]+)"')
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            match = pattern.search(line)
            if match:
                ids.add(match.group(1))
    return ids


def load_xlsx_article_ids(path: Path, candidate_columns: list[str]) -> set[str]:
    if not path.exists():
        return set()
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        header = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = next((header.index(col) for col in candidate_columns if col in header), -1)
        ids: set[str] = set()
        if idx >= 0:
            for cells in ws.iter_rows(min_row=2):
                value = cells[idx].value
                if value:
                    ids.add(str(value).strip())
        wb.close()
        return ids
    except Exception:
        return set()


def leakage_sets() -> dict[str, set[str]]:
    return {
        "structure_train_400": article_ids_from_jsonl(OLD160_TRAIN) | article_ids_from_jsonl(NEW240_TRAIN),
        "structure_dev_40": article_ids_from_jsonl(OLD40_DEV),
        "independent_frozen_60_article_ids_only": article_ids_only_by_regex(FROZEN60_TEST),
        "legacy80_diagnostic": article_ids_from_jsonl(LEGACY80),
        "historical_new200": article_ids_from_jsonl(ROOT / "outputs/task1_step5_7f_structure_training_v1/confirmed_eventmentions_200.jsonl"),
        "historical_new300": article_ids_from_jsonl(NEW240_TRAIN) | article_ids_only_by_regex(FROZEN60_TEST),
        "classifier_train": article_ids_from_jsonl(ROOT / "outputs/task1_step5_6d_classifier_retraining/train_articles.jsonl"),
        "classifier_dev": article_ids_from_jsonl(ROOT / "outputs/task1_step5_6d_classifier_retraining/dev_articles.jsonl"),
        "classifier_holdout": article_ids_from_jsonl(ROOT / "outputs/task1_step5_6d_classifier_retraining/holdout_articles.jsonl"),
        "gold1029_article_ids_only_from_blind_candidate_manifest": article_ids_only_by_regex(ROOT / "outputs/task1_frozen_gold_blind_round1/blind_candidate_v2_top20.jsonl")
        | article_ids_only_by_regex(ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl"),
        "official_unlabeled_test_100k": load_xlsx_article_ids(ROOT / "data/raw/news_without_tags_test.xlsx", ["article_id", "article_file_id", "news_id"]),
        "task2_manual_audit": article_ids_from_jsonl(ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent/primary_samples.jsonl"),
    }


def leakage_check(rows: list[dict[str, Any]]) -> dict[str, Any]:
    ids = {str(row.get("article_id", "") or "") for row in rows if row.get("article_id")}
    contents = {str(row.get("content_sha256", "") or content_hash(row.get("content", ""))) for row in rows}
    titles = {title_norm(row.get("title", "")) for row in rows if row.get("title")}
    sets = leakage_sets()
    intersections = {name: sorted(ids & values) for name, values in sets.items()}
    blocking = {name: values for name, values in intersections.items() if name in BLOCKING_INTERSECTION_NAMES and values}
    duplicate_article_count = len(rows) - len(ids)
    content_counter = Counter(str(row.get("content_sha256", "") or content_hash(row.get("content", ""))) for row in rows)
    duplicated_content_hashes = sorted([key for key, count in content_counter.items() if count > 1])
    return {
        "new_article_id_count": len(ids),
        "new_content_sha256_count": len(contents),
        "new_normalized_title_count": len(titles),
        "duplicate_article_row_count": duplicate_article_count,
        "duplicated_content_hashes": duplicated_content_hashes[:100],
        "intersection_counts": {name: len(values) for name, values in intersections.items()},
        "intersections_sample": {name: values[:20] for name, values in intersections.items() if values},
        "blocking_intersections": {name: values[:50] for name, values in blocking.items()},
        "blocking": bool(blocking),
        "frozen60_labels_or_structure_read": False,
        "gold1029_labels_read": False,
    }


def choose_exact_dev_groups(groups: list[tuple[str, list[dict[str, Any]]]], target: int = 50) -> set[str]:
    random.Random(2026).shuffle(groups)
    buckets: dict[int, tuple[int, int | None]] = {0: (-1, None)}
    for idx, (_, rows) in enumerate(groups):
        size = len(rows)
        for total in sorted(list(buckets.keys()), reverse=True):
            nxt = total + size
            if nxt <= target and nxt not in buckets:
                buckets[nxt] = (total, idx)
        if target in buckets:
            break
    if target not in buckets:
        raise RuntimeError(f"cannot create exact dev split with target rows={target}")
    chosen_idx: set[int] = set()
    total = target
    while total:
        prev, idx = buckets[total]
        if idx is None:
            break
        chosen_idx.add(idx)
        total = prev
    return {groups[idx][0] for idx in chosen_idx}


def split_new500(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("article_id", ""))].append(row)
    groups = list(grouped.items())
    dev_ids = choose_exact_dev_groups(groups, 50)
    train_rows = [row for aid, items in groups for row in items if aid not in dev_ids]
    dev_rows = [row for aid, items in groups for row in items if aid in dev_ids]
    train_rows.sort(key=lambda row: str(row.get("candidate_record_id", "")))
    dev_rows.sort(key=lambda row: str(row.get("candidate_record_id", "")))
    split = {
        "seed": 2026,
        "split_method": "article_id_grouped_exact_450_50",
        "new_train_count": len(train_rows),
        "new_dev_count": len(dev_rows),
        "new_train_article_count": len({row["article_id"] for row in train_rows}),
        "new_dev_article_count": len({row["article_id"] for row in dev_rows}),
        "article_overlap": sorted({row["article_id"] for row in train_rows} & {row["article_id"] for row in dev_rows}),
        "event_name_distribution_train": dict(Counter(row.get("event_name", "") for row in train_rows)),
        "event_name_distribution_dev": dict(Counter(row.get("event_name", "") for row in dev_rows)),
        "direction_distribution_train": dict(Counter(row.get("direction", "") for row in train_rows)),
        "direction_distribution_dev": dict(Counter(row.get("direction", "") for row in dev_rows)),
        "event_exists_distribution_train": dict(Counter(row.get("event_exists", "") for row in train_rows)),
        "event_exists_distribution_dev": dict(Counter(row.get("event_exists", "") for row in dev_rows)),
    }
    if len(train_rows) != 450 or len(dev_rows) != 50 or split["article_overlap"]:
        raise RuntimeError(f"split failed: {split}")
    return train_rows, dev_rows, split


def old400_rows() -> list[dict[str, Any]]:
    rows = read_jsonl(OLD160_TRAIN) + read_jsonl(NEW240_TRAIN)
    if len(rows) != 400:
        raise RuntimeError(f"old400 count mismatch: {len(rows)}")
    return rows


def old40_rows() -> list[dict[str, Any]]:
    rows = read_jsonl(OLD40_DEV)
    if len(rows) != 40:
        raise RuntimeError(f"old40 count mismatch: {len(rows)}")
    return rows


def write_outputs(confirmed: list[dict[str, Any]], new_train: list[dict[str, Any]], new_dev: list[dict[str, Any]], split: dict[str, Any], validation: dict[str, Any], leakage: dict[str, Any]) -> dict[str, str]:
    old_train = old400_rows()
    old_dev = old40_rows()
    merged_train = old_train + new_train
    merged_dev = old_dev + new_dev
    paths = {
        "confirmed": OUT_DIR / "task1_chatgpt500_human_confirmed.jsonl",
        "new_train": OUT_DIR / "new_train_450.jsonl",
        "new_dev": OUT_DIR / "new_dev_50.jsonl",
        "merged_train": OUT_DIR / "merged_structure_train_850.jsonl",
        "merged_dev": OUT_DIR / "merged_structure_dev_90.jsonl",
        "split_manifest": OUT_DIR / "split_manifest.json",
        "validation_report": OUT_DIR / "validation_report.json",
    }
    write_jsonl(paths["confirmed"], confirmed)
    write_jsonl(paths["new_train"], new_train)
    write_jsonl(paths["new_dev"], new_dev)
    write_jsonl(paths["merged_train"], merged_train)
    write_jsonl(paths["merged_dev"], merged_dev)
    split_manifest = {
        **split,
        "merged_structure_train_count": len(merged_train),
        "merged_structure_dev_count": len(merged_dev),
        "independent_frozen_test_count": 60,
        "independent_frozen_test_sha256": sha256_file(FROZEN60_TEST),
        "independent_frozen_test_unchanged": sha256_file(FROZEN60_TEST) == EXPECTED_FROZEN60_SHA,
        "independent_frozen_test_labels_read_for_split": False,
        "legacy80_use": "historical_diagnostic_only",
        "data_statement": "AI辅助、人工确认数据，不是独立人工金标",
    }
    write_json(paths["split_manifest"], split_manifest)
    write_json(paths["validation_report"], {"validation": validation, "leakage": leakage, "split": split_manifest})
    sha_paths = [path for path in paths.values()] + [FROZEN60_TEST]
    sha_report = {rel(path): sha256_file(path) for path in sha_paths if path.exists()}
    write_sha_txt(OUT_DIR / "SHA256SUMS.txt", sha_report)
    return {name: rel(path) for name, path in paths.items()}


def evaluate_scheme(rows: list[dict[str, Any]], predictions: list[dict[str, Any]], scheme_name: str) -> dict[str, Any]:
    metrics = evaluate_predictions(rows, predictions, scheme_name)
    checks = validate_offsets(rows, predictions)
    metrics.update(checks)
    metrics["argument_identification_f1"] = metrics.get("argument_identification_prf", {}).get("f1", 0.0)
    metrics["composite_score"] = round(
        0.15 * metrics.get("trigger_relaxed_f1", 0.0)
        + 0.15 * metrics.get("event_span_character_overlap_f1", 0.0)
        + 0.10 * metrics.get("subject_normalized_accuracy", 0.0)
        + 0.10 * metrics.get("event_time_normalized_date_accuracy", 0.0)
        + 0.10 * metrics.get("object_normalized_accuracy", 0.0)
        + 0.10 * metrics.get("direction_accuracy", 0.0)
        + 0.20 * metrics.get("argument_role_end_to_end_f1", 0.0)
        + 0.10 * metrics.get("article_exact_match_accuracy", 0.0),
        6,
    )
    return metrics


def train_scheme(key: str, train_rows: list[dict[str, Any]], dev_rows: list[dict[str, Any]]) -> tuple[CandidateRankerExtractor, dict[str, Any]]:
    cfg = SCHEME_CONFIGS[key]
    exp_dir = OUT_DIR / f"experiment_{key}"
    exp_dir.mkdir(parents=True, exist_ok=True)
    start = time.time()
    model, train_meta = build_extractor(train_rows, cfg["candidate_scheme"], cfg["thresholds"])
    model_path = exp_dir / "model.joblib"
    model.save(model_path)
    predictions = predict_batch(model, dev_rows)
    pred_path = exp_dir / "dev_predictions.jsonl"
    write_jsonl(pred_path, predictions)
    metrics = evaluate_scheme(dev_rows, predictions, cfg["name"])
    metrics.update(
        {
            "scheme_key": key,
            "scheme_name": cfg["name"],
            "description": cfg["description"],
            "training_backend": "reused_step5_7k_sklearn_tfidf_candidate_ranker",
            "candidate_scheme": cfg["candidate_scheme"],
            "thresholds": cfg["thresholds"],
            "train_seconds": round(time.time() - start, 3),
            "checkpoint": rel(model_path),
            "checkpoint_sha256": sha256_file(model_path),
            "peak_gpu_memory_mb": 0,
            "train_meta": train_meta,
        }
    )
    write_json(exp_dir / "training_log.json", {"config": cfg, "train_meta": train_meta, "metrics": metrics})
    write_json(exp_dir / "dev_metrics.json", metrics)
    return model, metrics


def allowed_inference_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "event_mention_id": row.get("event_mention_id", ""),
            "article_id": row.get("article_id", ""),
            "title": row.get("title", ""),
            "content": row.get("content", ""),
            "event_name": row.get("event_name", ""),
        }
        for row in rows
    ]


def prediction_sha(preds: list[dict[str, Any]]) -> str:
    return stable_json_sha(preds)


def predict_model(model: CandidateRankerExtractor, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return predict_batch(model, allowed_inference_rows(rows))


def negative_leakage_tests(model: CandidateRankerExtractor, dev_rows: list[dict[str, Any]]) -> dict[str, Any]:
    base_preds = predict_model(model, dev_rows)
    base_sha = prediction_sha(base_preds)
    polluted: list[dict[str, Any]] = []
    for idx, row in enumerate(dev_rows):
        item = dict(row)
        item["chatgpt_trigger"] = "FORBIDDEN"
        item["human_trigger"] = "FORBIDDEN"
        item["ai_trigger"] = "FORBIDDEN"
        item["reviewer_note"] = f"note_{len(dev_rows)-idx}"
        item["reference_raw_text"] = "FORBIDDEN"
        item["quality_grade"] = "Z"
        item["training_eligible"] = False
        polluted.append(item)
    variants = {
        "delete_human_fields": [{k: v for k, v in row.items() if not k.startswith("human_")} for row in polluted],
        "delete_chatgpt_fields": [{k: v for k, v in row.items() if not k.startswith("chatgpt_")} for row in polluted],
        "shuffle_reviewer_fields": [dict(row, reviewer_note=polluted[-idx - 1].get("reviewer_note", "")) for idx, row in enumerate(polluted)],
        "delete_reference_fields": [{k: v for k, v in row.items() if not k.startswith("reference")} for row in polluted],
        "gold_file_unavailable_inference": allowed_inference_rows(polluted),
    }
    tests: dict[str, Any] = {}
    for name, rows in variants.items():
        preds = predict_model(model, rows)
        tests[name] = {"passed": prediction_sha(preds) == base_sha, "prediction_sha": prediction_sha(preds)}
    repeat_preds = predict_model(model, dev_rows)
    tests["repeat_determinism"] = {"passed": prediction_sha(repeat_preds) == base_sha, "prediction_sha": prediction_sha(repeat_preds)}
    offset_checks = validate_offsets(dev_rows, base_preds)
    tests["offsets_from_content"] = {"passed": offset_checks["offset_illegal_count"] == 0 and offset_checks["trigger_not_in_event_span_count"] == 0, **offset_checks}
    scanned_files = [
        ROOT / "src/event_extraction/task1_supervised_structure_extractor.py",
        Path(__file__),
    ]
    banned = ["human_", "chatgpt_", "ai_", "reviewer_", "reference_raw_text", "quality_grade", "training_eligible", "confirmation_note"]
    module_hits = {}
    for path in scanned_files:
        text = path.read_text(encoding="utf-8") if path.exists() else ""
        if path.name == Path(__file__).name:
            marker = "def predict_model"
            text = text[text.find(marker) :] if marker in text else ""
        hits = {item: text.count(item) for item in banned if text.count(item)}
        if hits:
            module_hits[rel(path)] = hits
    # The orchestration script contains banned terms in test scaffolding; the extractor module must remain clean.
    extractor_text = (ROOT / "src/event_extraction/task1_supervised_structure_extractor.py").read_text(encoding="utf-8")
    extractor_hits = {item: extractor_text.count(item) for item in banned if extractor_text.count(item)}
    tests["inference_code_forbidden_field_scan"] = {"passed": not extractor_hits, "hits": extractor_hits, "orchestration_scan_for_trace": module_hits}
    all_passed = all(result.get("passed") for result in tests.values())
    return {"base_prediction_sha": base_sha, "tests": tests, "all_passed": all_passed}


def dev_gate_results(metrics: dict[str, Any], leakage_passed: bool) -> dict[str, Any]:
    gates = {f"{key}_ge_{threshold}": metrics.get(key, 0.0) >= threshold for key, threshold in DEV_GATES_N.items()}
    gates["offset_illegal_0"] = metrics.get("offset_illegal_count", 0) == 0
    gates["trigger_span_containment_0"] = metrics.get("trigger_not_in_event_span_count", 0) == 0
    gates["json_parse_failure_0"] = metrics.get("json_parse_failure_count", 0) == 0
    gates["illegal_direction_0"] = metrics.get("illegal_direction_count", 0) == 0
    gates["leakage_tests_all_passed"] = leakage_passed
    return gates


def frozen_gate_results(metrics: dict[str, Any], leakage_passed: bool) -> dict[str, Any]:
    gates = {f"{key}_ge_{threshold}": metrics.get(key, 0.0) >= threshold for key, threshold in FROZEN_GATES_N.items()}
    gates["offset_illegal_0"] = metrics.get("offset_illegal_count", 0) == 0
    gates["json_parse_failure_0"] = metrics.get("json_parse_failure_count", 0) == 0
    gates["illegal_direction_0"] = metrics.get("illegal_direction_count", 0) == 0
    gates["leakage_tests_all_passed"] = leakage_passed
    return gates


def run_training(backup_dir: Path) -> dict[str, Any]:
    train_rows = read_jsonl(OUT_DIR / "merged_structure_train_850.jsonl")
    dev_rows = read_jsonl(OUT_DIR / "merged_structure_dev_90.jsonl")
    if len(train_rows) != 850 or len(dev_rows) != 90:
        raise RuntimeError(f"merged training counts mismatch: {len(train_rows)}/{len(dev_rows)}")
    train_start = time.time()
    dev_metrics: list[dict[str, Any]] = []
    models: dict[str, CandidateRankerExtractor] = {}
    for key in ["A", "B", "C", "D"]:
        model, metrics = train_scheme(key, train_rows, dev_rows)
        models[key] = model
        dev_metrics.append(metrics)
    best = max(dev_metrics, key=lambda item: item.get("composite_score", 0.0))
    best_model = models[best["scheme_key"]]
    leakage = negative_leakage_tests(best_model, dev_rows)
    gates = dev_gate_results(best, leakage["all_passed"])
    dev_passed = all(gates.values())
    frozen = {"labels_read": False, "status": "not_run_due_to_dev_gate_failed", "metrics": {}, "gate_results": {}, "passed": False}
    frozen_model_manifest: dict[str, Any] = {
        "frozen": False,
        "reason": "dev_gate_failed",
        "best_scheme_key": best["scheme_key"],
        "best_checkpoint": best["checkpoint"],
        "best_checkpoint_sha256": best["checkpoint_sha256"],
        "frozen_before_independent_test": False,
    }
    status = "structure_model_dev_gate_failed"
    if dev_passed:
        frozen_dir = OUT_DIR / "frozen_best_model"
        frozen_dir.mkdir(parents=True, exist_ok=True)
        src_model = ROOT / best["checkpoint"]
        dst_model = frozen_dir / "model.joblib"
        shutil.copy2(src_model, dst_model)
        frozen_model_manifest = {
            "frozen": True,
            "frozen_at": now(),
            "best_scheme_key": best["scheme_key"],
            "best_scheme": best["scheme_name"],
            "model": rel(dst_model),
            "model_sha256": sha256_file(dst_model),
            "config": SCHEME_CONFIGS[best["scheme_key"]],
            "frozen_before_independent_test": True,
            "allowed_inference_fields": ["content", "event_name", "title"],
        }
        write_json(OUT_DIR / "frozen_model_manifest.json", frozen_model_manifest)
        if sha256_file(FROZEN60_TEST) != EXPECTED_FROZEN60_SHA:
            raise RuntimeError("frozen60 SHA mismatch before independent evaluation")
        frozen_rows = read_jsonl(FROZEN60_TEST)
        frozen_preds = predict_model(best_model, frozen_rows)
        write_jsonl(OUT_DIR / "independent_frozen60_predictions.jsonl", frozen_preds)
        frozen_metrics = evaluate_scheme(frozen_rows, frozen_preds, f"{best['scheme_name']}_independent_frozen60")
        frozen_gates = frozen_gate_results(frozen_metrics, leakage["all_passed"])
        frozen_passed = all(frozen_gates.values())
        frozen = {"labels_read": True, "status": "completed_once", "metrics": frozen_metrics, "gate_results": frozen_gates, "passed": frozen_passed}
        status = "ready_for_10k_annotation_expansion" if frozen_passed else "structure_training_strategy_revision_required"
    else:
        write_json(OUT_DIR / "frozen_model_manifest.json", frozen_model_manifest)
        (OUT_DIR / "independent_frozen60_predictions.jsonl").write_text("", encoding="utf-8")
    training_report = {
        "generated_at": now(),
        "status": status,
        "training": {
            "executed": True,
            "data_statement": "AI辅助、人工确认数据，不是独立人工金标",
            "train_count": len(train_rows),
            "dev_count": len(dev_rows),
            "independent_frozen_test_count": 60,
            "schemes_run": list(SCHEME_CONFIGS),
            "backend": "reused_step5_7k_sklearn_tfidf_candidate_ranker",
            "gpu_training_seconds": 0,
            "peak_gpu_memory_mb": 0,
            "wall_seconds": round(time.time() - train_start, 3),
        },
        "dev_metrics": dev_metrics,
        "selection": {
            "best_scheme_key": best["scheme_key"],
            "best_scheme": best["scheme_name"],
            "best_checkpoint": best["checkpoint"],
            "best_checkpoint_sha256": best["checkpoint_sha256"],
            "best_composite_score": best["composite_score"],
            "dev_gate_passed": dev_passed,
            "dev_gate_results": gates,
            "selection_data": "merged_structure_dev_90_only",
        },
        "negative_leakage_tests": leakage,
        "independent_frozen_60": frozen,
        "step5_7k_comparison": read_json(STEP57K_REPORT).get("dev_metrics", []) if STEP57K_REPORT.exists() else [],
        "decision": {
            "allow_10k_annotation_expansion": status == "ready_for_10k_annotation_expansion",
            "allow_1029_shadow": False,
        },
        "protection": {
            "independent_frozen_60_labels_read": frozen["labels_read"],
            "independent_frozen_60_labels_read_before_dev_gate": False,
            "gold1029_labels_read": False,
            "news_with_tags_test_read": False,
            "local_32b_calls": 0,
            "external_llm_calls": 0,
            "gpt55_calls": 0,
            "vllm_started": False,
            "formal_task1_task2_modified": False,
        },
        "backup_dir": str(backup_dir),
    }
    return training_report


def write_reports(confirm_report: dict[str, Any], training_report: dict[str, Any]) -> None:
    write_json(REPORT_CONFIRM_JSON, confirm_report)
    lines = [
        "# Task1 Step 5.7-N ChatGPT500 确认报告",
        "",
        f"- status: `{confirm_report['status']}`",
        f"- input_sha_match: {confirm_report['input_package']['sha256_matches']}",
        f"- strict_validation_valid: {confirm_report['strict_validation']['valid']}",
        f"- diff_valid: {confirm_report['repair_diff']['valid']}",
        f"- confirmed_rows: {confirm_report.get('outputs', {}).get('confirmed_count', 0)}",
        f"- new_train/new_dev: {confirm_report.get('split', {}).get('new_train_count', 0)}/{confirm_report.get('split', {}).get('new_dev_count', 0)}",
        f"- merged_train/dev/frozen60: {confirm_report.get('split', {}).get('merged_structure_train_count', 0)}/{confirm_report.get('split', {}).get('merged_structure_dev_count', 0)}/60",
        "",
        "本数据为 AI辅助、人工确认数据，不是独立人工金标。本轮未调用32B、GPT-5.5或外部LLM。",
    ]
    REPORT_CONFIRM_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    write_json(REPORT_TRAIN_JSON, training_report)
    t_lines = [
        "# Task1 Step 5.7-N 结构训练报告",
        "",
        f"- status: `{training_report['status']}`",
        f"- train/dev/frozen60: {training_report['training']['train_count']}/{training_report['training']['dev_count']}/{training_report['training']['independent_frozen_test_count']}",
        f"- best_scheme: `{training_report['selection']['best_scheme']}`",
        f"- dev_gate_passed: {training_report['selection']['dev_gate_passed']}",
        f"- frozen60_labels_read: {training_report['independent_frozen_60']['labels_read']}",
        f"- allow_10k_annotation_expansion: {training_report['decision']['allow_10k_annotation_expansion']}",
        "",
        "未读取1029 gold test标签，未运行1029 Shadow，未修改正式Task1/Task2输出。",
    ]
    REPORT_TRAIN_MD.write_text("\n".join(t_lines) + "\n", encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "strict_validation"
    for key, value in confirm_report["strict_validation"].items():
        if key != "error_details":
            ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws2 = wb.create_sheet("repair_diff")
    for key, value in confirm_report["repair_diff"].items():
        ws2.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws3 = wb.create_sheet("dev_metrics")
    keys = sorted({key for row in training_report.get("dev_metrics", []) for key in row if key != "direction_confusion_matrix"})
    ws3.append(keys)
    for item in training_report.get("dev_metrics", []):
        ws3.append([json.dumps(item.get(key), ensure_ascii=False) if isinstance(item.get(key), (dict, list)) else item.get(key, "") for key in keys])
    ws4 = wb.create_sheet("dev_gate")
    for key, value in training_report["selection"].get("dev_gate_results", {}).items():
        ws4.append([key, value])
    ws5 = wb.create_sheet("leakage_tests")
    ws5.append(["test", "passed", "details"])
    for key, value in training_report["negative_leakage_tests"].get("tests", {}).items():
        ws5.append([key, value.get("passed"), json.dumps(value, ensure_ascii=False)])
    ws6 = wb.create_sheet("frozen60")
    for key, value in training_report["independent_frozen_60"].items():
        ws6.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 44
        sheet.column_dimensions["B"].width = 120
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def write_final_shas(paths: Iterable[Path]) -> None:
    output_paths = [
        OUT_DIR / "task1_chatgpt500_human_confirmed.jsonl",
        OUT_DIR / "new_train_450.jsonl",
        OUT_DIR / "new_dev_50.jsonl",
        OUT_DIR / "merged_structure_train_850.jsonl",
        OUT_DIR / "merged_structure_dev_90.jsonl",
        OUT_DIR / "split_manifest.json",
        OUT_DIR / "validation_report.json",
        OUT_DIR / "frozen_model_manifest.json",
        REPORT_CONFIRM_JSON,
        REPORT_CONFIRM_MD,
        REPORT_TRAIN_JSON,
        REPORT_TRAIN_MD,
        DETAILS_XLSX,
    ]
    output_paths.extend(sorted(path for path in OUT_DIR.glob("experiment_*/*") if path.is_file()))
    snapshot = {rel(path): sha256_file(path) for path in output_paths if path.exists() and path.is_file()}
    write_sha_txt(OUT_DIR / "SHA256SUMS.txt", snapshot)
    write_sha_txt(OUT_DIR / "protected_sha256_after.txt", sha_snapshot(paths))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--backup-dir", type=Path, required=True)
    args = parser.parse_args()
    backup_dir = args.backup_dir
    backup_dir.mkdir(parents=True, exist_ok=False)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if process_check():
        raise RuntimeError(f"same-task process appears active: {process_check()}")
    if not INPUT_PACKAGE.exists():
        raise FileNotFoundError(INPUT_PACKAGE)
    package_sha = sha256_file(INPUT_PACKAGE)
    if package_sha != EXPECTED_INPUT_SHA:
        raise RuntimeError(f"input package SHA mismatch: {package_sha}")
    if sha256_file(FROZEN60_TEST) != EXPECTED_FROZEN60_SHA:
        raise RuntimeError("frozen60 SHA mismatch")
    protected_paths = FORMAL_PROTECTED + FROZEN_PROTECTED
    protected_before = sha_snapshot(protected_paths)
    write_sha_txt(backup_dir / "protected_sha256_before.txt", protected_before)
    write_sha_txt(OUT_DIR / "protected_sha256_before.txt", protected_before)
    extract_report = safe_extract(INPUT_PACKAGE, EXTRACT_DIR)
    new_jsonl = EXTRACT_DIR / NEW_JSONL_REL
    strict_manifest = read_json(EXTRACT_DIR / f"{PACKAGE_ROOT}/STRICT_REPAIR_MANIFEST.json")
    manifest_report = repair_manifest_check(strict_manifest)
    if not manifest_report["valid"]:
        raise RuntimeError(f"STRICT_REPAIR_MANIFEST validation failed: {manifest_report}")
    rows_raw = read_jsonl(new_jsonl)
    official64 = official_event_names()
    strict_validation = validate_chatgpt_rows(rows_raw, official64)
    old_rows = read_jsonl(OLD_FAILED_JSONL)
    repair_diff = diff_fixed_vs_old(old_rows, rows_raw, strict_manifest)
    confirmed_at = now()
    status = "chatgpt500_confirmation_strict_validation_failed"
    confirm_report_base = {
        "generated_at": now(),
        "status": status,
        "input_package": {
            "path": rel(INPUT_PACKAGE),
            "sha256": package_sha,
            "sha256_expected": EXPECTED_INPUT_SHA,
            "sha256_matches": package_sha == EXPECTED_INPUT_SHA,
        },
        "extract": extract_report,
        "strict_repair_manifest": manifest_report,
        "strict_validation": strict_validation,
        "repair_diff": repair_diff,
        "llm_calls": {"local_32b": 0, "external_llm": 0, "gpt55": 0},
        "protection": {
            "gold1029_labels_read": False,
            "news_with_tags_test_read": False,
            "formal_task1_task2_modified": False,
        },
        "backup_dir": str(backup_dir),
    }
    if not strict_validation["valid"] or not repair_diff["valid"]:
        empty_training = {
            "generated_at": now(),
            "status": "not_executed_due_to_chatgpt500_validation_failure",
            "training": {"executed": False, "train_count": 0, "dev_count": 0, "independent_frozen_test_count": 60},
            "selection": {"dev_gate_passed": False, "best_scheme": ""},
            "dev_metrics": [],
            "negative_leakage_tests": {"all_passed": False, "tests": {}},
            "independent_frozen_60": {"labels_read": False, "metrics": {}, "passed": False},
            "decision": {"allow_10k_annotation_expansion": False},
            "protection": {"formal_task1_task2_modified": False, "gold1029_labels_read": False, "local_32b_calls": 0, "external_llm_calls": 0, "gpt55_calls": 0},
            "backup_dir": str(backup_dir),
        }
        write_reports({**confirm_report_base, "outputs": {"confirmed_count": 0}, "split": {}}, empty_training)
        raise RuntimeError("strict validation or fixed-vs-old diff failed; stopped before confirmation/training")
    confirmed = [normalize_chatgpt_row(row, confirmed_at) for row in rows_raw]
    leakage = leakage_check(confirmed)
    if leakage["blocking"]:
        confirm_report_base["status"] = "chatgpt500_confirmation_blocked_by_data_leakage"
        empty_training = {
            "generated_at": now(),
            "status": "not_executed_due_to_data_leakage",
            "training": {"executed": False, "train_count": 0, "dev_count": 0, "independent_frozen_test_count": 60},
            "selection": {"dev_gate_passed": False, "best_scheme": ""},
            "dev_metrics": [],
            "negative_leakage_tests": {"all_passed": False, "tests": {}},
            "independent_frozen_60": {"labels_read": False, "metrics": {}, "passed": False},
            "decision": {"allow_10k_annotation_expansion": False},
            "protection": {"formal_task1_task2_modified": False, "gold1029_labels_read": False, "local_32b_calls": 0, "external_llm_calls": 0, "gpt55_calls": 0},
            "backup_dir": str(backup_dir),
        }
        write_reports({**confirm_report_base, "leakage": leakage, "outputs": {"confirmed_count": 0}, "split": {}}, empty_training)
        raise RuntimeError(f"blocking leakage detected: {leakage['blocking_intersections']}")
    new_train, new_dev, split = split_new500(confirmed)
    output_paths = write_outputs(confirmed, new_train, new_dev, split, strict_validation, leakage)
    split_manifest = read_json(OUT_DIR / "split_manifest.json")
    training_report = run_training(backup_dir)
    protected_after = sha_snapshot(protected_paths)
    protected_match = protected_before == protected_after
    training_report["protection"]["protected_sha_match"] = protected_match
    training_report["protection"]["formal_task1_task2_modified"] = not protected_match
    if not protected_match:
        training_report["status"] = "structure_training_strategy_revision_required"
    confirm_report = {
        **confirm_report_base,
        "status": "chatgpt500_confirmation_passed",
        "leakage": leakage,
        "outputs": {
            **output_paths,
            "confirmed_count": len(confirmed),
            "reviewer_confirmed_yes_count": sum(1 for row in confirmed if row.get("reviewer_confirmed") == "yes"),
            "annotation_source": ANNOTATION_SOURCE,
            "data_statement": "AI辅助、人工确认数据，不是独立人工金标",
        },
        "split": split_manifest,
        "protection": {
            **confirm_report_base["protection"],
            "protected_sha_match": protected_match,
            "formal_task1_task2_modified": not protected_match,
        },
    }
    write_reports(confirm_report, training_report)
    write_final_shas(protected_paths)
    write_json(backup_dir / "run_manifest.json", {"confirmation": confirm_report, "training": training_report})
    write_sha_txt(backup_dir / "protected_sha256_after.txt", protected_after)
    print(
        json.dumps(
            {
                "status": training_report["status"],
                "input_sha_matches": True,
                "strict_validation_errors": strict_validation["error_count"],
                "repair_diff_valid": repair_diff["valid"],
                "new_train_dev": [split_manifest["new_train_count"], split_manifest["new_dev_count"]],
                "merged_train_dev_frozen": [split_manifest["merged_structure_train_count"], split_manifest["merged_structure_dev_count"], 60],
                "dev_gate_passed": training_report["selection"]["dev_gate_passed"],
                "frozen60_labels_read": training_report["independent_frozen_60"]["labels_read"],
                "allow_10k": training_report["decision"]["allow_10k_annotation_expansion"],
                "protected_sha_match": protected_match,
                "backup_dir": str(backup_dir),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
