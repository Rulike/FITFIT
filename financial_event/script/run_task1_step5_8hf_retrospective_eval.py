#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.8-HF retrospective official-test evaluation.

This script is read-only with respect to formal Task1/Task2 outputs.
It builds gold-free inputs from ``news_with_tags_test.xlsx`` using only the
five allowed input columns, runs frozen inference, freezes predictions, and
only then reads gold labels for evaluation.
"""

from __future__ import annotations

import hashlib
import json
import pickle
import random
import re
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import numpy as np
import pandas as pd
import yaml
from openpyxl import load_workbook
from sklearn.metrics import (
    accuracy_score,
    adjusted_rand_score,
    confusion_matrix,
    f1_score,
    precision_recall_fscore_support,
)

from scripts.build_task1_official_format_shadow import official_enums_and_maps
from scripts.run_task1_step5_8e_final_optimization import (
    deterministic_duplication,
    evidence_quote,
    impact_text,
    pairwise_dup_metrics,
    predict_emotion,
    row_text,
)
from scripts.train_task1_step5_8hc_hierarchical_model import (
    head_classes,
    predict_multitask,
    predict_official64_fused,
)
from src.routing.task1_candidate_v2 import Task1CandidateGeneratorV2, read_jsonl


OUT_DIR = ROOT / "outputs/task1_step5_8hf_retrospective_test"
REPORT_JSON = ROOT / "reports/task1_step5_8hf_retrospective_official_test_report.json"
REPORT_MD = ROOT / "reports/task1_step5_8hf_retrospective_official_test_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_8hf_retrospective_official_test_details.xlsx"
CONFUSION_XLSX = ROOT / "reports/task1_step5_8hf_event_name_confusion_matrix.xlsx"
ERRORS_XLSX = ROOT / "reports/task1_step5_8hf_all_tags_errors.xlsx"
IMPACT_AUDIT_XLSX = ROOT / "reports/task1_step5_8hf_impact_human_audit.xlsx"
RUN_STATE = OUT_DIR / "run_state.json"
CHECKPOINT = OUT_DIR / "checkpoint_manifest.json"
SHA_SUMS = OUT_DIR / "SHA256SUMS.txt"
FREEZE_MANIFEST = OUT_DIR / "final_freeze_manifest.json"
PRED_JSONL = OUT_DIR / "official_test_predictions_frozen.jsonl"

RAW_TEST_TAGGED = ROOT / "data/raw/news_with_tags_test.xlsx"
RAW_TEST_UNLABELED = ROOT / "data/raw/news_without_tags_test.xlsx"
RAW_TRAIN = ROOT / "data/raw/news_with_tags_train.xlsx"
OFFICIAL_TRAIN_JSONL = ROOT / "data/processed/labeled/news_train.jsonl"
OFFICIAL_TRAIN_PAIRS = ROOT / "data/processed/labeled/event_pairs_train.jsonl"
TASK1_CANDIDATE_CONFIG = ROOT / "configs/task1_candidate_v2.yaml"
TAXONOMY_JSON = ROOT / "configs/task1_event_taxonomy_hierarchical_v1.json"
HIER_MODEL_PKL = ROOT / "outputs/task1_step5_8hc_hierarchical_event/models/best_hierarchical_model.pkl"
EMOTION_MODEL_PKL = ROOT / "outputs/task1_step5_8f_final_candidate/frozen_emotion_duplication/event_emotion_model.pkl"
HIST_EVAL_JSON = ROOT / "outputs/task1_independent_gold_eval_v1/task1_independent_gold_evaluation.json"
HIST_EVAL_MD = ROOT / "outputs/task1_independent_gold_eval_v1/task1_independent_gold_evaluation.md"
OFFICIAL_COMPANY_SCHEMA_JSON = ROOT / "data/raw/事件类型_标的.json"
OFFICIAL_INDUSTRY_SCHEMA_JSON = ROOT / "data/raw/事件类型_行业.json"
FORMAL_PROTECTED = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "reports/task1_task2_final_manifest.json",
    ROOT / "reports/TASK1_TASK2_FINAL_ACCEPTANCE.md",
]

INPUT_COLS = [
    "article_file_id",
    "article_title",
    "article_source",
    "content",
    "article_publish_time",
]
GOLD_MAIN_COLS = [
    "trading_code",
    "secu_abbr",
    "industry_code",
    "industry_name",
    "event_name",
    "event_emotion",
    "event_impact_analysis",
    "duplication_id",
]
GOLD_DUP_COLS = ["duplication_id"]
LEGAL_STATUSES = {
    "retrospective_official_test_evaluation_complete",
    "frozen_model_missing_required_tag_outputs",
    "evaluation_contract_failure",
    "blocked_by_environment_failure",
}


class ConstantHead:
    def __init__(self, label: str = "") -> None:
        self.classes_ = np.array([label], dtype=object)

    def fit(self, X: Any, y: Any, sample_weight: Any | None = None) -> "ConstantHead":
        return self

    def predict_proba(self, X: Any) -> np.ndarray:
        return np.ones((X.shape[0], 1), dtype=float)


try:
    import __main__ as _main  # type: ignore

    if not hasattr(_main, "ConstantHead"):
        _main.ConstantHead = ConstantHead
except Exception:
    pass


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def read_jsonl_local(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    return rows


def norm_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def norm_code(value: Any, width: int | None = 6) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    if width and text.isdigit():
        text = text.zfill(width)
    return text


def read_headers(path: Path) -> dict[str, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    headers = {}
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        headers[sheet] = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def ensure_dir() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backup = ROOT / "backups" / f"task1_step5_8hf_retrospective_{time.strftime('%Y%m%d_%H%M%S')}"
    backup.mkdir(parents=True, exist_ok=True)
    write_json(
        RUN_STATE,
        {
            "status": "running",
            "stage": "initialized",
            "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "backup_dir": str(backup),
            "model_frozen": True,
            "taxonomy_frozen": True,
            "thresholds_frozen": True,
            "official_test_labels_read": False,
            "llm_calls": 0,
        },
    )
    return backup


def update_state(stage: str, **extra: Any) -> None:
    state = read_json(RUN_STATE) if RUN_STATE.exists() else {}
    state.update({"stage": stage, "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"), **extra})
    write_json(RUN_STATE, state)


def protected_sha(path: Path) -> dict[str, str]:
    paths = [p for p in FORMAL_PROTECTED if p.exists()]
    shas = {str(p): sha256_file(p) for p in sorted(set(paths))}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(f"{v}  {k}" for k, v in sorted(shas.items())) + "\n", encoding="utf-8")
    return shas


def load_prediction_inputs() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for sheet in read_headers(RAW_TEST_TAGGED):
        df = pd.read_excel(RAW_TEST_TAGGED, sheet_name=sheet, usecols=INPUT_COLS, dtype=object).fillna("")
        for idx, record in enumerate(df.to_dict("records")):
            row = {k: ("" if pd.isna(v) else v) for k, v in record.items()}
            row["sheet_name"] = sheet
            row["sheet_row_index"] = idx + 2
            row["article_file_id"] = str(row.get("article_file_id") or "")
            row["article_title"] = str(row.get("article_title") or "")
            row["article_source"] = str(row.get("article_source") or "")
            row["content"] = str(row.get("content") or "")
            row["article_publish_time"] = str(row.get("article_publish_time") or "")
            row["content_sha256"] = hashlib.sha256(norm_text(row["content"]).encode("utf-8")).hexdigest()
            row["title_sha256"] = hashlib.sha256(norm_text(row["article_title"]).encode("utf-8")).hexdigest()
            rows.append(row)
    return rows


def load_gold_rows() -> dict[str, list[dict[str, Any]]]:
    headers = read_headers(RAW_TEST_TAGGED)
    gold: dict[str, list[dict[str, Any]]] = {}
    for sheet in headers:
        if sheet in {"个股新闻", "行业新闻"}:
            usecols = [col for col in INPUT_COLS + GOLD_MAIN_COLS if col in headers[sheet]]
        else:
            usecols = [col for col in INPUT_COLS + GOLD_DUP_COLS if col in headers[sheet]]
        df = pd.read_excel(RAW_TEST_TAGGED, sheet_name=sheet, usecols=usecols, dtype=object).fillna("")
        gold[sheet] = []
        for idx, record in enumerate(df.to_dict("records")):
            row = {k: ("" if pd.isna(v) else v) for k, v in record.items()}
            row["sheet_name"] = sheet
            row["sheet_row_index"] = idx + 2
            row["article_file_id"] = str(row.get("article_file_id") or "")
            row["article_title"] = str(row.get("article_title") or "")
            row["article_source"] = str(row.get("article_source") or "")
            row["content"] = str(row.get("content") or "")
            row["article_publish_time"] = str(row.get("article_publish_time") or "")
            gold[sheet].append(row)
    return gold


def load_train_data() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    train_articles = read_jsonl(OFFICIAL_TRAIN_JSONL)
    train_pairs = read_jsonl(OFFICIAL_TRAIN_PAIRS)
    train_pairs = [
        {
            **row,
            "event_emotion": str(row.get("event_emotion") or row.get("emotion") or ""),
            "event_impact_analysis": str(row.get("event_impact_analysis") or row.get("impact_analysis") or ""),
        }
        for row in train_pairs
    ]
    return train_articles, train_pairs


def load_models() -> tuple[dict[str, Any], Any, dict[str, Any], Task1CandidateGeneratorV2, dict[str, Any], list[dict[str, Any]], dict[str, str]]:
    with HIER_MODEL_PKL.open("rb") as handle:
        hier = pickle.load(handle)
    official_model_path = Path(hier["official_model_path"])
    with official_model_path.open("rb") as handle:
        official_bundle = pickle.load(handle)
    with EMOTION_MODEL_PKL.open("rb") as handle:
        emotion_bundle = pickle.load(handle)
    train_articles, train_pairs = load_train_data()
    with TASK1_CANDIDATE_CONFIG.open("r", encoding="utf-8") as handle:
        candidate_cfg = yaml.safe_load(handle) or {}
    candidate_cfg = candidate_cfg if isinstance(candidate_cfg, dict) else {}
    generator = Task1CandidateGeneratorV2.from_paths(
        ROOT / "data/processed/event_schema.jsonl",
        train_articles,
        train_pairs,
        alias_path=ROOT / "configs/target_aliases.json",
        config=candidate_cfg,
    )
    official_event_names, official_emotions, company_map, industry_map = official_enums_and_maps()
    company44_labels, industry20_labels = load_official_scope_labels()
    official_model = official_bundle["model"] if isinstance(official_bundle, dict) else official_bundle
    emotion_bundle = {
        **emotion_bundle,
        "model": emotion_bundle["model"] if isinstance(emotion_bundle, dict) and "model" in emotion_bundle else emotion_bundle,
    }
    return hier, official_model, emotion_bundle, generator, {
        "official_event_names": sorted(official_event_names),
        "official_emotions": sorted(official_emotions),
        "company44_event_names": company44_labels,
        "industry20_event_names": industry20_labels,
        "company_map": company_map,
        "industry_map": industry_map,
    }, train_pairs, {}


def load_hier_artifact_only() -> dict[str, Any]:
    with HIER_MODEL_PKL.open("rb") as handle:
        return pickle.load(handle)


def load_official_scope_labels() -> tuple[list[str], list[str]]:
    company_rows = json.loads(OFFICIAL_COMPANY_SCHEMA_JSON.read_text(encoding="utf-8"))
    industry_rows = json.loads(OFFICIAL_INDUSTRY_SCHEMA_JSON.read_text(encoding="utf-8"))
    company_labels = sorted({str(event.get("event_name") or "") for row in company_rows for event in row.get("event_schema", []) if event.get("event_name")})
    industry_labels = sorted({str(event.get("event_name") or "") for row in industry_rows for event in row.get("event_schema", []) if event.get("event_name")})
    return company_labels, industry_labels


def build_emotion_prior(train_pairs: list[dict[str, Any]]) -> dict[str, str]:
    grouped: dict[str, Counter[str]] = defaultdict(Counter)
    for row in train_pairs:
        grouped[str(row.get("event_name") or "")][str(row.get("event_emotion") or "")] += 1
    return {event: cnt.most_common(1)[0][0] for event, cnt in grouped.items() if cnt}


def pick_candidate_row(candidates: list[dict[str, Any]], official_event_name: str) -> dict[str, Any]:
    matched = [row for row in candidates if str(row.get("event_name") or "") == official_event_name]
    if matched:
        return max(matched, key=lambda row: float(row.get("candidate_score") or 0.0))
    return candidates[0] if candidates else {}


def predict_targets(
    article: dict[str, Any],
    official_event_name: str,
    generator: Task1CandidateGeneratorV2,
    company_map: dict[str, dict[str, str]],
    industry_map: dict[str, str],
) -> dict[str, Any]:
    cand_rows, _ = generator.generate(
        {
            "article_id": article["article_file_id"],
            "title": article["article_title"],
            "publish_time": article["article_publish_time"],
            "source": article["article_source"],
            "content": article["content"],
        },
        include_all64=True,
    )
    chosen = pick_candidate_row(cand_rows, official_event_name)
    target_type = str(chosen.get("target_type") or "missing_prediction")
    target_code = norm_code(chosen.get("target_code"), 6)
    target_name = str(chosen.get("target_name") or "missing_prediction")
    trading_code = ""
    secu_abbr = ""
    industry_code = ""
    industry_name = ""
    if target_type == "company":
        mapped = company_map.get(target_code) or {}
        trading_code = str(mapped.get("trading_code") or target_code or "missing_prediction")
        secu_abbr = str(mapped.get("secu_abbr") or target_name or "missing_prediction")
        industry_code = str(mapped.get("industry_code") or "")
        industry_name = str(mapped.get("industry_name") or "")
    elif target_type == "industry":
        industry_code = target_code
        industry_name = str(industry_map.get(target_code, target_name or "missing_prediction"))
    else:
        target_type = "missing_prediction"
        trading_code = "missing_prediction"
        secu_abbr = "missing_prediction"
        industry_code = "missing_prediction"
        industry_name = "missing_prediction"
    return {
        "target_type": target_type,
        "target_code": target_code,
        "target_name": target_name,
        "trading_code": trading_code,
        "secu_abbr": secu_abbr,
        "industry_code": industry_code,
        "industry_name": industry_name,
        "candidate_score": float(chosen.get("candidate_score") or 0.0),
        "candidate_rank": int(chosen.get("candidate_rank") or 0),
        "candidate_source": str(chosen.get("candidate_source") or ""),
        "route_count": len(cand_rows),
    }


def row_key(row: dict[str, Any]) -> tuple[str, int]:
    return str(row["sheet_name"]), int(row["sheet_row_index"])


def predict_rows(
    inputs: list[dict[str, Any]],
    hier_artifact: dict[str, Any],
    official_model: Any,
    emotion_bundle: dict[str, Any],
    generator: Task1CandidateGeneratorV2,
    maps: dict[str, Any],
    train_pairs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    taxonomy = hier_artifact["taxonomy"]
    vectorizer = hier_artifact["vectorizer"]
    suite = hier_artifact["suite"]
    texts = [row_text({"article_title": r["article_title"], "article_source": r["article_source"], "content": r["content"], "target_type": "", "secu_abbr": "", "industry_name": ""}) for r in inputs]
    X = vectorizer.transform(texts)
    official_probs = official_model.predict_proba(texts)
    official_pred, official_top, official_conf = predict_official64_fused(
        official_model,
        vectorizer,
        taxonomy,
        inputs,
        hierarchical_suite=suite,
    )
    hier_pred, hier_top, hier_conf = predict_multitask(
        suite,
        taxonomy,
        X,
        include_official_bonus=True,
        official_model=official_model,
        official_probs=official_probs,
    )
    emotion_model = emotion_bundle["model"]
    emotion_kind = emotion_bundle.get("kind", "text")
    emotion_prior = build_emotion_prior(train_pairs)

    rows: list[dict[str, Any]] = []
    for idx, base in enumerate(inputs):
        hier_leaf = hier_pred[idx]
        leaf_meta = taxonomy["leaf_by_name"].get(hier_leaf, {})
        target_pred = predict_targets(base, official_pred[idx], generator, maps["company_map"], maps["industry_map"])
        row = {
            "sheet_name": base["sheet_name"],
            "sheet_row_index": base["sheet_row_index"],
            "article_file_id": base["article_file_id"],
            "article_title": base["article_title"],
            "article_source": base["article_source"],
            "article_publish_time": base["article_publish_time"],
            "content": base["content"],
            "content_sha256": base["content_sha256"],
            "title_sha256": base["title_sha256"],
            "hierarchical_leaf_name": hier_leaf,
            "event_layer": str(leaf_meta.get("event_layer") or "missing_prediction"),
            "event_parent": str(leaf_meta.get("l1_parent_name") or "missing_prediction"),
            "event_group": str(leaf_meta.get("l2_group_name") or "missing_prediction"),
            "official_event_name": str(leaf_meta.get("official_event_name") or ""),
            "hierarchical_confidence": round(float(hier_conf[idx]), 6),
            "confidence": round(float(official_conf[idx]), 6),
            "event_name": official_pred[idx],
            **target_pred,
        }
        if not row["official_event_name"] and row["event_name"] in maps["official_event_names"]:
            row["official_event_name"] = row["event_name"]
        if not row["official_event_name"]:
            row["official_event_name"] = "missing_prediction"
        rows.append(row)

    emotion_rows = [dict(r, event_name=r["event_name"]) for r in rows]
    event_emotion = predict_emotion(emotion_model, emotion_prior, emotion_rows, emotion_kind, train_pairs)
    for row, emo in zip(rows, event_emotion):
        row["event_emotion"] = str(emo or "missing_prediction")
        quote = evidence_quote(row)
        row["event_impact_analysis"] = impact_text(row, row["event_name"], row["event_emotion"], quote)
    dup_pred = deterministic_duplication(rows, [r["event_name"] for r in rows])
    for row, dup in zip(rows, dup_pred):
        row["duplication_id"] = str(dup or "")
    for row in rows:
        status = []
        for key in ["target_type", "event_name", "event_emotion", "event_impact_analysis"]:
            if not str(row.get(key) or "").strip() or str(row.get(key)) == "missing_prediction":
                status.append(f"{key}:missing_prediction")
        row["prediction_status"] = "success" if not status else "partial"
        row["prediction_notes"] = "|".join(status)
    return rows


def load_gold_for_eval() -> dict[tuple[str, int], dict[str, Any]]:
    gold = load_gold_rows()
    out: dict[tuple[str, int], dict[str, Any]] = {}
    for sheet, rows in gold.items():
        for row in rows:
            out[(sheet, int(row["sheet_row_index"]))] = row
    return out


def sheet_gold_type(sheet: str) -> str:
    if "行业" in sheet:
        return "industry"
    return "company"


def normalize_emotion(v: Any) -> str:
    return str(v or "").strip()


def tuple_repr(row: dict[str, Any], is_gold: bool = False) -> str:
    if row.get("target_type") == "company" or row.get("trading_code") or (is_gold and row.get("sheet_name", "").startswith("个股")):
        return f"company|{norm_code(row.get('trading_code') or row.get('target_code'))}|{norm_text(row.get('secu_abbr') or row.get('target_name'))}"
    return f"industry|{norm_code(row.get('industry_code') or row.get('target_code'))}|{norm_text(row.get('industry_name') or row.get('target_name'))}"


def compute_target_metrics(pred_rows: list[dict[str, Any]], gold_rows: list[dict[str, Any]], sheets: set[str]) -> dict[str, Any]:
    pred = []
    gold = []
    target_types = []
    for row in pred_rows:
        key = row_key(row)
        g = gold_rows[key]
        pred.append(tuple_repr(row))
        gold.append(tuple_repr({**g, "sheet_name": row["sheet_name"]}, is_gold=True))
        target_types.append(sheet_gold_type(row["sheet_name"]))
    labels = sorted(set(gold) | set(pred))
    p, r, f, s = precision_recall_fscore_support(gold, pred, labels=labels, zero_division=0)
    cm = confusion_matrix(gold, pred, labels=labels)
    return {
        "support": len(gold),
        "accuracy": round(float(accuracy_score(gold, pred)), 6),
        "micro_precision": round(float(precision_recall_fscore_support(gold, pred, labels=labels, average="micro", zero_division=0)[0]), 6),
        "micro_recall": round(float(precision_recall_fscore_support(gold, pred, labels=labels, average="micro", zero_division=0)[1]), 6),
        "micro_f1": round(float(precision_recall_fscore_support(gold, pred, labels=labels, average="micro", zero_division=0)[2]), 6),
        "macro_f1": round(float(f1_score(gold, pred, labels=labels, average="macro", zero_division=0)), 6),
        "weighted_f1": round(float(f1_score(gold, pred, labels=labels, average="weighted", zero_division=0)), 6),
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
        "confusion_matrix": {labels[i]: {labels[j]: int(cm[i, j]) for j in range(len(labels))} for i in range(len(labels))},
    }


def compute_target_type_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    y_true = []
    y_pred = []
    missing = wrong = extra = 0
    for row in pred_rows:
        key = row_key(row)
        g = gold_rows[key]
        gt = sheet_gold_type(row["sheet_name"])
        pr = str(row.get("target_type") or "missing_prediction")
        y_true.append(gt)
        y_pred.append(pr)
        if not pr or pr == "missing_prediction":
            missing += 1
        elif any(sep in pr for sep in ["|", ",", "、", ";", "/", "；"]):
            extra += 1
        elif pr != gt:
            wrong += 1
    labels = sorted(set(y_true) | set(y_pred))
    p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=labels, zero_division=0)
    cm = confusion_matrix(y_true, y_pred, labels=labels)
    return {
        "support": len(y_true),
        "accuracy": round(float(accuracy_score(y_true, y_pred)), 6),
        "micro_precision": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[0]), 6),
        "micro_recall": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[1]), 6),
        "micro_f1": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[2]), 6),
        "macro_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="macro", zero_division=0)), 6),
        "missing_prediction_count": missing,
        "wrong_prediction_count": wrong,
        "extra_prediction_count": extra,
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
        "confusion_matrix": {labels[i]: {labels[j]: int(cm[i, j]) for j in range(len(labels))} for i in range(len(labels))},
    }


def code_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]], sheet_filter: str) -> dict[str, Any]:
    y_true = []
    y_pred = []
    missing = wrong = over = 0
    for row in pred_rows:
        if row["sheet_name"] != sheet_filter:
            continue
        gold = gold_rows[row_key(row)]
        if sheet_filter == "个股新闻":
            g = norm_code(gold.get("trading_code"))
            p = norm_code(row.get("trading_code"))
        else:
            g = norm_code(gold.get("industry_code"))
            p = norm_code(row.get("industry_code"))
        y_true.append(g)
        y_pred.append(p)
        if not p or p == "missing_prediction":
            missing += 1
        elif any(sep in p for sep in ["|", ",", "、", ";", "/", "；"]):
            over += 1
        elif p != g:
            wrong += 1
    labels = sorted(set(y_true) | set(y_pred))
    p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=labels, zero_division=0)
    cm = confusion_matrix(y_true, y_pred, labels=labels)
    return {
        "support": len(y_true),
        "accuracy": round(float(accuracy_score(y_true, y_pred)), 6),
        "macro_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="macro", zero_division=0)), 6),
        "missing_prediction_count": missing,
        "wrong_prediction_count": wrong,
        "multi_prediction_count": over,
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
        "confusion_matrix": {labels[i]: {labels[j]: int(cm[i, j]) for j in range(len(labels))} for i in range(len(labels))},
    }


def text_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]], field: str, sheet_filter: set[str]) -> dict[str, Any]:
    y_true = []
    y_pred = []
    missing = wrong = extra = 0
    for row in pred_rows:
        if row["sheet_name"] not in sheet_filter:
            continue
        gold = gold_rows[row_key(row)]
        g = norm_text(gold.get(field))
        p = norm_text(row.get(field))
        y_true.append(g)
        y_pred.append(p)
        if not p or p == "missing_prediction":
            missing += 1
        elif any(sep in p for sep in ["|", ",", "、", ";", "/", "；"]):
            extra += 1
        elif p != g:
            wrong += 1
    labels = sorted(set(y_true) | set(y_pred))
    p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=labels, zero_division=0)
    return {
        "support": len(y_true),
        "exact_accuracy": round(float(accuracy_score(y_true, y_pred)), 6),
        "macro_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="macro", zero_division=0)), 6),
        "normalized_accuracy": round(float(accuracy_score(y_true, y_pred)), 6),
        "missing_prediction_count": missing,
        "wrong_prediction_count": wrong,
        "extra_prediction_count": extra,
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
    }


def event_metrics(
    pred_rows: list[dict[str, Any]],
    gold_rows: dict[tuple[str, int], dict[str, Any]],
    official_labels: list[str],
    company44_labels: list[str],
    industry20_labels: list[str],
) -> dict[str, Any]:
    y_true = []
    y_pred = []
    main_rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    for row in main_rows:
        gold = gold_rows[row_key(row)]
        y_true.append(str(gold.get("event_name") or ""))
        y_pred.append(str(row.get("event_name") or ""))
    company_true = [str(gold_rows[row_key(row)].get("event_name") or "") for row in main_rows if row["sheet_name"] == "个股新闻"]
    company_pred = [str(row.get("event_name") or "") for row in main_rows if row["sheet_name"] == "个股新闻"]
    industry_true = [str(gold_rows[row_key(row)].get("event_name") or "") for row in main_rows if row["sheet_name"] == "行业新闻"]
    industry_pred = [str(row.get("event_name") or "") for row in main_rows if row["sheet_name"] == "行业新闻"]
    labels = sorted(set(official_labels) | set(y_true) | set(y_pred))
    p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=labels, zero_division=0)
    cm = confusion_matrix(y_true, y_pred, labels=labels)
    observed = sorted(set(y_true))
    return {
        "support": len(y_true),
        "micro_precision": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[0]), 6),
        "micro_recall": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[1]), 6),
        "micro_f1": round(float(precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)[2]), 6),
        "fixed_64_macro_f1": round(float(f1_score(y_true, y_pred, labels=official_labels, average="macro", zero_division=0)), 6),
        "observed_label_macro_f1": round(float(f1_score(y_true, y_pred, labels=observed, average="macro", zero_division=0)), 6),
        "weighted_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="weighted", zero_division=0)), 6),
        "company44_macro_f1": round(float(f1_score(company_true, company_pred, labels=company44_labels, average="macro", zero_division=0)), 6) if company_true else 0.0,
        "industry20_macro_f1": round(float(f1_score(industry_true, industry_pred, labels=industry20_labels, average="macro", zero_division=0)), 6) if industry_true else 0.0,
        "company_row_micro_f1": round(float(f1_score(company_true, company_pred, average="micro", zero_division=0)), 6) if company_true else 0.0,
        "industry_row_micro_f1": round(float(f1_score(industry_true, industry_pred, average="micro", zero_division=0)), 6) if industry_true else 0.0,
        "article_exact_match": round(float(accuracy_score(y_true, y_pred)), 6),
        "at_least_one_hit_rate": round(float(np.mean([g == p for g, p in zip(y_true, y_pred)])), 6),
        "complete_miss_count": int(sum(1 for g, p in zip(y_true, y_pred) if not p)),
        "multi_prediction_count": int(sum(1 for p in y_pred if any(sep in p for sep in ["|", ",", "、", ";", "/"]))),
        "misclassification_count": int(sum(1 for g, p in zip(y_true, y_pred) if g != p)),
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
        "confusion_matrix": {labels[i]: {labels[j]: int(cm[i, j]) for j in range(len(labels))} for i in range(len(labels))},
    }


def target_event_joint_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    y_true = []
    y_pred = []
    target_correct = []
    event_correct = []
    for row in rows:
        gold = gold_rows[row_key(row)]
        gold_target = tuple_repr({**gold, "sheet_name": row["sheet_name"]}, is_gold=True)
        pred_target = tuple_repr(row)
        gold_event = str(gold.get("event_name") or "")
        pred_event = str(row.get("event_name") or "")
        y_true.append(f"{gold_target}|{gold_event}")
        y_pred.append(f"{pred_target}|{pred_event}")
        target_correct.append(pred_target == gold_target)
        event_correct.append(pred_event == gold_event)
    labels = sorted(set(y_true) | set(y_pred))
    p_micro, r_micro, f_micro, _ = precision_recall_fscore_support(y_true, y_pred, labels=labels, average="micro", zero_division=0)
    _, _, f_macro, _ = precision_recall_fscore_support(y_true, y_pred, labels=labels, average="macro", zero_division=0)
    return {
        "support": len(rows),
        "micro_precision": round(float(p_micro), 6),
        "micro_recall": round(float(r_micro), 6),
        "micro_f1": round(float(f_micro), 6),
        "macro_f1": round(float(f_macro), 6),
        "article_exact_match": round(float(accuracy_score(y_true, y_pred)), 6) if rows else 0.0,
        "target_correct_event_wrong_count": int(sum(t and not e for t, e in zip(target_correct, event_correct))),
        "event_correct_target_wrong_count": int(sum(e and not t for t, e in zip(target_correct, event_correct))),
        "both_wrong_count": int(sum((not t) and (not e) for t, e in zip(target_correct, event_correct))),
        "both_correct_count": int(sum(t and e for t, e in zip(target_correct, event_correct))),
    }


def article_set_metrics(
    pred_rows: list[dict[str, Any]],
    gold_rows: dict[tuple[str, int], dict[str, Any]],
    include_event: bool = False,
) -> dict[str, Any]:
    pred_sets: dict[str, set[str]] = defaultdict(set)
    gold_sets: dict[str, set[str]] = defaultdict(set)
    for row in pred_rows:
        if row["sheet_name"] not in {"个股新闻", "行业新闻"}:
            continue
        gold = gold_rows[row_key(row)]
        article_id = str(row.get("article_file_id") or f"{row['sheet_name']}:{row['sheet_row_index']}")
        pred_item = tuple_repr(row)
        gold_item = tuple_repr({**gold, "sheet_name": row["sheet_name"]}, is_gold=True)
        if include_event:
            pred_item = f"{pred_item}|{row.get('event_name') or ''}"
            gold_item = f"{gold_item}|{gold.get('event_name') or ''}"
        pred_sets[article_id].add(pred_item)
        gold_sets[article_id].add(gold_item)
    article_ids = sorted(set(pred_sets) | set(gold_sets))
    tp = fp = fn = exact = at_least_one = 0
    p_values = []
    r_values = []
    f_values = []
    for article_id in article_ids:
        pred_set = pred_sets.get(article_id, set())
        gold_set = gold_sets.get(article_id, set())
        inter = pred_set & gold_set
        tp += len(inter)
        fp += len(pred_set - gold_set)
        fn += len(gold_set - pred_set)
        exact += int(pred_set == gold_set)
        at_least_one += int(bool(inter))
        precision = len(inter) / len(pred_set) if pred_set else (1.0 if not gold_set else 0.0)
        recall = len(inter) / len(gold_set) if gold_set else (1.0 if not pred_set else 0.0)
        f1 = 0.0 if not precision or not recall else 2 * precision * recall / (precision + recall)
        p_values.append(precision)
        r_values.append(recall)
        f_values.append(f1)
    micro_p = tp / (tp + fp) if tp + fp else 0.0
    micro_r = tp / (tp + fn) if tp + fn else 0.0
    micro_f1 = 0.0 if not micro_p or not micro_r else 2 * micro_p * micro_r / (micro_p + micro_r)
    return {
        "article_count": len(article_ids),
        "micro_precision": round(float(micro_p), 6),
        "micro_recall": round(float(micro_r), 6),
        "micro_f1": round(float(micro_f1), 6),
        "macro_precision": round(float(np.mean(p_values)), 6) if p_values else 0.0,
        "macro_recall": round(float(np.mean(r_values)), 6) if r_values else 0.0,
        "macro_f1": round(float(np.mean(f_values)), 6) if f_values else 0.0,
        "article_exact_match": round(float(exact / len(article_ids)), 6) if article_ids else 0.0,
        "at_least_one_hit_rate": round(float(at_least_one / len(article_ids)), 6) if article_ids else 0.0,
    }


def event_emotion_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    y_true = []
    y_pred = []
    missing = wrong = extra = 0
    main_rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    for row in main_rows:
        gold = gold_rows[row_key(row)]
        gt = normalize_emotion(gold.get("event_emotion"))
        pr = normalize_emotion(row.get("event_emotion"))
        y_true.append(gt)
        y_pred.append(pr)
        if not pr or pr == "missing_prediction":
            missing += 1
        elif any(sep in pr for sep in ["|", ",", "、", ";", "/", "；"]):
            extra += 1
        elif pr != gt:
            wrong += 1
    labels = sorted(set(y_true) | set(y_pred))
    p, r, f, s = precision_recall_fscore_support(y_true, y_pred, labels=labels, zero_division=0)
    cm = confusion_matrix(y_true, y_pred, labels=labels)
    invalid = sum(1 for p_ in y_pred if p_ not in labels)
    return {
        "support": len(y_true),
        "accuracy": round(float(accuracy_score(y_true, y_pred)), 6),
        "macro_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="macro", zero_division=0)), 6),
        "weighted_f1": round(float(f1_score(y_true, y_pred, labels=labels, average="weighted", zero_division=0)), 6),
        "invalid_enum_count": int(invalid),
        "missing_prediction_count": int(missing),
        "wrong_prediction_count": int(wrong),
        "extra_prediction_count": int(extra),
        "per_class": [
            {"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])}
            for i in range(len(labels))
        ],
        "confusion_matrix": {labels[i]: {labels[j]: int(cm[i, j]) for j in range(len(labels))} for i in range(len(labels))},
    }


def rouge_scores(ref: str, hyp: str) -> tuple[float, float, float]:
    ref = str(ref or "")
    hyp = str(hyp or "")
    if not ref or not hyp:
        return 0.0, 0.0, 0.0

    def grams(text: str, n: int) -> Counter[str]:
        if len(text) < n:
            return Counter({text: 1}) if text else Counter()
        return Counter(text[i : i + n] for i in range(len(text) - n + 1))

    def f1_score_ng(ref_cnt: Counter[str], hyp_cnt: Counter[str]) -> float:
        if not ref_cnt or not hyp_cnt:
            return 0.0
        overlap = sum(min(ref_cnt[k], hyp_cnt[k]) for k in ref_cnt.keys() & hyp_cnt.keys())
        prec = overlap / sum(hyp_cnt.values()) if hyp_cnt else 0.0
        rec = overlap / sum(ref_cnt.values()) if ref_cnt else 0.0
        return 0.0 if not prec or not rec else 2 * prec * rec / (prec + rec)

    def lcs(a: str, b: str) -> int:
        if not a or not b:
            return 0
        prev = [0] * (len(b) + 1)
        for ch_a in a:
            cur = [0]
            for j, ch_b in enumerate(b, 1):
                cur.append(prev[j - 1] + 1 if ch_a == ch_b else max(prev[j], cur[-1]))
            prev = cur
        return prev[-1]

    rouge1 = f1_score_ng(grams(ref, 1), grams(hyp, 1))
    rouge2 = f1_score_ng(grams(ref, 2), grams(hyp, 2))
    lcs_len = lcs(ref, hyp)
    rougeL = 0.0 if not ref or not hyp else (2 * lcs_len) / (len(ref) + len(hyp))
    return round(float(rouge1), 6), round(float(rouge2), 6), round(float(rougeL), 6)


def impact_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    ref_texts = []
    hyp_texts = []
    subject_ok = []
    event_ok = []
    emotion_ok = []
    evidence_ok = []
    unsupported = []
    missing = wrong = extra = 0
    lengths = []
    for row in rows:
        gold = gold_rows[row_key(row)]
        ref = str(gold.get("event_impact_analysis") or "")
        hyp = str(row.get("event_impact_analysis") or "")
        ref_texts.append(ref)
        hyp_texts.append(hyp)
        lengths.append(len(hyp))
        if not hyp or hyp == "missing_prediction":
            missing += 1
        target = str(gold.get("secu_abbr") or gold.get("industry_name") or "")
        subject_hit = bool(target and target in hyp)
        event_hit = bool(str(gold.get("event_name") or "") in hyp)
        emotion_hit = bool(str(gold.get("event_emotion") or "") in hyp)
        subject_ok.append(subject_hit)
        event_ok.append(event_hit)
        emotion_ok.append(emotion_hit)
        quote = re.search(r"原文证据(?:为)?[：:，, ]*[“\"](.+?)[”\"]", hyp)
        evidence = quote.group(1).strip() if quote else ""
        evidence_hit = bool(evidence and evidence in str(gold.get("content") or ""))
        evidence_ok.append(evidence_hit)
        unsupported_hit = bool(hyp and evidence and evidence not in str(gold.get("content") or ""))
        unsupported.append(unsupported_hit)
        if hyp and (not subject_hit or not event_hit or not emotion_hit or not evidence_hit):
            wrong += 1
        if unsupported_hit:
            extra += 1
    rouge1 = []
    rouge2 = []
    rougeL = []
    for ref, hyp in zip(ref_texts, hyp_texts):
        r1, r2, rl = rouge_scores(ref, hyp)
        rouge1.append(r1)
        rouge2.append(r2)
        rougeL.append(rl)
    return {
        "support": len(rows),
        "nonempty_coverage": round(float(np.mean([bool(t) for t in hyp_texts])), 6) if hyp_texts else 0.0,
        "missing_prediction_count": int(missing),
        "wrong_prediction_count": int(wrong),
        "extra_prediction_count": int(extra),
        "subject_consistency_rate": round(float(np.mean(subject_ok)), 6) if subject_ok else 0.0,
        "event_name_consistency_rate": round(float(np.mean(event_ok)), 6) if event_ok else 0.0,
        "emotion_consistency_rate": round(float(np.mean(emotion_ok)), 6) if emotion_ok else 0.0,
        "evidence_quote_support_rate": round(float(np.mean(evidence_ok)), 6) if evidence_ok else 0.0,
        "unsupported_fact_rate": round(float(np.mean(unsupported)), 6) if unsupported else 0.0,
        "rouge_1": round(float(np.mean(rouge1)), 6) if rouge1 else 0.0,
        "rouge_2": round(float(np.mean(rouge2)), 6) if rouge2 else 0.0,
        "rouge_l": round(float(np.mean(rougeL)), 6) if rougeL else 0.0,
        "length_mean": round(float(np.mean(lengths)), 3) if lengths else 0.0,
        "length_p50": round(float(np.median(lengths)), 3) if lengths else 0.0,
        "length_p95": round(float(np.quantile(lengths, 0.95)), 3) if lengths else 0.0,
        "template_duplicate_rate": round(float(1 - len(set(hyp_texts)) / max(1, len(hyp_texts))), 6),
    }


def bcubed_metrics(gold_cluster: list[str], pred_cluster: list[str]) -> dict[str, float]:
    if not gold_cluster:
        return {"bcubed_precision": 0.0, "bcubed_recall": 0.0, "bcubed_f1": 0.0}
    gold_members: dict[str, set[int]] = defaultdict(set)
    pred_members: dict[str, set[int]] = defaultdict(set)
    for idx, (g, p) in enumerate(zip(gold_cluster, pred_cluster)):
        gold_members[g].add(idx)
        pred_members[p].add(idx)
    precisions = []
    recalls = []
    for idx, (g, p) in enumerate(zip(gold_cluster, pred_cluster)):
        overlap = len(gold_members[g] & pred_members[p])
        precisions.append(overlap / len(pred_members[p]) if pred_members[p] else 0.0)
        recalls.append(overlap / len(gold_members[g]) if gold_members[g] else 0.0)
    precision = float(np.mean(precisions)) if precisions else 0.0
    recall = float(np.mean(recalls)) if recalls else 0.0
    f1 = 0.0 if not precision or not recall else 2 * precision * recall / (precision + recall)
    return {
        "bcubed_precision": round(precision, 6),
        "bcubed_recall": round(recall, 6),
        "bcubed_f1": round(float(f1), 6),
    }


def duplication_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    rows = []
    pred_dup = []
    gold_dup = []
    missing = wrong = extra = 0
    for idx, row in enumerate(pred_rows):
        gold = gold_rows[row_key(row)]
        gold_id = str(gold.get("duplication_id") or "")
        pred_id = str(row.get("duplication_id") or "")
        if not pred_id or pred_id == "missing_prediction":
            missing += 1
        elif any(sep in pred_id for sep in ["|", ",", "、", ";", "/", "；"]):
            extra += 1
        elif pred_id != gold_id:
            wrong += 1
        eval_row = {**row, "row_id": idx, "duplication_id": gold_id}
        if row["sheet_name"] == "个股新闻":
            eval_row.update({
                "trading_code": norm_code(gold.get("trading_code")),
                "secu_abbr": str(gold.get("secu_abbr") or ""),
                "industry_code": "",
                "industry_name": "",
                "event_name": str(gold.get("event_name") or ""),
            })
        elif row["sheet_name"] == "行业新闻":
            eval_row.update({
                "trading_code": "",
                "secu_abbr": "",
                "industry_code": norm_code(gold.get("industry_code")),
                "industry_name": str(gold.get("industry_name") or ""),
                "event_name": str(gold.get("event_name") or ""),
            })
        rows.append(eval_row)
        pred_dup.append(pred_id)
        gold_dup.append(gold_id or f"singleton_gold_{idx}")
    pred_cluster_for_metrics = [p or f"singleton_pred_{idx}" for idx, p in enumerate(pred_dup)]
    gold_clusters = Counter(gold_dup)
    pred_clusters = Counter(pred_cluster_for_metrics)
    out = pairwise_dup_metrics(rows, pred_dup)
    out.update(bcubed_metrics(gold_dup, pred_cluster_for_metrics))
    out.update(
        {
            "gold_cluster_count": len(gold_clusters),
            "pred_cluster_count": len(pred_clusters),
            "gold_multi_doc_cluster_count": int(sum(1 for v in gold_clusters.values() if v > 1)),
            "pred_multi_doc_cluster_count": int(sum(1 for v in pred_clusters.values() if v > 1)),
            "missing_prediction_count": int(missing),
            "wrong_prediction_count": int(wrong),
            "extra_prediction_count": int(extra),
            "blank_prediction_as_singleton_count": int(sum(1 for p in pred_dup if not p)),
            "validation_insufficient": bool(sum(1 for v in gold_clusters.values() if v > 1) < 3),
        }
    )
    return out


def full_record_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    exact = []
    any_tag_correct = []
    target_correct = []
    event_correct = []
    emotion_correct = []
    field_error_counts: Counter[str] = Counter()
    field_total_counts: Counter[str] = Counter()
    for row in rows:
        gold = gold_rows[row_key(row)]
        target_ok = tuple_repr(row) == tuple_repr({**gold, "sheet_name": row["sheet_name"]}, is_gold=True)
        event_ok = str(row.get("event_name") or "") == str(gold.get("event_name") or "")
        emotion_ok = str(row.get("event_emotion") or "") == str(gold.get("event_emotion") or "")
        dup_ok = norm_text(row.get("duplication_id")) == norm_text(gold.get("duplication_id"))
        target_type_ok = str(row.get("target_type") or "") == sheet_gold_type(row["sheet_name"])
        target_correct.append(target_ok)
        event_correct.append(event_ok)
        emotion_correct.append(emotion_ok)
        discrete_pred = [
            str(row.get("target_type") or ""),
            str(row.get("trading_code") or ""),
            str(row.get("secu_abbr") or ""),
            str(row.get("industry_code") or ""),
            str(row.get("industry_name") or ""),
            str(row.get("event_name") or ""),
            str(row.get("event_emotion") or ""),
            str(row.get("duplication_id") or ""),
        ]
        discrete_gold = [
            sheet_gold_type(row["sheet_name"]),
            str(gold.get("trading_code") or ""),
            str(gold.get("secu_abbr") or ""),
            str(gold.get("industry_code") or ""),
            str(gold.get("industry_name") or ""),
            str(gold.get("event_name") or ""),
            str(gold.get("event_emotion") or ""),
            str(gold.get("duplication_id") or ""),
        ]
        field_checks = {
            "target_type": target_type_ok,
            "target_entity": target_ok,
            "event_name": event_ok,
            "event_emotion": emotion_ok,
            "duplication_id": dup_ok,
        }
        for field, ok in field_checks.items():
            field_total_counts[field] += 1
            if not ok:
                field_error_counts[field] += 1
        any_tag_correct.append(any(field_checks.values()))
        exact.append(all(norm_text(a) == norm_text(b) for a, b in zip(discrete_pred, discrete_gold)))
    total_field_errors = sum(field_error_counts.values())
    return {
        "support": len(rows),
        "any_tag_correct_article_rate": round(float(np.mean(any_tag_correct)), 6) if any_tag_correct else 0.0,
        "target_and_event_event_emotion_exact_rate": round(float(np.mean([t and e and m for t, e, m in zip(target_correct, event_correct, emotion_correct)])), 6) if rows else 0.0,
        "target_event_emotion_exact_rate": round(float(np.mean([t and e and m for t, e, m in zip(target_correct, event_correct, emotion_correct)])), 6) if rows else 0.0,
        "all_discrete_tags_exact_rate": round(float(np.mean(exact)), 6) if exact else 0.0,
        "complete_record_exact_match_accuracy": round(float(np.mean(exact)), 6) if exact else 0.0,
        "target_event_exact_rate": round(float(np.mean([t and e for t, e in zip(target_correct, event_correct)])), 6) if rows else 0.0,
        "target_exact_rate": round(float(np.mean(target_correct)), 6) if target_correct else 0.0,
        "event_exact_rate": round(float(np.mean(event_correct)), 6) if event_correct else 0.0,
        "emotion_exact_rate": round(float(np.mean(emotion_correct)), 6) if emotion_correct else 0.0,
        "field_error_counts": dict(field_error_counts),
        "field_error_rates": {field: round(float(field_error_counts[field] / field_total_counts[field]), 6) for field in sorted(field_total_counts)},
        "field_error_contribution_ratio": {
            field: round(float(count / total_field_errors), 6) if total_field_errors else 0.0
            for field, count in sorted(field_error_counts.items())
        },
    }


def input_preservation_metrics(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    fields = ["article_title", "article_source", "article_publish_time", "content"]
    result: dict[str, Any] = {"support": len(pred_rows)}
    for field in fields:
        checks = [norm_text(row.get(field)) == norm_text(gold_rows[row_key(row)].get(field)) for row in pred_rows]
        result[f"{field}_preservation_rate"] = round(float(np.mean(checks)), 6) if checks else 0.0
    return result


def target_field_breakdown(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> dict[str, Any]:
    company_rows = [row for row in pred_rows if row["sheet_name"] == "个股新闻"]
    industry_rows = [row for row in pred_rows if row["sheet_name"] == "行业新闻"]
    dup_company_rows = [row for row in pred_rows if row["sheet_name"] == "个股重复新闻"]
    dup_industry_rows = [row for row in pred_rows if row["sheet_name"] == "行业重复新闻"]

    def field_stats(rows: list[dict[str, Any]], gold_field: str, pred_field: str) -> dict[str, Any]:
        y_true = []
        y_pred = []
        missing = wrong = over = 0
        for row in rows:
            gold = gold_rows[row_key(row)]
            g = norm_code(gold.get(gold_field), 6) if "code" in gold_field else norm_text(gold.get(gold_field))
            p = norm_code(row.get(pred_field), 6) if "code" in pred_field else norm_text(row.get(pred_field))
            y_true.append(g)
            y_pred.append(p)
            if not p or p == "missing_prediction":
                missing += 1
            elif any(sep in p for sep in ["|", ",", "、", ";", "/", "；"]):
                over += 1
            elif g and p != g:
                wrong += 1
        return {
            "support": len(y_true),
            "accuracy": round(float(accuracy_score(y_true, y_pred)), 6) if y_true else 0.0,
            "missing_prediction_count": missing,
            "wrong_prediction_count": wrong,
            "multi_prediction_count": over,
        }

    return {
        "main_target_type": field_stats(company_rows + industry_rows, "sheet_name", "target_type"),
        "company": {
            "trading_code": field_stats(company_rows, "trading_code", "trading_code"),
            "secu_abbr_exact": field_stats(company_rows, "secu_abbr", "secu_abbr"),
            "secu_abbr_normalized_accuracy": round(float(np.mean([norm_text(row.get("secu_abbr")) == norm_text(gold_rows[row_key(row)].get("secu_abbr")) for row in company_rows])), 6) if company_rows else 0.0,
        },
        "industry": {
            "industry_code": field_stats(industry_rows, "industry_code", "industry_code"),
            "industry_name_exact": field_stats(industry_rows, "industry_name", "industry_name"),
            "industry_name_normalized_accuracy": round(float(np.mean([norm_text(row.get("industry_name")) == norm_text(gold_rows[row_key(row)].get("industry_name")) for row in industry_rows])), 6) if industry_rows else 0.0,
        },
        "duplicate_sheet_target_type": {
            "company_rows": len(dup_company_rows),
            "industry_rows": len(dup_industry_rows),
        },
    }


def build_error_rows(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> list[dict[str, Any]]:
    errors = []
    for row in pred_rows:
        gold = gold_rows[row_key(row)]
        if row["sheet_name"] in {"个股新闻", "行业新闻"}:
            target_ok = tuple_repr(row) == tuple_repr({**gold, "sheet_name": row["sheet_name"]}, is_gold=True)
            event_ok = str(row.get("event_name") or "") == str(gold.get("event_name") or "")
            emotion_ok = str(row.get("event_emotion") or "") == str(gold.get("event_emotion") or "")
            impact_ok = norm_text(row.get("event_impact_analysis")) == norm_text(gold.get("event_impact_analysis"))
        else:
            target_ok = str(row.get("target_type") or "") == sheet_gold_type(row["sheet_name"])
            event_ok = True
            emotion_ok = True
            impact_ok = True
        dup_ok = norm_text(row.get("duplication_id")) == norm_text(gold.get("duplication_id"))
        if not (target_ok and event_ok and emotion_ok and impact_ok and dup_ok):
            errors.append(
                {
                    "sheet_name": row["sheet_name"],
                    "sheet_row_index": row["sheet_row_index"],
                    "article_file_id": row["article_file_id"],
                    "gold_target_type": sheet_gold_type(row["sheet_name"]),
                    "pred_target_type": row.get("target_type"),
                    "gold_trading_code": gold.get("trading_code", ""),
                    "pred_trading_code": row.get("trading_code", ""),
                    "gold_secu_abbr": gold.get("secu_abbr", ""),
                    "pred_secu_abbr": row.get("secu_abbr", ""),
                    "gold_industry_code": gold.get("industry_code", ""),
                    "pred_industry_code": row.get("industry_code", ""),
                    "gold_industry_name": gold.get("industry_name", ""),
                    "pred_industry_name": row.get("industry_name", ""),
                    "gold_event_name": gold.get("event_name", ""),
                    "pred_event_name": row.get("event_name", ""),
                    "gold_event_emotion": gold.get("event_emotion", ""),
                    "pred_event_emotion": row.get("event_emotion", ""),
                    "gold_duplication_id": gold.get("duplication_id", ""),
                    "pred_duplication_id": row.get("duplication_id", ""),
                    "error_reason": "|".join(
                        [
                            name
                            for name, ok in [
                                ("target", target_ok),
                                ("event", event_ok),
                                ("emotion", emotion_ok),
                                ("impact", impact_ok),
                                ("duplication", dup_ok),
                            ]
                            if not ok
                        ]
                    ),
                }
            )
    return errors


def summarize_history(current_metrics: dict[str, Any]) -> dict[str, Any]:
    hist = read_json(HIST_EVAL_JSON)
    current_model_path = str(HIER_MODEL_PKL)
    historical_source = hist["sources"]["formal_prediction_files"]
    same_model = False
    hist_metrics = hist["partial_metrics"]
    hist_match = hist["matchability"]
    return {
        "historical_status": hist["evaluation_status"],
        "historical_acceptance_conclusion": hist["acceptance_conclusion"],
        "historical_formal_prediction_files": historical_source,
        "historical_gold_files": hist["sources"]["raw_gold_files"],
        "historical_gold_fields_read": [
            "schema_id",
            "target_type",
            "target_code",
            "target_name",
            "event_name",
            "emotion",
            "impact_analysis",
            "duplication_id",
        ],
        "historical_model_same_as_current": same_model,
        "historical_partial_metrics": hist_metrics,
        "historical_matchability": hist_match,
        "current_vs_historical_note": "historical metrics use only 25 overlapping gold-test articles; current evaluation covers the full labeled official test workbook, so direct deltas are scope-mismatched.",
        "current_vs_historical_delta": {
            "event_name_micro_f1_delta": round(float(current_metrics["event_name"]["micro_f1"] - hist_metrics["raw_event_name_field_metrics"]["f1"]), 6),
            "event_name_macro_f1_delta": round(float(current_metrics["event_name"]["fixed_64_macro_f1"] - hist_metrics["raw_event_name_field_metrics"]["f1"]), 6),
        },
        "historical_model_same_as_current_reason": "historical evaluation was performed on final_task1 outputs, not the frozen Step 5.8-HC pickle.",
        "current_model_path": current_model_path,
    }


def write_reports(report: dict[str, Any], details: dict[str, pd.DataFrame]) -> None:
    write_json(REPORT_JSON, report)
    REPORT_MD.write_text(
        "\n".join(
            [
                "# Task1 Step 5.8-HF Retrospective Official Test Evaluation",
                "",
                f"- Status: `{report['status']}`",
                f"- Non-blind: `{report['non_blind']}`",
                f"- Test rows: `{report['prediction_inventory']['total_rows']}`",
                f"- Official event_name micro/macro: `{report['event_name']['micro_f1']}` / `{report['event_name']['fixed_64_macro_f1']}`",
                f"- Target tuple exact: `{report['target_tuple']['accuracy']}`",
                f"- Target+event joint micro/macro: `{report['target_event_joint']['micro_f1']}` / `{report['target_event_joint']['macro_f1']}`",
                f"- Emotion accuracy/macro: `{report['event_emotion']['accuracy']}` / `{report['event_emotion']['macro_f1']}`",
                f"- Duplication pairwise F1: `{report['duplication']['pairwise_f1']}`",
                f"- Historical evaluation model same as current: `{report['history']['historical_model_same_as_current']}`",
                f"- Formal SHA unchanged: `{report['formal_sha_match']}`",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(DETAILS_XLSX, engine="openpyxl") as writer:
        for name, df in details.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)


def write_confusion_matrix(report: dict[str, Any]) -> None:
    labels = report["event_name"]["labels"]
    cm = report["event_name"]["confusion_matrix"]
    df = pd.DataFrame(cm).reindex(index=labels, columns=labels).fillna(0).astype(int)
    per_class = pd.DataFrame(report["event_name"]["per_class"])
    with pd.ExcelWriter(CONFUSION_XLSX, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="confusion_matrix")
        per_class.to_excel(writer, sheet_name="per_class", index=False)


def write_error_workbook(errors: list[dict[str, Any]]) -> None:
    df = pd.DataFrame(errors)
    with pd.ExcelWriter(ERRORS_XLSX, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="all_tags_errors", index=False)


def write_impact_audit(pred_rows: list[dict[str, Any]], gold_rows: dict[tuple[str, int], dict[str, Any]]) -> None:
    main_rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    rows = []
    stratified = defaultdict(list)
    for row in main_rows:
        gold = gold_rows[row_key(row)]
        key = (
            "target_ok" if tuple_repr(row) == tuple_repr({**gold, "sheet_name": row["sheet_name"]}, is_gold=True) else "target_bad",
            "event_ok" if str(row.get("event_name")) == str(gold.get("event_name")) else "event_bad",
            "emotion_ok" if str(row.get("event_emotion")) == str(gold.get("event_emotion")) else "emotion_bad",
        )
        stratified[key].append(row)
    buckets = list(stratified.values())
    rng = random.Random(2026)
    sample_rows = []
    for bucket in buckets:
        if not bucket:
            continue
        sample_rows.extend(rng.sample(bucket, k=min(len(bucket), max(1, 100 // max(1, len(buckets))))))
    if len(sample_rows) < 100:
        remainder = [row for row in main_rows if row not in sample_rows]
        sample_rows.extend(rng.sample(remainder, k=min(len(remainder), 100 - len(sample_rows))))
    sample_rows = sample_rows[:100]
    for row in sample_rows:
        gold = gold_rows[row_key(row)]
        rows.append(
            {
                "sheet_name": row["sheet_name"],
                "sheet_row_index": row["sheet_row_index"],
                "article_file_id": row["article_file_id"],
                "article_title": row["article_title"],
                "gold_target": gold.get("secu_abbr") or gold.get("industry_name"),
                "pred_target": row.get("secu_abbr") or row.get("industry_name"),
                "gold_event_name": gold.get("event_name"),
                "pred_event_name": row.get("event_name"),
                "gold_event_emotion": gold.get("event_emotion"),
                "pred_event_emotion": row.get("event_emotion"),
                "gold_impact_analysis": gold.get("event_impact_analysis"),
                "pred_impact_analysis": row.get("event_impact_analysis"),
                "evidence_quote": evidence_quote(row),
                "human_supported": "",
                "human_note": "",
            }
        )
    with pd.ExcelWriter(IMPACT_AUDIT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="impact_audit_100", index=False)


def main() -> int:
    start = time.time()
    backup_dir = ensure_dir()
    before_sha = protected_sha(backup_dir / "protected_sha256_before.txt")
    headers_labeled = read_headers(RAW_TEST_TAGGED)
    headers_unlabeled = read_headers(RAW_TEST_UNLABELED)
    if not headers_unlabeled:
        raise RuntimeError("news_without_tags_test.xlsx has no readable sheets")
    input_cols = [c for c in headers_unlabeled[next(iter(headers_unlabeled))] if c in INPUT_COLS]
    if input_cols != INPUT_COLS:
        raise RuntimeError(f"unexpected unlabeled input contract: {input_cols}")
    update_state("prediction_inputs_loaded", input_cols=input_cols, labeled_sheets=list(headers_labeled), unlabeled_sheets=list(headers_unlabeled))
    inputs = load_prediction_inputs()
    company44_labels, industry20_labels = load_official_scope_labels()
    freeze_manifest: dict[str, Any]
    if PRED_JSONL.exists() and FREEZE_MANIFEST.exists():
        freeze_manifest = read_json(FREEZE_MANIFEST)
        pred_rows = read_jsonl_local(PRED_JSONL)
        pred_sha = sha256_file(PRED_JSONL)
        if freeze_manifest.get("prediction_sha256") and freeze_manifest["prediction_sha256"] != pred_sha:
            raise RuntimeError("existing frozen prediction SHA does not match manifest")
        if len(pred_rows) != len(inputs):
            raise RuntimeError(f"frozen prediction row count mismatch: {len(pred_rows)} vs {len(inputs)}")
        hier_artifact = load_hier_artifact_only()
        maps = {
            "official_event_names": sorted(hier_artifact["taxonomy"]["official64_names"]),
            "company44_event_names": company44_labels,
            "industry20_event_names": industry20_labels,
        }
        update_state("frozen_predictions_loaded_for_retro_eval", prediction_rows=len(pred_rows), prediction_sha256=pred_sha)
    else:
        hier_artifact, official_model, emotion_bundle, generator, maps, train_pairs, _ = load_models()
        if not isinstance(emotion_bundle, dict) or "model" not in emotion_bundle:
            raise RuntimeError("unexpected emotion bundle structure")
        update_state("predicting", row_count=len(inputs))
        pred_rows = predict_rows(inputs, hier_artifact, official_model, emotion_bundle, generator, maps, train_pairs)
        write_jsonl(PRED_JSONL, pred_rows)
        pred_sha = sha256_file(PRED_JSONL)
        freeze_manifest = {
            "status": "frozen",
            "model_frozen": True,
            "thresholds_frozen": True,
            "taxonomy_frozen": True,
            "official_test_labels_read": False,
            "prediction_file": str(PRED_JSONL),
            "prediction_sha256": pred_sha,
            "model_sha256": sha256_file(HIER_MODEL_PKL),
            "official_model_sha256": sha256_file(Path(hier_artifact["official_model_path"])),
            "emotion_model_sha256": sha256_file(EMOTION_MODEL_PKL),
            "taxonomy_sha256": sha256_file(TAXONOMY_JSON),
            "candidate_config_sha256": sha256_file(TASK1_CANDIDATE_CONFIG),
            "formal_task1_sha256_before": before_sha,
        }
        write_json(FREEZE_MANIFEST, freeze_manifest)
        update_state("gold_freeze_complete", prediction_rows=len(pred_rows), prediction_sha256=pred_sha)

    gold_rows = load_gold_for_eval()
    if not gold_rows:
        raise RuntimeError("no gold rows loaded for evaluation")
    update_state("gold_labels_read", official_test_labels_read=True)
    after_sha = protected_sha(backup_dir / "protected_sha256_after.txt")
    formal_sha_match = before_sha == after_sha

    main_pred = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    target_tuple = compute_target_metrics(main_pred, gold_rows, {"个股新闻", "行业新闻"})
    target_article_set_report = article_set_metrics(pred_rows, gold_rows, include_event=False)
    target_type_metrics = compute_target_type_metrics(pred_rows, gold_rows)
    company_code_metrics = code_metrics(pred_rows, gold_rows, "个股新闻")
    industry_code_metrics = code_metrics(pred_rows, gold_rows, "行业新闻")
    company_name_metrics = text_metrics([row for row in pred_rows if row["sheet_name"] == "个股新闻"], gold_rows, "secu_abbr", {"个股新闻"})
    industry_name_metrics = text_metrics([row for row in pred_rows if row["sheet_name"] == "行业新闻"], gold_rows, "industry_name", {"行业新闻"})
    event_metrics_report = event_metrics(
        pred_rows,
        gold_rows,
        hier_artifact["taxonomy"]["official64_names"],
        maps["company44_event_names"],
        maps["industry20_event_names"],
    )
    target_event_joint_report = target_event_joint_metrics(pred_rows, gold_rows)
    target_event_article_set_report = article_set_metrics(pred_rows, gold_rows, include_event=True)
    emotion_metrics_report = event_emotion_metrics(pred_rows, gold_rows)
    impact_metrics_report = impact_metrics(pred_rows, gold_rows)
    duplication_report = duplication_metrics(pred_rows, gold_rows)
    full_record_report = full_record_metrics(pred_rows, gold_rows)
    input_preservation_report = input_preservation_metrics(pred_rows, gold_rows)
    event_layer_rows = [row for row in pred_rows if row["sheet_name"] in {"个股新闻", "行业新闻"}]
    event_layer_metrics = {
        "support": len(event_layer_rows),
        "coverage": round(float(np.mean([bool(r.get("event_layer")) for r in event_layer_rows])), 6) if event_layer_rows else 0.0,
        "parent_child_consistency": round(float(np.mean([
            bool(r.get("event_layer")) and bool(r.get("event_parent")) and bool(r.get("event_group"))
            for r in event_layer_rows
        ])), 6) if event_layer_rows else 0.0,
        "illegal_leaf_count": int(sum(1 for r in event_layer_rows if str(r.get("event_layer")) == "missing_prediction" or str(r.get("event_parent")) == "missing_prediction")),
    }
    field_error_summary = {
        "target_type": {
            "missing_prediction_count": target_type_metrics["missing_prediction_count"],
            "wrong_prediction_count": target_type_metrics["wrong_prediction_count"],
            "extra_prediction_count": target_type_metrics["extra_prediction_count"],
        },
        "trading_code": {
            "missing_prediction_count": company_code_metrics["missing_prediction_count"],
            "wrong_prediction_count": company_code_metrics["wrong_prediction_count"],
            "extra_prediction_count": company_code_metrics["multi_prediction_count"],
        },
        "secu_abbr": {
            "missing_prediction_count": company_name_metrics["missing_prediction_count"],
            "wrong_prediction_count": company_name_metrics["wrong_prediction_count"],
            "extra_prediction_count": company_name_metrics["extra_prediction_count"],
        },
        "industry_code": {
            "missing_prediction_count": industry_code_metrics["missing_prediction_count"],
            "wrong_prediction_count": industry_code_metrics["wrong_prediction_count"],
            "extra_prediction_count": industry_code_metrics["multi_prediction_count"],
        },
        "industry_name": {
            "missing_prediction_count": industry_name_metrics["missing_prediction_count"],
            "wrong_prediction_count": industry_name_metrics["wrong_prediction_count"],
            "extra_prediction_count": industry_name_metrics["extra_prediction_count"],
        },
        "event_name": {
            "missing_prediction_count": event_metrics_report["complete_miss_count"],
            "wrong_prediction_count": event_metrics_report["misclassification_count"],
            "extra_prediction_count": event_metrics_report["multi_prediction_count"],
        },
        "event_emotion": {
            "missing_prediction_count": emotion_metrics_report["missing_prediction_count"],
            "wrong_prediction_count": emotion_metrics_report["wrong_prediction_count"],
            "extra_prediction_count": emotion_metrics_report["extra_prediction_count"],
        },
        "event_impact_analysis": {
            "missing_prediction_count": impact_metrics_report["missing_prediction_count"],
            "wrong_prediction_count": impact_metrics_report["wrong_prediction_count"],
            "extra_prediction_count": impact_metrics_report["extra_prediction_count"],
        },
        "duplication_id": {
            "missing_prediction_count": duplication_report["missing_prediction_count"],
            "wrong_prediction_count": duplication_report["wrong_prediction_count"],
            "extra_prediction_count": duplication_report["extra_prediction_count"],
        },
    }

    history = summarize_history({
        "event_name": event_metrics_report,
    })

    report = {
        "status": "retrospective_official_test_evaluation_complete",
        "non_blind": True,
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "prediction_inventory": {
            "total_rows": len(pred_rows),
            "sheet_counts": dict(Counter(r["sheet_name"] for r in pred_rows)),
            "main_labeled_rows": len([r for r in pred_rows if r["sheet_name"] in {"个股新闻", "行业新闻"}]),
            "duplicate_rows": len([r for r in pred_rows if r["sheet_name"] in {"个股重复新闻", "行业重复新闻"}]),
            "successful_prediction_rows": int(sum(1 for r in pred_rows if r.get("prediction_status") == "success")),
            "partial_prediction_rows": int(sum(1 for r in pred_rows if r.get("prediction_status") == "partial")),
            "failed_prediction_rows": int(sum(1 for r in pred_rows if r.get("prediction_status") == "failed")),
            "prediction_sha256": pred_sha,
        },
        "contract": {
            "unlabeled_headers": headers_unlabeled,
            "tagged_headers": headers_labeled,
            "input_cols": input_cols,
            "prediction_usecols_only": True,
        },
        "target_type": target_type_metrics,
        "target_tuple": {**target_tuple, "article_set_metrics": target_article_set_report},
        "target_event_joint": {**target_event_joint_report, "article_set_metrics": target_event_article_set_report},
        "company_fields": {
            "trading_code": company_code_metrics,
            "secu_abbr": company_name_metrics,
            "secu_abbr_normalized_accuracy": company_name_metrics["normalized_accuracy"],
        },
        "industry_fields": {
            "industry_code": industry_code_metrics,
            "industry_name": industry_name_metrics,
            "industry_name_normalized_accuracy": industry_name_metrics["normalized_accuracy"],
        },
        "event_name": {
            **event_metrics_report,
            "labels": hier_artifact["taxonomy"]["official64_names"],
        },
        "event_layer": event_layer_metrics,
        "event_emotion": emotion_metrics_report,
        "event_impact_analysis": impact_metrics_report,
        "duplication": duplication_report,
        "field_error_summary": field_error_summary,
        "full_record": full_record_report,
        "input_preservation": input_preservation_report,
        "history": history,
        "coverage": {
            "main_rows_with_full_gold": len(main_pred),
            "duplicate_rows_with_dup_gold": len([r for r in pred_rows if r["sheet_name"] in {"个股重复新闻", "行业重复新闻"}]),
            "missing_prediction_fields": {
                key: int(sum(1 for row in pred_rows if str(row.get(key) or "") in {"", "missing_prediction"}))
                for key in ["target_type", "event_name", "event_emotion", "event_impact_analysis", "duplication_id"]
            },
        },
        "model": {
            "hierarchical_model": str(HIER_MODEL_PKL),
            "official_model": str(hier_artifact["official_model_path"]),
            "emotion_model": str(EMOTION_MODEL_PKL),
            "candidate_config": str(TASK1_CANDIDATE_CONFIG),
            "taxonomy": str(TAXONOMY_JSON),
            "current_model_same_as_historical": history["historical_model_same_as_current"],
        },
        "formal_sha_before": before_sha,
        "formal_sha_after": after_sha,
        "formal_sha_match": formal_sha_match,
        "llm_calls": 0,
        "external_llm_calls": 0,
        "local_32b_calls": 0,
        "backup_dir": str(backup_dir),
        "elapsed_wall_seconds": round(time.time() - start, 3),
        "recomputed_from_existing_frozen_predictions": bool(freeze_manifest.get("official_test_labels_read") and PRED_JSONL.exists()),
    }

    if report["status"] not in LEGAL_STATUSES:
        raise AssertionError(report["status"])

    errors = build_error_rows(pred_rows, gold_rows)
    write_error_workbook(errors)
    write_impact_audit(pred_rows, gold_rows)
    write_confusion_matrix(report)

    details_tables = {
        "prediction_inventory": pd.DataFrame([report["prediction_inventory"]]),
        "target_type": pd.DataFrame(target_type_metrics["per_class"]),
        "target_tuple": pd.DataFrame(target_tuple["per_class"]),
        "target_article_set": pd.DataFrame([target_article_set_report]),
        "target_event_joint": pd.DataFrame([target_event_joint_report]),
        "target_event_article_set": pd.DataFrame([target_event_article_set_report]),
        "company_fields": pd.DataFrame([company_code_metrics, company_name_metrics]),
        "industry_fields": pd.DataFrame([industry_code_metrics, industry_name_metrics]),
        "event_name_per_class": pd.DataFrame(report["event_name"]["per_class"]),
        "event_emotion_per_class": pd.DataFrame(report["event_emotion"]["per_class"]),
        "impact_audit_summary": pd.DataFrame([impact_metrics_report]),
        "duplication": pd.DataFrame([duplication_report]),
        "field_error_summary": pd.DataFrame([field_error_summary]),
        "full_record": pd.DataFrame([full_record_report]),
        "input_preservation": pd.DataFrame([input_preservation_report]),
        "history": pd.DataFrame([history]),
        "coverage": pd.DataFrame([report["coverage"]]),
        "event_layer": pd.DataFrame([event_layer_metrics]),
    }
    write_reports(report, details_tables)

    checkpoint = {
        "completed": True,
        "status": report["status"],
        "prediction_file": str(PRED_JSONL),
        "prediction_sha256": pred_sha,
        "report_json": str(REPORT_JSON),
        "report_md": str(REPORT_MD),
        "details_xlsx": str(DETAILS_XLSX),
        "confusion_matrix_xlsx": str(CONFUSION_XLSX),
        "errors_xlsx": str(ERRORS_XLSX),
        "impact_audit_xlsx": str(IMPACT_AUDIT_XLSX),
        "official_test_labels_read": True,
        "formal_sha_match": formal_sha_match,
        "updated_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
    }
    write_json(CHECKPOINT, checkpoint)
    write_json(FREEZE_MANIFEST, {**freeze_manifest, "official_test_labels_read": True, "official_test_evaluation_count": 1, "post_test_tuning_allowed": False})
    sha_targets = [
        RUN_STATE,
        CHECKPOINT,
        FREEZE_MANIFEST,
        PRED_JSONL,
        REPORT_JSON,
        REPORT_MD,
        DETAILS_XLSX,
        CONFUSION_XLSX,
        ERRORS_XLSX,
        IMPACT_AUDIT_XLSX,
    ]
    SHA_SUMS.write_text("\n".join(f"{sha256_file(p)}  {p}" for p in sha_targets if p.exists()) + "\n", encoding="utf-8")
    update_state("completed", status=report["status"], prediction_rows=len(pred_rows), prediction_sha256=pred_sha, formal_sha_match=formal_sha_match)
    print(json.dumps({"status": report["status"], "prediction_rows": len(pred_rows), "formal_sha_match": formal_sha_match, "backup_dir": str(backup_dir)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
