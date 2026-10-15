#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build official-format Task1 shadow outputs without promoting formal files."""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs/task1_step5_8b_external_transfer"
XLSX_OUT = OUT_DIR / "task1_official_format_shadow.xlsx"
JSONL_OUT = OUT_DIR / "task1_official_format_shadow.jsonl"
PARQUET_OUT = OUT_DIR / "task1_official_format_shadow.parquet"
REPORT_JSON = ROOT / "reports/task1_step5_8b_official_format_contract_report.json"
REPORT_MD = ROOT / "reports/task1_step5_8b_official_format_contract_report.md"

RAW_TRAIN = ROOT / "data/raw/news_with_tags_train.xlsx"
FINAL_TRAIN = ROOT / "outputs/final_task1/final_events_train.jsonl"
FINAL_TEST = ROOT / "outputs/final_task1/final_events_test.jsonl"
RAW_UNLABELED = [
    ROOT / "data/raw/news_without_tags_train.xlsx",
    ROOT / "data/raw/news_without_tags_test.xlsx",
]
FORMAL_PROTECTED = [
    FINAL_TRAIN,
    FINAL_TEST,
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
]


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def official_headers() -> dict[str, list[str]]:
    wb = load_workbook(RAW_TRAIN, read_only=True, data_only=True)
    headers: dict[str, list[str]] = {}
    for sheet in ["个股新闻", "行业新闻"]:
        ws = wb[sheet]
        headers[sheet] = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def official_enums_and_maps() -> tuple[set[str], set[str], dict[str, dict[str, str]], dict[str, str]]:
    event_names: set[str] = set()
    emotions: set[str] = set()
    company_map: dict[str, dict[str, str]] = {}
    industry_map: dict[str, str] = {}
    for sheet in ["个股新闻", "行业新闻"]:
        df = pd.read_excel(RAW_TRAIN, sheet_name=sheet)
        event_names.update(str(x).strip() for x in df["event_name"].dropna())
        emotions.update(str(x).strip() for x in df["event_emotion"].dropna())
        if sheet == "个股新闻":
            for _, row in df.iterrows():
                code = str(row.get("trading_code", "")).split(".")[0].zfill(6)
                if not code or code == "000nan":
                    continue
                company_map[code] = {
                    "trading_code": code,
                    "secu_abbr": str(row.get("secu_abbr", "") or ""),
                    "industry_code": str(row.get("industry_code", "") or "").split(".")[0],
                    "industry_name": str(row.get("industry_name", "") or ""),
                }
        else:
            for _, row in df.iterrows():
                code = str(row.get("industry_code", "") or "").split(".")[0]
                if code:
                    industry_map[code] = str(row.get("industry_name", "") or "")
    return event_names, emotions, company_map, industry_map


def collect_raw_articles(ids: set[str]) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    for path in RAW_UNLABELED:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {name: i for i, name in enumerate(header)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            aid = str(row[idx["article_file_id"]] or "")
            if aid in ids:
                out[aid] = {
                    "article_file_id": aid,
                    "article_title": row[idx["article_title"]],
                    "article_source": row[idx["article_source"]],
                    "content": row[idx["content"]],
                    "article_publish_time": row[idx["article_publish_time"]],
                }
        wb.close()
    return out


def content_key(text: str) -> str:
    norm = re.sub(r"\s+", "", str(text or ""))
    return hashlib.sha256(norm.encode("utf-8")).hexdigest()


def evidence_quote(content: str, target_name: str, title: str) -> str:
    text = str(content or "")
    sentences = [s.strip() for s in re.split(r"(?<=[。！？!?；;])|\n", text) if s.strip()]
    for sent in sentences:
        if target_name and target_name in sent:
            return sent[:180]
    return (sentences[0] if sentences else str(title or ""))[:180]


def build_impact_analysis(row: dict[str, Any], emotion: str, quote: str) -> str:
    target = row.get("target_name") or row.get("event_name") or "相关主体"
    event_name = row.get("event_name") or "相关事件"
    return f"原文证据：{quote}。据此，该新闻涉及{target}的“{event_name}”事项；当前按{emotion}影响处理，未从正文中加入额外未证实金额、客户或产品信息。"


def normalize_code(value: Any, width: int | None = None) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    if width and text and text.isdigit():
        text = text.zfill(width)
    return text


def event_name_from_schema_id(schema_id: Any, official_event_names: set[str]) -> str:
    text = str(schema_id or "")
    for name in sorted(official_event_names, key=len, reverse=True):
        if text.endswith("_" + name):
            return name
    return ""


def clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if callable(getattr(value, "isoformat")) else str(value)
    if pd.isna(value):
        return ""
    return value


def clean_row(row: dict[str, Any]) -> dict[str, Any]:
    return {key: clean_value(value) for key, value in row.items()}


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    before = {str(path): sha256_file(path) for path in FORMAL_PROTECTED if path.exists()}
    headers = official_headers()
    official_event_names, official_emotions, company_map, industry_map = official_enums_and_maps()
    neutral = "中性" if "中性" in official_emotions else sorted(official_emotions)[0]

    final_rows = read_jsonl(FINAL_TRAIN) + read_jsonl(FINAL_TEST)
    needed_ids = {str(row.get("article_id")) for row in final_rows}
    raw = collect_raw_articles(needed_ids)

    content_groups: dict[str, list[str]] = defaultdict(list)
    for aid, meta in raw.items():
        content_groups[content_key(str(meta.get("content") or ""))].append(aid)
    duplication = {}
    for ids in content_groups.values():
        if len(ids) > 1:
            canonical = sorted(ids)[0]
            for aid in ids:
                duplication[aid] = canonical

    company_rows: list[dict[str, Any]] = []
    industry_rows: list[dict[str, Any]] = []
    union_rows: list[dict[str, Any]] = []
    mapping_conflicts = 0
    raw_missing = 0
    illegal_event_name = 0
    illegal_emotion = 0
    unchanged_errors = 0
    unsupported_fact_count = 0

    for pred in final_rows:
        aid = str(pred.get("article_id") or "")
        meta = raw.get(aid)
        if not meta:
            raw_missing += 1
            meta = {
                "article_file_id": aid,
                "article_title": pred.get("title", ""),
                "article_publish_time": pred.get("publish_time", ""),
                "article_source": "",
                "content": "",
            }
        event_name = str(pred.get("event_name") or "")
        if event_name not in official_event_names:
            recovered = event_name_from_schema_id(pred.get("schema_id"), official_event_names)
            if recovered:
                event_name = recovered
        if event_name not in official_event_names:
            illegal_event_name += 1
        emotion = neutral
        if emotion not in official_emotions:
            illegal_emotion += 1
        quote = evidence_quote(str(meta.get("content") or ""), str(pred.get("target_name") or ""), str(meta.get("article_title") or ""))
        analysis = build_impact_analysis(pred, emotion, quote)
        if quote and quote not in str(meta.get("content") or "") and quote != str(meta.get("article_title") or "")[:180]:
            unsupported_fact_count += 1

        base = {
            "article_file_id": meta["article_file_id"],
            "article_title": meta["article_title"],
            "article_publish_time": meta["article_publish_time"],
            "article_source": meta["article_source"],
            "content": meta["content"],
            "event_name": event_name,
            "event_emotion": emotion,
            "event_impact_analysis": analysis,
            "duplication_id": duplication.get(aid, ""),
        }
        if str(meta["article_file_id"]) != aid:
            unchanged_errors += 1
        if pred.get("target_type") == "company":
            code = normalize_code(pred.get("target_code"), 6)
            cmap = company_map.get(code)
            if not cmap:
                mapping_conflicts += 1
                cmap = {"trading_code": code, "secu_abbr": pred.get("target_name", ""), "industry_code": "", "industry_name": ""}
            out = {**base, **cmap}
            company_rows.append(clean_row({col: out.get(col, "") for col in headers["个股新闻"]}))
            union_rows.append(clean_row({col: out.get(col, "") for col in headers["个股新闻"]}))
        else:
            code = normalize_code(pred.get("target_code"))
            official_ind_name = industry_map.get(code, "")
            if official_ind_name and official_ind_name != str(pred.get("target_name", "")):
                mapping_conflicts += 1
            out = {**base, "industry_code": code, "industry_name": official_ind_name or pred.get("target_name", "")}
            industry_rows.append(clean_row({col: out.get(col, "") for col in headers["行业新闻"]}))
            union = {**{"trading_code": "", "secu_abbr": ""}, **out}
            union_rows.append(clean_row({col: union.get(col, "") for col in headers["个股新闻"]}))

    with pd.ExcelWriter(XLSX_OUT, engine="openpyxl") as writer:
        pd.DataFrame(company_rows, columns=headers["个股新闻"]).to_excel(writer, sheet_name="个股新闻", index=False)
        pd.DataFrame(industry_rows, columns=headers["行业新闻"]).to_excel(writer, sheet_name="行业新闻", index=False)

    with JSONL_OUT.open("w", encoding="utf-8") as handle:
        for row in union_rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")
    pd.DataFrame(union_rows, columns=headers["个股新闻"]).to_parquet(PARQUET_OUT, index=False)

    # Consistency check across normalized union rows.
    json_rows = read_jsonl(JSONL_OUT)
    parquet_rows = pd.read_parquet(PARQUET_OUT).fillna("").to_dict("records")
    formats_consistent = len(json_rows) == len(parquet_rows) == len(union_rows)
    for a, b in zip(json_rows, parquet_rows):
        if {k: str(v or "") for k, v in a.items()} != {k: str(v or "") for k, v in b.items()}:
            formats_consistent = False
            break

    forbidden = {"unmapped_external_event", "source_event_type", "trigger", "event_span", "arguments", "argument_roles", "field_provenance", "ai_assisted_silver"}
    forbidden_hits = sorted(forbidden & set(union_rows[0].keys())) if union_rows else []
    after = {str(path): sha256_file(path) for path in FORMAL_PROTECTED if path.exists()}

    report = {
        "status": "official_format_shadow_built",
        "generated_at": pd.Timestamp.now().isoformat(),
        "official_headers": headers,
        "official_emotion_values": sorted(official_emotions),
        "selected_shadow_emotion_fallback": neutral,
        "row_counts": {
            "company_rows": len(company_rows),
            "industry_rows": len(industry_rows),
            "total_rows": len(union_rows),
            "jsonl_rows": len(json_rows),
            "parquet_rows": len(parquet_rows),
        },
        "validation": {
            "row_count_consistent": formats_consistent,
            "illegal_event_name": illegal_event_name,
            "illegal_event_emotion": illegal_emotion,
            "company_or_industry_mapping_conflicts": mapping_conflicts,
            "raw_article_missing": raw_missing,
            "original_article_id_changed": unchanged_errors,
            "unsupported_fact_count_by_template_audit": unsupported_fact_count,
            "forbidden_internal_field_hits": forbidden_hits,
            "formal_sha_match": before == after,
        },
        "outputs": {
            "xlsx": str(XLSX_OUT),
            "jsonl": str(JSONL_OUT),
            "parquet": str(PARQUET_OUT),
        },
        "formal_sha_before": before,
        "formal_sha_after": after,
        "notes": [
            "This is a format shadow, not a promoted formal Task1 result.",
            "event_emotion is a conservative legal-enum fallback because no official emotion model is trained in this step.",
            "event_impact_analysis is deterministic and evidence-quote constrained.",
            "Internal extraction/training fields are intentionally excluded.",
        ],
    }
    REPORT_JSON.parent.mkdir(parents=True, exist_ok=True)
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    md = [
        "# Task1 Step 5.8-B Official Format Contract",
        "",
        f"- Status: `{report['status']}`",
        f"- Shadow XLSX: `{XLSX_OUT}`",
        f"- Shadow JSONL: `{JSONL_OUT}`",
        f"- Shadow Parquet: `{PARQUET_OUT}`",
        f"- Total rows: `{len(union_rows)}`",
        f"- Illegal event_name: `{illegal_event_name}`",
        f"- Illegal event_emotion: `{illegal_emotion}`",
        f"- Mapping conflicts: `{mapping_conflicts}`",
        f"- Raw article missing: `{raw_missing}`",
        f"- Forbidden internal field hits: `{forbidden_hits}`",
        f"- Formal SHA match: `{before == after}`",
        "",
        "## Field Separation",
        "1. 模型内部抽取结构：仅允许在intermediate/debug产物中出现。",
        "2. 外部数据训练结构：source_dataset/source_event_type等不得进入正式交付。",
        "3. 官方最终交付结构：严格使用官方Excel列名和顺序。",
    ]
    REPORT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "rows": len(union_rows), "validation": report["validation"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
