#!/usr/bin/env python3
"""Run the Step 4.5-B final Task2 Cached DEA candidate on the 260-row queue."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import random
import sqlite3
import subprocess
import sys
import time
import urllib.request
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.event_credibility import dea, fulltext  # noqa: E402
from src.event_credibility.retrieval_index import load_index_metadata, sqlite_path  # noqa: E402


DEFAULT_CONFIG = ROOT / "configs/task2_final_cached_dea_v1.yaml"
DEFAULT_INPUT = ROOT / "data/processed/event_credibility/llm_review_queue.jsonl"
DEFAULT_OUTPUT_DIR = ROOT / "outputs/event_credibility/final_cached_dea_v1_260"
DEFAULT_REPORT_JSON = ROOT / "reports/task2_final_cached_dea_v1_260_report.json"
DEFAULT_REPORT_MD = ROOT / "reports/task2_final_cached_dea_v1_260_report.md"
DEFAULT_COMPARISON_JSON = ROOT / "reports/task2_final_cached_dea_v1_260_comparison.json"
DEFAULT_COMPARISON_MD = ROOT / "reports/task2_final_cached_dea_v1_260_comparison.md"
DEFAULT_COMPARISON_XLSX = ROOT / "reports/task2_final_cached_dea_v1_260_changes.xlsx"
FORMAL_OUTPUTS = [
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
]
VALID_VERDICTS = {"support", "contradict", "insufficient", "unclear"}
REPORT_FILENAMES = {
    "task2_final_cached_dea_v1_260_report.json",
    "task2_final_cached_dea_v1_260_report.md",
    "task2_final_cached_dea_v1_260_comparison.json",
    "task2_final_cached_dea_v1_260_comparison.md",
    "task2_final_cached_dea_v1_260_changes.xlsx",
}


def resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def norm(value: Any) -> str:
    return fulltext.normalize_text(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def hash_paths(paths: list[Path]) -> dict[str, str]:
    return {str(path.relative_to(ROOT) if path.is_relative_to(ROOT) else path): sha256(path) for path in paths if path.exists()}


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(path)
    rows = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(f"{path}:{line_number} must contain a JSON object")
        rows.append(value)
    return rows


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, allow_nan=False) + "\n")


def load_config(path: Path) -> dict[str, Any]:
    config = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(config, dict):
        raise ValueError("config must contain a mapping")
    required_flags = {
        "dea_enabled": True,
        "safety_gate_enabled": False,
        "isr_enabled": False,
        "claim_contract_enabled": False,
        "score_fusion_enabled": False,
        "strict_json_validation": True,
        "checkpoint_enabled": True,
        "resume_enabled": True,
    }
    for key, expected in required_flags.items():
        if bool(config.get(key)) is not expected:
            raise ValueError(f"invalid final config flag {key}: {config.get(key)!r}")
    dea_config = config.get("dea") or {}
    if dea_config.get("trigger_policy") != "v2":
        raise ValueError("final Cached DEA must use the already verified v2 trigger policy")
    if float(dea_config.get("confidence_threshold", 0)) != 0.75:
        raise ValueError("confidence_threshold must remain 0.75")
    if float(dea_config.get("low_confidence_threshold", 0)) != 0.65:
        raise ValueError("low_confidence_threshold must remain 0.65")
    return config


def set_runtime_env(config: dict[str, Any], workers: int) -> dict[str, Any]:
    os.environ["TASK2_LLM_BASE_URL"] = str(config["api_base"])
    os.environ["TASK2_LLM_MODEL"] = str(config["model_name"])
    os.environ["TASK2_LLM_WIRE_API"] = str(config.get("wire_api", "chat_completions"))
    os.environ["TASK2_LLM_ENDPOINT_STYLE"] = str(config.get("endpoint_style", "openai_v1"))
    os.environ["TASK2_LLM_TIMEOUT"] = str(config.get("timeout_seconds", 180))
    os.environ["TASK2_LLM_MAX_TOKENS"] = str(config.get("max_tokens", 512))
    os.environ["TASK2_LLM_TEMPERATURE"] = str(config.get("temperature", 0))
    os.environ["TASK2_LLM_MAX_RETRIES"] = str(config.get("max_retries", 1))
    os.environ["TASK2_LLM_FORMAT_RETRIES"] = str(config.get("format_correction_retries", 1))
    os.environ["TASK2_LLM_WORKERS"] = str(workers)
    return fulltext.llm_runtime_config()


def fetch_model_info(config: dict[str, Any]) -> dict[str, Any]:
    url = fulltext.build_endpoint_url(str(config["api_base"]), "models", endpoint_style=str(config.get("endpoint_style", "openai_v1")))
    with urllib.request.urlopen(url, timeout=15) as response:
        data = json.loads(response.read().decode("utf-8", errors="replace"))
    names = [norm(item.get("id")) for item in data.get("data", []) if isinstance(item, dict)]
    model_name = norm(config["model_name"])
    if model_name not in names:
        raise RuntimeError(f"model {model_name!r} not available; found {names}")
    return {"status": response.status, "model_names": names, "raw": data}


def nvidia_smi_snapshot() -> dict[str, Any]:
    try:
        output = subprocess.check_output(["nvidia-smi"], text=True, stderr=subprocess.STDOUT, timeout=15)
    except Exception as exc:
        return {"ok": False, "error": f"{type(exc).__name__}: {exc}"}
    return {"ok": True, "output": output}


def validate_retrieval_index(config: dict[str, Any]) -> dict[str, Any]:
    retrieval = config.get("retrieval") or {}
    source = resolve(retrieval.get("article_index", "data/processed/event_credibility/fulltext_article_index.parquet"))
    index_dir = resolve(retrieval.get("sqlite_index_dir", "cache/event_credibility/dea_retrieval_index"))
    db_path = sqlite_path(index_dir)
    if not db_path.exists():
        raise FileNotFoundError(db_path)
    meta = load_index_metadata(index_dir)
    conn = sqlite3.connect(f"file:{db_path}?mode=ro&immutable=1", uri=True)
    conn.execute("pragma query_only=ON")
    query_only = conn.execute("pragma query_only").fetchone()[0]
    doc_count = int(conn.execute("select count(*) from docs").fetchone()[0])
    fts_hit_count = len(conn.execute("select rowid from docs_fts where docs_fts match ? limit 3", ("subject_text:汇川技术",)).fetchall())
    conn.close()
    if query_only != 1 or fts_hit_count <= 0:
        raise RuntimeError("SQLite FTS5 read-only probe failed")
    return {
        "source_path": str(source),
        "sqlite_path": str(db_path),
        "metadata_source_rows": int(meta.get("source_rows") or 0),
        "docs_count": doc_count,
        "index_version": meta.get("index_version", ""),
        "query_only": bool(query_only),
        "fts5_probe_hits": fts_hit_count,
    }


def safe_evidence_items(row: dict[str, Any]) -> list[dict[str, Any]]:
    items = fulltext.normalize_llm_evidence_items(row.get("evidence_items") or [])
    sanitized: list[dict[str, Any]] = []
    for item in items:
        sanitized.append(
            {
                "evidence_id": norm(item.get("evidence_id")),
                "article_id": norm(item.get("article_id")),
                "trace_id": norm(item.get("trace_id") or item.get("article_id") or item.get("evidence_id")),
                "title": norm(item.get("title") or item.get("article_title")),
                "source": norm(item.get("source") or item.get("article_source")),
                "source_type": norm(item.get("source_type")),
                "publish_time": norm(item.get("publish_time") or item.get("article_publish_time")),
                "sentence": norm(item.get("sentence") or item.get("content") or item.get("evidence_text")),
                "sentence_score": item.get("sentence_score", 0),
                "evidence_source": norm(item.get("evidence_source") or "cluster_pool"),
            }
        )
    return sanitized


def load_input_rows(path: Path, required_count: int | None = None, limit: int | None = None, seed: int = 20260809) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    raw_rows = read_jsonl(path)
    if required_count and len(raw_rows) != required_count:
        raise ValueError(f"expected {required_count} input rows, got {len(raw_rows)}")
    ids = [norm(row.get("cluster_id")) for row in raw_rows]
    if any(not cid for cid in ids):
        raise ValueError("input contains rows without cluster_id")
    duplicates = sorted(cid for cid, count in Counter(ids).items() if count > 1)
    if duplicates:
        raise ValueError(f"duplicate cluster_id: {duplicates[:10]}")
    selected_indices = list(range(len(raw_rows)))
    if limit is not None and limit > 0 and limit < len(raw_rows):
        selected_indices = sorted(random.Random(seed).sample(selected_indices, limit))
    selected: list[dict[str, Any]] = []
    missing_claim = 0
    missing_evidence = 0
    untraceable_evidence = 0
    for output_index, original_index in enumerate(selected_indices):
        row = dict(raw_rows[original_index])
        row.pop("llm_prompt", None)
        row.pop("credibility_score", None)
        row.pop("credibility_level", None)
        row.pop("risk_reasons", None)
        row.pop("supporting_reasons", None)
        row["input_index"] = output_index
        row["source_input_index"] = original_index
        row["event_claim"] = fulltext.build_event_claim(row)
        row["event_time"] = fulltext.build_event_time_value(row)
        row["evidence_items"] = safe_evidence_items(row)
        if not norm(row.get("event_claim")):
            missing_claim += 1
        if not row["evidence_items"]:
            missing_evidence += 1
        for item in row["evidence_items"]:
            if not (
                norm(item.get("evidence_id"))
                and norm(item.get("source"))
                and norm(item.get("publish_time"))
                and (norm(item.get("title")) or norm(item.get("sentence")))
            ):
                untraceable_evidence += 1
        selected.append(row)
    if missing_claim or missing_evidence:
        raise ValueError(f"invalid input rows: missing_claim={missing_claim}, missing_evidence={missing_evidence}")
    manifest = {
        "input_path": str(path),
        "source_row_count": len(raw_rows),
        "selected_count": len(selected),
        "selected_source_indices": selected_indices,
        "cluster_ids": [row["cluster_id"] for row in selected],
        "explicit_input_index_present": all("input_index" in row for row in raw_rows),
        "input_index_policy": "assigned from selected row order; source_input_index preserves original 260-row line number",
        "claim_policy": "event_claim built deterministically from event_claim/claim_text/claim or target_name+event_name; llm_prompt ignored",
        "untraceable_evidence_record_count": untraceable_evidence,
    }
    if untraceable_evidence:
        raise ValueError(f"untraceable evidence records: {untraceable_evidence}")
    return selected, manifest


def guard_output_dir(output_dir: Path, resume: bool, dry_run: bool = False) -> None:
    if dry_run:
        return
    protected_names = {
        "input_manifest.json",
        "round1_results.jsonl",
        "dea_results.jsonl",
        "final_results.jsonl",
        "retrieval_traces.jsonl",
        "request_metrics.jsonl",
        "failed_records.jsonl",
        "checkpoint.json",
        "run_manifest.json",
    }
    if output_dir.exists() and not resume:
        existing = [item for item in output_dir.iterdir() if item.name in protected_names and item.is_file()]
        if existing:
            raise FileExistsError(f"{output_dir} already contains final run artifacts; pass --resume")
    output_dir.mkdir(parents=True, exist_ok=True)


def load_article_index(config: dict[str, Any]) -> dea.FulltextArticleIndex:
    retrieval = config.get("retrieval") or {}
    index = dea.FulltextArticleIndex(
        resolve(retrieval.get("article_index", "data/processed/event_credibility/fulltext_article_index.parquet")),
        members_path=resolve(retrieval.get("members", "outputs/event_clustering/event_cluster_members.parquet")),
    )
    index.load()
    if not index.available:
        raise RuntimeError(f"retrieval index unavailable: {index.error}")
    return index


def build_dea_config(config: dict[str, Any]) -> dea.DEAConfig:
    settings = config.get("dea") or {}
    return dea.DEAConfig(
        initial_evidence_limit=int(settings.get("initial_evidence_limit", 4)),
        additional_evidence_limit=int(settings.get("additional_evidence_limit", 4)),
        confidence_threshold=float(settings.get("confidence_threshold", 0.75)),
        low_confidence_threshold=float(settings.get("low_confidence_threshold", 0.65)),
        max_rounds=int(settings.get("max_rounds", 2)),
        trigger_policy=str(settings.get("trigger_policy", "v2")),
        compact_prompt=bool(settings.get("compact_prompt", True)),
    )


def round_tokens(round_result: dict[str, Any]) -> tuple[int, int, int]:
    prompt = int(round_result.get("effective_prompt_tokens") or 0)
    total = int(round_result.get("effective_total_tokens") or 0)
    completion = max(0, total - prompt)
    return prompt, completion, total


def standardize_result(raw: dict[str, Any]) -> dict[str, Any]:
    round1 = raw.get("round1") or {}
    round2 = raw.get("round2") or {}
    acquisition = raw.get("acquisition") or {}
    final_evidence = raw.get("final_evidence") or []
    final_evidence_ids = [norm(item.get("evidence_id")) for item in final_evidence if isinstance(item, dict) and norm(item.get("evidence_id"))]
    round1_prompt, round1_completion, round1_total = round_tokens(round1)
    round2_prompt, round2_completion, round2_total = round_tokens(round2)
    model_latency = float(round1.get("latency_seconds") or 0.0) + float(round2.get("latency_seconds") or 0.0)
    retrieval_latency = float(acquisition.get("retrieval_latency_seconds") or 0.0)
    final_result = {
        "verdict": norm(raw.get("final_verdict")),
        "confidence": float(raw.get("final_confidence") or 0.0),
        "reason": norm(raw.get("final_reason")),
        "evidence_ids": final_evidence_ids,
    }
    return {
        "input_index": int(raw.get("input_index") or 0),
        "cluster_id": norm(raw.get("cluster_id")),
        "round1_result": {
            "verdict": norm(round1.get("verdict")),
            "confidence": float(round1.get("confidence") or 0.0),
            "reason": norm(round1.get("reason")),
            "evidence_ids": round1.get("evidence_ids") or [],
            "status": norm(round1.get("status")),
            "error_type": norm(round1.get("error_type")),
        },
        "dea_triggered": bool(raw.get("dea_triggered")),
        "dea_trigger_reasons": raw.get("trigger_reasons") or [],
        "retrieval_backend": norm(raw.get("retrieval_mode") or acquisition.get("retrieval_mode") or "not_triggered"),
        "retrieval_query_count": len(acquisition.get("queries") or []),
        "retrieved_evidence_ids": list(acquisition.get("new_evidence_ids") or []),
        "added_evidence_ids": list(acquisition.get("new_evidence_ids") or []),
        "final_evidence_ids": final_evidence_ids,
        "final_result": final_result,
        "selected_result_source": norm(raw.get("final_source")),
        "logical_request_count": int(raw.get("llm_call_count") or (1 + int(bool(round2)))),
        "http_attempt_count": int(round1.get("attempt_count") or 0) + int(round2.get("attempt_count") or 0),
        "retry_count": int(round1.get("format_retry_count") or 0) + int(round2.get("format_retry_count") or 0),
        "retrieval_latency_seconds": retrieval_latency,
        "model_latency_seconds": round(model_latency, 6),
        "total_latency_seconds": float(raw.get("latency_seconds") or round(model_latency + retrieval_latency, 6)),
        "prompt_tokens": round1_prompt + round2_prompt,
        "completion_tokens": round1_completion + round2_completion,
        "total_tokens": round1_total + round2_total,
        "status": norm(raw.get("status")),
        "error_type": norm(raw.get("error_type")),
    }


def traceable_evidence(item: dict[str, Any]) -> bool:
    return bool(
        norm(item.get("evidence_id"))
        and norm(item.get("trace_id") or item.get("article_id"))
        and norm(item.get("source") or item.get("article_source"))
        and norm(item.get("publish_time") or item.get("article_publish_time"))
        and (norm(item.get("title") or item.get("article_title")) or norm(item.get("sentence") or item.get("content")))
    )


def validate_final_outputs(rows: list[dict[str, Any]], input_rows: list[dict[str, Any]], raw_by_cluster: dict[str, dict[str, Any]]) -> dict[str, Any]:
    expected_ids = [norm(row.get("cluster_id")) for row in input_rows]
    actual_ids = [norm(row.get("cluster_id")) for row in rows]
    duplicate_success = len(actual_ids) - len(set(actual_ids))
    invalid_verdict = 0
    invalid_confidence = 0
    illegal_evidence = 0
    untraceable = 0
    cross_cluster = 0
    final_failed = 0
    oom = 0
    for row in rows:
        verdict = norm((row.get("final_result") or {}).get("verdict"))
        if verdict not in VALID_VERDICTS:
            invalid_verdict += 1
        confidence = (row.get("final_result") or {}).get("confidence")
        if not isinstance(confidence, (int, float)) or not math.isfinite(float(confidence)) or not 0.0 <= float(confidence) <= 1.0:
            invalid_confidence += 1
        raw = raw_by_cluster.get(norm(row.get("cluster_id")), {})
        allowed = {norm(item.get("evidence_id")) for item in raw.get("final_evidence", []) if isinstance(item, dict)}
        for eid in row.get("final_evidence_ids") or []:
            if norm(eid) not in allowed:
                illegal_evidence += 1
        for item in raw.get("final_evidence", []) or []:
            if isinstance(item, dict) and not traceable_evidence(item):
                untraceable += 1
            embedded_cluster = norm(item.get("cluster_id")) if isinstance(item, dict) else ""
            if embedded_cluster and embedded_cluster != norm(row.get("cluster_id")):
                cross_cluster += 1
        if row.get("status") != "success":
            final_failed += 1
        if "oom" in json.dumps(raw, ensure_ascii=False).lower() or "out of memory" in json.dumps(raw, ensure_ascii=False).lower():
            oom += 1
    return {
        "input_count": len(input_rows),
        "final_result_count": len(rows),
        "cluster_id_order_match": actual_ids == expected_ids,
        "missing_cluster_ids": sorted(set(expected_ids) - set(actual_ids)),
        "extra_cluster_ids": sorted(set(actual_ids) - set(expected_ids)),
        "duplicate_success_record_count": duplicate_success,
        "input_index_continuous": [row.get("input_index") for row in rows] == list(range(len(rows))),
        "invalid_verdict_count": invalid_verdict,
        "invalid_confidence_count": invalid_confidence,
        "illegal_evidence_id_count": illegal_evidence,
        "cross_cluster_reference_count": cross_cluster,
        "untraceable_evidence_count": untraceable,
        "final_failed_count": final_failed,
        "cuda_oom_count": oom,
    }


def pctl(values: list[float], percentile: float) -> float:
    clean = sorted(float(value) for value in values if math.isfinite(float(value)))
    if not clean:
        return 0.0
    position = (len(clean) - 1) * percentile
    low = int(math.floor(position))
    high = int(math.ceil(position))
    if low == high:
        return round(clean[low], 6)
    return round(clean[low] + (clean[high] - clean[low]) * (position - low), 6)


def distribution(values: list[str]) -> dict[str, int]:
    return dict(Counter(values))


def summarize_run(final_rows: list[dict[str, Any]], raw_rows: list[dict[str, Any]], validation: dict[str, Any], elapsed: float, resume_used: bool) -> dict[str, Any]:
    round1_verdicts = [norm(row.get("round1_result", {}).get("verdict")) for row in final_rows]
    final_verdicts = [norm(row.get("final_result", {}).get("verdict")) for row in final_rows]
    trigger_reasons = Counter(reason for row in final_rows for reason in row.get("dea_trigger_reasons", []))
    added_counts = [len(row.get("added_evidence_ids") or []) for row in final_rows]
    new_source_count = sum(int((raw.get("acquisition") or {}).get("new_source_count") or 0) for raw in raw_rows)
    changed = sum(r1 != final for r1, final in zip(round1_verdicts, final_verdicts))
    return {
        "completed_count": len(final_rows),
        "round1_verdict_distribution": distribution(round1_verdicts),
        "final_verdict_distribution": distribution(final_verdicts),
        "dea_trigger_count": sum(1 for row in final_rows if row.get("dea_triggered")),
        "dea_trigger_ratio": round(sum(1 for row in final_rows if row.get("dea_triggered")) / max(1, len(final_rows)), 6),
        "dea_trigger_reason_counts": dict(trigger_reasons),
        "round1_direct_output_count": sum(1 for row in final_rows if not row.get("dea_triggered")),
        "dea_result_count": sum(1 for row in final_rows if row.get("dea_triggered")),
        "round1_final_verdict_change_count": changed,
        "insufficient_to_support_count": sum(1 for r1, final in zip(round1_verdicts, final_verdicts) if r1 == "insufficient" and final == "support"),
        "insufficient_to_contradict_count": sum(1 for r1, final in zip(round1_verdicts, final_verdicts) if r1 == "insufficient" and final == "contradict"),
        "insufficient_kept_count": sum(1 for r1, final in zip(round1_verdicts, final_verdicts) if r1 == "insufficient" and final == "insufficient"),
        "support_or_contradict_to_insufficient_count": sum(1 for r1, final in zip(round1_verdicts, final_verdicts) if r1 in {"support", "contradict"} and final == "insufficient"),
        "new_evidence_total_count": sum(added_counts),
        "samples_with_new_evidence_count": sum(1 for count in added_counts if count > 0),
        "new_independent_source_count_sum": new_source_count,
        "latency": {
            "retrieval_avg": round(sum(row["retrieval_latency_seconds"] for row in final_rows) / max(1, len(final_rows)), 6),
            "retrieval_p50": pctl([row["retrieval_latency_seconds"] for row in final_rows], 0.50),
            "retrieval_p95": pctl([row["retrieval_latency_seconds"] for row in final_rows], 0.95),
            "model_avg": round(sum(row["model_latency_seconds"] for row in final_rows) / max(1, len(final_rows)), 6),
            "model_p50": pctl([row["model_latency_seconds"] for row in final_rows], 0.50),
            "model_p95": pctl([row["model_latency_seconds"] for row in final_rows], 0.95),
            "total_avg": round(sum(row["total_latency_seconds"] for row in final_rows) / max(1, len(final_rows)), 6),
            "total_p50": pctl([row["total_latency_seconds"] for row in final_rows], 0.50),
            "total_p95": pctl([row["total_latency_seconds"] for row in final_rows], 0.95),
        },
        "request_cost": {
            "logical_request_avg": round(sum(row["logical_request_count"] for row in final_rows) / max(1, len(final_rows)), 6),
            "http_attempt_avg": round(sum(row["http_attempt_count"] for row in final_rows) / max(1, len(final_rows)), 6),
            "prompt_tokens_total": sum(row["prompt_tokens"] for row in final_rows),
            "completion_tokens_total": sum(row["completion_tokens"] for row in final_rows),
            "total_tokens_total": sum(row["total_tokens"] for row in final_rows),
            "prompt_tokens_avg": round(sum(row["prompt_tokens"] for row in final_rows) / max(1, len(final_rows)), 6),
            "completion_tokens_avg": round(sum(row["completion_tokens"] for row in final_rows) / max(1, len(final_rows)), 6),
            "total_tokens_avg": round(sum(row["total_tokens"] for row in final_rows) / max(1, len(final_rows)), 6),
            "format_correction_count": sum(row["retry_count"] for row in final_rows),
        },
        "errors": {
            "timeout_count": sum("timeout" in row.get("error_type", "").lower() for row in final_rows),
            "http_5xx_count": sum("5xx" in row.get("error_type", "").lower() for row in final_rows),
            "parse_failure_count": sum("parse" in row.get("error_type", "").lower() for row in final_rows),
            **validation,
        },
        "checkpoint_resume": {"resume_used": bool(resume_used)},
        "wall_clock_seconds": round(elapsed, 3),
    }


def write_derived_outputs(output_dir: Path, raw_rows: list[dict[str, Any]], input_rows: list[dict[str, Any]], elapsed: float, resume_used: bool) -> tuple[list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    by_cluster: dict[str, dict[str, Any]] = {}
    duplicate_success: list[str] = []
    for raw in raw_rows:
        cid = norm(raw.get("cluster_id"))
        if raw.get("status") == "success" and cid in by_cluster and by_cluster[cid].get("status") == "success":
            duplicate_success.append(cid)
        by_cluster[cid] = raw
    final_rows = [standardize_result(by_cluster[norm(row.get("cluster_id"))]) for row in input_rows if norm(row.get("cluster_id")) in by_cluster]
    final_rows.sort(key=lambda row: row["input_index"])
    validation = validate_final_outputs(final_rows, input_rows, by_cluster)
    validation["duplicate_success_record_count"] += len(duplicate_success)
    raw_sorted = [by_cluster[norm(row.get("cluster_id"))] for row in input_rows if norm(row.get("cluster_id")) in by_cluster]
    raw_sorted.sort(key=lambda row: int(row.get("input_index") or 0))
    round1_rows = [
        {
            "input_index": row["input_index"],
            "cluster_id": row["cluster_id"],
            **row["round1_result"],
        }
        for row in final_rows
    ]
    retrieval_rows = [
        {
            "input_index": row["input_index"],
            "cluster_id": row["cluster_id"],
            "retrieval_backend": row["retrieval_backend"],
            "retrieval_query_count": row["retrieval_query_count"],
            "retrieved_evidence_ids": row["retrieved_evidence_ids"],
            "added_evidence_ids": row["added_evidence_ids"],
            "retrieval_latency_seconds": row["retrieval_latency_seconds"],
            "queries": (by_cluster[row["cluster_id"]].get("acquisition") or {}).get("queries", []),
        }
        for row in final_rows
    ]
    request_rows = [
        {
            "input_index": row["input_index"],
            "cluster_id": row["cluster_id"],
            "logical_request_count": row["logical_request_count"],
            "http_attempt_count": row["http_attempt_count"],
            "retry_count": row["retry_count"],
            "retrieval_latency_seconds": row["retrieval_latency_seconds"],
            "model_latency_seconds": row["model_latency_seconds"],
            "total_latency_seconds": row["total_latency_seconds"],
            "prompt_tokens": row["prompt_tokens"],
            "completion_tokens": row["completion_tokens"],
            "total_tokens": row["total_tokens"],
            "status": row["status"],
            "error_type": row["error_type"],
        }
        for row in final_rows
    ]
    failed_rows = [row for row in final_rows if row.get("status") != "success"]
    write_jsonl(output_dir / "round1_results.jsonl", round1_rows)
    write_jsonl(output_dir / "final_results.jsonl", final_rows)
    write_jsonl(output_dir / "retrieval_traces.jsonl", retrieval_rows)
    write_jsonl(output_dir / "request_metrics.jsonl", request_rows)
    write_jsonl(output_dir / "failed_records.jsonl", failed_rows)
    # Keep DEA rows in canonical input order once all completed or on resume finalization.
    write_jsonl(output_dir / "dea_results.jsonl", raw_sorted)
    checkpoint_payload = {
        "checkpoint_time": datetime.now().isoformat(timespec="seconds"),
        "output_dir": str(output_dir),
        "completed_cluster_ids": [row["cluster_id"] for row in final_rows],
        "success_cluster_ids": [row["cluster_id"] for row in final_rows if row.get("status") == "success"],
        "failed_cluster_ids": [row["cluster_id"] for row in failed_rows],
        "resume_command": f"python scripts/run_task2_final_cached_dea_v1.py --output-dir {output_dir} --resume",
    }
    write_json(output_dir / "checkpoint.json", checkpoint_payload)
    summary = summarize_run(final_rows, raw_sorted, validation, elapsed, resume_used=resume_used)
    return final_rows, validation, summary


def compare_with_old_formal(final_rows: list[dict[str, Any]], old_path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    old_rows = {norm(row.get("cluster_id")): row for row in read_jsonl(old_path)}
    changes: list[dict[str, Any]] = []
    matrix: dict[str, Counter[str]] = {}
    matched = 0
    evidence_changed = 0
    confidence_abs_diffs: list[float] = []
    for row in final_rows:
        cid = row["cluster_id"]
        old = old_rows.get(cid)
        if not old:
            continue
        matched += 1
        old_verdict = norm(old.get("verdict"))
        new_verdict = norm(row.get("final_result", {}).get("verdict"))
        matrix.setdefault(old_verdict, Counter())[new_verdict] += 1
        old_ids = set(old.get("supporting_evidence_ids") or []) | set(old.get("contradicting_evidence_ids") or []) | set(old.get("evidence_ids") or [])
        new_ids = set(row.get("final_evidence_ids") or [])
        if old_ids != new_ids:
            evidence_changed += 1
        try:
            confidence_abs_diffs.append(abs(float(old.get("confidence") or 0.0) - float(row.get("final_result", {}).get("confidence") or 0.0)))
        except Exception:
            pass
        if old_verdict != new_verdict or old_ids != new_ids:
            changes.append(
                {
                    "cluster_id": cid,
                    "old_verdict": old_verdict,
                    "new_verdict": new_verdict,
                    "old_confidence": old.get("confidence"),
                    "new_confidence": row.get("final_result", {}).get("confidence"),
                    "old_evidence_ids": sorted(old_ids),
                    "new_evidence_ids": sorted(new_ids),
                    "dea_triggered": row.get("dea_triggered"),
                    "selected_result_source": row.get("selected_result_source"),
                }
            )
    summary = {
        "old_path": str(old_path),
        "matched_count": matched,
        "missing_in_new_count": len(set(old_rows) - {row["cluster_id"] for row in final_rows}),
        "new_not_in_old_count": len({row["cluster_id"] for row in final_rows} - set(old_rows)),
        "verdict_consistency_rate": round(sum(1 for row in final_rows if norm(old_rows.get(row["cluster_id"], {}).get("verdict")) == norm(row.get("final_result", {}).get("verdict"))) / max(1, matched), 6),
        "verdict_change_count": sum(1 for row in final_rows if row["cluster_id"] in old_rows and norm(old_rows[row["cluster_id"]].get("verdict")) != norm(row.get("final_result", {}).get("verdict"))),
        "confidence_mean_abs_diff": round(sum(confidence_abs_diffs) / max(1, len(confidence_abs_diffs)), 6),
        "evidence_ids_change_count": evidence_changed,
        "new_evidence_total_count": sum(len(row.get("added_evidence_ids") or []) for row in final_rows),
        "old_success_new_failure_count": sum(1 for row in final_rows if row["cluster_id"] in old_rows and row.get("status") != "success"),
        "verdict_transition_matrix": {old: dict(counter) for old, counter in matrix.items()},
        "changes_count": len(changes),
    }
    return summary, changes


def write_comparison_reports(summary: dict[str, Any], changes: list[dict[str, Any]], json_path: Path, md_path: Path, xlsx_path: Path) -> None:
    write_json(json_path, {"summary": summary, "changes": changes})
    lines = [
        "# Task2 final Cached DEA v1 260 comparison",
        "",
        "只读比较新候选输出与旧正式 `llm_review_results.jsonl`，未覆盖旧正式输出。",
        "",
        f"- 匹配数量：{summary['matched_count']}",
        f"- verdict一致率：{summary['verdict_consistency_rate']:.4f}",
        f"- verdict变化数量：{summary['verdict_change_count']}",
        f"- evidence_ids变化数量：{summary['evidence_ids_change_count']}",
        f"- confidence平均绝对差：{summary['confidence_mean_abs_diff']}",
        "",
    ]
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "变化样本"
    headers = list(changes[0]) if changes else ["暂无变化"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in changes:
        sheet.append([json.dumps(row.get(header), ensure_ascii=False) if isinstance(row.get(header), (dict, list)) else row.get(header, "") for header in headers])
    for idx, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(idx)].width = min(60, max(12, len(str(header)) + 2))
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(xlsx_path)


def write_run_report(report: dict[str, Any], path_json: Path, path_md: Path) -> None:
    write_json(path_json, report)
    summary = report["summary"]
    lines = [
        "# Step 4.5-B final Cached DEA v1 260 report",
        "",
        "本轮启用32B单轮严格盲审 + 条件触发Cached DEA；禁用Safety Gate、ISR、Claim Contract和可信分融合。",
        "",
        f"- 完成数量：{summary['completed_count']}",
        f"- DEA触发：{summary['dea_trigger_count']} ({summary['dea_trigger_ratio']:.2%})",
        f"- Round1分布：{summary['round1_verdict_distribution']}",
        f"- Final分布：{summary['final_verdict_distribution']}",
        f"- verdict变化：{summary['round1_final_verdict_change_count']}",
        f"- 新增证据：{summary['new_evidence_total_count']}，使用新增证据样本：{summary['samples_with_new_evidence_count']}",
        f"- 逻辑请求均值：{summary['request_cost']['logical_request_avg']}",
        f"- HTTP尝试均值：{summary['request_cost']['http_attempt_avg']}",
        f"- 总tokens：{summary['request_cost']['total_tokens_total']}",
        f"- 平均/P50/P95端到端耗时：{summary['latency']['total_avg']} / {summary['latency']['total_p50']} / {summary['latency']['total_p95']}",
        f"- 错误统计：{summary['errors']}",
        "",
        "正式输出未覆盖。",
    ]
    path_md.parent.mkdir(parents=True, exist_ok=True)
    path_md.write_text("\n".join(lines) + "\n", encoding="utf-8")


def output_sha256(output_dir: Path) -> dict[str, str]:
    files = [
        "input_manifest.json",
        "round1_results.jsonl",
        "dea_results.jsonl",
        "final_results.jsonl",
        "retrieval_traces.jsonl",
        "request_metrics.jsonl",
        "failed_records.jsonl",
        "checkpoint.json",
    ]
    return {name: sha256(output_dir / name) for name in files if (output_dir / name).exists()}


def run(args: argparse.Namespace) -> dict[str, Any]:
    config_path = resolve(args.config)
    config = load_config(config_path)
    workers = int(args.workers or config.get("workers", 4))
    output_dir = resolve(args.output_dir)
    input_path = resolve(args.input)
    report_json = resolve(args.report_json)
    report_md = resolve(args.report_md)
    comparison_json = resolve(args.comparison_json)
    comparison_md = resolve(args.comparison_md)
    comparison_xlsx = resolve(args.comparison_xlsx)
    required_count = int((config.get("input") or {}).get("required_count") or 260)
    canary_seed = int((config.get("input") or {}).get("canary_seed") or 20260809)
    limit = args.limit
    if args.canary and limit is None:
        limit = int((config.get("input") or {}).get("canary_count") or 5)
    model_info = fetch_model_info(config)
    retrieval_info = validate_retrieval_index(config)
    rows, input_manifest = load_input_rows(input_path, required_count=required_count, limit=limit, seed=canary_seed)
    if args.dry_run:
        return {
            "dry_run": True,
            "config": str(config_path),
            "workers": workers,
            "input_manifest": input_manifest,
            "model_info": model_info,
            "retrieval_info": retrieval_info,
        }
    guard_output_dir(output_dir, resume=bool(args.resume), dry_run=False)
    for path in [report_json, report_md, comparison_json, comparison_md, comparison_xlsx]:
        if path.exists() and not args.resume and path.name in REPORT_FILENAMES:
            raise FileExistsError(f"{path} exists; pass --resume to regenerate")
    formal_sha_before = hash_paths(FORMAL_OUTPUTS)
    run_started = time.perf_counter()
    write_json(output_dir / "input_manifest.json", input_manifest)
    runtime_config = set_runtime_env(config, workers)
    article_index = load_article_index(config)
    dea_config = build_dea_config(config)
    raw_dea_path = output_dir / "dea_results.jsonl"
    dea_summary = dea.run_dea_batch(
        rows,
        raw_dea_path,
        runtime_config,
        dea_config,
        mode="dea",
        workers=workers,
        resume=bool(args.resume),
        checkpoint_every=int(config.get("checkpoint_interval") or 10),
        article_index=article_index,
    )
    raw_rows = read_jsonl(raw_dea_path) if raw_dea_path.exists() else []
    final_rows, validation, summary = write_derived_outputs(output_dir, raw_rows, rows, time.perf_counter() - run_started, resume_used=bool(args.resume))
    formal_sha_after = hash_paths(FORMAL_OUTPUTS)
    if formal_sha_before != formal_sha_after:
        raise RuntimeError("formal outputs changed during final Cached DEA run")
    comparison_summary: dict[str, Any] = {}
    changes: list[dict[str, Any]] = []
    if not args.canary and len(rows) == required_count:
        comparison_summary, changes = compare_with_old_formal(final_rows, FORMAL_OUTPUTS[0])
        write_comparison_reports(comparison_summary, changes, comparison_json, comparison_md, comparison_xlsx)
    config_hash = sha256(config_path)
    run_manifest = {
        "step": "4.5-B",
        "status": "completed" if validation["final_failed_count"] == 0 else "completed_with_failures",
        "canary": bool(args.canary),
        "config_path": str(config_path),
        "config_sha256": config_hash,
        "input_path": str(input_path),
        "output_dir": str(output_dir),
        "workers": workers,
        "model_info": model_info,
        "gpu_snapshot": nvidia_smi_snapshot(),
        "retrieval_info": retrieval_info,
        "dea_summary": dea_summary,
        "summary": summary,
        "comparison_summary": comparison_summary,
        "formal_sha256_before": formal_sha_before,
        "formal_sha256_after": formal_sha_after,
        "formal_outputs_unchanged": formal_sha_before == formal_sha_after,
        "output_sha256": output_sha256(output_dir),
        "model_call_count": sum(row["logical_request_count"] for row in final_rows),
        "http_attempt_count": sum(row["http_attempt_count"] for row in final_rows),
        "score_fusion_enabled": False,
        "safety_gate_enabled": False,
        "isr_enabled": False,
        "claim_contract_enabled": False,
        "completed_at": datetime.now().isoformat(timespec="seconds"),
    }
    write_json(output_dir / "run_manifest.json", run_manifest)
    report = {
        "step": "4.5-B",
        "config_sha256": config_hash,
        "output_dir": str(output_dir),
        "input_manifest": input_manifest,
        "model_info": model_info,
        "retrieval_info": retrieval_info,
        "summary": summary,
        "validation": validation,
        "comparison_summary": comparison_summary,
        "formal_outputs_unchanged": formal_sha_before == formal_sha_after,
        "formal_sha256_before": formal_sha_before,
        "formal_sha256_after": formal_sha_after,
        "model_call_count": run_manifest["model_call_count"],
        "http_attempt_count": run_manifest["http_attempt_count"],
        "disabled_components": {
            "safety_gate": True,
            "isr": True,
            "claim_contract": True,
            "score_fusion": True,
        },
    }
    if not args.canary:
        write_run_report(report, report_json, report_md)
    return report


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--workers", type=int, default=0)
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--resume", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--canary", action="store_true")
    parser.add_argument("--report-json", type=Path, default=DEFAULT_REPORT_JSON)
    parser.add_argument("--report-md", type=Path, default=DEFAULT_REPORT_MD)
    parser.add_argument("--comparison-json", type=Path, default=DEFAULT_COMPARISON_JSON)
    parser.add_argument("--comparison-md", type=Path, default=DEFAULT_COMPARISON_MD)
    parser.add_argument("--comparison-xlsx", type=Path, default=DEFAULT_COMPARISON_XLSX)
    return parser


def main(argv: list[str] | None = None) -> int:
    report = run(build_parser().parse_args(argv))
    printable = {
        "dry_run": report.get("dry_run", False),
        "output_dir": report.get("output_dir", ""),
        "completed_count": report.get("summary", {}).get("completed_count"),
        "dea_trigger_count": report.get("summary", {}).get("dea_trigger_count"),
        "model_call_count": report.get("model_call_count"),
        "formal_outputs_unchanged": report.get("formal_outputs_unchanged"),
    }
    print(json.dumps(printable, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
