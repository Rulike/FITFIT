#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-J structure expansion confirmation and freeze package."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
INPUT_XLSX = ROOT / "reports/task1_step5_7i_structure_expansion_ai_audited.xlsx"
CONFIRMED_XLSX = ROOT / "reports/task1_step5_7j_structure_expansion_human_confirmed.xlsx"
OUT_DIR = ROOT / "outputs/task1_step5_7j_structure_confirmation"
REPORT_JSON = ROOT / "reports/task1_step5_7j_structure_confirmation_report.json"
REPORT_MD = ROOT / "reports/task1_step5_7j_structure_confirmation_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7j_structure_confirmation_details.xlsx"
PLAN_MD = ROOT / "reports/task1_step5_7j_structure_training_plan.md"

STATUS_READY = "structure_training_dataset_ready"
STATUS_FAILED = "structure_annotation_confirmation_failed"
DIRECTIONS = {"positive", "negative", "neutral"}
ARG_ROLES = {"subject_company", "event_time", "object_target", "trigger"}
NOTE_APPEND = "用户已逐条审核并确认AI辅助结构标注无误。"

HUMAN_FIELDS = [
    "human_trigger",
    "human_event_span",
    "human_subject_company",
    "human_event_time",
    "human_event_time_normalized",
    "human_object_target",
    "human_direction",
    "human_arguments_json",
    "human_argument_roles_json",
]


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_sha_txt(path: Path, snapshot: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(f"{digest}  {name}" for name, digest in sorted(snapshot.items())) + "\n", encoding="utf-8")


def official_event_names() -> set[str]:
    return {str(row.get("event_name", "")).strip() for row in read_jsonl(ROOT / "data/processed/event_schema.jsonl") if str(row.get("event_name", "")).strip()}


def article_ids(path: Path) -> set[str]:
    ids: set[str] = set()
    for row in read_jsonl(path):
        aid = row.get("article_id") or row.get("article_file_id")
        if aid:
            ids.add(str(aid))
    return ids


def protected_paths(extra_paths: list[Path] | None = None) -> list[Path]:
    paths = [
        ROOT / "outputs/final_task1/final_events_train.jsonl",
        ROOT / "outputs/final_task1/final_events_test.jsonl",
        ROOT / "outputs/final_task1/final_events_train.xlsx",
        ROOT / "outputs/final_task1/final_events_test.xlsx",
        ROOT / "outputs/final_task1/task1_completion_certificate.json",
        ROOT / "outputs/event_credibility/llm_review_results.jsonl",
        ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
        ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
        ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
        ROOT / "reports/task1_step5_7e_eventmention_human_confirmed.xlsx",
        ROOT / "outputs/task1_step5_7f_structure_training_v1/confirmed_eventmentions_200.jsonl",
        ROOT / "outputs/task1_step5_7f_structure_training_v1/train_eventmentions.jsonl",
        ROOT / "outputs/task1_step5_7f_structure_training_v1/dev_eventmentions.jsonl",
        ROOT / "outputs/task1_step5_7f_structure_training_v1/legacy80_diagnostic_eventmentions.jsonl",
        ROOT / "outputs/task1_step5_7f_structure_training_v1/split_manifest.json",
        ROOT / "outputs/task1_step5_7i_structure_expansion/sampling_manifest.json",
        ROOT / "outputs/task1_step5_7i_structure_expansion/split_manifest.json",
        ROOT / "outputs/task1_step5_7i_structure_expansion/annotation_drafts_300.jsonl",
        ROOT / "outputs/task1_step5_7i_structure_expansion/expansion_train_240_source.jsonl",
        ROOT / "outputs/task1_step5_7i_structure_expansion/independent_frozen_test_60_source.jsonl",
        ROOT / "reports/task1_step5_7i_structure_expansion_report.json",
        ROOT / "reports/task1_step5_7i_structure_expansion_ai_audited.xlsx",
        ROOT / "configs/task1_candidate_v2.yaml",
        ROOT / "outputs/task1_step5_5_candidate_v2_dev/manifest.json",
        ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
    ]
    if extra_paths:
        paths.extend(extra_paths)
    return sorted({path for path in paths if path.exists()}, key=rel)


def sha_snapshot(paths: list[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in paths if path.exists()}


def process_check() -> list[str]:
    proc = subprocess.run(
        ["bash", "-lc", "ps -eo pid,ppid,stat,etime,cmd | rg -i 'step5_7j|structure_confirmation|vllm|DianJin-R1-32B|8000' || true"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )
    skip = ["rg -i", "build_task1_step5_7j_structure_confirmation.py"]
    return [line for line in proc.stdout.splitlines() if not any(token in line for token in skip)]


def parse_json_cell(value: Any, field: str, errors: list[str], row_id: str) -> Any:
    if value is None or str(value).strip() == "":
        errors.append(f"{row_id}:{field}:empty")
        return None
    try:
        return json.loads(str(value))
    except Exception as exc:
        errors.append(f"{row_id}:{field}:json_parse_failed:{type(exc).__name__}")
        return None


def valid_span(content: str, span: Any, field: str, row_id: str, errors: list[str], allow_empty: bool = True) -> dict[str, Any]:
    if not isinstance(span, dict):
        errors.append(f"{row_id}:{field}:not_object")
        return {"text": "", "start": -1, "end": -1}
    text = str(span.get("text", "") or "")
    try:
        start = int(span.get("start", -1))
        end = int(span.get("end", -1))
    except Exception:
        errors.append(f"{row_id}:{field}:offset_not_int")
        return {"text": text, "start": -1, "end": -1}
    out = {"text": text, "start": start, "end": end}
    if "normalized" in span:
        out["normalized"] = str(span.get("normalized", "") or "")
    if not text:
        if not allow_empty:
            errors.append(f"{row_id}:{field}:empty_not_allowed")
        if start != -1 or end != -1:
            errors.append(f"{row_id}:{field}:empty_offset_not_minus1")
        return out
    if not (0 <= start <= end <= len(content)) or content[start:end] != text:
        errors.append(f"{row_id}:{field}:offset_text_mismatch")
    return out


def extract_rows_from_workbook(path: Path) -> tuple[list[dict[str, Any]], list[str], dict[str, Any]]:
    wb = load_workbook(path, read_only=False, data_only=False)
    if "300条主审核表" not in wb.sheetnames:
        return [], ["missing_sheet:300条主审核表"], {}
    ws = wb["300条主审核表"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {str(h): i + 1 for i, h in enumerate(headers) if h}
    required = [
        "event_mention_id",
        "article_id",
        "split",
        "event_name",
        "title",
        "content",
        "publish_time",
        "target_name",
        "source",
        *HUMAN_FIELDS,
        "reviewer_confirmed",
        "reviewer_note",
    ]
    missing = [field for field in required if field not in idx]
    if missing:
        return [], [f"missing_columns:{'|'.join(missing)}"], {"headers": headers}
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for r in range(2, ws.max_row + 1):
        eid = ws.cell(r, idx["event_mention_id"]).value
        if eid in (None, ""):
            continue
        row_id = str(eid)
        base = {field: ws.cell(r, idx[field]).value for field in required if field in idx}
        content = str(base.get("content") or "")
        trigger = valid_span(content, parse_json_cell(base.get("human_trigger"), "human_trigger", errors, row_id), "human_trigger", row_id, errors, allow_empty=False)
        event_span = valid_span(content, parse_json_cell(base.get("human_event_span"), "human_event_span", errors, row_id), "human_event_span", row_id, errors, allow_empty=False)
        subject = valid_span(content, parse_json_cell(base.get("human_subject_company"), "human_subject_company", errors, row_id), "human_subject_company", row_id, errors, allow_empty=True)
        event_time = valid_span(content, parse_json_cell(base.get("human_event_time"), "human_event_time", errors, row_id), "human_event_time", row_id, errors, allow_empty=True)
        event_time["normalized"] = str(base.get("human_event_time_normalized") or event_time.get("normalized") or "未明确")
        obj = valid_span(content, parse_json_cell(base.get("human_object_target"), "human_object_target", errors, row_id), "human_object_target", row_id, errors, allow_empty=True)
        direction = str(base.get("human_direction") or "").strip()
        if direction not in DIRECTIONS:
            errors.append(f"{row_id}:human_direction:illegal:{direction}")
        args = parse_json_cell(base.get("human_arguments_json"), "human_arguments_json", errors, row_id)
        roles = parse_json_cell(base.get("human_argument_roles_json"), "human_argument_roles_json", errors, row_id)
        if not isinstance(args, list):
            errors.append(f"{row_id}:human_arguments_json:not_list")
            args = []
        if not isinstance(roles, list):
            errors.append(f"{row_id}:human_argument_roles_json:not_list")
            roles = []
        clean_args = []
        for i, arg in enumerate(args):
            if not isinstance(arg, dict):
                errors.append(f"{row_id}:argument_{i}:not_object")
                continue
            role = str(arg.get("role", "") or "")
            if role not in ARG_ROLES:
                errors.append(f"{row_id}:argument_{i}:illegal_role:{role}")
            span = valid_span(content, arg, f"argument_{i}", row_id, errors, allow_empty=False)
            clean_args.append({"text": span["text"], "role": role, "start": span["start"], "end": span["end"]})
        for i, arg in enumerate(roles):
            if not isinstance(arg, dict):
                errors.append(f"{row_id}:argument_role_{i}:not_object")
                continue
            role = str(arg.get("role", "") or "")
            if role not in ARG_ROLES:
                errors.append(f"{row_id}:argument_role_{i}:illegal_role:{role}")
            valid_span(content, arg, f"argument_role_{i}", row_id, errors, allow_empty=False)
        if trigger.get("text") and event_span.get("text"):
            if not (event_span["start"] <= trigger["start"] < trigger["end"] <= event_span["end"]):
                errors.append(f"{row_id}:trigger_not_inside_event_span")
        note = str(base.get("reviewer_note") or "")
        rows.append(
            {
                "event_mention_id": row_id,
                "article_id": str(base.get("article_id") or ""),
                "split": str(base.get("split") or ""),
                "event_name": str(base.get("event_name") or ""),
                "title": str(base.get("title") or ""),
                "content": content,
                "publish_time": str(base.get("publish_time") or ""),
                "target_name": str(base.get("target_name") or ""),
                "source": str(base.get("source") or ""),
                "trigger": trigger,
                "event_span": event_span,
                "subject_company": subject,
                "event_time": event_time,
                "object_target": obj,
                "direction": direction,
                "arguments": clean_args,
                "argument_roles": roles,
                "annotation_source": "ai_assisted_human_confirmed",
                "annotation_status": "confirmed",
                "reviewer_confirmed": "yes",
                "reviewer_note": note + (" " if note and NOTE_APPEND not in note else "") + (NOTE_APPEND if NOTE_APPEND not in note else ""),
                "confirmed_at": "",
                "source_workbook": rel(path),
                "label_description": "AI辅助、人工确认标注；不是独立人工金标。",
            }
        )
    return rows, sorted(set(errors)), {"headers": headers, "sheet_rows": ws.max_row, "sheet_cols": ws.max_column}


def update_confirmed_workbook(input_path: Path, output_path: Path, confirmed_at: str) -> None:
    if output_path.exists():
        raise FileExistsError(output_path)
    shutil.copy2(input_path, output_path)
    wb = load_workbook(output_path)
    for sheet in ["300条主审核表", "expansion_train_240", "independent_frozen_test_60"]:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        idx = {str(h): i + 1 for i, h in enumerate(headers) if h}
        for field in ["annotation_source", "confirmed_at"]:
            if field not in idx:
                col = ws.max_column + 1
                ws.cell(1, col, field)
                idx[field] = col
        for r in range(2, ws.max_row + 1):
            if "event_mention_id" in idx and not ws.cell(r, idx["event_mention_id"]).value:
                continue
            if "reviewer_confirmed" in idx:
                ws.cell(r, idx["reviewer_confirmed"]).value = "yes"
            ws.cell(r, idx["annotation_source"]).value = "ai_assisted_human_confirmed"
            ws.cell(r, idx["confirmed_at"]).value = confirmed_at
            if "reviewer_note" in idx:
                old = str(ws.cell(r, idx["reviewer_note"]).value or "")
                ws.cell(r, idx["reviewer_note"]).value = old + (" " if old and NOTE_APPEND not in old else "") + (NOTE_APPEND if NOTE_APPEND not in old else "")
    wb.save(output_path)


def validate_confirmed_rows(rows: list[dict[str, Any]], official_names: set[str]) -> dict[str, Any]:
    errors: list[str] = []
    if len(rows) != 300:
        errors.append(f"row_count:{len(rows)}")
    ids = [row["event_mention_id"] for row in rows]
    if len(ids) != len(set(ids)):
        errors.append("duplicate_event_mention_id")
    if any(not row["article_id"] for row in rows):
        errors.append("blank_article_id")
    split_counts = Counter(row["split"] for row in rows)
    if split_counts.get("expansion_train", 0) != 240 or split_counts.get("independent_frozen_test", 0) != 60:
        errors.append(f"bad_split_counts:{dict(split_counts)}")
    train_articles = {row["article_id"] for row in rows if row["split"] == "expansion_train"}
    test_articles = {row["article_id"] for row in rows if row["split"] == "independent_frozen_test"}
    if train_articles & test_articles:
        errors.append("train_test_article_overlap")
    invalid_names = sorted({row["event_name"] for row in rows if row["event_name"] not in official_names})
    if invalid_names:
        errors.append("event_name_outside_64:" + "|".join(invalid_names))
    if any(row.get("reviewer_confirmed") != "yes" for row in rows):
        errors.append("reviewer_confirmed_not_all_yes")
    field_coverage = {}
    for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
        nonempty = sum(1 for row in rows if row[field].get("text"))
        field_coverage[field] = {"nonempty": nonempty, "coverage": nonempty / max(1, len(rows))}
    return {"valid": not errors, "errors": sorted(set(errors)), "split_counts": dict(split_counts), "field_coverage": field_coverage}


def intersections(rows: list[dict[str, Any]]) -> dict[str, Any]:
    sampled = {row["article_id"] for row in rows}
    sources = {
        "legacy80": article_ids(ROOT / "outputs/task1_step5_7f_structure_training_v1/legacy80_diagnostic_eventmentions.jsonl"),
        "new200": article_ids(ROOT / "outputs/task1_step5_7f_structure_training_v1/confirmed_eventmentions_200.jsonl"),
        "classifier_dev": article_ids(ROOT / "outputs/task1_step5_6d_classifier_retraining/dev_articles.jsonl"),
        "classifier_holdout": article_ids(ROOT / "outputs/task1_step5_6d_classifier_retraining/holdout_articles.jsonl"),
        "gold1029_blind_article_ids_only": article_ids(ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl"),
    }
    return {name: {"count": len(sampled & ids), "article_ids": sorted(sampled & ids)[:100]} for name, ids in sources.items()}


def blocking_intersections(inter: dict[str, Any]) -> dict[str, Any]:
    blocking_names = {"legacy80", "new200", "classifier_holdout", "gold1029_blind_article_ids_only"}
    return {name: payload for name, payload in inter.items() if name in blocking_names and payload.get("count", 0)}


def freeze_data(rows: list[dict[str, Any]], confirmed_at: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    frozen_rows = []
    for row in rows:
        out = dict(row)
        out["confirmed_at"] = confirmed_at
        if row["split"] == "independent_frozen_test":
            out["frozen"] = True
            out["frozen_policy"] = "Do not read labels for training, threshold selection, early stopping, model/prompt selection, or error-driven changes before final model/config freeze."
        else:
            out["frozen"] = False
        frozen_rows.append(out)
    train = [row for row in frozen_rows if row["split"] == "expansion_train"]
    test = [row for row in frozen_rows if row["split"] == "independent_frozen_test"]
    write_jsonl(OUT_DIR / "expansion_train_confirmed.jsonl", train)
    write_jsonl(OUT_DIR / "independent_frozen_test_confirmed.jsonl", test)
    return train, test, {
        "split_unit": "article_id",
        "expansion_train_count": len(train),
        "independent_frozen_test_count": len(test),
        "expansion_train_article_count": len({row["article_id"] for row in train}),
        "independent_frozen_test_article_count": len({row["article_id"] for row in test}),
        "article_overlap": sorted({row["article_id"] for row in train} & {row["article_id"] for row in test}),
        "independent_frozen_test_sha256": sha256_file(OUT_DIR / "independent_frozen_test_confirmed.jsonl"),
        "expansion_train_sha256": sha256_file(OUT_DIR / "expansion_train_confirmed.jsonl"),
        "frozen_test_policy": "The 60 labels are locked and may be read only once for final evaluation after model/config freeze.",
    }


def existing_data_audit(new_train: list[dict[str, Any]], new_test: list[dict[str, Any]]) -> dict[str, Any]:
    old_train = read_jsonl(ROOT / "outputs/task1_step5_7f_structure_training_v1/train_eventmentions.jsonl")
    old_dev = read_jsonl(ROOT / "outputs/task1_step5_7f_structure_training_v1/dev_eventmentions.jsonl")
    old80 = read_jsonl(ROOT / "outputs/task1_step5_7f_structure_training_v1/legacy80_diagnostic_eventmentions.jsonl")
    sets = {
        "old160_train": {row["article_id"] for row in old_train},
        "old40_dev": {row["article_id"] for row in old_dev},
        "new240_train": {row["article_id"] for row in new_train},
        "new60_test": {row["article_id"] for row in new_test},
        "legacy80_diag": {row["article_id"] for row in old80},
    }
    overlap = {}
    keys = list(sets)
    for i, a in enumerate(keys):
        for b in keys[i + 1 :]:
            overlap[f"{a}__{b}"] = sorted(sets[a] & sets[b])[:100]
    return {
        "old160_train_count": len(old_train),
        "old40_dev_count": len(old_dev),
        "legacy80_diagnostic_count": len(old80),
        "planned_structure_train_count": len(old_train) + len(new_train),
        "planned_structure_dev_count": len(old_dev),
        "planned_independent_frozen_test_count": len(new_test),
        "legacy80_use": "history_diagnostic_only_not_model_selection",
        "article_overlaps": {k: {"count": len(v), "article_ids": v} for k, v in overlap.items()},
        "leakage_policy": {
            "train": "old160 + new240 = 400",
            "dev": "old40 only",
            "independent_frozen_test": "new60 only; labels hidden until final frozen evaluation",
            "forbidden_in_inference": ["ai_*", "human_*", "reference_raw_text", "reviewer_*", "quality_grade", "validation_errors", "gold labels", "label-derived candidates or offsets"],
        },
    }


def write_reports(report: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    REPORT_JSON.parent.mkdir(parents=True, exist_ok=True)
    write_json(REPORT_JSON, report)
    split = report.get("split_manifest") or {}
    validation = report.get("validation") or {}
    split_counts = validation.get("split_counts") or {}
    train_count = split.get("expansion_train_count", split_counts.get("expansion_train", 0))
    test_count = split.get("independent_frozen_test_count", split_counts.get("independent_frozen_test", 0))
    test_sha = split.get("independent_frozen_test_sha256", "not_generated")
    lines = [
        "# Task1 Step 5.7-J 结构标注确认报告",
        "",
        f"- status: `{report['status']}`",
        f"- confirmed_count: {validation.get('row_count', len(rows))}",
        f"- expansion_train: {train_count}",
        f"- independent_frozen_test: {test_count}",
        f"- independent_frozen_test_sha256: `{test_sha}`",
        f"- validation_errors: {len(validation.get('errors', []))}",
        "",
        "本数据为AI辅助、人工确认标注，不是独立人工金标。",
    ]
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if report["status"] == STATUS_READY:
        plan_text = (
            "# Task1 Step 5.7-J 下一轮结构抽取训练方案\n\n"
            "状态达到 `structure_training_dataset_ready` 后再启动训练。本轮不训练模型。\n\n"
            "## 数据用途\n\n"
            "- structure_train: 原160条确认训练集 + 新240条确认扩展集，共400条。\n"
            "- structure_dev: 原40条确认开发集。\n"
            "- independent_frozen_test: 新60条，仅最终模型和配置完全冻结后评估一次。\n"
            "- 旧80条: 仅历史诊断，不参与模型选择。\n\n"
            "## 推荐模型\n\n"
            "使用MacBERT或现有中文编码器共享backbone；trigger、event_span、subject、time、object使用独立token/span head；direction使用分类head；arguments使用token classification或span+role head。所有offset由token位置映射回原content。\n\n"
            "## 泄漏约束\n\n"
            "推理阶段只允许读取content、title、publish_time和冻结event_name。严禁读取ai_*、human_*、reviewer_*、reference_raw_text、quality_grade、validation_errors、gold标签或任何标签派生候选/offset。必须加入清空、打乱、删除ai/human/reference字段的负向泄漏测试。\n\n"
            "## 预计训练时间\n\n"
            "400条规模下，MacBERT多任务模型单卡或多卡微调预计约30-90分钟，具体取决于max_length、batch size和是否加入CRF/span pair head。\n"
        )
    else:
        plan_text = (
            "# Task1 Step 5.7-J 下一轮结构抽取训练方案\n\n"
            "当前状态为 `structure_annotation_confirmation_failed`，训练方案未解锁；不得启动训练、模型选择或读取独立冻结测试标签。\n\n"
            "必须先修正确认工作簿中的校验失败项，并重新通过5.7-J确认冻结流程。通过后推荐使用MacBERT或现有中文编码器共享backbone，配合trigger、event_span、subject、time、object独立token/span head，direction分类head，以及arguments的token classification或span+role head。推理阶段仍只能读取content、title、publish_time和冻结event_name，并必须保留ai/human/reference字段负向泄漏测试。\n"
        )
    PLAN_MD.write_text(plan_text, encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    summary_rows = []
    for key, value in report["validation"].items():
        summary_rows.append({"section": "validation", "key": key, "value": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value})
    for key, value in report["intersections"].items():
        summary_rows.append({"section": "intersections", "key": key, "value": json.dumps(value, ensure_ascii=False)})
    headers = ["section", "key", "value"]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, c).fill = PatternFill("solid", fgColor="1F4E78")
    for r, row in enumerate(summary_rows, 2):
        for c, h in enumerate(headers, 1):
            ws.cell(r, c, row[h])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["C"].width = 100
    ws2 = wb.create_sheet("confirmed_rows")
    row_headers = ["event_mention_id", "article_id", "split", "event_name", "direction", "trigger", "event_span", "subject_company", "event_time", "object_target"]
    for c, h in enumerate(row_headers, 1):
        ws2.cell(1, c, h).font = Font(bold=True, color="FFFFFF")
        ws2.cell(1, c).fill = PatternFill("solid", fgColor="1F4E78")
    for r, row in enumerate(rows, 2):
        values = {
            "event_mention_id": row["event_mention_id"],
            "article_id": row["article_id"],
            "split": row["split"],
            "event_name": row["event_name"],
            "direction": row["direction"],
            "trigger": json.dumps(row["trigger"], ensure_ascii=False),
            "event_span": json.dumps(row["event_span"], ensure_ascii=False),
            "subject_company": json.dumps(row["subject_company"], ensure_ascii=False),
            "event_time": json.dumps(row["event_time"], ensure_ascii=False),
            "object_target": json.dumps(row["object_target"], ensure_ascii=False),
        }
        for c, h in enumerate(row_headers, 1):
            ws2.cell(r, c, values[h])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-xlsx", default=str(INPUT_XLSX), help="AI-assisted audited workbook to confirm.")
    args = parser.parse_args()
    input_xlsx = Path(args.input_xlsx)
    if not input_xlsx.is_absolute():
        input_xlsx = ROOT / input_xlsx
    if not input_xlsx.exists():
        raise FileNotFoundError(input_xlsx)
    if OUT_DIR.exists() and any(OUT_DIR.iterdir()):
        raise FileExistsError(f"refusing to overwrite existing 5.7-J output directory: {OUT_DIR}")
    if REPORT_JSON.exists() or CONFIRMED_XLSX.exists():
        raise FileExistsError("refusing to overwrite existing 5.7-J report/workbook")
    proc = process_check()
    if proc:
        raise RuntimeError(f"unexpected related process running: {proc}")
    backup_dir = Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7j_structure_confirmation_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup_dir.mkdir(parents=True, exist_ok=False)
    OUT_DIR.mkdir(parents=True, exist_ok=False)
    protected_before = sha_snapshot(protected_paths([input_xlsx]))
    write_sha_txt(OUT_DIR / "protected_sha256_before.txt", protected_before)
    write_json(backup_dir / "protected_sha256_before.json", protected_before)
    confirmed_at = now()
    rows, row_errors, workbook_meta = extract_rows_from_workbook(input_xlsx)
    official_names = official_event_names()
    validation = validate_confirmed_rows(rows, official_names)
    validation["row_count"] = len(rows)
    validation["workbook_parse_errors"] = row_errors
    if row_errors:
        validation["valid"] = False
        validation["errors"] = sorted(set(validation["errors"] + row_errors))
    inter = intersections(rows)
    blocked_intersections = blocking_intersections(inter)
    if blocked_intersections:
        validation["valid"] = False
        validation["errors"].append("blocking_historical_or_blind_article_intersection")
    status = STATUS_READY if validation["valid"] else STATUS_FAILED
    if status != STATUS_READY:
        report = {
            "generated_at": confirmed_at,
            "status": status,
            "validation": validation,
            "intersections": inter,
            "blocking_intersections": blocked_intersections,
            "workbook_meta": workbook_meta,
            "source_input_workbook": rel(input_xlsx),
            "source_input_workbook_sha256": sha256_file(input_xlsx),
            "backup_dir": str(backup_dir),
            "llm_calls": {"local_32b": 0, "external_llm": 0},
            "vllm_started": False,
        }
        write_reports(report, rows)
        write_json(OUT_DIR / "validation_report.json", validation)
        print(json.dumps({"status": status, "errors": validation["errors"][:20]}, ensure_ascii=False, indent=2))
        return 1
    update_confirmed_workbook(input_xlsx, CONFIRMED_XLSX, confirmed_at)
    for row in rows:
        row["confirmed_at"] = confirmed_at
    train, test, split_manifest = freeze_data(rows, confirmed_at)
    write_json(OUT_DIR / "split_manifest.json", split_manifest)
    validation["confirmed_workbook"] = rel(CONFIRMED_XLSX)
    validation["confirmed_workbook_sha256"] = sha256_file(CONFIRMED_XLSX)
    write_json(OUT_DIR / "validation_report.json", validation)
    existing_audit = existing_data_audit(train, test)
    protected_after = sha_snapshot(protected_paths([input_xlsx]))
    write_sha_txt(OUT_DIR / "protected_sha256_after.txt", protected_after)
    protected_match = protected_before == protected_after
    if not protected_match:
        status = STATUS_FAILED
    report = {
        "generated_at": confirmed_at,
        "status": status,
        "artifact_type": "ai_assisted_human_confirmed_annotation",
        "independent_manual_gold": False,
        "validation": validation,
        "source_input_workbook": rel(input_xlsx),
        "source_input_workbook_sha256": sha256_file(input_xlsx),
        "split_manifest": split_manifest,
        "intersections": inter,
        "blocking_intersections": blocked_intersections,
        "existing_data_audit": existing_audit,
        "protection": {
            "protected_sha256_before": protected_before,
            "protected_sha256_after": protected_after,
            "protected_sha256_match": protected_match,
            "formal_task1_task2_modified": not protected_match,
            "news_with_tags_test_read": False,
            "gold1029_labels_read": False,
        },
        "llm_calls": {"local_32b": 0, "external_llm": 0},
        "vllm_started": False,
        "backup_dir": str(backup_dir),
    }
    write_reports(report, rows)
    output_paths = [
        CONFIRMED_XLSX,
        OUT_DIR / "expansion_train_confirmed.jsonl",
        OUT_DIR / "independent_frozen_test_confirmed.jsonl",
        OUT_DIR / "split_manifest.json",
        OUT_DIR / "validation_report.json",
        REPORT_JSON,
        REPORT_MD,
        DETAILS_XLSX,
        PLAN_MD,
    ]
    write_sha_txt(OUT_DIR / "SHA256SUMS.txt", {rel(path): sha256_file(path) for path in output_paths if path.exists()})
    write_json(backup_dir / "run_manifest.json", report)
    print(json.dumps({"status": status, "train": len(train), "test": len(test), "report": rel(REPORT_JSON), "backup_dir": str(backup_dir)}, ensure_ascii=False, indent=2))
    return 0 if status == STATUS_READY else 1


if __name__ == "__main__":
    raise SystemExit(main())
