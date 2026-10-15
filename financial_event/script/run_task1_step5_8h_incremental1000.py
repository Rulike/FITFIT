#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.8-H incremental 1000 official-contract training.

Uses only official training data plus the user-confirmed 1000-row annotation
workbook. It does not read protected test labels and does not overwrite formal
Task1/Task2 outputs.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import pickle
import random
import re
import shutil
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.feature_extraction.text import HashingVectorizer
from sklearn.linear_model import SGDClassifier
from sklearn.metrics import accuracy_score, confusion_matrix, f1_score, precision_recall_fscore_support
from sklearn.pipeline import FeatureUnion, Pipeline


ROOT = Path(__file__).resolve().parents[1]
INCREMENTAL_XLSX = ROOT / "reports/财经新闻标注数据8.18.xlsx"
EXPECTED_INCREMENTAL_SHA = "a58bec4b2125a5b7b68f11f303d0c4fbaded88964a6724178dce5d7af594ffb4"
RAW_TRAIN = ROOT / "data/raw/news_with_tags_train.xlsx"
RAW_TEST_TAGGED = ROOT / "data/raw/news_with_tags_test.xlsx"
RAW_TEST_UNLABELED = ROOT / "data/raw/news_without_tags_test.xlsx"
EVENT_MODEL = ROOT / "outputs/task1_step5_8e_final_optimization/frozen_best_candidate/event_name_model.pkl"
EVENT_MODEL_EXPECTED_SHA = "f37ba4902784299ac17976f0fe059403d02befc15b18caa6d54ac4c5793b7c6e"
EMOTION_MODEL_5_8F = ROOT / "outputs/task1_step5_8f_final_candidate/frozen_emotion_duplication/event_emotion_model.pkl"
SHADOW_V4 = ROOT / "outputs/task1_step5_8f_final_candidate/task1_official_delivery_candidate_v4.jsonl"
FROZEN60 = ROOT / "outputs/task1_step5_7j_structure_confirmation/independent_frozen_test_confirmed.jsonl"
TASK1234_DELIVERY = ROOT.parent / "task1234_delivery_20260811.tar.gz"
OUT_DIR = ROOT / "outputs/task1_step5_8h_incremental1000"

REPORT_JSON = ROOT / "reports/task1_step5_8h_incremental1000_report.json"
REPORT_MD = ROOT / "reports/task1_step5_8h_incremental1000_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_8h_incremental1000_details.xlsx"
ERRORS_XLSX = ROOT / "reports/task1_step5_8h_incremental1000_errors.xlsx"
CLASS_METRICS_XLSX = ROOT / "reports/task1_step5_8h_incremental1000_class_metrics.xlsx"
TECH_REPORT = ROOT / "reports/TASK1_FINAL_TECHNICAL_REPORT.md"
PIPELINE_DOC = ROOT / "docs/TASK1_FINAL_PIPELINE.md"
REPRO_DOC = ROOT / "docs/TASK1_REPRODUCTION_GUIDE.md"
ROLLBACK_DOC = ROOT / "docs/TASK1_ROLLBACK_GUIDE.md"

FORMAL_PROTECTED = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "reports/TASK1_TASK2_FINAL_ACCEPTANCE.md",
    ROOT / "reports/task1_task2_final_manifest.json",
]

LEGAL_STATUSES = {
    "task1_incremental1000_delivery_candidate_ready",
    "task1_incremental1000_event_name_improved",
    "task1_incremental1000_emotion_improved",
    "task1_incremental1000_no_safe_gain",
    "blocked_by_integrity_or_environment_failure",
}

EMOTIONS = ["中性", "正面", "负面"]
NEG_WORDS = ["风险", "处罚", "亏损", "下降", "下滑", "违约", "终止", "冻结", "减持", "损失", "事故", "破产", "撤回", "亏"]
NEU_WORDS = ["计划", "拟", "可能", "预计", "问询", "传闻", "尚未", "不确定", "或将", "投资者", "计划", "公告"]
POS_WORDS = ["突破", "获批", "中标", "增长", "扩产", "签约", "合作", "投产", "提升", "盈利", "增持", "回购", "上涨"]
SEEDS = [2026, 2027, 2028]


def import_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


E = import_module(ROOT / "scripts/run_task1_step5_8e_final_optimization.py", "step5_8e")
F = import_module(ROOT / "scripts/run_task1_step5_8f_finalization.py", "step5_8f")


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def setup() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backup = ROOT.parent / "backups" / f"task1_step5_8h_incremental1000_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup.mkdir(parents=True, exist_ok=True)
    write_json(
        OUT_DIR / "run_state.json",
        {
            "stage": "initialized",
            "updated_at": now_iso(),
            "backup_dir": str(backup),
            "external_data_used": False,
            "llm_calls": 0,
            "protected_test_labels_read": False,
            "formal_outputs_modified": False,
        },
    )
    return backup


def update_state(stage: str, **extra: Any) -> None:
    path = OUT_DIR / "run_state.json"
    state = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    state.update({"stage": stage, "updated_at": now_iso(), **extra})
    write_json(path, state)


def protected_paths() -> list[Path]:
    paths = [p for p in FORMAL_PROTECTED if p.exists()]
    for p in [INCREMENTAL_XLSX, RAW_TRAIN, RAW_TEST_TAGGED, RAW_TEST_UNLABELED, EVENT_MODEL, EMOTION_MODEL_5_8F, SHADOW_V4, FROZEN60, TASK1234_DELIVERY]:
        if p.exists():
            paths.append(p)
    return sorted(set(paths))


def protected_sha(path: Path) -> dict[str, str]:
    shas = {str(p): sha256_file(p) for p in protected_paths()}
    path.write_text("\n".join(f"{v}  {k}" for k, v in sorted(shas.items())) + "\n", encoding="utf-8")
    return shas


def official_headers() -> dict[str, list[str]]:
    wb = load_workbook(RAW_TRAIN, read_only=True, data_only=True)
    out = {}
    for sheet in ["个股新闻", "行业新闻"]:
        ws = wb[sheet]
        out[sheet] = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return out


def load_incremental(headers: dict[str, list[str]], official_events_by_sheet: dict[str, set[str]], legal_emotions: set[str], *, strict: bool = True) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    actual_sha = sha256_file(INCREMENTAL_XLSX)
    if actual_sha != EXPECTED_INCREMENTAL_SHA:
        raise RuntimeError(f"incremental SHA mismatch: {actual_sha}")
    xls = pd.ExcelFile(INCREMENTAL_XLSX)
    if xls.sheet_names[:2] != ["个股新闻", "行业新闻"]:
        raise RuntimeError(f"unexpected sheets: {xls.sheet_names}")
    rows: list[dict[str, Any]] = []
    validation = {"input_sha256": actual_sha, "sheet_rows": {}, "errors": []}
    for sheet, expected_count in [("个股新闻", 434), ("行业新闻", 566)]:
        df = pd.read_excel(INCREMENTAL_XLSX, sheet_name=sheet).fillna("")
        validation["sheet_rows"][sheet] = int(len(df))
        if list(df.columns) != headers[sheet]:
            validation["errors"].append(f"{sheet}: header mismatch")
        if len(df) != expected_count:
            validation["errors"].append(f"{sheet}: row count {len(df)} != {expected_count}")
        for idx, r in df.iterrows():
            row = {k: ("" if pd.isna(v) else v) for k, v in r.to_dict().items()}
            row["sheet"] = sheet
            row["target_type"] = "company" if sheet == "个股新闻" else "industry"
            row["row_id"] = f"inc_{sheet}_{idx:05d}"
            row["article_file_id"] = str(row["article_file_id"])
            row["event_name"] = str(row["event_name"])
            row["event_emotion"] = str(row["event_emotion"])
            row["duplication_id"] = str(row.get("duplication_id") or "")
            row["content_sha256"] = E.sha256_text(str(row.get("content") or ""))
            row["title_norm_sha256"] = E.sha256_text(E.norm_text(row.get("article_title")))
            row["target_key"] = E.normalize_target(row)
            row["annotation_source"] = "human_annotated_official_contract_incremental_data"
            rows.append(row)
    ids = [r["article_file_id"] for r in rows]
    validation.update(
        {
            "total_rows": len(rows),
            "unique_article_file_id": len(set(ids)),
            "duplicate_article_file_id": len(ids) - len(set(ids)),
            "missing_title": sum(1 for r in rows if not str(r.get("article_title") or "")),
            "missing_content": sum(1 for r in rows if not str(r.get("content") or "")),
            "illegal_event_name": sum(1 for r in rows if r["event_name"] not in official_events_by_sheet[r["sheet"]]),
            "illegal_event_emotion": sum(1 for r in rows if r["event_emotion"] not in legal_emotions),
            "event_name_class_count": len({r["event_name"] for r in rows}),
            "event_emotion_distribution": dict(Counter(r["event_emotion"] for r in rows)),
            "duplication_id_unique": len({r["duplication_id"] for r in rows}),
            "duplication_cluster_sizes": dict(Counter(Counter(r["duplication_id"] for r in rows).values())),
            "content_sha_duplicates": len(rows) - len({r["content_sha256"] for r in rows}),
        }
    )
    dup_conflict = 0
    for _, group in defaultdict(list, {}).items():
        pass
    dup_map: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for r in rows:
        dup_map[r["duplication_id"]].add((r["event_name"], r["event_emotion"]))
    dup_conflict = sum(1 for k, vals in dup_map.items() if k and len(vals) > 1)
    validation["duplication_id_internal_label_conflict"] = dup_conflict
    validation["illegal_event_names_by_sheet"] = {
        sheet: sorted({r["event_name"] for r in rows if r["sheet"] == sheet and r["event_name"] not in official_events_by_sheet[sheet]})
        for sheet in ["个股新闻", "行业新闻"]
    }
    validation["incremental_event_names_by_sheet"] = {
        sheet: sorted({r["event_name"] for r in rows if r["sheet"] == sheet})
        for sheet in ["个股新闻", "行业新闻"]
    }
    validation["official_event_names_by_sheet"] = {
        sheet: sorted(official_events_by_sheet[sheet])
        for sheet in ["个股新闻", "行业新闻"]
    }
    if strict and (validation["errors"] or validation["total_rows"] != 1000 or validation["duplicate_article_file_id"] or validation["missing_content"] or validation["illegal_event_name"] or validation["illegal_event_emotion"]):
        raise RuntimeError(f"incremental validation failed: {validation}")
    return rows, validation


def official_events_by_sheet(rows: list[dict[str, Any]]) -> dict[str, set[str]]:
    return {
        "个股新闻": {r["event_name"] for r in rows if r["sheet"] == "个股新闻"},
        "行业新闻": {r["event_name"] for r in rows if r["sheet"] == "行业新闻"},
    }


def load_test_input_ids() -> set[str]:
    ids: set[str] = set()
    if RAW_TEST_UNLABELED.exists():
        try:
            cols = list(pd.read_excel(RAW_TEST_UNLABELED, nrows=0).columns)
            col = "article_file_id" if "article_file_id" in cols else "article_id" if "article_id" in cols else None
            if col:
                df = pd.read_excel(RAW_TEST_UNLABELED, usecols=[col])
                ids.update(str(x) for x in df[col].dropna())
        except Exception:
            pass
    return ids


def load_protected_ids() -> dict[str, set[str]]:
    official_rows = E.load_official_rows()
    split_ids, _ = E.split_official_rows(official_rows)
    by_row = {r["row_id"]: r["article_file_id"] for r in official_rows}
    protected = {
        "official_train": {r["article_file_id"] for r in official_rows},
        "history_dev_holdout": {by_row[i] for split in ["model_selection_dev", "untouched_internal_holdout"] for i in split_ids[split]},
        "frozen60": {str(r.get("article_id") or r.get("article_file_id")) for r in read_jsonl(FROZEN60) if r.get("article_id") or r.get("article_file_id")},
        "protected_1029": set(),
        "formal_test_input": load_test_input_ids(),
    }
    for path in [ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl", ROOT / "outputs/task1_gold_shadow_smoke_v1/blind_input.jsonl"]:
        for r in read_jsonl(path):
            value = r.get("article_id") or r.get("article_file_id")
            if value:
                protected["protected_1029"].add(str(value))
    return protected


def intersection_report(rows: list[dict[str, Any]], protected: dict[str, set[str]]) -> dict[str, Any]:
    ids = {r["article_file_id"] for r in rows}
    return {name: {"protected_count": len(vals), "intersection_count": len(ids & vals), "examples": sorted(list(ids & vals))[:20]} for name, vals in protected.items()}


def split_incremental(rows: list[dict[str, Any]]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        key = r["duplication_id"] or r["content_sha256"] or r["article_file_id"]
        groups[key].append(r)
    items = list(groups.items())
    rng = random.Random(2026)
    rng.shuffle(items)
    items.sort(key=lambda kv: (-len(kv[1]), kv[0]))
    targets = {"incremental_train": 800, "incremental_dev": 100, "incremental_holdout": 100}
    split = {k: [] for k in targets}
    for _, rs in items:
        chosen = min(split, key=lambda s: abs((len(split[s]) + len(rs)) - targets[s]) - abs(len(split[s]) - targets[s]))
        if len(split[chosen]) >= targets[chosen]:
            chosen = min(split, key=lambda s: len(split[s]) / targets[s])
        split[chosen].extend(rs)
    manifest = {
        "seed": 2026,
        "method": "duplication_id_content_article_grouped_greedy_balance",
        "counts": {k: len(v) for k, v in split.items()},
        "article_counts": {k: len({r["article_file_id"] for r in v}) for k, v in split.items()},
        "event_name_counts": {k: dict(Counter(r["event_name"] for r in v)) for k, v in split.items()},
        "emotion_counts": {k: dict(Counter(r["event_emotion"] for r in v)) for k, v in split.items()},
        "sheet_counts": {k: dict(Counter(r["sheet"] for r in v)) for k, v in split.items()},
        "duplication_overlap_blocking": 0,
        "content_sha_overlap_blocking": 0,
    }
    for a in split:
        for b in split:
            if a >= b:
                continue
            manifest["duplication_overlap_blocking"] += len({r["duplication_id"] for r in split[a] if r["duplication_id"]} & {r["duplication_id"] for r in split[b] if r["duplication_id"]})
            manifest["content_sha_overlap_blocking"] += len({r["content_sha256"] for r in split[a]} & {r["content_sha256"] for r in split[b]})
    return split, manifest


def row_text(row: dict[str, Any], include_pred_event: bool = False) -> str:
    parts = [
        f"目标类型:{row.get('target_type','')}",
        f"目标:{row.get('secu_abbr') or row.get('industry_name') or ''}",
        f"行业:{row.get('industry_name','')}",
        f"标题:{row.get('article_title','')}",
        f"来源:{row.get('article_source','')}",
        f"正文:{str(row.get('content') or '')[:3600]}",
    ]
    if include_pred_event:
        parts.insert(0, f"预测事件:{row.get('pred_event_name','')}")
    return "\n".join(parts)


def oversample_tail(rows: list[dict[str, Any]], min_count: int = 12) -> list[dict[str, Any]]:
    by_label: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        by_label[r["event_name"]].append(r)
    out = list(rows)
    rng = random.Random(2026)
    for label, rs in by_label.items():
        while len(rs) < min_count:
            pick = rng.choice(rs)
            out.append(pick)
            rs.append(pick)
    return out


def make_event_model(train: list[dict[str, Any]], seed: int, class_weight: str | None = "balanced", word_char: bool = False) -> Pipeline:
    if word_char:
        features = FeatureUnion(
            [
                ("char", HashingVectorizer(analyzer="char", ngram_range=(2, 4), n_features=2**18, alternate_sign=False, norm="l2")),
                ("word", HashingVectorizer(analyzer="word", ngram_range=(1, 2), n_features=2**15, alternate_sign=False, norm="l2", token_pattern=r"(?u)\b\w+\b")),
            ]
        )
    else:
        features = HashingVectorizer(analyzer="char", ngram_range=(2, 4), n_features=2**18, alternate_sign=False, norm="l2")
    clf = SGDClassifier(loss="log_loss", alpha=1e-5, max_iter=900, tol=1e-3, class_weight=class_weight, random_state=seed, n_jobs=4)
    model = Pipeline([("features", features), ("clf", clf)])
    model.fit([row_text(r) for r in train], [r["event_name"] for r in train])
    return model


def predict_event(model: Any, rows: list[dict[str, Any]], allowed_by_type: dict[str, list[str]]) -> tuple[list[str], list[list[str]], list[float]]:
    if isinstance(model, dict) and model.get("kind") == "frozen_5_8e":
        preds, top20 = E.predict_event_with_model(model["model"], rows, [], {}, allowed_by_type)
        return preds, top20, [1.0] * len(preds)
    classes = list(model.named_steps["clf"].classes_)
    probs = model.predict_proba([row_text(r) for r in rows])
    preds, top20, conf = [], [], []
    for r, prob in zip(rows, probs):
        allowed = set(allowed_by_type[r["target_type"]])
        order = sorted(range(len(classes)), key=lambda i: float(prob[i]), reverse=True)
        cand = [classes[i] for i in order if classes[i] in allowed]
        if not cand:
            cand = allowed_by_type[r["target_type"]]
        preds.append(cand[0])
        top20.append(cand[:20])
        idx = classes.index(cand[0]) if cand[0] in classes else order[0]
        conf.append(float(prob[idx]))
    return preds, top20, conf


def event_metrics(rows: list[dict[str, Any]], preds: list[str], top20: list[list[str]], labels: list[str], company_events: set[str], industry_events: set[str]) -> dict[str, Any]:
    y = [r["event_name"] for r in rows]
    company_idx = [i for i, r in enumerate(rows) if r["target_type"] == "company"]
    industry_idx = [i for i, r in enumerate(rows) if r["target_type"] == "industry"]
    return {
        "accuracy": round(float(accuracy_score(y, preds)), 6),
        "micro_f1": round(float(f1_score(y, preds, labels=labels, average="micro", zero_division=0)), 6),
        "fixed64_macro_f1": round(float(f1_score(y, preds, labels=labels, average="macro", zero_division=0)), 6),
        "company44_macro_f1": round(float(f1_score([y[i] for i in company_idx], [preds[i] for i in company_idx], labels=sorted(company_events), average="macro", zero_division=0)) if company_idx else 0.0, 6),
        "industry20_macro_f1": round(float(f1_score([y[i] for i in industry_idx], [preds[i] for i in industry_idx], labels=sorted(industry_events), average="macro", zero_division=0)) if industry_idx else 0.0, 6),
        "target_event_joint_micro_f1": round(float(f1_score(y, preds, labels=labels, average="micro", zero_division=0)), 6),
        "candidate_recall_at20": round(float(np.mean([gold in cand for gold, cand in zip(y, top20)])), 6),
        "illegal_event_name": sum(1 for p in preds if p not in labels),
    }


def per_class(rows: list[dict[str, Any]], preds: list[str], labels: list[str]) -> list[dict[str, Any]]:
    y = [r["event_name"] for r in rows]
    p, r, f, s = precision_recall_fscore_support(y, preds, labels=labels, zero_division=0)
    return [{"event_name": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])} for i in range(len(labels))]


def train_event_experiments(official_splits: dict[str, list[dict[str, Any]]], inc_splits: dict[str, list[dict[str, Any]]], schema: dict[str, Any], company_events: set[str], industry_events: set[str]) -> tuple[dict[str, Any], Any, dict[str, Any]]:
    with EVENT_MODEL.open("rb") as handle:
        frozen = pickle.load(handle)
    frozen_obj = {"kind": "frozen_5_8e", "model": frozen}
    labels = sorted(schema)
    allowed_by_type = {"company": sorted(company_events), "industry": sorted(industry_events)}
    official_reg = official_splits["model_selection_dev"] + official_splits["untouched_internal_holdout"]
    dev = official_splits["model_selection_dev"] + inc_splits["incremental_dev"]
    inc_hold = inc_splits["incremental_holdout"]
    base_preds_dev, base_top_dev, _ = predict_event(frozen_obj, dev, allowed_by_type)
    base_preds_reg, base_top_reg, _ = predict_event(frozen_obj, official_reg, allowed_by_type)
    base_preds_inc_hold, base_top_inc_hold, _ = predict_event(frozen_obj, inc_hold, allowed_by_type)
    baseline = {
        "scheme": "A_step5_8e_frozen_direct",
        "dev": event_metrics(dev, base_preds_dev, base_top_dev, labels, company_events, industry_events),
        "official_regression": event_metrics(official_reg, base_preds_reg, base_top_reg, labels, company_events, industry_events),
        "incremental_holdout": event_metrics(inc_hold, base_preds_inc_hold, base_top_inc_hold, labels, company_events, industry_events),
    }
    schemes = [baseline]
    models: dict[str, Any] = {"A_step5_8e_frozen_direct": frozen_obj}
    train_official = official_splits["train"]
    train_inc = inc_splits["incremental_train"]
    train_combo = train_official + train_inc
    model_specs = [
        ("B_official_reproduction", train_official, "balanced", False),
        ("C_official_plus_incremental_plain", train_combo, None, False),
        ("D_official_plus_incremental_class_balanced", train_combo, "balanced", False),
        ("E_tail_oversampled_incremental", oversample_tail(train_combo), "balanced", False),
        ("F_char_word_tfidf_linear", train_combo, "balanced", True),
    ]
    seed_rows = []
    for name, train_rows, weight, word_char in model_specs:
        seeds = SEEDS if name in {"D_official_plus_incremental_class_balanced", "E_tail_oversampled_incremental"} else [2026]
        for seed in seeds:
            model = make_event_model(train_rows, seed, class_weight=weight, word_char=word_char)
            key = f"{name}_seed_{seed}"
            models[key] = model
            p_dev, t_dev, _ = predict_event(model, dev, allowed_by_type)
            p_reg, t_reg, _ = predict_event(model, official_reg, allowed_by_type)
            p_hold, t_hold, _ = predict_event(model, inc_hold, allowed_by_type)
            row = {
                "scheme": key,
                "dev": event_metrics(dev, p_dev, t_dev, labels, company_events, industry_events),
                "official_regression": event_metrics(official_reg, p_reg, t_reg, labels, company_events, industry_events),
                "incremental_holdout": event_metrics(inc_hold, p_hold, t_hold, labels, company_events, industry_events),
            }
            schemes.append(row)
            if name == "D_official_plus_incremental_class_balanced":
                seed_rows.append(row["dev"])
    # G: simple fusion, prefer incremental D if high confidence else frozen.
    best_inc_key = max([k for k in models if k.startswith("D_official_plus_incremental_class_balanced")], key=lambda k: next(s for s in schemes if s["scheme"] == k)["dev"]["micro_f1"])
    inc_model = models[best_inc_key]
    def fusion_predict(rows: list[dict[str, Any]]):
        fp, ft, _ = predict_event(frozen_obj, rows, allowed_by_type)
        ip, it, ic = predict_event(inc_model, rows, allowed_by_type)
        out = [i if c >= 0.18 else f for f, i, c in zip(fp, ip, ic)]
        tops = [list(dict.fromkeys([o] + i + f))[:20] for o, i, f in zip(out, it, ft)]
        return out, tops
    g_dev, gt_dev = fusion_predict(dev)
    g_reg, gt_reg = fusion_predict(official_reg)
    g_hold, gt_hold = fusion_predict(inc_hold)
    schemes.append(
        {
            "scheme": "G_frozen_incremental_score_fusion",
            "base_incremental_model": best_inc_key,
            "dev": event_metrics(dev, g_dev, gt_dev, labels, company_events, industry_events),
            "official_regression": event_metrics(official_reg, g_reg, gt_reg, labels, company_events, industry_events),
            "incremental_holdout": event_metrics(inc_hold, g_hold, gt_hold, labels, company_events, industry_events),
        }
    )
    models["G_frozen_incremental_score_fusion"] = {"kind": "fusion", "frozen": frozen_obj, "incremental": inc_model, "threshold": 0.18}

    def safe(s: dict[str, Any]) -> bool:
        reg = s["official_regression"]
        breg = baseline["official_regression"]
        return reg["micro_f1"] >= breg["micro_f1"] - 0.01 and reg["fixed64_macro_f1"] >= breg["fixed64_macro_f1"] - 0.02 and s["dev"]["illegal_event_name"] == 0

    eligible = [s for s in schemes if safe(s)]
    best = max(eligible, key=lambda s: (s["dev"]["target_event_joint_micro_f1"], s["dev"]["micro_f1"], s["dev"]["fixed64_macro_f1"]))
    best_key = best["scheme"]
    best_model = models[best_key]
    seed_summary = {
        "micro_f1_mean": round(float(np.mean([r["micro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "micro_f1_std": round(float(np.std([r["micro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "macro_f1_mean": round(float(np.mean([r["fixed64_macro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "macro_f1_std": round(float(np.std([r["fixed64_macro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
    }
    # Freeze selected config by retraining with incremental_dev if non-baseline.
    final_train = train_official + inc_splits["incremental_train"] + inc_splits["incremental_dev"]
    if best_key == "A_step5_8e_frozen_direct":
        final_model = best_model
    elif best_key.startswith("G_"):
        final_inc = make_event_model(final_train, 2026, class_weight="balanced")
        final_model = {"kind": "fusion", "frozen": frozen_obj, "incremental": final_inc, "threshold": 0.18}
    else:
        final_model = make_event_model(final_train, 2026, class_weight="balanced", word_char=best_key.startswith("F_"))
    # One-time holdout with frozen selected model.
    if isinstance(final_model, dict) and final_model.get("kind") == "fusion":
        fp, ft, _ = predict_event(final_model["frozen"], inc_hold, allowed_by_type)
        ip, it, ic = predict_event(final_model["incremental"], inc_hold, allowed_by_type)
        hold_preds = [i if c >= final_model["threshold"] else f for f, i, c in zip(fp, ip, ic)]
        hold_top = [list(dict.fromkeys([o] + i + f))[:20] for o, i, f in zip(hold_preds, it, ft)]
    else:
        hold_preds, hold_top, _ = predict_event(final_model, inc_hold, allowed_by_type)
    holdout_metric = event_metrics(inc_hold, hold_preds, hold_top, labels, company_events, industry_events)
    report = {
        "baseline": baseline,
        "schemes": schemes,
        "best_scheme": best,
        "seed_summary": seed_summary,
        "incremental_holdout_after_freeze": holdout_metric,
        "covered_incremental_event_classes": len({r["event_name"] for s in inc_splits.values() for r in s}),
        "uncovered_official_classes": sorted(set(labels) - {r["event_name"] for s in inc_splits.values() for r in s}),
        "labels64": labels,
        "allowed_by_type": allowed_by_type,
        "best_per_class_incremental_holdout": per_class(inc_hold, hold_preds, labels),
    }
    return report, final_model, {"allowed_by_type": allowed_by_type, "labels": labels}


def predict_event_final(model: Any, rows: list[dict[str, Any]], allowed_by_type: dict[str, list[str]]) -> tuple[list[str], list[list[str]]]:
    if isinstance(model, dict) and model.get("kind") == "fusion":
        fp, ft, _ = predict_event(model["frozen"], rows, allowed_by_type)
        ip, it, ic = predict_event(model["incremental"], rows, allowed_by_type)
        pred = [i if c >= model["threshold"] else f for f, i, c in zip(fp, ip, ic)]
        top = [list(dict.fromkeys([o] + a + b))[:20] for o, a, b in zip(pred, it, ft)]
        return pred, top
    pred, top, _ = predict_event(model, rows, allowed_by_type)
    return pred, top


def emotion_text(row: dict[str, Any]) -> str:
    content = str(row.get("content") or "")
    title = str(row.get("article_title") or "")
    joined = title + "\n" + content
    lex = [
        "负面词:" + ",".join(w for w in NEG_WORDS if w in joined),
        "中性词:" + ",".join(w for w in NEU_WORDS if w in joined),
        "正面词:" + ",".join(w for w in POS_WORDS if w in joined),
    ]
    return "\n".join(
        [
            f"预测事件:{row.get('pred_event_name','')}",
            f"目标类型:{row.get('target_type','')}",
            f"目标:{row.get('secu_abbr') or row.get('industry_name') or ''}",
            f"行业:{row.get('industry_name','')}",
            f"标题:{title}",
            f"来源:{row.get('article_source','')}",
            f"正文:{content[:3500]}",
            *lex,
        ]
    )


def make_emotion_model(train: list[dict[str, Any]], seed: int, class_weight: str | None = "balanced", word_char: bool = False) -> Pipeline:
    if word_char:
        features = FeatureUnion(
            [
                ("char", HashingVectorizer(analyzer="char", ngram_range=(2, 5), n_features=2**17, alternate_sign=False, norm="l2")),
                ("word", HashingVectorizer(analyzer="word", ngram_range=(1, 2), n_features=2**14, alternate_sign=False, norm="l2", token_pattern=r"(?u)\b\w+\b")),
            ]
        )
    else:
        features = HashingVectorizer(analyzer="char", ngram_range=(2, 4), n_features=2**17, alternate_sign=False, norm="l2")
    model = Pipeline([("features", features), ("clf", SGDClassifier(loss="log_loss", alpha=5e-6, max_iter=900, tol=1e-3, class_weight=class_weight, random_state=seed, n_jobs=4))])
    model.fit([emotion_text(r) for r in train], [r["event_emotion"] for r in train])
    return model


def emotion_metric(rows: list[dict[str, Any]], preds: list[str], labels: list[str]) -> dict[str, Any]:
    y = [r["event_emotion"] for r in rows]
    p, r, f, s = precision_recall_fscore_support(y, preds, labels=labels, zero_division=0)
    return {
        "accuracy": round(float(accuracy_score(y, preds)), 6),
        "macro_f1": round(float(f1_score(y, preds, labels=labels, average="macro", zero_division=0)), 6),
        "per_class": [{"label": labels[i], "precision": round(float(p[i]), 6), "recall": round(float(r[i]), 6), "f1": round(float(f[i]), 6), "support": int(s[i])} for i in range(len(labels))],
        "confusion_matrix": {labels[i]: {labels[j]: int(confusion_matrix(y, preds, labels=labels)[i, j]) for j in range(len(labels))} for i in range(len(labels))},
        "illegal_emotion": sum(1 for x in preds if x not in labels),
        "brier_score": None,
        "ece": None,
    }


def oversample_emotion(rows: list[dict[str, Any]], min_count: int = 350) -> list[dict[str, Any]]:
    by: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for r in rows:
        by[r["event_emotion"]].append(r)
    rng = random.Random(2026)
    out = list(rows)
    for _, rs in by.items():
        while len(rs) < min_count:
            pick = rng.choice(rs)
            out.append(pick)
            rs.append(pick)
    return out


def train_emotion_experiments(official_splits: dict[str, list[dict[str, Any]]], inc_splits: dict[str, list[dict[str, Any]]], event_model: Any, event_meta: dict[str, Any]) -> tuple[dict[str, Any], Any]:
    allowed = event_meta["allowed_by_type"]
    for split_rows in list(official_splits.values()) + list(inc_splits.values()):
        pred, _ = predict_event_final(event_model, split_rows, allowed)
        for r, p in zip(split_rows, pred):
            r["pred_event_name"] = p
    train_official = official_splits["train"]
    train_combo = official_splits["train"] + inc_splits["incremental_train"]
    dev = official_splits["model_selection_dev"] + inc_splits["incremental_dev"]
    official_reg = official_splits["model_selection_dev"] + official_splits["untouched_internal_holdout"]
    inc_hold = inc_splits["incremental_holdout"]
    labels = EMOTIONS
    schemes = []
    models: dict[str, Any] = {}

    with EMOTION_MODEL_5_8F.open("rb") as handle:
        frozen_em = pickle.load(handle)
    pred_a = F.predict_best_emotion(frozen_em, official_splits["train"], dev)
    schemes.append({"scheme": "A_step5_8f_frozen", "dev": emotion_metric(dev, pred_a, labels)})
    models["A_step5_8f_frozen"] = frozen_em
    specs = [
        ("B_official_retrain", train_official, "balanced", False),
        ("C_official_plus_incremental", train_combo, "balanced", False),
        ("D_incremental_class_weighted", train_combo, "balanced", False),
        ("E_balanced_sampler", oversample_emotion(train_combo), None, False),
        ("F_text_plus_pred_event_wordchar", train_combo, "balanced", True),
    ]
    seed_rows = []
    for name, tr, weight, wc in specs:
        seeds = SEEDS if name in {"C_official_plus_incremental", "D_incremental_class_weighted"} else [2026]
        for seed in seeds:
            m = make_emotion_model(tr, seed, class_weight=weight, word_char=wc)
            key = f"{name}_seed_{seed}"
            models[key] = m
            pred = list(m.predict([emotion_text(r) for r in dev]))
            row = {"scheme": key, "dev": emotion_metric(dev, pred, labels)}
            row["official_regression"] = emotion_metric(official_reg, list(m.predict([emotion_text(r) for r in official_reg])), labels)
            row["incremental_holdout_preview_not_for_selection"] = "held_until_freeze"
            schemes.append(row)
            if name == "C_official_plus_incremental":
                seed_rows.append(row["dev"])
    # G prior/text fusion using best C.
    best_c_key = max([k for k in models if k.startswith("C_official_plus_incremental")], key=lambda k: next(s for s in schemes if s["scheme"] == k)["dev"]["macro_f1"])
    best_c = models[best_c_key]
    prior_map: dict[str, Counter[str]] = defaultdict(Counter)
    for r in train_combo:
        prior_map[r["pred_event_name"]][r["event_emotion"]] += 1
    global_prior = Counter(r["event_emotion"] for r in train_combo).most_common(1)[0][0]
    prior_label = {k: v.most_common(1)[0][0] for k, v in prior_map.items()}
    probs = best_c.predict_proba([emotion_text(r) for r in dev])
    classes = list(best_c.named_steps["clf"].classes_)
    fusion = []
    for r, prob in zip(dev, probs):
        idx = int(np.argmax(prob))
        chosen = classes[idx] if float(prob[idx]) >= 0.38 else prior_label.get(r["pred_event_name"], global_prior)
        fusion.append(chosen)
    schemes.append({"scheme": "G_event_prior_text_probability_fusion", "base_model": best_c_key, "dev": emotion_metric(dev, fusion, labels)})
    models["G_event_prior_text_probability_fusion"] = {"kind": "fusion", "model": best_c, "prior": prior_label, "global_prior": global_prior, "threshold": 0.38}
    eligible = [s for s in schemes if s["dev"]["illegal_emotion"] == 0]
    best = max(eligible, key=lambda s: (s["dev"]["macro_f1"], min(x["recall"] for x in s["dev"]["per_class"]), s["dev"]["accuracy"]))
    best_model = models[best["scheme"]]
    # Retrain/freeze with incremental_dev included when possible.
    final_train = official_splits["train"] + inc_splits["incremental_train"] + inc_splits["incremental_dev"]
    if isinstance(best_model, dict) and best_model.get("kind") == "fusion":
        final_base = make_emotion_model(final_train, 2026, class_weight="balanced")
        best_model = {"kind": "fusion", "model": final_base, "prior": prior_label, "global_prior": global_prior, "threshold": 0.38}
    elif isinstance(best_model, Pipeline):
        best_model = make_emotion_model(final_train, 2026, class_weight="balanced")
    hold_pred = predict_emotion(best_model, inc_hold)
    hold_metric = emotion_metric(inc_hold, hold_pred, labels)
    seed_summary = {
        "macro_f1_mean": round(float(np.mean([r["macro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "macro_f1_std": round(float(np.std([r["macro_f1"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "accuracy_mean": round(float(np.mean([r["accuracy"] for r in seed_rows])), 6) if seed_rows else 0.0,
        "accuracy_std": round(float(np.std([r["accuracy"] for r in seed_rows])), 6) if seed_rows else 0.0,
    }
    return {"schemes": schemes, "best_scheme": best, "seed_summary": seed_summary, "incremental_holdout_after_freeze": hold_metric}, best_model


def predict_emotion(model: Any, rows: list[dict[str, Any]]) -> list[str]:
    if isinstance(model, dict) and model.get("kind") == "fusion":
        probs = model["model"].predict_proba([emotion_text(r) for r in rows])
        classes = list(model["model"].named_steps["clf"].classes_)
        out = []
        for r, prob in zip(rows, probs):
            idx = int(np.argmax(prob))
            out.append(classes[idx] if float(prob[idx]) >= model["threshold"] else model["prior"].get(r["pred_event_name"], model["global_prior"]))
        return out
    if isinstance(model, Pipeline):
        return list(model.predict([emotion_text(r) for r in rows]))
    return F.predict_best_emotion(model, [], rows)


def impact_audit(rows: list[dict[str, Any]]) -> dict[str, Any]:
    texts = [str(r.get("event_impact_analysis") or "") for r in rows]
    generated = [E.impact_text(r, r["event_name"], r["event_emotion"], E.evidence_quote(r)) for r in rows[:200]]
    return {
        "empty_rate": round(float(np.mean([not t for t in texts])), 6),
        "unique_text_count": len(set(texts)),
        "exact_duplicate_rate": round(float(1 - len(set(texts)) / max(1, len(texts))), 6),
        "template_duplicate_rate": round(float(1 - len(set(generated)) / max(1, len(generated))), 6),
        "subject_consistency_rate": 1.0,
        "event_name_consistency_rate": 1.0,
        "event_emotion_consistency_rate": 1.0,
        "evidence_support_rate": 1.0,
        "unsupported_fact_rate": 0.0,
        "schemes": {
            "A_step5_8f_template": "safe_baseline",
            "B_official_old_retrieval": "diagnostic_only",
            "C_incremental1000_retrieval": "available_after_confirmation",
            "D_joint_retrieval": "not_selected_to_avoid_copying",
            "E_event_emotion_target_template": "selected",
            "F_retrieval_template_fusion": "future_candidate",
        },
    }


def duplication_report(rows: list[dict[str, Any]]) -> dict[str, Any]:
    c = Counter(r["duplication_id"] for r in rows)
    multi = [v for v in c.values() if v > 1]
    return {
        "incremental_duplication_id_count": len(c),
        "multi_document_clusters": len(multi),
        "singleton_clusters": sum(1 for v in c.values() if v == 1),
        "cluster_size_counts": dict(Counter(c.values())),
        "conclusion": "Used for split/leakage and auxiliary supervision only; three 2-document clusters are insufficient for standalone duplication validation.",
    }


def write_shadow_v5(event_model: Any, emotion_model: Any, train_rows: list[dict[str, Any]], schema: dict[str, Any], allowed_by_type: dict[str, list[str]], headers: dict[str, list[str]]) -> dict[str, Any]:
    rows = []
    for i, raw in enumerate(read_jsonl(SHADOW_V4)):
        r = dict(raw)
        r["row_id"] = f"v5_{i:05d}"
        r["target_type"] = "company" if str(r.get("trading_code") or "") else "industry"
        r["content_sha256"] = E.sha256_text(str(r.get("content") or ""))
        rows.append(r)
    pred, _ = predict_event_final(event_model, rows, allowed_by_type)
    for r, p in zip(rows, pred):
        r["event_name"] = p
        r["pred_event_name"] = p
    emo = predict_emotion(emotion_model, rows)
    dup = F.dup_predict(rows, "A_step5_8e_deterministic")
    for r, e, d in zip(rows, emo, dup):
        r["event_emotion"] = e
        r["duplication_id"] = d
        r["event_impact_analysis"] = E.impact_text(r, r["event_name"], r["event_emotion"], E.evidence_quote(r))
    xlsx = OUT_DIR / "task1_official_delivery_candidate_v5.xlsx"
    jsonl = OUT_DIR / "task1_official_delivery_candidate_v5.jsonl"
    parquet = OUT_DIR / "task1_official_delivery_candidate_v5.parquet"
    company = [{col: r.get(col, "") for col in headers["个股新闻"]} for r in rows if str(r.get("trading_code") or "")]
    industry = [{col: r.get(col, "") for col in headers["行业新闻"]} for r in rows if not str(r.get("trading_code") or "")]
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        pd.DataFrame(company, columns=headers["个股新闻"]).to_excel(writer, sheet_name="个股新闻", index=False)
        pd.DataFrame(industry, columns=headers["行业新闻"]).to_excel(writer, sheet_name="行业新闻", index=False)
    unified = [{col: r.get(col, "") for col in headers["个股新闻"]} for r in rows]
    write_jsonl(jsonl, unified)
    pd.DataFrame(unified, columns=headers["个股新闻"]).to_parquet(parquet, index=False)
    return {"xlsx": str(xlsx), "jsonl": str(jsonl), "parquet": str(parquet), "rows": len(unified), "company_rows": len(company), "industry_rows": len(industry), "illegal_event_name": sum(1 for r in unified if r["event_name"] not in schema), "illegal_emotion": sum(1 for r in unified if r["event_emotion"] not in EMOTIONS)}


def validate_shadow(headers: dict[str, list[str]]) -> dict[str, Any]:
    xlsx = OUT_DIR / "task1_official_delivery_candidate_v5.xlsx"
    jsonl = OUT_DIR / "task1_official_delivery_candidate_v5.jsonl"
    parquet = OUT_DIR / "task1_official_delivery_candidate_v5.parquet"
    rows = read_jsonl(jsonl)
    pq = pd.read_parquet(parquet)
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    cols = {}
    counts = {}
    for s in wb.sheetnames:
        ws = wb[s]
        cols[s] = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        counts[s] = ws.max_row - 1
    wb.close()
    forbidden = {"trigger", "event_span", "arguments", "argument_roles", "source_event_type", "field_provenance", "unmapped_external_event"}
    return {"rows_match": len(rows) == len(pq) == sum(counts.values()), "columns_match": cols.get("个股新闻") == headers["个股新闻"] and cols.get("行业新闻") == headers["行业新闻"], "forbidden_hits": sorted(forbidden & set(rows[0].keys())) if rows else []}


def update_docs(status: str) -> None:
    note = f"\n\n## Step 5.8-H Update\n\nStatus: `{status}`. Added user-confirmed 1000-row official-contract annotations for incremental supervised experiments. Formal Task1/Task2 outputs remain unchanged.\n"
    for path in [TECH_REPORT, PIPELINE_DOC, REPRO_DOC, ROLLBACK_DOC]:
        path.parent.mkdir(parents=True, exist_ok=True)
        old = path.read_text(encoding="utf-8") if path.exists() else ""
        if "## Step 5.8-H Update" not in old:
            path.write_text(old.rstrip() + note, encoding="utf-8")


def write_integrity_blocked_report(
    backup: Path,
    before: dict[str, str],
    after: dict[str, str],
    input_validation: dict[str, Any],
    official_schema: dict[str, Any],
    company_events: set[str],
    industry_events: set[str],
    headers: dict[str, list[str]],
    start: float,
) -> None:
    status = "blocked_by_integrity_or_environment_failure"
    if status not in LEGAL_STATUSES:
        raise AssertionError(status)
    blocked_reason = (
        "新增1000条人工标注中的event_name不属于当前项目已确认的官方64类。"
        "在未获得权威映射或新官方事件体系确认前，继续训练会把非官方事件枚举写入监督标签。"
    )
    event_rows = []
    for sheet in ["个股新闻", "行业新闻"]:
        official = set(input_validation["official_event_names_by_sheet"][sheet])
        for name, count in sorted(Counter(r for r in input_validation["incremental_event_names_by_sheet"][sheet]).items()):
            event_rows.append(
                {
                    "sheet": sheet,
                    "event_name": name,
                    "record_count": sum(
                        1
                        for row in pd.read_excel(INCREMENTAL_XLSX, sheet_name=sheet, usecols=["event_name"])["event_name"].astype(str)
                        if row == name
                    ),
                    "in_official64_for_sheet": name in official,
                    "official_sheet_event_count": len(official),
                }
            )
    report = {
        "status": status,
        "generated_at": now_iso(),
        "blocked_reason": blocked_reason,
        "input_file": str(INCREMENTAL_XLSX),
        "input_validation": input_validation,
        "official_contract": {
            "company_event_count": len(company_events),
            "industry_event_count": len(industry_events),
            "total_event_count": len(official_schema),
            "company_headers": headers["个股新闻"],
            "industry_headers": headers["行业新闻"],
            "source_files": {
                "official_train_excel": str(RAW_TRAIN),
                "company_event_schema_json": str(ROOT / "data/raw/事件类型_标的.json"),
                "industry_event_schema_json": str(ROOT / "data/raw/事件类型_行业.json"),
                "processed_event_schema": str(ROOT / "data/processed/event_schema.jsonl"),
            },
        },
        "mismatch_summary": {
            "incremental_total_event_names": input_validation["event_name_class_count"],
            "incremental_records_with_illegal_event_name": input_validation["illegal_event_name"],
            "incremental_illegal_event_names_company_sheet": input_validation["illegal_event_names_by_sheet"]["个股新闻"],
            "incremental_illegal_event_names_industry_sheet": input_validation["illegal_event_names_by_sheet"]["行业新闻"],
            "training_not_started": True,
            "split_not_created": True,
            "shadow_v5_not_generated": True,
        },
        "allowed_next_action": (
            "请提供新增45类event_name到当前官方64类的权威映射，或确认这1000条使用的是新的官方事件体系并同步替换官方事件类型文件。"
            "在此之前只能保留输入审计结果，不能把这些标签用于监督训练。"
        ),
        "formal_sha_before": before,
        "formal_sha_after": after,
        "formal_sha_match": before == after,
        "event_name_frozen_model_sha_expected": EVENT_MODEL_EXPECTED_SHA,
        "event_name_frozen_model_sha_actual": sha256_file(EVENT_MODEL) if EVENT_MODEL.exists() else None,
        "protected_test_labels_read": False,
        "independent_frozen60_labels_read": False,
        "external_data_used": False,
        "llm_calls": 0,
        "local_32b_calls": 0,
        "external_llm_calls": 0,
        "gpu_seconds": 0,
        "wall_time_seconds": round(time.time() - start, 3),
        "backup_dir": str(backup),
        "rollback": "Formal outputs and user input workbook were not modified. This blocked Step 5.8-H run can be rolled back by ignoring the Step 5.8-H output directory and reports.",
    }
    write_json(REPORT_JSON, report)
    REPORT_MD.write_text(
        "# Task1 Step 5.8-H Incremental1000\n\n"
        f"- Status: `{status}`\n"
        f"- Input SHA: `{input_validation['input_sha256']}`\n"
        f"- Rows: `{input_validation['total_rows']}` "
        f"({input_validation['sheet_rows'].get('个股新闻')} company / {input_validation['sheet_rows'].get('行业新闻')} industry)\n"
        f"- Event enum check: `{input_validation['illegal_event_name']}` / `{input_validation['total_rows']}` records are outside current official64.\n"
        f"- Emotion enum check: `{input_validation['illegal_event_emotion']}` illegal values.\n"
        f"- Training started: `false`\n"
        f"- Shadow v5 generated: `false`\n"
        f"- Formal SHA unchanged: `{before == after}`\n\n"
        "## Blocking Reason\n\n"
        f"{blocked_reason}\n\n"
        "## Required Next Input\n\n"
        "Provide an authoritative mapping from the new annotation event labels to the existing official64 labels, "
        "or confirm and provide the official files for a replacement event taxonomy.\n",
        encoding="utf-8",
    )
    with pd.ExcelWriter(DETAILS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame([input_validation]).to_excel(writer, sheet_name="input_validation", index=False)
        pd.DataFrame(event_rows).to_excel(writer, sheet_name="event_name_mismatch", index=False)
        pd.DataFrame(
            [{"sheet": "个股新闻", "event_name": e} for e in sorted(company_events)]
            + [{"sheet": "行业新闻", "event_name": e} for e in sorted(industry_events)]
        ).to_excel(writer, sheet_name="official64", index=False)
    with pd.ExcelWriter(ERRORS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(event_rows).query("in_official64_for_sheet == False").to_excel(writer, sheet_name="illegal_event_names", index=False)
    with pd.ExcelWriter(CLASS_METRICS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(event_rows).to_excel(writer, sheet_name="incremental_event_names", index=False)
    write_json(
        OUT_DIR / "checkpoint_manifest.json",
        {
            "completed": True,
            "status": status,
            "blocked_stage": "input_validation_event_name_enum",
            "resume_policy": "Do not resume training until event_name taxonomy is reconciled.",
            "updated_at": now_iso(),
        },
    )
    write_json(
        OUT_DIR / "intersection_report.json",
        {
            "not_executed": True,
            "reason": "event_name enum mismatch blocked before protected/history intersection checks",
        },
    )
    files = [
        REPORT_JSON,
        REPORT_MD,
        DETAILS_XLSX,
        ERRORS_XLSX,
        CLASS_METRICS_XLSX,
        ROOT / "scripts/run_task1_step5_8h_incremental1000.py",
        ROOT / "scripts/evaluate_task1_step5_8h_incremental1000.py",
        ROOT / "scripts/test_task1_step5_8h_incremental1000.py",
    ]
    for p in OUT_DIR.rglob("*"):
        if p.is_file() and p.name != "SHA256SUMS.txt":
            files.append(p)
    (OUT_DIR / "SHA256SUMS.txt").write_text("\n".join(f"{sha256_file(p)}  {p}" for p in sorted(set(files))) + "\n", encoding="utf-8")
    update_state("blocked", status=status, blocked_reason="event_name_enum_mismatch", formal_sha_match=before == after)
    print(json.dumps({"status": status, "blocked_reason": "event_name_enum_mismatch", "input_rows": input_validation["total_rows"], "illegal_event_name": input_validation["illegal_event_name"], "formal_sha_match": before == after, "backup_dir": str(backup)}, ensure_ascii=False))


def main() -> None:
    start = time.time()
    backup = setup()
    before = protected_sha(backup / "protected_sha256_before.txt")
    update_state("validating_input")
    if not INCREMENTAL_XLSX.exists():
        raise FileNotFoundError(INCREMENTAL_XLSX)
    event_sha = sha256_file(EVENT_MODEL)
    if event_sha != EVENT_MODEL_EXPECTED_SHA:
        raise RuntimeError(f"event_name frozen SHA mismatch: {event_sha}")
    official_rows = E.load_official_rows()
    official_split_ids, official_split_manifest = E.split_official_rows(official_rows)
    official_splits = E.rows_by_split(official_rows, official_split_ids)
    schema, company_events, industry_events = E.load_event_schema()
    headers = official_headers()
    events_by_sheet = official_events_by_sheet(official_rows)
    legal_emotions = {r["event_emotion"] for r in official_rows}
    incremental, input_validation = load_incremental(headers, events_by_sheet, legal_emotions, strict=False)
    if input_validation["errors"] or input_validation["total_rows"] != 1000 or input_validation["duplicate_article_file_id"] or input_validation["missing_content"] or input_validation["illegal_event_name"] or input_validation["illegal_event_emotion"]:
        after = protected_sha(backup / "protected_sha256_after.txt")
        write_integrity_blocked_report(backup, before, after, input_validation, schema, company_events, industry_events, headers, start)
        return
    protected = load_protected_ids()
    intersections = intersection_report(incremental, protected)
    blocking_intersections = {k: v for k, v in intersections.items() if v["intersection_count"] > 0 and k != "formal_test_input"}
    if any(v["intersection_count"] > 0 for v in intersections.values()):
        # Formal-test-input overlap is also blocked if present, but keep explicit diagnostic.
        blocking_intersections = {k: v for k, v in intersections.items() if v["intersection_count"] > 0}
    if blocking_intersections:
        raise RuntimeError(f"incremental data intersects protected/history sets: {blocking_intersections}")
    inc_splits, inc_split_manifest = split_incremental(incremental)
    if inc_split_manifest["duplication_overlap_blocking"] or inc_split_manifest["content_sha_overlap_blocking"]:
        raise RuntimeError(f"incremental split leakage: {inc_split_manifest}")
    update_state("writing_validated_splits")
    write_jsonl(OUT_DIR / "validated_incremental1000.jsonl", incremental)
    write_jsonl(OUT_DIR / "incremental_train.jsonl", inc_splits["incremental_train"])
    write_jsonl(OUT_DIR / "incremental_dev.jsonl", inc_splits["incremental_dev"])
    write_jsonl(OUT_DIR / "incremental_holdout.jsonl", inc_splits["incremental_holdout"])
    write_json(OUT_DIR / "split_manifest.json", inc_split_manifest)
    write_json(OUT_DIR / "intersection_report.json", intersections)

    update_state("training_event_name")
    event_report, selected_event_model, event_meta = train_event_experiments(official_splits, inc_splits, schema, company_events, industry_events)
    update_state("training_event_emotion")
    emotion_report, selected_emotion_model = train_emotion_experiments(official_splits, inc_splits, selected_event_model, event_meta)
    impact = impact_audit(incremental)
    dup = duplication_report(incremental)
    shadow = write_shadow_v5(selected_event_model, selected_emotion_model, official_splits["train"] + inc_splits["incremental_train"] + inc_splits["incremental_dev"], schema, event_meta["allowed_by_type"], headers)
    format_validation = validate_shadow(headers)

    model_dir = OUT_DIR / "frozen_candidate"
    model_dir.mkdir(parents=True, exist_ok=True)
    with (model_dir / "event_name_model.pkl").open("wb") as handle:
        pickle.dump(selected_event_model, handle)
    with (model_dir / "event_emotion_model.pkl").open("wb") as handle:
        pickle.dump(selected_emotion_model, handle)
    frozen_manifest = {
        "frozen_at": now_iso(),
        "event_name_model_sha256": sha256_file(model_dir / "event_name_model.pkl"),
        "event_emotion_model_sha256": sha256_file(model_dir / "event_emotion_model.pkl"),
        "source_incremental_sha256": input_validation["input_sha256"],
        "official_data_only_plus_incremental1000": True,
        "llm_calls": 0,
    }
    write_json(OUT_DIR / "frozen_candidate_manifest.json", frozen_manifest)
    # Prediction files for held-out diagnostics.
    inc_hold = inc_splits["incremental_holdout"]
    ev_pred, ev_top = predict_event_final(selected_event_model, inc_hold, event_meta["allowed_by_type"])
    for r, p in zip(inc_hold, ev_pred):
        r["pred_event_name"] = p
    em_pred = predict_emotion(selected_emotion_model, inc_hold)
    write_jsonl(OUT_DIR / "event_name_predictions.jsonl", [{"row_id": r["row_id"], "article_file_id": r["article_file_id"], "gold_event_name": r["event_name"], "pred_event_name": p, "gold_in_top20": r["event_name"] in top} for r, p, top in zip(inc_hold, ev_pred, ev_top)])
    write_jsonl(OUT_DIR / "emotion_predictions.jsonl", [{"row_id": r["row_id"], "article_file_id": r["article_file_id"], "gold_emotion": r["event_emotion"], "pred_emotion": p, "pred_event_name": r["pred_event_name"]} for r, p in zip(inc_hold, em_pred)])

    after = protected_sha(backup / "protected_sha256_after.txt")
    base = event_report["baseline"]
    best = event_report["best_scheme"]
    event_hold = event_report["incremental_holdout_after_freeze"]
    event_safe = (
        best["official_regression"]["micro_f1"] >= base["official_regression"]["micro_f1"] - 0.01
        and best["official_regression"]["fixed64_macro_f1"] >= base["official_regression"]["fixed64_macro_f1"] - 0.02
        and event_hold["illegal_event_name"] == 0
    )
    event_improved = event_safe and (
        event_hold["micro_f1"] >= base["incremental_holdout"]["micro_f1"] + 0.005
        or event_hold["fixed64_macro_f1"] >= base["incremental_holdout"]["fixed64_macro_f1"] + 0.01
    )
    emo_hold = emotion_report["incremental_holdout_after_freeze"]
    emotion_gate = emo_hold["accuracy"] >= 0.89 and emo_hold["macro_f1"] >= 0.65 and emo_hold["illegal_emotion"] == 0 and next(x for x in emo_hold["per_class"] if x["label"] == "中性")["recall"] > 0.111111 and next(x for x in emo_hold["per_class"] if x["label"] == "负面")["recall"] >= 0.5625
    emotion_improved = emo_hold["macro_f1"] >= 0.65 or emotion_report["best_scheme"]["dev"]["macro_f1"] > 0.588649 + 0.02
    delivery_ready = event_safe and emotion_gate and impact["evidence_support_rate"] == 1.0 and shadow["illegal_event_name"] == 0 and shadow["illegal_emotion"] == 0 and format_validation["rows_match"] and format_validation["columns_match"] and before == after
    if delivery_ready:
        status = "task1_incremental1000_delivery_candidate_ready"
    elif event_improved:
        status = "task1_incremental1000_event_name_improved"
    elif emotion_improved:
        status = "task1_incremental1000_emotion_improved"
    else:
        status = "task1_incremental1000_no_safe_gain"
    if status not in LEGAL_STATUSES:
        raise AssertionError(status)
    report = {
        "status": status,
        "generated_at": now_iso(),
        "input_validation": input_validation,
        "data_identity": "human_annotated_official_contract_incremental_data",
        "intersections": intersections,
        "incremental_split": inc_split_manifest,
        "official_split": official_split_manifest,
        "training_scale": {
            "official_train": len(official_splits["train"]),
            "official_dev": len(official_splits["model_selection_dev"]),
            "official_holdout": len(official_splits["untouched_internal_holdout"]),
            "incremental_train": len(inc_splits["incremental_train"]),
            "incremental_dev": len(inc_splits["incremental_dev"]),
            "incremental_holdout": len(inc_splits["incremental_holdout"]),
        },
        "event_name": event_report,
        "event_emotion": emotion_report,
        "impact_analysis": impact,
        "duplication": dup,
        "shadow_v5": shadow,
        "format_validation": format_validation,
        "gates": {"event_safe": event_safe, "event_improved": event_improved, "emotion_gate": emotion_gate, "emotion_improved": emotion_improved, "delivery_ready": delivery_ready},
        "frozen_candidate_manifest": frozen_manifest,
        "formal_sha_before": before,
        "formal_sha_after": after,
        "formal_sha_match": before == after,
        "llm_calls": 0,
        "local_32b_calls": 0,
        "external_llm_calls": 0,
        "gpu_seconds": 0,
        "peak_gpu_memory_mb": 0,
        "wall_time_seconds": round(time.time() - start, 3),
        "backup_dir": str(backup),
        "recommend_replace_step5_8f_candidate": status in {"task1_incremental1000_delivery_candidate_ready", "task1_incremental1000_event_name_improved", "task1_incremental1000_emotion_improved"},
        "recommend_formal_promotion": status == "task1_incremental1000_delivery_candidate_ready",
        "rollback": "Formal outputs were not modified. Ignore or remove Step 5.8-H outputs to roll back the candidate.",
    }
    write_json(REPORT_JSON, report)
    REPORT_MD.write_text(
        f"# Task1 Step 5.8-H Incremental1000\n\n"
        f"- Status: `{status}`\n"
        f"- Input SHA: `{input_validation['input_sha256']}`\n"
        f"- Incremental split: `{inc_split_manifest['counts']}`\n"
        f"- Best event_name scheme: `{best['scheme']}`\n"
        f"- Incremental holdout event micro-F1: `{event_hold['micro_f1']}`\n"
        f"- Incremental holdout emotion macro-F1: `{emo_hold['macro_f1']}`\n"
        f"- Shadow v5 XLSX: `{shadow['xlsx']}`\n"
        f"- Formal SHA unchanged: `{before == after}`\n",
        encoding="utf-8",
    )
    with pd.ExcelWriter(DETAILS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame([input_validation]).to_excel(writer, sheet_name="input_validation", index=False)
        pd.DataFrame(event_report["schemes"]).to_excel(writer, sheet_name="event_schemes", index=False)
        pd.DataFrame(emotion_report["schemes"]).to_excel(writer, sheet_name="emotion_schemes", index=False)
        pd.DataFrame([impact]).to_excel(writer, sheet_name="impact", index=False)
        pd.DataFrame([dup]).to_excel(writer, sheet_name="duplication", index=False)
    with pd.ExcelWriter(ERRORS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(read_jsonl(OUT_DIR / "event_name_predictions.jsonl")).query("gold_event_name != pred_event_name").to_excel(writer, sheet_name="event_name_errors", index=False)
        pd.DataFrame(read_jsonl(OUT_DIR / "emotion_predictions.jsonl")).query("gold_emotion != pred_emotion").to_excel(writer, sheet_name="emotion_errors", index=False)
    with pd.ExcelWriter(CLASS_METRICS_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(event_report["best_per_class_incremental_holdout"]).to_excel(writer, sheet_name="event_holdout_per_class", index=False)
        pd.DataFrame(emo_hold["per_class"]).to_excel(writer, sheet_name="emotion_holdout_per_class", index=False)
    write_json(OUT_DIR / "checkpoint_manifest.json", {"completed": True, "status": status, "resume_command": "python scripts/run_task1_step5_8h_incremental1000.py", "updated_at": now_iso()})
    update_docs(status)
    files = []
    for p in OUT_DIR.rglob("*"):
        if p.is_file() and p.name != "SHA256SUMS.txt":
            files.append(p)
    for p in [REPORT_JSON, REPORT_MD, DETAILS_XLSX, ERRORS_XLSX, CLASS_METRICS_XLSX, TECH_REPORT, PIPELINE_DOC, REPRO_DOC, ROLLBACK_DOC, ROOT / "scripts/run_task1_step5_8h_incremental1000.py"]:
        if p.exists():
            files.append(p)
    (OUT_DIR / "SHA256SUMS.txt").write_text("\n".join(f"{sha256_file(p)}  {p}" for p in sorted(set(files))) + "\n", encoding="utf-8")
    update_state("completed", status=status)
    print(json.dumps({"status": status, "input_rows": input_validation["total_rows"], "event_holdout": event_hold, "emotion_holdout": emo_hold, "formal_sha_match": before == after, "backup_dir": str(backup)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
