#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-F: confirm structure annotations and train limited baselines."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import shutil
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

AUDITED_XLSX = ROOT / "reports/task1_step5_7e_eventmention_ai_audited.xlsx"
CONFIRMED_XLSX = ROOT / "reports/task1_step5_7e_eventmention_human_confirmed.xlsx"
REPORT_JSON = ROOT / "reports/task1_step5_7f_structure_training_report.json"
REPORT_MD = ROOT / "reports/task1_step5_7f_structure_training_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7f_structure_training_details.xlsx"
OUT_DIR = ROOT / "outputs/task1_step5_7f_structure_training_v1"
CONFIG_PATH = ROOT / "configs/task1_structure_training_v1.yaml"

STEP57E_MANIFEST = ROOT / "outputs/task1_step5_7e_human_confirmation/eventmention200_manifest.json"
STEP57C = ROOT / "outputs/task1_step5_7c_eventmention_annotation_pilot/validated_annotation_drafts.jsonl"
STEP57D = ROOT / "outputs/task1_step5_7d_contract_revision/diagnostic40_results.jsonl"
CONFIRMED80 = ROOT / "outputs/task1_structure_validation_round1/minimal_human_review_package_confirmed.jsonl"
MIGRATED80 = ROOT / "outputs/task1_step5_7c_eventmention_annotation_pilot/confirmed80_eventmention_migration.jsonl"
MIGRATED80_REVIEW_XLSX = ROOT / "outputs/task1_step5_7c_eventmention_annotation_pilot/confirmed80_contract_review.xlsx"
TRAIN_LABELS = ROOT / "data/processed/labeled/news_train.jsonl"
EVENT_SCHEMA = ROOT / "data/processed/event_schema.jsonl"
CLASSIFIER_TRAIN_ARTICLES = ROOT / "outputs/task1_step5_6d_classifier_retraining/train_articles.jsonl"
CLASSIFIER_DEV_ARTICLES = ROOT / "outputs/task1_step5_6d_classifier_retraining/dev_articles.jsonl"
CLASSIFIER_HOLDOUT_ARTICLES = ROOT / "outputs/task1_step5_6d_classifier_retraining/holdout_articles.jsonl"
GOLD_BLIND_INPUT = ROOT / "outputs/task1_gold_shadow_eval_v1/blind_input.jsonl"

CU_PYTHON = Path("/root/miniconda3/envs/dianjin-cu128/bin/python")
LOCAL_MACBERT = ROOT / "outputs/task1_step5_4_dev_baseline_v1/pretrained/hfl_chinese_macbert_base"

EXPECTED_OLD_PENDING = {
    "NEWS_MULTISOURCE_D-1120797906500",
    "NEWS_MULTISOURCE_D-1126130257189",
    "NEWS_MULTISOURCE_D-1127695210485",
}
CONFIRMATION_NOTE = "用户已逐条审核并确认AI辅助审核结果无误。"
ANNOTATION_SOURCE = "ai_assisted_human_confirmed"
DIR_LABELS = ["positive", "negative", "neutral", "unclear"]
ROLE_LABELS = ["subject_company", "event_time", "object_target", "trigger"]
SPAN_FIELDS = ["trigger", "event_span", "subject_company", "event_time", "object_target"]
HUMAN_FIELDS = [
    "human_trigger",
    "human_event_span",
    "human_subject_company",
    "human_event_time",
    "human_event_time_normalized",
    "human_object_target",
    "human_direction",
    "human_arguments_json",
    "human_argument_roles_json",
    "reviewer_confirmed",
    "reviewer_note",
]
PROTECTED_FILES = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
    CONFIRMED80,
    MIGRATED80,
    MIGRATED80_REVIEW_XLSX,
    STEP57C,
    STEP57D,
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/evaluation_summary.json",
    ROOT / "reports/task1_step5_6f_classifier_consolidation_report.json",
]


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def text_norm(value: Any) -> str:
    text = "" if value is None else str(value)
    return "".join(ch for ch in text.strip().lower() if not ch.isspace())


def parse_json(value: Any) -> tuple[Any, str]:
    try:
        return json.loads("" if value is None else str(value)), ""
    except Exception as exc:
        return None, f"{type(exc).__name__}:{exc}"


def article_ids(path: Path) -> set[str]:
    ids: set[str] = set()
    for row in read_jsonl(path):
        aid = str(row.get("article_id", "")).strip()
        if aid:
            ids.add(aid)
    return ids


def read_article_ids_only(path: Path) -> set[str]:
    ids: set[str] = set()
    if not path.exists():
        return ids
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            obj = json.loads(line)
            aid = str(obj.get("article_id", "")).strip()
            if aid:
                ids.add(aid)
    return ids


def official_event_names() -> set[str]:
    names: set[str] = set()
    label_count = 0
    for row in read_jsonl(TRAIN_LABELS):
        for label in row.get("labels") or []:
            name = str(label.get("event_name", "")).strip()
            if name:
                names.add(name)
                label_count += 1
    if len(names) != 64 or label_count != 2081:
        raise RuntimeError(f"unexpected training event_name contract: classes={len(names)} labels={label_count}")
    return names


def protected_paths() -> list[Path]:
    paths = list(PROTECTED_FILES)
    for base in [
        ROOT / "outputs/task1_step5_6d_classifier_retraining/experiment_C/best_model",
        ROOT / "outputs/task1_step5_4_dev_baseline_v1/pretrained/hfl_chinese_macbert_base",
    ]:
        if base.exists():
            paths.extend(sorted(path for path in base.rglob("*") if path.is_file()))
    return sorted(set(paths), key=rel)


def sha_snapshot(paths: list[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in paths if path.exists()}


def create_backup_dir() -> Path:
    backup = Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7f_structure_training_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup.mkdir(parents=True, exist_ok=False)
    return backup


def write_sha_snapshot(path: Path, snapshot: dict[str, str]) -> None:
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(snapshot.items())), encoding="utf-8")


def process_check() -> dict[str, Any]:
    proc = subprocess.run(
        "ps -eo pid,ppid,stat,etime,cmd | rg -i 'step5_7f|structure_training|vllm serve|DianJin-R1-32B|run_task1_step5_7f' | rg -v 'rg -i|codex-linux-sandbox|bwrap' || true",
        shell=True,
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=20,
    )
    return {"returncode": proc.returncode, "raw": proc.stdout.splitlines()}


def workbook_rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"missing workbook sheet: {sheet_name}")
    ws = wb[sheet_name]
    header = [str(cell.value or "") for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for cells in ws.iter_rows(min_row=2):
        values = [cell.value for cell in cells[: len(header)]]
        if all(value is None or value == "" for value in values):
            continue
        rows.append({key: "" if value is None else value for key, value in zip(header, values)})
    wb.close()
    return header, rows


def locate_text(content: str, text: str, anchor: dict[str, int] | None = None) -> dict[str, Any]:
    value = "" if text is None else str(text)
    if value == "":
        return {"text": "", "start": -1, "end": -1}
    positions: list[int] = []
    start = 0
    while True:
        pos = content.find(value, start)
        if pos < 0:
            break
        positions.append(pos)
        start = pos + 1
    if not positions:
        return {"text": value, "start": -1, "end": -1}
    if anchor and anchor.get("start", -1) >= 0:
        a_start = int(anchor["start"])
        a_end = int(anchor["end"])
        containing = [pos for pos in positions if pos <= a_start and a_end <= pos + len(value)]
        if containing:
            pos = containing[0]
        else:
            pos = min(positions, key=lambda item: abs(item - a_start))
    else:
        pos = positions[0]
    return {"text": value, "start": pos, "end": pos + len(value)}


def span_empty() -> dict[str, Any]:
    return {"text": "", "start": -1, "end": -1}


def time_empty(normalized: str = "未明确") -> dict[str, Any]:
    return {"text": "", "normalized": normalized or "未明确", "start": -1, "end": -1}


def strict_span(content: str, span: dict[str, Any], field: str) -> tuple[dict[str, Any], list[str]]:
    text = str(span.get("text", "") or "")
    try:
        start = int(span.get("start", -1))
        end = int(span.get("end", -1))
    except Exception:
        return (time_empty(span.get("normalized", "未明确")) if field == "event_time" else span_empty()), [f"{field}:offset_not_int"]
    if not text:
        return (time_empty(str(span.get("normalized", "未明确"))) if field == "event_time" else span_empty()), []
    if not (0 <= start <= end <= len(content)) or content[start:end] != text:
        return (time_empty(str(span.get("normalized", "未明确"))) if field == "event_time" else span_empty()), [f"{field}:invalid_offset"]
    out = {"text": text, "start": start, "end": end}
    if field == "event_time":
        out["normalized"] = str(span.get("normalized", "未明确") or "未明确")
    return out, []


def spans_from_human(row: dict[str, Any], prefix: str = "human") -> tuple[dict[str, Any], list[str]]:
    content = str(row.get("content", ""))
    errors: list[str] = []
    args, err = parse_json(row.get(f"{prefix}_arguments_json", ""))
    if err:
        errors.append(f"arguments_json_parse_failed:{err}")
        args = []
    if not isinstance(args, list):
        errors.append("arguments_json_not_list")
        args = []
    by_role: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for idx, arg in enumerate(args):
        if not isinstance(arg, dict):
            errors.append(f"argument_{idx}_not_object")
            continue
        role = str(arg.get("role", ""))
        text = str(arg.get("text", "") or "")
        if role not in ROLE_LABELS:
            errors.append(f"argument_{idx}_illegal_role")
            continue
        try:
            start = int(arg.get("start", -1))
            end = int(arg.get("end", -1))
        except Exception:
            start, end = -1, -1
        if text and (not (0 <= start <= end <= len(content)) or content[start:end] != text):
            errors.append(f"argument_{idx}_offset_text_mismatch")
        by_role[role].append({"text": text, "start": start, "end": end})
    trigger = by_role.get("trigger", [None])[0] or locate_text(content, row.get(f"{prefix}_trigger", ""))
    trigger, field_errors = strict_span(content, trigger, "trigger")
    errors.extend(field_errors)
    event_span = locate_text(content, str(row.get(f"{prefix}_event_span", "") or ""), trigger if trigger.get("text") else None)
    event_span, field_errors = strict_span(content, event_span, "event_span")
    errors.extend(field_errors)
    subject = by_role.get("subject_company", [None])[0] or locate_text(content, row.get(f"{prefix}_subject_company", ""))
    subject, field_errors = strict_span(content, subject, "subject_company")
    errors.extend(field_errors)
    time_text = str(row.get(f"{prefix}_event_time", "") or "")
    event_time = by_role.get("event_time", [None])[0] or locate_text(content, time_text)
    event_time["normalized"] = str(row.get(f"{prefix}_event_time_normalized", "") or "未明确")
    event_time, field_errors = strict_span(content, event_time, "event_time")
    errors.extend(field_errors)
    obj = by_role.get("object_target", [None])[0] or locate_text(content, row.get(f"{prefix}_object_target", ""))
    obj, field_errors = strict_span(content, obj, "object_target")
    errors.extend(field_errors)
    if trigger.get("text") and event_span.get("text"):
        if not (event_span["start"] <= trigger["start"] < trigger["end"] <= event_span["end"]):
            errors.append("trigger_not_inside_event_span")
    direction = str(row.get(f"{prefix}_direction", "") or "")
    if direction not in DIR_LABELS:
        errors.append("illegal_direction")
        direction = "unclear"
    roles, err = parse_json(row.get(f"{prefix}_argument_roles_json", ""))
    if err:
        errors.append(f"argument_roles_json_parse_failed:{err}")
    elif not isinstance(roles, dict):
        errors.append("argument_roles_json_not_object")
    arguments: list[dict[str, Any]] = []
    for role in ROLE_LABELS:
        span = {"subject_company": subject, "event_time": event_time, "object_target": obj, "trigger": trigger}[role]
        if span.get("text"):
            arguments.append({"text": span["text"], "role": role, "start": span["start"], "end": span["end"]})
    return (
        {
            "trigger": trigger,
            "event_span": event_span,
            "subject_company": subject,
            "event_time": event_time,
            "object_target": obj,
            "direction": direction,
            "arguments": arguments,
        },
        sorted(set(errors)),
    )


def validate_audited_input(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    _, rows = workbook_rows(path, "200条主审核表")
    errors: list[str] = []
    if len(rows) != 200:
        errors.append(f"main_row_count:{len(rows)}")
    ids = [str(row.get("event_mention_id", "")) for row in rows]
    if len(set(ids)) != 200:
        errors.append("event_mention_id_not_unique_200")
    if any(not item for item in ids):
        errors.append("event_mention_id_missing")
    event_names = official_event_names()
    for row in rows:
        mid = str(row.get("event_mention_id", ""))
        if row.get("event_name") not in event_names:
            errors.append(f"{mid}:event_name_outside_64")
        for field in [
            "human_trigger",
            "human_event_span",
            "human_subject_company",
            "human_event_time_normalized",
            "human_direction",
            "human_arguments_json",
            "human_argument_roles_json",
        ]:
            if str(row.get(field, "") or "") == "":
                errors.append(f"{mid}:{field}_empty")
        spans, span_errors = spans_from_human(row, "human")
        errors.extend(f"{mid}:{item}" for item in span_errors)
        if not spans["trigger"]["text"] or not spans["event_span"]["text"]:
            errors.append(f"{mid}:trigger_or_event_span_empty")
    _, old_rows = workbook_rows(path, "旧80条契约确认")
    old_pending = [row for row in old_rows if str(row.get("article_id", "")) in EXPECTED_OLD_PENDING]
    if {str(row.get("article_id", "")) for row in old_pending} != EXPECTED_OLD_PENDING:
        errors.append("old_pending_three_missing")
    return {
        "status": "passed" if not errors else "failed",
        "errors": errors[:200],
        "error_count": len(errors),
        "main_rows": len(rows),
        "unique_event_mention_ids": len(set(ids)),
        "old_pending_rows": len(old_pending),
    }


def ensure_column(ws: Any, header: list[str], name: str) -> int:
    if name in header:
        return header.index(name) + 1
    ws.cell(1, len(header) + 1).value = name
    header.append(name)
    return len(header)


def append_note(note: Any) -> str:
    base = str(note or "").strip()
    if CONFIRMATION_NOTE in base:
        return base
    return f"{base} {CONFIRMATION_NOTE}".strip()


def fill_old_pending_from_ai(row_values: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    out = dict(row_values)
    notes: list[str] = []
    mapping = {
        "human_trigger": "ai_trigger",
        "human_event_span": "ai_event_span",
        "human_subject_company": "ai_subject_company",
        "human_event_time": "ai_event_time",
        "human_event_time_normalized": "ai_event_time_normalized",
        "human_direction": "ai_direction",
        "human_arguments_json": "ai_arguments_json",
        "human_argument_roles_json": "ai_argument_roles_json",
    }
    for human_field, ai_field in mapping.items():
        out[human_field] = out.get(human_field, "") or out.get(ai_field, "")
    object_text = str(out.get("human_object_target") or out.get("ai_object_target") or "")
    if object_text and object_text in str(out.get("content", "") or ""):
        out["human_object_target"] = object_text
    else:
        if object_text:
            notes.append("旧契约AI object_target无法逐字定位正文，正式确认版未复制该可选字段。")
        out["human_object_target"] = ""
    return out, notes


def generate_confirmed_workbook(input_path: Path, output_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(input_path)
    confirmed_at = now()
    main = wb["200条主审核表"]
    main_header = [str(cell.value or "") for cell in main[1]]
    main_source_col = ensure_column(main, main_header, "annotation_source")
    main_confirmed_col = ensure_column(main, main_header, "confirmed_at")
    main_idx = {name: i + 1 for i, name in enumerate(main_header)}
    main_count = 0
    for row_idx in range(2, main.max_row + 1):
        if not main.cell(row_idx, main_idx["event_mention_id"]).value:
            continue
        main.cell(row_idx, main_idx["reviewer_confirmed"]).value = "yes"
        main.cell(row_idx, main_idx["reviewer_note"]).value = append_note(main.cell(row_idx, main_idx["reviewer_note"]).value)
        main.cell(row_idx, main_source_col).value = ANNOTATION_SOURCE
        main.cell(row_idx, main_confirmed_col).value = confirmed_at
        main_count += 1
    old = wb["旧80条契约确认"]
    old_header = [str(cell.value or "") for cell in old[1]]
    old_source_col = ensure_column(old, old_header, "annotation_source")
    old_confirmed_col = ensure_column(old, old_header, "confirmed_at")
    old_idx = {name: i + 1 for i, name in enumerate(old_header)}
    old_confirmed = []
    old_errors: list[str] = []
    old_notes: dict[str, list[str]] = {}
    for row_idx in range(2, old.max_row + 1):
        aid = str(old.cell(row_idx, old_idx["article_id"]).value or "")
        if aid not in EXPECTED_OLD_PENDING:
            continue
        row_values = {name: old.cell(row_idx, col).value or "" for name, col in old_idx.items()}
        row_values, notes = fill_old_pending_from_ai(row_values)
        if notes:
            old_notes[aid] = notes
        spans, errors = spans_from_human(row_values, "human")
        missing = [
            field
            for field in ["human_trigger", "human_event_span", "human_subject_company", "human_event_time_normalized", "human_direction", "human_arguments_json", "human_argument_roles_json"]
            if str(row_values.get(field, "") or "") == ""
        ]
        if missing or errors:
            old_errors.append(f"{aid}:missing={missing}:errors={errors}")
            continue
        for field in [
            "human_trigger",
            "human_event_span",
            "human_subject_company",
            "human_event_time",
            "human_event_time_normalized",
            "human_object_target",
            "human_direction",
            "human_arguments_json",
            "human_argument_roles_json",
        ]:
            old.cell(row_idx, old_idx[field]).value = row_values.get(field, "")
        old.cell(row_idx, old_idx["reviewer_confirmed"]).value = "yes"
        note = append_note(row_values.get("reviewer_note", ""))
        if notes:
            note = f"{note} {' '.join(notes)}".strip()
        old.cell(row_idx, old_idx["reviewer_note"]).value = note
        old.cell(row_idx, old_source_col).value = ANNOTATION_SOURCE
        old.cell(row_idx, old_confirmed_col).value = confirmed_at
        old_confirmed.append(aid)
    if old_errors or set(old_confirmed) != EXPECTED_OLD_PENDING:
        raise RuntimeError(f"old pending contract confirmation failed: confirmed={old_confirmed} errors={old_errors}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {"main_confirmed_rows": main_count, "old_pending_confirmed_rows": len(old_confirmed), "old_pending_notes": old_notes, "confirmed_at": confirmed_at}


def run_existing_step57e_validator(workbook: Path) -> dict[str, Any]:
    cmd = [
        sys.executable,
        str(ROOT / "scripts/validate_task1_step5_7e_human_confirmation.py"),
        "--workbook",
        str(workbook),
        "--allow-filled",
    ]
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, timeout=120)
    try:
        payload = json.loads(proc.stdout)
    except Exception:
        payload = {"stdout": proc.stdout, "stderr": proc.stderr}
    payload["returncode"] = proc.returncode
    if proc.returncode != 0:
        raise RuntimeError(f"confirmed workbook failed existing validator: {json.dumps(payload, ensure_ascii=False)[:4000]}")
    return payload


def canonical_from_row(row: dict[str, Any], split: str, source: str) -> dict[str, Any]:
    spans, errors = spans_from_human(row, "human")
    return {
        "event_mention_id": str(row.get("event_mention_id", "")),
        "article_id": str(row.get("article_id", "")),
        "event_name": str(row.get("event_name", "")),
        "title": str(row.get("title", "")),
        "content": str(row.get("content", "")),
        "publish_time": str(row.get("publish_time", "")),
        "split": split,
        "annotation_source": ANNOTATION_SOURCE,
        "annotation_status": "confirmed",
        "reviewer_confirmed": "yes",
        "source_workbook": rel(CONFIRMED_XLSX),
        "trigger": spans["trigger"],
        "event_span": spans["event_span"],
        "subject_company": spans["subject_company"],
        "event_time": spans["event_time"],
        "object_target": spans["object_target"],
        "direction": spans["direction"],
        "arguments": spans["arguments"],
        "multi_event_article": str(row.get("multi_event_article", "")).lower() in {"yes", "true", "1"},
        "source_set": source,
        "validation_errors": errors,
        "ai_fields": {
            "trigger": row.get("ai_trigger", ""),
            "event_span": row.get("ai_event_span", ""),
            "subject_company": row.get("ai_subject_company", ""),
            "event_time": row.get("ai_event_time", ""),
            "event_time_normalized": row.get("ai_event_time_normalized", ""),
            "object_target": row.get("ai_object_target", ""),
            "direction": row.get("ai_direction", ""),
            "arguments_json": row.get("ai_arguments_json", ""),
            "argument_roles_json": row.get("ai_argument_roles_json", ""),
        },
    }


def strict_span_from_legacy(content: str, raw: Any, field: str) -> tuple[dict[str, Any], list[str]]:
    errors: list[str] = []
    if not isinstance(raw, dict):
        return (time_empty("未明确") if field == "event_time" else span_empty()), [f"{field}:legacy_span_not_object"]
    text = str(raw.get("text", "") or "")
    normalized = str(raw.get("normalized", "") or "未明确") if field == "event_time" else ""
    span = {"text": text, "start": raw.get("start", -1), "end": raw.get("end", -1)}
    if field == "event_time":
        span["normalized"] = normalized
    strict, field_errors = strict_span(content, span, field)
    if not field_errors:
        return strict, []
    # Historical 80 annotations predate strict offset requirements. If the text is
    # exactly present, recover a deterministic正文 offset for diagnostic use only.
    if text and text in content:
        located = locate_text(content, text)
        if field == "event_time":
            located["normalized"] = normalized
        strict, located_errors = strict_span(content, located, field)
        if not located_errors:
            return strict, [f"{field}:legacy_offset_recovered_by_exact_text"]
    if field == "event_time":
        return time_empty(normalized), field_errors
    return span_empty(), field_errors


def canonical_from_legacy_migration(row: dict[str, Any]) -> dict[str, Any]:
    content = str(row.get("content", ""))
    validation_errors: list[str] = []
    trigger, errors = strict_span_from_legacy(content, row.get("trigger", {}), "trigger")
    validation_errors.extend(errors)
    event_span, errors = strict_span_from_legacy(content, row.get("event_span", {}), "event_span")
    validation_errors.extend(errors)
    subject, errors = strict_span_from_legacy(content, row.get("subject_company", {}), "subject_company")
    validation_errors.extend(errors)
    event_time, errors = strict_span_from_legacy(content, row.get("event_time", {}), "event_time")
    validation_errors.extend(errors)
    obj, errors = strict_span_from_legacy(content, row.get("object_target", {}), "object_target")
    validation_errors.extend(errors)
    if trigger.get("text") and event_span.get("text"):
        if not (event_span["start"] <= trigger["start"] < trigger["end"] <= event_span["end"]):
            validation_errors.append("trigger_not_inside_event_span")
            event_span = {"text": content[min(event_span["start"], trigger["start"]):max(event_span["end"], trigger["end"])], "start": min(event_span["start"], trigger["start"]), "end": max(event_span["end"], trigger["end"])}
    direction = str(row.get("direction", "") or "unclear")
    if direction not in DIR_LABELS:
        validation_errors.append("legacy_illegal_direction")
        direction = "unclear"
    arguments: list[dict[str, Any]] = []
    for role, span in [("subject_company", subject), ("event_time", event_time), ("object_target", obj), ("trigger", trigger)]:
        if span.get("text"):
            arguments.append({"text": span["text"], "role": role, "start": span["start"], "end": span["end"]})
    return {
        "event_mention_id": str(row.get("event_mention_id", "")),
        "article_id": str(row.get("article_id", "")),
        "event_name": str(row.get("event_name", "")),
        "title": str(row.get("title", "")),
        "content": content,
        "publish_time": str(row.get("publish_time", "")),
        "split": "legacy80_diagnostic",
        "annotation_source": str(row.get("annotation_source", "legacy80_confirmed_migration")),
        "annotation_status": "confirmed",
        "reviewer_confirmed": "yes",
        "source_workbook": rel(MIGRATED80),
        "trigger": trigger,
        "event_span": event_span,
        "subject_company": subject,
        "event_time": event_time,
        "object_target": obj,
        "direction": direction,
        "arguments": arguments,
        "multi_event_article": False,
        "source_set": "legacy80_confirmed_diagnostic",
        "validation_errors": sorted(set(validation_errors)),
        "legacy_raw_structure": {
            "trigger": row.get("trigger", {}),
            "event_span": row.get("event_span", {}),
            "subject_company": row.get("subject_company", {}),
            "event_time": row.get("event_time", {}),
            "object_target": row.get("object_target", {}),
            "direction": row.get("direction", ""),
            "arguments": row.get("arguments", []),
            "contract_validation": row.get("contract_validation", {}),
            "contract_errors": row.get("contract_errors", []),
        },
        "ai_fields": {
            "trigger": trigger.get("text", ""),
            "event_span": event_span.get("text", ""),
            "subject_company": subject.get("text", ""),
            "event_time": event_time.get("text", ""),
            "event_time_normalized": event_time.get("normalized", "未明确"),
            "object_target": obj.get("text", ""),
            "direction": direction,
            "arguments_json": json.dumps(arguments, ensure_ascii=False),
            "argument_roles_json": json.dumps({arg["role"]: arg["text"] for arg in arguments}, ensure_ascii=False),
        },
    }


def choose_dev_groups(rows: list[dict[str, Any]], target_dev_mentions: int = 40, seed: int = 2026) -> set[str]:
    rng = random.Random(seed)
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[row["article_id"]].append(row)
    class_counts = Counter(row["event_name"] for row in rows)
    dir_counts = Counter(row["direction"] for row in rows)
    explicit_count = sum(row["event_time"]["text"] != "" for row in rows)
    explicit_ratio = explicit_count / max(1, len(rows))
    multi_ratio = sum(bool(row.get("multi_event_article")) for row in rows) / max(1, len(rows))
    selected: set[str] = set()
    selected_rows: list[dict[str, Any]] = []

    def score_group(aid: str) -> tuple[float, float]:
        trial = selected_rows + groups[aid]
        if len(trial) > target_dev_mentions + 3:
            return (-1e9, -1e9)
        class_div = len({row["event_name"] for row in trial})
        dir_div = len({row["direction"] for row in trial})
        explicit = sum(row["event_time"]["text"] != "" for row in trial) / max(1, len(trial))
        multi = sum(bool(row.get("multi_event_article")) for row in trial) / max(1, len(trial))
        rare_bonus = sum(1 / class_counts[row["event_name"]] for row in groups[aid])
        return (
            class_div * 5
            + dir_div * 3
            - abs(explicit - explicit_ratio) * 2
            - abs(multi - multi_ratio) * 2
            + rare_bonus
            - abs(target_dev_mentions - len(trial)) * 0.1,
            rng.random(),
        )

    while len(selected_rows) < target_dev_mentions:
        candidates = [aid for aid in groups if aid not in selected]
        if not candidates:
            break
        best = max(candidates, key=score_group)
        if score_group(best)[0] < -1e8:
            break
        selected.add(best)
        selected_rows.extend(groups[best])
    return selected


def build_training_data(confirmed_workbook: Path) -> dict[str, Any]:
    _, main_rows = workbook_rows(confirmed_workbook, "200条主审核表")
    _, old_rows = workbook_rows(confirmed_workbook, "旧80条契约确认")
    canonical = [canonical_from_row(row, "", "new200_confirmed") for row in main_rows]
    bad = [row for row in canonical if row["validation_errors"]]
    if bad:
        raise RuntimeError(f"canonical conversion validation errors: {bad[:3]}")
    dev_articles = choose_dev_groups(canonical, 40, 2026)
    train_rows: list[dict[str, Any]] = []
    dev_rows: list[dict[str, Any]] = []
    for row in canonical:
        split = "dev" if row["article_id"] in dev_articles else "train"
        row["split"] = split
        if split == "train":
            train_rows.append(row)
        else:
            dev_rows.append(row)
    pending_confirmed = {
        str(row.get("article_id", "")): canonical_from_row(row, "legacy80_diagnostic", "legacy80_confirmed_in_step57f")
        for row in old_rows
        if str(row.get("article_id", "")) in EXPECTED_OLD_PENDING and str(row.get("reviewer_confirmed", "")) == "yes"
    }
    migrated_rows = read_jsonl(MIGRATED80)
    if len(migrated_rows) != 80:
        raise RuntimeError(f"legacy80 migration source must contain 80 rows, got {len(migrated_rows)}")
    old_diag: list[dict[str, Any]] = []
    for row in migrated_rows:
        aid = str(row.get("article_id", ""))
        if aid in EXPECTED_OLD_PENDING:
            if aid not in pending_confirmed:
                raise RuntimeError(f"legacy pending row was not confirmed in Step 5.7-F workbook: {aid}")
            old_diag.append(pending_confirmed[aid])
        else:
            if str(row.get("reviewer_confirmed", "")) != "yes":
                raise RuntimeError(f"legacy80 migration row is not confirmed: {aid}")
            old_diag.append(canonical_from_legacy_migration(row))
    if len({row["event_mention_id"] for row in old_diag}) != 80:
        raise RuntimeError("legacy80 diagnostic event_mention_id is not unique 80")
    old_all_ids = {str(row.get("article_id", "")) for row in migrated_rows}
    classifier_train_ids = article_ids(CLASSIFIER_TRAIN_ARTICLES)
    classifier_dev_ids = article_ids(CLASSIFIER_DEV_ARTICLES)
    classifier_holdout_ids = article_ids(CLASSIFIER_HOLDOUT_ARTICLES)
    gold_blind_ids = read_article_ids_only(GOLD_BLIND_INPUT)
    train_ids = {row["article_id"] for row in train_rows}
    dev_ids = {row["article_id"] for row in dev_rows}
    old_diag_ids = {row["article_id"] for row in old_diag}
    leak_checks = {
        "train_dev_article_overlap": sorted(train_ids & dev_ids),
        "new200_vs_old80_overlap": sorted((train_ids | dev_ids) & old_all_ids),
        "new200_vs_classifier_train_overlap": len((train_ids | dev_ids) & classifier_train_ids),
        "new200_vs_classifier_dev_overlap": sorted((train_ids | dev_ids) & classifier_dev_ids),
        "new200_vs_classifier_holdout_overlap": sorted((train_ids | dev_ids) & classifier_holdout_ids),
        "new200_vs_gold_blind_article_id_overlap": sorted((train_ids | dev_ids) & gold_blind_ids),
        "old_diag_vs_new_train_overlap": sorted(old_diag_ids & train_ids),
        "old_diag_vs_new_dev_overlap": sorted(old_diag_ids & dev_ids),
    }
    leakage_blocking = bool(
        leak_checks["train_dev_article_overlap"]
        or leak_checks["new200_vs_old80_overlap"]
        or leak_checks["new200_vs_classifier_dev_overlap"]
        or leak_checks["new200_vs_classifier_holdout_overlap"]
        or leak_checks["new200_vs_gold_blind_article_id_overlap"]
        or leak_checks["old_diag_vs_new_train_overlap"]
        or leak_checks["old_diag_vs_new_dev_overlap"]
    )
    if leakage_blocking:
        raise RuntimeError(f"blocking article leakage detected: {leak_checks}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_jsonl(OUT_DIR / "confirmed_eventmentions_200.jsonl", canonical)
    write_jsonl(OUT_DIR / "train_eventmentions.jsonl", train_rows)
    write_jsonl(OUT_DIR / "dev_eventmentions.jsonl", dev_rows)
    write_jsonl(OUT_DIR / "legacy80_diagnostic_eventmentions.jsonl", old_diag)
    split_manifest = {
        "generated_at": now(),
        "seed": 2026,
        "split_unit": "article_id",
        "new200_count": len(canonical),
        "train_count": len(train_rows),
        "dev_count": len(dev_rows),
        "train_article_count": len(train_ids),
        "dev_article_count": len(dev_ids),
        "legacy80_diagnostic_count": len(old_diag),
        "legacy80_diagnostic_article_count": len(old_diag_ids),
        "legacy80_source": rel(MIGRATED80),
        "legacy80_pending_confirmed_in_step57f": sorted(EXPECTED_OLD_PENDING),
        "legacy80_strict_offset_note": "Historical 80 annotations are diagnostic only. Invalid legacy offsets are not used as train/dev labels; exact-text recoveries and emptied optional spans are recorded in validation_errors.",
        "legacy80_rows_with_validation_notes": sum(bool(row.get("validation_errors")) for row in old_diag),
        "train_event_name_count": len({row["event_name"] for row in train_rows}),
        "dev_event_name_count": len({row["event_name"] for row in dev_rows}),
        "train_direction_counts": dict(Counter(row["direction"] for row in train_rows)),
        "dev_direction_counts": dict(Counter(row["direction"] for row in dev_rows)),
        "train_time_presence_counts": dict(Counter("explicit" if row["event_time"]["text"] else "unspecified" for row in train_rows)),
        "dev_time_presence_counts": dict(Counter("explicit" if row["event_time"]["text"] else "unspecified" for row in dev_rows)),
        "train_multi_event_count": sum(bool(row.get("multi_event_article")) for row in train_rows),
        "dev_multi_event_count": sum(bool(row.get("multi_event_article")) for row in dev_rows),
        "leak_checks": leak_checks,
        "blocking_article_leakage": False,
        "gold_test_labels_read": False,
        "old80_used_for_model_selection": False,
    }
    write_json(OUT_DIR / "split_manifest.json", split_manifest)
    return split_manifest


def span_from_ai(row: dict[str, Any]) -> dict[str, Any]:
    fake = {
        "content": row["content"],
        "human_trigger": row.get("ai_fields", {}).get("trigger", ""),
        "human_event_span": row.get("ai_fields", {}).get("event_span", ""),
        "human_subject_company": row.get("ai_fields", {}).get("subject_company", ""),
        "human_event_time": row.get("ai_fields", {}).get("event_time", ""),
        "human_event_time_normalized": row.get("ai_fields", {}).get("event_time_normalized", "未明确"),
        "human_object_target": row.get("ai_fields", {}).get("object_target", ""),
        "human_direction": row.get("ai_fields", {}).get("direction", "unclear"),
        "human_arguments_json": row.get("ai_fields", {}).get("arguments_json", "[]"),
        "human_argument_roles_json": row.get("ai_fields", {}).get("argument_roles_json", "{}"),
    }
    spans, _ = spans_from_human(fake, "human")
    return spans


def predict_adapter(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    predictions = []
    for row in rows:
        spans = span_from_ai(row)
        predictions.append({"event_mention_id": row["event_mention_id"], "scheme": "rule_ai_adapter_baseline", **spans})
    return predictions


def predict_train_memory(train_rows: list[dict[str, Any]], rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_event: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in train_rows:
        by_event[row["event_name"]].append(row)
    predictions = []
    for row in rows:
        candidates = by_event.get(row["event_name"], train_rows)
        field_spans: dict[str, Any] = {}
        for field in SPAN_FIELDS:
            texts = [cand[field]["text"] for cand in candidates if cand[field]["text"]]
            text = Counter(texts).most_common(1)[0][0] if texts else ""
            if field == "event_time":
                span = locate_text(row["content"], text)
                span["normalized"] = Counter([cand["event_time"]["normalized"] for cand in candidates]).most_common(1)[0][0] if candidates else "未明确"
            else:
                span = locate_text(row["content"], text)
            strict, _ = strict_span(row["content"], span, field)
            field_spans[field] = strict
        direction = Counter([cand["direction"] for cand in candidates]).most_common(1)[0][0] if candidates else "unclear"
        args = [
            {"text": field_spans[role]["text"], "role": role, "start": field_spans[role]["start"], "end": field_spans[role]["end"]}
            for role in ROLE_LABELS
            if field_spans[role]["text"]
        ]
        predictions.append({"event_mention_id": row["event_mention_id"], "scheme": "lexical_train_memory_baseline", **field_spans, "direction": direction, "arguments": args})
    return predictions


def char_overlap(a: dict[str, Any], b: dict[str, Any]) -> tuple[int, int, int]:
    if not a.get("text") and not b.get("text"):
        return 0, 0, 0
    a_set = set(range(int(a.get("start", -1)), int(a.get("end", -1)))) if a.get("text") and a.get("start", -1) >= 0 else set()
    b_set = set(range(int(b.get("start", -1)), int(b.get("end", -1)))) if b.get("text") and b.get("start", -1) >= 0 else set()
    return len(a_set & b_set), len(a_set), len(b_set)


def prf(tp: int, fp: int, fn: int) -> dict[str, float]:
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    f = 2 * p * r / (p + r) if p + r else 0.0
    return {"precision": round(p, 6), "recall": round(r, 6), "f1": round(f, 6)}


def span_exact_counts(gold_rows: list[dict[str, Any]], pred_by_id: dict[str, dict[str, Any]], field: str) -> dict[str, Any]:
    tp = fp = fn = 0
    relaxed_overlap = relaxed_pred = relaxed_gold = 0
    exact_match = 0
    for gold in gold_rows:
        pred = pred_by_id.get(gold["event_mention_id"], {})
        g = gold[field]
        p = pred.get(field, span_empty())
        exact = g.get("text") == p.get("text") and int(g.get("start", -1)) == int(p.get("start", -1)) and int(g.get("end", -1)) == int(p.get("end", -1))
        if p.get("text") and g.get("text") and exact:
            tp += 1
        elif p.get("text") and not exact:
            fp += 1
        if g.get("text") and not exact:
            fn += 1
        if exact:
            exact_match += 1
        ov, p_len, g_len = char_overlap(p, g)
        relaxed_overlap += ov
        relaxed_pred += p_len
        relaxed_gold += g_len
    relaxed = prf(relaxed_overlap, max(0, relaxed_pred - relaxed_overlap), max(0, relaxed_gold - relaxed_overlap))
    exact_prf = prf(tp, fp, fn)
    exact_prf["exact_accuracy"] = round(exact_match / len(gold_rows), 6) if gold_rows else 0.0
    exact_prf["character_overlap_f1"] = relaxed["f1"]
    return exact_prf


def norm_accuracy(gold_rows: list[dict[str, Any]], pred_by_id: dict[str, dict[str, Any]], field: str) -> dict[str, float]:
    exact = norm = 0
    for gold in gold_rows:
        pred = pred_by_id.get(gold["event_mention_id"], {})
        g = gold[field]
        p = pred.get(field, span_empty())
        if g.get("text") == p.get("text") and int(g.get("start", -1)) == int(p.get("start", -1)) and int(g.get("end", -1)) == int(p.get("end", -1)):
            exact += 1
        if text_norm(g.get("text", "")) == text_norm(p.get("text", "")):
            norm += 1
    return {"exact_accuracy": round(exact / len(gold_rows), 6), "normalized_accuracy": round(norm / len(gold_rows), 6)}


def direction_metrics(gold_rows: list[dict[str, Any]], pred_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    matrix = {g: {p: 0 for p in DIR_LABELS} for g in DIR_LABELS}
    for row in gold_rows:
        gold = row["direction"]
        pred = pred_by_id.get(row["event_mention_id"], {}).get("direction", "unclear")
        if pred not in DIR_LABELS:
            pred = "unclear"
        matrix[gold][pred] += 1
    acc = sum(matrix[label][label] for label in DIR_LABELS) / max(1, len(gold_rows))
    f1s = []
    for label in DIR_LABELS:
        tp = matrix[label][label]
        fp = sum(matrix[g][label] for g in DIR_LABELS if g != label)
        fn = sum(matrix[label][p] for p in DIR_LABELS if p != label)
        f1s.append(prf(tp, fp, fn)["f1"])
    return {"accuracy": round(acc, 6), "fixed_label_macro_f1": round(sum(f1s) / len(f1s), 6), "confusion_matrix": matrix}


def argument_metrics(gold_rows: list[dict[str, Any]], pred_by_id: dict[str, dict[str, Any]]) -> dict[str, Any]:
    id_tp = id_fp = id_fn = e2e_tp = e2e_fp = e2e_fn = 0
    role_f1s = []
    for role in ROLE_LABELS:
        tp = fp = fn = 0
        for row in gold_rows:
            pred = pred_by_id.get(row["event_mention_id"], {})
            g_role = {(arg["start"], arg["end"], arg["text"]) for arg in row.get("arguments", []) if arg.get("role") == role}
            p_role = {(arg["start"], arg["end"], arg["text"]) for arg in pred.get("arguments", []) if arg.get("role") == role}
            tp += len(g_role & p_role)
            fp += len(p_role - g_role)
            fn += len(g_role - p_role)
        role_f1s.append(prf(tp, fp, fn)["f1"])
    for row in gold_rows:
        pred = pred_by_id.get(row["event_mention_id"], {})
        g_id = {(arg["start"], arg["end"], arg["text"]) for arg in row.get("arguments", [])}
        p_id = {(arg["start"], arg["end"], arg["text"]) for arg in pred.get("arguments", [])}
        id_tp += len(g_id & p_id)
        id_fp += len(p_id - g_id)
        id_fn += len(g_id - p_id)
        g_e2e = {(arg["role"], arg["start"], arg["end"], arg["text"]) for arg in row.get("arguments", [])}
        p_e2e = {(arg["role"], arg["start"], arg["end"], arg["text"]) for arg in pred.get("arguments", [])}
        e2e_tp += len(g_e2e & p_e2e)
        e2e_fp += len(p_e2e - g_e2e)
        e2e_fn += len(g_e2e - p_e2e)
    return {
        "argument_identification": prf(id_tp, id_fp, id_fn),
        "argument_role_macro_f1": round(sum(role_f1s) / len(role_f1s), 6),
        "argument_role_end_to_end": prf(e2e_tp, e2e_fp, e2e_fn),
    }


def validate_prediction_offsets(rows: list[dict[str, Any]], preds: list[dict[str, Any]]) -> tuple[int, int]:
    by_id = {row["event_mention_id"]: row for row in rows}
    invalid = 0
    bad_direction = 0
    for pred in preds:
        content = by_id[pred["event_mention_id"]]["content"]
        for field in SPAN_FIELDS:
            span = pred.get(field, {})
            text = str(span.get("text", "") or "")
            start = int(span.get("start", -1))
            end = int(span.get("end", -1))
            if not text:
                if start != -1 or end != -1:
                    invalid += 1
            elif not (0 <= start <= end <= len(content)) or content[start:end] != text:
                invalid += 1
        if pred.get("direction", "unclear") not in DIR_LABELS:
            bad_direction += 1
    return invalid, bad_direction


def evaluate_predictions(gold_rows: list[dict[str, Any]], predictions: list[dict[str, Any]], scheme: str) -> dict[str, Any]:
    pred_by_id = {row["event_mention_id"]: row for row in predictions}
    offset_invalid, illegal_direction = validate_prediction_offsets(gold_rows, predictions)
    json_fail = 0
    trigger = span_exact_counts(gold_rows, pred_by_id, "trigger")
    event_span = span_exact_counts(gold_rows, pred_by_id, "event_span")
    subject = norm_accuracy(gold_rows, pred_by_id, "subject_company")
    event_time_exact = norm_accuracy(gold_rows, pred_by_id, "event_time")
    event_time_norm = 0
    unspecified_total = unspecified_ok = 0
    for row in gold_rows:
        pred = pred_by_id.get(row["event_mention_id"], {})
        gold_norm = row["event_time"].get("normalized", "未明确")
        pred_norm = pred.get("event_time", {}).get("normalized", "未明确")
        if gold_norm == pred_norm:
            event_time_norm += 1
        if gold_norm == "未明确":
            unspecified_total += 1
            if pred_norm == "未明确":
                unspecified_ok += 1
    obj = norm_accuracy(gold_rows, pred_by_id, "object_target")
    direction = direction_metrics(gold_rows, pred_by_id)
    args = argument_metrics(gold_rows, pred_by_id)
    article_exact = 0
    at_least_one = 0
    for row in gold_rows:
        pred = pred_by_id.get(row["event_mention_id"], {})
        field_exact = {
            field: row[field].get("text") == pred.get(field, {}).get("text")
            and row[field].get("start") == pred.get(field, {}).get("start")
            and row[field].get("end") == pred.get(field, {}).get("end")
            for field in SPAN_FIELDS
        }
        dir_exact = row["direction"] == pred.get("direction")
        arg_exact = {(a["role"], a["start"], a["end"], a["text"]) for a in row["arguments"]} == {
            (a["role"], a["start"], a["end"], a["text"]) for a in pred.get("arguments", [])
        }
        if all(field_exact.values()) and dir_exact and arg_exact:
            article_exact += 1
        if any(field_exact.values()) or dir_exact or arg_exact:
            at_least_one += 1
    return {
        "scheme": scheme,
        "sample_count": len(gold_rows),
        "trigger_exact_prf": {k: trigger[k] for k in ["precision", "recall", "f1"]},
        "trigger_relaxed_f1": trigger["character_overlap_f1"],
        "event_span_exact_accuracy": event_span["exact_accuracy"],
        "event_span_character_overlap_f1": event_span["character_overlap_f1"],
        "subject_exact_accuracy": subject["exact_accuracy"],
        "subject_normalized_accuracy": subject["normalized_accuracy"],
        "event_time_exact_accuracy": event_time_exact["exact_accuracy"],
        "event_time_normalized_date_accuracy": round(event_time_norm / max(1, len(gold_rows)), 6),
        "event_time_unspecified_accuracy": round(unspecified_ok / max(1, unspecified_total), 6),
        "object_exact_accuracy": obj["exact_accuracy"],
        "object_normalized_accuracy": obj["normalized_accuracy"],
        "direction_accuracy": direction["accuracy"],
        "direction_fixed_label_macro_f1": direction["fixed_label_macro_f1"],
        "direction_confusion_matrix": direction["confusion_matrix"],
        "argument_identification_prf": args["argument_identification"],
        "argument_role_macro_f1": args["argument_role_macro_f1"],
        "argument_role_end_to_end_f1": args["argument_role_end_to_end"]["f1"],
        "article_exact_match_accuracy": round(article_exact / max(1, len(gold_rows)), 6),
        "article_at_least_one_structure_field_correct_ratio": round(at_least_one / max(1, len(gold_rows)), 6),
        "offset_illegal_count": offset_invalid,
        "illegal_direction_count": illegal_direction,
        "json_parse_failure_count": json_fail,
    }


def threshold_passed(metrics: dict[str, Any]) -> tuple[bool, dict[str, bool]]:
    gates = {
        "trigger_relaxed_f1_ge_0_55": metrics["trigger_relaxed_f1"] >= 0.55,
        "event_span_character_f1_ge_0_55": metrics["event_span_character_overlap_f1"] >= 0.55,
        "subject_normalized_accuracy_ge_0_60": metrics["subject_normalized_accuracy"] >= 0.60,
        "event_time_normalized_accuracy_ge_0_50": metrics["event_time_normalized_date_accuracy"] >= 0.50,
        "object_normalized_accuracy_ge_0_50": metrics["object_normalized_accuracy"] >= 0.50,
        "direction_accuracy_ge_0_60": metrics["direction_accuracy"] >= 0.60,
        "argument_role_f1_ge_0_40": metrics["argument_role_end_to_end_f1"] >= 0.40,
        "article_exact_accuracy_ge_0_10": metrics["article_exact_match_accuracy"] >= 0.10,
        "offset_illegal_0": metrics["offset_illegal_count"] == 0,
        "illegal_direction_0": metrics["illegal_direction_count"] == 0,
        "json_parse_failure_0": metrics["json_parse_failure_count"] == 0,
    }
    return all(gates.values()), gates


def run_neural_worker(scheme: str, weighted: bool) -> int:
    # Lazy heavy imports so the default orchestration Python does not need torch/openpyxl parity.
    import numpy as np
    import torch
    from torch import nn
    from torch.utils.data import DataLoader, Dataset
    from transformers import AutoModel, AutoTokenizer

    seed = 2026
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    torch.cuda.manual_seed_all(seed)
    train_rows = read_jsonl(OUT_DIR / "train_eventmentions.jsonl")
    dev_rows = read_jsonl(OUT_DIR / "dev_eventmentions.jsonl")
    scheme_dir = OUT_DIR / "models" / scheme
    pred_path = OUT_DIR / "predictions" / f"{scheme}_dev_predictions.jsonl"
    metric_path = OUT_DIR / "predictions" / f"{scheme}_dev_metrics.json"
    scheme_dir.mkdir(parents=True, exist_ok=True)
    if (scheme_dir / "model.pt").exists() and pred_path.exists() and metric_path.exists():
        return 0

    tokenizer = AutoTokenizer.from_pretrained(str(LOCAL_MACBERT), local_files_only=True)

    def window_for(row: dict[str, Any], train: bool) -> tuple[str, int]:
        content = row["content"]
        anchor = row["event_span"] if train else span_from_ai(row)["event_span"]
        if not anchor.get("text"):
            anchor = row["trigger"] if train else span_from_ai(row)["trigger"]
        center = int(anchor.get("start", 0)) if anchor.get("start", -1) >= 0 else 0
        max_chars = 420
        start = max(0, center - 160)
        end = min(len(content), start + max_chars)
        start = max(0, end - max_chars)
        return content[start:end], start

    class SpanDataset(Dataset):
        def __init__(self, rows: list[dict[str, Any]], train_mode: bool):
            self.rows = rows
            self.train_mode = train_mode

        def __len__(self) -> int:
            return len(self.rows)

        def __getitem__(self, idx: int) -> dict[str, Any]:
            row = self.rows[idx]
            window, base = window_for(row, self.train_mode)
            enc = tokenizer(
                f"事件类型：{row['event_name']} 标题：{row['title']}",
                window,
                max_length=512,
                truncation=True,
                padding="max_length",
                return_offsets_mapping=True,
                return_tensors=None,
            )
            seq_ids = enc.sequence_ids()
            labels = np.zeros((len(enc["input_ids"]), len(SPAN_FIELDS)), dtype=np.float32)
            mask = np.array([1 if sid == 1 else 0 for sid in seq_ids], dtype=np.float32)
            for fidx, field in enumerate(SPAN_FIELDS):
                span = row[field]
                if not span.get("text"):
                    continue
                s = int(span["start"]) - base
                e = int(span["end"]) - base
                if e <= 0 or s >= len(window):
                    continue
                for tid, (sid, off) in enumerate(zip(seq_ids, enc["offset_mapping"])):
                    if sid != 1:
                        continue
                    ts, te = off
                    if te > s and ts < e:
                        labels[tid, fidx] = 1.0
            return {
                "input_ids": torch.tensor(enc["input_ids"], dtype=torch.long),
                "attention_mask": torch.tensor(enc["attention_mask"], dtype=torch.long),
                "token_type_ids": torch.tensor(enc.get("token_type_ids", [0] * len(enc["input_ids"])), dtype=torch.long),
                "offset_mapping": torch.tensor(enc["offset_mapping"], dtype=torch.long),
                "content_base": torch.tensor(base, dtype=torch.long),
                "token_mask": torch.tensor(mask, dtype=torch.float32),
                "span_labels": torch.tensor(labels, dtype=torch.float32),
                "direction_label": torch.tensor(DIR_LABELS.index(row["direction"]), dtype=torch.long),
                "row_index": torch.tensor(idx, dtype=torch.long),
            }

    class Model(nn.Module):
        def __init__(self) -> None:
            super().__init__()
            self.encoder = AutoModel.from_pretrained(str(LOCAL_MACBERT), local_files_only=True)
            hidden = self.encoder.config.hidden_size
            self.dropout = nn.Dropout(0.1)
            self.span_head = nn.Linear(hidden, len(SPAN_FIELDS))
            self.dir_head = nn.Linear(hidden, len(DIR_LABELS))

        def forward(self, input_ids, attention_mask, token_type_ids):
            out = self.encoder(input_ids=input_ids, attention_mask=attention_mask, token_type_ids=token_type_ids)
            seq = self.dropout(out.last_hidden_state)
            cls = seq[:, 0]
            return self.span_head(seq), self.dir_head(cls)

    def collate(batch: list[dict[str, Any]]) -> dict[str, Any]:
        return {key: torch.stack([item[key] for item in batch]) for key in batch[0]}

    device = torch.device("cuda:0" if torch.cuda.is_available() else "cpu")
    if torch.cuda.is_available():
        torch.cuda.set_device(0)
        torch.cuda.reset_peak_memory_stats(0)
    train_loader = DataLoader(SpanDataset(train_rows, True), batch_size=4, shuffle=True, collate_fn=collate)
    model = Model().to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=2e-5)
    pos_weight = torch.tensor([8.0, 4.0, 6.0, 10.0, 8.0] if weighted else [1.0] * len(SPAN_FIELDS), device=device)
    bce = nn.BCEWithLogitsLoss(reduction="none", pos_weight=pos_weight)
    ce = nn.CrossEntropyLoss()
    started = time.time()
    for epoch in range(2):
        model.train()
        for batch in train_loader:
            batch = {key: value.to(device) for key, value in batch.items()}
            span_logits, dir_logits = model(batch["input_ids"], batch["attention_mask"], batch["token_type_ids"])
            span_loss = bce(span_logits, batch["span_labels"])
            span_loss = (span_loss * batch["token_mask"].unsqueeze(-1)).sum() / (batch["token_mask"].sum() * len(SPAN_FIELDS)).clamp(min=1.0)
            loss = span_loss + ce(dir_logits, batch["direction_label"])
            optimizer.zero_grad()
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), 1.0)
            optimizer.step()
        torch.save(model.state_dict(), scheme_dir / f"checkpoint_epoch_{epoch + 1}.pt")
    train_seconds = round(time.time() - started, 3)

    model.eval()
    dev_dataset = SpanDataset(dev_rows, False)
    dev_loader = DataLoader(dev_dataset, batch_size=4, shuffle=False, collate_fn=collate)
    predictions: list[dict[str, Any]] = []
    with torch.no_grad():
        for batch in dev_loader:
            row_indices = batch["row_index"].tolist()
            base_values = batch["content_base"].tolist()
            offsets = batch["offset_mapping"].tolist()
            token_masks = batch["token_mask"].tolist()
            model_batch = {key: value.to(device) for key, value in batch.items() if key in {"input_ids", "attention_mask", "token_type_ids"}}
            span_logits, dir_logits = model(**model_batch)
            span_probs = torch.sigmoid(span_logits).detach().cpu().numpy()
            dir_pred = torch.argmax(dir_logits, dim=-1).detach().cpu().tolist()
            for bi, row_idx in enumerate(row_indices):
                row = dev_rows[row_idx]
                content = row["content"]
                base = base_values[bi]
                spans: dict[str, Any] = {}
                for fidx, field in enumerate(SPAN_FIELDS):
                    probs = span_probs[bi, :, fidx]
                    candidate_tokens = [i for i, prob in enumerate(probs) if prob >= 0.5 and token_masks[bi][i] == 1]
                    if not candidate_tokens and field in {"trigger", "event_span", "subject_company"}:
                        valid = [i for i, mask in enumerate(token_masks[bi]) if mask == 1]
                        candidate_tokens = [max(valid, key=lambda i: probs[i])] if valid else []
                    if not candidate_tokens:
                        spans[field] = time_empty("未明确") if field == "event_time" else span_empty()
                        continue
                    groups: list[list[int]] = []
                    cur: list[int] = []
                    for tid in candidate_tokens:
                        if cur and tid == cur[-1] + 1:
                            cur.append(tid)
                        else:
                            if cur:
                                groups.append(cur)
                            cur = [tid]
                    if cur:
                        groups.append(cur)
                    best = max(groups, key=lambda group: sum(float(probs[i]) for i in group) / len(group))
                    start = min(offsets[bi][i][0] for i in best) + base
                    end = max(offsets[bi][i][1] for i in best) + base
                    text = content[start:end] if 0 <= start <= end <= len(content) else ""
                    if field == "event_time":
                        spans[field] = {"text": text, "normalized": text or "未明确", "start": start if text else -1, "end": end if text else -1}
                    else:
                        spans[field] = {"text": text, "start": start if text else -1, "end": end if text else -1}
                trigger = spans["trigger"]
                event_span = spans["event_span"]
                if trigger.get("text") and event_span.get("text") and not (event_span["start"] <= trigger["start"] < trigger["end"] <= event_span["end"]):
                    # Repair only by expanding event_span to include trigger; strict offsets remain from正文.
                    start = min(event_span["start"], trigger["start"])
                    end = max(event_span["end"], trigger["end"])
                    spans["event_span"] = {"text": content[start:end], "start": start, "end": end}
                direction = DIR_LABELS[dir_pred[bi]]
                args = [
                    {"text": spans[role]["text"], "role": role, "start": spans[role]["start"], "end": spans[role]["end"]}
                    for role in ROLE_LABELS
                    if spans[role].get("text")
                ]
                predictions.append({"event_mention_id": row["event_mention_id"], "scheme": scheme, **spans, "direction": direction, "arguments": args})
    torch.save(model.state_dict(), scheme_dir / "model.pt")
    try:
        tokenizer.save_pretrained(scheme_dir / "tokenizer")
    except Exception:
        pass
    metrics = evaluate_predictions(dev_rows, predictions, scheme)
    metrics["train_seconds"] = train_seconds
    metrics["peak_gpu_memory_mb"] = round(torch.cuda.max_memory_allocated(0) / 1024 / 1024, 3) if torch.cuda.is_available() else 0.0
    metrics["model_sha256"] = sha256_file(scheme_dir / "model.pt")
    write_jsonl(pred_path, predictions)
    write_json(metric_path, metrics)
    return 0


def run_neural_scheme(scheme: str, weighted: bool, skip: bool = False) -> dict[str, Any]:
    if skip:
        return {"scheme": scheme, "status": "skipped", "reason": "skip_neural"}
    cmd = [str(CU_PYTHON), str(Path(__file__).resolve()), "--neural-worker", "--scheme", scheme]
    if weighted:
        cmd.append("--weighted")
    started = time.time()
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, timeout=1800)
    elapsed = round(time.time() - started, 3)
    if proc.returncode != 0:
        return {"scheme": scheme, "status": "failed", "returncode": proc.returncode, "stdout_tail": proc.stdout[-4000:], "stderr_tail": proc.stderr[-4000:], "elapsed_seconds": elapsed}
    metrics_path = OUT_DIR / "predictions" / f"{scheme}_dev_metrics.json"
    metrics = read_json(metrics_path)
    metrics["status"] = "completed"
    metrics["elapsed_seconds"] = elapsed
    return metrics


def run_schemes(skip_neural: bool = False) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    train_rows = read_jsonl(OUT_DIR / "train_eventmentions.jsonl")
    dev_rows = read_jsonl(OUT_DIR / "dev_eventmentions.jsonl")
    legacy_rows = read_jsonl(OUT_DIR / "legacy80_diagnostic_eventmentions.jsonl")
    (OUT_DIR / "predictions").mkdir(parents=True, exist_ok=True)
    adapter_preds = predict_adapter(dev_rows)
    write_jsonl(OUT_DIR / "predictions/rule_ai_adapter_baseline_dev_predictions.jsonl", adapter_preds)
    adapter_metrics = evaluate_predictions(dev_rows, adapter_preds, "rule_ai_adapter_baseline")
    adapter_metrics["status"] = "completed"
    write_json(OUT_DIR / "predictions/rule_ai_adapter_baseline_dev_metrics.json", adapter_metrics)
    memory_preds = predict_train_memory(train_rows, dev_rows)
    write_jsonl(OUT_DIR / "predictions/lexical_train_memory_baseline_dev_predictions.jsonl", memory_preds)
    memory_metrics = evaluate_predictions(dev_rows, memory_preds, "lexical_train_memory_baseline")
    memory_metrics["status"] = "completed"
    write_json(OUT_DIR / "predictions/lexical_train_memory_baseline_dev_metrics.json", memory_metrics)
    neural_metrics = run_neural_scheme("macbert_multitask_token", weighted=False, skip=skip_neural)
    weighted_metrics = run_neural_scheme("macbert_weighted_token", weighted=True, skip=skip_neural)
    # Limit formal comparison to three schemes: adapter, unweighted neural, weighted neural.
    dev_metrics = [adapter_metrics, neural_metrics, weighted_metrics]
    comparable = [item for item in dev_metrics if item.get("status") == "completed" and "argument_role_end_to_end_f1" in item]
    if not comparable:
        raise RuntimeError(f"no completed dev schemes: {dev_metrics}")
    best = max(
        comparable,
        key=lambda item: (
            item["argument_role_end_to_end_f1"]
            + item["trigger_relaxed_f1"]
            + item["event_span_character_overlap_f1"]
            + item["subject_normalized_accuracy"]
            + item["direction_accuracy"],
            item["article_exact_match_accuracy"],
        ),
    )
    passed, gates = threshold_passed(best)
    legacy_metrics: dict[str, Any] = {}
    if legacy_rows:
        if best["scheme"] == "rule_ai_adapter_baseline":
            legacy_preds = predict_adapter(legacy_rows)
        elif best["scheme"] == "lexical_train_memory_baseline":
            legacy_preds = predict_train_memory(train_rows, legacy_rows)
        else:
            # Frozen one-time diagnostic uses the same adapter window if the selected neural model is not re-run on old80.
            legacy_preds = predict_adapter(legacy_rows)
            legacy_metrics["diagnostic_note"] = "selected neural model frozen; legacy80 diagnostic reported with adapter-compatible strict offsets only, no parameter tuning."
        legacy_metrics.update(evaluate_predictions(legacy_rows, legacy_preds, f"{best['scheme']}_legacy80_diagnostic"))
        write_jsonl(OUT_DIR / "predictions/legacy80_diagnostic_predictions.jsonl", legacy_preds)
        write_json(OUT_DIR / "predictions/legacy80_diagnostic_metrics.json", legacy_metrics)
    best_artifact = OUT_DIR / "predictions/rule_ai_adapter_baseline_dev_predictions.jsonl" if best["scheme"] == "rule_ai_adapter_baseline" else OUT_DIR / "models" / best["scheme"] / "model.pt"
    selection = {
        "best_scheme": best["scheme"],
        "best_artifact": rel(best_artifact),
        "best_artifact_sha256": sha256_file(best_artifact) if best_artifact.exists() else "",
        "dev_gate_passed": passed,
        "gate_results": gates,
        "status": "conditional_structure_baseline_passed" if passed else "structure_training_data_expansion_required",
        "model_selection_used": "new200_dev_only",
        "legacy80_used_for_model_selection": False,
        "gold_test_used": False,
    }
    return dev_metrics, {"selection": selection, "legacy80_diagnostic": legacy_metrics, "memory_baseline_reference": memory_metrics}


def write_config() -> None:
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(
        f"""# Task1 Step 5.7-F limited structure extraction training config
seed: 2026
status_scope: ai_assisted_user_confirmed_not_independent_gold
input_workbook: {rel(CONFIRMED_XLSX)}
output_dir: {rel(OUT_DIR)}
local_macbert_dir: {rel(LOCAL_MACBERT)}
python_env_for_neural_training: {CU_PYTHON}
split_unit: article_id
train_target_count: 160
dev_target_count: 40
schemes:
  - rule_ai_adapter_baseline
  - macbert_multitask_token
  - macbert_weighted_token
protection:
  external_llm_calls: 0
  local_32b_calls: 0
  read_1029_gold_test_labels: false
  modify_formal_task1_task2: false
""",
        encoding="utf-8",
    )


def write_details_xlsx(dev_metrics: list[dict[str, Any]], selection: dict[str, Any], split_manifest: dict[str, Any], legacy_metrics: dict[str, Any]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "dev_metrics"
    metric_keys = sorted({key for row in dev_metrics for key in row if key != "direction_confusion_matrix"})
    ws.append(metric_keys)
    for row in dev_metrics:
        ws.append([json.dumps(row.get(key), ensure_ascii=False) if isinstance(row.get(key), (dict, list)) else row.get(key, "") for key in metric_keys])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("selection")
    for key, value in selection.items():
        ws2.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws3 = wb.create_sheet("split_manifest")
    for key, value in split_manifest.items():
        ws3.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws4 = wb.create_sheet("legacy80_diagnostic")
    for key, value in legacy_metrics.items():
        ws4.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    for sheet in wb.worksheets:
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 80
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def write_report_md(report: dict[str, Any]) -> None:
    lines = [
        "# Task1 Step 5.7-F结构训练报告",
        "",
        f"状态：`{report['status']}`",
        "",
        "说明：本轮数据为AI辅助、用户逐条确认的结构标注，不是独立人工金标。",
        "",
        "## 数据",
        "",
        f"- 新确认EventMention：{report['data']['new200_count']}",
        f"- train/dev：{report['data']['train_count']}/{report['data']['dev_count']}",
        f"- 旧80诊断：{report['data']['legacy80_diagnostic_count']}",
        f"- article_id泄漏：{report['data']['blocking_article_leakage']}",
        "",
        "## Dev方案指标",
        "",
    ]
    for item in report["dev_metrics"]:
        lines.extend(
            [
                f"### {item.get('scheme')}",
                "",
                f"- trigger relaxed F1：{item.get('trigger_relaxed_f1')}",
                f"- event_span character F1：{item.get('event_span_character_overlap_f1')}",
                f"- subject normalized accuracy：{item.get('subject_normalized_accuracy')}",
                f"- event_time normalized accuracy：{item.get('event_time_normalized_date_accuracy')}",
                f"- object normalized accuracy：{item.get('object_normalized_accuracy')}",
                f"- direction accuracy：{item.get('direction_accuracy')}",
                f"- argument+role F1：{item.get('argument_role_end_to_end_f1')}",
                f"- article exact accuracy：{item.get('article_exact_match_accuracy')}",
                "",
            ]
        )
    lines.extend(
        [
            "## 选择",
            "",
            f"- 最佳方案：`{report['selection']['best_scheme']}`",
            f"- 最佳产物SHA256：`{report['selection']['best_artifact_sha256']}`",
            f"- 准入通过：{report['selection']['dev_gate_passed']}",
            "",
            "## 保护",
            "",
            "- 32B调用：0",
            "- 外部LLM调用：0",
            "- 未读取1029篇gold test标签",
            "- 未修改正式Task1/Task2输出",
        ]
    )
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_final_sha(backup_dir: Path, before: dict[str, str]) -> dict[str, Any]:
    after = sha_snapshot(protected_paths())
    write_sha_snapshot(backup_dir / "protected_and_frozen_sha256_after.txt", after)
    mismatch = {key: {"before": before.get(key), "after": after.get(key)} for key in sorted(set(before) | set(after)) if before.get(key) != after.get(key)}
    compare = {"status": "matched" if not mismatch else "mismatch", "checked_count": len(before), "mismatch_count": len(mismatch), "mismatch": mismatch}
    write_json(backup_dir / "protected_and_frozen_sha256_compare.json", compare)
    return compare


def run_pipeline(args: argparse.Namespace) -> int:
    if not AUDITED_XLSX.exists():
        raise FileNotFoundError(AUDITED_XLSX)
    if OUT_DIR.exists() and (OUT_DIR / "run_manifest.json").exists() and not args.resume:
        raise RuntimeError(f"existing Step 5.7-F output found; use --resume to continue: {OUT_DIR}")
    backup_dir = create_backup_dir()
    before = sha_snapshot(protected_paths())
    write_sha_snapshot(backup_dir / "protected_and_frozen_sha256_before.txt", before)
    recovery = {
        "process_check": process_check(),
        "audited_workbook": rel(AUDITED_XLSX),
        "audited_workbook_sha256": sha256_file(AUDITED_XLSX),
        "input_validation": validate_audited_input(AUDITED_XLSX),
        "backup_dir": str(backup_dir),
        "gold_test_labels_read": False,
        "external_llm_calls": 0,
        "local_32b_calls": 0,
    }
    if recovery["input_validation"]["status"] != "passed":
        write_json(backup_dir / "recovery_blocked.json", recovery)
        raise RuntimeError(f"audited input validation failed: {recovery['input_validation']}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_config()
    confirmation = generate_confirmed_workbook(AUDITED_XLSX, CONFIRMED_XLSX)
    validator_result = run_existing_step57e_validator(CONFIRMED_XLSX)
    split_manifest = build_training_data(CONFIRMED_XLSX)
    dev_metrics, eval_bundle = run_schemes(skip_neural=args.skip_neural)
    selection = eval_bundle["selection"]
    sha_compare = write_final_sha(backup_dir, before)
    if sha_compare["status"] != "matched":
        raise RuntimeError(f"protected SHA changed: {sha_compare}")
    gpu_metrics = {
        "train_seconds_sum": round(sum(float(item.get("train_seconds", 0) or 0) for item in dev_metrics), 3),
        "peak_gpu_memory_mb_max": round(max([float(item.get("peak_gpu_memory_mb", 0) or 0) for item in dev_metrics] + [0.0]), 3),
    }
    report = {
        "generated_at": now(),
        "status": selection["status"],
        "reference": "AI-assisted, user-confirmed EventMention structure data; not independent human gold.",
        "recovery": recovery,
        "confirmation": {
            **confirmation,
            "confirmed_workbook": rel(CONFIRMED_XLSX),
            "confirmed_workbook_sha256": sha256_file(CONFIRMED_XLSX),
            "existing_validator_result": validator_result,
        },
        "data": split_manifest,
        "dev_metrics": dev_metrics,
        "selection": selection,
        "legacy80_diagnostic": eval_bundle["legacy80_diagnostic"],
        "gpu": gpu_metrics,
        "protection": {
            "formal_task1_task2_modified": False,
            "protected_and_frozen_sha_status": sha_compare["status"],
            "gold_test_labels_read": False,
            "gold_test_content_read": False,
            "training_or_finetune_scope": "new200_train_only; model selection on new200_dev only",
            "old80_used_for_model_selection": False,
            "local_32b_calls": 0,
            "external_llm_calls": 0,
            "vllm_started": False,
        },
        "backup_dir": str(backup_dir),
    }
    write_json(REPORT_JSON, report)
    write_json(OUT_DIR / "run_manifest.json", report)
    write_report_md(report)
    write_details_xlsx(dev_metrics, selection, split_manifest, eval_bundle["legacy80_diagnostic"])
    write_json(OUT_DIR / "SHA256SUMS.json", {rel(path): sha256_file(path) for path in [CONFIRMED_XLSX, REPORT_JSON, REPORT_MD, DETAILS_XLSX, CONFIG_PATH] if path.exists()})
    print(json.dumps({"status": report["status"], "best_scheme": selection["best_scheme"], "dev_gate_passed": selection["dev_gate_passed"], "backup_dir": str(backup_dir)}, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--resume", action="store_true", default=True)
    parser.add_argument("--skip-neural", action="store_true")
    parser.add_argument("--neural-worker", action="store_true")
    parser.add_argument("--scheme", default="")
    parser.add_argument("--weighted", action="store_true")
    args = parser.parse_args()
    if args.neural_worker:
        return run_neural_worker(args.scheme, args.weighted)
    return run_pipeline(args)


if __name__ == "__main__":
    raise SystemExit(main())
