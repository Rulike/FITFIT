#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-V sentence selector and constrained PP-UIE decoding.

This step freezes the Step 5.7-U 8800-step PP-UIE checkpoint as a short-field
model.  It does not train the full 850-row structure model, does not read Dev90
labels, frozen60 labels, or 1029 gold labels, and does not call any LLM.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import random
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import joblib
import paddle
import paddle.nn.functional as F
from difflib import SequenceMatcher
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from paddlenlp.transformers import AutoModelForCausalLM, AutoTokenizer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_task1_step5_7f_structure_training import evaluate_predictions  # noqa: E402
from scripts.run_task1_step5_7s_ppuie_multisample_memory import (  # noqa: E402
    DIRS,
    FIELD_QUERIES,
    MODEL_DIR,
    NO_ENTITY,
    ROLES,
    TRAIN850,
    encode_shifted,
    eval_struct,
    field_target,
    first_generated_line,
    generate_answer,
    parser_preservation,
    pick_memory,
    render_official_prompt,
    sha256_file,
)
from scripts.run_task1_step5_7u_event_span_hybrid import (  # noqa: E402
    FORMAL_PROTECTED,
    normalize_time_u,
    split_sentences,
)


OUT_DIR = ROOT / "outputs/task1_step5_7v_sentence_selector"
REPORT_JSON = ROOT / "reports/task1_step5_7v_sentence_selector.json"
REPORT_MD = ROOT / "reports/task1_step5_7v_sentence_selector.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7v_sentence_selector_details.xlsx"
CONSTRAINED_MD = ROOT / "reports/task1_step5_7v_constrained_decoding_audit.md"
U_CKPT = ROOT / "outputs/task1_step5_7u_event_span_hybrid/checkpoints/U_original_joint_8800"
U_RESULT = U_CKPT / "experiment_result.json"
U_REPORT = ROOT / "reports/task1_step5_7u_event_span_hybrid.json"
TASK1234 = Path("/opt/data/private/JinRong/kechuang/task1234_delivery_20260811.tar.gz")
DEV90 = ROOT / "outputs/task1_step5_7n_chatgpt500_confirmation/merged_structure_dev_90.jsonl"
FROZEN60 = ROOT / "outputs/task1_step5_7j_structure_confirmation/independent_frozen_test_confirmed.jsonl"

SHORT_FIELDS = ["trigger", "subject_company", "event_time", "object_target"]
ALL_SHORT_FOR_REPRO = ["trigger", "subject_company", "event_time", "object_target"]
ACTION_WORDS = [
    "发布", "推出", "上线", "启动", "落地", "签署", "签约", "合作", "中标", "获批", "增长",
    "下降", "上涨", "下跌", "增产", "减产", "扩产", "投产", "停产", "涨价", "降价", "突破",
    "创新", "并购", "收购", "投资", "融资", "上市", "开工", "竣工", "交付", "成立", "建设",
    "升级", "调整", "召回", "处罚", "亏损", "盈利", "扭亏", "预增", "预减", "供货", "采购",
    "恢复", "维持", "宣布", "完成", "提高", "减少", "增加", "扩大", "降低", "提升", "进入",
    "投建", "招标", "备案", "通过", "实施", "推进", "加快", "实现", "计划", "预计",
]
NEGATIVE_WORDS = ["风险", "违约", "亏损", "下跌", "下降", "下滑", "处罚", "召回", "减产", "停产", "延期", "债务", "危机", "制裁", "终止", "取消"]
POSITIVE_WORDS = ["增长", "上涨", "突破", "获批", "中标", "签约", "扩产", "投产", "盈利", "提升", "创新", "拓展", "完成", "交付", "发布", "推出", "升级"]
NEUTRAL_WORDS = ["发布", "宣布", "计划", "预计", "拟", "称", "表示", "讨论", "披露", "召开"]


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def sha_snapshot(paths: Iterable[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in paths if path.exists() and path.is_file()}


def write_sha_txt(path: Path, snapshot: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(snapshot.items())), encoding="utf-8")


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def validate_sentences(rows: list[dict[str, Any]]) -> dict[str, Any]:
    bad = []
    total = 0
    for row in rows:
        content = row.get("content", "")
        for sent in split_sentences(content):
            total += 1
            if content[sent["start"] : sent["end"]] != sent["text"]:
                bad.append({"event_mention_id": row.get("event_mention_id"), "sentence_id": sent.get("sentence_id")})
    return {"sentence_count": total, "bad_offset_count": len(bad), "bad_offsets": bad[:20]}


def field_span_empty(field: str = "") -> dict[str, Any]:
    if field == "event_time":
        return {"text": "", "start": -1, "end": -1, "normalized": "未明确"}
    return {"text": "", "start": -1, "end": -1}


def event_terms(event_name: str) -> list[str]:
    raw = str(event_name or "")
    terms = [raw]
    for sep in ["与", "及", "和", "、", "/", "+", "（", "）", "(", ")", "，", ","]:
        parts = []
        for item in terms:
            parts.extend(item.split(sep))
        terms.extend(parts)
    terms.extend(re.findall(r"[\u4e00-\u9fff]{2,6}", raw))
    return list(dict.fromkeys(t.strip() for t in terms if len(t.strip()) >= 2))


def make_windows(row: dict[str, Any], max_size: int = 3) -> list[dict[str, Any]]:
    content = row.get("content", "")
    sentences = split_sentences(content)
    windows: list[dict[str, Any]] = []
    seen: set[tuple[int, int]] = set()
    for i, sent in enumerate(sentences):
        for size in range(1, max_size + 1):
            j = i + size - 1
            if j >= len(sentences):
                continue
            start, end = sentences[i]["start"], sentences[j]["end"]
            if (start, end) in seen:
                continue
            seen.add((start, end))
            text = content[start:end]
            windows.append(
                {
                    "candidate_id": f"W{i + 1:04d}_{size}",
                    "sentence_ids": [sentences[k]["sentence_id"] for k in range(i, j + 1)],
                    "window_size": size,
                    "text": text,
                    "start": start,
                    "end": end,
                    "sentence_index": i,
                }
            )
    if not windows and content:
        windows.append({"candidate_id": "W0001_1", "sentence_ids": ["S0001"], "window_size": 1, "text": content, "start": 0, "end": len(content), "sentence_index": 0})
    return windows


def overlap_len(a_start: int, a_end: int, b_start: int, b_end: int) -> int:
    return max(0, min(a_end, b_end) - max(a_start, b_start))


def char_f1_span(gold: dict[str, Any], pred: dict[str, Any]) -> float:
    gs, ge = int(gold.get("start", -1)), int(gold.get("end", -1))
    ps, pe = int(pred.get("start", -1)), int(pred.get("end", -1))
    if gs < 0 or ge <= gs or ps < 0 or pe <= ps:
        return 0.0
    ov = overlap_len(gs, ge, ps, pe)
    if ov <= 0:
        return 0.0
    p = ov / max(1, pe - ps)
    r = ov / max(1, ge - gs)
    return 2 * p * r / max(1e-12, p + r)


def gold_window(row: dict[str, Any], max_size: int = 3) -> dict[str, Any]:
    gold = row.get("event_span", {})
    gs, ge = int(gold.get("start", -1)), int(gold.get("end", -1))
    best = None
    best_key = (-1.0, -999999)
    for win in make_windows(row, max_size):
        contains = win["start"] <= gs and ge <= win["end"]
        f1 = char_f1_span(gold, win)
        key = (1.0 if contains else f1, -len(win["text"]))
        if key > best_key:
            best_key = key
            best = dict(win)
            best["contains_gold_span"] = contains
            best["gold_char_f1"] = round(f1, 6)
    return best or {"candidate_id": "", "sentence_ids": [], "text": "", "start": -1, "end": -1, "contains_gold_span": False, "gold_char_f1": 0.0}


def build_lexicons(rows: list[dict[str, Any]]) -> dict[str, Any]:
    lex: dict[str, Any] = {
        "trigger_by_event": defaultdict(Counter),
        "subject_by_event": defaultdict(Counter),
        "object_by_event": defaultdict(Counter),
        "time_terms": Counter(),
        "event_support": Counter(),
    }
    for row in rows:
        ev = row.get("event_name", "")
        if not ev:
            continue
        lex["event_support"][ev] += 1
        for field, key in [("trigger", "trigger_by_event"), ("subject_company", "subject_by_event"), ("object_target", "object_by_event")]:
            text = str(row.get(field, {}).get("text", "") or "").strip()
            if text:
                lex[key][ev][text] += 1
        text = str(row.get("event_time", {}).get("text", "") or "").strip()
        if text:
            lex["time_terms"][text] += 1
    return lex


def title_overlap(title: str, text: str) -> float:
    chars = {ch for ch in str(title or "") if "\u4e00" <= ch <= "\u9fff"}
    if not chars:
        return 0.0
    return sum(1 for ch in chars if ch in text) / max(1, len(chars))


def rule_score(row: dict[str, Any], win: dict[str, Any], lex: dict[str, Any]) -> float:
    text = win["text"]
    ev = row.get("event_name", "")
    score = 0.0
    for term in event_terms(ev):
        if term in text:
            score += min(2.0, len(term) * 0.18)
    for term, cnt in lex["trigger_by_event"].get(ev, {}).most_common(20):
        if term and term in text:
            score += 3.0 + math.log1p(cnt)
    for term, cnt in lex["subject_by_event"].get(ev, {}).most_common(20):
        if term and term in text:
            score += 1.1 + 0.2 * math.log1p(cnt)
    for term, cnt in lex["object_by_event"].get(ev, {}).most_common(20):
        if term and term in text:
            score += 0.9 + 0.2 * math.log1p(cnt)
    score += 0.25 * sum(1 for word in ACTION_WORDS if word in text)
    if re.search(r"\d{4}年|\d{1,2}月\d{1,2}日|近日|今日|昨日|当日|上半年|季度", text):
        score += 0.35
    score += 0.8 * title_overlap(row.get("title", ""), text)
    score += max(0.0, 0.4 - 0.035 * win.get("sentence_index", 0))
    if any(bad in text for bad in ["免责声明", "风险提示", "仅供参考", "文章来源", "责任编辑"]):
        score -= 2.0
    if len(text) < 8:
        score -= 1.0
    if win.get("window_size") == 2:
        score -= 0.05
    elif win.get("window_size") == 3:
        score -= 0.2
    return score


def rank_rule(row: dict[str, Any], lex: dict[str, Any]) -> list[dict[str, Any]]:
    wins = [dict(w) for w in make_windows(row)]
    for win in wins:
        win["score"] = rule_score(row, win, lex)
        win["selector"] = "A_rule_ranking"
    wins.sort(key=lambda x: (-x["score"], len(x["text"]), x["start"]))
    return wins


def pair_text(row: dict[str, Any], win: dict[str, Any]) -> str:
    return f"事件名称：{row.get('event_name','')}。标题：{row.get('title','')}。候选句：{win.get('text','')}"


def train_tfidf_selector(train_rows: list[dict[str, Any]], model_path: Path) -> dict[str, Any]:
    if model_path.exists():
        model = joblib.load(model_path)
        return {"trained": False, "model_path": rel(model_path), "model_sha256": sha256_file(model_path), "sample_count": None, "positive_count": None}
    xs: list[str] = []
    ys: list[int] = []
    for row in train_rows:
        gw = gold_window(row)
        best_id = gw.get("candidate_id")
        for win in make_windows(row):
            xs.append(pair_text(row, win))
            ys.append(1 if win["candidate_id"] == best_id else 0)
    if len(set(ys)) < 2:
        raise RuntimeError("sentence selector training labels have one class")
    model = Pipeline(
        [
            ("tfidf", TfidfVectorizer(analyzer="char", ngram_range=(2, 4), min_df=1, max_features=120000)),
            ("clf", LogisticRegression(max_iter=300, class_weight="balanced", random_state=2026)),
        ]
    )
    model.fit(xs, ys)
    model_path.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(model, model_path)
    return {"trained": True, "model_path": rel(model_path), "model_sha256": sha256_file(model_path), "sample_count": len(xs), "positive_count": int(sum(ys))}


def rank_tfidf(row: dict[str, Any], model: Any) -> list[dict[str, Any]]:
    wins = [dict(w) for w in make_windows(row)]
    if not wins:
        return []
    probs = model.predict_proba([pair_text(row, win) for win in wins])[:, 1]
    for win, prob in zip(wins, probs):
        win["score"] = float(prob)
        win["selector"] = "B_tfidf_lightweight_cross_encoder"
    wins.sort(key=lambda x: (-x["score"], len(x["text"]), x["start"]))
    return wins


def load_ppuie_checkpoint(dtype: str = "float16") -> tuple[Any, Any]:
    tokenizer = AutoTokenizer.from_pretrained(str(U_CKPT))
    model = AutoModelForCausalLM.from_pretrained(str(U_CKPT), dtype=dtype)
    model.eval()
    return tokenizer, model


def parameter_count(model: Any) -> int:
    return int(sum(int(p.numel()) for p in model.parameters()))


def ppuie_yes_no_loss(model: Any, tokenizer: Any, prompt: str, answer: str) -> float:
    enc = encode_shifted(tokenizer, prompt, answer, max_length=1024)
    ids = paddle.to_tensor([enc["input_ids"]], dtype="int64")
    labels = paddle.to_tensor(enc["labels"], dtype="int64")
    with paddle.no_grad():
        out = model(input_ids=ids, use_cache=False)
        logits = out.logits if hasattr(out, "logits") else out[0]
        loss = F.cross_entropy(logits.reshape([-1, logits.shape[-1]]).astype("float32"), labels, ignore_index=-100)
    return float(loss.numpy())


def relevance_prompt(tokenizer: Any, row: dict[str, Any], win: dict[str, Any]) -> str:
    raw = (
        "你是事件句判断助手。请判断候选句是否直接描述给定事件，只回答“是”或“否”。\n"
        f"事件名称：{row.get('event_name','')}\n"
        f"标题：{row.get('title','')}\n"
        f"候选句：{win.get('text','')}\n"
        "**回答开始**\n"
    )
    if tokenizer.chat_template is not None:
        return tokenizer.apply_chat_template(raw, tokenize=False)
    return raw


def rank_ppuie_scoring(row: dict[str, Any], model: Any, tokenizer: Any, prefilter: list[dict[str, Any]], max_candidates: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    selected = [dict(win) for win in prefilter[:max_candidates]]
    start = time.time()
    scored = []
    for win in selected:
        prompt = relevance_prompt(tokenizer, row, win)
        yes = ppuie_yes_no_loss(model, tokenizer, prompt, "是")
        no = ppuie_yes_no_loss(model, tokenizer, prompt, "否")
        item = dict(win)
        item["loss_yes"] = round(yes, 6)
        item["loss_no"] = round(no, 6)
        item["score"] = no - yes
        item["selector"] = "C_frozen_ppuie_candidate_scoring"
        scored.append(item)
    scored.sort(key=lambda x: (-x["score"], len(x["text"]), x["start"]))
    return scored, {"candidate_count": len(scored), "forward_count": len(scored) * 2, "latency_seconds": round(time.time() - start, 4)}


def selected_window_pred(row: dict[str, Any], win: dict[str, Any], scheme: str) -> dict[str, Any]:
    return {
        "event_mention_id": row.get("event_mention_id", ""),
        "article_id": row.get("article_id", ""),
        "scheme": scheme,
        "trigger": field_span_empty("trigger"),
        "event_span": {"text": win.get("text", ""), "start": win.get("start", -1), "end": win.get("end", -1)},
        "subject_company": field_span_empty("subject_company"),
        "event_time": field_span_empty("event_time"),
        "object_target": field_span_empty("object_target"),
        "direction": "neutral",
        "arguments": [],
        "argument_roles": [],
    }


def selector_metrics(rows: list[dict[str, Any]], rankings: dict[str, list[dict[str, Any]]], scheme: str) -> dict[str, Any]:
    preds = []
    top1 = top3 = top5 = illegal = containment = 0
    gold_single = gold_cross = single_hit = cross_hit = 0
    multi_total = multi_hit = 0
    avg_candidates = []
    latencies = []
    for row in rows:
        ranked = rankings.get(row["event_mention_id"], [])
        avg_candidates.append(len(ranked))
        gw = gold_window(row)
        gold_is_single = len(gw.get("sentence_ids", [])) == 1
        if gold_is_single:
            gold_single += 1
        else:
            gold_cross += 1
        if row.get("multi_event_article"):
            multi_total += 1
        def hit(win: dict[str, Any]) -> bool:
            return bool(win) and (win.get("start", 10**9) <= row.get("event_span", {}).get("start", -1) and row.get("event_span", {}).get("end", -1) <= win.get("end", -1))
        if ranked and hit(ranked[0]):
            top1 += 1
            if gold_is_single:
                single_hit += 1
            else:
                cross_hit += 1
            if row.get("multi_event_article"):
                multi_hit += 1
        if any(hit(win) for win in ranked[:3]):
            top3 += 1
        if any(hit(win) for win in ranked[:5]):
            top5 += 1
        chosen = ranked[0] if ranked else {"text": "", "start": -1, "end": -1}
        pred = selected_window_pred(row, chosen, scheme)
        trig = row.get("trigger", {})
        if chosen and chosen.get("start", 10**9) <= trig.get("start", -1) and trig.get("end", -1) <= chosen.get("end", -1):
            containment += 1
        content = row.get("content", "")
        span = pred["event_span"]
        if span["text"] and content[span["start"] : span["end"]] != span["text"]:
            illegal += 1
        preds.append(pred)
        if ranked and "latency_seconds" in ranked[0]:
            latencies.append(float(ranked[0]["latency_seconds"]))
    m = evaluate_predictions(rows, preds, scheme)
    n = max(1, len(rows))
    return {
        "sample_count": len(rows),
        "sentence_top1_accuracy": round(top1 / n, 6),
        "sentence_top3_recall": round(top3 / n, 6),
        "sentence_top5_recall": round(top5 / n, 6),
        "single_sentence_recall_top1": round(single_hit / max(1, gold_single), 6),
        "cross_sentence_recall_top1": round(cross_hit / max(1, gold_cross), 6),
        "multi_event_article_top1": round(multi_hit / max(1, multi_total), 6),
        "gold_single_count": gold_single,
        "gold_cross_count": gold_cross,
        "multi_event_article_count": multi_total,
        "average_candidate_count": round(sum(avg_candidates) / max(1, len(avg_candidates)), 3),
        "event_span_exact": m.get("event_span_exact_accuracy", 0.0),
        "event_span_char_f1": m.get("event_span_character_overlap_f1", 0.0),
        "trigger_containment": round(containment / n, 6),
        "illegal_offset_count": illegal,
        "average_latency_seconds": round(sum(latencies) / max(1, len(latencies)), 6) if latencies else 0.0,
        "raw_metrics": m,
    }


def window_size_metrics(rows: list[dict[str, Any]], base_rankings: dict[str, list[dict[str, Any]]], scheme_prefix: str) -> dict[str, Any]:
    result = {}
    for size in [1, 2, 3]:
        sized = {}
        for row in rows:
            items = [win for win in base_rankings.get(row["event_mention_id"], []) if win.get("window_size") == size]
            if not items:
                # Fall back to the first candidate of any size to keep the output defined.
                items = base_rankings.get(row["event_mention_id"], [])[:1]
            sized[row["event_mention_id"]] = items
        result[f"{size}_sentence_window"] = selector_metrics(rows, sized, f"{scheme_prefix}_window{size}")
    return result


def locate_exact(content: str, quote: str, preferred_start: int | None = None, preferred_end: int | None = None) -> dict[str, Any]:
    text = str(quote or "").strip()
    if not text or text == NO_ENTITY:
        return field_span_empty()
    positions = [m.start() for m in re.finditer(re.escape(text), content)]
    if not positions:
        return {"text": "", "start": -1, "end": -1, "failed_quote": text}
    if preferred_start is not None and preferred_end is not None:
        inside = [p for p in positions if preferred_start <= p and p + len(text) <= preferred_end]
        if len(inside) == 1:
            return {"text": text, "start": inside[0], "end": inside[0] + len(text), "resolution": "unique_in_selected_window"}
        if inside:
            p = min(inside, key=lambda x: abs(x - preferred_start))
            return {"text": text, "start": p, "end": p + len(text), "resolution": "multiple_in_selected_window_nearest_start", "ambiguous_count": len(inside)}
    if len(positions) == 1:
        p = positions[0]
        return {"text": text, "start": p, "end": p + len(text), "resolution": "unique_in_full_content"}
    p = min(positions)
    return {"text": text, "start": p, "end": p + len(text), "resolution": "multiple_in_full_content_first", "ambiguous_count": len(positions)}


def longest_common_substring(a: str, b: str) -> str:
    match = SequenceMatcher(None, a, b).find_longest_match(0, len(a), 0, len(b))
    return a[match.a : match.a + match.size]


def restricted_repair(content: str, window: dict[str, Any], generated: str, field: str) -> dict[str, Any]:
    raw = str(generated or "").strip()
    if not raw or raw == NO_ENTITY:
        return field_span_empty(field)
    exact = locate_exact(content, raw, window.get("start"), window.get("end"))
    if exact.get("text"):
        if field == "event_time":
            exact["normalized"] = "未明确"
        exact["repair_status"] = "exact"
        exact["original_generated_text"] = raw
        return exact
    win_text = window.get("text", "")
    compact_raw = re.sub(r"\s+", "", raw)
    compact_win = re.sub(r"\s+", "", win_text)
    if compact_raw and compact_raw in compact_win:
        # Map by searching the original window after whitespace removal fallback.
        for n in range(len(raw), 1, -1):
            for i in range(0, len(raw) - n + 1):
                sub = raw[i : i + n].strip()
                if len(sub) >= 2 and sub in win_text:
                    loc = locate_exact(content, sub, window.get("start"), window.get("end"))
                    loc["repair_status"] = "substring_from_generated"
                    loc["original_generated_text"] = raw
                    loc["repaired_quote"] = sub
                    if field == "event_time":
                        loc["normalized"] = "未明确"
                    return loc
    lcs = longest_common_substring(raw, win_text).strip()
    ratio = len(lcs) / max(1, len(raw))
    if len(lcs) >= 2 and ratio >= 0.5:
        loc = locate_exact(content, lcs, window.get("start"), window.get("end"))
        if loc.get("text"):
            loc["repair_status"] = "restricted_lcs_ngram"
            loc["original_generated_text"] = raw
            loc["repaired_quote"] = lcs
            loc["repair_similarity"] = round(ratio, 6)
            if field == "event_time":
                loc["normalized"] = "未明确"
            return loc
    return {"text": "", "start": -1, "end": -1, "repair_status": "unresolvable", "original_generated_text": raw, "failed_quote": raw, **({"normalized": "未明确"} if field == "event_time" else {})}


def direction_rule(text: str, event_name: str) -> str:
    source = f"{event_name} {text}"
    neg = sum(1 for word in NEGATIVE_WORDS if word in source)
    pos = sum(1 for word in POSITIVE_WORDS if word in source)
    neu = sum(1 for word in NEUTRAL_WORDS if word in source)
    if neg > pos and neg >= 1:
        return "negative"
    if pos > neg and pos >= 1:
        return "positive"
    if neu >= 1:
        return "neutral"
    return "neutral"


def train_direction_classifier(train_rows: list[dict[str, Any]], model_path: Path) -> dict[str, Any]:
    if model_path.exists():
        return {"trained": False, "model_path": rel(model_path), "model_sha256": sha256_file(model_path)}
    xs, ys = [], []
    for row in train_rows:
        direction = row.get("direction", "neutral")
        if direction not in DIRS:
            continue
        gw = gold_window(row)
        xs.append(pair_text(row, gw))
        ys.append(direction)
    model = Pipeline(
        [
            ("tfidf", TfidfVectorizer(analyzer="char", ngram_range=(2, 4), min_df=1, max_features=80000)),
            ("clf", LogisticRegression(max_iter=300, class_weight="balanced", random_state=2026)),
        ]
    )
    model.fit(xs, ys)
    model_path.parent.mkdir(parents=True, exist_ok=True)
    joblib.dump(model, model_path)
    return {"trained": True, "model_path": rel(model_path), "model_sha256": sha256_file(model_path), "sample_count": len(xs), "label_distribution": dict(Counter(ys))}


def predict_direction(row: dict[str, Any], win: dict[str, Any], model: Any | None) -> str:
    if model is not None:
        pred = str(model.predict([pair_text(row, win)])[0])
        if pred in DIRS:
            return pred
    return direction_rule(win.get("text", ""), row.get("event_name", ""))


def predict_short_fields_on_windows(
    rows: list[dict[str, Any]],
    selected: dict[str, dict[str, Any]],
    model: Any,
    tokenizer: Any,
    direction_model: Any | None,
) -> tuple[list[dict[str, Any]], dict[str, Any], list[dict[str, Any]]]:
    preds: list[dict[str, Any]] = []
    details: list[dict[str, Any]] = []
    raw_total = raw_located = constrained_total = constrained_located = 0
    repairs = Counter()
    repair_correct = repair_wrong = 0
    started = time.time()
    max_new = {"trigger": 32, "subject_company": 48, "event_time": 32, "object_target": 64}
    for row in rows:
        win = selected[row["event_mention_id"]]
        content = row.get("content", "")
        pred: dict[str, Any] = {
            "event_mention_id": row.get("event_mention_id", ""),
            "article_id": row.get("article_id", ""),
            "scheme": "V_actual_sentence_selector_ppuie_short_fields",
            "event_span": {"text": win.get("text", ""), "start": win.get("start", -1), "end": win.get("end", -1)},
        }
        row_detail = {"event_mention_id": row.get("event_mention_id"), "selected_candidate": win, "fields": {}}
        for field in SHORT_FIELDS:
            prompt_row = dict(row)
            prompt_row["content"] = win.get("text", "")
            out = generate_answer(model, tokenizer, render_official_prompt(tokenizer, prompt_row, field), max_new_tokens=max_new[field])
            raw = first_generated_line(out["text"])
            raw_total += 1
            if not raw or raw == NO_ENTITY:
                raw_located += 1
            exact = locate_exact(content, raw, win.get("start"), win.get("end"))
            if exact.get("text"):
                raw_located += 1
            repaired = restricted_repair(content, win, raw, field)
            if field == "event_time":
                repaired["normalized"] = normalize_time_u(repaired.get("text", ""), row.get("publish_time", ""))
            pred[field] = repaired if repaired.get("text") else field_span_empty(field)
            constrained_total += 1
            if not raw or raw == NO_ENTITY:
                constrained_located += 1
            elif repaired.get("text") and content[repaired["start"] : repaired["end"]] == repaired["text"]:
                constrained_located += 1
            repairs[repaired.get("repair_status", "empty")] += 1
            gold_text = row.get(field, {}).get("text", "") or ""
            if repaired.get("repair_status") not in {"exact", "empty"} and repaired.get("text"):
                if repaired.get("text") == gold_text:
                    repair_correct += 1
                else:
                    repair_wrong += 1
            row_detail["fields"][field] = {"raw": raw, "generation": out, "resolved": pred[field], "gold_for_eval": gold_text}
        if not pred["trigger"].get("text") or not (
            pred["event_span"]["start"] <= pred["trigger"].get("start", -1) and pred["trigger"].get("end", -1) <= pred["event_span"]["end"]
        ):
            # Do not fabricate a trigger.  This is counted by the downstream gate.
            row_detail["trigger_containment_error"] = True
        pred["direction"] = predict_direction(row, win, direction_model)
        args = [
            {"text": pred[role]["text"], "role": role, "start": pred[role]["start"], "end": pred[role]["end"]}
            for role in ROLES
            if pred.get(role, {}).get("text")
        ]
        pred["arguments"] = args
        pred["argument_roles"] = args
        preds.append(pred)
        details.append(row_detail)
    summary = {
        "raw_parser_preservation": round(raw_located / max(1, raw_total), 6),
        "constrained_parser_preservation": round(constrained_located / max(1, constrained_total), 6),
        "raw_total": raw_total,
        "raw_located": raw_located,
        "constrained_total": constrained_total,
        "constrained_located": constrained_located,
        "repair_status_counts": dict(repairs),
        "repair_count": int(sum(v for k, v in repairs.items() if k not in {"exact", "empty"})),
        "correct_repair_count": repair_correct,
        "wrong_repair_count": repair_wrong,
        "unresolvable_count": repairs.get("unresolvable", 0),
        "elapsed_seconds": round(time.time() - started, 4),
    }
    return preds, summary, details


def final_metrics(rows: list[dict[str, Any]], preds: list[dict[str, Any]], selector_metrics_obj: dict[str, Any], constrained: dict[str, Any]) -> dict[str, Any]:
    m = eval_struct(rows, preds, "V_actual_combination")
    text_ok = norm_ok = 0
    for row, pred in zip(rows, preds):
        text_ok += int((row.get("event_time", {}).get("text", "") or "") == (pred.get("event_time", {}).get("text", "") or ""))
        norm_ok += int((row.get("event_time", {}).get("normalized", "未明确") or "未明确") == (pred.get("event_time", {}).get("normalized", "未明确") or "未明确"))
    n = max(1, len(rows))
    containment_errors = sum(
        1
        for pred in preds
        if pred.get("trigger", {}).get("text")
        and not (pred["event_span"]["start"] <= pred["trigger"]["start"] and pred["trigger"]["end"] <= pred["event_span"]["end"])
    )
    metrics = {
        **m,
        "actual_sentence_top1_accuracy": selector_metrics_obj.get("sentence_top1_accuracy", 0.0),
        "event_time_text_accuracy": round(text_ok / n, 6),
        "event_time_normalized_accuracy": round(norm_ok / n, 6),
        "trigger_span_containment_error_count": containment_errors,
        "constrained_parser_preservation": constrained.get("constrained_parser_preservation", 0.0),
    }
    metrics["gate"] = {
        "actual_sentence_top1_ge_0_85": metrics.get("actual_sentence_top1_accuracy", 0.0) >= 0.85,
        "event_span_char_ge_0_85": metrics.get("event_span_character_overlap_f1", 0.0) >= 0.85,
        "trigger_relaxed_ge_0_90": metrics.get("trigger_relaxed_f1", 0.0) >= 0.90,
        "subject_norm_ge_0_85": metrics.get("subject_normalized_accuracy", 0.0) >= 0.85,
        "time_text_ge_0_80": metrics.get("event_time_text_accuracy", 0.0) >= 0.80,
        "time_norm_ge_0_80": metrics.get("event_time_normalized_accuracy", 0.0) >= 0.80,
        "object_norm_ge_0_80": metrics.get("object_normalized_accuracy", 0.0) >= 0.80,
        "direction_ge_0_90": metrics.get("direction_accuracy", 0.0) >= 0.90,
        "argument_role_ge_0_75": metrics.get("argument_role_end_to_end_f1", 0.0) >= 0.75,
        "constrained_parser_ge_0_99": metrics.get("constrained_parser_preservation", 0.0) >= 0.99,
        "illegal_offset_0": metrics.get("offset_illegal_count", 0) == 0,
        "article_first_stage_ge_0_60": metrics.get("article_first_stage_exact_accuracy", 0.0) >= 0.60,
    }
    return metrics


def reproduce_short_fields(rows: list[dict[str, Any]], model: Any, tokenizer: Any, u_result: dict[str, Any]) -> dict[str, Any]:
    stored_by_id = {pred.get("event_mention_id"): pred for pred in u_result.get("predictions", [])}
    preds: list[dict[str, Any]] = []
    raw_answers_for_parser: list[dict[str, str]] = []
    mismatches = []
    started = time.time()
    for row in rows:
        pred: dict[str, Any] = {
            "event_mention_id": row.get("event_mention_id", ""),
            "article_id": row.get("article_id", ""),
            "scheme": "V_short_field_reproduction",
            "event_span": field_span_empty("event_span"),
        }
        raw_answers: dict[str, str] = {}
        for field in SHORT_FIELDS:
            out = generate_answer(model, tokenizer, render_official_prompt(tokenizer, row, field), max_new_tokens={"trigger": 32, "subject_company": 48, "event_time": 32, "object_target": 64}[field])
            answer = first_generated_line(out["text"])
            raw_answers[field] = answer
            located = locate_exact(row.get("content", ""), answer)
            if field == "event_time":
                located["normalized"] = normalize_time_u(located.get("text", ""), row.get("publish_time", ""))
            pred[field] = located if located.get("text") else field_span_empty(field)
        pred["direction"] = stored_by_id.get(row["event_mention_id"], {}).get("direction", "neutral")
        args = [
            {"text": pred[role]["text"], "role": role, "start": pred[role]["start"], "end": pred[role]["end"]}
            for role in ROLES
            if pred.get(role, {}).get("text")
        ]
        pred["arguments"] = args
        pred["argument_roles"] = args
        preds.append(pred)
        raw_answers_for_parser.append(raw_answers)
        stored = stored_by_id.get(row["event_mention_id"], {})
        for field in SHORT_FIELDS:
            got = pred.get(field, {}).get("text", "")
            expected = stored.get(field, {}).get("text", "")
            if got != expected:
                mismatches.append({"event_mention_id": row["event_mention_id"], "field": field, "stored": expected, "reproduced": got})
    metrics = eval_struct(rows, preds, "V_short_field_reproduction")
    return {
        "sample_count": len(rows),
        "reproduction_mismatch_count": len(mismatches),
        "reproduction_error_zero": len(mismatches) == 0,
        "mismatches": mismatches[:50],
        "metrics": metrics,
        "parser_preservation_rate": parser_preservation(rows, raw_answers_for_parser, SHORT_FIELDS),
        "elapsed_seconds": round(time.time() - started, 4),
        "predictions": preds,
    }


def write_reports(report: dict[str, Any], details: list[dict[str, Any]], constrained_details: list[dict[str, Any]]) -> None:
    write_json(REPORT_JSON, report)
    lines = [
        "# Task1 Step 5.7-V 真实事件句选择器与原文约束解码",
        "",
        f"- status: `{report['status']}`",
        f"- frozen checkpoint: `{report['frozen_short_field_model']['checkpoint_path']}`",
        f"- frozen checkpoint SHA: `{report['frozen_short_field_model']['checkpoint_sha256']}`",
        f"- short-field reproduction mismatch count: `{report['short_field_reproduction']['reproduction_mismatch_count']}`",
        f"- best selector: `{report['best_selector_name']}`",
        f"- best selector Top1/Top3/Top5: `{report['best_selector_metrics']['sentence_top1_accuracy']}` / `{report['best_selector_metrics']['sentence_top3_recall']}` / `{report['best_selector_metrics']['sentence_top5_recall']}`",
        f"- final event_span char F1: `{report['final_combination_metrics'].get('event_span_character_overlap_f1')}`",
        f"- final trigger relaxed F1: `{report['final_combination_metrics'].get('trigger_relaxed_f1')}`",
        f"- constrained parser preservation: `{report['constrained_decoding']['constrained_parser_preservation']}`",
        f"- allow 850 training: `{report['recommendation']['allow_850_training']}`",
        "",
        "本轮未读取90 Dev、冻结60、1029 gold，未调用32B/GPT/外部LLM，未修改正式Task1/Task2输出。",
    ]
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    CONSTRAINED_MD.write_text(
        "\n".join(
            [
                "# Task1 Step 5.7-V 原文约束解码审计",
                "",
                f"- raw parser preservation: `{report['constrained_decoding']['raw_parser_preservation']}`",
                f"- constrained parser preservation: `{report['constrained_decoding']['constrained_parser_preservation']}`",
                f"- repair status counts: `{json.dumps(report['constrained_decoding']['repair_status_counts'], ensure_ascii=False)}`",
                f"- repair count: `{report['constrained_decoding']['repair_count']}`",
                f"- correct/wrong/unresolvable: `{report['constrained_decoding']['correct_repair_count']}` / `{report['constrained_decoding']['wrong_repair_count']}` / `{report['constrained_decoding']['unresolvable_count']}`",
                f"- illegal offset count: `{report['final_combination_metrics'].get('offset_illegal_count')}`",
                "",
                "无法可靠定位的输出被置空；未使用模型二次生成或gold字段修复。",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["key", "value"])
    for key in ["status", "frozen_short_field_model", "selector_metrics", "window_size_metrics", "final_combination_metrics", "protection"]:
        ws.append([key, json.dumps(report.get(key, {}), ensure_ascii=False)[:32000]])
    ws2 = wb.create_sheet("selector_details")
    ws2.append(["event_mention_id", "selector", "rank", "candidate_id", "sentence_ids", "window_size", "score", "gold_contains", "char_f1", "text"])
    for item in details:
        for rank, cand in enumerate(item.get("ranked", [])[:8], start=1):
            ws2.append([item["event_mention_id"], item["selector"], rank, cand.get("candidate_id"), ",".join(cand.get("sentence_ids", [])), cand.get("window_size"), cand.get("score"), cand.get("contains_gold"), cand.get("gold_char_f1"), cand.get("text")])
    ws3 = wb.create_sheet("constrained_decode")
    ws3.append(["event_mention_id", "field", "raw", "resolved_text", "start", "end", "repair_status", "gold_for_eval"])
    for row in constrained_details:
        for field, info in row.get("fields", {}).items():
            resolved = info.get("resolved", {})
            ws3.append([row.get("event_mention_id"), field, info.get("raw"), resolved.get("text"), resolved.get("start"), resolved.get("end"), resolved.get("repair_status"), info.get("gold_for_eval")])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.column_dimensions["A"].width = 28
        sheet.column_dimensions["B"].width = 36
        sheet.column_dimensions["J"].width = 100
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def load_macbert_selector_report() -> dict[str, Any]:
    path = OUT_DIR / "macbert_selector_report.json"
    if not path.exists():
        return {
            "run": False,
            "reason": "macbert selector report not found; run scripts/run_task1_step5_7v_macbert_sentence_selector.py in an environment with torch/transformers",
        }
    data = read_json(path)
    return {
        "run": True,
        "selector": data.get("selector"),
        "model_dir": data.get("model_dir"),
        "base_model_dir": data.get("base_model_dir"),
        "base_model_sha256": data.get("base_model_sha256"),
        "training": data.get("training", {}),
        "metrics": data.get("metrics", {}),
        "gpu": data.get("gpu", {}),
        "rankings_path": data.get("rankings_path"),
        "protection": data.get("protection", {}),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dtype", default="float16")
    parser.add_argument("--ppuie-score-max-candidates", type=int, default=10)
    parser.add_argument("--skip-ppuie-selector-c", action="store_true")
    parser.add_argument("--backup-dir", default="")
    args = parser.parse_args()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backup_dir = Path(args.backup_dir) if args.backup_dir else Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7v_sentence_selector_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    protected_paths = [
        *FORMAL_PROTECTED,
        TASK1234,
        MODEL_DIR / "model_state.pdparams",
        TRAIN850,
        U_CKPT / "model_state.pdparams",
        U_RESULT,
        U_REPORT,
    ]
    before = sha_snapshot(protected_paths)
    write_sha_txt(backup_dir / "protected_sha256_before.txt", before)
    write_sha_txt(OUT_DIR / "protected_sha256_before.txt", before)

    if not U_RESULT.exists() or not (U_CKPT / "model_state.pdparams").exists():
        raise FileNotFoundError(f"missing Step 5.7-U checkpoint: {U_CKPT}")
    rows = read_jsonl(TRAIN850)
    memory32 = pick_memory(rows, 32)
    memory32_ids = {row["event_mention_id"] for row in memory32}
    train_for_selector = [
        row for row in rows
        if row.get("event_mention_id") not in memory32_ids
        and row.get("event_exists", "yes") == "yes"
        and row.get("event_span", {}).get("text")
    ]
    write_jsonl(OUT_DIR / "fixed_memory32_rows_metadata_only.jsonl", [{"event_mention_id": r["event_mention_id"], "article_id": r["article_id"], "event_name": r["event_name"]} for r in memory32])
    sentence_validation = validate_sentences(memory32)
    lex = build_lexicons(train_for_selector)

    prompt_payload = {"official_ie_prompt": "see scripts/run_task1_step5_7s_ppuie_multisample_memory.py", "field_queries": FIELD_QUERIES, "short_fields": SHORT_FIELDS, "event_span_generation": "disabled"}
    prompt_file = OUT_DIR / "frozen_short_field_prompt_contract.json"
    write_json(prompt_file, prompt_payload)
    frozen_model = {
        "checkpoint_path": rel(U_CKPT),
        "checkpoint_sha256": sha256_file(U_CKPT / "model_state.pdparams"),
        "config_sha256": sha256_file(U_CKPT / "config.json"),
        "generation_config_sha256": sha256_file(U_CKPT / "generation_config.json"),
        "prompt_contract_path": rel(prompt_file),
        "prompt_sha256": sha256_file(prompt_file),
        "base_model_path": str(MODEL_DIR),
        "base_model_sha256": sha256_file(MODEL_DIR / "model_state.pdparams"),
        "training_steps": read_json(U_RESULT).get("steps"),
        "event_span_output_disabled": True,
    }

    print(json.dumps({"event": "load_frozen_ppuie", "checkpoint": frozen_model["checkpoint_path"]}, ensure_ascii=False), flush=True)
    tokenizer, model = load_ppuie_checkpoint(args.dtype)
    frozen_model["parameter_count"] = parameter_count(model)
    u_result = read_json(U_RESULT)
    reproduction = reproduce_short_fields(memory32, model, tokenizer, u_result)

    selector_b_path = OUT_DIR / "models/selector_b_tfidf_cross_encoder.joblib"
    selector_b_train = train_tfidf_selector(train_for_selector, selector_b_path)
    selector_b = joblib.load(selector_b_path)
    direction_model_path = OUT_DIR / "models/direction_tfidf_classifier.joblib"
    direction_train = train_direction_classifier(train_for_selector, direction_model_path)
    direction_model = joblib.load(direction_model_path)

    selector_details: list[dict[str, Any]] = []
    rankings: dict[str, dict[str, list[dict[str, Any]]]] = {"A_rule_ranking": {}, "B_tfidf_lightweight_cross_encoder": {}, "C_frozen_ppuie_candidate_scoring": {}}
    for row in memory32:
        a_rank = rank_rule(row, lex)
        b_rank = rank_tfidf(row, selector_b)
        rankings["A_rule_ranking"][row["event_mention_id"]] = a_rank
        rankings["B_tfidf_lightweight_cross_encoder"][row["event_mention_id"]] = b_rank
    c_meta = {"run": False, "reason": ""}
    if args.skip_ppuie_selector_c:
        c_meta = {"run": False, "reason": "skipped_by_argument"}
    else:
        c_forward = 0
        c_latency = 0.0
        for row in memory32:
            pre = rankings["B_tfidf_lightweight_cross_encoder"][row["event_mention_id"]]
            c_rank, meta = rank_ppuie_scoring(row, model, tokenizer, pre, args.ppuie_score_max_candidates)
            rankings["C_frozen_ppuie_candidate_scoring"][row["event_mention_id"]] = c_rank
            c_forward += int(meta["forward_count"])
            c_latency += float(meta["latency_seconds"])
        c_meta = {"run": True, "forward_count": c_forward, "latency_seconds": round(c_latency, 4), "max_candidates_per_row": args.ppuie_score_max_candidates}

    selector_metrics_obj = {}
    for name, rankmap in rankings.items():
        if not rankmap:
            selector_metrics_obj[name] = {"run": False}
            continue
        selector_metrics_obj[name] = {"run": True, **selector_metrics(memory32, rankmap, name)}
        for row in memory32:
            annotated = []
            for cand in rankmap.get(row["event_mention_id"], [])[:8]:
                c = dict(cand)
                c["contains_gold"] = c.get("start", 10**9) <= row.get("event_span", {}).get("start", -1) and row.get("event_span", {}).get("end", -1) <= c.get("end", -1)
                c["gold_char_f1"] = round(char_f1_span(row.get("event_span", {}), c), 6)
                annotated.append(c)
            selector_details.append({"event_mention_id": row["event_mention_id"], "selector": name, "ranked": annotated})

    runnable_selectors = {k: v for k, v in selector_metrics_obj.items() if v.get("run")}
    best_selector_name = max(
        runnable_selectors,
        key=lambda k: (
            runnable_selectors[k].get("sentence_top1_accuracy", 0.0),
            runnable_selectors[k].get("event_span_char_f1", 0.0),
            runnable_selectors[k].get("sentence_top3_recall", 0.0),
        ),
    )
    best_rankings = rankings[best_selector_name]
    selected = {eid: items[0] for eid, items in best_rankings.items() if items}
    window_metrics = window_size_metrics(memory32, best_rankings, best_selector_name)

    final_preds, constrained_summary, constrained_details = predict_short_fields_on_windows(memory32, selected, model, tokenizer, direction_model)
    final = final_metrics(memory32, final_preds, selector_metrics_obj[best_selector_name], constrained_summary)
    write_jsonl(OUT_DIR / "final_combination_predictions.jsonl", final_preds)
    write_jsonl(OUT_DIR / "constrained_decoding_details.jsonl", constrained_details)

    if all(final["gate"].values()):
        status = "ready_for_sentence_selector_and_ppuie_850_training"
    elif not final["gate"]["actual_sentence_top1_ge_0_85"] or not final["gate"]["event_span_char_ge_0_85"]:
        status = "event_sentence_selector_failed"
    elif not final["gate"]["constrained_parser_ge_0_99"] or not final["gate"]["illegal_offset_0"]:
        status = "constrained_decoding_failed"
    else:
        status = "ppuie_short_field_generalization_not_ready"

    after = sha_snapshot(protected_paths)
    protection = {
        "protected_sha_match": before == after,
        "formal_task1_task2_modified": before != after,
        "dev90_labels_read": False,
        "independent_frozen60_labels_read": False,
        "gold1029_labels_read": False,
        "train850_full_structure_training_run": False,
        "ppuie_short_field_weight_updates": 0,
        "ppuie_15b_tested": False,
        "local_32b_calls": 0,
        "gpt_or_external_llm_calls": 0,
        "task1234_sha_before": before.get(rel(TASK1234)),
        "task1234_sha_after": after.get(rel(TASK1234)),
    }
    macbert_report = load_macbert_selector_report()
    report = {
        "generated_at": now(),
        "status": status,
        "backup_dir": str(backup_dir),
        "input_counts": {"train850_total": len(rows), "fixed32_count": len(memory32), "selector_train_count_excluding_fixed32": len(train_for_selector)},
        "sentence_validation": sentence_validation,
        "frozen_short_field_model": frozen_model,
        "short_field_reproduction": {k: v for k, v in reproduction.items() if k != "predictions"},
        "selector_training": {"B_tfidf_auxiliary": selector_b_train, "direction_classifier": direction_train, "B_macbert_cross_encoder": macbert_report},
        "selector_metrics_fair_comparison": {
            "A_rule_ranking": selector_metrics_obj.get("A_rule_ranking"),
            "B_macbert_lightweight_cross_encoder": macbert_report.get("metrics") if macbert_report.get("run") else {"run": False, "reason": macbert_report.get("reason")},
            "C_frozen_ppuie_candidate_scoring": selector_metrics_obj.get("C_frozen_ppuie_candidate_scoring"),
        },
        "selector_c_ppuie_scoring": c_meta,
        "selector_metrics": selector_metrics_obj,
        "best_selector_name": best_selector_name,
        "best_selector_metrics": selector_metrics_obj[best_selector_name],
        "window_size_metrics": window_metrics,
        "constrained_decoding": constrained_summary,
        "final_combination_metrics": final,
        "recommendation": {
            "allow_850_training": status == "ready_for_sentence_selector_and_ppuie_850_training",
            "next_stage": "train sentence selector and short-field PP-UIE on 850, then evaluate Dev90; do not read frozen60 until Dev gate passes",
            "estimated_850_training": {"short_field_model": "reuse frozen PP-UIE architecture with 850-row training only in next step", "expected_gpu_memory_gb": "18-24", "expected_time_hours": "4-8 for split short-field plus selector experiments"},
            "worth_expanding_10k": False,
        },
        "protection": protection,
    }
    write_reports(report, selector_details, constrained_details)
    write_sha_txt(OUT_DIR / "protected_sha256_after.txt", after)
    write_sha_txt(backup_dir / "protected_sha256_after.txt", after)
    write_json(backup_dir / "run_manifest.json", report)
    sha_targets = [
        REPORT_JSON,
        REPORT_MD,
        DETAILS_XLSX,
        CONSTRAINED_MD,
        OUT_DIR / "final_combination_predictions.jsonl",
        OUT_DIR / "constrained_decoding_details.jsonl",
        selector_b_path,
        direction_model_path,
        prompt_file,
    ]
    write_sha_txt(OUT_DIR / "SHA256SUMS.txt", sha_snapshot(sha_targets))
    print(json.dumps({"status": status, "best_selector": best_selector_name, "top1": report["best_selector_metrics"].get("sentence_top1_accuracy"), "event_span_char_f1": final.get("event_span_character_overlap_f1"), "trigger_relaxed_f1": final.get("trigger_relaxed_f1"), "protected_sha_match": protection["protected_sha_match"], "backup_dir": str(backup_dir)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
