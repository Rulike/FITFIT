#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build Task4 severity candidates by re-entering targetable raw/shadow into Task1.

This is an inference/export script only. It does not train models and does not
modify Task1/Task3/Task4 core artifacts or existing human Gold.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import pickle
import random
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.task1_step5_8hj_common import (  # noqa: E402
    allowed_official_events_by_type,
    build_event_pair_text,
    build_training_rows,
    event_schema_map,
    load_candidate_generator,
    predict_target_ranker,
)
from scripts.run_task1_step5_7k_macbert_worker import (  # noqa: E402
    as_span as structure_as_span,
    empty_span as structure_empty_span,
    load_checkpoint as load_structure_checkpoint,
    predict_rows as predict_structure_rows,
)


SEED = 20260824
TARGETABLE_XLSX = ROOT / "outputs/task4_targetable_raw_v1/targetable_corpus.xlsx"
FULLTEXT_INDEX = ROOT / "data/processed/event_credibility/fulltext_article_index.parquet"
SHADOW_MANIFEST = ROOT / "outputs/task34_raw_shadow_inference/eligible_document_manifest.parquet"
GOLD_XLSX = ROOT / "outputs/task4 golddata/task4_severity_training_data_v1.xlsx"
TASK3_RELEASE = ROOT / "outputs/task3_kg_v3_release"

TARGET_MODEL = ROOT / "outputs/task1_step5_8hj_training/models/target/target_word_seed_2026/target_model.pkl"
EVENT_MODEL = ROOT / "outputs/task1_step5_8hj_training/models/event_name/pair_ce_f8_seed_2028/final_model"
STRUCTURE_CHECKPOINT = ROOT / "outputs/task1_step5_7k_structure_extractor/experiment_C/model.pt"

OUT_DIR = ROOT / "outputs/task1_severity_event_expansion_v1"
REPORT_MD = ROOT / "reports/task1_severity_event_expansion_v1.md"

FORMAL_COLUMNS = [
    "sample_id",
    "target_name",
    "target_type",
    "event_name",
    "event_type",
    "title",
    "evidence_sentence",
    "local_context",
    "financial_facts",
    "publish_time",
    "severity_label",
    "risk_intensity",
    "annotation_confidence",
    "annotation_basis",
    "label_tier",
    "label_source",
    "target_local_status",
    "split_status",
    "provenance",
    "provenance_status",
    "supervision_eligible",
]

HUMAN_COLUMNS = {"severity_label", "risk_intensity", "annotation_confidence", "annotation_basis"}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HUMAN_FILL = PatternFill("solid", fgColor="FFF2CC")
META_FILL = PatternFill("solid", fgColor="E2F0D9")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

SENTENCE_END = set("。！？!?；;\n")
NUM_RE = re.compile(r"\d+(?:\.\d+)?(?:%|亿元|亿元人民币|万|万元|亿|元|倍|期|个|家|笔|天)")

RISK_FAMILIES = {
    "cashflow_debt": [
        "现金流",
        "经营现金流",
        "流动性",
        "资金链",
        "偿债",
        "债务",
        "违约",
        "逾期",
        "展期",
        "兑付",
        "本息",
        "被执行",
        "强制执行",
        "资产负债率",
        "利息保障",
    ],
    "performance_profit": [
        "亏损",
        "净亏损",
        "预亏",
        "利润下降",
        "利润下滑",
        "净利润下降",
        "营收下降",
        "业绩下滑",
        "扣非",
        "毛利率下降",
    ],
    "compliance": [
        "处罚",
        "行政处罚",
        "立案",
        "调查",
        "监管",
        "问询",
        "警示函",
        "纪律处分",
        "违法",
        "违规",
    ],
    "litigation_safety": [
        "诉讼",
        "仲裁",
        "起诉",
        "纠纷",
        "安全事故",
        "事故",
        "爆炸",
        "伤亡",
    ],
    "delisting_warning": ["退市", "风险警示", "*ST", "ST ", "终止上市"],
    "shutdown_default_extension": ["停产", "停工", "违约", "债务展期", "延期兑付", "无法兑付", "未兑付"],
    "asset_loss": ["资产损失", "重大损失", "计提减值", "资产减值", "信用减值", "商誉减值"],
    "going_concern": ["持续经营", "持续经营不确定性", "资不抵债", "破产", "重整", "债务重组"],
}
RISK_TERMS = sorted({term for terms in RISK_FAMILIES.values() for term in terms}, key=len, reverse=True)
FORCE_EVENT_LABELS = {
    "现金流风险",
    "市场竞争",
    "市场竞争加剧",
    "行业下游的需求变化",
    "大宗商品景气与价格波动加剧",
    "发生地缘政治事件",
    "地缘政治",
    "资本市场运作",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def compact(value: Any) -> str:
    return re.sub(r"\s+", "", clean(value)).lower()


def norm_code(value: Any) -> str:
    text = clean(value)
    if text.endswith(".0"):
        text = text[:-2]
    return text.zfill(6) if text.isdigit() and len(text) <= 6 else text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_id(prefix: str, *parts: Any) -> str:
    payload = "||".join(clean(x) for x in parts)
    return f"{prefix}_{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:24]}"


def json_default(obj: Any) -> Any:
    if isinstance(obj, (pd.Timestamp,)):
        return obj.isoformat()
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        if math.isnan(float(obj)):
            return None
        return float(obj)
    if isinstance(obj, (np.ndarray,)):
        return obj.tolist()
    return str(obj)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=json_default) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, default=json_default) + "\n")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def parse_ids(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, float) and pd.isna(value):
        return []
    if isinstance(value, (list, tuple, set, np.ndarray, pd.Series)):
        return [clean(x) for x in value if clean(x)]
    text = clean(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return [clean(x) for x in parsed if clean(x)]
        return [clean(parsed)] if clean(parsed) else []
    except Exception:
        if text.startswith("[") and text.endswith("]"):
            parts = [p.strip().strip("\"'") for p in text[1:-1].split(",")]
            return [p for p in parts if p]
        return [text]


def sentence_spans(content: str) -> list[dict[str, Any]]:
    spans: list[dict[str, Any]] = []
    start = 0
    for idx, ch in enumerate(content):
        if ch in SENTENCE_END:
            end = idx + 1
            segment = content[start:end]
            lead = len(segment) - len(segment.lstrip())
            trail = len(segment) - len(segment.rstrip())
            real_start = start + lead
            real_end = end - trail
            if real_start < real_end:
                spans.append({"sentence_id": len(spans), "start": real_start, "end": real_end, "text": content[real_start:real_end]})
            start = end
    if start < len(content):
        segment = content[start:]
        lead = len(segment) - len(segment.lstrip())
        trail = len(segment) - len(segment.rstrip())
        real_start = start + lead
        real_end = len(content) - trail
        if real_start < real_end:
            spans.append({"sentence_id": len(spans), "start": real_start, "end": real_end, "text": content[real_start:real_end]})
    return spans or ([{"sentence_id": 0, "start": 0, "end": len(content), "text": content}] if content else [])


def find_sentence_index(spans: list[dict[str, Any]], start: int, end: int) -> int | None:
    if start is None or end is None or start < 0 or end <= start:
        return None
    for sent in spans:
        if sent["start"] <= start < end <= sent["end"]:
            return int(sent["sentence_id"])
    return None


def context_window(spans: list[dict[str, Any]], idx: int, radius: int = 1) -> str:
    selected = [s["text"] for s in spans if idx - radius <= int(s["sentence_id"]) <= idx + radius]
    return " ".join(selected)


def event_terms(event_name: str) -> list[str]:
    terms = [event_name]
    for part in re.split(r"[（(、/，,及与和等\s]+", event_name):
        part = part.strip("）) ")
        if len(part) >= 2:
            terms.append(part)
    for gram in re.findall(r"[\u4e00-\u9fff]{2,4}", event_name):
        if len(gram) >= 2:
            terms.append(gram)
    out: list[str] = []
    seen: set[str] = set()
    for term in terms:
        if term and term not in seen:
            seen.add(term)
            out.append(term)
    return out


def risk_hits(text: str) -> dict[str, list[str]]:
    hits: dict[str, list[str]] = {}
    for family, terms in RISK_FAMILIES.items():
        found = [term for term in terms if term in text]
        if found:
            hits[family] = found
    return hits


def risk_score(text: str) -> tuple[int, dict[str, list[str]]]:
    hits = risk_hits(text)
    score = 0
    for family, terms in hits.items():
        score += 4 if family in {"cashflow_debt", "delisting_warning", "going_concern", "shutdown_default_extension"} else 3
        score += min(len(terms), 4)
    if NUM_RE.search(text) and hits:
        score += 2
    if any(x in text for x in ["重大", "严重", "巨额", "大幅", "明显"]):
        score += 1
    return score, hits


def select_evidence(row: dict[str, Any], structure: dict[str, Any], target_name: str, event_name: str) -> dict[str, Any]:
    content = clean(row.get("structure_content")) or clean(row.get("content"))
    spans = sentence_spans(content)
    choices: list[tuple[int, int, str]] = []
    for field, weight in [("event_span", 8), ("trigger", 7), ("subject_company", 4), ("object_target", 3)]:
        span = structure.get(field) or {}
        idx = find_sentence_index(spans, int(span.get("start", -1) or -1), int(span.get("end", -1) or -1))
        if idx is not None:
            choices.append((weight, idx, field))
    event_kw = event_terms(event_name)
    for sent in spans:
        text = sent["text"]
        ctx = context_window(spans, int(sent["sentence_id"]), 1)
        score = 0
        if target_name and target_name in ctx:
            score += 5
        if any(term and term in ctx for term in event_kw):
            score += 3
        local_score, _ = risk_score(ctx)
        score += min(local_score, 8)
        if NUM_RE.search(ctx):
            score += 1
        if score > 0:
            choices.append((score, int(sent["sentence_id"]), "lexical_local"))
    if not choices:
        return {
            "evidence_sentence": "",
            "sentence_id": "",
            "local_context": "",
            "evidence_span": "",
            "local_window_used": "none",
        }
    _, idx, reason = max(choices, key=lambda item: (item[0], -(abs(item[1]))))
    sent = spans[idx]
    return {
        "evidence_sentence": sent["text"],
        "sentence_id": int(sent["sentence_id"]),
        "local_context": context_window(spans, idx, 1),
        "evidence_span": json.dumps({"start": sent["start"], "end": sent["end"]}, ensure_ascii=False),
        "local_window_used": reason,
    }


def structure_local_content(row: dict[str, Any], max_chars: int = 1800) -> str:
    content = clean(row.get("content"))
    if len(content) <= max_chars:
        return content
    spans = sentence_spans(content)
    target = clean(row.get("target_name"))
    event_name = clean(row.get("event_name"))
    event_kw = event_terms(event_name)
    scored: list[tuple[int, int]] = []
    for sent in spans:
        idx = int(sent["sentence_id"])
        ctx = context_window(spans, idx, 1)
        score = 0
        if target and target in ctx:
            score += 8
        if any(term and term in ctx for term in event_kw):
            score += 4
        local_risk, _ = risk_score(ctx)
        score += min(local_risk, 10)
        if NUM_RE.search(ctx):
            score += 1
        if score > 0:
            scored.append((score, idx))
    if not scored:
        return content[:max_chars]
    selected_ids: set[int] = set()
    for _, idx in sorted(scored, key=lambda item: (item[0], -item[1]), reverse=True):
        for j in range(max(0, idx - 1), min(len(spans), idx + 2)):
            selected_ids.add(j)
        snippet = " ".join(spans[j]["text"] for j in sorted(selected_ids))
        if len(snippet) >= max_chars:
            break
    snippet = " ".join(spans[j]["text"] for j in sorted(selected_ids))
    return snippet[:max_chars]


def install_fast_structure_decoder() -> None:
    def fast_choose_span_pair(start_logits: Any, end_logits: Any, exists_logit: Any, offsets: Any, mask: Any, content: str, field: str, mandatory: bool) -> tuple[dict[str, Any], float]:
        import torch

        valid = torch.nonzero(mask, as_tuple=False).view(-1)
        if valid.numel() == 0:
            return structure_empty_span(field), 0.0
        start_scores = torch.softmax(start_logits.masked_fill(~mask, -1e4), dim=-1)
        end_scores = torch.softmax(end_logits.masked_fill(~mask, -1e4), dim=-1)
        exists_score = float(torch.sigmoid(exists_logit).detach().cpu())
        if not mandatory and exists_score < 0.45:
            return structure_empty_span(field), exists_score
        max_tokens = 150 if field == "event_span" else 32
        best_s = int(valid[0].item())
        best_e = best_s
        best_score = -1.0
        for s_tensor in valid:
            s = int(s_tensor.item())
            end_valid = valid[(valid >= s) & (valid <= s + max_tokens)]
            if end_valid.numel() == 0:
                continue
            e_idx = int(torch.argmax(end_scores[end_valid]).item())
            e = int(end_valid[e_idx].item())
            score = float(start_scores[s] * end_scores[e]) * max(exists_score, 0.05)
            if score > best_score:
                best_s, best_e, best_score = s, e, score
        start = int(offsets[best_s][0])
        end = int(offsets[best_e][1])
        return structure_as_span(content, start, end, field), best_score

    predict_structure_rows.__globals__["choose_span_pair"] = fast_choose_span_pair


def parse_structure_text(structure: dict[str, Any], field: str) -> str:
    value = structure.get(field) or {}
    if isinstance(value, dict):
        return clean(value.get("text"))
    return clean(value)


def compute_grounding(row: dict[str, Any], structure: dict[str, Any], evidence: dict[str, Any]) -> dict[str, Any]:
    target = clean(row.get("target_name"))
    event = clean(row.get("event_name"))
    content = clean(row.get("content"))
    evidence_sentence = clean(evidence.get("evidence_sentence"))
    local_context = clean(evidence.get("local_context"))
    subject = parse_structure_text(structure, "subject_company")
    trigger = parse_structure_text(structure, "trigger")
    event_span = parse_structure_text(structure, "event_span")
    evidence_text = f"{evidence_sentence} {local_context}"

    target_in_content = bool(target and target in content)
    target_in_local = bool(target and target in evidence_text)
    subject_ok = True
    if subject and target:
        subject_ok = compact(target) in compact(subject) or compact(subject) in compact(target) or target in content

    event_local_terms = event_terms(event)
    event_signal = bool(event_span and event_span in evidence_text) or bool(trigger and trigger in evidence_text)
    if not event_signal:
        event_signal = any(term and term in evidence_text for term in event_local_terms)
    if not event_signal:
        event_signal = bool(risk_hits(evidence_text))

    evidence_nonempty = bool(evidence_sentence and len(evidence_sentence) >= 8 and evidence_sentence in content)
    provenance_valid = bool(clean(row.get("article_id")) and clean(row.get("source")) and clean(row.get("publish_time")) and content)
    target_grounded = bool(target_in_content and target_in_local and subject_ok)
    event_grounded = bool(event and evidence_nonempty and event_signal)
    evidence_grounded = bool(evidence_nonempty and target_grounded and event_grounded)

    reasons = []
    if not provenance_valid:
        reasons.append("provenance_or_content_missing")
    if not target_grounded:
        reasons.append("target_not_local_to_evidence")
    if not event_grounded:
        reasons.append("event_not_supported_by_local_evidence")
    if not evidence_grounded:
        reasons.append("evidence_not_grounded")
    return {
        "target_grounded": target_grounded,
        "event_grounded": event_grounded,
        "evidence_grounded": evidence_grounded,
        "provenance_valid": provenance_valid,
        "grounding_pass": bool(target_grounded and event_grounded and evidence_grounded and provenance_valid and row.get("train_only") and not row.get("validation_test_leakage")),
        "grounding_failure_reasons": reasons,
        "structure_subject": subject,
        "structure_trigger": trigger,
        "structure_event_span": event_span,
    }


def load_articles() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    targetable = pd.read_excel(TARGETABLE_XLSX, dtype=object).fillna("")
    full = pd.read_parquet(
        FULLTEXT_INDEX,
        columns=["dataset_split", "article_id", "article_title", "article_source", "article_publish_time", "content", "content_length"],
    )
    full["article_id"] = full["article_id"].map(clean)
    full_by_id = {row["article_id"]: row for row in full.drop_duplicates("article_id").to_dict("records")}

    shadow = pd.read_parquet(SHADOW_MANIFEST, columns=["canonical_document_id", "title", "content", "publish_time", "source", "original_ids"])
    shadow_map: dict[str, dict[str, Any]] = {}
    for rec in shadow.to_dict("records"):
        for aid in parse_ids(rec.get("original_ids")):
            shadow_map.setdefault(aid, rec)

    rows: list[dict[str, Any]] = []
    source_counts = Counter()
    for rec in targetable.to_dict("records"):
        aid = clean(rec.get("article_id"))
        direct = full_by_id.get(aid)
        mapped = shadow_map.get(aid)
        source_dataset = ""
        canonical_document_id = ""
        dataset_split = ""
        content = ""
        if direct and clean(direct.get("content")):
            content = clean(direct.get("content"))
            dataset_split = clean(direct.get("dataset_split"))
            source_dataset = "fulltext_article_index"
        elif mapped and clean(mapped.get("content")):
            content = clean(mapped.get("content"))
            canonical_document_id = clean(mapped.get("canonical_document_id"))
            dataset_split = "raw_shadow_train_only"
            source_dataset = "eligible_document_manifest"
        else:
            source_dataset = "missing_content"
        source_counts[source_dataset] += 1
        title = clean(rec.get("title")) or clean((direct or {}).get("article_title")) or clean((mapped or {}).get("title"))
        source = clean(rec.get("source")) or clean((direct or {}).get("article_source")) or clean((mapped or {}).get("source"))
        publish_time = clean(rec.get("publish_time")) or clean((direct or {}).get("article_publish_time")) or clean((mapped or {}).get("publish_time"))
        train_only = dataset_split in {"train", "labeled_train", "raw_shadow_train_only"}
        validation_test_leakage = dataset_split in {"test", "labeled_test", "validation", "val"}
        rows.append(
            {
                **rec,
                "article_id": aid,
                "article_file_id": aid,
                "article_title": title,
                "title": title,
                "article_source": source,
                "source": source,
                "article_publish_time": publish_time,
                "publish_time": publish_time,
                "content": content,
                "content_length": len(content),
                "dataset_split": dataset_split,
                "source_dataset": source_dataset,
                "canonical_document_id": canonical_document_id,
                "train_only": train_only,
                "validation_test_leakage": validation_test_leakage,
            }
        )
    meta = {
        "targetable_rows": len(targetable),
        "unique_article_id": int(targetable["article_id"].astype(str).nunique()),
        "source_dataset_counts": dict(source_counts),
        "content_available": int(sum(1 for r in rows if clean(r.get("content")))),
    }
    return rows, meta


def predict_events(
    rows: list[dict[str, Any]],
    generator: Any,
    batch_size: int,
    max_length: int,
    event_candidate_topk: int,
    device_arg: str,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    import torch
    from transformers import AutoModelForSequenceClassification, AutoTokenizer

    device = torch.device("cuda" if device_arg == "cuda" and torch.cuda.is_available() else "cpu")
    tokenizer = AutoTokenizer.from_pretrained(str(EVENT_MODEL), local_files_only=True)
    model = AutoModelForSequenceClassification.from_pretrained(str(EVENT_MODEL), local_files_only=True).to(device)
    model.eval()
    schema = event_schema_map()
    allowed = allowed_official_events_by_type()

    pair_texts: list[str] = []
    pair_meta: list[tuple[int, str, str, float, str]] = []
    candidate_counts: list[int] = []
    for idx, row in enumerate(rows):
        target_type = clean(row.get("target_type")) or "company"
        article = {
            "article_id": clean(row.get("article_id")),
            "title": clean(row.get("title")),
            "content": clean(row.get("content")),
            "publish_time": clean(row.get("publish_time")),
            "source": clean(row.get("source")),
        }
        generated, _ = generator.generate(article, include_all64=True)
        candidates = [cand for cand in generated if clean(cand.get("event_name")) in set(allowed.get(target_type, []))]
        if event_candidate_topk > 0 and len(candidates) > event_candidate_topk:
            forced = [cand for cand in candidates if clean(cand.get("event_name")) in FORCE_EVENT_LABELS]
            selected = candidates[:event_candidate_topk]
            seen = {clean(c.get("event_name")) for c in selected}
            for cand in forced:
                if clean(cand.get("event_name")) not in seen:
                    selected.append(cand)
                    seen.add(clean(cand.get("event_name")))
            candidates = selected
        candidate_counts.append(len(candidates))
        for cand in candidates:
            ev = clean(cand.get("event_name"))
            info = schema.get(ev, {})
            pair_texts.append(build_event_pair_text(row, ev, clean(info.get("raw_description") or info.get("description")), clean(info.get("raw_examples") or info.get("examples"))))
            pair_meta.append((idx, ev, clean(cand.get("schema_id")), float(cand.get("candidate_score") or 0.0), clean(cand.get("candidate_source"))))

    probs: list[float] = []
    logits_margin: list[float] = []
    with torch.no_grad():
        for start in range(0, len(pair_texts), batch_size):
            batch = pair_texts[start : start + batch_size]
            enc = tokenizer(batch, truncation=True, max_length=max_length, padding=True, return_tensors="pt")
            enc = {k: v.to(device) for k, v in enc.items()}
            logits = model(**enc).logits
            if logits.shape[-1] == 1:
                pos = torch.sigmoid(logits.view(-1))
                margin = logits.view(-1)
            else:
                pos = torch.softmax(logits, dim=-1)[:, 1]
                margin = logits[:, 1] - logits[:, 0]
            probs.extend(pos.detach().cpu().tolist())
            logits_margin.extend(margin.detach().cpu().tolist())

    grouped: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for meta, prob, margin in zip(pair_meta, probs, logits_margin):
        idx, ev, schema_id, candidate_score, source = meta
        grouped[idx].append(
            {
                "event_name": ev,
                "schema_id": schema_id,
                "event_score": float(prob),
                "event_logit_margin": float(margin),
                "candidate_score": candidate_score,
                "candidate_source": source,
            }
        )

    outputs: list[dict[str, Any]] = []
    for idx, row in enumerate(rows):
        scored = sorted(grouped.get(idx, []), key=lambda item: (item["event_score"], item["candidate_score"], item["event_name"]), reverse=True)
        best = scored[0] if scored else {}
        top5 = [item["event_name"] for item in scored[:5]]
        outputs.append(
            {
                "event_name": clean(best.get("event_name")),
                "schema_id": clean(best.get("schema_id")),
                "event_confidence": round(float(best.get("event_score") or 0.0), 6),
                "event_logit_margin": round(float(best.get("event_logit_margin") or 0.0), 6),
                "event_candidate_score": round(float(best.get("candidate_score") or 0.0), 6),
                "event_candidate_source": clean(best.get("candidate_source")),
                "event_top5": top5,
                "event_candidate_count": int(candidate_counts[idx] if idx < len(candidate_counts) else 0),
            }
        )
    meta = {
        "device": str(device),
        "pair_text_count": len(pair_texts),
        "batch_size": batch_size,
        "max_length": max_length,
        "event_candidate_topk": event_candidate_topk,
        "candidate_count_summary": {
            "min": int(min(candidate_counts) if candidate_counts else 0),
            "max": int(max(candidate_counts) if candidate_counts else 0),
            "mean": round(float(np.mean(candidate_counts)), 4) if candidate_counts else 0,
        },
    }
    return outputs, meta


def load_task3_lookup() -> dict[str, list[str]]:
    nodes_path = TASK3_RELEASE / "graph_nodes.jsonl"
    edges_path = TASK3_RELEASE / "graph_edges.jsonl"
    mention_keys: dict[str, str] = {}
    for row in read_jsonl(nodes_path):
        if clean(row.get("node_type")) != "EventMention":
            continue
        mention_id = clean(row.get("stable_id"))
        provisional = clean(row.get("provisional_event_type") or row.get("event_type"))
        if mention_id and provisional:
            mention_keys[mention_id] = provisional
    out: dict[str, list[str]] = defaultdict(list)
    for row in read_jsonl(edges_path):
        if clean(row.get("relation_type")) != "INSTANCE_OF_EVENT":
            continue
        src = clean(row.get("source_id"))
        tgt = clean(row.get("target_id"))
        key = mention_keys.get(src)
        if key and tgt:
            out[key].append(tgt)
    return {k: sorted(set(v)) for k, v in out.items()}


def load_gold_duplicate_keys() -> set[str]:
    keys: set[str] = set()
    if not GOLD_XLSX.exists():
        return keys
    for sheet in ["Gold_A_Verified", "Gold_B_Legacy"]:
        df = pd.read_excel(GOLD_XLSX, sheet_name=sheet, dtype=object).fillna("")
        for rec in df.to_dict("records"):
            title_key = compact(rec.get("title"))
            target_key = compact(rec.get("target_name"))
            event_key = compact(rec.get("event_name"))
            publish_key = clean(rec.get("publish_time"))[:10]
            evidence_key = compact(rec.get("evidence_sentence"))[:80]
            if title_key and target_key:
                keys.add(f"title|{target_key}|{event_key}|{title_key}")
            if target_key and event_key and evidence_key:
                keys.add(f"evidence|{target_key}|{event_key}|{publish_key}|{evidence_key}")
    return keys


def is_duplicate_with_gold(row: dict[str, Any], keys: set[str]) -> bool:
    target_key = compact(row.get("target_name"))
    event_key = compact(row.get("event_name"))
    title_key = compact(row.get("title"))
    publish_key = clean(row.get("publish_time"))[:10]
    evidence_key = compact(row.get("evidence_sentence"))[:80]
    return (
        f"title|{target_key}|{event_key}|{title_key}" in keys
        or f"evidence|{target_key}|{event_key}|{publish_key}|{evidence_key}" in keys
    )


def exact_remap_to_full_content(row: dict[str, Any], field: str) -> bool:
    content = clean(row.get("content"))
    text = clean(row.get(field))
    return bool(content and text and text in content)


def build_expansion_rows(
    rows: list[dict[str, Any]],
    target_preds: list[dict[str, Any]],
    event_preds: list[dict[str, Any]],
    structure_preds: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for base, target, event, structure in zip(rows, target_preds, event_preds, structure_preds):
        target_name = clean(target.get("target_name"))
        target_type = clean(target.get("target_type"))
        target_code = norm_code(target.get("target_code"))
        event_name = clean(event.get("event_name"))
        event_id = stable_id("T1E", base.get("article_id"), target_type, target_code, event_name)
        evidence = select_evidence(base, structure, target_name, event_name)
        merged = {
            **base,
            "sample_id": stable_id("T4S_T1", base.get("article_id"), target_type, target_code, event_name, evidence.get("evidence_sentence")),
            "event_id": event_id,
            "target_name": target_name,
            "target_type": target_type,
            "target_code": target_code,
            "entity_id": stable_id("ENT", target_type, target_code, target_name),
            "entity_mention": clean(target.get("matched_alias")) or target_name,
            "entity_confidence": round(float(target.get("score") or target.get("route_confidence") or 0.0), 6),
            "target_candidate_count": int(target.get("candidate_count") or 0),
            "event_name": event_name,
            "event_type": f"{target_type}_{target_code}_{event_name}" if target_type and target_code and event_name else "",
            "event_time": parse_structure_text(structure, "event_time") or "未明确",
            "event_status": clean(structure.get("direction")) or "unclear",
            "event_confidence": float(event.get("event_confidence") or 0.0),
            "event_candidate_count": int(event.get("event_candidate_count") or 0),
            "event_top5": event.get("event_top5") or [],
            "sentence_id": evidence.get("sentence_id"),
            "evidence_sentence": evidence.get("evidence_sentence"),
            "context_window": evidence.get("local_context"),
            "local_context": evidence.get("local_context"),
            "evidence_span": evidence.get("evidence_span"),
            "evidence_confidence": 1.0 if evidence.get("evidence_sentence") else 0.0,
            "local_window_used": evidence.get("local_window_used"),
            "structure_trigger": parse_structure_text(structure, "trigger"),
            "structure_event_span": parse_structure_text(structure, "event_span"),
            "structure_subject": parse_structure_text(structure, "subject_company"),
            "structure_object": parse_structure_text(structure, "object_target"),
        }
        grounding = compute_grounding(merged, structure, evidence)
        merged.update(grounding)
        local_text = clean(merged.get("local_context")) or clean(merged.get("evidence_sentence"))
        score, hits = risk_score(local_text)
        merged["high_risk_score"] = int(score)
        merged["high_risk_families"] = sorted(hits.keys())
        merged["high_risk_hits"] = hits
        merged["financial_fact_if_available"] = local_text if hits else ""
        merged["provenance"] = json.dumps(
            {
                "article_id": clean(merged.get("article_id")),
                "canonical_document_id": clean(merged.get("canonical_document_id")),
                "source": clean(merged.get("source")),
                "publish_time": clean(merged.get("publish_time")),
                "source_dataset": clean(merged.get("source_dataset")),
                "input": str(TARGETABLE_XLSX.relative_to(ROOT)),
                "target_model": str(TARGET_MODEL.relative_to(ROOT)),
                "event_model": str(EVENT_MODEL.relative_to(ROOT)),
                "structure_checkpoint": str(STRUCTURE_CHECKPOINT.relative_to(ROOT)),
            },
            ensure_ascii=False,
        )
        out.append(merged)
    return out


def attach_task3(rows: list[dict[str, Any]], task3_lookup: dict[str, list[str]]) -> list[dict[str, Any]]:
    for row in rows:
        key = clean(row.get("event_type"))
        matches = task3_lookup.get(key, [])
        row["canonical_match"] = bool(matches)
        row["canonical_event_id"] = matches[0] if matches else ""
        row["duplicate_status"] = "existing_event" if matches else "new_event_candidate"
    return rows


def deduplicate_and_select(rows: list[dict[str, Any]], max_rows: int) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    gold_keys = load_gold_duplicate_keys()
    eligible = [r for r in rows if r.get("grounding_pass")]
    seen_identity: set[str] = set()
    deduped: list[dict[str, Any]] = []
    duplicate_existing_gold = 0
    duplicate_within_batch = 0
    offset_not_remappable = 0
    for row in eligible:
        if is_duplicate_with_gold(row, gold_keys):
            row["candidate_exclusion_reason"] = "duplicate_with_existing_657_gold"
            duplicate_existing_gold += 1
            continue
        if not (
            exact_remap_to_full_content(row, "structure_trigger")
            and exact_remap_to_full_content(row, "structure_event_span")
            and exact_remap_to_full_content(row, "evidence_sentence")
        ):
            row["candidate_exclusion_reason"] = "offset_not_exactly_remappable_to_full_content"
            offset_not_remappable += 1
            continue
        identity = "|".join(
            [
                clean(row.get("target_type")),
                norm_code(row.get("target_code")),
                clean(row.get("event_name")),
                compact(row.get("evidence_sentence"))[:120],
            ]
        )
        if identity in seen_identity:
            row["candidate_exclusion_reason"] = "duplicate_within_new_pool"
            duplicate_within_batch += 1
            continue
        seen_identity.add(identity)
        row["candidate_exclusion_reason"] = ""
        deduped.append(row)

    high_risk = [r for r in deduped if int(r.get("high_risk_score") or 0) > 0]
    high_risk.sort(
        key=lambda r: (
            int(r.get("high_risk_score") or 0),
            float(r.get("entity_confidence") or 0.0),
            float(r.get("event_confidence") or 0.0),
            clean(r.get("publish_time")),
            clean(r.get("sample_id")),
        ),
        reverse=True,
    )

    selected: list[dict[str, Any]] = []
    target_counts: Counter[str] = Counter()
    source_counts: Counter[str] = Counter()
    family_counts: Counter[str] = Counter()
    for row in high_risk:
        target = clean(row.get("target_name"))
        source = clean(row.get("source"))
        families = row.get("high_risk_families") or ["unknown"]
        major_family = families[0] if families else "unknown"
        if target_counts[target] >= 10:
            continue
        if source_counts[source] >= 40:
            continue
        if family_counts[major_family] >= max(80, int(max_rows * 0.45)):
            continue
        selected.append(row)
        target_counts[target] += 1
        source_counts[source] += 1
        family_counts[major_family] += 1
        if len(selected) >= max_rows:
            break
    if len(selected) < min(max_rows, len(high_risk)):
        selected_ids = {r["sample_id"] for r in selected}
        for row in high_risk:
            if row["sample_id"] in selected_ids:
                continue
            selected.append(row)
            selected_ids.add(row["sample_id"])
            if len(selected) >= max_rows:
                break

    meta = {
        "grounding_eligible_before_dedup": len(eligible),
        "duplicate_with_existing_657_gold": duplicate_existing_gold,
        "duplicate_within_new_pool": duplicate_within_batch,
        "offset_not_exactly_remappable_to_full_content": offset_not_remappable,
        "deduped_qualified_pool": len(deduped),
        "high_risk_priority_candidates": len(high_risk),
        "selected_count": len(selected),
        "selected_target_top10": dict(Counter(clean(r.get("target_name")) for r in selected).most_common(10)),
        "selected_source_top10": dict(Counter(clean(r.get("source")) for r in selected).most_common(10)),
        "selected_high_risk_family_counts": dict(Counter(f for r in selected for f in (r.get("high_risk_families") or ["unknown"]))),
    }
    return selected, meta


def formal_annotation_rows(selected: list[dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for row in selected:
        out = {
            "sample_id": clean(row.get("sample_id")),
            "target_name": clean(row.get("target_name")),
            "target_type": clean(row.get("target_type")),
            "event_name": clean(row.get("event_name")),
            "event_type": clean(row.get("event_type")),
            "title": clean(row.get("title")),
            "evidence_sentence": clean(row.get("evidence_sentence")),
            "local_context": clean(row.get("local_context")),
            "financial_facts": clean(row.get("financial_fact_if_available")),
            "publish_time": clean(row.get("publish_time")),
            "severity_label": None,
            "risk_intensity": None,
            "annotation_confidence": None,
            "annotation_basis": None,
            "label_tier": "task1_grounded_train_only_candidate_v1",
            "label_source": "pending_human_annotation",
            "target_local_status": "auto_grounding_pass_pending_human_review",
            "split_status": "train_only",
            "provenance": clean(row.get("provenance")),
            "provenance_status": "valid",
            "supervision_eligible": "pending_human_severity_label",
        }
        rows.append(out)
    return pd.DataFrame(rows, columns=FORMAL_COLUMNS)


def add_dropdown(ws: Any, column_name: str, values: list[str], end_row: int) -> None:
    headers = [clean(cell.value) for cell in ws[1]]
    if column_name not in headers:
        return
    col = headers.index(column_name) + 1
    letter = get_column_letter(col)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{end_row}")


def style_workbook(path: Path, annotation_sheet: str = "annotation") -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = WRAP
        headers = [clean(cell.value) for cell in ws[1]]
        for idx, header in enumerate(headers, 1):
            width = 18
            if header in {"title"}:
                width = 42
            elif header in {"evidence_sentence", "local_context", "financial_facts", "provenance"}:
                width = 68
            elif header in {"sample_id", "event_type"}:
                width = 32
            elif header in HUMAN_COLUMNS:
                width = 20
            ws.column_dimensions[get_column_letter(idx)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = TOP if ws.title != annotation_sheet else WRAP
                header = clean(ws.cell(1, cell.column).value)
                if header in HUMAN_COLUMNS:
                    cell.fill = HUMAN_FILL
        if ws.title == annotation_sheet:
            end_row = max(ws.max_row + 5, 2)
            add_dropdown(ws, "severity_label", ["red", "orange", "yellow", "blue"], end_row)
            add_dropdown(ws, "risk_intensity", ["1", "2", "3", "4", "5"], end_row)
            add_dropdown(ws, "annotation_confidence", ["high", "medium", "low"], end_row)
        if ws.title == "sampling_meta":
            for cell in ws[1]:
                cell.fill = META_FILL
            ws.sheet_state = "hidden"
    wb.save(path)


def write_outputs(expansion_rows: list[dict[str, Any]], selected: list[dict[str, Any]], report: dict[str, Any]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    expansion_df = pd.DataFrame(expansion_rows)
    expansion_df.to_excel(OUT_DIR / "task1_event_expansion_result.xlsx", index=False)
    write_jsonl(OUT_DIR / "task1_event_expansion_result.jsonl", expansion_rows)

    audit_cols = [
        "sample_id",
        "article_id",
        "target_name",
        "target_type",
        "target_code",
        "event_name",
        "event_type",
        "canonical_match",
        "canonical_event_id",
        "duplicate_status",
        "grounding_pass",
        "high_risk_score",
        "title",
        "evidence_sentence",
        "provenance",
    ]
    pd.DataFrame([{c: r.get(c) for c in audit_cols} for r in expansion_rows]).to_excel(OUT_DIR / "task3_event_match_audit.xlsx", index=False)

    candidate_cols = [
        "sample_id",
        "article_id",
        "target_name",
        "target_type",
        "event_id",
        "event_name",
        "event_type",
        "event_time",
        "evidence_sentence",
        "financial_fact_if_available",
        "canonical_event_id",
        "duplicate_status",
        "provenance",
        "target_grounded",
        "event_grounded",
        "evidence_grounded",
        "high_risk_score",
        "high_risk_families",
    ]
    candidate_df = pd.DataFrame([{c: r.get(c) for c in candidate_cols} for r in selected])
    candidate_df = candidate_df.rename(columns={"evidence_sentence": "event_evidence"})
    candidate_df.to_excel(OUT_DIR / "severity_candidate_input.xlsx", index=False)

    annotation_df = formal_annotation_rows(selected)
    annotation_path = OUT_DIR / "task4_severity_human_annotation_candidates_200_from_task1_v1.xlsx"
    instructions = pd.DataFrame(
        [
            {"项目": "标注目标", "说明": "这些样本均来自 targetable raw/shadow 重新进入 Task1 后的 target-local 事件候选，用于补充 Task4 red/orange 稀缺人工 Gold。"},
            {"项目": "只填写", "说明": "只填写 annotation 页浅黄色列：severity_label、risk_intensity、annotation_confidence、annotation_basis。"},
            {"项目": "severity_label", "说明": "四选一：red/orange/yellow/blue。不要因为候选被抽中就默认 red/orange。"},
            {"项目": "risk_intensity", "说明": "1-5 的有序风险强度评分。5 最严重，1 最轻。"},
            {"项目": "annotation_basis", "说明": "简要写出依据，必须基于 target + event + evidence/local_context，不使用全文泛化风险词。"},
            {"项目": "禁止", "说明": "不要修改 sample_id、target、event、evidence、provenance；不要打开 hidden meta 做标签依据。"},
        ]
    )
    meta_cols = [
        "sample_id",
        "article_id",
        "canonical_document_id",
        "source_dataset",
        "dataset_split",
        "target_code",
        "entity_confidence",
        "event_confidence",
        "event_top5",
        "event_candidate_count",
        "structure_trigger",
        "structure_event_span",
        "structure_subject",
        "high_risk_score",
        "high_risk_families",
        "high_risk_hits",
        "canonical_match",
        "canonical_event_id",
        "duplicate_status",
        "grounding_failure_reasons",
    ]
    meta_df = pd.DataFrame([{c: r.get(c) for c in meta_cols} for r in selected])
    with pd.ExcelWriter(annotation_path, engine="openpyxl") as writer:
        instructions.to_excel(writer, sheet_name="instructions", index=False)
        annotation_df.to_excel(writer, sheet_name="annotation", index=False)
        meta_df.to_excel(writer, sheet_name="sampling_meta", index=False)
    style_workbook(annotation_path)

    write_json(OUT_DIR / "quality_report.json", report)
    md = [
        "# Task1 Severity Event Expansion V1",
        "",
        f"- 处理文章数量：{report['task1_input_articles']}",
        f"- Task1抽取事件数量：{report['task1_extracted_event_count']}",
        f"- target grounding：{report['grounding']['target_grounded_count']} / {report['task1_input_articles']} = {report['grounding']['target_grounded_rate']}",
        f"- event grounding：{report['grounding']['event_grounded_count']} / {report['task1_input_articles']} = {report['grounding']['event_grounded_rate']}",
        f"- evidence grounding：{report['grounding']['evidence_grounded_count']} / {report['task1_input_articles']} = {report['grounding']['evidence_grounded_rate']}",
        f"- provenance通过率：{report['grounding']['provenance_valid_count']} / {report['task1_input_articles']} = {report['grounding']['provenance_valid_rate']}",
        f"- 与Task3已有CanonicalEvent匹配数量：{report['task3']['canonical_match_count']}",
        f"- 新增事件候选数量：{report['task3']['new_event_candidate_count']}",
        f"- 可进入Severity人工标注数量：{report['selection']['selected_count']}",
        f"- 失败主要原因前三名：{report['failure_top3']}",
        f"- 是否达到Severity mining条件：{report['severity_mining_gate']}",
        "",
        "## 输出",
        f"- `{(OUT_DIR / 'task1_event_expansion_result.xlsx').relative_to(ROOT)}`",
        f"- `{(OUT_DIR / 'task1_event_expansion_result.jsonl').relative_to(ROOT)}`",
        f"- `{(OUT_DIR / 'task3_event_match_audit.xlsx').relative_to(ROOT)}`",
        f"- `{(OUT_DIR / 'severity_candidate_input.xlsx').relative_to(ROOT)}`",
        f"- `{annotation_path.relative_to(ROOT)}`",
        f"- `{(OUT_DIR / 'quality_report.json').relative_to(ROOT)}`",
    ]
    REPORT_MD.parent.mkdir(parents=True, exist_ok=True)
    REPORT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")


def validate_outputs(path: Path, selected_count: int) -> dict[str, Any]:
    wb = load_workbook(path, data_only=False)
    ws = wb["annotation"]
    headers = [clean(c.value) for c in ws[1]]
    df = pd.read_excel(path, sheet_name="annotation", dtype=object)
    human_nonempty = {col: int(df[col].notna().sum()) for col in HUMAN_COLUMNS if col in df.columns}
    qa = {
        "annotation_rows": int(len(df)),
        "expected_rows": int(selected_count),
        "sample_id_unique": bool(df["sample_id"].astype(str).is_unique) if "sample_id" in df.columns else False,
        "human_fields_empty": all(v == 0 for v in human_nonempty.values()),
        "human_nonempty": human_nonempty,
        "freeze_ok": ws.freeze_panes == "A2",
        "filter_ok": bool(ws.auto_filter.ref),
        "dropdown_count": len(ws.data_validations.dataValidation),
        "sampling_meta_hidden": wb["sampling_meta"].sheet_state == "hidden",
        "visible_forbidden_model_columns": sorted(set(headers) & {"high_risk_score", "event_confidence", "event_top5", "candidate_score"}),
    }
    qa["all_pass"] = (
        qa["annotation_rows"] == selected_count
        and qa["sample_id_unique"]
        and qa["human_fields_empty"]
        and qa["freeze_ok"]
        and qa["filter_ok"]
        and qa["dropdown_count"] >= 3
        and qa["sampling_meta_hidden"]
        and not qa["visible_forbidden_model_columns"]
    )
    return qa


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--max-rows", type=int, default=200)
    parser.add_argument("--event-batch-size", type=int, default=64)
    parser.add_argument("--event-max-length", type=int, default=384)
    parser.add_argument("--event-candidate-topk", type=int, default=32)
    parser.add_argument("--structure-batch-size", type=int, default=8)
    parser.add_argument("--device", choices=["cpu", "cuda"], default="cpu")
    args = parser.parse_args()

    random.seed(SEED)
    np.random.seed(SEED)
    started = time.time()
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    protected_before = {
        "targetable_sha256": sha256_file(TARGETABLE_XLSX),
        "gold_sha256": sha256_file(GOLD_XLSX),
        "task3_graph_nodes_sha256": sha256_file(TASK3_RELEASE / "graph_nodes.jsonl"),
        "task3_graph_edges_sha256": sha256_file(TASK3_RELEASE / "graph_edges.jsonl"),
    }

    print(json.dumps({"step": "load_articles"}, ensure_ascii=False), flush=True)
    rows, article_meta = load_articles()
    print(json.dumps({"step": "build_training_rows"}, ensure_ascii=False), flush=True)
    all_rows, rows_by_split, _, _ = build_training_rows()
    train_rows = rows_by_split["train"]
    generator = load_candidate_generator(train_rows)

    print(json.dumps({"step": "target_prediction", "rows": len(rows)}, ensure_ascii=False), flush=True)
    with TARGET_MODEL.open("rb") as handle:
        target_model = pickle.load(handle)
    target_preds = predict_target_ranker(target_model, rows, generator)
    for row, pred in zip(rows, target_preds):
        row["target_type"] = clean(pred.get("target_type"))
        row["target_code"] = norm_code(pred.get("target_code"))
        row["target_name"] = clean(pred.get("target_name"))

    print(json.dumps({"step": "event_prediction", "rows": len(rows)}, ensure_ascii=False), flush=True)
    event_preds, event_meta = predict_events(
        rows,
        generator,
        batch_size=args.event_batch_size,
        max_length=args.event_max_length,
        event_candidate_topk=args.event_candidate_topk,
        device_arg=args.device,
    )
    for row, event in zip(rows, event_preds):
        row["event_name"] = clean(event.get("event_name"))
    for row in rows:
        row["structure_content"] = structure_local_content(row)

    print(json.dumps({"step": "structure_extraction", "rows": len(rows)}, ensure_ascii=False), flush=True)
    import torch
    from transformers import AutoTokenizer

    install_fast_structure_decoder()
    structure_device = torch.device("cuda" if args.device == "cuda" and torch.cuda.is_available() else "cpu")
    structure_model, structure_data = load_structure_checkpoint(STRUCTURE_CHECKPOINT, structure_device)
    structure_tokenizer = AutoTokenizer.from_pretrained(structure_data["model_dir"], local_files_only=True, use_fast=True)
    structure_input = [
        {
            "event_mention_id": stable_id("TMP_EM", row.get("article_id"), row.get("target_type"), row.get("target_code"), row.get("event_name")),
            "article_id": row.get("article_id"),
            "content": row.get("structure_content"),
            "event_name": row.get("event_name"),
            "title": row.get("title"),
        }
        for row in rows
    ]
    structure_preds = predict_structure_rows(structure_model, structure_tokenizer, structure_input, structure_data["scheme"], structure_device, args.structure_batch_size)

    print(json.dumps({"step": "grounding_and_task3"}, ensure_ascii=False), flush=True)
    expansion_rows = build_expansion_rows(rows, target_preds, event_preds, structure_preds)
    task3_lookup = load_task3_lookup()
    expansion_rows = attach_task3(expansion_rows, task3_lookup)
    selected, selection_meta = deduplicate_and_select(expansion_rows, args.max_rows)

    failure_reasons = Counter(reason for row in expansion_rows for reason in (row.get("grounding_failure_reasons") or []))
    grounding = {
        "target_grounded_count": int(sum(1 for r in expansion_rows if r.get("target_grounded"))),
        "event_grounded_count": int(sum(1 for r in expansion_rows if r.get("event_grounded"))),
        "evidence_grounded_count": int(sum(1 for r in expansion_rows if r.get("evidence_grounded"))),
        "provenance_valid_count": int(sum(1 for r in expansion_rows if r.get("provenance_valid"))),
        "grounding_pass_count": int(sum(1 for r in expansion_rows if r.get("grounding_pass"))),
    }
    for key in ["target_grounded", "event_grounded", "evidence_grounded", "provenance_valid"]:
        count = grounding[f"{key}_count"]
        grounding[f"{key}_rate"] = round(count / max(1, len(expansion_rows)), 6)
    grounding["grounding_pass_rate"] = round(grounding["grounding_pass_count"] / max(1, len(expansion_rows)), 6)

    protected_after = {
        "targetable_sha256": sha256_file(TARGETABLE_XLSX),
        "gold_sha256": sha256_file(GOLD_XLSX),
        "task3_graph_nodes_sha256": sha256_file(TASK3_RELEASE / "graph_nodes.jsonl"),
        "task3_graph_edges_sha256": sha256_file(TASK3_RELEASE / "graph_edges.jsonl"),
    }
    protected_unchanged = protected_before == protected_after
    gate = (
        "pass"
        if grounding["target_grounded_rate"] >= 0.90
        and grounding["event_grounded_rate"] >= 0.70
        and grounding["evidence_grounded_rate"] >= 0.70
        and grounding["provenance_valid_rate"] >= 0.99
        else "not_pass"
    )

    report = {
        "status": "task1_severity_event_expansion_v1_complete",
        "seed": SEED,
        "runtime_seconds": round(time.time() - started, 3),
        "task1_input_articles": len(rows),
        "article_meta": article_meta,
        "task1_components_reused": {
            "target_model": str(TARGET_MODEL.relative_to(ROOT)),
            "target_prediction_function": "scripts.task1_step5_8hj_common.predict_target_ranker",
            "event_name_model": str(EVENT_MODEL.relative_to(ROOT)),
            "event_prediction_function": "pairwise candidate scoring with scripts.task1_step5_8hj_common.build_event_pair_text",
            "candidate_generator": "src.routing.task1_candidate_v2.Task1CandidateGeneratorV2",
            "structure_checkpoint": str(STRUCTURE_CHECKPOINT.relative_to(ROOT)),
            "structure_prediction_function": "scripts.run_task1_step5_7k_macbert_worker.predict_rows",
            "official_schema": "data/processed/event_schema.jsonl",
        },
        "event_inference_meta": event_meta,
        "structure_inference_meta": {
            "device": str(structure_device),
            "scheme": structure_data.get("scheme"),
            "batch_size": args.structure_batch_size,
        },
        "task1_extracted_event_count": int(sum(1 for r in expansion_rows if clean(r.get("event_name")))),
        "task1_success_count": int(sum(1 for r in expansion_rows if clean(r.get("target_name")) and clean(r.get("event_name")) and clean(r.get("evidence_sentence")))),
        "task1_failure_count": int(sum(1 for r in expansion_rows if not (clean(r.get("target_name")) and clean(r.get("event_name")) and clean(r.get("evidence_sentence"))))),
        "grounding": grounding,
        "task3": {
            "canonical_match_count": int(sum(1 for r in expansion_rows if r.get("canonical_match"))),
            "new_event_candidate_count": int(sum(1 for r in expansion_rows if not r.get("canonical_match"))),
            "duplicate_status_counts": dict(Counter(clean(r.get("duplicate_status")) for r in expansion_rows)),
        },
        "selection": selection_meta,
        "failure_top3": failure_reasons.most_common(3),
        "severity_mining_gate": gate,
        "protected_inputs_unchanged": protected_unchanged,
        "protected_before": protected_before,
        "protected_after": protected_after,
        "outputs": {
            "task1_event_expansion_result_xlsx": str((OUT_DIR / "task1_event_expansion_result.xlsx").relative_to(ROOT)),
            "task1_event_expansion_result_jsonl": str((OUT_DIR / "task1_event_expansion_result.jsonl").relative_to(ROOT)),
            "task3_event_match_audit": str((OUT_DIR / "task3_event_match_audit.xlsx").relative_to(ROOT)),
            "severity_candidate_input": str((OUT_DIR / "severity_candidate_input.xlsx").relative_to(ROOT)),
            "human_annotation_package": str((OUT_DIR / "task4_severity_human_annotation_candidates_200_from_task1_v1.xlsx").relative_to(ROOT)),
            "quality_report": str((OUT_DIR / "quality_report.json").relative_to(ROOT)),
            "report_md": str(REPORT_MD.relative_to(ROOT)),
        },
    }

    print(json.dumps({"step": "write_outputs", "selected": len(selected)}, ensure_ascii=False), flush=True)
    write_outputs(expansion_rows, selected, report)
    qa = validate_outputs(OUT_DIR / "task4_severity_human_annotation_candidates_200_from_task1_v1.xlsx", len(selected))
    report["annotation_package_qa"] = qa
    write_json(OUT_DIR / "quality_report.json", report)
    if not qa["all_pass"]:
        print(json.dumps({"status": "qa_failed", "qa": qa}, ensure_ascii=False, indent=2), flush=True)
        return 1
    print(json.dumps({"status": "complete", "selected": len(selected), "gate": gate, "qa": qa}, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
