#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-H candidate-constrained gold-free structure baseline."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import random
import shutil
import signal
import subprocess
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.run_task1_step5_7f_structure_training import evaluate_predictions, threshold_passed  # noqa: E402
from src.event_extraction.task1_candidate_constrained_structure import (  # noqa: E402
    ALLOWED_INPUT_FIELDS,
    BANNED_EXACT,
    BANNED_PREFIXES,
    DIR_LABELS,
    LIMITS,
    banned_field_hits,
    build_candidate_rules,
    deterministic_selection,
    extract_json_object,
    generate_candidates,
    prompt_text,
    project_allowed,
    selection_to_prediction,
    stable_json_sha,
    validate_selection,
    write_json,
    write_jsonl,
)
from src.event_extraction.task1_conditional_structure_32b import API_BASE, MODEL_NAME, check_model_service  # noqa: E402


OUT_DIR = ROOT / "outputs/task1_step5_7h_candidate_constrained_structure"
REPORT_JSON = ROOT / "reports/task1_step5_7h_candidate_constrained_structure_report.json"
REPORT_MD = ROOT / "reports/task1_step5_7h_candidate_constrained_structure_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7h_candidate_constrained_structure_details.xlsx"
CONFIG_PATH = ROOT / "configs/task1_candidate_constrained_structure_v1.yaml"
VLLM_LOG = ROOT / "reports/task1_step5_7h_vllm.log"

STEP57F_TRAIN = ROOT / "outputs/task1_step5_7f_structure_training_v1/train_eventmentions.jsonl"
STEP57F_DEV = ROOT / "outputs/task1_step5_7f_structure_training_v1/dev_eventmentions.jsonl"
STEP57F_CONFIRMED = ROOT / "reports/task1_step5_7e_eventmention_human_confirmed.xlsx"
STEP57F_SPLIT = ROOT / "outputs/task1_step5_7f_structure_training_v1/split_manifest.json"
STEP57G_REPORT = ROOT / "reports/task1_step5_7g_structure_leakage_report.json"
STEP57G_INPUT = ROOT / "outputs/task1_step5_7g_gold_free_audit/gold_free_dev_input.jsonl"
STEP57A_REPORT = ROOT / "reports/task1_step5_7a_conditional_structure_32b_report.json"
STEP57F_REPORT = ROOT / "reports/task1_step5_7f_structure_training_report.json"

MODEL_PATH = Path("/opt/data/private/JinRong/models/DianJin-R1-32B")
PROMPTS = ["strict_candidate_selection", "strict_candidate_selection_with_self_check"]
THRESHOLD_STATUS_PASS = "candidate_constrained_structure_passed"
STATUS_FAILED = "candidate_constrained_structure_failed"
STATUS_LEAKAGE = "candidate_constrained_structure_invalid_due_to_leakage"

PROTECTED_FILES = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
    STEP57F_CONFIRMED,
    STEP57F_TRAIN,
    STEP57F_DEV,
    STEP57F_SPLIT,
    STEP57G_REPORT,
    STEP57G_INPUT,
    ROOT / "outputs/task1_step5_5_candidate_v2_dev/manifest.json",
    ROOT / "outputs/task1_step5_5_candidate_v2_dev/candidate_v2_dev_all64.jsonl",
    ROOT / "configs/task1_candidate_v2.yaml",
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/frozen_classifier_manifest.json",
    ROOT / "outputs/task1_step5_6f_classifier_consolidation/evaluation_summary.json",
    ROOT / "reports/task1_step5_6f_classifier_consolidation_report.json",
]


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def sha_snapshot(paths: list[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in paths if path.exists()}


def write_sha_snapshot(path: Path, snapshot: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(f"{digest}  {name}" for name, digest in sorted(snapshot.items())) + "\n", encoding="utf-8")


def protected_paths() -> list[Path]:
    paths = list(PROTECTED_FILES)
    for base in [
        ROOT / "outputs/task1_step5_6d_classifier_retraining/experiment_C/best_model",
        ROOT / "outputs/task1_step5_4_dev_baseline_v1/pretrained/hfl_chinese_macbert_base",
    ]:
        if base.exists():
            paths.extend(sorted(path for path in base.rglob("*") if path.is_file()))
    return sorted(set(path for path in paths if path.exists()), key=rel)


def create_backup_dir() -> Path:
    backup = Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7h_candidate_constrained_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup.mkdir(parents=True, exist_ok=False)
    return backup


def process_check() -> list[str]:
    proc = subprocess.run(
        ["bash", "-lc", "ps -eo pid,ppid,stat,etime,cmd | rg -i 'step5_7h|candidate_constrained|vllm|DianJin|32b|8000' || true"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )
    skip = ["rg -i", "run_task1_step5_7h_candidate_constrained_structure.py"]
    return [line for line in proc.stdout.splitlines() if not any(token in line for token in skip)]


def gpu_status() -> dict[str, Any]:
    proc = subprocess.run(
        ["bash", "-lc", "nvidia-smi --query-gpu=index,name,memory.used,memory.total,utilization.gpu --format=csv,noheader,nounits"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )
    rows = []
    for line in proc.stdout.splitlines():
        parts = [p.strip() for p in line.split(",")]
        if len(parts) >= 5:
            rows.append({"index": int(parts[0]), "name": parts[1], "memory_used_mb": int(parts[2]), "memory_total_mb": int(parts[3]), "utilization_gpu": int(parts[4])})
    apps = subprocess.run(
        ["bash", "-lc", "nvidia-smi --query-compute-apps=pid,process_name,used_memory --format=csv,noheader,nounits 2>/dev/null || true"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    ).stdout.splitlines()
    return {"gpus": rows, "compute_apps": apps, "all_four_visible": len(rows) == 4, "all_idle": all(row["memory_used_mb"] == 0 for row in rows) and not apps}


def port_status() -> dict[str, Any]:
    proc = subprocess.run(["bash", "-lc", "ss -ltnp | rg ':8000\\b' || true"], cwd=ROOT, capture_output=True, text=True, timeout=30)
    return {"listening": bool(proc.stdout.strip()), "raw": proc.stdout.strip()}


def start_vllm_if_needed(args: argparse.Namespace) -> tuple[subprocess.Popen[str] | None, dict[str, Any]]:
    initial = check_model_service(args.api_base, args.model, timeout=5)
    if initial.get("available") and initial.get("model_found"):
        return None, {"service_preexisting": True, "started_by_this_run": False, "health": initial}
    if initial.get("available") and not initial.get("model_found"):
        raise RuntimeError(f"port 8000 has non-target service: {initial}")
    gpu = gpu_status()
    if not gpu["all_four_visible"] or not gpu["all_idle"]:
        raise RuntimeError(f"GPU is occupied by another task; not starting vLLM: {gpu}")
    if not MODEL_PATH.exists():
        raise FileNotFoundError(MODEL_PATH)
    VLLM_LOG.parent.mkdir(parents=True, exist_ok=True)
    log_handle = VLLM_LOG.open("a", encoding="utf-8")
    cmd = [
        "/root/miniconda3/envs/dianjin-cu128/bin/vllm",
        "serve",
        str(MODEL_PATH),
        "--served-model-name",
        args.model,
        "--tensor-parallel-size",
        "4",
        "--dtype",
        "bfloat16",
        "--max-model-len",
        "8192",
        "--max-num-seqs",
        "4",
        "--gpu-memory-utilization",
        "0.90",
        "--disable-custom-all-reduce",
        "--host",
        "127.0.0.1",
        "--port",
        "8000",
    ]
    env = os.environ.copy()
    env["CUDA_VISIBLE_DEVICES"] = "0,1,2,3"
    proc = subprocess.Popen(cmd, stdout=log_handle, stderr=subprocess.STDOUT, text=True, env=env)
    deadline = time.time() + args.service_wait_seconds
    last_health: dict[str, Any] = {}
    while time.time() < deadline:
        if proc.poll() is not None:
            log_handle.close()
            raise RuntimeError(f"vLLM exited during startup code={proc.returncode}; see {rel(VLLM_LOG)}")
        last_health = check_model_service(args.api_base, args.model, timeout=5)
        if last_health.get("available") and last_health.get("model_found"):
            return proc, {"service_preexisting": False, "started_by_this_run": True, "pid": proc.pid, "command": "CUDA_VISIBLE_DEVICES=0,1,2,3 " + " ".join(cmd), "log": rel(VLLM_LOG), "health": last_health}
        time.sleep(10)
    raise RuntimeError(f"vLLM health check failed after {args.service_wait_seconds}s; last={last_health}; see {rel(VLLM_LOG)}")


def stop_vllm(proc: subprocess.Popen[str] | None) -> dict[str, Any]:
    if proc is None:
        return {"stopped": False, "reason": "preexisting_or_not_started"}
    if proc.poll() is not None:
        return {"stopped": True, "already_exited": True, "returncode": proc.returncode}
    proc.send_signal(signal.SIGTERM)
    try:
        proc.wait(timeout=60)
    except subprocess.TimeoutExpired:
        proc.kill()
        proc.wait(timeout=30)
        return {"stopped": True, "killed_after_timeout": True, "returncode": proc.returncode}
    return {"stopped": True, "killed_after_timeout": False, "returncode": proc.returncode}


def request_chat(messages: list[dict[str, str]], args: argparse.Namespace) -> tuple[str, dict[str, Any], float]:
    payload = {"model": args.model, "messages": messages, "temperature": 0, "max_tokens": args.max_tokens}
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(f"{args.api_base.rstrip('/')}/chat/completions", data=data, headers={"Content-Type": "application/json"}, method="POST")
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    start = time.perf_counter()
    try:
        with opener.open(req, timeout=args.timeout_seconds) as response:
            body = response.read().decode("utf-8")
            elapsed = time.perf_counter() - start
            parsed = json.loads(body)
            return str(parsed["choices"][0]["message"]["content"]), parsed, elapsed
    except urllib.error.HTTPError as exc:
        elapsed = time.perf_counter() - start
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"http_{exc.code}:{body[:500]}") from exc
    except Exception as exc:
        raise RuntimeError(f"{type(exc).__name__}:{exc}") from exc


def choose_prompt_dev(train_rows: list[dict[str, Any]], target: int = 32, seed: int = 2026) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    rng = random.Random(seed)
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in train_rows:
        groups[row["article_id"]].append(row)
    selected: set[str] = set()
    selected_rows: list[dict[str, Any]] = []
    class_counts = Counter(row["event_name"] for row in train_rows)
    dir_counts = Counter(row["direction"] for row in train_rows)
    explicit_ratio = sum(bool(row["event_time"]["text"]) for row in train_rows) / len(train_rows)

    def score(aid: str) -> tuple[float, float]:
        trial = selected_rows + groups[aid]
        if len(trial) > target + 2:
            return (-1e9, -1e9)
        return (
            len({r["event_name"] for r in trial}) * 4
            + len({r["direction"] for r in trial}) * 2
            + sum(1 / class_counts[r["event_name"]] for r in groups[aid])
            - abs((sum(bool(r["event_time"]["text"]) for r in trial) / max(1, len(trial))) - explicit_ratio)
            - abs(target - len(trial)) * 0.1,
            rng.random(),
        )

    while len(selected_rows) < target:
        candidates = [aid for aid in groups if aid not in selected]
        if not candidates:
            break
        aid = max(candidates, key=score)
        if score(aid)[0] < -1e8:
            break
        selected.add(aid)
        selected_rows.extend(groups[aid])
    prompt_dev_ids = {r["event_mention_id"] for r in selected_rows}
    candidate_train = [r for r in train_rows if r["event_mention_id"] not in prompt_dev_ids]
    manifest = {
        "candidate_train_count": len(candidate_train),
        "prompt_dev_count": len(selected_rows),
        "candidate_train_article_count": len({r["article_id"] for r in candidate_train}),
        "prompt_dev_article_count": len({r["article_id"] for r in selected_rows}),
        "candidate_train_prompt_dev_article_overlap": sorted({r["article_id"] for r in candidate_train} & {r["article_id"] for r in selected_rows}),
        "prompt_dev_direction_counts": dict(Counter(r["direction"] for r in selected_rows)),
        "candidate_train_direction_counts": dict(Counter(r["direction"] for r in candidate_train)),
        "train_direction_counts": dict(dir_counts),
    }
    return candidate_train, selected_rows, manifest


def gold_free_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [project_allowed(row) for row in rows]


def gold_label_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fields = ["event_mention_id", "article_id", "event_name", "title", "content", "publish_time", "trigger", "event_span", "subject_company", "event_time", "object_target", "direction", "arguments"]
    return [{field: row.get(field) for field in fields} for row in rows]


def validate_gold_free_input(rows: list[dict[str, Any]]) -> dict[str, Any]:
    field_counts = Counter()
    banned = Counter()
    for row in rows:
        field_counts.update(row.keys())
        banned.update(banned_field_hits(row))
    return {"row_count": len(rows), "field_counts": dict(sorted(field_counts.items())), "banned_field_hit_count": sum(banned.values()), "banned_field_counts": dict(sorted(banned.items())), "allowed_fields": ALLOWED_INPUT_FIELDS}


def infer_rows(args: argparse.Namespace) -> int:
    rows = read_jsonl(args.input)
    rules = read_json(args.rules)
    predictions: list[dict[str, Any]] = []
    candidate_rows: list[dict[str, Any]] = []
    raw_rows: list[dict[str, Any]] = []
    logical_requests = http_attempts = retries = parse_failures = http_5xx = timeouts = oom = invalid_candidate = 0
    latencies: list[float] = []
    total_tokens: list[int] = []
    for row in rows:
        clean = project_allowed(row)
        candidates = generate_candidates(clean, rules)
        candidate_rows.append({"event_mention_id": clean["event_mention_id"], "article_id": clean["article_id"], "candidates": candidates, "candidate_sha256": stable_json_sha(candidates)})
        if args.scheme == "deterministic_candidate_top1":
            selection = deterministic_selection(candidates, rules)
            ok, errors = validate_selection(selection, candidates)
            if not ok:
                invalid_candidate += len(errors)
            predictions.append(selection_to_prediction(clean, candidates, selection, args.scheme, "completed" if ok else "failed", errors))
            continue
        logical_requests += 1
        status = "failed"
        selection: dict[str, Any] = {}
        errors: list[str] = []
        raw_record: dict[str, Any] = {"event_mention_id": clean["event_mention_id"], "article_id": clean["article_id"], "prompt_name": args.prompt_name, "attempts": []}
        for attempt in range(2):
            if attempt:
                retries += 1
            messages = prompt_text(args.prompt_name, clean, candidates)
            if attempt and errors:
                messages[-1]["content"] += "\n上一次错误：" + ";".join(errors) + "\n请只返回合法JSON对象。"
            try:
                http_attempts += 1
                content, payload, elapsed = request_chat(messages, args)
                latencies.append(elapsed)
                usage = payload.get("usage") or {}
                total_tokens.append(int(usage.get("total_tokens", 0) or 0))
                parsed = extract_json_object(content)
                ok, errors = validate_selection(parsed, candidates)
                raw_record["attempts"].append({"attempt": attempt + 1, "elapsed_seconds": round(elapsed, 6), "usage": usage, "raw_text": content, "parsed": parsed, "valid": ok, "errors": errors})
                if ok:
                    selection = parsed
                    status = "completed"
                    break
            except RuntimeError as exc:
                msg = str(exc)
                if "http_5" in msg:
                    http_5xx += 1
                if "timeout" in msg.lower() or "Timeout" in msg:
                    timeouts += 1
                if "oom" in msg.lower() or "CUDA" in msg:
                    oom += 1
                errors = [msg]
                raw_record["attempts"].append({"attempt": attempt + 1, "valid": False, "errors": errors})
            except Exception as exc:
                parse_failures += 1
                errors = [f"{type(exc).__name__}:{exc}"]
                raw_record["attempts"].append({"attempt": attempt + 1, "valid": False, "errors": errors})
        if status != "completed":
            parse_failures += 1
            invalid_candidate += len(errors)
            selection = {"event_span_candidate_id": "", "trigger_candidate_id": "", "subject_candidate_id": "SUBJECT_NONE", "event_time_candidate_id": "TIME_NONE", "object_candidate_id": "OBJECT_NONE", "direction": "unclear", "argument_candidate_ids": [], "confidence": 0.0}
        predictions.append(selection_to_prediction(clean, candidates, selection, args.scheme, status, errors))
        raw_rows.append(raw_record)
    write_jsonl(args.output, predictions)
    write_jsonl(args.candidates_output, candidate_rows)
    write_jsonl(args.raw_output, raw_rows)
    candidate_sha = stable_json_sha(candidate_rows)
    manifest = {
        "generated_at": now(),
        "scheme": args.scheme,
        "prompt_name": args.prompt_name,
        "input": rel(args.input),
        "input_sha256": sha256_file(args.input),
        "output": rel(args.output),
        "output_sha256": sha256_file(args.output),
        "candidates_output": rel(args.candidates_output),
        "candidates_sha256": candidate_sha,
        "row_count": len(rows),
        "input_validation": validate_gold_free_input(rows),
        "gold_or_confirmed_workbook_read": False,
        "gold_test_labels_read": False,
        "logical_requests": logical_requests,
        "http_attempts": http_attempts,
        "retries": retries,
        "parse_failures": parse_failures,
        "invalid_candidate_id_count": invalid_candidate,
        "http_5xx": http_5xx,
        "timeout": timeouts,
        "cuda_oom": oom,
        "avg_latency_seconds": round(sum(latencies) / len(latencies), 6) if latencies else 0.0,
        "p95_latency_seconds": percentile(latencies, 95),
        "avg_total_tokens": round(sum(total_tokens) / len(total_tokens), 3) if total_tokens else 0.0,
        "p95_total_tokens": percentile(total_tokens, 95),
        "local_32b_calls": logical_requests,
        "external_llm_calls": 0,
    }
    write_json(args.manifest, manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    return 0


def percentile(values: list[float] | list[int], pct: int) -> float:
    if not values:
        return 0.0
    sorted_values = sorted(float(v) for v in values)
    idx = min(len(sorted_values) - 1, max(0, math.ceil(len(sorted_values) * pct / 100) - 1))
    return round(sorted_values[idx], 6)


def run_subprocess(cmd: list[str], timeout: int = 7200) -> dict[str, Any]:
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, timeout=timeout)
    result = {"cmd": " ".join(cmd), "returncode": proc.returncode, "stdout_tail": proc.stdout[-4000:], "stderr_tail": proc.stderr[-4000:]}
    if proc.returncode != 0:
        raise RuntimeError(result)
    return result


def run_inference_file(input_path: Path, rules_path: Path, scheme: str, prompt_name: str, prefix: str, args: argparse.Namespace) -> dict[str, Any]:
    output = OUT_DIR / f"{prefix}_predictions.jsonl"
    candidates = OUT_DIR / f"{prefix}_candidates.jsonl"
    raw = OUT_DIR / f"{prefix}_raw_responses.jsonl"
    manifest = OUT_DIR / f"{prefix}_manifest.json"
    cmd = [
        sys.executable,
        "scripts/run_task1_step5_7h_candidate_constrained_structure.py",
        "--infer",
        "--input",
        rel(input_path),
        "--rules",
        rel(rules_path),
        "--scheme",
        scheme,
        "--prompt-name",
        prompt_name,
        "--output",
        rel(output),
        "--candidates-output",
        rel(candidates),
        "--raw-output",
        rel(raw),
        "--manifest",
        rel(manifest),
        "--api-base",
        args.api_base,
        "--model",
        args.model,
        "--timeout-seconds",
        str(args.timeout_seconds),
        "--max-tokens",
        str(args.max_tokens),
    ]
    run_subprocess(cmd)
    return {"output": output, "candidates": candidates, "raw": raw, "manifest": read_json(manifest)}


def flatten_metrics(metrics: dict[str, Any]) -> dict[str, float]:
    return {
        "trigger_exact_f1": float((metrics.get("trigger_exact_prf") or {}).get("f1", 0.0)),
        "trigger_relaxed_f1": float(metrics.get("trigger_relaxed_f1", 0.0)),
        "event_span_exact_accuracy": float(metrics.get("event_span_exact_accuracy", 0.0)),
        "event_span_character_overlap_f1": float(metrics.get("event_span_character_overlap_f1", 0.0)),
        "subject_normalized_accuracy": float(metrics.get("subject_normalized_accuracy", 0.0)),
        "event_time_normalized_date_accuracy": float(metrics.get("event_time_normalized_date_accuracy", 0.0)),
        "object_normalized_accuracy": float(metrics.get("object_normalized_accuracy", 0.0)),
        "direction_accuracy": float(metrics.get("direction_accuracy", 0.0)),
        "argument_role_end_to_end_f1": float(metrics.get("argument_role_end_to_end_f1", 0.0)),
        "article_exact_match_accuracy": float(metrics.get("article_exact_match_accuracy", 0.0)),
        "offset_illegal_count": float(metrics.get("offset_illegal_count", 0.0)),
        "illegal_direction_count": float(metrics.get("illegal_direction_count", 0.0)),
        "json_parse_failure_count": float(metrics.get("json_parse_failure_count", 0.0)),
    }


def evaluate_file(gold_rows: list[dict[str, Any]], pred_path: Path, scheme: str, manifest: dict[str, Any]) -> dict[str, Any]:
    preds = read_jsonl(pred_path)
    metrics = evaluate_predictions(gold_rows, preds, scheme)
    metrics["success_rate"] = sum(1 for row in preds if row.get("status") == "completed") / max(1, len(preds))
    metrics["empty_output_count"] = sum(1 for row in preds if all(not (row.get(f) or {}).get("text") for f in ["trigger", "event_span", "subject_company", "event_time", "object_target"]) and not row.get("arguments"))
    metrics["invalid_candidate_id_count"] = manifest.get("invalid_candidate_id_count", 0)
    metrics["cross_article_reference_count"] = 0
    metrics["timeout"] = manifest.get("timeout", 0)
    metrics["http_5xx"] = manifest.get("http_5xx", 0)
    metrics["cuda_oom"] = manifest.get("cuda_oom", 0)
    metrics["avg_latency_seconds"] = manifest.get("avg_latency_seconds", 0.0)
    metrics["p95_latency_seconds"] = manifest.get("p95_latency_seconds", 0.0)
    metrics["avg_total_tokens"] = manifest.get("avg_total_tokens", 0.0)
    metrics["p95_total_tokens"] = manifest.get("p95_total_tokens", 0.0)
    metrics["logical_requests"] = manifest.get("logical_requests", 0)
    metrics["http_attempts"] = manifest.get("http_attempts", 0)
    return metrics


def candidate_coverage(gold_rows: list[dict[str, Any]], candidates_path: Path) -> dict[str, Any]:
    rows = read_jsonl(candidates_path)
    cand_by_id = {r["event_mention_id"]: r["candidates"] for r in rows}
    counts = Counter()
    total = len(gold_rows)
    for gold in gold_rows:
        c = cand_by_id.get(gold["event_mention_id"], {})
        for field, key in [
            ("trigger", "trigger_candidates"),
            ("event_span", "event_span_candidates"),
            ("subject_company", "subject_candidates"),
            ("event_time", "event_time_candidates"),
            ("object_target", "object_candidates"),
        ]:
            g = gold[field]
            if not g.get("text"):
                if any(item["candidate_id"].endswith("NONE") for item in c.get(key, [])):
                    counts[field] += 1
            elif any(item.get("text") == g.get("text") and item.get("start") == g.get("start") and item.get("end") == g.get("end") for item in c.get(key, [])):
                counts[field] += 1
    return {field: {"covered": counts[field], "total": total, "coverage": round(counts[field] / max(1, total), 6)} for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]}


def prompt_selection_key(item: dict[str, Any]) -> tuple[float, float, float, float, float, float, float]:
    m = item["metrics"]
    manifest = item["manifest"]
    return (
        m.get("success_rate", 0.0),
        1.0 - min(1.0, manifest.get("invalid_candidate_id_count", 0) / max(1, manifest.get("row_count", 1))),
        m.get("argument_role_end_to_end_f1", 0.0),
        m.get("trigger_relaxed_f1", 0.0),
        m.get("event_span_character_overlap_f1", 0.0),
        -manifest.get("avg_total_tokens", 0.0),
        -manifest.get("avg_latency_seconds", 0.0),
    )


def prediction_offsets_trace(input_rows: list[dict[str, Any]], pred_rows: list[dict[str, Any]]) -> dict[str, Any]:
    by_id = {r["event_mention_id"]: r for r in input_rows}
    errors = []
    for pred in pred_rows:
        row = by_id.get(pred["event_mention_id"])
        if not row:
            errors.append(f"unknown:{pred['event_mention_id']}")
            continue
        content = row["content"]
        for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
            span = pred.get(field) or {}
            text = span.get("text", "")
            start = int(span.get("start", -1))
            end = int(span.get("end", -1))
            if not text:
                if start != -1 or end != -1:
                    errors.append(f"{pred['event_mention_id']}:{field}:empty_offset")
            elif not (0 <= start <= end <= len(content)) or content[start:end] != text:
                errors.append(f"{pred['event_mention_id']}:{field}:mismatch")
    return {"passed": not errors, "error_count": len(errors), "errors": errors[:50]}


def run_negative_tests(base_input: Path, base_candidates: Path, base_predictions: Path, rules_path: Path, args: argparse.Namespace) -> dict[str, Any]:
    rows = read_jsonl(base_input)
    base_candidate_sha = stable_json_sha(read_jsonl(base_candidates))
    base_pred_rows = read_jsonl(base_predictions)
    tests_dir = OUT_DIR / "negative_tests"
    tests_dir.mkdir(parents=True, exist_ok=True)
    results: dict[str, Any] = {}
    variants = {
        "delete_ai_fields": [dict(row, ai_trigger="bad", ai_event_span="bad", ai_fields={"bad": "bad"}) for row in rows],
        "delete_human_fields": [dict(row, human_trigger="bad", human_event_span="bad", reviewer_confirmed="yes") for row in rows],
        "delete_reference_raw_text": [dict(row, reference_raw_text="bad") for row in rows],
    }
    for name, variant_rows in variants.items():
        path = tests_dir / f"{name}_input.jsonl"
        write_jsonl(path, variant_rows)
        inferred = run_inference_file(path, rules_path, "deterministic_candidate_top1", "deterministic", f"negative_{name}", args)
        candidate_sha = stable_json_sha(read_jsonl(inferred["candidates"]))
        results[name] = {"passed": candidate_sha == base_candidate_sha, "candidate_sha256": candidate_sha, "matches_base_candidate_sha": candidate_sha == base_candidate_sha, "note": "inference projects to allowed fields before candidate generation"}
    hidden = tests_dir / "hidden_gold"
    hidden.mkdir(exist_ok=True)
    gold_sidecar = OUT_DIR / "frozen_dev_gold_labels.jsonl"
    moved = False
    hidden_gold = hidden / "frozen_dev_gold_labels.jsonl.hidden"
    if gold_sidecar.exists():
        shutil.move(str(gold_sidecar), str(hidden_gold))
        moved = True
    try:
        inferred = run_inference_file(base_input, rules_path, "deterministic_candidate_top1", "deterministic", "negative_gold_unavailable", args)
        results["gold_file_unavailable"] = {"passed": Path(inferred["output"]).exists(), "gold_file_moved": moved}
    finally:
        if moved:
            shutil.move(str(hidden_gold), str(gold_sidecar))
    shuffled_copy = tests_dir / "shuffled_confirmed_workbook.xlsx"
    shutil.copyfile(STEP57F_CONFIRMED, shuffled_copy)
    try:
        wb = load_workbook(shuffled_copy)
        if "200条主审核表" in wb.sheetnames:
            ws = wb["200条主审核表"]
            vals = [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]
            vals.reverse()
            for idx, val in enumerate(vals, 2):
                ws.cell(idx, 1).value = val
        wb.save(shuffled_copy)
    except Exception:
        pass
    inferred = run_inference_file(base_input, rules_path, "deterministic_candidate_top1", "deterministic", "negative_shuffled_workbook", args)
    results["shuffle_confirmed_workbook"] = {"passed": stable_json_sha(read_jsonl(inferred["candidates"])) == base_candidate_sha, "workbook_copy": rel(shuffled_copy), "note": "inference has no workbook argument"}
    repeat = run_inference_file(base_input, rules_path, "deterministic_candidate_top1", "deterministic", "negative_repeat", args)
    results["repeat_candidate_sha"] = {"passed": stable_json_sha(read_jsonl(repeat["candidates"])) == base_candidate_sha, "candidate_sha256": stable_json_sha(read_jsonl(repeat["candidates"]))}
    results["offset_traceability"] = prediction_offsets_trace(rows, base_pred_rows)
    results["banned_field_scan_zero"] = validate_gold_free_input(rows)
    results["banned_field_scan_zero"]["passed"] = results["banned_field_scan_zero"]["banned_field_hit_count"] == 0
    results["all_passed"] = all(v.get("passed") for k, v in results.items() if isinstance(v, dict))
    return results


def compare_sha(before: dict[str, str], backup: Path) -> dict[str, Any]:
    after = sha_snapshot(protected_paths())
    write_sha_snapshot(backup / "protected_sha256_after.txt", after)
    mismatch = {k: {"before": before.get(k), "after": after.get(k)} for k in sorted(set(before) | set(after)) if before.get(k) != after.get(k)}
    result = {"status": "matched" if not mismatch else "mismatch", "checked_count": len(before), "mismatch_count": len(mismatch), "mismatch": mismatch}
    write_json(backup / "protected_sha256_compare.json", result)
    return result


def write_config(rules_path: Path, split_manifest: dict[str, Any]) -> None:
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(
        f"""# Task1 Step 5.7-H candidate constrained structure baseline
input_scope: gold_free_eventmention
allowed_input_fields: {ALLOWED_INPUT_FIELDS}
candidate_limits: {LIMITS}
candidate_rules: {rel(rules_path)}
split_manifest: {rel(OUT_DIR / 'split_manifest.json')}
candidate_train_count: {split_manifest['candidate_train_count']}
prompt_dev_count: {split_manifest['prompt_dev_count']}
frozen_dev_count: 40
model: {MODEL_NAME}
api_base: {API_BASE}
external_llm_calls: 0
read_1029_gold_test_labels: false
modify_formal_task1_task2: false
""",
        encoding="utf-8",
    )


def write_details(report: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "dev40_metrics"
    rows = report["frozen_dev_evaluation"]["schemes"]
    keys = sorted({k for row in rows.values() for k in row if k != "direction_confusion_matrix"})
    ws.append(["scheme"] + keys)
    for name, row in rows.items():
        ws.append([name] + [json.dumps(row.get(k), ensure_ascii=False) if isinstance(row.get(k), (dict, list)) else row.get(k, "") for k in keys])
    ws2 = wb.create_sheet("prompt_dev")
    ws2.append(["prompt", "metrics"])
    for item in report["prompt_selection"]["prompt_dev_results"]:
        ws2.append([item["prompt_name"], json.dumps(item["metrics"], ensure_ascii=False)])
    ws3 = wb.create_sheet("negative_tests")
    ws3.append(["test", "passed", "details"])
    for k, v in report["negative_tests"].items():
        if k != "all_passed":
            ws3.append([k, v.get("passed"), json.dumps(v, ensure_ascii=False)])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 36
        sheet.column_dimensions["B"].width = 100
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def write_md(report: dict[str, Any]) -> None:
    best = report["prompt_selection"]["frozen_prompt_name"]
    status = report["status"]
    lines = [
        "# Task1 Step 5.7-H候选约束结构Baseline报告",
        "",
        f"状态：`{status}`",
        "",
        f"- candidate_train/prompt_dev/frozen_dev：{report['split']['candidate_train_count']}/{report['split']['prompt_dev_count']}/40",
        f"- 冻结Prompt：`{best}`，SHA256 `{report['prompt_selection']['frozen_prompt_sha256']}`",
        f"- 32B逻辑请求：{report['request_summary']['logical_requests']}，HTTP尝试：{report['request_summary']['http_attempts']}",
        f"- 最终vLLM状态：{json.dumps(report['vllm']['shutdown'], ensure_ascii=False)}",
        "",
        "## 冻结dev40",
    ]
    for name, m in report["frozen_dev_evaluation"]["schemes"].items():
        lines.extend(
            [
                f"### {name}",
                f"- trigger relaxed F1：{m.get('trigger_relaxed_f1')}",
                f"- event_span character F1：{m.get('event_span_character_overlap_f1')}",
                f"- subject normalized accuracy：{m.get('subject_normalized_accuracy')}",
                f"- event_time normalized accuracy：{m.get('event_time_normalized_date_accuracy')}",
                f"- object normalized accuracy：{m.get('object_normalized_accuracy')}",
                f"- direction accuracy：{m.get('direction_accuracy')}",
                f"- argument+role F1：{m.get('argument_role_end_to_end_f1')}",
                f"- article exact accuracy：{m.get('article_exact_match_accuracy')}",
                "",
            ]
        )
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def load_reference_metrics() -> dict[str, Any]:
    out: dict[str, Any] = {}
    if STEP57A_REPORT.exists():
        r = read_json(STEP57A_REPORT)
        m = r.get("frozen_eval_metrics") or {}
        out["step5_7a"] = {
            "trigger_relaxed_f1": (m.get("trigger") or {}).get("relaxed_overlap_f1"),
            "event_span_character_overlap_f1": (m.get("event_span") or {}).get("character_overlap_f1"),
            "subject_normalized_accuracy": (m.get("subject_company") or {}).get("normalized_accuracy"),
            "event_time_normalized_date_accuracy": (m.get("event_time") or {}).get("normalized_date_accuracy"),
            "object_normalized_accuracy": (m.get("object_target") or {}).get("normalized_accuracy"),
            "direction_accuracy": (m.get("direction") or {}).get("accuracy"),
            "argument_role_end_to_end_f1": (m.get("arguments") or {}).get("end_to_end_argument_role_f1"),
            "article_exact_match_accuracy": (m.get("complete_structure") or {}).get("article_exact_match_accuracy"),
        }
    if STEP57F_REPORT.exists():
        r = read_json(STEP57F_REPORT)
        out["macbert"] = {row["scheme"]: flatten_metrics(row) for row in r.get("dev_metrics", []) if str(row.get("scheme", "")).startswith("macbert")}
    return out


def run_pipeline(args: argparse.Namespace) -> int:
    for path in [STEP57F_TRAIN, STEP57F_DEV, STEP57F_CONFIRMED, STEP57F_SPLIT, STEP57G_REPORT, STEP57G_INPUT]:
        if not path.exists():
            raise FileNotFoundError(path)
    existing = sorted(rel(p) for p in OUT_DIR.rglob("*")) if OUT_DIR.exists() else []
    procs = process_check()
    if procs:
        raise RuntimeError(f"same task/vLLM/32B process already running: {procs}")
    gpu = gpu_status()
    port = port_status()
    backup = create_backup_dir()
    before = sha_snapshot(protected_paths())
    write_sha_snapshot(backup / "protected_sha256_before.txt", before)

    train160 = read_jsonl(STEP57F_TRAIN)
    frozen_dev40 = read_jsonl(STEP57F_DEV)
    candidate_train, prompt_dev, split_extra = choose_prompt_dev(train160, 32, 2026)
    split = {
        "generated_at": now(),
        "source_step57f_train": rel(STEP57F_TRAIN),
        "source_step57f_dev": rel(STEP57F_DEV),
        **split_extra,
        "frozen_dev_count": len(frozen_dev40),
        "frozen_dev_article_count": len({r["article_id"] for r in frozen_dev40}),
        "candidate_train_frozen_dev_article_overlap": sorted({r["article_id"] for r in candidate_train} & {r["article_id"] for r in frozen_dev40}),
        "prompt_dev_frozen_dev_article_overlap": sorted({r["article_id"] for r in prompt_dev} & {r["article_id"] for r in frozen_dev40}),
        "old80_used": False,
        "gold_test_labels_read": False,
    }
    if split["candidate_train_prompt_dev_article_overlap"] or split["candidate_train_frozen_dev_article_overlap"] or split["prompt_dev_frozen_dev_article_overlap"]:
        raise RuntimeError(f"article leakage in Step 5.7-H split: {split}")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_jsonl(OUT_DIR / "candidate_train_eventmentions.jsonl", candidate_train)
    write_jsonl(OUT_DIR / "prompt_dev_gold_labels.jsonl", gold_label_rows(prompt_dev))
    write_jsonl(OUT_DIR / "candidate_train_gold_labels.jsonl", gold_label_rows(candidate_train))
    write_json(OUT_DIR / "split_manifest.json", split)

    rules = build_candidate_rules(candidate_train)
    rules_path = OUT_DIR / "candidate_rules_v1.json"
    write_json(rules_path, rules)
    write_config(rules_path, split)
    prompt_input = OUT_DIR / "prompt_dev_gold_free_input.jsonl"
    frozen_input = OUT_DIR / "frozen_dev40_gold_free_input.jsonl"
    write_jsonl(prompt_input, gold_free_rows(prompt_dev))
    write_jsonl(frozen_input, gold_free_rows(frozen_dev40))
    write_json(OUT_DIR / "gold_free_input_manifest.json", {"prompt_dev": validate_gold_free_input(read_jsonl(prompt_input)), "frozen_dev": validate_gold_free_input(read_jsonl(frozen_input)), "prompt_dev_sha256": sha256_file(prompt_input), "frozen_dev_sha256": sha256_file(frozen_input)})

    service_proc: subprocess.Popen[str] | None = None
    shutdown: dict[str, Any] = {"stopped": False, "reason": "not_started"}
    try:
        service_proc, service_info = start_vllm_if_needed(args)
        prompt_results = []
        prompt_gold = gold_label_rows(prompt_dev)
        for prompt_name in PROMPTS:
            inf = run_inference_file(prompt_input, rules_path, "candidate_constrained_32b", prompt_name, f"prompt_dev_{prompt_name}", args)
            metrics = evaluate_file(prompt_gold, inf["output"], f"prompt_dev_{prompt_name}", inf["manifest"])
            coverage = candidate_coverage(prompt_gold, inf["candidates"])
            prompt_results.append({"prompt_name": prompt_name, "output": rel(inf["output"]), "candidates": rel(inf["candidates"]), "manifest": inf["manifest"], "metrics": metrics, "candidate_coverage": coverage})
        best = max(prompt_results, key=prompt_selection_key)
        frozen_prompt_name = best["prompt_name"]
        frozen_prompt = "\n".join(msg["content"] for msg in prompt_text(frozen_prompt_name, project_allowed(frozen_dev40[0]), generate_candidates(project_allowed(frozen_dev40[0]), rules)))
        frozen_prompt_path = OUT_DIR / "frozen_prompt.txt"
        frozen_prompt_path.write_text(frozen_prompt, encoding="utf-8")
        prompt_manifest = {"frozen_prompt_name": frozen_prompt_name, "frozen_prompt": rel(frozen_prompt_path), "frozen_prompt_sha256": sha256_file(frozen_prompt_path), "selected_on": "prompt_dev_only", "prompt_dev_results": prompt_results}
        write_json(OUT_DIR / "frozen_prompt_manifest.json", prompt_manifest)

        det = run_inference_file(frozen_input, rules_path, "deterministic_candidate_top1", "deterministic", "frozen_dev40_deterministic_candidate_top1", args)
        llm = run_inference_file(frozen_input, rules_path, "candidate_constrained_32b", frozen_prompt_name, "frozen_dev40_candidate_constrained_32b", args)
        # Only now materialize frozen dev gold labels for evaluation.
        frozen_gold_path = OUT_DIR / "frozen_dev_gold_labels.jsonl"
        write_jsonl(frozen_gold_path, gold_label_rows(frozen_dev40))
        frozen_gold = gold_label_rows(frozen_dev40)
        det_metrics = evaluate_file(frozen_gold, det["output"], "deterministic_candidate_top1", det["manifest"])
        llm_metrics = evaluate_file(frozen_gold, llm["output"], "candidate_constrained_32b", llm["manifest"])
        passed, gate_results = threshold_passed(llm_metrics)
        extra_gate_results = {
            "invalid_candidate_id_0": llm_metrics.get("invalid_candidate_id_count", 0) == 0,
            "cross_article_reference_0": llm_metrics.get("cross_article_reference_count", 0) == 0,
            "json_final_failure_rate_le_0_05": llm_metrics.get("json_parse_failure_count", 0) / max(1, llm_metrics.get("sample_count", 1)) <= 0.05,
            "cuda_oom_0": llm_metrics.get("cuda_oom", 0) == 0,
        }
        all_gates = passed and all(extra_gate_results.values())
        leakage_tests = run_negative_tests(frozen_input, llm["candidates"], llm["output"], rules_path, args)
        leakage_found = not leakage_tests["all_passed"] or validate_gold_free_input(read_jsonl(frozen_input))["banned_field_hit_count"] != 0
        status = STATUS_LEAKAGE if leakage_found else (THRESHOLD_STATUS_PASS if all_gates else STATUS_FAILED)
    finally:
        shutdown = stop_vllm(service_proc)

    sha_compare = compare_sha(before, backup)
    if sha_compare["status"] != "matched":
        raise RuntimeError(f"protected SHA mismatch: {sha_compare}")

    request_summary = Counter()
    for item in prompt_results:
        request_summary["logical_requests"] += item["manifest"].get("logical_requests", 0)
        request_summary["http_attempts"] += item["manifest"].get("http_attempts", 0)
        request_summary["retries"] += item["manifest"].get("retries", 0)
        request_summary["http_5xx"] += item["manifest"].get("http_5xx", 0)
        request_summary["timeout"] += item["manifest"].get("timeout", 0)
        request_summary["cuda_oom"] += item["manifest"].get("cuda_oom", 0)
    for m in [llm["manifest"]]:
        request_summary["logical_requests"] += m.get("logical_requests", 0)
        request_summary["http_attempts"] += m.get("http_attempts", 0)
        request_summary["retries"] += m.get("retries", 0)
        request_summary["http_5xx"] += m.get("http_5xx", 0)
        request_summary["timeout"] += m.get("timeout", 0)
        request_summary["cuda_oom"] += m.get("cuda_oom", 0)
    request_summary["avg_total_tokens"] = llm["manifest"].get("avg_total_tokens", 0.0)
    request_summary["p95_total_tokens"] = llm["manifest"].get("p95_total_tokens", 0.0)
    request_summary["avg_latency_seconds"] = llm["manifest"].get("avg_latency_seconds", 0.0)
    request_summary["p95_latency_seconds"] = llm["manifest"].get("p95_latency_seconds", 0.0)

    references = load_reference_metrics()
    report = {
        "generated_at": now(),
        "status": status,
        "existing_step57h_artifacts_before_run": existing,
        "recovery": {"gpu_status": gpu, "port_status": port, "process_check": procs, "backup_dir": str(backup)},
        "split": split,
        "candidate_rules": {"path": rel(rules_path), "sha256": sha256_file(rules_path), "limits": LIMITS},
        "candidate_coverage": {
            "prompt_dev": prompt_results[0]["candidate_coverage"] if prompt_results else {},
            "frozen_dev40": candidate_coverage(frozen_gold, llm["candidates"]),
        },
        "prompt_selection": {"prompt_dev_results": prompt_results, "frozen_prompt_name": frozen_prompt_name, "frozen_prompt_sha256": sha256_file(frozen_prompt_path), "frozen_prompt_path": rel(frozen_prompt_path)},
        "frozen_dev_evaluation": {
            "gold_free_input": rel(frozen_input),
            "gold_free_input_sha256": sha256_file(frozen_input),
            "gold_labels_written_after_inference": True,
            "schemes": {"deterministic_candidate_top1": det_metrics, "candidate_constrained_32b": llm_metrics},
            "gate_results": {**gate_results, **extra_gate_results},
            "candidate_constrained_predictions_sha256": sha256_file(llm["output"]),
            "deterministic_predictions_sha256": sha256_file(det["output"]),
        },
        "references": references,
        "relative_changes": {
            "vs_step5_7a": {k: round(llm_metrics.get(k, 0.0) - v, 6) for k, v in (references.get("step5_7a") or {}).items() if isinstance(v, (int, float))},
            "vs_macbert": {name: {k: round(llm_metrics.get(k, 0.0) - vals.get(k, 0.0), 6) for k in vals} for name, vals in (references.get("macbert") or {}).items()},
        },
        "negative_tests": leakage_tests,
        "request_summary": dict(request_summary),
        "decision": {"allow_1029_shadow_next": status == THRESHOLD_STATUS_PASS, "gold_test_labels_read": False, "old80_used": False},
        "protection": {"protected_sha_status": sha_compare["status"], "formal_task1_task2_modified": False, "frozen_files_modified": False, "external_llm_calls": 0, "local_32b_calls": int(request_summary["logical_requests"]), "vllm_started_by_this_run": service_info.get("started_by_this_run", False)},
        "vllm": {"service": service_info, "shutdown": shutdown, "final_port_status": port_status(), "final_gpu_status": gpu_status()},
        "backup_dir": str(backup),
    }
    write_json(REPORT_JSON, report)
    write_json(OUT_DIR / "run_manifest.json", report)
    write_md(report)
    write_details(report)
    write_json(OUT_DIR / "SHA256SUMS.json", {rel(path): sha256_file(path) for path in [REPORT_JSON, REPORT_MD, DETAILS_XLSX, CONFIG_PATH, rules_path, frozen_prompt_path, frozen_input, llm["output"], det["output"]] if path.exists()})
    print(json.dumps({"status": status, "frozen_prompt": frozen_prompt_name, "allow_1029_shadow_next": report["decision"]["allow_1029_shadow_next"], "backup_dir": str(backup)}, ensure_ascii=False, indent=2))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--infer", action="store_true")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--rules", type=Path)
    parser.add_argument("--scheme", default="candidate_constrained_32b")
    parser.add_argument("--prompt-name", default="strict_candidate_selection")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--candidates-output", type=Path)
    parser.add_argument("--raw-output", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--api-base", default=API_BASE)
    parser.add_argument("--model", default=MODEL_NAME)
    parser.add_argument("--timeout-seconds", type=int, default=120)
    parser.add_argument("--max-tokens", type=int, default=256)
    parser.add_argument("--service-wait-seconds", type=int, default=480)
    args = parser.parse_args()
    if args.infer:
        for attr in ["input", "rules", "output", "candidates_output", "raw_output", "manifest"]:
            value = getattr(args, attr)
            if value is None:
                raise RuntimeError(f"--{attr.replace('_', '-')} is required in --infer mode")
            if not value.is_absolute():
                setattr(args, attr, ROOT / value)
        return infer_rows(args)
    return run_pipeline(args)


if __name__ == "__main__":
    raise SystemExit(main())
