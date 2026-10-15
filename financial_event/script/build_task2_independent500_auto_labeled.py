#!/usr/bin/env python3
"""Build an auto-labeled silver version of the Task2 independent500 audit package."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.build_task2_independent500_human_audit import (  # noqa: E402
    ALLOWED_VALUES,
    HUMAN_FIELDS,
    STRESS_TERMS,
    count_terms,
    row_text_for_signals,
    target_match_ratio,
    time_mismatch_ratio,
)


DEFAULT_SOURCE_WORKBOOK = ROOT / "reports/task2_independent500_human_audit.xlsx"
DEFAULT_OUTPUT_WORKBOOK = ROOT / "reports/task2_independent500_auto_labeled.xlsx"
DEFAULT_SOURCE_MANIFEST = ROOT / "reports/task2_independent500_manifest.json"
DEFAULT_OUTPUT_MANIFEST = ROOT / "reports/task2_independent500_auto_labeled_manifest.json"
DEFAULT_SOURCE_SAMPLES = ROOT / "outputs/event_credibility/task2_independent500_v1/independent_samples.jsonl"
DEFAULT_SOURCE_SCORES = ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet"
DEFAULT_LABELS_JSONL = ROOT / "reports/task2_independent500_auto_labeled_labels.jsonl"
DEFAULT_REPORT_JSON = ROOT / "reports/task2_independent500_auto_labeled_report.json"
DEFAULT_REPORT_MD = ROOT / "reports/task2_independent500_auto_labeled_report.md"
DEFAULT_GUIDE = ROOT / "reports/task2_independent500_auto_labeled_guide.md"


def norm(value: Any) -> str:
    if value is None:
        return ""
    try:
        if isinstance(value, float) and math.isnan(value):
            return ""
    except Exception:
        pass
    return re.sub(r"\s+", " ", str(value)).strip()


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def rel(path: Path) -> str:
    return str(path.relative_to(ROOT))


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, allow_nan=False) + "\n")


def safe_int(value: Any, default: int = 0) -> int:
    try:
        if value is None:
            return default
        if isinstance(value, float) and math.isnan(value):
            return default
        return int(value)
    except Exception:
        return default


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None:
            return default
        if isinstance(value, float) and math.isnan(value):
            return default
        out = float(value)
        if math.isfinite(out):
            return out
    except Exception:
        pass
    return default


def text_contains_any(text: str, terms: tuple[str, ...]) -> bool:
    return any(term and term in text for term in terms)


def score_lookup(df: pd.DataFrame) -> dict[str, dict[str, Any]]:
    return {norm(row["cluster_id"]): row for row in df.to_dict("records")}


def sample_lookup(samples: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {norm(row["cluster_id"]): row for row in samples}


def effective_event_term(sample: dict[str, Any]) -> str:
    event_name = norm(sample.get("event_name"))
    if event_name:
        return event_name
    schema = norm(sample.get("schema_id"))
    if "_" in schema:
        return schema.rsplit("_", 1)[-1]
    return schema


def select_evidence_ids(sample: dict[str, Any], score_row: dict[str, Any], verdict: str, max_ids: int = 3) -> str:
    items = [item for item in sample.get("evidence_items", []) if isinstance(item, dict)]
    if not items:
        return ""

    target = norm(sample.get("target_name"))
    target_code = norm(sample.get("target_code"))
    event_term = effective_event_term(sample)
    schema_id = norm(sample.get("schema_id"))
    contradiction_terms = STRESS_TERMS["contradiction_terms"]

    ranked: list[tuple[float, dict[str, Any]]] = []
    for item in items:
        text = norm(f"{item.get('title', '')} {item.get('evidence_text', '')}")
        score = safe_float(item.get("sentence_score"))
        if verdict == "support":
            if target and target in text:
                score += 3.0
            if event_term and event_term in text:
                score += 2.0
            if schema_id and schema_id in text:
                score += 1.0
            if target_code and target_code in text:
                score += 1.0
            if norm(item.get("source_type")) == "official":
                score += 1.0
        elif verdict == "contradict":
            if text_contains_any(text, contradiction_terms):
                score += 3.5
            if target and target in text:
                score += 1.5
            if event_term and event_term in text:
                score += 1.0
            if norm(item.get("source_type")) == "official":
                score += 0.5
        else:
            if target and target in text:
                score += 1.5
            if event_term and event_term in text:
                score += 1.0
        ranked.append((score, item))

    ranked.sort(key=lambda pair: pair[0], reverse=True)
    seen: set[str] = set()
    chosen: list[str] = []
    for _score, item in ranked:
        eid = norm(item.get("evidence_id"))
        if not eid or eid in seen:
            continue
        seen.add(eid)
        chosen.append(eid)
        if len(chosen) >= max_ids:
            break
    return ";".join(chosen)


def auto_verdict(sample: dict[str, Any], score_row: dict[str, Any]) -> str:
    verdict = norm(score_row.get("fulltext_evidence_status"))
    if verdict in {"support", "contradict", "insufficient", "unclear"}:
        return verdict
    support_count = safe_int(score_row.get("fulltext_support_count"))
    contradict_count = safe_int(score_row.get("fulltext_contradict_count"))
    if contradict_count > 0:
        return "contradict"
    if support_count > 0:
        return "support"
    if not norm(sample.get("event_claim")):
        return "unclear"
    return "insufficient"


def time_match_label(sample: dict[str, Any]) -> str:
    ratio = time_mismatch_ratio(sample)
    if ratio <= 0.2:
        return "yes"
    if ratio >= 0.4:
        return "no"
    return "unclear"


def source_independence_label(score_row: dict[str, Any]) -> str:
    return "yes" if safe_int(score_row.get("fulltext_distinct_source_count")) >= 2 else "no"


def official_source_label(score_row: dict[str, Any]) -> str:
    return "yes" if safe_int(score_row.get("fulltext_official_support_count")) > 0 else "no"


def investor_question_label(sample: dict[str, Any], score_row: dict[str, Any]) -> str:
    text = row_text_for_signals(sample)
    if not text_contains_any(text, STRESS_TERMS["investor_question_terms"]):
        return "no"
    if safe_int(score_row.get("fulltext_official_support_count")) > 0:
        return "no"
    if safe_int(score_row.get("fulltext_support_count")) > 1:
        return "no"
    if safe_int(score_row.get("fulltext_distinct_source_count")) > 1:
        return "no"
    if safe_int(score_row.get("fulltext_contradict_count")) > 0:
        return "no"
    return "yes"


def market_signal_label(sample: dict[str, Any], score_row: dict[str, Any]) -> str:
    text = row_text_for_signals(sample)
    if not text_contains_any(text, STRESS_TERMS["market_signal_terms"]):
        return "no"
    if safe_int(score_row.get("fulltext_official_support_count")) > 0:
        return "no"
    if safe_int(score_row.get("fulltext_support_count")) > 1:
        return "no"
    if safe_int(score_row.get("fulltext_distinct_source_count")) > 1:
        return "no"
    if safe_int(score_row.get("fulltext_contradict_count")) > 0:
        return "no"
    return "yes"


def contradiction_present_label(score_row: dict[str, Any]) -> str:
    return "yes" if safe_int(score_row.get("fulltext_contradict_count")) > 0 or norm(score_row.get("fulltext_evidence_status")) == "contradict" else "no"


def confidence_label(verdict: str, score_row: dict[str, Any]) -> str:
    support_count = safe_int(score_row.get("fulltext_support_count"))
    contradict_count = safe_int(score_row.get("fulltext_contradict_count"))
    distinct_sources = safe_int(score_row.get("fulltext_distinct_source_count"))
    official_support_count = safe_int(score_row.get("fulltext_official_support_count"))
    credibility_score = safe_float(score_row.get("fulltext_credibility_score"))
    support_score = safe_float(score_row.get("fulltext_support_score"))
    contradiction_score = safe_float(score_row.get("fulltext_contradiction_score"))
    credibility_level = norm(score_row.get("fulltext_credibility_level"))

    if verdict == "support":
        if (
            support_count >= 4
            and distinct_sources >= 2
            and (official_support_count > 0 or support_score >= 0.6 or credibility_score >= 80 or credibility_level == "高可信")
        ):
            return "high"
        if support_count >= 2 or distinct_sources >= 2 or support_score >= 0.25 or credibility_score >= 30:
            return "medium"
        return "low"
    if verdict == "contradict":
        if contradict_count >= 2 or contradiction_score >= 0.45 or credibility_level == "疑似虚假或已反转" or credibility_score <= 20:
            return "high"
        if contradict_count >= 1 or contradiction_score >= 0.2:
            return "medium"
        return "low"
    if verdict == "unclear":
        return "low"
    return "low"


def primary_error_type(verdict: str) -> str:
    if verdict == "contradict":
        return "contradiction"
    if verdict == "insufficient":
        return "evidence_insufficient"
    if verdict == "unclear":
        return "claim_unclear"
    return "none"


def reviewer_notes(
    sample: dict[str, Any],
    score_row: dict[str, Any],
    verdict: str,
    time_match: str,
    source_independence: str,
    official_source_present: str,
    investor_question_only: str,
    market_signal_only: str,
    contradiction_present: str,
) -> str:
    parts = [
        "auto_silver_v1",
        f"status={verdict}",
        f"verif={norm(score_row.get('fulltext_verification_status')) or 'n/a'}",
        f"support={safe_int(score_row.get('fulltext_support_count'))}",
        f"contra={safe_int(score_row.get('fulltext_contradict_count'))}",
        f"official={safe_int(score_row.get('fulltext_official_support_count'))}",
        f"src={safe_int(score_row.get('fulltext_distinct_source_count'))}",
        f"cred={safe_float(score_row.get('fulltext_credibility_score')):.2f}",
        f"tmatch={time_match}",
        f"src_ind={source_independence}",
        f"official_src={official_source_present}",
        f"inv_q={investor_question_only}",
        f"market={market_signal_only}",
        f"contra_flag={contradiction_present}",
        f"target_ratio={target_match_ratio(sample):.2f}",
    ]
    llm_verdict = norm(score_row.get("llm_verdict"))
    if llm_verdict:
        parts.append(f"llm={llm_verdict}")
    return "; ".join(parts)[:300]


def build_labels(samples: list[dict[str, Any]], scores: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for sample in samples:
        cid = norm(sample.get("cluster_id"))
        score_row = scores.get(cid, {})
        verdict = auto_verdict(sample, score_row)
        time_match = time_match_label(sample)
        source_independence = source_independence_label(score_row)
        official_source_present = official_source_label(score_row)
        investor_question_only = investor_question_label(sample, score_row)
        market_signal_only = market_signal_label(sample, score_row)
        contradiction_present = contradiction_present_label(score_row)
        evidence_sufficiency = "yes" if verdict in {"support", "contradict"} else "no"
        evidence_relevance = evidence_sufficiency
        label = {
            "cluster_id": cid,
            "sample_group": norm(sample.get("sample_group")),
            "sample_index": safe_int(sample.get("sample_index")),
            "human_verdict": verdict,
            "human_confidence": confidence_label(verdict, score_row),
            "human_evidence_ids": select_evidence_ids(sample, score_row, verdict),
            "claim_clarity": "yes",
            "subject_match": "yes" if target_match_ratio(sample) >= 0.5 else "unclear",
            "time_match": time_match,
            "evidence_sufficiency": evidence_sufficiency,
            "evidence_relevance": evidence_relevance,
            "source_independence": source_independence,
            "official_source_present": official_source_present,
            "investor_question_only": investor_question_only,
            "market_signal_only": market_signal_only,
            "contradiction_present": contradiction_present,
            "primary_error_type": primary_error_type(verdict),
            "reviewer_notes": reviewer_notes(
                sample,
                score_row,
                verdict,
                time_match,
                source_independence,
                official_source_present,
                investor_question_only,
                market_signal_only,
                contradiction_present,
            ),
            "reviewer_confirmed": "no",
            "fulltext_evidence_status": norm(score_row.get("fulltext_evidence_status")),
            "fulltext_verification_status": norm(score_row.get("fulltext_verification_status")),
            "fulltext_support_count": safe_int(score_row.get("fulltext_support_count")),
            "fulltext_contradict_count": safe_int(score_row.get("fulltext_contradict_count")),
            "fulltext_official_support_count": safe_int(score_row.get("fulltext_official_support_count")),
            "fulltext_distinct_source_count": safe_int(score_row.get("fulltext_distinct_source_count")),
            "fulltext_support_score": round(safe_float(score_row.get("fulltext_support_score")), 4),
            "fulltext_contradiction_score": round(safe_float(score_row.get("fulltext_contradiction_score")), 4),
            "fulltext_credibility_score": round(safe_float(score_row.get("fulltext_credibility_score")), 2),
            "fulltext_credibility_level": norm(score_row.get("fulltext_credibility_level")),
            "llm_verdict": norm(score_row.get("llm_verdict")),
            "llm_confidence": round(safe_float(score_row.get("llm_confidence")), 4),
        }
        out.append(label)
    return out


def fill_instruction_sheet(ws, manifest_path: Path) -> None:
    headers = [norm(cell.value) for cell in ws[1]]
    try:
        item_col = headers.index("item") + 1
        value_col = headers.index("value") + 1
    except ValueError:
        return
    updates = {
        "任务": "Task2 Step 6.1 独立500条事件簇证据可信判断自动银标",
        "判断对象": "事件簇/事件主张与证据之间的可信判断，不是单篇新闻真假。",
        "Primary": "400条；用于自然分布自动银标参考。",
        "Stress": "100条；困难/高风险场景，只能单独报告，不并入自然总体准确率。",
        "禁止": "不要把本包当作人工金标；reviewer_confirmed 已由脚本回填为 no。",
        "人工字段": "浅橙色列已自动填写；结果仅供 silver 参考，不是人工金标。",
        "manifest": rel(manifest_path),
        "status": "task2_independent500_auto_silver_labeled",
    }
    for row in range(2, ws.max_row + 1):
        key = norm(ws.cell(row, item_col).value)
        if key in updates:
            ws.cell(row, value_col).value = updates[key]


def fill_workbook(
    source_workbook: Path,
    output_workbook: Path,
    samples_by_cluster: dict[str, dict[str, Any]],
    labels_by_cluster: dict[str, dict[str, Any]],
    manifest_path: Path,
) -> None:
    wb = load_workbook(source_workbook)
    fill_instruction_sheet(wb["填写说明"], manifest_path)
    for sheet_name in ("Primary人工审核", "Stress人工审核"):
        ws = wb[sheet_name]
        headers = [norm(cell.value) for cell in ws[1]]
        col_idx = {header: idx + 1 for idx, header in enumerate(headers)}
        cluster_col = col_idx.get("cluster_id")
        if not cluster_col:
            continue
        for row in range(2, ws.max_row + 1):
            cid = norm(ws.cell(row, cluster_col).value)
            if not cid or cid not in labels_by_cluster:
                continue
            sample = samples_by_cluster.get(cid, {})
            label = labels_by_cluster[cid]
            if "event_name" in col_idx and not norm(ws.cell(row, col_idx["event_name"]).value):
                ws.cell(row, col_idx["event_name"]).value = effective_event_term(sample)
            if "event_time" in col_idx and not norm(ws.cell(row, col_idx["event_time"]).value):
                fallback_time = norm(sample.get("event_time")) or norm(sample.get("earliest_publish_time")) or norm(sample.get("latest_publish_time")) or "未明确"
                ws.cell(row, col_idx["event_time"]).value = fallback_time
            for field in HUMAN_FIELDS:
                if field in col_idx:
                    ws.cell(row, col_idx[field]).value = label.get(field, "")
    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_workbook)


def render_report(report: dict[str, Any]) -> str:
    lines = [
        "# Task2 Independent500 Auto-Labeled Silver Report",
        "",
        f"- Status: `{report['status']}`",
        f"- Output workbook: `{report['output_workbook']}`",
        f"- Labels JSONL: `{report['labels_jsonl']}`",
        f"- Source workbook: `{report['source_workbook']}`",
        "",
        "## Verdict Counts",
        "",
        "| verdict | count |",
        "|---|---:|",
    ]
    for verdict, count in report["verdict_counts"].items():
        lines.append(f"| {verdict} | {count} |")
    lines.extend(
        [
            "",
            "## Auxiliary Counts",
            "",
            f"- time_match: {report['time_match_counts']}",
            f"- source_independence: {report['source_independence_counts']}",
            f"- official_source_present: {report['official_source_counts']}",
            f"- investor_question_only: {report['investor_question_counts']}",
            f"- market_signal_only: {report['market_signal_counts']}",
            f"- contradiction_present: {report['contradiction_counts']}",
            "",
            "## Notes",
            "",
            "- `human_verdict` follows the deterministic fulltext evidence status first.",
            "- `reviewer_confirmed` is fixed to `no`; this is auto-silver, not human gold.",
            "- The original 47 gold labels remain unchanged.",
        ]
    )
    return "\n".join(lines) + "\n"


def build_report(
    labels: list[dict[str, Any]],
    samples: list[dict[str, Any]],
    source_workbook: Path,
    output_workbook: Path,
    labels_jsonl: Path,
    manifest_path: Path,
) -> dict[str, Any]:
    verdict_counts = Counter(label["human_verdict"] for label in labels)
    confidence_counts = Counter(label["human_confidence"] for label in labels)
    time_match_counts = Counter(label["time_match"] for label in labels)
    source_independence_counts = Counter(label["source_independence"] for label in labels)
    official_source_counts = Counter(label["official_source_present"] for label in labels)
    investor_question_counts = Counter(label["investor_question_only"] for label in labels)
    market_signal_counts = Counter(label["market_signal_only"] for label in labels)
    contradiction_counts = Counter(label["contradiction_present"] for label in labels)
    group_counts = Counter(label["sample_group"] for label in labels)
    score_status_counts = Counter(label["fulltext_evidence_status"] or "unknown" for label in labels)
    error_counts = Counter(label["primary_error_type"] for label in labels)
    source_status_counts = Counter(
        "blank" if not norm(sample.get("event_name")) else "named" for sample in samples
    )
    report = {
        "status": "task2_independent500_auto_silver_labeled",
        "generated_at": now_iso(),
        "source_workbook": rel(source_workbook),
        "output_workbook": rel(output_workbook),
        "labels_jsonl": rel(labels_jsonl),
        "manifest": rel(manifest_path),
        "counts": {
            "Primary": group_counts.get("Primary", 0),
            "Stress": group_counts.get("Stress", 0),
            "total": len(labels),
        },
        "verdict_counts": dict(verdict_counts),
        "confidence_counts": dict(confidence_counts),
        "time_match_counts": dict(time_match_counts),
        "source_independence_counts": dict(source_independence_counts),
        "official_source_counts": dict(official_source_counts),
        "investor_question_counts": dict(investor_question_counts),
        "market_signal_counts": dict(market_signal_counts),
        "contradiction_counts": dict(contradiction_counts),
        "primary_error_type_counts": dict(error_counts),
        "source_event_name_counts": dict(source_status_counts),
        "score_status_counts": dict(score_status_counts),
        "rule_summary": {
            "verdict_source": "fulltext_evidence_status from event_credibility_scores_fulltext.parquet",
            "confidence_source": "score counts, distinct sources, official support, and credibility score",
            "time_match_rule": "time_mismatch_ratio <= 0.2 => yes; >= 0.4 => no; otherwise unclear",
            "source_independence_rule": "distinct source count >= 2 => yes else no",
            "official_source_rule": "official support count > 0 => yes else no",
            "investor_market_rules": "keyword hit plus weak evidence and no official support => yes",
        },
    }
    return report


def build_guide(report: dict[str, Any]) -> str:
    counts = report["counts"]
    return (
        "# Task2 Independent500 Auto-Labeled Silver Guide\n\n"
        "本包是 Task2 Step 6.1 独立500样本的全自动银标回填，不是人工金标。\n\n"
        f"- Primary: {counts['Primary']} 条\n"
        f"- Stress: {counts['Stress']} 条\n"
        f"- `reviewer_confirmed` 已固定为 `no`\n"
        f"- `human_verdict` 由 fulltext evidence status 自动映射\n"
        f"- `time_match`、`source_independence`、`official_source_present` 等为规则信号\n\n"
        "原始 47 条人工金标保持不变；本文件只提供 500 条自动扩展结果。\n"
    )


def build_package(args: argparse.Namespace) -> dict[str, Any]:
    source_workbook = args.source_workbook
    output_workbook = args.output_workbook
    source_manifest = args.source_manifest
    output_manifest = args.output_manifest
    labels_jsonl = args.labels_jsonl
    report_json = args.report_json
    report_md = args.report_md
    guide = args.guide

    if not source_workbook.exists():
        raise FileNotFoundError(source_workbook)
    if not source_manifest.exists():
        raise FileNotFoundError(source_manifest)
    if not args.overwrite:
        for path in (output_workbook, output_manifest, labels_jsonl, report_json, report_md, guide):
            if path.exists():
                raise FileExistsError(f"refusing to overwrite existing file: {path}")

    samples = read_jsonl(args.samples)
    scores_df = pd.read_parquet(args.scores)
    scores = score_lookup(scores_df)
    samples_by_cluster = sample_lookup(samples)
    labels = build_labels(samples, scores)
    labels_by_cluster = {label["cluster_id"]: label for label in labels}

    source_manifest_data = read_json(source_manifest)
    fill_workbook(source_workbook, output_workbook, samples_by_cluster, labels_by_cluster, output_manifest)

    report = build_report(labels, samples, source_workbook, output_workbook, labels_jsonl, output_manifest)
    write_jsonl(labels_jsonl, labels)
    write_json(report_json, report)
    report_md.write_text(render_report(report), encoding="utf-8")
    guide.write_text(build_guide(report), encoding="utf-8")

    output_manifest_data = dict(source_manifest_data)
    output_manifest_data.update(
        {
            "status": "task2_independent500_auto_silver_labeled",
            "generated_at": now_iso(),
            "step": "Task2 Step 6.1 independent500_v1 auto_silver",
            "project_root": str(ROOT),
            "labeling_mode": "auto_silver",
            "auto_labeling": {
                "generated": True,
                "model_calls": 0,
                "reason": "deterministic_rule_based_fulltext_auto_labeling",
                "verdict_source": "fulltext_evidence_status",
            },
            "human_fields_blank_on_generation": False,
            "outputs": {
                **source_manifest_data["outputs"],
                "workbook": rel(output_workbook),
                "guide": rel(guide),
                "manifest": rel(output_manifest),
                "labels_jsonl": rel(labels_jsonl),
                "report_json": rel(report_json),
                "report_md": rel(report_md),
            },
        }
    )
    output_manifest_data["outputs"]["workbook_sha256"] = sha256(output_workbook)
    output_manifest_data["outputs"]["guide_sha256"] = sha256(guide)
    output_manifest_data["outputs"]["labels_jsonl_sha256"] = sha256(labels_jsonl)
    output_manifest_data["outputs"]["report_json_sha256"] = sha256(report_json)
    output_manifest_data["outputs"]["report_md_sha256"] = sha256(report_md)
    write_json(output_manifest, output_manifest_data)

    return {
        "status": output_manifest_data["status"],
        "counts": report["counts"],
        "verdict_counts": report["verdict_counts"],
        "confidence_counts": report["confidence_counts"],
        "time_match_counts": report["time_match_counts"],
        "source_independence_counts": report["source_independence_counts"],
        "official_source_counts": report["official_source_counts"],
        "investor_question_counts": report["investor_question_counts"],
        "market_signal_counts": report["market_signal_counts"],
        "contradiction_counts": report["contradiction_counts"],
        "primary_error_type_counts": report["primary_error_type_counts"],
        "output_workbook": rel(output_workbook),
        "manifest": rel(output_manifest),
        "labels_jsonl": rel(labels_jsonl),
        "report_json": rel(report_json),
        "report_md": rel(report_md),
        "guide": rel(guide),
    }


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-workbook", type=Path, default=DEFAULT_SOURCE_WORKBOOK)
    parser.add_argument("--output-workbook", type=Path, default=DEFAULT_OUTPUT_WORKBOOK)
    parser.add_argument("--source-manifest", type=Path, default=DEFAULT_SOURCE_MANIFEST)
    parser.add_argument("--output-manifest", type=Path, default=DEFAULT_OUTPUT_MANIFEST)
    parser.add_argument("--samples", type=Path, default=DEFAULT_SOURCE_SAMPLES)
    parser.add_argument("--scores", type=Path, default=DEFAULT_SOURCE_SCORES)
    parser.add_argument("--labels-jsonl", type=Path, default=DEFAULT_LABELS_JSONL)
    parser.add_argument("--report-json", type=Path, default=DEFAULT_REPORT_JSON)
    parser.add_argument("--report-md", type=Path, default=DEFAULT_REPORT_MD)
    parser.add_argument("--guide", type=Path, default=DEFAULT_GUIDE)
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    for attr in ("source_workbook", "output_workbook", "source_manifest", "output_manifest", "samples", "scores", "labels_jsonl", "report_json", "report_md", "guide"):
        value = getattr(args, attr)
        if not value.is_absolute():
            setattr(args, attr, ROOT / value)
    summary = build_package(args)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
