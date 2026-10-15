#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the offline Task1 Step 5.7-E human-confirmation package."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.event_extraction.task1_eventmention_contract import (  # noqa: E402
    ARGUMENT_ROLES,
    DIRECTION_LABELS,
    HUMAN_REVIEW_FIELDS,
    SPAN_FIELDS,
    empty_span,
    empty_time,
    validate_eventmention_record,
)


STEP57C_DIR = ROOT / "outputs/task1_step5_7c_eventmention_annotation_pilot"
STEP57D_DIR = ROOT / "outputs/task1_step5_7d_contract_revision"
CONFIRMED80 = ROOT / "outputs/task1_structure_validation_round1/minimal_human_review_package_confirmed.jsonl"
MIGRATED80 = STEP57C_DIR / "confirmed80_eventmention_migration.jsonl"
TRAIN_LABELS = ROOT / "data/processed/labeled/news_train.jsonl"
EVENT_SCHEMA = ROOT / "data/processed/event_schema.jsonl"
TRAIN_PAIRS = ROOT / "outputs/task1_step5_6d_classifier_retraining/train_gold_pairs.jsonl"

OUT_DIR = ROOT / "outputs/task1_step5_7e_human_confirmation"
REPORT_XLSX = ROOT / "reports/task1_step5_7e_eventmention_human_confirmation.xlsx"
REPORT_GUIDE = ROOT / "reports/task1_step5_7e_eventmention_human_confirmation_guide.md"
REPORT_MANIFEST = ROOT / "reports/task1_step5_7e_eventmention_human_confirmation_manifest.json"

EXPECTED_OLD_PENDING = {
    "NEWS_MULTISOURCE_D-1120797906500",
    "NEWS_MULTISOURCE_D-1126130257189",
    "NEWS_MULTISOURCE_D-1127695210485",
}

BASE_FIELDS = [
    "event_mention_id",
    "article_id",
    "event_name",
    "title",
    "content",
    "publish_time",
    "suggestion_source",
    "original_quality_grade",
    "needs_manual_correction",
    "validation_errors",
]
AI_FIELDS = [
    "ai_trigger",
    "ai_event_span",
    "ai_subject_company",
    "ai_event_time",
    "ai_event_time_normalized",
    "ai_object_target",
    "ai_direction",
    "ai_arguments_json",
    "ai_argument_roles_json",
]
HUMAN_FIELDS = list(HUMAN_REVIEW_FIELDS)
EXTRA_FIELDS = [
    "selected_quality_grade",
    "multi_event_article",
    "reference_raw_text",
]
MAIN_COLUMNS = BASE_FIELDS + AI_FIELDS + HUMAN_FIELDS + EXTRA_FIELDS

BLUE = "D9EAF7"
ORANGE = "FCE4D6"
RED = "F4CCCC"
YELLOW = "FFF2CC"
HEADER = "1F4E78"
WHITE = "FFFFFF"


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def official_event_names() -> tuple[set[str], dict[str, Any]]:
    schema_rows = read_jsonl(EVENT_SCHEMA)
    schema_names = {str(row.get("event_name", "")) for row in schema_rows if row.get("event_name")}
    train_names: set[str] = set()
    target_counts: Counter[str] = Counter()
    label_count = 0
    with TRAIN_LABELS.open("r", encoding="utf-8") as handle:
        for line in handle:
            row = json.loads(line)
            for label in row.get("labels") or []:
                name = str(label.get("event_name", ""))
                if name:
                    train_names.add(name)
                    target_counts[str(label.get("target_type", ""))] += 1
                    label_count += 1
    if schema_names != train_names:
        raise RuntimeError(
            "official event classes mismatch between event_schema.jsonl and training labels: "
            f"schema_only={sorted(schema_names - train_names)} train_only={sorted(train_names - schema_names)}"
        )
    if len(train_names) != 64 or label_count != 2081:
        raise RuntimeError(f"unexpected training label contract: labels={label_count}, classes={len(train_names)}")
    return train_names, {
        "class_count": len(train_names),
        "label_count": label_count,
        "target_type_counts": dict(sorted(target_counts.items())),
        "source": rel(TRAIN_LABELS),
        "source_sha256": sha256_file(TRAIN_LABELS),
        "schema_source": rel(EVENT_SCHEMA),
        "schema_sha256": sha256_file(EVENT_SCHEMA),
    }


def raw_text_by_id() -> dict[str, str]:
    out: dict[str, str] = {}
    path = STEP57C_DIR / "raw_responses.jsonl"
    for row in read_jsonl(path):
        value = str(row.get("final_content", "") or "")
        if not value:
            value = json.dumps(row.get("final_payload") or {}, ensure_ascii=False)
        out[str(row["event_mention_id"])] = value
    return out


def span_from_row(row: dict[str, Any], field: str) -> dict[str, Any]:
    value = row.get(field)
    if field == "event_time":
        default = empty_time()
    else:
        default = empty_span()
    if not isinstance(value, dict):
        return default
    if field == "event_time":
        return {
            "text": str(value.get("text", "") or ""),
            "normalized": str(value.get("normalized", "未明确") or "未明确"),
            "start": int(value.get("start", -1)),
            "end": int(value.get("end", -1)),
        }
    return {
        "text": str(value.get("text", "") or ""),
        "start": int(value.get("start", -1)),
        "end": int(value.get("end", -1)),
    }


def strict_span(content: str, value: dict[str, Any], field: str) -> tuple[dict[str, Any], list[str]]:
    span = dict(value)
    text = str(span.get("text", "") or "")
    try:
        start = int(span.get("start", -1))
        end = int(span.get("end", -1))
    except Exception:
        start, end = -1, -1
    span["start"] = start
    span["end"] = end
    if not text:
        span["start"] = -1
        span["end"] = -1
        if field == "event_time":
            span["normalized"] = str(span.get("normalized", "未明确") or "未明确")
        return span, []
    if not (0 <= start <= end <= len(content)) or content[start:end] != text:
        clean = empty_time() if field == "event_time" else empty_span()
        if field == "event_time":
            clean["normalized"] = str(span.get("normalized", "未明确") or "未明确")
        return clean, [f"{field}:invalid_offset_or_text"]
    return span, []


def sanitize_spans(row: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    content = str(row.get("content", ""))
    spans: dict[str, Any] = {}
    errors: list[str] = []
    for field in SPAN_FIELDS:
        clean, field_errors = strict_span(content, span_from_row(row, field), field)
        spans[field] = clean
        errors.extend(field_errors)
    trigger = spans["trigger"]
    event_span = spans["event_span"]
    if trigger["text"] and event_span["text"]:
        inside = (
            event_span["start"] <= trigger["start"] < trigger["end"] <= event_span["end"]
        )
        if not inside:
            spans["trigger"] = empty_span()
            errors.append("trigger:not_inside_event_span")
    direction = str(row.get("direction", "unclear"))
    if direction not in DIRECTION_LABELS:
        direction = "unclear"
        errors.append("direction:invalid_enum")
    spans["direction"] = direction
    arguments: list[dict[str, Any]] = []
    for role in ARGUMENT_ROLES:
        span = spans.get(role)
        if isinstance(span, dict) and span.get("text"):
            arguments.append(
                {
                    "text": span["text"],
                    "role": role,
                    "start": int(span["start"]),
                    "end": int(span["end"]),
                }
            )
    spans["arguments"] = arguments
    return spans, sorted(set(errors))


def serialise_arguments(arguments: list[dict[str, Any]]) -> tuple[str, str]:
    argument_json = json.dumps(arguments, ensure_ascii=False, separators=(",", ":"))
    role_map = {str(item.get("role", "")): str(item.get("text", "")) for item in arguments}
    role_json = json.dumps(role_map, ensure_ascii=False, separators=(",", ":"))
    return argument_json, role_json


def validate_source_consistency(base: dict[str, Any], diagnostic: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for field in ["article_id", "title", "content", "publish_time"]:
        if str(diagnostic.get(field, "")) != str(base.get(field, "")):
            errors.append(f"step5_7d_{field}_mismatch")
    if str(diagnostic.get("event_name", "")) != str(base.get("event_name", "")):
        errors.append("step5_7d_event_name_mismatch")
    return errors


def build_merged_rows() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    event_names, label_meta = official_event_names()
    frozen_rows = read_jsonl(STEP57C_DIR / "validated_annotation_drafts.jsonl")
    raw_by_id = raw_text_by_id()
    diagnostic_rows = read_jsonl(STEP57D_DIR / "diagnostic40_results.jsonl")
    diagnostic_by_id = {str(row["event_mention_id"]): row for row in diagnostic_rows}
    if len(frozen_rows) != 200:
        raise RuntimeError(f"Step 5.7-C frozen row count is {len(frozen_rows)}, expected 200")
    frozen_ids = [str(row.get("event_mention_id", "")) for row in frozen_rows]
    if len(set(frozen_ids)) != 200 or any(not item for item in frozen_ids):
        raise RuntimeError("Step 5.7-C event_mention_id sequence is not complete and unique")

    rows: list[dict[str, Any]] = []
    source_counts: Counter[str] = Counter()
    grade_counts: Counter[str] = Counter()
    selected_grade_counts: Counter[str] = Counter()
    correction_count = 0
    multi_count = 0
    all_errors: Counter[str] = Counter()

    for frozen in frozen_rows:
        event_name = str(frozen.get("event_name", ""))
        if event_name not in event_names:
            raise RuntimeError(f"frozen event_name is outside the 64 training classes: {event_name}")
        mid = str(frozen["event_mention_id"])
        diagnostic = diagnostic_by_id.get(mid)
        original_grade = str(frozen.get("draft_status", "D"))
        selected = frozen
        source = "step5_7c_unvalidated_reference"
        selected_grade = original_grade
        needs_correction = "yes"
        source_errors: list[str] = []

        if original_grade == "A":
            source = "step5_7c_A_preserved"
            needs_correction = "no"
        elif diagnostic is not None and str(diagnostic.get("draft_status", "")) in {"A", "B"}:
            consistency_errors = validate_source_consistency(frozen, diagnostic)
            if not consistency_errors:
                selected = diagnostic
                source = "step5_7d_sentence_contract"
                selected_grade = str(diagnostic["draft_status"])
                needs_correction = "no"
            else:
                source_errors.extend(consistency_errors)

        spans, span_errors = sanitize_spans(selected)
        validation_errors = list(frozen.get("contract_errors") or [])
        if selected is not frozen:
            validation_errors.extend(diagnostic.get("sentence_contract_errors") or [])
        validation_errors.extend(source_errors)
        validation_errors.extend(span_errors)
        validation_errors = sorted(set(str(item) for item in validation_errors if str(item)))

        base = {
            "event_mention_id": mid,
            "article_id": str(frozen.get("article_id", "")),
            "event_name": event_name,
            "schema_id": str(frozen.get("schema_id", "")),
            "target_type": str(frozen.get("target_type", "")),
            "target_code": str(frozen.get("target_code", "")),
            "target_name": str(frozen.get("target_name", "")),
            "title": str(frozen.get("title", "")),
            "content": str(frozen.get("content", "")),
            "publish_time": str(frozen.get("publish_time", "")),
        }
        arguments_json, roles_json = serialise_arguments(spans["arguments"])
        row = {
            **base,
            "suggestion_source": source,
            "original_quality_grade": original_grade,
            "needs_manual_correction": needs_correction,
            "validation_errors": validation_errors,
            "ai_trigger": spans["trigger"]["text"],
            "ai_event_span": spans["event_span"]["text"],
            "ai_subject_company": spans["subject_company"]["text"],
            "ai_event_time": spans["event_time"]["text"],
            "ai_event_time_normalized": spans["event_time"].get("normalized", "未明确"),
            "ai_object_target": spans["object_target"]["text"],
            "ai_direction": spans["direction"],
            "ai_arguments_json": arguments_json,
            "ai_argument_roles_json": roles_json,
            "selected_quality_grade": selected_grade,
            "multi_event_article": bool(frozen.get("multi_event_article", False)),
            "reference_raw_text": raw_by_id.get(mid, json.dumps(frozen.get("ai_raw_text_fields", {}), ensure_ascii=False)),
            **{field: "" for field in HUMAN_FIELDS},
            "ai_offsets": {
                field: spans[field]
                for field in SPAN_FIELDS
            },
            "ai_arguments": spans["arguments"],
            "step5_7d_available": diagnostic is not None,
            "step5_7d_quality_grade": str(diagnostic.get("draft_status", "")) if diagnostic else "",
            "step5_7d_sentence_contract_errors": list(diagnostic.get("sentence_contract_errors") or []) if diagnostic else [],
            "artifact_type": "ai_assisted_review_template",
        }
        if row["needs_manual_correction"] == "yes":
            correction_count += 1
        if row["multi_event_article"]:
            multi_count += 1
        source_counts[source] += 1
        grade_counts[original_grade] += 1
        selected_grade_counts[selected_grade] += 1
        for error in validation_errors:
            all_errors[error] += 1
        validation_row = {
            **base,
            "event_mention_id": mid,
            "trigger": spans["trigger"],
            "event_span": spans["event_span"],
            "subject_company": spans["subject_company"],
            "event_time": spans["event_time"],
            "object_target": spans["object_target"],
            "direction": spans["direction"],
            "arguments": spans["arguments"],
        }
        validation = validate_eventmention_record(validation_row, event_names)
        if not validation["valid"]:
            raise RuntimeError(f"sanitized suggestion is invalid for {mid}: {validation['errors']}")
        rows.append(row)

    manifest = {
        "generated_at": now(),
        "status": "ready_for_human_review",
        "artifact_type": "ai_assisted_review_template_only",
        "gold": False,
        "confirmed": False,
        "training_ready": False,
        "reference": "Step 5.7-E unified human-confirmation workbook and machine-readable templates; AI assistance only.",
        "source_contract": {
            "step5_7c_validated_sha256": sha256_file(STEP57C_DIR / "validated_annotation_drafts.jsonl"),
            "step5_7c_raw_responses_sha256": sha256_file(STEP57C_DIR / "raw_responses.jsonl"),
            "step5_7d_results_sha256": sha256_file(STEP57D_DIR / "diagnostic40_results.jsonl"),
            "step5_7d_raw_responses_sha256": sha256_file(STEP57D_DIR / "diagnostic40_raw_responses.jsonl"),
            "confirmed80_sha256": sha256_file(CONFIRMED80),
            "migrated80_sha256": sha256_file(MIGRATED80),
            "train_pairs_sha256": sha256_file(TRAIN_PAIRS),
            "official_event_classes": label_meta,
        },
        "counts": {
            "eventmention_count": len(rows),
            "unique_event_mention_id_count": len({row["event_mention_id"] for row in rows}),
            "unique_article_count": len({row["article_id"] for row in rows}),
            "multi_event_mention_count": multi_count,
            "needs_manual_correction_count": correction_count,
            "original_quality_grade_counts": dict(sorted(grade_counts.items())),
            "selected_quality_grade_counts": dict(sorted(selected_grade_counts.items())),
            "suggestion_source_counts": dict(sorted(source_counts.items())),
            "diagnostic40_count": len(diagnostic_rows),
            "diagnostic40_used_A_or_B_count": sum(1 for row in rows if row["suggestion_source"] == "step5_7d_sentence_contract"),
        },
        "old_contract_review": {
            "scope": "three_pending_rows_from_old_80_contract",
            "event_mention_ids": sorted(EXPECTED_OLD_PENDING),
            "reviewer_confirmed_initial": "",
            "included_in_main_200": False,
        },
        "validation_errors": dict(sorted(all_errors.items())),
        "protection": {
            "gold_test_labels_read": False,
            "gold_test_content_read": False,
            "training_or_finetune_executed": False,
            "vllm_started": False,
            "local_32b_calls": 0,
            "external_llm_calls": 0,
            "formal_task1_task2_modified": False,
            "confirmed80_modified": False,
            "step5_7c_modified": False,
            "step5_7d_modified": False,
        },
    }
    return rows, manifest


def template_row(row: dict[str, Any]) -> dict[str, Any]:
    return {key: row.get(key, "") for key in MAIN_COLUMNS} | {
        "schema_id": row.get("schema_id", ""),
        "target_type": row.get("target_type", ""),
        "target_code": row.get("target_code", ""),
        "target_name": row.get("target_name", ""),
        "ai_offsets": row.get("ai_offsets", {}),
        "ai_arguments": row.get("ai_arguments", []),
        "step5_7d_available": row.get("step5_7d_available", False),
        "step5_7d_quality_grade": row.get("step5_7d_quality_grade", ""),
        "step5_7d_sentence_contract_errors": row.get("step5_7d_sentence_contract_errors", []),
        "artifact_type": "ai_assisted_review_template",
        "gold": False,
        "confirmed": False,
    }


def build_old_contract_rows() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    migrated = read_jsonl(MIGRATED80)
    if len(migrated) != 80:
        raise RuntimeError(f"migrated old contract rows={len(migrated)}, expected 80")
    pending = []
    display = []
    for row in migrated:
        aid = str(row.get("article_id", ""))
        spans = {field: span_from_row(row, field) for field in SPAN_FIELDS}
        arguments_json, roles_json = serialise_arguments(row.get("arguments") or [])
        base = {
            "event_mention_id": str(row.get("event_mention_id", "")),
            "article_id": aid,
            "event_name": str(row.get("event_name", "")),
            "title": str(row.get("title", "")),
            "content": str(row.get("content", "")),
            "publish_time": str(row.get("publish_time", "")),
            "suggestion_source": "legacy_80_migrated_reference",
            "original_quality_grade": "legacy_confirmed" if row.get("annotation_status") == "confirmed" else "legacy_pending",
            "needs_manual_correction": "yes" if aid in EXPECTED_OLD_PENDING else "no",
            "validation_errors": list(row.get("contract_errors") or []),
            "ai_trigger": spans["trigger"]["text"],
            "ai_event_span": spans["event_span"]["text"],
            "ai_subject_company": spans["subject_company"]["text"],
            "ai_event_time": spans["event_time"]["text"],
            "ai_event_time_normalized": spans["event_time"].get("normalized", "未明确"),
            "ai_object_target": spans["object_target"]["text"],
            "ai_direction": str(row.get("direction", "unclear")),
            "ai_arguments_json": arguments_json,
            "ai_argument_roles_json": roles_json,
            "selected_quality_grade": "legacy",
            "multi_event_article": aid in EXPECTED_OLD_PENDING,
            "reference_raw_text": json.dumps(row.get("legacy_structure", {}), ensure_ascii=False),
            **{field: "" for field in HUMAN_FIELDS},
            "legacy_annotation_status": str(row.get("annotation_status", "")),
            "legacy_reviewer_confirmed": str(row.get("reviewer_confirmed", "")),
            "requires_reconfirmation": "yes" if aid in EXPECTED_OLD_PENDING else "no",
            "alignment_category": str(row.get("alignment_category", "")),
            "migration_note": str(row.get("migration_note", "")),
            "ai_offsets": spans,
            "ai_arguments": row.get("arguments") or [],
            "artifact_type": "ai_assisted_review_template",
            "gold": False,
            "confirmed": False,
        }
        display.append(base)
        if aid in EXPECTED_OLD_PENDING:
            pending.append(base)
    if {row["article_id"] for row in pending} != EXPECTED_OLD_PENDING:
        raise RuntimeError("old pending contract rows do not match the required three article IDs")
    return pending, display


def display_value(value: Any) -> Any:
    if isinstance(value, list):
        return "|".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def write_table(
    ws: Any,
    rows: list[dict[str, Any]],
    columns: list[str],
    *,
    color_rows: dict[int, str] | None = None,
    ai_columns: set[str] | None = None,
    human_columns: set[str] | None = None,
    max_width: int = 42,
) -> None:
    color_rows = color_rows or {}
    ai_columns = ai_columns or set()
    human_columns = human_columns or set()
    ws.append(columns)
    header_fill = PatternFill("solid", fgColor=HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(rows, start=2):
        ws.append([display_value(row.get(column, "")) for column in columns])
        row_fill = PatternFill("solid", fgColor=color_rows.get(index, "FFFFFF"))
        for col_idx, column in enumerate(columns, start=1):
            cell = ws.cell(index, col_idx)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = row_fill
            if column in ai_columns:
                cell.fill = PatternFill("solid", fgColor=BLUE)
            elif column in human_columns:
                cell.fill = PatternFill("solid", fgColor=ORANGE)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    for col_idx, column in enumerate(columns, start=1):
        width = min(max_width, max(12, len(column) + 2))
        if column in {"content", "reference_raw_text"}:
            width = max_width
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def add_review_validations(ws: Any, columns: list[str], start_row: int, end_row: int) -> None:
    direction_col = get_column_letter(columns.index("human_direction") + 1)
    confirmed_col = get_column_letter(columns.index("reviewer_confirmed") + 1)
    direction = DataValidation(type="list", formula1='"positive,negative,neutral,unclear"', allow_blank=True)
    confirmed = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(direction)
    ws.add_data_validation(confirmed)
    direction.add(f"{direction_col}{start_row}:{direction_col}{end_row}")
    confirmed.add(f"{confirmed_col}{start_row}:{confirmed_col}{end_row}")


def write_workbook(rows: list[dict[str, Any]], pending_old: list[dict[str, Any]], display_old: list[dict[str, Any]], path: Path) -> None:
    wb = Workbook()
    instruction = wb.active
    instruction.title = "使用说明"
    instructions = [
        ["字段", "说明"],
        ["包类型", "AI辅助审核模板，仅供人工确认；不是gold，不是confirmed，不得直接用于训练。"],
        ["主键", "以Step 5.7-C冻结的event_mention_id为唯一主键，200条顺序固定。"],
        ["event_name", "冻结训练集64类监督标签；不得用模型文本覆盖，不生成自由event_name。"],
        ["AI建议", "浅蓝色。C/D非法字段已清空，原始模型文本保存在reference_raw_text。"],
        ["人工填写", "浅橙色。人工字段初始为空；reviewer_confirmed初始为空。"],
        ["审核顺序", "优先处理B/C/D及多事件文章；确认前核对正文逐字片段、offset、direction和论元JSON。"],
        ["direction", "只能填positive、negative、neutral、unclear。"],
        ["reviewer_confirmed", "只能填yes或no；200条主表完成后才可全部填yes。"],
        ["旧80条", "77条历史确认仅作迁移参考；指定的3条旧契约记录必须重新人工确认。"],
        ["保护边界", "本包未读取1029篇gold test正文或标签，未调用32B/外部LLM，未训练模型，不修改正式输出。"],
    ]
    for row in instructions:
        instruction.append(row)
    instruction.freeze_panes = "A2"
    instruction.column_dimensions["A"].width = 24
    instruction.column_dimensions["B"].width = 110
    for cell in instruction[1]:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=HEADER)
    for row in instruction.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    colors: dict[int, str] = {}
    for index, row in enumerate(rows, start=2):
        if row["original_quality_grade"] in {"C", "D"} or row["needs_manual_correction"] == "yes":
            colors[index] = RED
        if row.get("multi_event_article"):
            colors[index] = YELLOW
    main = wb.create_sheet("200条主审核表")
    write_table(main, rows, MAIN_COLUMNS, color_rows=colors, ai_columns=set(AI_FIELDS), human_columns=set(HUMAN_FIELDS), max_width=48)
    add_review_validations(main, MAIN_COLUMNS, 2, len(rows) + 1)

    low = wb.create_sheet("A级低风险记录")
    low_rows = [row for row in rows if row["original_quality_grade"] == "A" and row["suggestion_source"] == "step5_7c_A_preserved"]
    write_table(low, low_rows, MAIN_COLUMNS, color_rows={}, ai_columns=set(AI_FIELDS), human_columns=set(HUMAN_FIELDS), max_width=48)
    add_review_validations(low, MAIN_COLUMNS, 2, max(2, len(low_rows) + 1))

    high = wb.create_sheet("B／C／D高风险记录")
    high_rows = [row for row in rows if row not in low_rows]
    high_colors = {index: RED for index in range(2, len(high_rows) + 2)}
    write_table(high, high_rows, MAIN_COLUMNS, color_rows=high_colors, ai_columns=set(AI_FIELDS), human_columns=set(HUMAN_FIELDS), max_width=48)
    add_review_validations(high, MAIN_COLUMNS, 2, max(2, len(high_rows) + 1))

    multi = wb.create_sheet("多事件文章")
    multi_rows = [row for row in rows if row.get("multi_event_article")]
    write_table(multi, multi_rows, MAIN_COLUMNS, color_rows={i: YELLOW for i in range(2, len(multi_rows) + 2)}, ai_columns=set(AI_FIELDS), human_columns=set(HUMAN_FIELDS), max_width=48)
    add_review_validations(multi, MAIN_COLUMNS, 2, max(2, len(multi_rows) + 1))

    offsets = wb.create_sheet("字段与Offset详情")
    offset_columns = [
        "event_mention_id",
        "article_id",
        "event_name",
        "suggestion_source",
        "original_quality_grade",
        "field",
        "text",
        "start",
        "end",
        "content_slice_matches",
        "trigger_inside_event_span",
        "validation_errors",
    ]
    offset_rows: list[dict[str, Any]] = []
    for row in rows:
        for field in SPAN_FIELDS:
            span = row["ai_offsets"].get(field, empty_time() if field == "event_time" else empty_span())
            content = row["content"]
            matched = bool(span.get("text")) and content[span["start"]:span["end"]] == span["text"] if span.get("start", -1) >= 0 else not span.get("text")
            inside = True
            if field == "trigger" and row["ai_offsets"]["event_span"].get("text") and span.get("text"):
                event_span = row["ai_offsets"]["event_span"]
                inside = event_span["start"] <= span["start"] < span["end"] <= event_span["end"]
            offset_rows.append(
                {
                    "event_mention_id": row["event_mention_id"],
                    "article_id": row["article_id"],
                    "event_name": row["event_name"],
                    "suggestion_source": row["suggestion_source"],
                    "original_quality_grade": row["original_quality_grade"],
                    "field": field,
                    "text": span.get("text", ""),
                    "start": span.get("start", -1),
                    "end": span.get("end", -1),
                    "content_slice_matches": "yes" if matched else "no",
                    "trigger_inside_event_span": "yes" if inside else "no",
                    "validation_errors": "|".join(row["validation_errors"]),
                }
            )
    write_table(offsets, offset_rows, offset_columns, max_width=48)

    old = wb.create_sheet("旧80条契约确认")
    old_columns = [
        "event_mention_id",
        "article_id",
        "event_name",
        "title",
        "content",
        "publish_time",
        "legacy_annotation_status",
        "legacy_reviewer_confirmed",
        "requires_reconfirmation",
        "alignment_category",
        "migration_note",
        "ai_trigger",
        "ai_event_span",
        "ai_subject_company",
        "ai_event_time",
        "ai_event_time_normalized",
        "ai_object_target",
        "ai_direction",
        "ai_arguments_json",
        "ai_argument_roles_json",
        "human_trigger",
        "human_event_span",
        "human_subject_company",
        "human_event_time",
        "human_event_time_normalized",
        "human_object_target",
        "human_direction",
        "human_arguments_json",
        "human_argument_roles_json",
        "reviewer_confirmed",
        "reviewer_note",
    ]
    old_colors = {index: RED for index, row in enumerate(display_old, start=2) if row["requires_reconfirmation"] == "yes"}
    write_table(old, display_old, old_columns, color_rows=old_colors, ai_columns=set(old_columns[11:20]), human_columns=set(old_columns[20:]), max_width=48)
    add_review_validations(old, old_columns, 2, len(display_old) + 1)

    coverage = wb.create_sheet("类别覆盖统计")
    coverage_rows: list[dict[str, Any]] = []
    for name in sorted({row["event_name"] for row in rows} | official_event_names()[0]):
        matching = [row for row in rows if row["event_name"] == name]
        coverage_rows.append(
            {
                "event_name": name,
                "selected_count": len(matching),
                "original_A": sum(row["original_quality_grade"] == "A" for row in matching),
                "original_B": sum(row["original_quality_grade"] == "B" for row in matching),
                "original_C": sum(row["original_quality_grade"] == "C" for row in matching),
                "original_D": sum(row["original_quality_grade"] == "D" for row in matching),
                "step5_7d_used": sum(row["suggestion_source"] == "step5_7d_sentence_contract" for row in matching),
                "needs_manual_correction": sum(row["needs_manual_correction"] == "yes" for row in matching),
            }
        )
    write_table(coverage, coverage_rows, list(coverage_rows[0].keys()), max_width=30)

    errors = wb.create_sheet("错误类型统计")
    error_counts: Counter[str] = Counter()
    for row in rows:
        error_counts.update(row["validation_errors"])
    error_rows = [{"error_type": key, "count": value} for key, value in sorted(error_counts.items())]
    if not error_rows:
        error_rows = [{"error_type": "", "count": 0}]
    write_table(errors, error_rows, ["error_type", "count"], max_width=80)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def guide_text(manifest: dict[str, Any]) -> str:
    counts = manifest["counts"]
    return f"""# Task1 Step 5.7-E EventMention人工确认指南

## 包性质

本目录和工作簿是 **AI辅助审核模板**，不是gold、不是confirmed，也不是可直接训练的数据。审核完成前不得训练模型或扩大到500/800条。

200条主表以Step 5.7-C冻结的`event_mention_id`作为唯一主键，顺序固定。`event_name`来自训练集固定64类，构建过程中没有使用模型回显覆盖它，也没有生成自由类别。

## 建议来源

- `step5_7c_A_preserved`：保留5.7-C的A级输出，不自动改写。
- `step5_7d_sentence_contract`：仅对5.7-D诊断40条中的A/B记录使用句子契约结果。
- `step5_7c_unvalidated_reference`：其余C/D保留5.7-C原始模型文本作为参考，并标记需要人工修正。

非法offset或非正文字段已清空；原始模型文本保存在`reference_raw_text`。正文、标题、文章ID和冻结`event_name`不可修改。

## 人工填写规则

人工字段初始为空。人工填写时：

1. 核对AI文本是否逐字属于对应`content`。
2. 核对`human_event_span`是正文连续片段，`human_trigger`位于该片段内。
3. `human_direction`只能填写`positive`、`negative`、`neutral`或`unclear`。
4. `human_arguments_json`和`human_argument_roles_json`必须是可解析JSON。
5. 主表200条全部完成后，`reviewer_confirmed`才可全部填写`yes`或必要时填写`no`。
6. 旧80条中的3条待确认记录必须单独确认，不能继承或自动设置`reviewer_confirmed=yes`。

## 当前统计

- 200条：{counts["eventmention_count"]}
- 唯一EventMention：{counts["unique_event_mention_id_count"]}
- 原始质量等级：{json.dumps(manifest["counts"]["original_quality_grade_counts"], ensure_ascii=False)}
- 建议来源：{json.dumps(manifest["counts"]["suggestion_source_counts"], ensure_ascii=False)}
- 需要人工修正：{counts["needs_manual_correction_count"]}
- 多事件文章记录：{counts["multi_event_mention_count"]}
- 旧契约待确认：3条

## 校验命令

空模板：

```bash
python scripts/validate_task1_step5_7e_human_confirmation.py
```

填写后：

```bash
python scripts/validate_task1_step5_7e_human_confirmation.py \\
  --allow-filled --workbook reports/task1_step5_7e_eventmention_human_confirmation.xlsx
```

本轮保护状态：未训练、未启动vLLM、32B调用0、外部LLM调用0、未读取1029篇gold test标签，正式Task1/Task2和5.7-C/5.7-D冻结输出未修改。
"""


def existing_package_state() -> str:
    package_files = [
        OUT_DIR / "eventmention200_review_template.jsonl",
        OUT_DIR / "eventmention200_ai_suggestions.jsonl",
        OUT_DIR / "eventmention200_manifest.json",
        OUT_DIR / "contract_review3_template.jsonl",
        OUT_DIR / "SHA256SUMS.txt",
        REPORT_XLSX,
        REPORT_GUIDE,
        REPORT_MANIFEST,
    ]
    present = [path for path in package_files if path.exists()]
    if not present:
        return "absent"
    if len(present) == len(package_files):
        return "complete"
    return "partial"


def write_sha256sums(files: list[Path], path: Path) -> dict[str, str]:
    rows = {rel(file): sha256_file(file) for file in files if file.exists()}
    path.write_text("".join(f"{digest}  {name}\n" for name, digest in sorted(rows.items())), encoding="utf-8")
    return rows


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--allow-existing-complete", action="store_true")
    args = parser.parse_args()

    state = existing_package_state()
    if state == "complete" and not args.allow_existing_complete:
        print(json.dumps({"status": "existing_complete_package", "action": "no_overwrite"}, ensure_ascii=False, indent=2))
        return 0
    if state == "partial":
        print(json.dumps({"status": "partial_package_detected", "action": "resume_and_complete_missing_outputs"}, ensure_ascii=False))

    rows, manifest = build_merged_rows()
    pending_old, display_old = build_old_contract_rows()
    if len(rows) != 200:
        raise SystemExit(f"merged row count={len(rows)}, expected 200")
    if len(pending_old) != 3:
        raise SystemExit(f"old pending row count={len(pending_old)}, expected 3")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    review_rows = [template_row(row) for row in rows]
    suggestion_rows = [
        {
            key: row[key]
            for key in [
                "event_mention_id",
                "article_id",
                "event_name",
                "title",
                "content",
                "publish_time",
                "suggestion_source",
                "original_quality_grade",
                "selected_quality_grade",
                "needs_manual_correction",
                "validation_errors",
                "ai_trigger",
                "ai_event_span",
                "ai_subject_company",
                "ai_event_time",
                "ai_event_time_normalized",
                "ai_object_target",
                "ai_direction",
                "ai_arguments_json",
                "ai_argument_roles_json",
                "ai_offsets",
                "ai_arguments",
                "reference_raw_text",
                "multi_event_article",
                "step5_7d_available",
                "step5_7d_quality_grade",
                "step5_7d_sentence_contract_errors",
                "artifact_type",
            ]
        }
        for row in rows
    ]
    write_jsonl(OUT_DIR / "eventmention200_review_template.jsonl", review_rows)
    write_jsonl(OUT_DIR / "eventmention200_ai_suggestions.jsonl", suggestion_rows)
    write_jsonl(OUT_DIR / "contract_review3_template.jsonl", pending_old)

    manifest["backup_dir"] = Path("/tmp/task1_step5_7e_backup_dir.txt").read_text(encoding="utf-8").strip() if Path("/tmp/task1_step5_7e_backup_dir.txt").exists() else ""
    manifest["output_paths"] = {
        "workbook": rel(REPORT_XLSX),
        "guide": rel(REPORT_GUIDE),
        "manifest": rel(REPORT_MANIFEST),
        "review_template": rel(OUT_DIR / "eventmention200_review_template.jsonl"),
        "ai_suggestions": rel(OUT_DIR / "eventmention200_ai_suggestions.jsonl"),
        "old_contract_review3": rel(OUT_DIR / "contract_review3_template.jsonl"),
    }
    write_workbook(rows, pending_old, display_old, REPORT_XLSX)
    REPORT_GUIDE.write_text(guide_text(manifest), encoding="utf-8")
    write_json(REPORT_MANIFEST, manifest)
    write_json(OUT_DIR / "eventmention200_manifest.json", manifest)

    artifacts = [
        OUT_DIR / "eventmention200_review_template.jsonl",
        OUT_DIR / "eventmention200_ai_suggestions.jsonl",
        OUT_DIR / "eventmention200_manifest.json",
        OUT_DIR / "contract_review3_template.jsonl",
        REPORT_XLSX,
        REPORT_GUIDE,
        REPORT_MANIFEST,
    ]
    output_hashes = {rel(file): sha256_file(file) for file in artifacts if file.exists() and file not in {REPORT_MANIFEST, OUT_DIR / "eventmention200_manifest.json"}}
    manifest["output_sha256"] = output_hashes
    write_json(REPORT_MANIFEST, manifest)
    write_json(OUT_DIR / "eventmention200_manifest.json", manifest)
    output_hashes = write_sha256sums(artifacts, OUT_DIR / "SHA256SUMS.txt")

    print(json.dumps(
        {
            "status": manifest["status"],
            "eventmention_count": len(rows),
            "old_contract_review3_count": len(pending_old),
            "suggestion_source_counts": manifest["counts"]["suggestion_source_counts"],
            "original_quality_grade_counts": manifest["counts"]["original_quality_grade_counts"],
            "needs_manual_correction_count": manifest["counts"]["needs_manual_correction_count"],
            "multi_event_mention_count": manifest["counts"]["multi_event_mention_count"],
            "output_sha256_count": len(output_hashes),
        },
        ensure_ascii=False,
        indent=2,
    ))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
