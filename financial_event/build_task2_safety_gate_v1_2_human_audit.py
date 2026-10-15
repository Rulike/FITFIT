#!/usr/bin/env python3
"""Build the blank human-audit workbook for Safety Gate v1.2 independent validation."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


DEFAULT_OUTPUT_DIR = ROOT / "outputs/event_credibility/dea_safety_gate_v1_2_independent"
DEFAULT_WORKBOOK = ROOT / "reports/task2_safety_gate_v1_2_independent_human_audit.xlsx"
DEFAULT_GUIDE = ROOT / "reports/task2_safety_gate_v1_2_independent_human_audit_guide.md"
DEFAULT_MANIFEST = ROOT / "reports/task2_safety_gate_v1_2_independent_manifest.json"

MAIN_SHEETS = {
    "primary": "Primary人工审核",
    "contradiction_stress": "Contradiction压力审核",
}

HUMAN_FIELDS = [
    "human_verdict",
    "human_evidence_sufficient",
    "human_subject_match",
    "human_time_match",
    "human_error_type",
    "human_notes",
    "reviewer_confirmed",
]

ALLOWED_VALUES = {
    "human_verdict": ["support", "contradict", "insufficient", "unclear"],
    "human_evidence_sufficient": ["yes", "partial", "no", "uncertain"],
    "human_subject_match": ["yes", "no", "uncertain"],
    "human_time_match": ["yes", "no", "uncertain", "not_applicable"],
    "human_error_type": [
        "none",
        "claim_unclear",
        "evidence_insufficient",
        "subject_mismatch",
        "time_mismatch",
        "investor_question_only",
        "market_signal_only",
        "contradiction",
        "model_error",
        "gate_error",
        "other",
    ],
    "reviewer_confirmed": ["yes", "no"],
}

MAIN_COLUMNS = [
    "sample_group",
    "cluster_id",
    "claim",
    "company/subject",
    "schema_id",
    "event_name",
    "event_time",
    "baseline_verdict",
    "round1_verdict",
    "cached_verdict",
    "safety_gate_v1_2_verdict",
    "safety_gate_v1_2_selected_source",
    "gate_blocked",
    "gate_reasons",
    "cited_evidence_ids",
    "evidence摘要",
    "来源",
    "发布时间",
    "AI辅助建议_verdict",
    "AI辅助建议_confidence",
    "AI辅助建议_evidence_ids",
    "AI辅助建议_reason",
    *HUMAN_FIELDS,
]

FORMAL_OUTPUTS = [
    Path("outputs/event_credibility/llm_review_results.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.parquet"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.xlsx"),
]


def norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def resolve(path: str | Path) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def hash_files(paths: list[Path]) -> dict[str, str]:
    return {str(path): sha256(resolve(path)) if resolve(path).exists() else "missing" for path in paths}


def unique_by_cluster(rows: list[dict[str, Any]], name: str) -> dict[str, dict[str, Any]]:
    out: dict[str, dict[str, Any]] = {}
    duplicates: list[str] = []
    for row in rows:
        cid = norm(row.get("cluster_id"))
        if not cid:
            continue
        if cid in out:
            duplicates.append(cid)
        out[cid] = row
    if duplicates:
        raise ValueError(f"{name} duplicate cluster_id: {sorted(set(duplicates))[:10]}")
    return out


def stringify(value: Any) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return norm(value)


def evidence_text(item: dict[str, Any]) -> str:
    return norm(item.get("sentence") or item.get("text") or item.get("content") or item.get("证据文本"))


def evidence_source(item: dict[str, Any]) -> str:
    return norm(item.get("source") or item.get("article_source") or item.get("来源"))


def evidence_time(item: dict[str, Any]) -> str:
    return norm(item.get("publish_time") or item.get("article_publish_time") or item.get("发布时间"))


def evidence_id(item: dict[str, Any]) -> str:
    return norm(item.get("evidence_id") or item.get("id"))


def sample_evidence_items(sample: dict[str, Any]) -> list[dict[str, Any]]:
    items = sample.get("evidence_items", [])
    return [dict(item) for item in items if isinstance(item, dict)] if isinstance(items, list) else []


def evidence_summary(items: list[dict[str, Any]], limit: int = 4) -> str:
    parts = []
    for item in items[:limit]:
        text = evidence_text(item)
        parts.append(f"{evidence_id(item)}｜{evidence_source(item)}｜{evidence_time(item)}｜{text[:180]}")
    if len(items) > limit:
        parts.append(f"... 另有{len(items) - limit}条证据")
    return "\n".join(parts)


def join_unique(values: list[Any]) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        item = norm(value)
        if item and item not in seen:
            seen.add(item)
            out.append(item)
    return ";".join(out)


def model_suggestion(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "AI辅助建议_verdict": norm(row.get("verdict")),
        "AI辅助建议_confidence": row.get("confidence", ""),
        "AI辅助建议_evidence_ids": stringify(row.get("evidence_ids") or []),
        "AI辅助建议_reason": norm(row.get("reason")),
    }


def build_main_rows(data: dict[str, Any], sample_group: str) -> list[dict[str, Any]]:
    samples = [
        row
        for row in data["samples"]
        if norm(row.get("sample_group")) == sample_group
    ]
    baseline = unique_by_cluster(data["baseline"], "baseline")
    round1 = unique_by_cluster(data["round1"], "round1")
    cached = unique_by_cluster(data["cached"], "cached")
    gate = unique_by_cluster(data["gate"], "safety_gate")
    model = unique_by_cluster(data["model"], "model_request_log")
    rows: list[dict[str, Any]] = []
    for sample in sorted(samples, key=lambda row: int(row.get("input_index", 10**9))):
        cid = norm(sample.get("cluster_id"))
        evidence_items = sample_evidence_items(sample)
        gate_row = gate.get(cid, {})
        row = {
            "sample_group": sample_group,
            "cluster_id": cid,
            "claim": norm(sample.get("event_claim") or sample.get("claim_text")),
            "company/subject": norm(sample.get("target_name") or sample.get("subject")),
            "schema_id": norm(sample.get("schema_id")),
            "event_name": norm(sample.get("event_name")),
            "event_time": norm(sample.get("event_time")),
            "baseline_verdict": norm(baseline.get(cid, {}).get("verdict")),
            "round1_verdict": norm(round1.get(cid, {}).get("verdict")),
            "cached_verdict": norm(cached.get(cid, {}).get("final_verdict")),
            "safety_gate_v1_2_verdict": norm(gate_row.get("safety_gate_v1_2_verdict")),
            "safety_gate_v1_2_selected_source": norm(gate_row.get("safety_gate_v1_2_selected_source")),
            "gate_blocked": bool(gate_row.get("gate_blocked")),
            "gate_reasons": stringify(gate_row.get("gate_reasons") or []),
            "cited_evidence_ids": stringify(gate_row.get("cited_evidence_ids") or []),
            "evidence摘要": evidence_summary(evidence_items),
            "来源": join_unique([evidence_source(item) for item in evidence_items]),
            "发布时间": join_unique([evidence_time(item) for item in evidence_items]),
        }
        row.update(model_suggestion(model.get(cid, {})))
        for field in HUMAN_FIELDS:
            row[field] = ""
        rows.append({column: row.get(column, "") for column in MAIN_COLUMNS})
    return rows


def build_evidence_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    gate = unique_by_cluster(data["gate"], "safety_gate")
    rows: list[dict[str, Any]] = []
    for sample in sorted(data["samples"], key=lambda row: int(row.get("input_index", 10**9))):
        cid = norm(sample.get("cluster_id"))
        gate_row = gate.get(cid, {})
        cited = set(map(norm, gate_row.get("cited_evidence_ids") or []))
        valid_cited = set(map(norm, gate_row.get("valid_cited_evidence_ids") or []))
        for item in sample_evidence_items(sample):
            eid = evidence_id(item)
            rows.append(
                {
                    "sample_group": norm(sample.get("sample_group")),
                    "input_index": sample.get("input_index", ""),
                    "cluster_id": cid,
                    "evidence_id": eid,
                    "document_id": norm(item.get("document_id") or item.get("article_id") or item.get("trace_id")),
                    "source": evidence_source(item),
                    "publish_time": evidence_time(item),
                    "title": norm(item.get("title") or item.get("article_title")),
                    "evidence_text": evidence_text(item),
                    "traceable": bool(eid and evidence_source(item) and evidence_text(item)),
                    "cited_by_cached": eid in cited,
                    "valid_cited_by_safety_gate": eid in valid_cited,
                    "evidence_source": norm(item.get("evidence_source") or item.get("retrieval_source")),
                }
            )
    return rows


def build_comparison_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    samples = unique_by_cluster(data["samples"], "samples")
    baseline = unique_by_cluster(data["baseline"], "baseline")
    round1 = unique_by_cluster(data["round1"], "round1")
    cached = unique_by_cluster(data["cached"], "cached")
    gate = unique_by_cluster(data["gate"], "safety_gate")
    model = unique_by_cluster(data["model"], "model_request_log")
    rows = []
    for cid in sorted(samples, key=lambda key: int(samples[key].get("input_index", 10**9))):
        gate_row = gate.get(cid, {})
        model_row = model.get(cid, {})
        rows.append(
            {
                "input_index": samples[cid].get("input_index", ""),
                "sample_group": samples[cid].get("sample_group", ""),
                "cluster_id": cid,
                "baseline_verdict": baseline.get(cid, {}).get("verdict", ""),
                "baseline_confidence": baseline.get(cid, {}).get("confidence", ""),
                "round1_verdict": round1.get(cid, {}).get("verdict", ""),
                "round1_confidence": round1.get(cid, {}).get("confidence", ""),
                "cached_verdict": cached.get(cid, {}).get("final_verdict", ""),
                "cached_confidence": cached.get(cid, {}).get("final_confidence", ""),
                "safety_gate_v1_2_verdict": gate_row.get("safety_gate_v1_2_verdict", ""),
                "selected_result_source": gate_row.get("safety_gate_v1_2_selected_source", ""),
                "gate_blocked": gate_row.get("gate_blocked", ""),
                "gate_reasons": stringify(gate_row.get("gate_reasons") or []),
                "AI辅助建议_verdict": model_row.get("verdict", ""),
                "AI辅助建议_confidence": model_row.get("confidence", ""),
                "AI辅助建议_evidence_ids": stringify(model_row.get("evidence_ids") or []),
                "AI辅助建议_reason": model_row.get("reason", ""),
                "model_status": model_row.get("status", ""),
                "model_error_type": model_row.get("error_type", ""),
            }
        )
    return rows


def flatten_manifest(manifest: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def visit(prefix: str, value: Any) -> None:
        if isinstance(value, dict) and len(value) <= 30:
            for key, item in value.items():
                visit(f"{prefix}.{key}" if prefix else str(key), item)
        else:
            rows.append({"item": prefix, "value": stringify(value)})

    visit("", manifest)
    return rows


def guide_rows() -> list[dict[str, str]]:
    return [
        {
            "field": field,
            "allowed_values": ",".join(ALLOWED_VALUES.get(field, [])),
            "description": {
                "human_verdict": "人工确认最终证据结论。",
                "human_evidence_sufficient": "证据是否足以支持人工结论。",
                "human_subject_match": "证据主体与主张主体是否匹配。",
                "human_time_match": "证据时间与事件时间是否匹配。",
                "human_error_type": "若模型或门控错误，标注主要错误类型。",
                "human_notes": "人工备注。",
                "reviewer_confirmed": "复核完成后填写 yes。",
            }.get(field, ""),
        }
        for field in HUMAN_FIELDS
    ]


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title[:31])
    if not rows:
        rows = [{"empty": ""}]
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([stringify(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, 1):
        max_len = max([len(stringify(header))] + [len(stringify(row.get(header, ""))) for row in rows[:100]])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 80)


def write_workbook(path: Path, workbook_rows: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name in (
        "Primary人工审核",
        "Contradiction压力审核",
        "证据明细",
        "模型结果对比",
        "运行统计",
        "填写说明",
    ):
        write_sheet(wb, sheet_name, workbook_rows[sheet_name])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def load_inputs(output_dir: Path) -> dict[str, Any]:
    output_dir = resolve(output_dir)
    data = {
        "samples": read_jsonl(output_dir / "independent_samples.jsonl"),
        "baseline": read_jsonl(output_dir / "baseline_results.jsonl"),
        "round1": read_jsonl(output_dir / "round1_results.jsonl"),
        "cached": read_jsonl(output_dir / "cached_dea_results.jsonl"),
        "gate": read_jsonl(output_dir / "safety_gate_v1_2_results.jsonl"),
        "model": read_jsonl(output_dir / "model_request_log.jsonl"),
        "run_manifest": read_json(output_dir / "run_manifest.json"),
    }
    expected = {norm(row.get("cluster_id")) for row in data["samples"]}
    for name in ("baseline", "round1", "cached", "gate", "model"):
        got = {norm(row.get("cluster_id")) for row in data[name]}
        if got != expected:
            raise ValueError(f"{name} cluster_id set does not match samples: missing={sorted(expected-got)[:5]}, extra={sorted(got-expected)[:5]}")
    return data


def build_workbook_payload(output_dir: Path) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any]]:
    data = load_inputs(output_dir)
    workbook_rows = {
        "Primary人工审核": build_main_rows(data, "primary"),
        "Contradiction压力审核": build_main_rows(data, "contradiction_stress"),
        "证据明细": build_evidence_rows(data),
        "模型结果对比": build_comparison_rows(data),
        "运行统计": flatten_manifest(data["run_manifest"]),
        "填写说明": guide_rows(),
    }
    counts = {
        "primary_count": len(workbook_rows["Primary人工审核"]),
        "contradiction_stress_count": len(workbook_rows["Contradiction压力审核"]),
        "total_count": len(workbook_rows["Primary人工审核"]) + len(workbook_rows["Contradiction压力审核"]),
        "evidence_rows": len(workbook_rows["证据明细"]),
        "human_fields_blank": all(
            not norm(row.get(field))
            for sheet in ("Primary人工审核", "Contradiction压力审核")
            for row in workbook_rows[sheet]
            for field in HUMAN_FIELDS
        ),
    }
    manifest = {
        "step": "4.4-H-A",
        "output_dir": str(resolve(output_dir)),
        "counts": counts,
        "input_manifest": data["run_manifest"],
        "formal_output_sha256_after_workbook": hash_files(FORMAL_OUTPUTS),
        "human_fields": HUMAN_FIELDS,
        "allowed_values": ALLOWED_VALUES,
        "ai_assistance_is_not_gold": True,
        "isr_started": False,
    }
    return workbook_rows, manifest


def render_guide(manifest: dict[str, Any]) -> str:
    lines = [
        "# Safety Gate v1.2 Independent Human Audit Guide",
        "",
        "本工作簿用于50条独立验证样本的人工确认。AI辅助建议仅供参考，不是金标。",
        "",
        "- Primary样本: 40条自然分布样本，用于后续主要泛化评估。",
        "- Contradiction压力样本: 10条否定关键词召回样本，单独报告，不与Primary混合声称自然分布准确率。",
        "- 本步骤未进入ISR，不更新正式260条复核和fulltext评分输出。",
        "",
        "## 人工字段枚举",
        "",
        "| Field | Allowed values |",
        "|---|---|",
    ]
    for field, values in ALLOWED_VALUES.items():
        lines.append(f"| {field} | {', '.join(values)} |")
    lines.extend(
        [
            "",
            "填写前这些字段必须为空：",
            ", ".join(HUMAN_FIELDS),
            "",
            f"样本计数: {json.dumps(manifest['counts'], ensure_ascii=False)}",
            "",
        ]
    )
    return "\n".join(lines)


def ensure_outputs(paths: list[Path], overwrite: bool) -> None:
    existing = [path for path in paths if path.exists()]
    if existing and not overwrite:
        raise FileExistsError(f"refusing to overwrite existing audit outputs: {[str(path) for path in existing]}")


def build_audit_outputs(args: argparse.Namespace) -> dict[str, Any]:
    workbook_path = resolve(args.workbook)
    guide_path = resolve(args.guide)
    manifest_path = resolve(args.manifest)
    ensure_outputs([workbook_path, guide_path, manifest_path], args.overwrite)
    rows, manifest = build_workbook_payload(args.output_dir)
    write_workbook(workbook_path, rows)
    guide_path.parent.mkdir(parents=True, exist_ok=True)
    guide_path.write_text(render_guide(manifest), encoding="utf-8")
    manifest.update(
        {
            "workbook": str(workbook_path),
            "guide": str(guide_path),
            "manifest": str(manifest_path),
            "workbook_sha256": sha256(workbook_path),
            "guide_sha256": sha256(guide_path),
        }
    )
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--guide", type=Path, default=DEFAULT_GUIDE)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    manifest = build_audit_outputs(args)
    print(
        json.dumps(
            {
                "workbook": manifest["workbook"],
                "guide": manifest["guide"],
                "manifest": manifest["manifest"],
                "counts": manifest["counts"],
                "human_fields_blank": manifest["counts"]["human_fields_blank"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
