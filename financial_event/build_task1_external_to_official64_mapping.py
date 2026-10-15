#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build explicit DocFEE/ChFinAnn -> official 64-class mapping audit."""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
REPORT_JSON = ROOT / "reports/task1_external_to_official64_mapping.json"
REPORT_MD = ROOT / "reports/task1_external_to_official64_mapping.md"
REPORT_XLSX = ROOT / "reports/task1_external_to_official64_mapping_review.xlsx"
CONFIG_YAML = ROOT / "configs/task1_external_to_official64_mapping_v1.yaml"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def load_official_schema() -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    rows = read_jsonl(ROOT / "data/processed/event_schema.jsonl")
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        by_name[str(row["event_name"])].append(row)
    return rows, by_name


def rep_candidate(by_name: dict[str, list[dict[str, Any]]], event_name: str, relation: str, confidence: float, reason: str) -> dict[str, Any]:
    rows = by_name.get(event_name, [])
    row = rows[0] if rows else {}
    return {
        "schema_id": row.get("schema_id"),
        "target_type": row.get("target_type"),
        "event_name": event_name,
        "relation": relation,
        "confidence": confidence,
        "mapping_reason": reason,
        "definition_evidence": str(row.get("raw_description", ""))[:800],
    }


def source_definitions() -> dict[tuple[str, str], dict[str, Any]]:
    docfee_schema = read_json(ROOT / "data/external/docfee_v3/schema.json")
    docfee_defs: dict[tuple[str, str], dict[str, Any]] = {}
    zh_to_en = {
        "破产清算": "Bankruptcy Liquidation",
        "重大安全事故": "Major Safety Incident",
        "股东减持": "Shareholder Reduction",
        "股权质押": "Equity Pledge",
        "股东增持": "Shareholder Increase",
        "股权冻结": "Equity Freeze",
        "重大对外赔付": "Major External Compensation",
        "重大资产损失": "Major Asset Loss",
        "高层死亡": "Senior Executive Death",
    }
    for zh, en in zh_to_en.items():
        props = ((docfee_schema.get("properties") or {}).get(en) or {}).get("properties") or {}
        docfee_defs[("docfee", zh)] = {
            "definition_source": "data/external/docfee_v3/schema.json",
            "source_definition": en,
            "roles": sorted(props.keys()),
            "subject_type": "listed company announcement",
            "core_behavior": zh,
            "time_semantics": [r for r in props if "Date" in r],
            "necessary_arguments": sorted(props.keys())[:],
        }
    event_type_py = (ROOT / "third_party/Doc2EDAG/dee/event_type.py").read_text(encoding="utf-8")
    ch_defs = {
        "EquityFreeze": {
            "roles": ["EquityHolder", "FrozeShares", "LegalInstitution", "TotalHoldingShares", "TotalHoldingRatio", "StartDate", "EndDate", "UnfrozeDate"],
            "key_arguments": ["EquityHolder", "FrozeShares", "LegalInstitution"],
            "core_behavior": "股权冻结",
        },
        "EquityRepurchase": {
            "roles": ["CompanyName", "HighestTradingPrice", "LowestTradingPrice", "RepurchasedShares", "ClosingDate", "RepurchaseAmount"],
            "key_arguments": ["CompanyName"],
            "core_behavior": "股份回购",
        },
        "EquityUnderweight": {
            "roles": ["EquityHolder", "TradedShares", "StartDate", "EndDate", "LaterHoldingShares", "AveragePrice"],
            "key_arguments": ["EquityHolder", "TradedShares"],
            "core_behavior": "股东减持",
        },
        "EquityOverweight": {
            "roles": ["EquityHolder", "TradedShares", "StartDate", "EndDate", "LaterHoldingShares", "AveragePrice"],
            "key_arguments": ["EquityHolder", "TradedShares"],
            "core_behavior": "股东增持",
        },
        "EquityPledge": {
            "roles": ["Pledger", "PledgedShares", "Pledgee", "TotalHoldingShares", "TotalHoldingRatio", "TotalPledgedShares", "StartDate", "EndDate", "ReleasedDate"],
            "key_arguments": ["Pledger", "PledgedShares", "Pledgee"],
            "core_behavior": "股权质押",
        },
    }
    out = dict(docfee_defs)
    for name, item in ch_defs.items():
        out[("chfinann", name)] = {
            "definition_source": "third_party/Doc2EDAG/dee/event_type.py",
            "source_definition": f"{name}: {item['core_behavior']}; key fields {', '.join(item['key_arguments'])}",
            "roles": item["roles"],
            "subject_type": "listed company announcement, distant supervision",
            "core_behavior": item["core_behavior"],
            "time_semantics": [r for r in item["roles"] if "Date" in r],
            "necessary_arguments": item["key_arguments"],
            "source_excerpt": event_type_py[event_type_py.find(f"class {name}Event") : event_type_py.find(f"class {name}Event") + 900] if f"class {name}Event" in event_type_py else "",
        }
    return out


def build_mapping(by_name: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    defs = source_definitions()
    specs = [
        ("docfee", "破产清算", [("现金流风险", "partial_overlap", 0.62, "Both are severe corporate risk events, but official 现金流风险 is defined around public debt repayment pressure; bankruptcy liquidation is a legal liquidation state and is not fully covered.")]),
        ("docfee", "股东减持", [("资本市场运作", "partial_overlap", 0.48, "Shareholder reduction is a capital-market related transaction, but official 资本市场运作 focuses on the target company's own capital operations such as restructuring/refinancing/equity incentives.")]),
        ("docfee", "股东增持", [("资本市场运作", "partial_overlap", 0.48, "Shareholder increase is related to equity market actions but is not listed in official examples and may concern shareholders rather than the target company's own operation.")]),
        ("docfee", "股权冻结", [("现金流风险", "partial_overlap", 0.42, "Equity freeze may signal legal/financial risk, but official 现金流风险 requires debt repayment stress."), ("资本市场运作", "partial_overlap", 0.32, "Equity legal status is not a company capital operation in the official definition.")]),
        ("docfee", "股权质押", [("现金流风险", "partial_overlap", 0.40, "Equity pledge can correlate with liquidity risk but does not itself satisfy debt payment default/extension conditions."), ("资本市场运作", "partial_overlap", 0.34, "Pledge is a shareholder financing/collateral action, not a listed company capital operation as defined.")]),
        ("docfee", "重大安全事故", []),
        ("docfee", "重大对外赔付", [("现金流风险", "partial_overlap", 0.28, "Large compensation can affect cash flow, but official event condition is specific debt repayment difficulty.")]),
        ("docfee", "重大资产损失", [("现金流风险", "partial_overlap", 0.30, "Asset loss may affect solvency but is broader than official debt repayment risk.")]),
        ("docfee", "高层死亡", []),
        ("chfinann", "EquityOverweight", [("资本市场运作", "partial_overlap", 0.45, "Shareholder increase overlaps capital-market behavior but is not the official target company's own restructuring/refinancing/incentive action.")]),
        ("chfinann", "EquityUnderweight", [("资本市场运作", "partial_overlap", 0.45, "Shareholder reduction overlaps capital-market behavior but official definition is different in subject and event condition.")]),
        ("chfinann", "EquityFreeze", [("现金流风险", "partial_overlap", 0.42, "Equity freeze can indicate risk but not necessarily official debt repayment stress."), ("资本市场运作", "partial_overlap", 0.32, "Legal freeze is not an official capital operation action.")]),
        ("chfinann", "EquityPledge", [("现金流风险", "partial_overlap", 0.40, "Equity pledge can indicate liquidity pressure but does not entail official debt repayment difficulty."), ("资本市场运作", "partial_overlap", 0.34, "Pledge is collateral financing, not a listed company capital operation per official definition.")]),
        ("chfinann", "EquityRepurchase", [("资本市场运作", "partial_overlap", 0.58, "Share repurchase is a company capital market action, but official definition/examples do not explicitly include repurchase and target-specific schema is for securities company capital operations.")]),
    ]
    rows = []
    for source, event_type, candidates in specs:
        src_def = defs[(source, event_type)]
        cand_rows = [rep_candidate(by_name, name, rel, conf, reason) for name, rel, conf, reason in candidates]
        relation = "none" if not cand_rows else "partial_overlap"
        rows.append(
            {
                "source_dataset": source,
                "source_event_type": event_type,
                "source_definition": src_def,
                "candidate_official_events": cand_rows,
                "selected_schema_id": None,
                "selected_event_name": None,
                "usable_for_official_classification": False,
                "usable_for_shared_extraction_pretraining": True,
                "requires_review": True,
                "mapping_relation": relation,
                "classification_loss_weight": 0.0,
                "mapping_status": "extraction_only" if relation in {"partial_overlap", "none"} else "review_required",
            }
        )
    return rows


def count_external_samples() -> dict[str, Any]:
    docfee_train = sum(1 for _ in (ROOT / "data/external/docfee_v3/train.jsonl").open(encoding="utf-8"))
    docfee_test = sum(1 for _ in (ROOT / "data/external/docfee_v3/test.jsonl").open(encoding="utf-8"))
    ch_counts = {}
    for split in ["train", "dev", "test", "sample_train"]:
        data = read_json(ROOT / f"data/external/chfinann/Data/{split}.json")
        ch_counts[split] = len(data)
    return {
        "docfee": {"train": docfee_train, "test_reserved": docfee_test},
        "chfinann": ch_counts,
        "classification_usable_train_samples": 0,
        "extraction_only_train_samples": docfee_train + ch_counts["train"],
        "rejected_train_samples": 0,
        "reserved_external_test_samples": docfee_test + ch_counts["test"],
    }


def relation_stats(rows: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    out: dict[str, Counter[str]] = defaultdict(Counter)
    for row in rows:
        relation = row["mapping_relation"]
        out[row["source_dataset"]][relation] += 1
        exact = any(c["relation"] == "exact_equivalent" for c in row["candidate_official_events"])
        narrower = any(c["relation"] == "source_narrower" for c in row["candidate_official_events"])
        if exact:
            out[row["source_dataset"]]["exact_equivalent"] += 1
        if narrower:
            out[row["source_dataset"]]["source_narrower"] += 1
        if relation == "partial_overlap":
            out[row["source_dataset"]]["broader_or_partial"] += 1
        if relation == "none":
            out[row["source_dataset"]]["unmapped_none"] += 1
    return {k: dict(v) for k, v in out.items()}


def write_yaml(rows: list[dict[str, Any]]) -> None:
    CONFIG_YAML.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "version: task1_external_to_official64_mapping_v1",
        "policy:",
        "  exact_equivalent_min_confidence: 0.90",
        "  source_narrower_min_confidence: 0.90",
        "  source_narrower_max_classification_loss_weight: 0.5",
        "  partial_overlap_classification_loss_weight: 0.0",
        "  external_test_used_for_training: false",
        "mappings:",
    ]
    for row in rows:
        lines.append(f"  - source_dataset: {row['source_dataset']}")
        lines.append(f"    source_event_type: {row['source_event_type']}")
        lines.append(f"    selected_schema_id: {row['selected_schema_id'] if row['selected_schema_id'] else 'null'}")
        lines.append(f"    selected_event_name: {row['selected_event_name'] if row['selected_event_name'] else 'null'}")
        lines.append(f"    mapping_relation: {row['mapping_relation']}")
        lines.append(f"    usable_for_official_classification: {str(row['usable_for_official_classification']).lower()}")
        lines.append(f"    usable_for_shared_extraction_pretraining: {str(row['usable_for_shared_extraction_pretraining']).lower()}")
        lines.append(f"    classification_loss_weight: {row['classification_loss_weight']}")
        lines.append("    candidate_official_events:")
        if not row["candidate_official_events"]:
            lines.append("      []")
        for cand in row["candidate_official_events"]:
            lines.append(f"      - schema_id: {cand['schema_id']}")
            lines.append(f"        event_name: {cand['event_name']}")
            lines.append(f"        relation: {cand['relation']}")
            lines.append(f"        confidence: {cand['confidence']}")
            lines.append(f"        mapping_reason: {json.dumps(cand['mapping_reason'], ensure_ascii=False)}")
    CONFIG_YAML.write_text("\n".join(lines) + "\n", encoding="utf-8")


def format_xlsx(path: Path) -> None:
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for idx in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(idx)].width = min(60, max(14, len(str(ws.cell(1, idx).value or "")) + 4))
    wb.save(path)


def main() -> None:
    _, by_name = load_official_schema()
    rows = build_mapping(by_name)
    sample_counts = count_external_samples()
    stats = relation_stats(rows)
    target_type_counts = Counter()
    for row in rows:
        for cand in row["candidate_official_events"]:
            if cand.get("target_type"):
                target_type_counts[cand["target_type"]] += 1
    conflicts = {
        "one_to_many_candidate_source_types": sum(1 for r in rows if len(r["candidate_official_events"]) > 1),
        "selected_multi_source_same_official": 0,
        "hard_label_selected_count": sum(1 for r in rows if r["selected_event_name"]),
    }
    report = {
        "generated_at": datetime.now().replace(microsecond=0).isoformat(),
        "status": "external_dataset_schema_incompatible",
        "reason": "No external event type satisfies exact_equivalent/source_narrower with confidence >=0.90 against official 64 definitions. All 14 types are extraction_only or none for official classification.",
        "official64_event_count": len(by_name),
        "mapping_policy": {
            "exact_equivalent": "requires bidirectional definition coverage and confidence >=0.90",
            "source_narrower": "allowed only if source is clear subclass without official-class conflict and confidence >=0.90; classification_loss_weight<=0.5",
            "partial_overlap_or_none": "no hard official label; extraction-only pretraining allowed",
        },
        "mappings": rows,
        "relation_stats": stats,
        "mapped_to_company_official_candidates": target_type_counts.get("company", 0),
        "mapped_to_industry_official_candidates": target_type_counts.get("industry", 0),
        "sample_usage_counts": sample_counts,
        "conflict_audit": conflicts,
        "ai_mapping_used": False,
        "training_allowed": False,
        "next_required_step": "Manual review of mapping workbook or official clarification. Do not start full transfer training until at least one exact/source_narrower mapping is approved or training is limited to extraction-only pretraining.",
    }
    REPORT_JSON.parent.mkdir(parents=True, exist_ok=True)
    REPORT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_yaml(rows)

    flat_rows = []
    for row in rows:
        candidates = row["candidate_official_events"] or [{"schema_id": None, "event_name": None, "relation": "none", "confidence": 0.0, "mapping_reason": "No candidate official event.", "definition_evidence": ""}]
        for cand in candidates:
            flat_rows.append(
                {
                    "source_dataset": row["source_dataset"],
                    "source_event_type": row["source_event_type"],
                    "source_core_behavior": row["source_definition"]["core_behavior"],
                    "source_roles": " | ".join(row["source_definition"]["roles"]),
                    "candidate_schema_id": cand.get("schema_id"),
                    "candidate_event_name": cand.get("event_name"),
                    "relation": cand.get("relation"),
                    "confidence": cand.get("confidence"),
                    "mapping_reason": cand.get("mapping_reason"),
                    "definition_evidence": cand.get("definition_evidence"),
                    "selected_schema_id": row["selected_schema_id"],
                    "selected_event_name": row["selected_event_name"],
                    "usable_for_official_classification": row["usable_for_official_classification"],
                    "usable_for_shared_extraction_pretraining": row["usable_for_shared_extraction_pretraining"],
                    "classification_loss_weight": row["classification_loss_weight"],
                    "requires_review": row["requires_review"],
                    "mapping_status": row["mapping_status"],
                }
            )
    with pd.ExcelWriter(REPORT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(flat_rows).to_excel(writer, sheet_name="mapping_review", index=False)
        pd.DataFrame([{"source_dataset": k, **v} for k, v in stats.items()]).to_excel(writer, sheet_name="relation_stats", index=False)
        pd.DataFrame([sample_counts]).to_excel(writer, sheet_name="sample_usage", index=False)
    format_xlsx(REPORT_XLSX)

    md = [
        "# Task1 External to Official64 Mapping",
        "",
        f"- Status: `{report['status']}`",
        f"- Official 64 event_name count: `{report['official64_event_count']}`",
        f"- AI mapping used: `{report['ai_mapping_used']}`",
        f"- Training allowed: `{report['training_allowed']}`",
        "",
        "## Summary",
        report["reason"],
        "",
        "## Relation Counts",
    ]
    for dataset, values in stats.items():
        md.append(f"- `{dataset}`: {values}")
    md.extend(
        [
            "",
            "## Sample Usage",
            f"- Classification-usable train samples: `{sample_counts['classification_usable_train_samples']}`",
            f"- Extraction-only train samples: `{sample_counts['extraction_only_train_samples']}`",
            f"- Rejected train samples: `{sample_counts['rejected_train_samples']}`",
            "",
            "## Per-Type Mapping",
        ]
    )
    for row in rows:
        md.append(f"### {row['source_dataset']} / {row['source_event_type']}")
        md.append(f"- Source behavior: {row['source_definition']['core_behavior']}")
        md.append(f"- Roles: {', '.join(row['source_definition']['roles'])}")
        if row["candidate_official_events"]:
            for cand in row["candidate_official_events"]:
                md.append(f"- Candidate `{cand['event_name']}` (`{cand['schema_id']}`): `{cand['relation']}`, confidence `{cand['confidence']}`. {cand['mapping_reason']}")
        else:
            md.append("- Candidate: none")
        md.append("- Selected official label: `null`; usable_for_official_classification=`false`; usable_for_shared_extraction_pretraining=`true`.")
    md.append("")
    md.append("## Decision")
    md.append(report["next_required_step"])
    REPORT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "json": str(REPORT_JSON), "xlsx": str(REPORT_XLSX), "yaml": str(CONFIG_YAML)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
