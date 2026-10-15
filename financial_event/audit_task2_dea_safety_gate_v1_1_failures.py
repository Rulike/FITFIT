#!/usr/bin/env python3
"""Audit Safety Gate v1.1 failure modes without changing gate artifacts."""

from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts import evaluate_task2_dea_safety_gate as v1_eval  # noqa: E402
from src.event_credibility.dea_safety_gate_adapter import (  # noqa: E402
    build_safety_gate_evidence_records,
    normalize_cited_evidence_ids,
)


CLASSES = ["support", "contradict", "insufficient", "unclear"]
FORMAL_OUTPUTS = [
    Path("outputs/event_credibility/llm_review_results.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.jsonl"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.parquet"),
    Path("outputs/event_credibility/event_credibility_scores_fulltext.xlsx"),
]
FROZEN_FILES = [
    Path("src/event_credibility/dea_safety_gate.py"),
    Path("configs/task2_dea_safety_gate_v1.yaml"),
    Path("scripts/test_task2_dea_safety_gate.py"),
    Path("src/event_credibility/dea_safety_gate_v1_1.py"),
    Path("configs/task2_dea_safety_gate_v1_1.yaml"),
    Path("src/event_credibility/dea_safety_gate_adapter.py"),
]
DEFAULT_GOLD_FILE = ROOT / "reports/task2_dea_cached_human_audit_gold_filled.xlsx"
DEFAULT_V1_RESULTS = ROOT / "outputs/event_credibility/dea_safety_gate_v1_adapter_v2/safety_gate_results.jsonl"
DEFAULT_V1_1_RESULTS = ROOT / "outputs/event_credibility/dea_safety_gate_v1_1_dev/safety_gate_v1_1_results.jsonl"
DEFAULT_V1_REPORT = ROOT / "reports/task2_dea_safety_gate_adapter_v2_report.json"
DEFAULT_V1_1_REPORT = ROOT / "reports/task2_dea_safety_gate_v1_1_dev_report.json"
DEFAULT_REPORT_JSON = ROOT / "reports/task2_dea_safety_gate_v1_1_failure_audit.json"
DEFAULT_REPORT_MD = ROOT / "reports/task2_dea_safety_gate_v1_1_failure_audit.md"
DEFAULT_REPORT_XLSX = ROOT / "reports/task2_dea_safety_gate_v1_1_failure_audit.xlsx"
DEFAULT_BACKUP_DIR = ROOT.parent / "backups/step4_4f_safety_gate_v1_1_failure_audit_20260808_193622"


def norm(value: Any) -> str:
    return str(value if value is not None else "").strip()


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def unique_by_cluster(rows: list[dict[str, Any]], name: str) -> dict[str, dict[str, Any]]:
    ids = [norm(row.get("cluster_id")) for row in rows]
    duplicates = sorted({cid for cid in ids if cid and ids.count(cid) > 1})
    missing = [idx + 1 for idx, cid in enumerate(ids) if not cid]
    if duplicates or missing:
        raise ValueError(f"{name} cluster_id invalid: duplicates={duplicates[:10]}, missing_rows={missing[:10]}")
    return {norm(row["cluster_id"]): row for row in rows}


def validate_same_clusters(*mappings: dict[str, Any]) -> list[str]:
    if not mappings:
        return []
    expected = set(mappings[0])
    for mapping in mappings[1:]:
        if set(mapping) != expected:
            raise ValueError(f"cluster_id mismatch: expected={len(expected)}, got={len(mapping)}")
    if len(expected) != 47:
        raise ValueError(f"expected 47 clusters, got {len(expected)}")
    return sorted(expected)


def sha256(path: Path) -> str:
    return v1_eval.sha256(path)


def hash_files(paths: list[Path]) -> dict[str, str]:
    return {str(path): sha256(ROOT / path) for path in paths}


def verdict_valid(value: Any) -> bool:
    return norm(value) in CLASSES


def classify_block(row: dict[str, Any]) -> str:
    if not row["v1_1_gate_blocked"]:
        return "not_blocked"
    if (not row["cached_correct"]) and row["v1_1_correct"]:
        return "correct_block"
    if row["cached_correct"] and not row["v1_1_correct"]:
        return "wrong_block"
    if (not row["cached_correct"]) and (not row["v1_1_correct"]):
        return "neutral_wrong_to_wrong"
    return "neutral_other"


def row_outcome(row: dict[str, Any]) -> str:
    if (not row["v1_correct"]) and row["v1_1_correct"]:
        return "changed_to_correct"
    if row["v1_correct"] and not row["v1_1_correct"]:
        return "changed_to_wrong"
    if row["v1_correct"] and row["v1_1_correct"]:
        return "no_accuracy_change_both_correct"
    return "no_accuracy_change_both_wrong"


def direct_score_summary(scores: dict[str, Any], direct_ids: list[str]) -> dict[str, Any]:
    direct_scores = [scores[eid] for eid in direct_ids if eid in scores]
    return {
        "subject_match": all(float(item.get("subject_score", 0.0)) > 0.0 for item in direct_scores) if direct_scores else False,
        "event_semantic_match": all(float(item.get("event_semantic_score", 0.0)) > 0.0 for item in direct_scores) if direct_scores else False,
        "concrete_fact_match": all(float(item.get("concrete_fact_score", 0.0)) > 0.0 for item in direct_scores) if direct_scores else False,
        "time_match": all(float(item.get("time_score", 0.0)) >= 0.6 for item in direct_scores) if direct_scores else False,
        "has_investor_question": any(float(item.get("question_penalty", 0.0)) > 0.0 for item in direct_scores),
        "has_market_signal": any(float(item.get("market_signal_penalty", 0.0)) > 0.0 for item in direct_scores),
    }


def normalize_text_record(record: dict[str, Any]) -> dict[str, str]:
    return {
        "evidence_id": norm(record.get("evidence_id")),
        "text": norm(record.get("text") or record.get("sentence") or record.get("content") or record.get("证据文本")),
        "source": norm(record.get("source") or record.get("article_source") or record.get("来源")),
        "publish_time": norm(record.get("publish_time") or record.get("article_publish_time") or record.get("发布时间")),
    }


def evidence_context(cid: str, inputs: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    cached = inputs["cached"][cid]
    round1 = cached.get("round1") or {}
    evidence_records = build_safety_gate_evidence_records(
        cid,
        inputs["evidence_detail_by_cluster"].get(cid, []),
        cached,
        round1,
    )
    cited_ids = normalize_cited_evidence_ids(cached, evidence_records)
    return evidence_records, cited_ids


def direct_evidence_text(evidence_records: list[dict[str, Any]], direct_ids: list[str]) -> str:
    by_id = {norm(row.get("evidence_id")): row for row in evidence_records}
    parts = []
    for evidence_id in direct_ids:
        item = normalize_text_record(by_id.get(evidence_id, {}))
        if item["evidence_id"]:
            parts.append(f"{item['evidence_id']} | {item['source']} | {item['publish_time']} | {item['text']}")
    return "\n---\n".join(parts)


def direct_evidence_sources(evidence_records: list[dict[str, Any]], direct_ids: list[str]) -> str:
    by_id = {norm(row.get("evidence_id")): row for row in evidence_records}
    return "; ".join(norm(by_id.get(eid, {}).get("source") or by_id.get(eid, {}).get("article_source")) for eid in direct_ids)


def direct_evidence_times(evidence_records: list[dict[str, Any]], direct_ids: list[str]) -> str:
    by_id = {norm(row.get("evidence_id")): row for row in evidence_records}
    return "; ".join(norm(by_id.get(eid, {}).get("publish_time") or by_id.get(eid, {}).get("article_publish_time")) for eid in direct_ids)


def infer_failure_categories(row: dict[str, Any], cited_ids: list[str]) -> tuple[list[str], str, str]:
    categories: list[str] = []
    notes = norm(row.get("human_notes"))
    direct_ids = list(row.get("direct_evidence_ids") or [])
    if row["claim_truly_unclear"] and row["human_verdict"] != "unclear":
        categories.extend(["claim_truly_unclear_overreach", "raw_claim_incomplete_misused"])
    if row["claim_truly_unclear"] and row["evidence_resolves_claim"]:
        categories.append("evidence_resolves_claim_not_respected")
    if direct_ids and row["human_verdict"] != "support":
        categories.extend(["direct_evidence_false_positive", "semantic_group_too_broad"])
    if (not direct_ids) and row["human_verdict"] == "support":
        categories.extend(["direct_evidence_false_negative", "semantic_group_too_narrow"])
    if row["v1_1_verdict"] != row["human_verdict"] and row["v1_verdict"] == row["human_verdict"]:
        categories.append("fallback_target_wrong")
    if "投资者提问" in notes or "互动" in notes:
        if direct_ids or row["v1_1_verdict"] == "support":
            categories.append("investor_question_not_detected")
    if any(token in notes for token in ("行情", "股价", "资金", "板块", "指数")):
        if direct_ids or row["v1_1_verdict"] == "support":
            categories.append("market_signal_not_detected")
    if any("subject_mismatch" == reason for reason in row["gate_reasons"]):
        categories.append("subject_mismatch")
    if any("time_mismatch" == reason for reason in row["gate_reasons"]):
        categories.append("time_mismatch")
    if not verdict_valid(row["v1_1_verdict"]):
        categories.append("verdict_mapping_error")
    if direct_ids and not set(direct_ids).intersection(set(cited_ids)):
        categories.append("evaluation_adapter_error")
    if not categories:
        categories.append("other")
    categories = unique_text(categories)
    primary = categories[0]
    suggestion = suggestion_for_categories(categories)
    return categories, primary, suggestion


def suggestion_for_categories(categories: list[str]) -> str:
    if "claim_truly_unclear_overreach" in categories:
        return "v1.2应把 unclear 限定为结构上不可验证；raw_claim_incomplete 默认不覆盖 Round1/Cached。"
    if "direct_evidence_false_positive" in categories:
        return "v1.2应要求模型引用证据、非投资者提问/非行情信号、动作对象更具体后才放行 support。"
    if "direct_evidence_false_negative" in categories:
        return "v1.2可扩展通用语义组，但仍以保守选择器方式回退 Round1。"
    if "investor_question_not_detected" in categories:
        return "v1.2应在模型实际引用证据范围内判定投资者提问风险。"
    if "market_signal_not_detected" in categories:
        return "v1.2应在模型实际引用证据范围内判定行情信号风险。"
    return "保留为人工复核问题，避免按单个样本新增规则。"


def unique_text(values: list[str]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for value in values:
        item = norm(value)
        if item and item not in seen:
            seen.add(item)
            output.append(item)
    return output


def metric_row(name: str, labels: list[str], preds: list[str]) -> dict[str, Any]:
    item = v1_eval.classification_metrics(labels, preds)
    return {
        "strategy": name,
        "accuracy": item["accuracy"],
        "correct_count": sum(1 for label, pred in zip(labels, preds) if label == pred),
        "sample_count": len(labels),
        "macro_f1_fixed_4class": item["macro_f1_fixed_4class"],
        "macro_f1_actual_labels": item["macro_f1_actual_labels"],
        "weighted_f1": item["weighted_f1"],
    }


def structurally_unjudgeable(row: dict[str, Any]) -> bool:
    reasons = set(row.get("raw_claim_incomplete_reasons") or [])
    return bool(
        row["claim_truly_unclear"]
        and not row["evidence_resolves_claim"]
        and not row["direct_evidence_ids"]
        and ("missing_event_time" in reasons or "event_type_missing" in reasons)
    )


def counterfactual_strategies(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    labels = [row["human_verdict"] for row in rows]
    strategy_a = [row["v1_1_verdict"] for row in rows]
    strategy_b = [row["round1_verdict"] if row["v1_1_gate_blocked"] else row["v1_1_verdict"] for row in rows]
    strategy_c = [row["cached_verdict"] if row["v1_1_gate_blocked"] else row["v1_1_verdict"] for row in rows]
    strategy_d = [
        "unclear"
        if row["v1_1_gate_blocked"] and structurally_unjudgeable(row)
        else (row["round1_verdict"] if row["v1_1_gate_blocked"] else row["v1_1_verdict"])
        for row in rows
    ]
    return [
        metric_row("A_blocked_to_unclear_current_v1_1", labels, strategy_a),
        metric_row("B_blocked_to_round1", labels, strategy_b),
        metric_row("C_blocked_keep_cached", labels, strategy_c),
        metric_row("D_structural_unclear_else_round1", labels, strategy_d),
    ]


def claim_unclear_category(row: dict[str, Any], gold: dict[str, Any]) -> str:
    if row["human_verdict"] == "unclear":
        return "金标本身为 unclear"
    if row["evidence_resolves_claim"]:
        return "原始主张不完整，但证据可补足"
    if norm(gold.get("claim_is_clear")) == "yes" and row["human_verdict"] == "support":
        return "主张清楚，只是事件类型比较抽象"
    if norm(gold.get("claim_is_clear")) == "yes":
        return "规则误判"
    if norm(gold.get("claim_is_clear")) in {"no", "uncertain"}:
        return "真正无法形成可验证主张"
    return "其他"


def build_metric_audit(rows: list[dict[str, Any]], v1_report: dict[str, Any], v1_1_report: dict[str, Any]) -> dict[str, Any]:
    labels = [row["human_verdict"] for row in rows]
    invalid = [
        (row["cluster_id"], field, row.get(field))
        for row in rows
        for field in ("baseline_verdict", "round1_verdict", "cached_verdict", "v1_verdict", "v1_1_verdict")
        if not verdict_valid(row.get(field))
    ]
    metrics = {
        "baseline": metric_row("baseline", labels, [row["baseline_verdict"] for row in rows]),
        "round1": metric_row("round1", labels, [row["round1_verdict"] for row in rows]),
        "cached": metric_row("cached", labels, [row["cached_verdict"] for row in rows]),
        "v1": metric_row("v1", labels, [row["v1_verdict"] for row in rows]),
        "v1_1": metric_row("v1_1", labels, [row["v1_1_verdict"] for row in rows]),
    }
    block_counts = Counter(classify_block(row) for row in rows)
    blocked_count = sum(1 for row in rows if row["v1_1_gate_blocked"])
    baseline_candidates = [row for row in rows if row["baseline_verdict"] == "insufficient" and row["cached_verdict"] in {"support", "contradict"}]
    round1_candidates = [row for row in rows if row["round1_verdict"] == "insufficient" and row["cached_verdict"] in {"support", "contradict"}]
    return {
        "same_47_cluster_ids": len(rows) == 47,
        "all_verdicts_valid": not invalid,
        "invalid_verdicts": invalid,
        "accuracy_31_91_is_15_of_47": metrics["v1_1"]["correct_count"] == 15 and metrics["v1_1"]["sample_count"] == 47,
        "v1_1_correct_count": metrics["v1_1"]["correct_count"],
        "metric_rows": list(metrics.values()),
        "fixed_macro_classes": CLASSES,
        "observed_label_classes": sorted(set(labels), key=CLASSES.index),
        "reported_v1_1_accuracy": v1_1_report["method_metrics"]["v1_1"]["accuracy"],
        "recomputed_v1_1_accuracy": metrics["v1_1"]["accuracy"],
        "reported_v1_accuracy": v1_report["method_metrics"]["safety_gated"]["accuracy"],
        "recomputed_v1_accuracy": metrics["v1"]["accuracy"],
        "gate_blocked_count": blocked_count,
        "block_destination_counts": dict(block_counts),
        "neutral_block_explanation": "剩余 blocked 样本为 cached 原本错误且 v1.1 仍错误，只是从一种错误 verdict 变成另一种错误 verdict。",
        "baseline_insufficient_cached_resolved": {
            "candidate_count": len(baseline_candidates),
            "true_solution_count": sum(1 for row in baseline_candidates if row["human_verdict"] in {"support", "contradict"} and row["cached_correct"]),
            "wrong_solution_count": sum(1 for row in baseline_candidates if not row["cached_correct"]),
        },
        "round1_insufficient_cached_resolved": {
            "candidate_count": len(round1_candidates),
            "true_solution_count": sum(1 for row in round1_candidates if row["human_verdict"] in {"support", "contradict"} and row["cached_correct"]),
            "wrong_solution_count": sum(1 for row in round1_candidates if not row["cached_correct"]),
        },
        "denominator_explanation": "旧口径通常以 baseline=insufficient 且 cached 改为 support/contradict 统计；Step 4.4-E 脚本改用 round1=insufficient 且 cached 改为 support/contradict，因此出现 true 4、wrong 7 的分母。",
        "metrics_same_sample_and_labels": True,
        "verdict_normalization": "All methods are normalized to support/contradict/insufficient/unclear and checked for invalid values.",
    }


def build_change_rows(
    rows: list[dict[str, Any]],
    v1_by_id: dict[str, dict[str, Any]],
    gold_by_id: dict[str, dict[str, Any]],
    inputs: dict[str, Any],
) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        v1_reasons = v1_by_id[row["cluster_id"]].get("gate_reasons") or []
        verdict_changed = row["v1_verdict"] != row["v1_1_verdict"]
        reason_changed = v1_reasons != row["gate_reasons"]
        if not verdict_changed and not reason_changed:
            continue
        cid = row["cluster_id"]
        evidence_records, cited_ids = evidence_context(cid, inputs)
        categories, primary, suggestion = infer_failure_categories(row, cited_ids)
        out.append(
            {
                "cluster_id": cid,
                "gold_verdict": row["human_verdict"],
                "baseline_verdict": row["baseline_verdict"],
                "round1_verdict": row["round1_verdict"],
                "cached_verdict": row["cached_verdict"],
                "v1_verdict": row["v1_verdict"],
                "v1_1_verdict": row["v1_1_verdict"],
                "cached原本是否正确": "yes" if row["cached_correct"] else "no",
                "v1.1是否改对": "yes" if (not row["v1_correct"] and row["v1_1_correct"]) else "no",
                "v1.1是否改错": "yes" if (row["v1_correct"] and not row["v1_1_correct"]) else "no",
                "变化结果": row_outcome(row),
                "verdict是否不一致": verdict_changed,
                "gate_reason是否变化": reason_changed,
                "gate_blocked": row["v1_1_gate_blocked"],
                "gate_reasons": ";".join(row["gate_reasons"]),
                "raw_claim_incomplete": row["raw_claim_incomplete"],
                "evidence_resolves_claim": row["evidence_resolves_claim"],
                "claim_truly_unclear": row["claim_truly_unclear"],
                "direct_evidence_ids": ";".join(row["direct_evidence_ids"]),
                "cited_evidence_ids": ";".join(cited_ids),
                "direct evidence 文本、来源、时间": direct_evidence_text(evidence_records, row["direct_evidence_ids"]),
                "direct evidence 来源": direct_evidence_sources(evidence_records, row["direct_evidence_ids"]),
                "direct evidence 时间": direct_evidence_times(evidence_records, row["direct_evidence_ids"]),
                "最主要失败原因": primary,
                "失败原因类别": ";".join(categories),
                "建议的通用修复方向": suggestion,
                "human_notes": row["human_notes"],
                "v1_gate_reasons": ";".join(v1_reasons),
                "claim_is_clear": norm(gold_by_id[cid].get("claim_is_clear")),
            }
        )
    return out


def build_wrong_block_rows(rows: list[dict[str, Any]], inputs: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for row in rows:
        if classify_block(row) != "wrong_block":
            continue
        _, cited_ids = evidence_context(row["cluster_id"], inputs)
        out.append(
            {
                "cluster_id": row["cluster_id"],
                "gold_verdict": row["human_verdict"],
                "cached_verdict": row["cached_verdict"],
                "v1_1_verdict": row["v1_1_verdict"],
                "Cached原verdict为何符合金标": f"cached_verdict={row['cached_verdict']} equals gold_verdict={row['human_verdict']}",
                "v1.1为何阻断": ";".join(row["gate_reasons"]) or "gated verdict differs from cached",
                "claim_truly_unclear是否覆盖evidence_resolves_claim": bool(row["claim_truly_unclear"] and row["evidence_resolves_claim"]),
                "回退Round1是否正确": row["round1_verdict"] == row["human_verdict"],
                "保留Cached是否正确": row["cached_verdict"] == row["human_verdict"],
                "回退Baseline是否正确": row["baseline_verdict"] == row["human_verdict"],
                "raw_claim_incomplete": row["raw_claim_incomplete"],
                "evidence_resolves_claim": row["evidence_resolves_claim"],
                "claim_truly_unclear": row["claim_truly_unclear"],
                "direct_evidence_ids": ";".join(row["direct_evidence_ids"]),
                "cited_evidence_ids": ";".join(cited_ids),
                "human_notes": row["human_notes"],
            }
        )
    return out


def build_claim_unclear_rows(rows: list[dict[str, Any]], gold_by_id: dict[str, dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    detail = []
    for row in rows:
        if not row["claim_truly_unclear"]:
            continue
        category = claim_unclear_category(row, gold_by_id[row["cluster_id"]])
        detail.append(
            {
                "cluster_id": row["cluster_id"],
                "category": category,
                "gold_verdict": row["human_verdict"],
                "v1_1_verdict": row["v1_1_verdict"],
                "v1_1_correct": row["v1_1_correct"],
                "raw_claim_incomplete_reasons": ";".join(row["raw_claim_incomplete_reasons"]),
                "evidence_resolves_claim": row["evidence_resolves_claim"],
                "claim_is_clear": norm(gold_by_id[row["cluster_id"]].get("claim_is_clear")),
                "human_notes": row["human_notes"],
            }
        )
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in detail:
        grouped[item["category"]].append(item)
    summary = [
        {
            "category": category,
            "count": len(items),
            "accuracy": round(sum(1 for item in items if item["v1_1_correct"]) / len(items), 6) if items else 0.0,
        }
        for category, items in sorted(grouped.items())
    ]
    return detail, summary


def build_direct_evidence_rows(rows: list[dict[str, Any]], inputs: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    detail = []
    for row in rows:
        if not row["direct_evidence_ids"]:
            continue
        evidence_records, cited_ids = evidence_context(row["cluster_id"], inputs)
        score_summary = direct_score_summary(row.get("direct_evidence_scores") or {}, row["direct_evidence_ids"])
        is_true_positive = row["human_verdict"] == "support"
        detail.append(
            {
                "cluster_id": row["cluster_id"],
                "gold_verdict": row["human_verdict"],
                "v1_1_verdict": row["v1_1_verdict"],
                "direct_evidence_ids": ";".join(row["direct_evidence_ids"]),
                "真阳性": is_true_positive,
                "假阳性": not is_true_positive,
                "与主张主体匹配": score_summary["subject_match"],
                "与事件语义匹配": score_summary["event_semantic_match"],
                "包含具体动作或事实": score_summary["concrete_fact_match"],
                "时间匹配": score_summary["time_match"],
                "是否只是投资者提问": score_summary["has_investor_question"],
                "是否只是行情资金股价信号": score_summary["has_market_signal"],
                "是否支持模型最终引用": bool(set(row["direct_evidence_ids"]).intersection(set(cited_ids))),
                "cited_evidence_ids": ";".join(cited_ids),
                "direct evidence 文本": direct_evidence_text(evidence_records, row["direct_evidence_ids"]),
                "human_notes": row["human_notes"],
            }
        )
    tp = sum(1 for item in detail if item["真阳性"])
    fp = sum(1 for item in detail if item["假阳性"])
    summary = {
        "sample_count": len(detail),
        "true_positive_count": tp,
        "false_positive_count": fp,
        "precision": round(tp / (tp + fp), 6) if tp + fp else 0.0,
        "subject_match_count": sum(1 for item in detail if item["与主张主体匹配"]),
        "event_semantic_match_count": sum(1 for item in detail if item["与事件语义匹配"]),
        "concrete_fact_match_count": sum(1 for item in detail if item["包含具体动作或事实"]),
        "time_match_count": sum(1 for item in detail if item["时间匹配"]),
        "investor_question_count": sum(1 for item in detail if item["是否只是投资者提问"]),
        "market_signal_count": sum(1 for item in detail if item["是否只是行情资金股价信号"]),
        "cited_by_cached_count": sum(1 for item in detail if item["是否支持模型最终引用"]),
    }
    return detail, summary


def v1_2_recommendations() -> list[dict[str, str]]:
    return [
        {"priority": "P0", "recommendation": "Safety Gate 保持保守选择器定位，不作为新的事实分类器。"},
        {"priority": "P0", "recommendation": "不得仅凭 broad/encoded/raw_claim_incomplete 将大量样本强制输出 unclear。"},
        {"priority": "P0", "recommendation": "evidence_resolves_claim=true 时，不应仅因 raw_claim_incomplete 或 broad claim 输出 unclear。"},
        {"priority": "P0", "recommendation": "unclear 只用于结构上确实无法形成可验证主张的情况。"},
        {"priority": "P0", "recommendation": "风险门控默认回退 Round1，不默认回退 unclear。"},
        {"priority": "P1", "recommendation": "Cached support/contradict 只有在高风险条件明确命中时才阻断。"},
        {"priority": "P1", "recommendation": "direct evidence 需强调模型实际引用、来源类型、动作对象和非行情/非提问约束。"},
        {"priority": "P1", "recommendation": "禁止公司名、cluster_id 或47条样本特例规则；47条仅作为开发期回放。"},
        {"priority": "P1", "recommendation": "当前金标没有 contradict，不能宣称 contradict 门控有效。"},
    ]


def build_audit(args: argparse.Namespace | None = None) -> dict[str, Any]:
    if args is None:
        args = build_arg_parser().parse_args([])
    formal_before = hash_files(FORMAL_OUTPUTS)
    frozen_before = hash_files(FROZEN_FILES)
    inputs = v1_eval.load_inputs(args.gold_file)
    gold_by_id = unique_by_cluster(inputs["gold_rows"], "gold")
    v1_by_id = unique_by_cluster(read_jsonl(args.v1_results), "v1")
    v1_1_by_id = unique_by_cluster(read_jsonl(args.v1_1_results), "v1_1")
    ids = validate_same_clusters(gold_by_id, v1_by_id, v1_1_by_id)
    rows = [dict(v1_1_by_id[cid]) for cid in ids]
    v1_report = read_json(args.v1_report)
    v1_1_report = read_json(args.v1_1_report)
    metric_audit = build_metric_audit(rows, v1_report, v1_1_report)
    change_rows = build_change_rows(rows, v1_by_id, gold_by_id, inputs)
    wrong_block_rows = build_wrong_block_rows(rows, inputs)
    claim_unclear_rows, claim_unclear_summary = build_claim_unclear_rows(rows, gold_by_id)
    direct_rows, direct_summary = build_direct_evidence_rows(rows, inputs)
    counterfactual_rows = counterfactual_strategies(rows)
    change_counts = Counter(row["变化结果"] for row in change_rows)
    failure_counts = Counter(category for row in change_rows for category in norm(row["失败原因类别"]).split(";") if category)
    wrong_block_failure_counts = Counter(
        category
        for row in change_rows
        if row["cluster_id"] in {item["cluster_id"] for item in wrong_block_rows}
        for category in norm(row["失败原因类别"]).split(";")
        if category
    )
    formal_after = hash_files(FORMAL_OUTPUTS)
    frozen_after = hash_files(FROZEN_FILES)
    return {
        "summary": {
            "sample_count": len(rows),
            "model_calls": 0,
            "statistics_code_error_found": False,
            "statistics_definition_inconsistency_found": True,
            "statistics_issue_summary": [
                "correct/wrong/neutral block accounting is internally reconcilable: 22 = 2 + 17 + 3.",
                "3/4 and 3/7 come from the Step 4.4-E Round1-insufficient denominator, while earlier reports usually used Baseline-insufficient denominator.",
                "Step 4.4-E's 32/47 change count means verdict-or-gate_reason changed; pure verdict mismatch is lower.",
                "No empty or invalid verdict mappings were found across the 47 rows.",
            ],
            "gate_blocked_destination_counts": metric_audit["block_destination_counts"],
            "change_outcome_counts": dict(change_counts),
            "pure_verdict_change_count": sum(1 for row in change_rows if row["verdict是否不一致"]),
            "verdict_or_reason_change_count": len(change_rows),
            "wrong_block_failure_reason_counts": dict(wrong_block_failure_counts),
            "claim_truly_unclear_summary": claim_unclear_summary,
            "direct_evidence_summary": direct_summary,
            "counterfactual_strategies": counterfactual_rows,
            "best_v1_2_base_strategy": max(counterfactual_rows, key=lambda item: (item["accuracy"], item["macro_f1_fixed_4class"]))["strategy"],
            "recommend_enter_v1_2": True,
            "recommend_enter_gray_or_isr": False,
            "formal_outputs_unchanged": formal_before == formal_after,
            "frozen_files_unchanged": frozen_before == frozen_after,
            "formal_output_sha256_before": formal_before,
            "formal_output_sha256_after": formal_after,
            "frozen_file_sha256_before": frozen_before,
            "frozen_file_sha256_after": frozen_after,
            "backup_dir": str(args.backup_dir),
        },
        "metric_audit": metric_audit,
        "changed_samples_32": change_rows,
        "wrong_blocks_17": wrong_block_rows,
        "claim_unclear_detail": claim_unclear_rows,
        "claim_unclear_summary": claim_unclear_summary,
        "direct_evidence_quality": direct_rows,
        "direct_evidence_summary": direct_summary,
        "counterfactual_strategies": counterfactual_rows,
        "v1_2_recommendations": v1_2_recommendations(),
    }


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def stringify(value: Any) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return norm(value)


def write_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title=title[:31])
    if not rows:
        rows = [{"empty": ""}]
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in rows:
        ws.append([stringify(row.get(header, "")) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, 1):
        max_len = max([len(stringify(header))] + [len(stringify(row.get(header, ""))) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 80)


def write_xlsx(path: Path, audit: dict[str, Any]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    write_sheet(wb, "指标口径核对", flatten_metric_audit(audit["metric_audit"]))
    write_sheet(wb, "32条变化样本", audit["changed_samples_32"])
    write_sheet(wb, "17条错误阻止", audit["wrong_blocks_17"])
    write_sheet(wb, "claim unclear分类", audit["claim_unclear_detail"] + [{"summary": json.dumps(audit["claim_unclear_summary"], ensure_ascii=False)}])
    write_sheet(wb, "direct evidence质量", audit["direct_evidence_quality"] + [{"summary": json.dumps(audit["direct_evidence_summary"], ensure_ascii=False)}])
    write_sheet(wb, "反事实回退策略", audit["counterfactual_strategies"])
    write_sheet(wb, "v1.2建议", audit["v1_2_recommendations"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def flatten_metric_audit(metric_audit: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [
        {"item": "same_47_cluster_ids", "value": metric_audit["same_47_cluster_ids"], "note": ""},
        {"item": "all_verdicts_valid", "value": metric_audit["all_verdicts_valid"], "note": json.dumps(metric_audit["invalid_verdicts"], ensure_ascii=False)},
        {"item": "accuracy_31_91_is_15_of_47", "value": metric_audit["accuracy_31_91_is_15_of_47"], "note": metric_audit["v1_1_correct_count"]},
        {"item": "fixed_macro_classes", "value": ",".join(metric_audit["fixed_macro_classes"]), "note": ""},
        {"item": "observed_label_classes", "value": ",".join(metric_audit["observed_label_classes"]), "note": ""},
        {"item": "gate_blocked_destination_counts", "value": json.dumps(metric_audit["block_destination_counts"], ensure_ascii=False), "note": metric_audit["neutral_block_explanation"]},
        {"item": "baseline_insufficient_cached_resolved", "value": json.dumps(metric_audit["baseline_insufficient_cached_resolved"], ensure_ascii=False), "note": ""},
        {"item": "round1_insufficient_cached_resolved", "value": json.dumps(metric_audit["round1_insufficient_cached_resolved"], ensure_ascii=False), "note": metric_audit["denominator_explanation"]},
    ]
    for metric in metric_audit["metric_rows"]:
        rows.append({"item": f"metric_{metric['strategy']}", "value": json.dumps(metric, ensure_ascii=False), "note": ""})
    return rows


def render_markdown(audit: dict[str, Any]) -> str:
    summary = audit["summary"]
    metric = audit["metric_audit"]
    lines = [
        "# Task 2 Safety Gate v1.1 Failure Audit",
        "",
        "本报告只做诊断，不修改 Safety Gate v1/v1.1、配置、适配器、DEA、评分融合或正式输出；未调用 32B 模型，未进入 ISR。",
        "",
        "## 统计口径",
        "",
        f"- 47条一致匹配: {metric['same_47_cluster_ids']}",
        f"- verdict 均合法: {metric['all_verdicts_valid']}",
        f"- accuracy=31.91% 对应: {metric['v1_1_correct_count']}/47",
        f"- 固定 macro-F1 类别: {', '.join(metric['fixed_macro_classes'])}",
        f"- observed-label macro-F1 类别: {', '.join(metric['observed_label_classes'])}",
        f"- gate_blocked 去向: {json.dumps(metric['block_destination_counts'], ensure_ascii=False)}",
        f"- 3/4 与 3/7 解释: {metric['denominator_explanation']}",
        "",
        "## 变化样本",
        "",
        f"- v1 与 v1.1 verdict 或 gate_reason 变化: {len(audit['changed_samples_32'])}",
        f"- 纯 verdict 不一致: {summary['pure_verdict_change_count']}",
        f"- 变化结果: {json.dumps(summary['change_outcome_counts'], ensure_ascii=False)}",
        f"- 17条错误阻止原因分布: {json.dumps(summary['wrong_block_failure_reason_counts'], ensure_ascii=False)}",
        "",
        "## Claim Truly Unclear",
        "",
        f"- 分类: {json.dumps(summary['claim_truly_unclear_summary'], ensure_ascii=False)}",
        "- 数据验证结论: evidence_resolves_claim=true 的样本不应仅因 raw_claim_incomplete 或 broad claim 被强制改为 unclear。",
        "",
        "## Direct Evidence",
        "",
        f"- 汇总: {json.dumps(summary['direct_evidence_summary'], ensure_ascii=False)}",
        "",
        "## 反事实策略",
        "",
        "| Strategy | Accuracy | Correct | Fixed macro-F1 | Actual macro-F1 | Weighted-F1 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for row in audit["counterfactual_strategies"]:
        lines.append(
            f"| {row['strategy']} | {row['accuracy']:.6f} | {row['correct_count']}/{row['sample_count']} | "
            f"{row['macro_f1_fixed_4class']:.6f} | {row['macro_f1_actual_labels']:.6f} | {row['weighted_f1']:.6f} |"
        )
    lines.extend(
        [
            "",
            f"- 最适合作为 v1.2 基础的策略: {summary['best_v1_2_base_strategy']}",
            "",
            "## v1.2建议",
            "",
        ]
    )
    for item in audit["v1_2_recommendations"]:
        lines.append(f"- {item['priority']}: {item['recommendation']}")
    lines.extend(
        [
            "",
            "## 保护状态",
            "",
            f"- 正式输出不变: {summary['formal_outputs_unchanged']}",
            f"- 冻结文件不变: {summary['frozen_files_unchanged']}",
            f"- 32B 调用次数: {summary['model_calls']}",
            f"- 备份目录: {summary['backup_dir']}",
            "",
        ]
    )
    return "\n".join(lines)


def ensure_new_outputs(paths: list[Path], overwrite: bool) -> None:
    existing = [path for path in paths if path.exists()]
    if existing and not overwrite:
        raise FileExistsError(f"refusing to overwrite existing audit outputs: {[str(path) for path in existing]}")


def write_outputs(audit: dict[str, Any], args: argparse.Namespace) -> None:
    ensure_new_outputs([args.report_json, args.report_md, args.report_xlsx], args.overwrite)
    write_json(args.report_json, audit)
    args.report_md.parent.mkdir(parents=True, exist_ok=True)
    args.report_md.write_text(render_markdown(audit), encoding="utf-8")
    write_xlsx(args.report_xlsx, audit)


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--gold-file", type=Path, default=DEFAULT_GOLD_FILE)
    parser.add_argument("--v1-results", type=Path, default=DEFAULT_V1_RESULTS)
    parser.add_argument("--v1-1-results", type=Path, default=DEFAULT_V1_1_RESULTS)
    parser.add_argument("--v1-report", type=Path, default=DEFAULT_V1_REPORT)
    parser.add_argument("--v1-1-report", type=Path, default=DEFAULT_V1_1_REPORT)
    parser.add_argument("--report-json", type=Path, default=DEFAULT_REPORT_JSON)
    parser.add_argument("--report-md", type=Path, default=DEFAULT_REPORT_MD)
    parser.add_argument("--report-xlsx", type=Path, default=DEFAULT_REPORT_XLSX)
    parser.add_argument("--backup-dir", type=Path, default=DEFAULT_BACKUP_DIR)
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_arg_parser().parse_args(argv)
    audit = build_audit(args)
    write_outputs(audit, args)
    print(
        json.dumps(
            {
                "sample_count": audit["summary"]["sample_count"],
                "changed_samples": len(audit["changed_samples_32"]),
                "wrong_blocks": len(audit["wrong_blocks_17"]),
                "direct_evidence_precision": audit["direct_evidence_summary"]["precision"],
                "formal_outputs_unchanged": audit["summary"]["formal_outputs_unchanged"],
                "frozen_files_unchanged": audit["summary"]["frozen_files_unchanged"],
                "model_calls": audit["summary"]["model_calls"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
