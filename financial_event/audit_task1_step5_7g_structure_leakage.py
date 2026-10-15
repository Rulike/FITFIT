#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Task1 Step 5.7-G: leakage audit and gold-free retest."""

from __future__ import annotations

import hashlib
import inspect
import json
import random
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import scripts.run_task1_step5_7f_structure_training as step57f  # noqa: E402


OUT_DIR = ROOT / "outputs/task1_step5_7g_gold_free_audit"
REPORT_JSON = ROOT / "reports/task1_step5_7g_structure_leakage_report.json"
REPORT_MD = ROOT / "reports/task1_step5_7g_structure_leakage_report.md"
DETAILS_XLSX = ROOT / "reports/task1_step5_7g_structure_leakage_details.xlsx"

STEP57F_OUT = ROOT / "outputs/task1_step5_7f_structure_training_v1"
DEV_ROWS = STEP57F_OUT / "dev_eventmentions.jsonl"
STEP57F_PRED = STEP57F_OUT / "predictions/rule_ai_adapter_baseline_dev_predictions.jsonl"
STEP57F_REPORT = ROOT / "reports/task1_step5_7f_structure_training_report.json"
CONFIRMED_XLSX = ROOT / "reports/task1_step5_7e_eventmention_human_confirmed.xlsx"

GOLD_FREE_INPUT = OUT_DIR / "gold_free_dev_input.jsonl"
GOLD_FREE_INPUT_MANIFEST = OUT_DIR / "gold_free_dev_input_manifest.json"
GOLD_FREE_PRED = OUT_DIR / "gold_free_dev_predictions.jsonl"
GOLD_FREE_INFERENCE_MANIFEST = OUT_DIR / "gold_free_inference_manifest.json"
EVAL_GOLD = OUT_DIR / "evaluation_gold_dev_labels.jsonl"
GOLD_FREE_METRICS = OUT_DIR / "gold_free_dev_metrics.json"

ALLOWED_FIELDS = ["event_mention_id", "article_id", "title", "content", "publish_time", "event_name"]
BANNED_PREFIXES = ("ai_", "human_", "reviewer_")
BANNED_EXACT = {
    "ai_fields",
    "reference_raw_text",
    "original_quality_grade",
    "selected_quality_grade",
    "needs_manual_correction",
    "validation_errors",
    "annotation_source",
    "annotation_status",
    "confirmation_note",
    "confirmed_at",
    "source_workbook",
    "trigger",
    "event_span",
    "subject_company",
    "event_time",
    "object_target",
    "direction",
    "arguments",
    "legacy_raw_structure",
    "multi_event_article",
    "source_set",
}
NEGATIVE_TESTS = [
    "shuffle_human_fields",
    "clear_ai_fields",
    "delete_reference_raw_text",
    "replace_reviewer_note_and_quality_grade",
    "repeat_same_input_deterministic",
    "gold_file_unavailable",
    "prediction_offsets_from_content",
]
PROTECTED_FILES = [
    ROOT / "outputs/final_task1/final_events_train.jsonl",
    ROOT / "outputs/final_task1/final_events_test.jsonl",
    ROOT / "outputs/final_task1/final_events_train.xlsx",
    ROOT / "outputs/final_task1/final_events_test.xlsx",
    ROOT / "outputs/final_task1/task1_completion_certificate.json",
    ROOT / "outputs/event_credibility/llm_review_results.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.jsonl",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.parquet",
    ROOT / "outputs/event_credibility/event_credibility_scores_fulltext.xlsx",
    ROOT / "reports/task1_step5_7f_structure_training_report.json",
    ROOT / "reports/task1_step5_7f_structure_training_report.md",
    ROOT / "reports/task1_step5_7f_structure_training_details.xlsx",
    CONFIRMED_XLSX,
    DEV_ROWS,
    STEP57F_OUT / "train_eventmentions.jsonl",
    STEP57F_OUT / "confirmed_eventmentions_200.jsonl",
    STEP57F_OUT / "legacy80_diagnostic_eventmentions.jsonl",
    STEP57F_OUT / "split_manifest.json",
    STEP57F_OUT / "run_manifest.json",
    STEP57F_PRED,
    STEP57F_OUT / "models/macbert_multitask_token/model.pt",
    STEP57F_OUT / "models/macbert_weighted_token/model.pt",
    ROOT / "scripts/run_task1_step5_7f_structure_training.py",
    ROOT / "scripts/evaluate_task1_step5_7f_structure_training.py",
    ROOT / "scripts/test_task1_step5_7f_structure_training.py",
]


def now() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def rel(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        return str(path)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False) + "\n")


def protected_paths() -> list[Path]:
    return sorted({path for path in PROTECTED_FILES if path.exists()}, key=rel)


def sha_snapshot(paths: list[Path]) -> dict[str, str]:
    return {rel(path): sha256_file(path) for path in paths}


def write_sha_snapshot(path: Path, snapshot: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(f"{digest}  {name}" for name, digest in sorted(snapshot.items())) + "\n", encoding="utf-8")


def create_backup_dir() -> Path:
    backup = Path("/opt/data/private/JinRong/kechuang/backups") / f"task1_step5_7g_structure_leakage_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    backup.mkdir(parents=True, exist_ok=False)
    return backup


def process_check() -> list[str]:
    proc = subprocess.run(
        ["bash", "-lc", "ps -eo pid,ppid,stat,etime,cmd | rg -i 'step5_7g|audit_task1_step5_7g|run_task1_step5_7g|evaluate_task1_step5_7g|vllm|32b' || true"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )
    own = {"ps -eo", "rg -i", "audit_task1_step5_7g_structure_leakage.py"}
    return [line for line in proc.stdout.splitlines() if not any(token in line for token in own)]


def banned_fields(row: dict[str, Any]) -> list[str]:
    hits = []
    for key in row:
        if key in BANNED_EXACT or key.startswith(BANNED_PREFIXES):
            hits.append(key)
    return sorted(hits)


def project_gold_free(row: dict[str, Any]) -> dict[str, Any]:
    return {field: str(row.get(field, "") or "") for field in ALLOWED_FIELDS}


def build_gold_free_input(dev_rows: list[dict[str, Any]]) -> dict[str, Any]:
    if len(dev_rows) != 40:
        raise RuntimeError(f"Step 5.7-F dev rows must be 40, got {len(dev_rows)}")
    rows = [project_gold_free(row) for row in dev_rows]
    write_jsonl(GOLD_FREE_INPUT, rows)
    field_counts = Counter()
    banned_counter = Counter()
    for row in rows:
        field_counts.update(row.keys())
        banned_counter.update(banned_fields(row))
    bad_counts = dict(sorted(banned_counter.items()))
    if bad_counts:
        raise RuntimeError(f"gold-free input contains banned fields: {bad_counts}")
    manifest = {
        "generated_at": now(),
        "row_count": len(rows),
        "source_dev_rows": rel(DEV_ROWS),
        "dev_row_order_preserved": True,
        "article_ids": [row["article_id"] for row in rows],
        "allowed_fields": ALLOWED_FIELDS,
        "field_whitelist": ALLOWED_FIELDS,
        "field_counts": dict(sorted(field_counts.items())),
        "banned_field_hit_count": 0,
        "banned_field_counts": {},
        "gold_free_input": rel(GOLD_FREE_INPUT),
        "gold_free_input_sha256": sha256_file(GOLD_FREE_INPUT),
    }
    write_json(GOLD_FREE_INPUT_MANIFEST, manifest)
    return manifest


def gold_rows_from_dev(dev_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fields = [
        "event_mention_id",
        "article_id",
        "event_name",
        "title",
        "content",
        "publish_time",
        "trigger",
        "event_span",
        "subject_company",
        "event_time",
        "object_target",
        "direction",
        "arguments",
    ]
    return [{field: row.get(field) for field in fields} for row in dev_rows]


def run_cmd(cmd: list[str]) -> dict[str, Any]:
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, timeout=180)
    return {
        "cmd": " ".join(cmd),
        "returncode": proc.returncode,
        "stdout_tail": proc.stdout[-4000:],
        "stderr_tail": proc.stderr[-4000:],
    }


def run_inference(input_path: Path, output_path: Path, manifest_path: Path) -> dict[str, Any]:
    result = run_cmd([
        sys.executable,
        "scripts/run_task1_step5_7g_gold_free_structure.py",
        "--input",
        rel(input_path),
        "--output",
        rel(output_path),
        "--manifest",
        rel(manifest_path),
    ])
    if result["returncode"] != 0:
        raise RuntimeError(f"gold-free inference failed: {result}")
    manifest = read_json(manifest_path)
    return {"process": result, "manifest": manifest}


def run_evaluation() -> dict[str, Any]:
    result = run_cmd([
        sys.executable,
        "scripts/evaluate_task1_step5_7g_gold_free_structure.py",
        "--gold",
        rel(EVAL_GOLD),
        "--predictions",
        rel(GOLD_FREE_PRED),
        "--output",
        rel(GOLD_FREE_METRICS),
    ])
    if result["returncode"] != 0:
        raise RuntimeError(f"gold-free evaluation failed: {result}")
    return {"process": result, "metrics": read_json(GOLD_FREE_METRICS)}


def prediction_core_sha(path: Path) -> str:
    return sha256_file(path)


def write_variant(path: Path, rows: list[dict[str, Any]]) -> None:
    write_jsonl(path, rows)


def human_variant(base_rows: list[dict[str, Any]], dev_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    shuffled = list(dev_rows)
    random.Random(2026).shuffle(shuffled)
    out = []
    for base, donor in zip(base_rows, shuffled):
        row = dict(base)
        row.update(
            {
                "human_trigger": (donor.get("trigger") or {}).get("text", ""),
                "human_event_span": (donor.get("event_span") or {}).get("text", ""),
                "human_subject_company": (donor.get("subject_company") or {}).get("text", ""),
                "human_event_time": (donor.get("event_time") or {}).get("text", ""),
                "human_event_time_normalized": (donor.get("event_time") or {}).get("normalized", "未明确"),
                "human_object_target": (donor.get("object_target") or {}).get("text", ""),
                "human_direction": donor.get("direction", ""),
                "human_arguments_json": json.dumps(donor.get("arguments") or [], ensure_ascii=False),
                "human_argument_roles_json": json.dumps({arg["role"]: arg["text"] for arg in donor.get("arguments") or []}, ensure_ascii=False),
            }
        )
        out.append(row)
    return out


def ai_cleared_variant(base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for base in base_rows:
        row = dict(base)
        row.update(
            {
                "ai_trigger": "",
                "ai_event_span": "",
                "ai_subject_company": "",
                "ai_event_time": "",
                "ai_event_time_normalized": "",
                "ai_object_target": "",
                "ai_direction": "",
                "ai_arguments_json": "",
                "ai_argument_roles_json": "",
                "ai_fields": {},
            }
        )
        out.append(row)
    return out


def reference_variant(base_rows: list[dict[str, Any]], include_reference: bool) -> list[dict[str, Any]]:
    out = []
    for idx, base in enumerate(base_rows):
        row = dict(base)
        if include_reference:
            row["reference_raw_text"] = f"forbidden_reference_payload_{idx}"
        else:
            row.pop("reference_raw_text", None)
        out.append(row)
    return out


def quality_variant(base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    grades = ["A", "B", "C", "D"]
    for idx, base in enumerate(base_rows):
        row = dict(base)
        row.update(
            {
                "reviewer_note": f"negative leakage note {idx}",
                "reviewer_confirmed": "no" if idx % 2 else "yes",
                "original_quality_grade": grades[idx % len(grades)],
                "selected_quality_grade": grades[(idx + 2) % len(grades)],
                "needs_manual_correction": "yes" if idx % 2 else "no",
                "validation_errors": "injected_negative_test",
            }
        )
        out.append(row)
    return out


def validate_offsets_from_content(input_rows: list[dict[str, Any]], pred_rows: list[dict[str, Any]]) -> dict[str, Any]:
    by_id = {row["event_mention_id"]: row for row in input_rows}
    errors: list[str] = []
    for pred in pred_rows:
        row = by_id.get(pred.get("event_mention_id"))
        if not row:
            errors.append(f"unknown_event_mention_id:{pred.get('event_mention_id')}")
            continue
        content = row["content"]
        for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]:
            span = pred.get(field) or {}
            text = str(span.get("text", "") or "")
            start = int(span.get("start", -1))
            end = int(span.get("end", -1))
            if not text:
                if start != -1 or end != -1:
                    errors.append(f"{pred.get('event_mention_id')}:{field}:empty_with_nonempty_offset")
            elif not (0 <= start <= end <= len(content)) or content[start:end] != text:
                errors.append(f"{pred.get('event_mention_id')}:{field}:offset_mismatch")
    return {"passed": not errors, "error_count": len(errors), "errors": errors[:50]}


def run_negative_tests(base_rows: list[dict[str, Any]], dev_rows: list[dict[str, Any]], base_sha: str) -> dict[str, Any]:
    tests_dir = OUT_DIR / "negative_tests"
    tests_dir.mkdir(parents=True, exist_ok=True)
    results: dict[str, Any] = {}

    def run_variant(name: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
        input_path = tests_dir / f"{name}_input.jsonl"
        output_path = tests_dir / f"{name}_predictions.jsonl"
        manifest_path = tests_dir / f"{name}_manifest.json"
        write_variant(input_path, rows)
        run_inference(input_path, output_path, manifest_path)
        sha = prediction_core_sha(output_path)
        return {"passed": sha == base_sha, "prediction_sha256": sha, "matches_base": sha == base_sha, "manifest": read_json(manifest_path)}

    results["shuffle_human_fields"] = run_variant("shuffle_human_fields", human_variant(base_rows, dev_rows))
    results["clear_ai_fields"] = run_variant("clear_ai_fields", ai_cleared_variant(base_rows))
    with_ref = run_variant("reference_raw_text_present", reference_variant(base_rows, True))
    without_ref = run_variant("reference_raw_text_deleted", reference_variant(base_rows, False))
    results["delete_reference_raw_text"] = {
        "passed": with_ref["prediction_sha256"] == without_ref["prediction_sha256"] == base_sha,
        "with_reference_sha256": with_ref["prediction_sha256"],
        "without_reference_sha256": without_ref["prediction_sha256"],
        "matches_base": with_ref["prediction_sha256"] == without_ref["prediction_sha256"] == base_sha,
    }
    results["replace_reviewer_note_and_quality_grade"] = run_variant("replace_reviewer_note_and_quality_grade", quality_variant(base_rows))
    results["repeat_same_input_deterministic"] = run_variant("repeat_same_input_deterministic", base_rows)

    hidden_dir = tests_dir / "gold_unavailable"
    hidden_dir.mkdir(parents=True, exist_ok=True)
    hidden_gold = hidden_dir / "evaluation_gold_dev_labels.jsonl.hidden"
    moved = False
    if EVAL_GOLD.exists():
        shutil.move(str(EVAL_GOLD), str(hidden_gold))
        moved = True
    try:
        gold_unavailable = run_variant("gold_file_unavailable", base_rows)
        results["gold_file_unavailable"] = {**gold_unavailable, "gold_file_moved": moved}
    finally:
        if moved:
            shutil.move(str(hidden_gold), str(EVAL_GOLD))

    pred_rows = read_jsonl(GOLD_FREE_PRED)
    offset_result = validate_offsets_from_content(base_rows, pred_rows)
    results["prediction_offsets_from_content"] = offset_result
    results["all_passed"] = all(bool(results[name].get("passed")) for name in NEGATIVE_TESTS)
    return results


def compare_predictions_to_adapter(dev_rows: list[dict[str, Any]], pred_rows: list[dict[str, Any]]) -> dict[str, Any]:
    reproduced = step57f.predict_adapter(dev_rows)
    pred_by_id = {row["event_mention_id"]: row for row in pred_rows}
    mismatches = []
    for expected in reproduced:
        actual = pred_by_id.get(expected["event_mention_id"])
        comparable_actual = {key: actual.get(key) for key in expected if actual} if actual else None
        if comparable_actual != expected:
            mismatches.append(expected["event_mention_id"])
    return {
        "dev_prediction_count": len(pred_rows),
        "reproduced_from_ai_fields_count": len(reproduced) - len(mismatches),
        "all_predictions_reproduced_from_ai_fields": not mismatches and len(reproduced) == len(pred_rows),
        "mismatch_event_mention_ids": mismatches[:50],
    }


def legacy80_same_source_check() -> dict[str, Any]:
    legacy_rows = read_jsonl(STEP57F_OUT / "legacy80_diagnostic_eventmentions.jsonl")
    legacy_preds = read_jsonl(STEP57F_OUT / "predictions/legacy80_diagnostic_predictions.jsonl")
    copied = compare_predictions_to_adapter(legacy_rows, legacy_preds)
    ai_equals_gold = 0
    for row in legacy_rows:
        pred = step57f.span_from_ai(row)
        if all(pred[field] == row[field] for field in ["trigger", "event_span", "subject_company", "event_time", "object_target"]) and pred["direction"] == row["direction"] and pred["arguments"] == row["arguments"]:
            ai_equals_gold += 1
    return {
        **copied,
        "legacy_rows": len(legacy_rows),
        "ai_fields_reconstruct_gold_count": ai_equals_gold,
        "explanation": "旧80诊断rows的ai_fields由迁移后的结构标签/本轮确认结构生成，adapter再从ai_fields还原，评估端也使用同一批结构字段，因此1.0是同源一致性，不代表真实泛化能力。",
    }


def static_dependency_audit() -> dict[str, Any]:
    source_predict = inspect.getsource(step57f.predict_adapter)
    source_span_ai = inspect.getsource(step57f.span_from_ai)
    source_canonical = inspect.getsource(step57f.canonical_from_row)
    return {
        "implementation_files": [
            rel(ROOT / "scripts/run_task1_step5_7f_structure_training.py"),
            rel(ROOT / "scripts/evaluate_task1_step5_7f_structure_training.py"),
            rel(ROOT / "outputs/task1_step5_7f_structure_training_v1/predictions/rule_ai_adapter_baseline_dev_predictions.jsonl"),
            rel(ROOT / "outputs/task1_step5_7f_structure_training_v1/run_manifest.json"),
        ],
        "rule_ai_adapter_actual_read_fields": [
            "event_mention_id",
            "content",
            "ai_fields.trigger",
            "ai_fields.event_span",
            "ai_fields.subject_company",
            "ai_fields.event_time",
            "ai_fields.event_time_normalized",
            "ai_fields.object_target",
            "ai_fields.direction",
            "ai_fields.arguments_json",
            "ai_fields.argument_roles_json",
        ],
        "allowed_inference_fields": ALLOWED_FIELDS,
        "prohibited_fields_read": [
            "ai_fields.trigger",
            "ai_fields.event_span",
            "ai_fields.subject_company",
            "ai_fields.event_time",
            "ai_fields.event_time_normalized",
            "ai_fields.object_target",
            "ai_fields.direction",
            "ai_fields.arguments_json",
            "ai_fields.argument_roles_json",
        ],
        "source_columns_of_ai_fields": [
            "ai_trigger",
            "ai_event_span",
            "ai_subject_company",
            "ai_event_time",
            "ai_event_time_normalized",
            "ai_object_target",
            "ai_direction",
            "ai_arguments_json",
            "ai_argument_roles_json",
        ],
        "direct_human_fields_read_by_adapter": [],
        "reference_raw_text_read_by_adapter": False,
        "confirmed_workbook_loaded_before_step57f_inference": True,
        "dev_gold_labels_present_in_step57f_inference_rows": True,
        "rule_dictionary_or_threshold_training": "none",
        "dev_labels_used_for_rule_construction": False,
        "dev_labels_used_at_inference_via_ai_fields": True,
        "label_dependent_adapter": True,
        "source_snippets_sha256": {
            "predict_adapter": hashlib.sha256(source_predict.encode("utf-8")).hexdigest(),
            "span_from_ai": hashlib.sha256(source_span_ai.encode("utf-8")).hexdigest(),
            "canonical_from_row": hashlib.sha256(source_canonical.encode("utf-8")).hexdigest(),
        },
    }


def make_details_xlsx(report: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    for key, value in [
        ("status", report["status"]),
        ("label_dependent_adapter", report["audit"]["static_dependency"]["label_dependent_adapter"]),
        ("gold_free_dev_count", report["gold_free_input"]["row_count"]),
        ("gold_free_status", report["gold_free_evaluation"]["metrics"]["status"]),
        ("negative_tests_all_passed", report["negative_tests"]["all_passed"]),
    ]:
        ws.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws2 = wb.create_sheet("actual_read_fields")
    ws2.append(["field", "allowed"])
    allowed = set(ALLOWED_FIELDS)
    for field in report["audit"]["static_dependency"]["rule_ai_adapter_actual_read_fields"]:
        ws2.append([field, field in allowed])
    ws3 = wb.create_sheet("gold_free_metrics")
    metrics = report["gold_free_evaluation"]["metrics"]["gold_free_metrics"]
    for key, value in metrics.items():
        ws3.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    ws4 = wb.create_sheet("negative_tests")
    ws4.append(["test", "passed", "details"])
    for name, result in report["negative_tests"].items():
        if name == "all_passed":
            continue
        ws4.append([name, result.get("passed"), json.dumps(result, ensure_ascii=False)])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 100
    DETAILS_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DETAILS_XLSX)


def make_report_md(report: dict[str, Any]) -> None:
    m = report["gold_free_evaluation"]["metrics"]["gold_free_metrics"]
    lines = [
        "# Task1 Step 5.7-G结构泄漏审计报告",
        "",
        f"状态：`{report['status']}`",
        "",
        "## 结论",
        "",
        "- Step 5.7-F的rule_ai_adapter_baseline读取`ai_fields.*`，这些字段由`ai_*`列或历史确认结构派生，违反gold-free推理字段约束。",
        "- 旧80条1.0来自同源结构字段被写入`ai_fields`后再被adapter还原，不代表泛化能力。",
        "",
        "## Gold-free指标",
        "",
        f"- trigger relaxed F1：{m.get('trigger_relaxed_f1')}",
        f"- event_span character F1：{m.get('event_span_character_overlap_f1')}",
        f"- subject normalized accuracy：{m.get('subject_normalized_accuracy')}",
        f"- event_time normalized accuracy：{m.get('event_time_normalized_date_accuracy')}",
        f"- object normalized accuracy：{m.get('object_normalized_accuracy')}",
        f"- direction accuracy：{m.get('direction_accuracy')}",
        f"- argument+role F1：{m.get('argument_role_end_to_end_f1')}",
        f"- article exact accuracy：{m.get('article_exact_match_accuracy')}",
        f"- empty output count：{m.get('empty_output_count')}",
        "",
        "## 保护",
        "",
        "- 32B调用：0",
        "- 外部LLM调用：0",
        "- 未读取1029篇gold test标签",
        "- 未修改正式Task1/Task2输出或Step 5.7-F冻结产物",
    ]
    REPORT_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")


def compare_sha(before: dict[str, str], backup_dir: Path) -> dict[str, Any]:
    after = sha_snapshot(protected_paths())
    write_sha_snapshot(backup_dir / "protected_sha256_after.txt", after)
    mismatch = {key: {"before": before.get(key), "after": after.get(key)} for key in sorted(set(before) | set(after)) if before.get(key) != after.get(key)}
    result = {"status": "matched" if not mismatch else "mismatch", "checked_count": len(before), "mismatch_count": len(mismatch), "mismatch": mismatch}
    write_json(backup_dir / "protected_sha256_compare.json", result)
    return result


def main() -> int:
    for path in [DEV_ROWS, STEP57F_PRED, STEP57F_REPORT, CONFIRMED_XLSX]:
        if not path.exists():
            raise FileNotFoundError(path)
    existing = {
        "outputs": sorted(rel(path) for path in OUT_DIR.rglob("*")) if OUT_DIR.exists() else [],
        "reports": sorted(rel(path) for path in [REPORT_JSON, REPORT_MD, DETAILS_XLSX] if path.exists()),
    }
    process_lines = process_check()
    if process_lines:
        raise RuntimeError(f"another Step 5.7-G/vLLM/32B-like process is running: {process_lines}")
    backup_dir = create_backup_dir()
    before = sha_snapshot(protected_paths())
    write_sha_snapshot(backup_dir / "protected_sha256_before.txt", before)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dev_rows = read_jsonl(DEV_ROWS)
    step57f_preds = read_jsonl(STEP57F_PRED)
    step57f_report = read_json(STEP57F_REPORT)
    static_audit = static_dependency_audit()
    direct_copy = compare_predictions_to_adapter(dev_rows, step57f_preds)
    legacy_check = legacy80_same_source_check()
    gold_free_input_manifest = build_gold_free_input(dev_rows)
    base_inference = run_inference(GOLD_FREE_INPUT, GOLD_FREE_PRED, GOLD_FREE_INFERENCE_MANIFEST)
    base_pred_sha = sha256_file(GOLD_FREE_PRED)

    # Gold labels become a file only after the independent inference process has
    # completed and the prediction SHA has been frozen.
    write_jsonl(EVAL_GOLD, gold_rows_from_dev(dev_rows))
    eval_result = run_evaluation()
    base_rows = read_jsonl(GOLD_FREE_INPUT)
    negative_tests = run_negative_tests(base_rows, dev_rows, base_pred_sha)

    leakage_found = bool(static_audit["label_dependent_adapter"] or static_audit["prohibited_fields_read"])
    gold_free_pass = eval_result["metrics"]["status"] == "passed_thresholds"
    if leakage_found:
        status = "structure_baseline_invalid_due_to_label_leakage"
    elif not negative_tests["all_passed"] or not gold_free_pass:
        status = "gold_free_structure_baseline_failed"
    else:
        status = "gold_free_structure_baseline_passed"

    sha_compare = compare_sha(before, backup_dir)
    if sha_compare["status"] != "matched":
        raise RuntimeError(f"protected SHA changed: {sha_compare}")

    report = {
        "generated_at": now(),
        "status": status,
        "existing_step57g_artifacts_before_run": existing,
        "recovery": {
            "process_check": process_lines,
            "backup_dir": str(backup_dir),
            "step57f_report": rel(STEP57F_REPORT),
            "step57f_best_scheme": (step57f_report.get("selection") or {}).get("best_scheme"),
            "step57f_best_artifact": (step57f_report.get("selection") or {}).get("best_artifact"),
            "step57f_best_artifact_sha256": (step57f_report.get("selection") or {}).get("best_artifact_sha256"),
            "confirmed_workbook": rel(CONFIRMED_XLSX),
            "confirmed_workbook_sha256": sha256_file(CONFIRMED_XLSX),
            "dev_training_data": rel(DEV_ROWS),
            "dev_training_data_sha256": sha256_file(DEV_ROWS),
        },
        "audit": {
            "static_dependency": static_audit,
            "prediction_direct_copy_check": direct_copy,
            "legacy80_same_source_check": legacy_check,
        },
        "gold_free_input": gold_free_input_manifest,
        "gold_free_inference": {
            "base_prediction_sha256": base_pred_sha,
            "base_prediction_path": rel(GOLD_FREE_PRED),
            "manifest": base_inference["manifest"],
            "gold_sidecar_written_after_inference": True,
            "gold_sidecar": rel(EVAL_GOLD),
            "gold_sidecar_sha256": sha256_file(EVAL_GOLD),
        },
        "gold_free_evaluation": eval_result,
        "negative_tests": negative_tests,
        "decision": {
            "input_banned_field_hit_count_zero": gold_free_input_manifest["banned_field_hit_count"] == 0,
            "negative_tests_all_passed": negative_tests["all_passed"],
            "prediction_independent_of_confirmed_workbook_or_gold": True,
            "gold_free_metrics_pass_thresholds": gold_free_pass,
            "invalid_offsets_json_cross_article_zero": (
                eval_result["metrics"]["gold_free_metrics"].get("offset_illegal_count") == 0
                and eval_result["metrics"]["gold_free_metrics"].get("illegal_direction_count") == 0
                and eval_result["metrics"]["gold_free_metrics"].get("json_parse_failure_count") == 0
            ),
            "label_leakage_found": leakage_found,
            "allow_1029_shadow_next": status == "gold_free_structure_baseline_passed",
        },
        "protection": {
            "formal_task1_task2_modified": False,
            "step57f_frozen_artifacts_modified": False,
            "protected_sha_status": sha_compare["status"],
            "gold_test_labels_read": False,
            "local_32b_calls": 0,
            "external_llm_calls": 0,
            "vllm_started": False,
        },
        "backup_dir": str(backup_dir),
    }
    write_json(REPORT_JSON, report)
    write_json(OUT_DIR / "run_manifest.json", report)
    make_report_md(report)
    make_details_xlsx(report)
    write_json(OUT_DIR / "SHA256SUMS.json", {rel(path): sha256_file(path) for path in [GOLD_FREE_INPUT, GOLD_FREE_INPUT_MANIFEST, GOLD_FREE_PRED, GOLD_FREE_INFERENCE_MANIFEST, EVAL_GOLD, GOLD_FREE_METRICS, REPORT_JSON, REPORT_MD, DETAILS_XLSX] if path.exists()})
    print(json.dumps({"status": status, "label_leakage_found": leakage_found, "gold_free_pass": gold_free_pass, "backup_dir": str(backup_dir)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
